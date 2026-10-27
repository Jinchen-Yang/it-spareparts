"""06 工作簿调用契约：真实 PostgreSQL 等锁后不能沿用已撤销权限。

使用已签名 HTTP 登录、管理员权限修改、独立项目范围事务及 pg_blocking_pids
证明等待窗口；只有锁内回执分支测试直接调用服务，不替换鉴权或锁。
"""

from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
import json
from threading import Event
from time import monotonic
from types import SimpleNamespace
from uuid import uuid4

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import func, select, text

from app import config
from app.api import maintenance_project_master_workbook as master_api
from app.auth import hash_password, verify_token
from app.db import Base, SessionLocal, engine
from app.main import app
from app.models.maintenance_manager import MaintenanceCollectionMilestone
from app.models.maintenance_project import (
    MaintenanceProjectContract,
    MaintenanceProjectUserAssignment,
)
from app.models.maintenance_project_operations import (
    MaintenanceProjectOperationAudit,
    MaintenanceProjectWorkbookOperation,
    MaintenanceProjectWorkbookState,
    MaintenanceSiteIssue,
    MaintenanceSiteIssueLine,
)
from app.models.system import SysUser
from app.services import maintenance_project_master_workbook as master
from app.services import maintenance_project_operations as operations
from tests.test_maintenance_expense_collection_workbook import _client, _PASSWORD, _XLSX
from tests.test_maintenance_project_master_v2_editable import _make_project_with_line, _site_issue


_ADMIN_PASSWORD = "synthetic-site-reauthorization-admin-123"
_OVERRIDES = {
    "page_maintenance": True,
    "page_maintenance_beta": True,
    "data_purchase_cost": True,
    "data_profit": True,
    "action_maintenance_expense_collection_upload": True,
    "action_maintenance_project_manage": True,
    # 工作簿编辑资格不要求独立网页更正动作权。
    "action_maintenance_site_issue_manage": False,
}
_FACT_TABLES = (
    "f_maintenance_line", "f_project_expense", "maintenance_project",
    "maintenance_project_contract", "maintenance_project_audit_log",
    "maintenance_source_order_assignment", "maintenance_project_workbook_state",
    "maintenance_project_workbook_operation", "maintenance_project_operation_audit",
    "maintenance_site_issue", "maintenance_site_issue_line",
    "maintenance_site_issue_display_counter", "maintenance_site_issue_command",
    "maintenance_site_issue_return_event", "maintenance_return_obligation",
    "maintenance_collection_milestone",
)


def _snapshot():
    """快照覆盖 16 张表的全部列，包括版本、审计内容和时间戳。"""
    with SessionLocal() as db:
        return {
            name: sorted(
                json.dumps(dict(row), sort_keys=True, ensure_ascii=False, default=str)
                for row in db.execute(select(Base.metadata.tables[name])).mappings()
            )
            for name in _FACT_TABLES
        }


def _login(client, username, password):
    response = client.post("/api/auth/login", json={"username": username, "password": password})
    assert response.status_code == 200, response.text
    client.headers["Authorization"] = f"Bearer {response.json()['token']}"


def _post(case):
    return case.client.post(case.route + "/apply", files={"file": ("site-rights.xlsx", case.data, _XLSX)})


@contextmanager
def _case(db, monkeypatch, *, editor="uploader", mixed=False):
    monkeypatch.setattr(master_api.get_settings(), "maintenance_project_master_v2_enabled", True)
    assert config.ENABLE_RBAC is True
    project, part, _order, _line = _make_project_with_line(db)
    issue, site_line = _site_issue(db, project, part)
    username = f"site-rights-{editor}"
    overrides = dict(_OVERRIDES)
    if editor in ("manager", "sales"):
        overrides["action_maintenance_expense_collection_upload"] = False
        overrides["action_maintenance_project_manage"] = False
    client = _client(db, username=username, overrides=overrides)
    if editor in ("manager", "sales"):
        user = db.scalar(select(SysUser).where(SysUser.username == username))
        if editor == "manager":
            db.add(MaintenanceProjectUserAssignment(
                assignment_id=str(uuid4()), project_id=project.project_id,
                responsibility_type="primary_manager", user_id=user.id, version=1,
                assigned_by="synthetic-site-reauthorization",
                assignment_reason="Synthetic workbook manager qualification",
            ))
        else:
            user.salesperson_name = "Synthetic canonical salesperson"
            project.salesperson = user.salesperson_name
        db.commit()
        _login(client, username, _PASSWORD)
    admin_user = SysUser(
        username="site-rights-admin", role="admin", display_name="Synthetic admin",
        password_hash=hash_password(_ADMIN_PASSWORD), is_active=True,
    )
    db.add(admin_user)
    operations.get_or_create_workbook_state(db, project_id=project.project_id)
    milestone_id = None
    if mixed:
        contract = db.scalar(select(MaintenanceProjectContract).where(
            MaintenanceProjectContract.project_id == project.project_id,
        ))
        milestone_id = str(uuid4())
        db.add(MaintenanceCollectionMilestone(
            milestone_id=milestone_id, project_id=project.project_id,
            project_contract_id=contract.project_contract_id, sequence=1,
            planned_date=date(2026, 10, 31), planned_amount=Decimal("800.00"),
            completeness_state="complete", source="project_master_v2", version=1,
        ))
    db.commit()
    admin = TestClient(app)
    _login(admin, "site-rights-admin", _ADMIN_PASSWORD)
    route = f"/api/maintenance/projects/stable/{project.project_id}/master-workbook"
    sheets = [master.V2_SHEET_SITE]
    if mixed:
        sheets.append(master.V2_SHEET_PLAN)
    exported = client.get(route + ".xlsx", params={"sheets": ",".join(sheets)})
    assert exported.status_code == 200, exported.text
    workbook = load_workbook(BytesIO(exported.content))
    sheet = workbook[master.V2_SHEET_SITE]
    columns = {cell.value: cell.column for cell in sheet[1] if cell.value}
    sheet.cell(2, columns["领用数量"], "3.25")
    if mixed:
        sheet = workbook[master.V2_SHEET_PLAN]
        # V2 签名模板固定列：实体 ID 第 13 列，计划金额第 6 列。
        matched = [row for row in sheet.iter_rows(min_row=2) if row[12].value == milestone_id]
        assert len(matched) == 1
        matched[0][5].value = "725.50"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    data = output.getvalue()
    preview = client.post(route + "/validate", files={"file": ("site-rights.xlsx", data, _XLSX)})
    assert preview.status_code == 200, preview.text
    result = SimpleNamespace(
        project_id=project.project_id, issue_id=issue.issue_id,
        site_line_id=site_line.issue_line_id, milestone_id=milestone_id,
        username=username, overrides=overrides, client=client, admin=admin,
        route=route, data=data,
    )
    db.rollback()
    try:
        yield result
    finally:
        client.close()
        admin.close()


def _revoke(case, revocation):
    if revocation == "scope":
        # 用独立 ORM 事务更改范围，不代表公开挂靠接口使用相同锁序。
        with SessionLocal() as db:
            assignment = db.scalar(select(MaintenanceProjectUserAssignment).join(
                SysUser, SysUser.id == MaintenanceProjectUserAssignment.user_id,
            ).where(
                MaintenanceProjectUserAssignment.project_id == case.project_id,
                SysUser.username == case.username,
                MaintenanceProjectUserAssignment.archived_at.is_(None),
            ).with_for_update())
            assignment.archived_at = datetime.now(timezone.utc)
            assignment.archived_by = "site-rights-admin"
            assignment.archive_reason = "Synthetic post-lock authorization regression"
            assignment.version += 1
            db.commit()
    else:
        changed = case.admin.put(f"/api/accounts/{case.username}", json={
            "overrides": {**case.overrides, "action_maintenance_expense_collection_upload": False},
        })
        assert changed.status_code == 200, changed.text


def _restore(case, revocation):
    if revocation == "scope":
        with SessionLocal() as db:
            user = db.scalar(select(SysUser).where(SysUser.username == case.username))
            db.add(MaintenanceProjectUserAssignment(
                assignment_id=str(uuid4()), project_id=case.project_id,
                responsibility_type="viewer", user_id=user.id, version=1,
                assigned_by="site-rights-admin", assignment_reason="Synthetic restoration",
            ))
            db.commit()
    else:
        restored = case.admin.put(f"/api/accounts/{case.username}", json={"overrides": case.overrides})
        assert restored.status_code == 200, restored.text
    _login(case.client, case.username, _PASSWORD)


def _hold(holder, case, lock_kind):
    if lock_kind == "global_advisory":
        holder.execute(select(func.pg_advisory_xact_lock(config.DATA_CHANGE_ADVISORY_LOCK_KEY)))
        return "pg_advisory_xact_lock"
    model, column, identity = {
        "workbook_state": (MaintenanceProjectWorkbookState, MaintenanceProjectWorkbookState.project_id, case.project_id),
        "site_header": (MaintenanceSiteIssue, MaintenanceSiteIssue.issue_id, case.issue_id),
        "site_line": (MaintenanceSiteIssueLine, MaintenanceSiteIssueLine.issue_line_id, case.site_line_id),
        "late_milestone": (MaintenanceCollectionMilestone, MaintenanceCollectionMilestone.milestone_id, case.milestone_id),
    }[lock_kind]
    holder.execute(select(model).where(column == identity).with_for_update()).scalar_one()
    return model.__tablename__


def _apply_after_revocation_during_proven_wait(case, lock_kind, revocation):
    before = _snapshot()
    with SessionLocal() as holder, engine.connect().execution_options(isolation_level="AUTOCOMMIT") as observer:
        holder_pid = holder.scalar(text("SELECT pg_backend_pid()"))
        expected_query = _hold(holder, case, lock_kind)
        with ThreadPoolExecutor(max_workers=1) as pool:
            future = pool.submit(_post, case)
            try:
                deadline = monotonic() + 12
                waiter = None
                while monotonic() < deadline:
                    waiters = observer.execute(text("""
                        SELECT pid, wait_event_type, query FROM pg_stat_activity
                        WHERE datname = current_database()
                          AND :holder = ANY(pg_blocking_pids(pid))
                    """), {"holder": holder_pid}).mappings().all()
                    if waiters:
                        assert len(waiters) == 1, waiters
                        waiter = waiters[0]
                        break
                    assert not future.done(), f"apply returned before a real lock wait: {future.result().text}"
                    Event().wait(0.01)
                assert waiter is not None, "did not observe a PostgreSQL lock wait"
                assert waiter["wait_event_type"] == "Lock"
                assert expected_query in waiter["query"], waiter
                # 宽行 SELECT 会超过 PostgreSQL 默认 1024 字节的活动文本上限。
                # pg_blocking_pids 已证明实际等待，不要求被截断的 FOR UPDATE 尾部。
                assert waiter["query"].startswith("SELECT ")
                _revoke(case, revocation)
                assert not future.done()
                assert holder_pid in observer.scalar(text("SELECT pg_blocking_pids(:pid)"), {"pid": waiter["pid"]})
                assert _snapshot() == before
            finally:
                holder.rollback()
            response = future.result(timeout=30)
    assert response.status_code == (403 if revocation == "scope" else 401), response.text
    assert _snapshot() == before, "revoked request changed facts, versions, receipts or business audits"
    _restore(case, revocation)
    first = _post(case)
    assert first.status_code == 200, first.text
    assert first.json()["replayed"] is False
    with SessionLocal() as db:
        assert db.get(MaintenanceSiteIssueLine, case.site_line_id).quantity == Decimal("3.250")
        assert db.scalar(select(func.count()).select_from(MaintenanceProjectWorkbookOperation)) == 1
        assert db.scalar(select(func.count()).select_from(MaintenanceProjectOperationAudit).where(
            MaintenanceProjectOperationAudit.action == "apply_completed",
        )) == 1
        if case.milestone_id:
            assert db.get(MaintenanceCollectionMilestone, case.milestone_id).planned_amount == Decimal("725.50")
    after = _snapshot()
    replay = _post(case)
    assert replay.status_code == 200, replay.text
    assert replay.json() == {**first.json(), "replayed": True}
    assert _snapshot() == after


@pytest.mark.parametrize("lock_kind", ["global_advisory", "workbook_state", "site_header", "site_line"])
@pytest.mark.parametrize("revocation", ["action_and_token", "scope"])
def test_site_write_rechecks_after_real_lock_wait(db, monkeypatch, lock_kind, revocation):
    with _case(db, monkeypatch) as case:
        _apply_after_revocation_during_proven_wait(case, lock_kind, revocation)


@pytest.mark.parametrize("revocation", ["action_and_token", "scope"])
def test_site_and_plan_atomic_rollback_after_later_sheet_lock(db, monkeypatch, revocation):
    """里程碑等锁发生在 06 编辑后；提交前复验须同时回滚两个 sheet。"""
    with _case(db, monkeypatch, mixed=True) as case:
        _apply_after_revocation_during_proven_wait(case, "late_milestone", revocation)


@pytest.mark.parametrize("editor", ["manager", "sales"])
def test_workbook_editor_exceptions_do_not_require_page_site_action(db, monkeypatch, editor):
    with _case(db, monkeypatch, editor=editor) as case:
        applied = _post(case)
        assert applied.status_code == 200, applied.text
        assert applied.json()["replayed"] is False
        with SessionLocal() as current:
            assert current.get(MaintenanceSiteIssueLine, case.site_line_id).quantity == Decimal("3.250")


def test_shared_admin_password_cannot_write_site_rows(db, monkeypatch):
    with _case(db, monkeypatch) as case:
        # 使用真实共享口令登录，不伪造身份或覆盖依赖。
        _login(case.client, "admin", master_api.get_settings().admin_password)
        before = _snapshot()
        response = _post(case)
        assert response.status_code == 403, response.text
        assert "实名" in response.text
        assert _snapshot() == before


def test_locked_existing_receipt_rechecks_real_identity_before_replay(db, monkeypatch):
    """直接调用服务进入锁内回执分支，验证真实数据库身份复验。

    HTTP 通常会提前拒绝已撤销 token；保留先前预检计划以模拟并发应用后的分支。
    """
    with _case(db, monkeypatch) as case:
        ident = verify_token(case.client.headers["Authorization"].removeprefix("Bearer "))
        with SessionLocal() as service_db:
            plan = master.validate_project_master_v2(service_db, project_id=case.project_id, data=case.data)
            service_db.rollback()
            initial = _post(case)
            assert initial.status_code == 200, initial.text
            before = _snapshot()
            _revoke(case, "action_and_token")

            def reauthorize():
                return master_api._reauthorize_site_workbook(
                    service_db, ident=ident, plan=plan, operated_by=case.username,
                )

            with pytest.raises(HTTPException) as rejected:
                master.apply_project_master_v2(
                    service_db, plan, operated_by=case.username,
                    import_batch_id=str(uuid4()), reauthorize=reauthorize,
                )
            assert rejected.value.status_code == 401
            service_db.rollback()
            assert _snapshot() == before
