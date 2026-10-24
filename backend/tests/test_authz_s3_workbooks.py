"""Whole-file field pairing and object scope using real HTTP/PG/PDP."""
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.api import maintenance_project_workbook_v3 as api
from app.authz import runtime
from app.config import get_settings
from app.models.authorization import SysAuthorizationGrant
from app.services import maintenance_project_workbook_v3 as workbook
from tests.test_authz_s3_checklist import setup, cerbos_url  # noqa: F401


@pytest.fixture
def v3(db, setup, monkeypatch):
    uid, identity, pid, client, headers = setup
    monkeypatch.setattr(get_settings(), 'authz_enabled_resources',
                        ['maintenance_acceptance', 'maintenance_project_roundtrip'])
    runtime.activate_configured_resources(db.get_bind())
    grant = SysAuthorizationGrant(user_id=uid, resource='maintenance_project_roundtrip',
        actions=['download_v3'], scope='all_projects',
        field_permissions=['data_purchase_cost', 'data_profit'], version=1)
    db.add(grant)
    db.commit()
    client.app.include_router(api.router, prefix='/api')
    return uid, identity, pid, client, headers, grant.id


def test_v3_explicit_complete_grant_without_customer_field_or_old_role_gate(db, v3):
    _, _, pid, client, headers, _ = v3
    response = client.get(f'/api/maintenance/projects/stable/{pid}/workbook-v3.xlsx', headers=headers)
    assert response.status_code == 200, response.text[:200]
    book = load_workbook(BytesIO(response.content), read_only=True)
    try:
        assert len(book.sheetnames) >= 6
    finally:
        book.close()


def test_whole_workbook_cannot_union_finance_field_from_other_resource(db, v3, monkeypatch):
    uid, _, pid, client, headers, gid = v3
    grant = db.get(SysAuthorizationGrant, gid)
    grant.field_permissions = ['data_purchase_cost']
    grant.version += 1
    db.add(SysAuthorizationGrant(user_id=uid, resource='maintenance_finance_workbook',
        actions=['download'], scope='all_projects', field_permissions=['data_profit'], version=1))
    db.commit()
    calls = []
    def forbidden_builder(*args):
        calls.append(True)
        raise AssertionError('whole-file SQL reached without matching fields')
    monkeypatch.setattr(workbook, 'build_project_workbook', forbidden_builder)
    response = client.get(f'/api/maintenance/projects/stable/{pid}/workbook-v3.xlsx', headers=headers)
    assert response.status_code == 403, response.text
    assert calls == []


def test_grant_scope_is_paired_with_same_complete_fields(db, v3, monkeypatch):
    uid, _, pid, client, headers, gid = v3
    grant = db.get(SysAuthorizationGrant, gid)
    grant.scope = 'assigned_or_sales_projects'  # This actor has no relationship to this project.
    grant.version += 1
    db.add(SysAuthorizationGrant(user_id=uid, resource='maintenance_finance_workbook',
        actions=['download'], scope='all_projects', field_permissions=['data_profit'], version=1))
    db.commit()
    response = client.get(f'/api/maintenance/projects/stable/{pid}/workbook-v3.xlsx', headers=headers)
    assert response.status_code == 403, response.text


@pytest.fixture
def finance(db, setup, monkeypatch):
    from datetime import date
    from decimal import Decimal
    from uuid import uuid4
    from app.api import maintenance_expense_collection_workbook as finance_api
    from app.models.maintenance_project import MaintenanceProjectContract
    from tests.test_maintenance_expense_collection_workbook import _expense
    uid, identity, pid, client, headers = setup
    monkeypatch.setattr(get_settings(), 'authz_enabled_resources',
                        ['maintenance_acceptance', 'maintenance_finance_workbook'])
    runtime.activate_configured_resources(db.get_bind())
    grant = SysAuthorizationGrant(user_id=uid, resource='maintenance_finance_workbook',
        actions=['download', 'validate', 'apply'], scope='all_projects', field_permissions=['data_profit'], version=1)
    db.add(grant)
    db.add(MaintenanceProjectContract(project_contract_id=str(uuid4()), project_id=pid,
        contract_id='C-1', contract_no='HT-001', amount_inc_tax=Decimal('10000'), included_in_total=True,
        status_mapping_state='mapped', status_mapping_version='v1', effective_from=date(2026, 1, 1),
        source='ledger', version=1))
    db.commit()
    _expense(db)
    client.app.include_router(finance_api.router, prefix='/api')
    return uid, identity, pid, client, headers, grant.id


def test_finance_workbook_complete_download_validate_apply_uses_single_grant(db, finance):
    from decimal import Decimal
    from app.models.maintenance import FProjectExpense
    from app.services import maintenance_expense_collection_workbook as finance_workbook
    _, _, pid, client, headers, _ = finance
    base = f'/api/maintenance/projects/stable/{pid}/expense-collection-workbook'
    download = client.get(base + '.xlsx', headers=headers)
    assert download.status_code == 200, download.text[:200]
    book = load_workbook(BytesIO(download.content))
    try:
        book[finance_workbook.SHEET_EXPENSE].cell(row=2, column=8, value=200)
        output = BytesIO()
        book.save(output)
    finally:
        book.close()
    for action in ['validate', 'apply']:
        response = client.post(base + '/' + action, headers=headers,
            files={'file': ('finance.xlsx', output.getvalue())})
        assert response.status_code == 200, response.text
    db.expire_all()
    expense = db.scalar(select(FProjectExpense))
    assert expense.amount_ex_tax == Decimal('200.00')
    assert expense.amount_inc_tax == Decimal('226.00')


def test_finance_read_grant_cannot_validate_or_apply_even_old_admin_role(db, finance):
    from app.models.system import SysUser
    uid, _, pid, client, headers, gid = finance
    grant = db.get(SysAuthorizationGrant, gid)
    grant.actions = ['download']
    grant.version += 1
    db.get(SysUser, uid).role = 'admin'
    db.commit()
    base = f'/api/maintenance/projects/stable/{pid}/expense-collection-workbook'
    download = client.get(base + '.xlsx', headers=headers)
    assert download.status_code == 200
    for action in ['validate', 'apply']:
        response = client.post(base + '/' + action, headers=headers,
            files={'file': ('finance.xlsx', download.content)})
        assert response.status_code == 403, response.text
