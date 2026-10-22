"""D-32 §5：免返标注在项目总表中仅展示，不产生 Excel 写入意图。"""

import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.config import get_settings
from app.models.maintenance_project import MaintenanceProjectAuditLog
from tests.test_maintenance_project_management_api import _admin_client
from tests.test_maintenance_project_master_v2_editable import _make_project_with_line, _save
from tests.test_maintenance_project_note_workbook import _patch, _send


NOTE_LABEL = "免返标注（仅网页维护，此处只读；仅作提醒，返还要求以明细标记和规则为准）"
OLD_NOTE_LABEL = "免返标注（仅网页维护，此处只读）"


def _book(client, pid, *, v2):
    sheet = "01_项目概览" if v2 else "01_项目基础信息"
    response = client.get(
        f"/api/maintenance/projects/stable/{pid}/master-workbook.xlsx",
        params={"sheets": sheet},
    )
    assert response.status_code == 200, response.text
    return load_workbook(io.BytesIO(response.content), data_only=False), sheet


def _note_cell(wb, sheet):
    cells = [row[1] for row in wb[sheet] if row[0].value == NOTE_LABEL]
    assert len(cells) == 1, "01 应包含唯一的只读免返标注行"
    return cells[0]


@pytest.mark.parametrize("v2", [True, False], ids=["v2", "v1"])
@pytest.mark.parametrize(("value", "expected"), [
    (None, None),
    ("硬盘免返\n其他按合同", "硬盘免返\n其他按合同"),
    ("旧\x0b标注\ufffe\uffff\n第二行", "旧标注\n第二行"),
    ("=仅作提醒", "=仅作提醒"),
    ("-48V", "-48V"),
    ("+5V", "+5V"),
    ("- 硬盘免返", "- 硬盘免返"),
    ("-1", "-1"),
    ("\n第二行开头", "\n第二行开头"),
    ("\t制表开头", "\t制表开头"),
    ("\r回车开头", "\r回车开头"),
    ("@备注", "@备注"),
    ("=SUM(A1)", "=SUM(A1)"),
    ('=HYPERLINK("https://example.invalid","提醒")', '=HYPERLINK("https://example.invalid","提醒")'),
    ("#N/A", "#N/A"),
])
def test_download_displays_no_return_note_as_readonly_plain_text(
    db, monkeypatch, v2, value, expected,
):
    monkeypatch.setattr(get_settings(), "maintenance_project_master_v2_enabled", v2)
    client = _admin_client(db)
    project, *_ = _make_project_with_line(db)
    # 合成历史脏数据；正常网页入口拒绝不可表示的控制字符。
    project.no_return_note = value
    project.project_note = value
    db.commit()

    wb, sheet = _book(client, project.project_id, v2=v2)
    try:
        cell = _note_cell(wb, sheet)
        assert cell.value == expected
        assert cell.data_type != "f"
        if expected is not None:
            assert cell.data_type == "s"
        assert cell.alignment.wrap_text is True
        assert wb[sheet].cell(cell.row, 1).alignment.wrap_text is True
        assert cell.fill.patternType == "solid"
        assert wb[sheet].protection.sheet is False
        if v2:
            project_note = next(row[1] for row in wb[sheet] if row[0].value == "项目备注")
            assert cell.row == project_note.row + 1
            assert cell.fill != project_note.fill
            assert project_note.value == expected
            if expected is not None:
                assert project_note.data_type == "s"
    finally:
        wb.close()


def _edit_readonly_row(wb, sheet, mutation):
    cell = _note_cell(wb, sheet)
    if mutation == "missing":
        # 上一版本没有此行；签名元数据保持原样，不重签。
        wb[sheet].delete_rows(cell.row)
    elif mutation == "old_label":
        # 模拟已下载的旧标签，保留原签名；改值仍不能开启免返标注回写。
        wb[sheet].cell(cell.row, 1).value = OLD_NOTE_LABEL
        cell.value = "旧标签下也不应保存"
    elif mutation != "unchanged":
        cell.value = {
            "modified": "Excel 不应保存的标注",
            "blank": None,
            "formula": "=1+1",
            "escaped_control": "无须校验的_x000B_只读值",
        }[mutation]


def _audits(db, pid):
    return list(db.scalars(select(MaintenanceProjectAuditLog).where(
        MaintenanceProjectAuditLog.project_id == pid,
    )))


@pytest.mark.parametrize("v2", [True, False], ids=["v2", "v1"])
@pytest.mark.parametrize("mutation", [
    "modified", "blank", "missing", "unchanged", "formula", "escaped_control", "old_label",
])
def test_validate_and_apply_ignore_no_return_note_without_master_audit(
    db, monkeypatch, v2, mutation,
):
    monkeypatch.setattr(get_settings(), "maintenance_project_master_v2_enabled", v2)
    client = _admin_client(db)
    project, *_ = _make_project_with_line(db)
    project.no_return_note = "仅在网页维护\n硬盘免返"
    project.project_note = "独立的项目备注"
    db.commit()
    pid = project.project_id
    before = client.get(f"/api/maintenance/projects/stable/{pid}").json()["project"]
    no_return_default = project.no_return_default
    assert _audits(db, pid) == []
    wb, sheet = _book(client, pid, v2=v2)
    _edit_readonly_row(wb, sheet, mutation)
    data = _save(wb)
    wb.close()

    for action in ("validate", "apply"):
        response = _send(client, pid, data, action=action)
        assert response.status_code == 200, response.text
        if v2 and action == "validate":
            assert response.json()["changes"] == []
            assert response.json()["project_note_updates"] == 0
        after = client.get(f"/api/maintenance/projects/stable/{pid}").json()["project"]
        for key in ("no_return_note", "project_note", "version"):
            assert after[key] == before[key]
        db.refresh(project)
        assert project.no_return_default == no_return_default
        assert _audits(db, pid) == []


@pytest.mark.parametrize("mutation", ["modified", "missing", "unchanged", "old_label"])
def test_readonly_no_return_note_does_not_block_project_note_writeback(db, monkeypatch, mutation):
    monkeypatch.setattr(get_settings(), "maintenance_project_master_v2_enabled", True)
    client = _admin_client(db)
    project, *_ = _make_project_with_line(db)
    project.no_return_note = "保留免返标注"
    project.project_note = "原项目备注"
    db.commit()
    pid = project.project_id
    wb, sheet = _book(client, pid, v2=True)
    _edit_readonly_row(wb, sheet, mutation)
    next(row[1] for row in wb[sheet] if row[0].value == "项目备注").value = "新的项目备注"
    data = _save(wb)
    wb.close()

    preview = _send(client, pid, data, action="validate")
    assert preview.status_code == 200, preview.text
    assert preview.json()["project_note_updates"] == 1
    assert [item["field"] for item in preview.json()["changes"]] == ["项目备注"]
    assert _audits(db, pid) == []
    applied = _send(client, pid, data)
    assert applied.status_code == 200, applied.text
    after = client.get(f"/api/maintenance/projects/stable/{pid}").json()["project"]
    assert after["no_return_note"] == "保留免返标注"
    assert after["project_note"] == "新的项目备注"
    assert after["version"] == 2
    audits = _audits(db, pid)
    assert len(audits) == 1
    assert audits[0].before_json["no_return_note"] == audits[0].after_json["no_return_note"] == "保留免返标注"
    assert audits[0].before_json["project_note"] == "原项目备注"
    assert audits[0].after_json["project_note"] == "新的项目备注"


@pytest.mark.parametrize("v2", [True, False], ids=["v2", "v1"])
def test_old_download_never_overwrites_later_web_no_return_note(db, monkeypatch, v2):
    monkeypatch.setattr(get_settings(), "maintenance_project_master_v2_enabled", v2)
    client = _admin_client(db)
    project, *_ = _make_project_with_line(db)
    project.no_return_note = "下载时的标注"
    db.commit()
    pid = project.project_id
    wb, _sheet = _book(client, pid, v2=v2)
    data = _save(wb)
    wb.close()
    changed = _patch(client, pid, 1, no_return_note="网页后来更新的标注")
    assert changed.status_code == 200, changed.text
    for action in ("validate", "apply"):
        response = _send(client, pid, data, action=action)
        assert response.status_code == 200, response.text
    after = client.get(f"/api/maintenance/projects/stable/{pid}").json()["project"]
    assert after["no_return_note"] == "网页后来更新的标注"
    assert after["version"] == 2
    assert len(_audits(db, pid)) == 1  # 仅网页更新审计。
