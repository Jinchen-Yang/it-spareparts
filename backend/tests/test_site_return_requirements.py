"""Workbook return rules remain auditable without delivery IDs or PN allocation."""

import io
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models.maintenance_bad_return import MaintenanceReturnObligation
from app.models.maintenance_project_operations import (
    MaintenanceProjectOperationAudit,
    MaintenanceSiteIssueReturnEvent,
)
from app.services import maintenance_project_master_workbook as master
from app.services import maintenance_project_operations as operations
from app.services import maintenance_site_return_requirements as requirements
from app.services.maintenance_return_receipts import receipt_summary, register_receipt
from tests.test_maintenance_project_master_v2_editable import (
    _make_project_with_line,
    _save,
    _site_issue,
)


def _download(db, pid, line_id):
    wb = load_workbook(
        io.BytesIO(
            master.build_project_master_v2(
                db, project_id=pid, sheets=(master.V2_SHEET_SITE,)
            )
        )
    )
    ws = wb[master.V2_SHEET_SITE]
    headers = {c.value: c.column for c in ws[1]}
    row = next(
        r
        for r in range(2, ws.max_row + 1)
        if ws.cell(r, headers["实体ID"]).value == line_id
    )
    return wb, ws, headers, row


def _apply(db, pid, wb):
    plan = master.validate_project_master_v2(db, project_id=pid, data=_save(wb))
    result = master.apply_project_master_v2(
        db, plan, operated_by="real-operator", import_batch_id=str(uuid4())
    )
    return plan, result


def test_workbook_return_requirement_roundtrip_audit_and_receipt_independence(db):
    project, part, _, _ = _make_project_with_line(db)
    issue, line = _site_issue(db, project, part, qty="2")
    # These are the actual production workbook identities: no delivery ID.
    assert line.delivery_line_id is None
    register_receipt(
        db,
        project_id=project.project_id,
        pn="DIFFERENT-RETURN-PN",
        qty=5,
        operated_by="receiver",
    )
    db.commit()
    original_version = line.version
    for flag, expected_qty, expected_flag in [
        ("是", 2, False),
        ("否", 0, True),
        ("", None, None),
    ]:
        wb, ws, h, row = _download(db, project.project_id, line.issue_line_id)
        ws.cell(row, h["是否应返还"]).value = flag
        plan, _ = _apply(db, project.project_id, wb)
        db.refresh(line)
        assert line.no_return is expected_flag
        assert (
            receipt_summary(db, project_id=project.project_id)["project_total_qty"]
            == "5.000"
        )
        _, exported, eh, er = _download(db, project.project_id, line.issue_line_id)
        assert exported.cell(er, eh["应返数量"]).value == (
            expected_qty if expected_qty is not None else "—"
        )
        assert "逐行分配" in exported.cell(er, eh["返还单号"]).value
        if flag == "否":
            assert any("实际返还数量和历史均保留" in w for w in plan.warnings)
    assert line.version == original_version + 3
    audits = list(
        db.scalars(
            select(MaintenanceProjectOperationAudit).where(
                MaintenanceProjectOperationAudit.entity_id == line.issue_line_id,
                MaintenanceProjectOperationAudit.action
                == "return_obligation_corrected",
            )
        )
    )
    assert len(audits) == 3
    assert {a.operated_by for a in audits} == {"real-operator"}
    events = list(
        db.scalars(
            select(MaintenanceSiteIssueReturnEvent).where(
                MaintenanceSiteIssueReturnEvent.issue_id == issue.issue_id
            )
        )
    )
    assert len(events) == 3
    assert all(
        e.payload["schema_version"] == requirements.SCHEMA and e.consumed_at
        for e in events
    )
    assert list(db.scalars(select(MaintenanceReturnObligation))) == []
    # A correction must not make the existing panel void action unusable.
    operations.void_site_issue(
        db,
        project_id=project.project_id,
        issue_id=issue.issue_id,
        version=issue.version,
        idempotency_key=str(uuid4()),
        reason="撤销领用",
        operated_by="real-operator",
    )
    assert (
        receipt_summary(db, project_id=project.project_id)["project_total_qty"]
        == "5.000"
    )


@pytest.mark.parametrize("field", ["应返数量", "返还状态", "返还单号"])
def test_workbook_readonly_return_columns_cannot_be_silently_changed(db, field):
    project, part, _, _ = _make_project_with_line(db)
    _, line = _site_issue(db, project, part)
    wb, ws, h, row = _download(db, project.project_id, line.issue_line_id)
    ws.cell(row, h[field], "用户修改")
    with pytest.raises(master.WorkbookError) as error:
        master.validate_project_master_v2(
            db, project_id=project.project_id, data=_save(wb)
        )
    assert error.value.code == "readonly_field_changed"


def test_site_search_uses_part_identity_and_keeps_remark_separate(db):
    project, part, _, _ = _make_project_with_line(db)
    _, line = _site_issue(db, project, part)
    part.description = "当前型号描述"
    part.brand = "测试品牌"
    part.unit = "件"
    line.remark = "领用现场备注"
    line.pn = "历史PN文本"
    line.no_return = False
    db.commit()
    payload = operations.search_site_issues(
        db,
        project_id=project.project_id,
        q_text=None,
        workflow_statuses=["confirmed"],
        page=1,
        page_size=20,
    )
    actual = payload["rows"][0]["lines"][0]
    assert actual["pn"] == "历史PN文本"
    assert actual["description"] == "当前型号描述"
    assert actual["description_source"] == "current_master_data"
    assert actual["remark"] == "领用现场备注"
    assert actual["brand"] == "测试品牌"
    assert actual["unit"] == "件"
    assert actual["return_requirement"]["requirement_status"] == "required"


def test_export_actual_receipts_separately_and_reject_edited_snapshot(db):
    project, part, _, _ = _make_project_with_line(db)
    _, line = _site_issue(db, project, part)
    register_receipt(
        db,
        project_id=project.project_id,
        pn="=PLAIN-TEXT-PN",
        qty=5,
        operated_by="receiver",
    )
    db.commit()
    wb, _, _, _ = _download(db, project.project_id, line.issue_line_id)
    ledger = wb[master.V2_SHEET_RETURN_LEDGER]
    assert ledger["B2"].value == "5.000"
    assert ledger["B3"].value == "5.000"
    pn_cell = next(
        cell for row in ledger for cell in row if cell.value == "=PLAIN-TEXT-PN"
    )
    assert pn_cell.data_type == "s"
    # Another receipt after export is not an edit to the file's read-only snapshot.
    register_receipt(
        db, project_id=project.project_id, pn="NEW-PN", qty=2, operated_by="receiver"
    )
    db.commit()
    _apply(db, project.project_id, wb)
    assert (
        receipt_summary(db, project_id=project.project_id)["project_total_qty"]
        == "7.000"
    )
    ledger["B2"] = "500"
    with pytest.raises(master.WorkbookError) as error:
        master.validate_project_master_v2(
            db, project_id=project.project_id, data=_save(wb)
        )
    assert error.value.code == "readonly_field_changed"
