"""销售模板通过 M4 注册回调往返，保持领域编辑与权限约束。"""

from datetime import date
from io import BytesIO
import os
import shutil
import subprocess

from openpyxl import load_workbook
import pytest
from sqlalchemy import select

from app.models.fields_audience import FieldAnchorUsage
from app.models.presales import PresalesInquiry, PresalesInquiryLine, PresalesWorkbookExport
from app.services import audience, field_pool
from app.services.presales import common, inquiries, inquiry_workbook as workbook
from app.services.workbook import operations, registry
from app.services.workbook.registry import ApplyFailed, Conflict, Forbidden as WorkbookForbidden, WorkbookSpec
from app.services.workbook.schema import WorkbookError
from app.services.workflow.registry import Forbidden, InvalidTransition
from tests.test_presales_inquiries import _offer, presales_parts, presales_users


@pytest.fixture(autouse=True)
def workbook_registration(monkeypatch):
    monkeypatch.setattr(registry, "_specs", {})
    registry.register(workbook.INQUIRY_SCHEMA.name, WorkbookSpec(
        workbook.INQUIRY_SCHEMA, workbook.load_current, workbook.apply_inquiry, workbook.can_access))


def _ctx(db, user):
    return common.resolve_actor(db, user.id)[1]


def _edit(content, fn):
    book = load_workbook(BytesIO(content))
    sheet = book["询价明细"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    fn(sheet, columns)
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def _set(sheet, columns, row, **values):
    for header, value in values.items():
        sheet.cell(row, columns[header], value)


def _preview(db, user, content, inquiry=None):
    return operations.preview(db, workbook.INQUIRY_SCHEMA.name, content, ctx=_ctx(db, user),
        subject_type=inquiries.SUBJECT_TYPE if inquiry else None, subject_id=inquiry.id if inquiry else None)


def _apply(db, user, preview, *, deletes=None):
    return operations.apply(db, preview["operation_id"], _ctx(db, user), preview["plan_digest"],
                            confirm_deletes=deletes or [])


def test_empty_template_creates_draft_exact_match_and_pending_raw(db, presales_users, presales_parts):
    user, part = presales_users.sales, presales_parts[0]
    content = workbook.export_template(db, _ctx(db, user))
    def fill(sheet, cols):
        _set(sheet, cols, 2, **{"PN 原文": part.pn_std, "数量": "2.500", "期望交期": "2026-12-01"})
        _set(sheet, cols, 3, **{"PN 原文": "0000-UNKNOWN", "数量": 1})
    content = _edit(content, fill)
    preview = _preview(db, user, content)
    assert preview["plan"]["counts"]["added"] == 2
    assert db.scalars(select(PresalesInquiry)).all() == []
    result = _apply(db, user, preview)
    inquiry = db.get(PresalesInquiry, result["response"]["inquiry_id"])
    lines = inquiries.lines_for(db, inquiry.id)
    assert inquiry.status == "draft"
    assert lines[0].part_id == part.id and lines[0].expected_delivery == date(2026, 12, 1)
    assert lines[1].part_id is None and lines[1].pn_raw == "0000-UNKNOWN"
    replay = _preview(db, user, content)
    assert replay["replayed"] and replay["response"] == result["response"]
    assert len(db.scalars(select(PresalesInquiry)).all()) == 1


def test_delete_insert_modify_only_confirmed_deletions(db, presales_users, presales_parts):
    user = presales_users.sales
    inquiry = inquiries.create_draft(db, user.id, lines=[{"part_id": part.id} for part in presales_parts[:2]])
    original = inquiries.lines_for(db, inquiry.id)
    def fill(sheet, cols):
        sheet.delete_rows(3)
        _set(sheet, cols, 2, **{"数量": 7})
        _set(sheet, cols, 4, **{"PN 原文": presales_parts[2].pn_std, "数量": 9})
    content = _edit(workbook.export_inquiry(db, inquiry.id, _ctx(db, user)), fill)
    preview = _preview(db, user, content, inquiry)
    assert preview["plan"]["counts"] == {"added": 1, "modified": 1, "unchanged": 0, "deleted": 1, "conflict": 0}
    deletes = [row["delete_key"] for row in preview["plan"]["rows"] if row["status"] == "deleted"]
    _apply(db, user, preview, deletes=deletes)
    rows = inquiries.lines_for(db, inquiry.id)
    assert [(row.part_id, row.qty) for row in rows] == [(presales_parts[0].id, 7), (presales_parts[2].id, 9)]
    assert db.get(PresalesInquiryLine, original[1].id).status == "deleted"


def test_unconfirmed_delete_preserves_line_and_omitted_columns(db, presales_users, presales_parts):
    user = presales_users.sales
    inquiry = inquiries.create_draft(db, user.id, lines=[{"part_id": presales_parts[0].id, "qty": 4, "notes": "保留"},
                                                         {"part_id": presales_parts[1].id, "qty": 3}])
    content = workbook.export_inquiry(db, inquiry.id, _ctx(db, user), visible_columns={"lines": ["qty"]})
    book = load_workbook(BytesIO(content))
    assert [cell.value for cell in book["询价明细"][1]] == ["_row_token", "PN", "数量"]
    def fill(sheet, cols):
        sheet.delete_rows(3)
        _set(sheet, cols, 2, **{"数量": 8})
    preview = _preview(db, user, _edit(content, fill), inquiry)
    _apply(db, user, preview)
    rows = inquiries.lines_for(db, inquiry.id)
    assert len(rows) == 2 and rows[0].notes == "保留" and rows[0].qty == 8


def test_submitted_rejected_and_withdrawn_edit_uses_existing_stale_diff(db, presales_users, presales_parts):
    user = presales_users.sales
    inquiry = inquiries.create_draft(db, user.id, assignee_ids=[presales_users.buyer.id],
                                    lines=[{"part_id": presales_parts[0].id}])
    old = workbook.export_inquiry(db, inquiry.id, _ctx(db, user))
    inquiries.submit(db, inquiry.id, user.id, inquiry.version)
    with pytest.raises(InvalidTransition):
        workbook.export_inquiry(db, inquiry.id, _ctx(db, user))
    with pytest.raises(WorkbookForbidden):
        _preview(db, user, old, inquiry)
    offer = _offer(db, inquiry, inquiries.lines_for(db, inquiry.id)[0], presales_users.buyer)
    inquiries.withdraw(db, inquiry.id, user.id, inquiry.version)
    content = _edit(workbook.export_inquiry(db, inquiry.id, _ctx(db, user)),
                    lambda sheet, cols: _set(sheet, cols, 2, **{"数量": 5}))
    _apply(db, user, _preview(db, user, content, inquiry))
    assert inquiry.status == "revising" and offer.status == "published"
    inquiries.resubmit(db, inquiry.id, user.id, inquiry.version)
    assert offer.status == "stale"


def test_current_owner_and_parent_version_are_rechecked(db, presales_users, presales_parts):
    user = presales_users.sales
    inquiry = inquiries.create_draft(db, user.id, lines=[{"part_id": presales_parts[0].id}])
    with pytest.raises(Forbidden):
        workbook.export_inquiry(db, inquiry.id, _ctx(db, presales_users.other_sales))
    content = _edit(workbook.export_inquiry(db, inquiry.id, _ctx(db, user)),
                    lambda sheet, cols: _set(sheet, cols, 2, **{"数量": 9}))
    preview = _preview(db, user, content, inquiry)
    inquiries.update_draft(db, inquiry.id, user.id, inquiry.version, notes="服务端变更")
    with pytest.raises(Conflict):
        _apply(db, user, preview)


def test_existing_rows_cannot_target_another_inquiry_or_change_bound_pn(db, presales_users, presales_parts):
    user = presales_users.sales
    first = inquiries.create_draft(db, user.id, lines=[{"part_id": presales_parts[0].id}])
    second = inquiries.create_draft(db, user.id, lines=[{"part_id": presales_parts[1].id}])
    content = _edit(workbook.export_inquiry(db, first.id, _ctx(db, user)),
                    lambda sheet, cols: _set(sheet, cols, 2, **{"数量": 4}))
    preview = _preview(db, user, content, second)
    with pytest.raises(ApplyFailed):
        _apply(db, user, preview)
    content = _edit(workbook.export_inquiry(db, first.id, _ctx(db, user)),
                    lambda sheet, cols: _set(sheet, cols, 2, **{"PN 原文": presales_parts[1].pn_std}))
    with pytest.raises(ApplyFailed):
        _apply(db, user, _preview(db, user, content, first))
    assert inquiries.lines_for(db, first.id)[0].part_id == presales_parts[0].id


def test_custom_fields_are_saved_and_remembered_per_part(db, presales_users, presales_parts):
    user, part = presales_users.sales, presales_parts[0]
    inquiry = inquiries.create_draft(db, user.id, lines=[{"part_id": part.id}])
    def fill(sheet, cols):
        sheet.cell(1, sheet.max_column + 1, "+质保要求")
        sheet.cell(2, sheet.max_column, "三年")
    content = _edit(workbook.export_inquiry(db, inquiry.id, _ctx(db, user)), fill)
    _apply(db, user, _preview(db, user, content, inquiry))
    line = inquiries.lines_for(db, inquiry.id)[0]
    values = field_pool.read_values(db, [field_pool.Subject(inquiries.SUBJECT_TYPE, f"line:{line.id}", "line")],
                                    _ctx(db, user), common.bypass)[f"line:{line.id}"]
    assert values[0]["value"] == "三年"
    usage = db.scalar(select(FieldAnchorUsage))
    assert usage.anchor == f"part:{part.id}"
    definition = field_pool.list_defs(db, inquiries.SUBJECT_TYPE, "line", _ctx(db, user))[0]
    field_pool.update_field(db, definition, auto_include=False)
    next_inquiry = inquiries.create_draft(db, user.id, lines=[{"part_id": part.id}])
    book = load_workbook(BytesIO(workbook.export_inquiry(db, next_inquiry.id, _ctx(db, user))))
    assert "+质保要求" in [cell.value for cell in book["询价明细"][1]]
    assert book["询价明细"].cell(2, book["询价明细"].max_column).value is None


@pytest.mark.skipif(os.environ.get("RUN_SOFFICE_ROUNDTRIP") != "1" or not shutil.which("soffice"),
                    reason="需要启用 LibreOffice 往返")
def test_libreoffice_inquiry_template_roundtrip(db, presales_users, presales_parts, tmp_path):
    user = presales_users.sales
    content = _edit(workbook.export_template(db, _ctx(db, user)),
                    lambda sheet, cols: _set(sheet, cols, 2, **{"PN 原文": presales_parts[0].pn_std, "数量": 2}))
    source = tmp_path / "inquiry.xlsx"
    source.write_bytes(content)
    output = tmp_path / "converted"
    output.mkdir()
    completed = subprocess.run(["soffice", "-env:UserInstallation=file:///tmp/t16-lo-profile", "--headless",
                                "--convert-to", "xlsx", "--outdir", str(output), str(source)],
                               capture_output=True, text=True, timeout=60)
    assert completed.returncode == 0, completed.stderr
    result = _apply(db, user, _preview(db, user, (output / source.name).read_bytes()))
    assert result["response"]["line_count"] == 1


def test_private_fields_by_other_author_are_not_exported_or_overwritten(db, presales_users, presales_parts):
    user = presales_users.sales
    boss_ctx = _ctx(db, presales_users.boss)
    rule = audience.create_rule(db, "private", [], [], boss_ctx)
    definition = field_pool.define_field(db, subject_type=inquiries.SUBJECT_TYPE, level="line",
        name="内部参数", owner_scope="team", viewer=boss_ctx, default_rule_id=rule)
    inquiry = inquiries.create_draft(db, user.id, lines=[{"part_id": presales_parts[0].id, "notes": "保留备注"}])
    line = inquiries.lines_for(db, inquiry.id)[0]
    subject = field_pool.Subject(inquiries.SUBJECT_TYPE, f"line:{line.id}", "line")
    field_pool.save_values(db, subject, [{"field_def_id": definition.id, "value": "隐藏参数"}], boss_ctx)
    content = workbook.export_inquiry(db, inquiry.id, _ctx(db, user))
    book = load_workbook(BytesIO(content))
    text = str(list(book["询价明细"].values)) + str(list(book["_meta"].values))
    assert "隐藏参数" not in text
    changed = _edit(content, lambda sheet, cols: _set(sheet, cols, 2, **{"数量": 2}))
    _apply(db, user, _preview(db, user, changed, inquiry))
    assert inquiries.lines_for(db, inquiry.id)[0].notes == "保留备注"
    assert field_pool.read_values(db, [subject], boss_ctx, common.bypass)[subject.subject_id][0]["value"] == "隐藏参数"
    changed = _edit(content, lambda sheet, cols: _set(sheet, cols, 2, **{"+内部参数": "尝试覆盖"}))
    with pytest.raises(WorkbookForbidden):
        _apply(db, user, _preview(db, user, changed, inquiry))


def test_unchanged_workbook_does_not_bump_versions(db, presales_users, presales_parts):
    user = presales_users.sales
    inquiry = inquiries.create_draft(db, user.id, lines=[{"part_id": presales_parts[0].id}])
    before = inquiry.version
    content = workbook.export_inquiry(db, inquiry.id, _ctx(db, user))
    _apply(db, user, _preview(db, user, content, inquiry))
    assert inquiry.version == before


def test_custom_field_failure_rolls_back_all_rows(db, presales_users, presales_parts):
    user = presales_users.sales
    definition = field_pool.define_field(db, subject_type=inquiries.SUBJECT_TYPE, level="line",
                                         name="数值要求", value_type="number", viewer=_ctx(db, user))
    inquiry = inquiries.create_draft(db, user.id, lines=[{"part_id": part.id} for part in presales_parts[:2]])
    def fill(sheet, cols):
        _set(sheet, cols, 2, **{"数量": 8, "+数值要求": "123"})
        _set(sheet, cols, 3, **{"+数值要求": "不是数字"})
    content = _edit(workbook.export_inquiry(db, inquiry.id, _ctx(db, user)), fill)
    preview = _preview(db, user, content, inquiry)
    with pytest.raises(ApplyFailed):
        _apply(db, user, preview)
    rows = inquiries.lines_for(db, inquiry.id)
    assert rows[0].qty == 1
    values = field_pool.read_values(db, [field_pool.Subject(inquiries.SUBJECT_TYPE, f"line:{row.id}", "line")
                                       for row in rows], _ctx(db, user), common.bypass)
    assert all(not entries for entries in values.values())
    assert definition.id is not None


def test_added_rows_cannot_default_missing_required_quantity(db, presales_users):
    user = presales_users.sales
    content = workbook.export_template(db, _ctx(db, user), visible_columns={"lines": ["pn_raw"]})
    changed = _edit(content, lambda sheet, cols: _set(sheet, cols, 2, **{"PN 原文": "NEW-PN"}))
    preview = _preview(db, user, changed)
    assert preview["plan"]["errors"]
    with pytest.raises(WorkbookError):
        _apply(db, user, preview)
    assert db.scalars(select(PresalesInquiry)).all() == []


def test_history_replay_rejects_custom_field_after_visibility_narrows(db, presales_users, presales_parts):
    user, boss = presales_users.sales, presales_users.boss
    boss_ctx = _ctx(db, boss)
    definition = field_pool.define_field(db, subject_type=inquiries.SUBJECT_TYPE, level="line",
                                         name="公开后收窄", owner_scope="team", viewer=boss_ctx)
    inquiry = inquiries.create_draft(db, user.id, lines=[{"part_id": presales_parts[0].id}])
    line = inquiries.lines_for(db, inquiry.id)[0]
    subject = field_pool.Subject(inquiries.SUBJECT_TYPE, f"line:{line.id}", "line")
    field_pool.save_values(db, subject, [{"field_def_id": definition.id, "value": "先前公开内容"}], boss_ctx)
    content = _edit(workbook.export_inquiry(db, inquiry.id, _ctx(db, user)),
                    lambda sheet, cols: _set(sheet, cols, 2, **{"数量": 2}))
    preview = _preview(db, user, content, inquiry)
    _apply(db, user, preview)
    field_pool.save_values(db, subject, [{"field_def_id": definition.id, "value": "先前公开内容",
                                         "rule": {"mode": "private"}}], boss_ctx)
    with pytest.raises(WorkbookForbidden):
        _preview(db, user, content, inquiry)
    with pytest.raises(WorkbookForbidden):
        _apply(db, user, preview)


def test_template_requires_sales_inquiry_permission(db, presales_users):
    with pytest.raises(Forbidden):
        workbook.export_template(db, _ctx(db, presales_users.buyer))


@pytest.mark.parametrize("create_from_template", [False, True])
@pytest.mark.parametrize("rename", [False, True])
def test_replay_rechecks_private_fields_for_created_drafts_and_renamed_fields(
        db, presales_users, presales_parts, create_from_template, rename):
    user, boss = presales_users.sales, presales_users.boss
    boss_ctx = _ctx(db, boss)
    definition = field_pool.define_field(db, subject_type=inquiries.SUBJECT_TYPE, level="line",
                                         name="历史参数", owner_scope="team", viewer=boss_ctx)
    inquiry = None if create_from_template else inquiries.create_draft(
        db, user.id, lines=[{"part_id": presales_parts[0].id}])
    content = workbook.export_template(db, _ctx(db, user)) if inquiry is None else workbook.export_inquiry(
        db, inquiry.id, _ctx(db, user))
    def fill(sheet, cols):
        _set(sheet, cols, 2, **{"PN 原文": presales_parts[0].pn_std, "数量": 2, "+历史参数": "先前可见"})
    content = _edit(content, fill)
    preview = _preview(db, user, content, inquiry)
    result = _apply(db, user, preview)
    created = db.get(PresalesInquiry, result["response"]["inquiry_id"])
    line = inquiries.lines_for(db, created.id)[0]
    if rename:
        field_pool.update_field(db, definition, name="改名后的参数")
    subject = field_pool.Subject(inquiries.SUBJECT_TYPE, f"line:{line.id}", "line")
    field_pool.save_values(db, subject, [{"field_def_id": definition.id, "value": "先前可见",
                                         "rule": {"mode": "private"}}], boss_ctx)
    with pytest.raises(WorkbookForbidden):
        _preview(db, user, content, inquiry)
    with pytest.raises(WorkbookForbidden):
        _apply(db, user, preview)


def test_replay_keeps_line_note_visible_to_its_author(db, presales_users, presales_parts):
    user = presales_users.sales
    ctx = _ctx(db, user)
    inquiry = inquiries.create_draft(db, user.id, lines=[{"part_id": presales_parts[0].id, "notes": "历史备注"}])
    content = _edit(workbook.export_inquiry(db, inquiry.id, ctx),
                    lambda sheet, cols: _set(sheet, cols, 2, **{"数量": 2}))
    preview = _preview(db, user, content, inquiry)
    _apply(db, user, preview)
    line = inquiries.lines_for(db, inquiry.id)[0]
    line.notes_audience_rule_id = audience.create_rule(db, "exclude", [ctx.user_id], [],
                                                     _ctx(db, presales_users.boss))
    db.flush()
    assert _preview(db, user, content, inquiry)["replayed"] is True
    assert _apply(db, user, preview)["replayed"] is True


@pytest.mark.parametrize("rename", [False, True])
def test_first_deleted_row_preview_rechecks_exported_custom_field_access(db, presales_users, presales_parts, rename):
    user, boss_ctx = presales_users.sales, _ctx(db, presales_users.boss)
    definition = field_pool.define_field(db, subject_type=inquiries.SUBJECT_TYPE, level="line",
                                         name="导出过的参数", owner_scope="team", viewer=boss_ctx)
    inquiry = inquiries.create_draft(db, user.id, lines=[{"part_id": presales_parts[0].id}])
    line = inquiries.lines_for(db, inquiry.id)[0]
    subject = field_pool.Subject(inquiries.SUBJECT_TYPE, f"line:{line.id}", "line")
    field_pool.save_values(db, subject, [{"field_def_id": definition.id, "value": "不可泄漏的旧参数"}], boss_ctx)
    content = workbook.export_inquiry(db, inquiry.id, _ctx(db, user))
    exported = db.scalar(select(PresalesWorkbookExport))
    assert exported.scope["fields"] == [[line.id, definition.id]]
    assert "不可泄漏的旧参数" not in str(exported.scope)
    if rename:
        field_pool.update_field(db, definition, name="改名后收窄")
    field_pool.save_values(db, subject, [{"field_def_id": definition.id, "value": "不可泄漏的旧参数",
                                         "rule": {"mode": "private"}}], boss_ctx)
    deleted = _edit(content, lambda sheet, cols: sheet.delete_rows(2))
    with pytest.raises(WorkbookForbidden):
        _preview(db, user, deleted, inquiry)
    assert inquiries.lines_for(db, inquiry.id)[0].id == line.id


def test_unselected_custom_fields_do_not_create_export_authorization_obligations(db, presales_users, presales_parts):
    user, boss_ctx = presales_users.sales, _ctx(db, presales_users.boss)
    definition = field_pool.define_field(db, subject_type=inquiries.SUBJECT_TYPE, level="line",
                                         name="未选择的字段", owner_scope="team", viewer=boss_ctx)
    inquiry = inquiries.create_draft(db, user.id, lines=[{"part_id": presales_parts[0].id}])
    line = inquiries.lines_for(db, inquiry.id)[0]
    subject = field_pool.Subject(inquiries.SUBJECT_TYPE, f"line:{line.id}", "line")
    field_pool.save_values(db, subject, [{"field_def_id": definition.id, "value": "未导出"}], boss_ctx)
    content = workbook.export_inquiry(db, inquiry.id, _ctx(db, user), visible_columns={"lines": ["qty"]})
    assert db.scalar(select(PresalesWorkbookExport)).scope["fields"] == []
    field_pool.save_values(db, subject, [{"field_def_id": definition.id, "value": "未导出",
                                         "rule": {"mode": "private"}}], boss_ctx)
    content = _edit(content, lambda sheet, cols: _set(sheet, cols, 2, **{"数量": 3}))
    preview = _preview(db, user, content, inquiry)
    # 可以授权预览；外部字段变化仍由原有版本摘要保守标为冲突。
    assert preview["plan"]["counts"]["conflict"] == 1
    assert preview["plan"]["errors"] == []


@pytest.mark.parametrize("scoped", [False, True])
def test_other_exporters_private_baseline_is_rejected_before_first_preview(db, presales_users, presales_parts, scoped):
    user, boss_ctx = presales_users.sales, _ctx(db, presales_users.boss)
    rule = audience.create_rule(db, "private", [], [], boss_ctx)
    definition = field_pool.define_field(db, subject_type=inquiries.SUBJECT_TYPE, level="line",
        name="主管私密参数", owner_scope="team", viewer=boss_ctx, default_rule_id=rule)
    inquiry = inquiries.create_draft(db, user.id, lines=[{"part_id": presales_parts[0].id}])
    line = inquiries.lines_for(db, inquiry.id)[0]
    subject = field_pool.Subject(inquiries.SUBJECT_TYPE, f"line:{line.id}", "line")
    field_pool.save_values(db, subject, [{"field_def_id": definition.id, "value": "主管导出的私密内容"}], boss_ctx)
    content = workbook.export_inquiry(db, inquiry.id, boss_ctx)
    deleted = _edit(content, lambda sheet, cols: sheet.delete_rows(2))
    with pytest.raises(WorkbookForbidden):
        _preview(db, user, deleted, inquiry if scoped else None)


def test_export_scope_uses_the_same_field_snapshot_as_the_file(db, presales_users, presales_parts, monkeypatch):
    user, boss_ctx = presales_users.sales, _ctx(db, presales_users.boss)
    definition = field_pool.define_field(db, subject_type=inquiries.SUBJECT_TYPE, level="line",
                                         name="并发收窄的字段", owner_scope="team", viewer=boss_ctx)
    inquiry = inquiries.create_draft(db, user.id, lines=[{"part_id": presales_parts[0].id}])
    line = inquiries.lines_for(db, inquiry.id)[0]
    subject = field_pool.Subject(inquiries.SUBJECT_TYPE, f"line:{line.id}", "line")
    field_pool.save_values(db, subject, [{"field_def_id": definition.id, "value": "文件内的旧值"}], boss_ctx)
    original = workbook.export_workbook
    def export_and_narrow(*args, **kwargs):
        content = original(*args, **kwargs)
        field_pool.save_values(db, subject, [{"field_def_id": definition.id, "value": "文件内的旧值",
                                             "rule": {"mode": "private"}}], boss_ctx)
        return content
    monkeypatch.setattr(workbook, "export_workbook", export_and_narrow)
    content = workbook.export_inquiry(db, inquiry.id, _ctx(db, user))
    assert db.scalar(select(PresalesWorkbookExport)).scope["fields"] == [[line.id, definition.id]]
    with pytest.raises(WorkbookForbidden):
        _preview(db, user, _edit(content, lambda sheet, cols: sheet.delete_rows(2)), inquiry)


def test_clearing_own_custom_field_keeps_first_receipt_replayable(db, presales_users, presales_parts):
    user, ctx = presales_users.sales, _ctx(db, presales_users.sales)
    inquiry = inquiries.create_draft(db, user.id, lines=[{"part_id": presales_parts[0].id,
        "field_values": [{"new_name": "作者自己的字段", "value": "清空前"}]}])
    content = workbook.export_inquiry(db, inquiry.id, ctx)
    def clear(sheet, cols):
        sheet.cell(2, cols["+作者自己的字段"]).value = None
    content = _edit(content, clear)
    preview = _preview(db, user, content, inquiry)
    first = _apply(db, user, preview)
    replay = _preview(db, user, content, inquiry)
    assert replay["replayed"] is True and replay["response"] == first["response"]
