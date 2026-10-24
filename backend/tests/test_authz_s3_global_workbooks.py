"""Actual global file bytes, SQL field projection and current action contracts."""
import io
import json
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from sqlalchemy import event, select, update

from app.models.authorization import SysAuthorizationGrant
from app.models.maintenance import FMaintenanceOrder
from tests.test_authz_s3_source_cost import source_cost, give, captured  # noqa: F401
from tests.test_authz_s3_analysis import analysis_client  # noqa: F401
from tests.test_authz_s3_workspace import workspace  # noqa: F401
from tests.test_authz_s3_projects import canonical  # noqa: F401
from tests.test_authz_s3_project_foundation import projects  # noqa: F401
from tests.authz_pdp_helpers import cerbos_url  # noqa: F401


def test_order_export_sql_scope_and_fields_do_not_borrow_read_or_contract_grants(db, source_cost):
    client, headers, actor = source_cost
    give(db, actor, ['export_orders'])
    response, statements = captured(db, client, headers, '/api/maintenance/orders/export', {})
    assert response.status_code == 200, response.text
    workbook = load_workbook(io.BytesIO(response.content), read_only=True, data_only=True)
    rows = list(workbook['维保订单'].values)
    assert len(rows) == 2 and rows[1][2] == 'WBDD-S3-other'
    assert rows[1][6] is None
    lines = list(workbook['订单明细'].values)
    assert len(lines) == 2 and all(v is None for v in lines[1][13:34])
    for query in statements:
        assert not any(name in query for name in ['.cost_amount', '.unit_cost', '.end_customer', '.receiver_phone', '.receiver_address']), query
    assert client.get('/api/maintenance/lines', headers=headers, params={'project': 'same-source-project'}).status_code == 403
    assert response.headers['cache-control'] == 'no-store'


@pytest.mark.parametrize('path,params', [('/api/maintenance/export-workbook', {'contract': 'XSDD-S3'}),
    ('/api/maintenance/export-workbooks', {})])
def test_contract_files_require_whole_source_scope_and_complete_fields(db, source_cost, path, params):
    client, headers, actor = source_cost
    row = give(db, actor, ['export_workbook'], fields=['data_purchase_cost', 'data_profit'])
    assert client.get(path, headers=headers, params={**params, 'view': 'authorization'}).status_code == 403
    assert client.get(path, headers=headers, params=params).status_code == 403
    row.scope, row.version = 'all_sources', row.version + 1
    db.commit()
    response = client.get(path, headers=headers, params=params)
    assert response.status_code == 200, response.text
    if path.endswith('workbooks'):
        archive = ZipFile(io.BytesIO(response.content))
        data = archive.read(next(n for n in archive.namelist() if n.endswith('.xlsx')))
    else:
        data = response.content
    book = load_workbook(io.BytesIO(data), read_only=True)
    assert '备件明细-氚云' in book.sheetnames
    row.field_permissions, row.version = ['data_purchase_cost'], row.version + 1
    db.commit()
    assert client.get(path, headers=headers, params=params).status_code == 403


@pytest.mark.parametrize('path', ['/api/maintenance/roundtrip-template', '/api/maintenance/roundtrip-templates'])
def test_roundtrip_template_is_independent_and_whole_file_fields_are_mandatory(db, source_cost, path):
    client, headers, actor = source_cost
    # A grant for regular cost files cannot supply this module's fields/action.
    give(db, actor, ['export_workbook'], fields=['data_purchase_cost', 'data_profit', 'data_customer'], scope='all_sources')
    assert client.get(path, headers=headers).status_code == 403
    row = give(db, actor, ['template_download'], resource='maintenance_cost_roundtrip', scope='all_contracts',
        fields=['data_purchase_cost', 'data_profit'])
    assert client.get(path, headers=headers).status_code == 403
    row.field_permissions, row.version = ['data_purchase_cost', 'data_profit', 'data_customer'], row.version + 1
    db.commit()
    view = client.get(path, headers=headers, params={'view': 'authorization'})
    assert view.status_code == 200, view.text
    assert set(view.json()) == {'locator', 'authorization_context'}
    response = client.get(path, headers=headers)
    assert response.status_code == 200, response.text
    if path.endswith('templates'):
        archive = ZipFile(io.BytesIO(response.content))
        data = archive.read(next(n for n in archive.namelist() if n.endswith('.xlsx')))
    else:
        data = response.content
    book = load_workbook(io.BytesIO(data), read_only=True)
    assert '02_维保订单' in book.sheetnames and '05_人工成本回填' in book.sheetnames


def test_files_share_readonly_snapshot_and_revoke_before_headers_closes_spool(db, source_cost, monkeypatch):
    from sqlalchemy.orm import Session
    from sqlalchemy import text
    from app.services import maintenance_export
    from app.models.system import SysUser
    client, headers, actor = source_cost
    uid = actor.id
    give(db, actor, ['export_orders'])
    original = maintenance_export.build_workbook
    files = []
    def build(snapshot, *args, **kwargs):
        assert snapshot.scalar(text('SHOW transaction_isolation')) == 'repeatable read'
        assert snapshot.scalar(text('SHOW transaction_read_only')) == 'on'
        result = original(snapshot, *args, **kwargs)
        files.append(result)
        with Session(db.get_bind()) as administrator:
            administrator.execute(update(SysUser).where(SysUser.id == uid).values(token_version=SysUser.token_version + 1))
            administrator.commit()
        return result
    monkeypatch.setattr(maintenance_export, 'build_workbook', build)
    response = client.get('/api/maintenance/orders/export', headers=headers)
    assert response.status_code == 401, response.text
    assert len(files) == 1 and files[0].closed


def test_order_file_query_precondition_rejects_changed_date_selector(db, source_cost):
    client, headers, actor = source_cost
    give(db, actor, ['export_orders'])
    path = '/api/maintenance/orders/export'
    view = client.get(path, headers=headers, params={'view': 'authorization'})
    assert view.status_code == 200, view.text
    ctx = view.json()['authorization_context']
    precondition = {k: ctx[k] for k in ('catalog_version', 'authorization_epoch', 'context')}
    response = client.get(path, headers=headers, params={'date_from': '2026-01-01', 'date_to': '2026-01-31',
        'authorization_precondition': json.dumps(precondition)})
    assert response.status_code == 409, response.text


def test_recompute_is_independent_global_action_and_rolls_back_late_revocation(db, source_cost, monkeypatch):
    from sqlalchemy.orm import Session
    from app.services import maintenance_cost
    from app.models.maintenance import FMaintenanceLine
    client, headers, actor = source_cost
    row = give(db, actor, ['recompute'], fields=['data_purchase_cost'])
    path = '/api/maintenance/recompute'
    assert client.post(path, headers=headers, params={'view': 'authorization'}).status_code == 403
    row.scope, row.version = 'all_sources', row.version + 1
    db.commit()
    view = client.post(path, headers=headers, params={'view': 'authorization'})
    assert view.status_code == 200, view.text
    before = list(db.execute(select(FMaintenanceLine.id, FMaintenanceLine.cost_source, FMaintenanceLine.cost_amount).order_by(FMaintenanceLine.id)))
    resource_id = row.id
    original = maintenance_cost.recompute
    def revoke_after_staged_changes(transaction, **kwargs):
        result = original(transaction, **kwargs)
        with Session(db.get_bind()) as administrator:
            administrator.execute(update(SysAuthorizationGrant).where(SysAuthorizationGrant.id == resource_id)
                .values(actions=[], version=SysAuthorizationGrant.version + 1))
            administrator.commit()
        return result
    monkeypatch.setattr(maintenance_cost, 'recompute', revoke_after_staged_changes)
    response = client.post(path, headers=headers)
    assert response.status_code == 403, response.text
    db.expire_all()
    after = list(db.execute(select(FMaintenanceLine.id, FMaintenanceLine.cost_source, FMaintenanceLine.cost_amount).order_by(FMaintenanceLine.id)))
    assert before == after
    monkeypatch.setattr(maintenance_cost, 'recompute', original)
    db.execute(update(SysAuthorizationGrant).where(SysAuthorizationGrant.id == resource_id).values(actions=['recompute']))
    db.commit()
    response = client.post(path, headers=headers)
    assert response.status_code == 200, response.text
