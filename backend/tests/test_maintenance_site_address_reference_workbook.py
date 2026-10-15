"""D-29 download references stay scoped, permissioned and outside write identity."""

import io
import uuid
from datetime import date, datetime, timezone
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select

from app import config
from app.api import maintenance_project_master_workbook as master_api
from app.models.maintenance import (
    FMaintenanceLine, FMaintenanceOrder, MaintenanceDemandDeleteIntent,
    MaintenanceDemandTombstone,
)
from app.models.maintenance_project_operations import (
    MaintenanceSiteIssueLine, MaintenanceProjectOperationAudit,
)
from app.models.maintenance_source_assignment import MaintenanceSourceOrderAssignment
from app.models.system import SysUser
from app.services import maintenance_project_master_workbook as master
from app.services.maintenance_site_address_reference import (
    DEMAND_NOT_FOUND, MULTIPLE_ADDRESSES,
)
from tests.test_maintenance_expense_collection_workbook import _client, _PASSWORD
from tests.test_maintenance_project_master_v2_editable import (
    _append_site_line, _make_project_with_line, _site_issue,
)
from tests.test_maintenance_site_address_workbook import _apply, _headers, _row, _write


REFERENCE = master.V2_SITE_DEMAND_ADDRESS_COLUMN
SOURCE_ADDRESS = "需求单收货地址：合成园区A楼"
ACTUAL_ADDRESS = "实际领用地址：机房B楼"
_BASE = "/api/maintenance/projects/stable"


def _seed(db):
    project, part, order, demand = _make_project_with_line(db)
    order.order_no = "WBDD-20261012-0001"
    order.receiver_address = SOURCE_ADDRESS
    issue, line = _site_issue(db, project, part)
    issue.site_location = ACTUAL_ADDRESS
    db.commit()
    return project, part, order, demand, issue, line


def _extra_demand(db, project, part, *, address="另一收货地址", number=2):
    batch_id = db.scalar(select(FMaintenanceOrder.import_batch_id).limit(1))
    order = FMaintenanceOrder(
        raw_order_id=f"address-order-{uuid.uuid4()}",
        order_no=f"WBDD-20261012-{number:04d}",
        order_date=date(2026, 10, 12), data_status="已生效",
        receiver_address=address, import_batch_id=batch_id,
    )
    db.add(order)
    db.flush()
    demand = FMaintenanceLine(
        raw_line_id=f"address-line-{uuid.uuid4()}", order_id=order.id,
        line_no=1, part_id=part.id, pn_std=part.pn_std,
        qty=Decimal("1"), return_qty=Decimal("0"), is_active=True,
        import_batch_id=batch_id,
    )
    assignment = MaintenanceSourceOrderAssignment(
        assignment_id=str(uuid.uuid4()), project_id=project.project_id,
        source_order_id=order.raw_order_id, is_active=True, created_by="reference-test",
    )
    db.add_all([demand, assignment])
    db.commit()
    return order, demand, assignment


def _download(db, project, *, full=False, visible=True):
    data = master.build_project_master_v2(
        db, project_id=project.project_id,
        sheets=None if full else (master.V2_SHEET_SITE,),
        include_demand_addresses=visible,
    )
    book = load_workbook(io.BytesIO(data))
    return book, book[master.V2_SHEET_SITE]


def _reference(sheet, line):
    return sheet.cell(_row(sheet, line), _headers(sheet)[REFERENCE])


@pytest.mark.parametrize("full", [False, True])
def test_full_and_site_download_append_visible_locked_reference_without_overwriting_actual(db, full):
    project, _part, _order, _demand, _issue, line = _seed(db)
    _book, sheet = _download(db, project, full=full)
    assert sheet.cell(1, sheet.max_column).value == REFERENCE
    cell = _reference(sheet, line)
    assert cell.column == sheet.max_column
    assert cell.value == SOURCE_ADDRESS
    assert cell.protection.locked is True
    assert not sheet.column_dimensions[get_column_letter(cell.column)].hidden
    assert sheet.cell(_row(sheet, line), _headers(sheet)["领用地址"]).value == ACTUAL_ADDRESS


def test_same_project_stable_part_maps_multiple_addresses_and_explicit_demand_number(db):
    project, part, order, _demand, issue, line = _seed(db)
    other_order, _other_demand, _assignment = _extra_demand(db, project, part)
    _, sheet = _download(db, project)
    assert _reference(sheet, line).value == MULTIPLE_ADDRESSES
    line.demand_order_no = order.order_no
    db.commit()
    _, sheet = _download(db, project)
    assert _reference(sheet, line).value == SOURCE_ADDRESS
    line.demand_order_no = other_order.order_no
    db.commit()
    _, sheet = _download(db, project)
    assert _reference(sheet, line).value == "另一收货地址"
    db.refresh(issue)
    assert issue.site_location == ACTUAL_ADDRESS


def test_same_part_in_other_project_never_supplies_address_or_satisfies_explicit_reference(db):
    project, part, _order, _demand, _issue, line = _seed(db)
    other, _other_part, _other_order, _other_line = _make_project_with_line(db)
    cross_order, _, _ = _extra_demand(db, other, part, address="其他项目机密地址", number=3)
    _, sheet = _download(db, project)
    assert _reference(sheet, line).value == SOURCE_ADDRESS
    line.source_order_id = cross_order.raw_order_id
    db.commit()
    _, sheet = _download(db, project)
    assert _reference(sheet, line).value == DEMAND_NOT_FOUND


@pytest.mark.parametrize("invalid", ["inactive_line", "cancelled_head", "tombstone", "inactive_assignment"])
def test_ineligible_demand_is_excluded_even_without_cutover_and_explicit_hint_does_not_fallback(
    db, monkeypatch, invalid,
):
    monkeypatch.setattr(master_api.get_settings(), "maintenance_cutover_enabled", False)
    project, part, _order, _demand, _issue, line = _seed(db)
    other_order, other_line, assignment = _extra_demand(db, project, part)
    if invalid == "inactive_line":
        other_line.is_active = False
    elif invalid == "cancelled_head":
        other_order.data_status = "已取消"
    elif invalid == "inactive_assignment":
        assignment.is_active = False
        assignment.version += 1
        assignment.archived_by = "reference-test"
        assignment.archived_at = datetime.now(timezone.utc)
    else:
        now = datetime.now(timezone.utc)
        intent = MaintenanceDemandDeleteIntent(
            intent_id=str(uuid.uuid4()), idempotency_key=f"address-delete-{uuid.uuid4()}",
            request_digest="a" * 64, selection_digest="b" * 64,
            status="executed", reason="合成需求单已作废", operated_by="reference-test",
            header_count=1, line_count=1, created_at=now, expires_at=now,
        )
        db.add(intent)
        db.flush()
        db.add(MaintenanceDemandTombstone(
            source_order_id=other_order.raw_order_id, delete_intent_id=intent.intent_id,
            version_digest="d" * 64,
            deleted_by="reference-test", delete_reason="合成需求单已作废",
            deleted_at=datetime.now(timezone.utc), version=1,
        ))
    db.commit()
    _, sheet = _download(db, project)
    assert _reference(sheet, line).value == SOURCE_ADDRESS
    line.demand_order_no = other_order.order_no
    db.commit()
    _, sheet = _download(db, project)
    assert _reference(sheet, line).value == DEMAND_NOT_FOUND


def test_reference_source_change_does_not_change_site_base_token_quantity_or_cost(db):
    project, _part, order, _demand, issue, line = _seed(db)
    line.unit_cost = Decimal("20.00")
    line.cost_amount = Decimal("40.00")
    line.unit_cost_ex_tax = Decimal("20.00")
    line.cost_amount_ex_tax = Decimal("40.00")
    line.unit_cost_inc_tax = Decimal("22.60")
    line.cost_amount_inc_tax = Decimal("45.20")
    line.cost_source = "maint_demand"
    db.commit()
    _, before = _download(db, project)
    base = before.cell(_row(before, line), _headers(before)["基线令牌"]).value
    entity = before.cell(_row(before, line), _headers(before)["实体ID"]).value
    order.receiver_address = "收货地址变了，实际领用事实没变"
    db.commit()
    _, after = _download(db, project)
    assert _reference(after, line).value == order.receiver_address
    assert after.cell(_row(after, line), _headers(after)["基线令牌"]).value == base
    assert after.cell(_row(after, line), _headers(after)["实体ID"]).value == entity
    for row in (issue, line):
        db.refresh(row)
    assert issue.site_location == ACTUAL_ADDRESS
    assert line.quantity == Decimal("2")
    assert line.cost_amount_inc_tax == Decimal("45.20")


@pytest.mark.parametrize("mode", ["edit_reference", "remove_reference"])
def test_reference_edit_or_old_template_missing_column_roundtrips_without_repricing(db, monkeypatch, mode):
    project, _part, _order, _demand, issue, line = _seed(db)
    book, sheet = _download(db, project)
    if mode == "edit_reference":
        _reference(sheet, line).value = "用户改了参考列，不是实际领用地址"
    else:
        sheet.delete_cols(_headers(sheet)[REFERENCE])
    from app.services import maintenance_consumption_cost
    monkeypatch.setattr(maintenance_consumption_cost, "resolve_lines",
                        lambda *a, **kw: pytest.fail("参考列变化不应重新取价"))
    before = (issue.version, line.version, line.quantity, line.cost_amount_inc_tax)
    audit_count = db.scalar(select(func.count()).select_from(MaintenanceProjectOperationAudit))
    plan, _ = _apply(db, project, book)
    db.refresh(issue)
    db.refresh(line)
    assert plan.site_flags == ()
    assert not plan.conflicts
    assert (issue.version, line.version, line.quantity, line.cost_amount_inc_tax) == before
    assert issue.site_location == ACTUAL_ADDRESS
    assert db.scalar(select(func.count()).select_from(MaintenanceProjectOperationAudit)) == audit_count


def test_reference_is_not_a_site_deletion_identity_and_existing_delete_contract_still_works(db):
    project, part, _order, _demand, issue, first = _seed(db)
    second = _append_site_line(db, issue, part)
    book, sheet = _download(db, project)
    sheet.delete_rows(_row(sheet, second))
    _reference(sheet, first).value = "参考内容被改变"
    _apply(db, project, book)
    db.refresh(first)
    db.refresh(second)
    assert first.is_active is True
    assert second.is_active is False


def test_reference_only_value_on_an_otherwise_empty_row_is_ignored_not_an_issue(db):
    project, _part, _order, _demand, issue, line = _seed(db)
    book, sheet = _download(db, project)
    sheet.cell(sheet.max_row + 1, _headers(sheet)[REFERENCE], "粘贴到空白行的参考地址")
    before = db.scalar(select(func.count()).select_from(MaintenanceSiteIssueLine))
    plan, _ = _apply(db, project, book)
    db.refresh(line)
    db.refresh(issue)
    assert plan.site_flags == ()
    assert db.scalar(select(func.count()).select_from(MaintenanceSiteIssueLine)) == before
    assert line.is_active is True
    assert issue.site_location == ACTUAL_ADDRESS


def test_equals_prefixed_and_long_source_address_remains_full_plain_text(db):
    project, _part, order, _demand, _issue, line = _seed(db)
    address = "=合成库房\n" + "完整地址" * 100
    order.receiver_address = address
    db.commit()
    _, sheet = _download(db, project)
    cell = _reference(sheet, line)
    assert cell.value == address
    assert cell.data_type == "s"


def test_builder_defaults_to_no_source_address_visibility(db, monkeypatch):
    project, _part, _order, _demand, _issue, line = _seed(db)
    from app.services.maintenance_site_address_reference import AddressReferenceIndex
    monkeypatch.setattr(AddressReferenceIndex, "from_rows",
                        lambda *_: pytest.fail("无权限不得建立来源地址索引"))
    data = master.build_project_master_v2(db, project_id=project.project_id,
                                         sheets=(master.V2_SHEET_SITE,))
    sheet = load_workbook(io.BytesIO(data))[master.V2_SHEET_SITE]
    assert _reference(sheet, line).value == "无客户信息权限"
    assert SOURCE_ADDRESS not in [cell.value for row in sheet for cell in row]


@pytest.mark.parametrize("full", [False, True])
@pytest.mark.parametrize("customer", [False, True])
def test_real_http_download_uses_customer_permission_independently_of_profit_and_page(db, monkeypatch, full, customer):
    monkeypatch.setattr(config, "ENABLE_RBAC", True)
    monkeypatch.setattr(master_api.get_settings(), "maintenance_project_master_v2_enabled", True)
    project, _part, _order, _demand, _issue, line = _seed(db)
    client = _client(db, username=f"ref-{uuid.uuid4().hex[:10]}", overrides={
        "page_maintenance": True, "data_purchase_cost": True,
        "data_profit": True, "data_customer": customer,
    })
    response = client.get(f"{_BASE}/{project.project_id}/master-workbook.xlsx",
                          params=None if full else {"sheets": master.V2_SHEET_SITE})
    assert response.status_code == 200, response.text
    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook[master.V2_SHEET_SITE]
    assert _reference(sheet, line).value == (SOURCE_ADDRESS if customer else "无客户信息权限")
    if not customer:
        assert SOURCE_ADDRESS not in [cell.value for ws in workbook for row in ws for cell in row]


def test_real_http_customer_permission_does_not_bypass_project_scope(db, monkeypatch):
    monkeypatch.setattr(config, "ENABLE_RBAC", True)
    monkeypatch.setattr(master_api.get_settings(), "maintenance_project_master_v2_enabled", True)
    _project, _part, _order, _demand, _issue, _line = _seed(db)
    client = _client(db, username=f"ref-scope-{uuid.uuid4().hex[:8]}", overrides={
        "page_maintenance": True, "data_purchase_cost": True,
        "data_profit": True, "data_customer": True,
    })
    other, _other_part, _other_order, _other_line = _make_project_with_line(db)
    response = client.get(f"{_BASE}/{other.project_id}/master-workbook.xlsx",
                          params={"sheets": master.V2_SHEET_SITE})
    assert response.status_code == 403


def test_real_http_rbac_disabled_preserves_full_source_visibility(db, monkeypatch):
    monkeypatch.setattr(config, "ENABLE_RBAC", False)
    monkeypatch.setattr(master_api.get_settings(), "maintenance_project_master_v2_enabled", True)
    project, _part, _order, _demand, _issue, line = _seed(db)
    client = _client(db, username=f"ref-no-rbac-{uuid.uuid4().hex[:8]}", overrides={
        "page_maintenance": True, "data_purchase_cost": True,
        "data_profit": True, "data_customer": False,
    })
    response = client.get(f"{_BASE}/{project.project_id}/master-workbook.xlsx",
                          params={"sheets": master.V2_SHEET_SITE})
    assert response.status_code == 200, response.text
    sheet = load_workbook(io.BytesIO(response.content))[master.V2_SHEET_SITE]
    assert _reference(sheet, line).value == SOURCE_ADDRESS


def test_real_http_named_admin_can_reference_source_without_customer_override(db, monkeypatch):
    monkeypatch.setattr(config, "ENABLE_RBAC", True)
    monkeypatch.setattr(master_api.get_settings(), "maintenance_project_master_v2_enabled", True)
    project, _part, _order, _demand, _issue, line = _seed(db)
    username = f"ref-admin-{uuid.uuid4().hex[:8]}"
    client = _client(db, username=username, overrides={
        "page_maintenance": True, "data_purchase_cost": True,
        "data_profit": True, "data_customer": False,
    })
    user = db.scalar(select(SysUser).where(SysUser.username == username))
    user.role = "admin"
    db.commit()
    login = client.post("/api/auth/login", json={"username": username, "password": _PASSWORD})
    assert login.status_code == 200
    assert login.json()["role"] == "admin"
    client.headers["Authorization"] = f"Bearer {login.json()['token']}"
    response = client.get(f"{_BASE}/{project.project_id}/master-workbook.xlsx",
                          params={"sheets": master.V2_SHEET_SITE})
    assert response.status_code == 200, response.text
    sheet = load_workbook(io.BytesIO(response.content))[master.V2_SHEET_SITE]
    assert _reference(sheet, line).value == SOURCE_ADDRESS
