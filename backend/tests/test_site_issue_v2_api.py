"""Public workflow tests for server-owned site-consumption documents."""

import io
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from decimal import Decimal
from uuid import uuid4

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.exc import DBAPIError

from app import auth
from app.api import maintenance_project_operations
from app.auth import hash_password
from app.business_time import business_today
from app.models.dimensions import DimPart
from app.models.inventory import Inventory
from app.models.maintenance_bad_return import MaintenanceReturnObligation
from app.models.maintenance_project import (
    MaintenanceProject,
    MaintenanceProjectUserAssignment,
)
from app.models.maintenance_project_operations import (
    MaintenanceProjectOperationAudit,
    MaintenanceSiteIssue,
    MaintenanceSiteIssueCommand,
    MaintenanceSiteIssueDeliverySource,
    MaintenanceSiteIssueLine,
    MaintenanceSiteIssueReturnEvent,
)
from app.models.purchase import FPurchaseLine, FPurchaseOrder
from app.models.system import SysAccessLog, SysImportBatch, SysUser
from app.security import UserContext
from app.services import maintenance_boss_board as board
from app.services import maintenance_project_master_workbook as master
from app.services import maintenance_project_operations as operations


def _client(
    db,
    *,
    username: str = "site_issue_v2_admin",
    role: str = "admin",
    permissions: dict | None = None,
) -> TestClient:
    db.add(
        SysUser(
            username=username,
            role=role,
            display_name="合成现场领用管理员",
            password_hash=hash_password("synthetic-password-123"),
            permissions=permissions,
        )
    )
    db.commit()
    app = FastAPI()
    app.include_router(auth.router, prefix="/api")
    app.include_router(maintenance_project_operations.router, prefix="/api")
    app.include_router(maintenance_project_operations.site_issue_router, prefix="/api")
    client = TestClient(app)
    login = client.post(
        "/api/auth/login",
        json={"username": username, "password": "synthetic-password-123"},
    )
    assert login.status_code == 200, login.text
    client.headers["Authorization"] = f"Bearer {login.json()['token']}"
    return client


def _client_for_existing_user(db, *, username: str) -> TestClient:
    assert db.query(SysUser).filter_by(username=username, is_active=True).one_or_none()
    app = FastAPI()
    app.include_router(auth.router, prefix="/api")
    app.include_router(maintenance_project_operations.router, prefix="/api")
    app.include_router(maintenance_project_operations.site_issue_router, prefix="/api")
    client = TestClient(app)
    login = client.post(
        "/api/auth/login",
        json={"username": username, "password": "synthetic-password-123"},
    )
    assert login.status_code == 200, login.text
    client.headers["Authorization"] = f"Bearer {login.json()['token']}"
    return client


def _project(db, *, project_id: str) -> MaintenanceProject:
    row = MaintenanceProject(
        project_id=project_id,
        project_code=f"SYNTH-{project_id[-8:]}",
        display_name="合成现场领用项目",
        lifecycle_status="ongoing",
    )
    db.add(row)
    db.commit()
    return row


def _delivery_source(
    db,
    *,
    project: MaintenanceProject,
    delivery_line_id: str = "synthetic-delivery-line-001",
    quantity: str = "5",
    linked_purchase_line_id: int | None = None,
) -> MaintenanceSiteIssueDeliverySource:
    part = DimPart(pn_std=f"PN-{delivery_line_id[-8:]}")
    db.add(part)
    db.flush()
    row = MaintenanceSiteIssueDeliverySource(
        delivery_line_id=delivery_line_id,
        adapter_key="synthetic_delivery_v1",
        project_id=project.project_id,
        source_order_id=f"synthetic-order-{delivery_line_id[-24:]}",
        source_line_id=f"synthetic-line-{delivery_line_id[-25:]}",
        delivery_no=f"SYNTH-{delivery_line_id[-24:]}",
        delivery_date=date(2026, 8, 8),
        part_id=part.id,
        pn=part.pn_std,
        serial_number="SYNTH-SN-001",
        delivered_quantity=quantity,
        linked_purchase_line_id=linked_purchase_line_id,
        mapping_state="ready",
        mapping_version="synthetic-delivery-map-v1",
    )
    db.add(row)
    db.commit()
    return row


def _purchase_evidence(db, *, part: DimPart) -> FPurchaseLine:
    batch = SysImportBatch(
        filename="synthetic-site-issue-v2-purchase.xlsx",
        file_type="purchase",
        file_hash="synthetic-site-issue-v2-purchase",
        status="success",
    )
    db.add(batch)
    db.flush()
    order = FPurchaseOrder(
        raw_order_id="synthetic-site-issue-v2-purchase-order",
        order_no="SYNTH-PO-SITE-ISSUE-V2",
        order_date=date(2026, 8, 8),
        data_status="已生效",
        import_batch_id=batch.id,
        is_tax_inclusive=False,
    )
    db.add(order)
    db.flush()
    line = FPurchaseLine(
        raw_line_id="synthetic-site-issue-v2-purchase-line",
        order_id=order.id,
        part_id=part.id,
        pn_std=part.pn_std,
        qty="10",
        unit_price="20",
        import_batch_id=batch.id,
    )
    db.add(line)
    db.flush()
    return line


def _create_draft(
    client: TestClient,
    *,
    project_id: str,
    delivery_line_id: str,
    quantity: str,
    key: str,
) -> dict:
    response = client.post(
        f"/api/maintenance/site-issues/projects/{project_id}",
        json={
            "idempotency_key": key,
            "issue_date": "2026-08-09",
            "receiver": "合成接收人",
            "issued_by": "合成发出人",
            "site_location": "合成现场",
            "lines": [
                {"delivery_line_id": delivery_line_id, "quantity": quantity}
            ],
            "reason": "建立合成草稿",
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def _workbook_issue(
    db,
    *,
    project: MaintenanceProject,
    issue_no: str,
) -> tuple[MaintenanceSiteIssue, MaintenanceSiteIssueLine]:
    """06 表建的领用单形态：source=workbook、无发货来源、无幂等字段、无返还事件。"""
    part = DimPart(pn_std=f"PN-WB-{issue_no[-4:]}")
    db.add(part)
    db.flush()
    issue = MaintenanceSiteIssue(
        issue_id=str(uuid4()),
        project_id=project.project_id,
        issue_no=issue_no,
        issue_date=date(2026, 8, 20),
        raw_status="已确认",
        status_mapping_state="mapped",
        normalized_status="confirmed",
        status_mapping_version="workbook-manual-v1",
        source="workbook",
        import_batch_id="synthetic-master-workbook-batch",
        created_by="synthetic-uploader",
        version=1,
    )
    db.add(issue)
    db.flush()
    line = MaintenanceSiteIssueLine(
        issue_line_id=f"{issue.issue_id}:1",
        issue_id=issue.issue_id,
        line_no=1,
        part_id=part.id,
        pn=part.pn_std,
        quantity=Decimal("2"),
        no_return=False,
        is_active=True,
        # 成本结果七列同生同灭，且 unit_cost/cost_amount 是 ex_tax 的旧别名
        unit_cost=Decimal("50.00"),
        cost_amount=Decimal("100.00"),
        unit_cost_ex_tax=Decimal("50.00"),
        unit_cost_inc_tax=Decimal("56.50"),
        cost_amount_ex_tax=Decimal("100.00"),
        cost_amount_inc_tax=Decimal("113.00"),
        cost_source="manual",
        algorithm_version="workbook-manual-v1",
    )
    db.add(line)
    db.commit()
    return issue, line


_FULL_VIEW = UserContext(
    user_id="synthetic-full-view",
    role="admin",
    permissions={"page_maintenance": True, "data_purchase_cost": True, "data_profit": True},
    is_authenticated=True,
)


def _site_sheet_entity_ids(db, project_id: str) -> set[str]:
    """06 表导出里出现的领用行实体 ID（D-01 的「导出」读模型）。"""
    content = master.build_project_master_v2(
        db, project_id=project_id, sheets=(master.V2_SHEET_SITE,)
    )
    # 导出会锁项目工作簿状态行；释放它，随后走 API 的作废才能拿到同一把锁
    db.commit()
    worksheet = load_workbook(io.BytesIO(content))[master.V2_SHEET_SITE]
    headers = {cell.value: cell.column for cell in worksheet[1]}
    return {
        str(worksheet.cell(row, headers["实体ID"]).value or "")
        for row in range(2, worksheet.max_row + 1)
    }


def _workspace_requisition_cost(db, project_id: str) -> str:
    """项目工作台 metrics.site_requisition_known_cost（面板健康带「已领用」取数）。"""
    workspace = operations.project_workspace(
        db, project_id=project_id, as_of=business_today(), user_ctx=_FULL_VIEW
    )
    db.commit()
    return workspace["project"]["metrics"]["site_requisition_known_cost"]


def _board_requisition_cost(db, project_id: str) -> Decimal:
    """展示板项目卡 requisition_cost_inc_tax（maintenance_boss_board 已领用成本口径）。"""
    rows = board.projects(
        db, user_ctx=_FULL_VIEW, allowed_project_ids={project_id}
    )["rows"]
    db.commit()
    stat = next(row for row in rows if row["project_id"] == project_id)[
        "requisition_cost_inc_tax"
    ]
    assert stat["state"] == "ready", stat
    return Decimal(str(stat["value"]))


def test_create_site_issue_generates_identity_and_only_saves_a_draft(db):
    project = _project(db, project_id="project-site-issue-v2-create")
    delivery = _delivery_source(db, project=project)
    client = _client(db)
    request = {
        "idempotency_key": "synthetic-site-issue-create-001",
        "issue_date": "2026-08-09",
        "receiver": "合成接收人",
        "issued_by": "合成发出人",
        "site_location": "合成现场 A",
        "lines": [
            {
                "delivery_line_id": delivery.delivery_line_id,
                "quantity": "2",
            }
        ],
        "reason": "建立现场领用草稿",
    }

    created = client.post(
        f"/api/maintenance/site-issues/projects/{project.project_id}",
        json=request,
    )

    assert created.status_code == 201, created.text
    payload = created.json()
    assert payload["project_id"] == project.project_id
    assert payload["workflow_status"] == "draft"
    assert payload["issue_no"].startswith("LYD-20260809-")
    assert payload["idempotent_replay"] is False
    assert payload["lines"] == [
        {
            **payload["lines"][0],
            "delivery_line_id": delivery.delivery_line_id,
            "source_order_id": delivery.source_order_id,
            "source_line_id": delivery.source_line_id,
            "pn": delivery.pn,
            "quantity": "2.000",
            "cost_source": None,
            "cost_amount_ex_tax": None,
            "cost_amount_inc_tax": None,
        }
    ]
    assert payload["lines"][0]["issue_line_id"]

    saved = db.get(MaintenanceSiteIssue, payload["issue_id"])
    saved_line = db.get(MaintenanceSiteIssueLine, payload["lines"][0]["issue_line_id"])
    assert saved is not None and saved.normalized_status == "draft"
    assert saved_line is not None and saved_line.cost_source is None

    legacy_status_bypass = client.patch(
        f"/api/maintenance/projects/stable/site-issues/{payload['issue_id']}/status",
        json={
            "version": payload["version"],
            "raw_status": "confirmed",
            "normalized_status": "confirmed",
            "status_mapping_version": "legacy-status-route",
            "reason": "新版单据不能绕过确认命令",
        },
    )
    assert legacy_status_bypass.status_code == 400, legacy_status_bypass.text
    db.expire_all()
    assert db.get(MaintenanceSiteIssue, payload["issue_id"]).normalized_status == "draft"

    replayed = client.post(
        f"/api/maintenance/site-issues/projects/{project.project_id}",
        json=request,
    )
    assert replayed.status_code == 201, replayed.text
    assert replayed.json()["issue_id"] == payload["issue_id"]
    assert replayed.json()["idempotent_replay"] is True


def test_delivery_candidate_search_is_post_only_and_fails_closed_without_adapter_rows(db):
    project = _project(db, project_id="project-site-issue-v2-candidates")
    client = _client(db, username="site_issue_v2_candidates_admin")
    path = (
        f"/api/maintenance/site-issues/projects/{project.project_id}"
        "/candidates/search"
    )

    unavailable = client.post(path, json={"page": 1, "page_size": 50})
    assert unavailable.status_code == 200, unavailable.text
    assert unavailable.json() == {
        "adapter": {
            "key": "synthetic_delivery_v1",
            "state": "unavailable",
            "production_ready": False,
            "detail": "真实 WBDD/仓库发货适配器尚未接入，系统不会按项目名猜测",
        },
        "rows": [],
        "total": 0,
        "page": 1,
        "page_size": 50,
    }
    assert client.get(path).status_code == 405

    delivery = _delivery_source(
        db,
        project=project,
        delivery_line_id="synthetic-delivery-line-candidate",
        quantity="5",
    )
    available = client.post(path, json={"q": "candidate", "page": 1, "page_size": 20})
    assert available.status_code == 200, available.text
    payload = available.json()
    assert payload["adapter"] == {
        "key": "synthetic_delivery_v1",
        "state": "synthetic_ready",
        "production_ready": False,
        "detail": "当前仅启用稳定合成发货契约；真实适配器接入前不得用于生产确认",
    }
    assert payload["total"] == 1
    assert payload["rows"][0] == {
        "delivery_line_id": delivery.delivery_line_id,
        "source_order_id": delivery.source_order_id,
        "source_line_id": delivery.source_line_id,
        "delivery_no": delivery.delivery_no,
        "delivery_date": "2026-08-08",
        "part_id": delivery.part_id,
        "pn": delivery.pn,
        "serial_number": delivery.serial_number,
        "delivered_quantity": "5.000",
        "confirmed_quantity": "0.000",
        "available_quantity": "5.000",
        "mapping_state": "ready",
        "mapping_version": delivery.mapping_version,
    }


def test_site_issue_write_requires_dedicated_action_and_purchase_cost_permission(db):
    project = _project(db, project_id="project-site-issue-v2-permission")
    delivery = _delivery_source(
        db,
        project=project,
        delivery_line_id="synthetic-delivery-line-permission",
    )
    path = f"/api/maintenance/site-issues/projects/{project.project_id}"
    body = {
        "idempotency_key": "synthetic-site-issue-permission-denied",
        "issue_date": "2026-08-09",
        "receiver": "合成接收人",
        "issued_by": "合成发出人",
        "site_location": "合成现场",
        "lines": [{"delivery_line_id": delivery.delivery_line_id, "quantity": "1"}],
        "reason": "验证专用权限失败关闭",
    }
    denied = _client(
        db,
        username="site_issue_v2_permission_denied",
        role="boss",
        permissions={
            "page_maintenance": True,
            "data_purchase_cost": True,
            "action_maintenance_site_issue_manage": False,
        },
    ).post(path, json=body)
    assert denied.status_code == 403, denied.text
    assert db.query(MaintenanceSiteIssue).filter_by(project_id=project.project_id).count() == 0

    without_cost = _client(
        db,
        username="site_issue_v2_permission_without_cost",
        role="boss",
        permissions={
            "page_maintenance": True,
            "data_purchase_cost": False,
            "action_maintenance_site_issue_manage": True,
        },
    ).post(
        path,
        json={
            **body,
            "idempotency_key": "synthetic-site-issue-permission-without-cost",
        },
    )
    assert without_cost.status_code == 403, without_cost.text
    assert db.query(MaintenanceSiteIssue).filter_by(project_id=project.project_id).count() == 0

    allowed = _client(
        db,
        username="site_issue_v2_permission_allowed",
        role="boss",
        permissions={
            "page_maintenance": True,
            "data_purchase_cost": True,
            "action_maintenance_site_issue_manage": True,
        },
    ).post(
        path,
        json={
            **body,
            "idempotency_key": "synthetic-site-issue-permission-allowed",
            "reason": "明确授权后建立草稿",
        },
    )
    assert allowed.status_code == 201, allowed.text
    assert allowed.json()["workflow_status"] == "draft"


def test_preview_and_confirm_freeze_cost_emit_one_return_event_and_never_touch_inventory(db):
    project = _project(db, project_id="project-site-issue-v2-confirm")
    part = DimPart(pn_std="PN-SYNTH-SITE-ISSUE-CONFIRM")
    db.add(part)
    db.flush()
    purchase_line = _purchase_evidence(db, part=part)
    inventory = Inventory(
        raw_inventory_id="synthetic-inventory-site-issue-v2",
        part_id=part.id,
        pn_std=part.pn_std,
        warehouse="合成公司库",
        source_qty="9",
        manual_qty="8",
        is_qty_overridden=True,
    )
    db.add(inventory)
    delivery = MaintenanceSiteIssueDeliverySource(
        delivery_line_id="synthetic-delivery-line-confirm",
        adapter_key="synthetic_delivery_v1",
        project_id=project.project_id,
        source_order_id="synthetic-wbdd-order-confirm",
        source_line_id="synthetic-wbdd-line-confirm",
        delivery_no="SYNTH-DELIVERY-CONFIRM",
        delivery_date=date(2026, 8, 8),
        part_id=part.id,
        pn=part.pn_std,
        delivered_quantity="5",
        linked_purchase_line_id=purchase_line.id,
        mapping_state="ready",
        mapping_version="synthetic-delivery-map-v1",
    )
    db.add(delivery)
    db.commit()
    client = _client(db, username="site_issue_v2_confirm_admin")
    created = client.post(
        f"/api/maintenance/site-issues/projects/{project.project_id}",
        json={
            "idempotency_key": "synthetic-site-issue-create-confirm",
            "issue_date": "2026-08-09",
            "receiver": "合成接收人",
            "issued_by": "合成发出人",
            "site_location": "合成现场 B",
            "lines": [{"delivery_line_id": delivery.delivery_line_id, "quantity": "2"}],
            "reason": "建立待确认草稿",
        },
    )
    assert created.status_code == 201, created.text
    issue = created.json()
    inventory_before = (
        inventory.source_qty,
        inventory.manual_qty,
        inventory.is_qty_overridden,
    )

    preview = client.post(
        f"/api/maintenance/site-issues/{issue['issue_id']}/preview",
        json={"project_id": project.project_id, "version": issue["version"]},
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()["can_confirm"] is True
    assert preview.json()["inventory_effect"] == "none"
    assert preview.json()["lines"][0]["available_quantity"] == "5.000"
    assert preview.json()["lines"][0]["cost_source"] == "direct_purchase"
    assert preview.json()["lines"][0]["cost_amount_ex_tax"] == "40.00"

    command = {
        "project_id": project.project_id,
        "version": issue["version"],
        "idempotency_key": "synthetic-site-issue-confirm-command",
        "reason": "确认现场实际领用",
    }
    confirmed = client.post(
        f"/api/maintenance/site-issues/{issue['issue_id']}/confirm",
        json=command,
    )
    assert confirmed.status_code == 200, confirmed.text
    payload = confirmed.json()
    assert payload["workflow_status"] == "confirmed"
    assert payload["idempotent_replay"] is False
    assert payload["lines"][0]["cost_source"] == "direct_purchase"
    assert payload["lines"][0]["cost_amount_ex_tax"] == "40.00"
    assert payload["return_obligation_event"]["event_type"] == "return_obligation_created"
    assert payload["inventory_effect"] == "none"

    replayed = client.post(
        f"/api/maintenance/site-issues/{issue['issue_id']}/confirm",
        json=command,
    )
    assert replayed.status_code == 200, replayed.text
    assert replayed.json()["issue_id"] == issue["issue_id"]
    assert replayed.json()["idempotent_replay"] is True

    db.expire_all()
    unchanged = db.get(Inventory, inventory.id)
    assert unchanged is not None
    assert (
        unchanged.source_qty,
        unchanged.manual_qty,
        unchanged.is_qty_overridden,
    ) == inventory_before
    assert (
        db.query(MaintenanceSiteIssueReturnEvent)
        .filter_by(issue_id=issue["issue_id"])
        .count()
        == 1
    )
    assert (
        db.query(MaintenanceProjectOperationAudit)
        .filter_by(
            entity_type="site_issue",
            entity_id=issue["issue_id"],
            action="confirm",
        )
        .count()
        == 1
    )


def test_confirmation_rejects_one_overdrawn_line_atomically_and_keeps_all_costs_empty(db):
    project = _project(db, project_id="project-site-issue-v2-atomic")
    first = _delivery_source(
        db,
        project=project,
        delivery_line_id="synthetic-delivery-line-atomic-1",
        quantity="1",
    )
    second = _delivery_source(
        db,
        project=project,
        delivery_line_id="synthetic-delivery-line-atomic-2",
        quantity="5",
    )
    client = _client(db, username="site_issue_v2_atomic_admin")
    created = client.post(
        f"/api/maintenance/site-issues/projects/{project.project_id}",
        json={
            "idempotency_key": "synthetic-site-issue-create-atomic",
            "issue_date": "2026-08-09",
            "receiver": "合成接收人",
            "issued_by": "合成发出人",
            "site_location": "合成现场 C",
            "lines": [
                {"delivery_line_id": first.delivery_line_id, "quantity": "2"},
                {"delivery_line_id": second.delivery_line_id, "quantity": "1"},
            ],
            "reason": "建立整单原子性草稿",
        },
    )
    assert created.status_code == 201, created.text
    issue = created.json()

    rejected = client.post(
        f"/api/maintenance/site-issues/{issue['issue_id']}/confirm",
        json={
            "project_id": project.project_id,
            "version": issue["version"],
            "idempotency_key": "synthetic-site-issue-confirm-atomic",
            "reason": "验证任一行冲突整单失败",
        },
    )

    assert rejected.status_code == 409, rejected.text
    db.expire_all()
    saved = db.get(MaintenanceSiteIssue, issue["issue_id"])
    saved_lines = (
        db.query(MaintenanceSiteIssueLine)
        .filter_by(issue_id=issue["issue_id"])
        .order_by(MaintenanceSiteIssueLine.line_no)
        .all()
    )
    assert saved is not None and saved.normalized_status == "draft"
    assert all(line.cost_source is None and line.cost_amount is None for line in saved_lines)
    assert (
        db.query(MaintenanceSiteIssueReturnEvent)
        .filter_by(issue_id=issue["issue_id"])
        .count()
        == 0
    )


def test_missing_price_confirms_quantity_but_keeps_amount_null_and_opens_cost_gap(db):
    project = _project(db, project_id="project-site-issue-v2-missing-price")
    delivery = _delivery_source(
        db,
        project=project,
        delivery_line_id="synthetic-delivery-line-missing-price",
        quantity="3",
    )
    client = _client(db, username="site_issue_v2_missing_price_admin")
    created = client.post(
        f"/api/maintenance/site-issues/projects/{project.project_id}",
        json={
            "idempotency_key": "synthetic-site-issue-create-missing-price",
            "issue_date": "2026-08-09",
            "receiver": "合成接收人",
            "issued_by": "合成发出人",
            "site_location": "合成现场 D",
            "lines": [{"delivery_line_id": delivery.delivery_line_id, "quantity": "1"}],
            "reason": "建立缺价草稿",
        },
    )
    issue = created.json()

    confirmed = client.post(
        f"/api/maintenance/site-issues/{issue['issue_id']}/confirm",
        json={
            "project_id": project.project_id,
            "version": issue["version"],
            "idempotency_key": "synthetic-site-issue-confirm-missing-price",
            "reason": "确认数量事实并保留缺价",
        },
    )

    assert confirmed.status_code == 200, confirmed.text
    line = confirmed.json()["lines"][0]
    assert line["quantity"] == "1.000"
    assert line["cost_source"] is None
    assert line["cost_amount_ex_tax"] is None
    gaps = client.get(f"/api/maintenance/projects/stable/{project.project_id}/cost-gaps")
    assert gaps.status_code == 200, gaps.text
    assert gaps.json()["total"] == 1


def test_production_environment_fails_closed_for_synthetic_delivery_confirmation(db, monkeypatch):
    project = _project(db, project_id="project-site-issue-v2-prod-gate")
    delivery = _delivery_source(
        db,
        project=project,
        delivery_line_id="synthetic-delivery-line-prod-gate",
    )
    client = _client(db, username="site_issue_v2_prod_gate_admin")
    created = client.post(
        f"/api/maintenance/site-issues/projects/{project.project_id}",
        json={
            "idempotency_key": "synthetic-site-issue-create-prod-gate",
            "issue_date": "2026-08-09",
            "receiver": "合成接收人",
            "issued_by": "合成发出人",
            "site_location": "合成现场 E",
            "lines": [{"delivery_line_id": delivery.delivery_line_id, "quantity": "1"}],
            "reason": "建立生产闸门草稿",
        },
    )
    issue = created.json()
    monkeypatch.setattr(
        maintenance_project_operations.operations,
        "_site_issue_is_production_blocked",
        lambda: True,
    )

    preview = client.post(
        f"/api/maintenance/site-issues/{issue['issue_id']}/preview",
        json={"project_id": project.project_id, "version": issue["version"]},
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()["can_confirm"] is False

    rejected = client.post(
        f"/api/maintenance/site-issues/{issue['issue_id']}/confirm",
        json={
            "project_id": project.project_id,
            "version": issue["version"],
            "idempotency_key": "synthetic-site-issue-confirm-prod-gate",
            "reason": "生产环境不得使用合成适配器",
        },
    )
    assert rejected.status_code == 400, rejected.text
    db.expire_all()
    assert db.get(MaintenanceSiteIssue, issue["issue_id"]).normalized_status == "draft"


def test_draft_patch_replaces_lines_with_server_ids_and_search_is_post_only(db):
    project = _project(db, project_id="project-site-issue-v2-edit")
    first = _delivery_source(
        db,
        project=project,
        delivery_line_id="synthetic-delivery-line-edit-1",
    )
    second = _delivery_source(
        db,
        project=project,
        delivery_line_id="synthetic-delivery-line-edit-2",
    )
    client = _client(db, username="site_issue_v2_edit_admin")
    issue = _create_draft(
        client,
        project_id=project.project_id,
        delivery_line_id=first.delivery_line_id,
        quantity="1",
        key="synthetic-site-issue-create-edit",
    )
    old_line_id = issue["lines"][0]["issue_line_id"]

    edited = client.patch(
        f"/api/maintenance/site-issues/{issue['issue_id']}",
        json={
            "project_id": project.project_id,
            "version": issue["version"],
            "idempotency_key": "synthetic-site-issue-patch-draft",
            "receiver": "更新后的合成接收人",
            "lines": [{"delivery_line_id": second.delivery_line_id, "quantity": "2"}],
            "reason": "修改草稿",
        },
    )
    assert edited.status_code == 200, edited.text
    payload = edited.json()
    assert payload["workflow_status"] == "draft"
    assert payload["receiver"] == "更新后的合成接收人"
    assert payload["lines"][0]["issue_line_id"] != old_line_id
    assert payload["lines"][0]["delivery_line_id"] == second.delivery_line_id
    assert payload["lines"][0]["cost_source"] is None

    no_business_change = client.patch(
        f"/api/maintenance/site-issues/{issue['issue_id']}",
        json={
            "project_id": project.project_id,
            "version": payload["version"],
            "idempotency_key": "synthetic-site-issue-patch-noop",
            "reason": "只有原因不能制造更正事实",
        },
    )
    assert no_business_change.status_code == 400, no_business_change.text
    db.expire_all()
    assert db.get(MaintenanceSiteIssue, issue["issue_id"]).version == payload["version"]
    assert (
        db.query(MaintenanceSiteIssueCommand)
        .filter_by(idempotency_key="synthetic-site-issue-patch-noop")
        .count()
        == 0
    )

    replayed = client.patch(
        f"/api/maintenance/site-issues/{issue['issue_id']}",
        json={
            "project_id": project.project_id,
            "version": issue["version"],
            "idempotency_key": "synthetic-site-issue-patch-draft",
            "receiver": "更新后的合成接收人",
            "lines": [{"delivery_line_id": second.delivery_line_id, "quantity": "2"}],
            "reason": "修改草稿",
        },
    )
    assert replayed.status_code == 200, replayed.text
    assert replayed.json()["idempotent_replay"] is True

    searched = client.post(
        "/api/maintenance/site-issues/search",
        json={
            "project_id": project.project_id,
            "q": payload["issue_no"],
            "workflow_statuses": ["draft"],
            "page": 1,
            "page_size": 20,
        },
    )
    assert searched.status_code == 200, searched.text
    assert searched.json()["total"] == 1
    assert searched.json()["rows"][0]["issue_id"] == issue["issue_id"]
    assert client.get("/api/maintenance/site-issues/search").status_code == 405

    client_supplied_identity = client.post(
        f"/api/maintenance/site-issues/projects/{project.project_id}",
        json={
            "issue_no": "CLIENT-MUST-NOT-CONTROL",
            "idempotency_key": "synthetic-client-identity-rejected",
            "issue_date": "2026-08-09",
            "receiver": "合成接收人",
            "issued_by": "合成发出人",
            "site_location": "合成现场",
            "lines": [{"delivery_line_id": first.delivery_line_id, "quantity": "1"}],
            "reason": "客户端不得指定编号",
        },
    )
    assert client_supplied_identity.status_code == 422


def test_metadata_correction_keeps_line_identity_and_frozen_cost_evidence(db):
    project = _project(db, project_id="project-site-v2-frozen-metadata")
    part = DimPart(pn_std="PN-SYNTH-FROZEN-METADATA")
    db.add(part)
    db.flush()
    purchase = _purchase_evidence(db, part=part)
    delivery = MaintenanceSiteIssueDeliverySource(
        delivery_line_id="synthetic-delivery-line-frozen-metadata",
        adapter_key="synthetic_delivery_v1",
        project_id=project.project_id,
        source_order_id="synthetic-order-frozen-metadata",
        source_line_id="synthetic-line-frozen-metadata",
        delivery_no="SYNTH-FROZEN-METADATA",
        delivery_date=date(2026, 8, 8),
        part_id=part.id,
        pn=part.pn_std,
        delivered_quantity="5",
        linked_purchase_line_id=purchase.id,
        mapping_state="ready",
        mapping_version="synthetic-delivery-map-v1",
    )
    db.add(delivery)
    db.commit()
    client = _client(db, username="site_issue_v2_frozen_metadata_admin")
    draft = _create_draft(
        client,
        project_id=project.project_id,
        delivery_line_id=delivery.delivery_line_id,
        quantity="1",
        key="synthetic-site-issue-create-frozen-metadata",
    )
    confirmed_response = client.post(
        f"/api/maintenance/site-issues/{draft['issue_id']}/confirm",
        json={
            "project_id": project.project_id,
            "version": draft["version"],
            "idempotency_key": "synthetic-site-issue-confirm-frozen-metadata",
            "reason": "冻结初次确认成本证据",
        },
    )
    assert confirmed_response.status_code == 200, confirmed_response.text
    confirmed = confirmed_response.json()
    frozen_line = confirmed["lines"][0]
    assert frozen_line["unit_cost_ex_tax"] == "20.00"

    db.expire_all()
    purchase_row = db.get(FPurchaseLine, purchase.id)
    assert purchase_row is not None
    purchase_row.unit_price = "99"
    db.commit()

    corrected_response = client.patch(
        f"/api/maintenance/site-issues/{draft['issue_id']}",
        json={
            "project_id": project.project_id,
            "version": confirmed["version"],
            "idempotency_key": "synthetic-site-issue-correct-frozen-metadata",
            "receiver": "更正后的合成接收人",
            "lines": [
                {
                    "delivery_line_id": delivery.delivery_line_id,
                    "quantity": "1",
                }
            ],
            "reason": "仅更正接收人，不改变成本输入",
        },
    )
    assert corrected_response.status_code == 200, corrected_response.text
    corrected = corrected_response.json()
    corrected_line = corrected["lines"][0]
    assert corrected_line["issue_line_id"] == frozen_line["issue_line_id"]
    assert corrected_line["unit_cost_ex_tax"] == "20.00"
    assert corrected_line["reference_samples"] == frozen_line["reference_samples"]

    no_change_key = "synthetic-site-issue-correct-no-business-change"
    no_change = client.patch(
        f"/api/maintenance/site-issues/{draft['issue_id']}",
        json={
            "project_id": project.project_id,
            "version": corrected["version"],
            "idempotency_key": no_change_key,
            "receiver": "更正后的合成接收人",
            "reason": "相同内容不能制造新更正事实",
        },
    )
    assert no_change.status_code == 400, no_change.text
    db.expire_all()
    assert db.get(MaintenanceSiteIssue, draft["issue_id"]).version == corrected["version"]
    assert (
        db.query(MaintenanceSiteIssueCommand)
        .filter_by(idempotency_key=no_change_key)
        .count()
        == 0
    )


def test_confirmed_issue_can_be_corrected_then_fully_voided_without_registration(db):
    project = _project(db, project_id="project-site-issue-v2-correct-void")
    part = DimPart(pn_std="PN-SYNTH-CORRECT-VOID")
    db.add(part)
    db.flush()
    purchase = _purchase_evidence(db, part=part)
    inventory = Inventory(
        raw_inventory_id="synthetic-inventory-correct-void",
        part_id=part.id,
        pn_std=part.pn_std,
        warehouse="合成公司库",
        source_qty="10",
        manual_qty="7",
        is_qty_overridden=True,
    )
    db.add(inventory)
    delivery = MaintenanceSiteIssueDeliverySource(
        delivery_line_id="synthetic-delivery-line-correct-void",
        adapter_key="synthetic_delivery_v1",
        project_id=project.project_id,
        source_order_id="synthetic-order-correct-void",
        source_line_id="synthetic-line-correct-void",
        delivery_no="SYNTH-CORRECT-VOID",
        delivery_date=date(2026, 8, 8),
        part_id=part.id,
        pn=part.pn_std,
        delivered_quantity="5",
        linked_purchase_line_id=purchase.id,
        mapping_state="ready",
        mapping_version="synthetic-delivery-map-v1",
    )
    db.add(delivery)
    db.commit()
    before_inventory = (inventory.source_qty, inventory.manual_qty, inventory.is_qty_overridden)
    client = _client(db, username="site_issue_v2_correct_void_admin")
    draft = _create_draft(
        client,
        project_id=project.project_id,
        delivery_line_id=delivery.delivery_line_id,
        quantity="2",
        key="synthetic-site-issue-create-correct-void",
    )
    confirmed = client.post(
        f"/api/maintenance/site-issues/{draft['issue_id']}/confirm",
        json={
            "project_id": project.project_id,
            "version": draft["version"],
            "idempotency_key": "synthetic-site-issue-confirm-correct-void",
            "reason": "确认后再更正",
        },
    ).json()

    corrected_response = client.patch(
        f"/api/maintenance/site-issues/{draft['issue_id']}",
        json={
            "project_id": project.project_id,
            "version": confirmed["version"],
            "idempotency_key": "synthetic-site-issue-correct-command",
            "lines": [{"delivery_line_id": delivery.delivery_line_id, "quantity": "3"}],
            "reason": "现场核对后更正数量",
        },
    )
    assert corrected_response.status_code == 200, corrected_response.text
    corrected = corrected_response.json()
    assert corrected["workflow_status"] == "corrected"
    assert corrected["lines"][0]["quantity"] == "3.000"
    assert corrected["lines"][0]["cost_amount_ex_tax"] == "60.00"
    assert corrected["return_obligation_event"]["event_type"] == "return_obligation_corrected"

    voided_response = client.post(
        f"/api/maintenance/site-issues/{draft['issue_id']}/void",
        json={
            "project_id": project.project_id,
            "version": corrected["version"],
            "idempotency_key": "synthetic-site-issue-void-command",
            "reason": "作废错误现场领用",
        },
    )
    assert voided_response.status_code == 200, voided_response.text
    assert voided_response.json()["workflow_status"] == "void"
    # D-01：v2 整单作废同样把明细软作废，06 表导出随之排除（此前单头 void 行仍会被导出）
    db.expire_all()
    voided_lines = db.query(MaintenanceSiteIssueLine).filter_by(issue_id=draft["issue_id"]).all()
    assert voided_lines and all(item.is_active is False for item in voided_lines)
    assert not {item.issue_line_id for item in voided_lines} & _site_sheet_entity_ids(
        db, project.project_id
    )

    candidates = client.post(
        f"/api/maintenance/site-issues/projects/{project.project_id}/candidates/search",
        json={"page": 1, "page_size": 20},
    ).json()
    assert candidates["rows"][0]["confirmed_quantity"] == "0.000"
    assert candidates["rows"][0]["available_quantity"] == "5.000"
    obligation = db.query(MaintenanceReturnObligation).filter_by(
        issue_id=draft["issue_id"]
    ).one()
    assert obligation.obligation_id
    assert obligation.required_quantity == 0  # no standard category: pending
    assert obligation.source_quantity == 3
    assert obligation.is_active is False
    assert obligation.source_issue_version == voided_response.json()["version"]
    db.expire_all()
    unchanged = db.get(Inventory, inventory.id)
    assert (unchanged.source_qty, unchanged.manual_qty, unchanged.is_qty_overridden) == before_inventory


def test_projected_return_obligation_does_not_block_safe_full_void(db):
    project = _project(db, project_id="project-site-issue-v2-downstream")
    delivery = _delivery_source(
        db,
        project=project,
        delivery_line_id="synthetic-delivery-line-downstream",
    )
    client = _client(db, username="site_issue_v2_downstream_admin")
    draft = _create_draft(
        client,
        project_id=project.project_id,
        delivery_line_id=delivery.delivery_line_id,
        quantity="1",
        key="synthetic-site-issue-create-downstream",
    )
    confirmed = client.post(
        f"/api/maintenance/site-issues/{draft['issue_id']}/confirm",
        json={
            "project_id": project.project_id,
            "version": draft["version"],
            "idempotency_key": "synthetic-site-issue-confirm-downstream",
            "reason": "确认并模拟下游消费",
        },
    ).json()
    event = (
        db.query(MaintenanceSiteIssueReturnEvent)
        .filter_by(issue_id=draft["issue_id"], event_type="return_obligation_created")
        .one()
    )
    assert event.downstream_reference.startswith("maintenance-return-obligations:")
    assert event.consumed_at is not None

    voided = client.post(
        f"/api/maintenance/site-issues/{draft['issue_id']}/void",
        json={
            "project_id": project.project_id,
            "version": confirmed["version"],
            "idempotency_key": "synthetic-site-issue-void-downstream",
            "reason": "仅生成义务但尚未登记返还，可以整单撤回",
        },
    )
    assert voided.status_code == 200, voided.text
    assert voided.json()["workflow_status"] == "void"
    db.expire_all()
    obligation = db.query(MaintenanceReturnObligation).filter_by(
        issue_id=draft["issue_id"]
    ).one()
    assert obligation.is_active is False


def test_concurrent_confirmations_never_overdraw_one_delivery_balance(db):
    project = _project(db, project_id="project-site-issue-v2-concurrency")
    delivery = _delivery_source(
        db,
        project=project,
        delivery_line_id="synthetic-delivery-line-concurrency",
        quantity="5",
    )
    client = _client(db, username="site_issue_v2_concurrency_admin")
    first = _create_draft(
        client,
        project_id=project.project_id,
        delivery_line_id=delivery.delivery_line_id,
        quantity="4",
        key="synthetic-site-issue-create-concurrency-1",
    )
    second = _create_draft(
        client,
        project_id=project.project_id,
        delivery_line_id=delivery.delivery_line_id,
        quantity="4",
        key="synthetic-site-issue-create-concurrency-2",
    )

    def confirm(issue: dict, suffix: str):
        return client.post(
            f"/api/maintenance/site-issues/{issue['issue_id']}/confirm",
            json={
                "project_id": project.project_id,
                "version": issue["version"],
                "idempotency_key": f"synthetic-site-issue-confirm-concurrency-{suffix}",
                "reason": "并发确认余额测试",
            },
        )

    with ThreadPoolExecutor(max_workers=2) as executor:
        responses = list(
            executor.map(lambda pair: confirm(*pair), [(first, "1"), (second, "2")])
        )
    assert sorted(response.status_code for response in responses) == [200, 409]

    candidates = client.post(
        f"/api/maintenance/site-issues/projects/{project.project_id}/candidates/search",
        json={"page": 1, "page_size": 20},
    ).json()
    assert candidates["rows"][0]["confirmed_quantity"] == "4.000"
    assert candidates["rows"][0]["available_quantity"] == "1.000"


def test_search_rejects_oversized_body_without_reflecting_or_auditing_it(db):
    project = _project(db, project_id="project-site-issue-v2-search-limit")
    client = _client(db, username="site_issue_v2_search_limit_admin")
    before = (
        db.query(SysAccessLog)
        .filter_by(action="site_issue_search")
        .count()
    )
    sentinel = "SENSITIVE-SITE-ISSUE-SEARCH-" + ("x" * 300)

    rejected = client.post(
        "/api/maintenance/site-issues/search",
        json={
            "project_id": project.project_id,
            "q": sentinel,
            "workflow_statuses": ["draft"],
            "page": 1,
            "page_size": 20,
        },
    )

    assert rejected.status_code == 422, rejected.text
    assert sentinel not in rejected.text
    assert (
        db.query(SysAccessLog)
        .filter_by(action="site_issue_search")
        .count()
        == before
    )


def test_all_direct_issue_actions_fail_closed_for_a_different_project_id(db):
    owner = _project(db, project_id="project-site-issue-v2-owner")
    other = _project(db, project_id="project-site-issue-v2-other")
    delivery = _delivery_source(
        db,
        project=owner,
        delivery_line_id="synthetic-delivery-line-project-scope",
    )
    client = _client(db, username="site_issue_v2_scope_admin")
    issue = _create_draft(
        client,
        project_id=owner.project_id,
        delivery_line_id=delivery.delivery_line_id,
        quantity="1",
        key="synthetic-site-issue-create-project-scope",
    )
    common = {"project_id": other.project_id, "version": issue["version"]}

    preview = client.post(
        f"/api/maintenance/site-issues/{issue['issue_id']}/preview",
        json=common,
    )
    confirm = client.post(
        f"/api/maintenance/site-issues/{issue['issue_id']}/confirm",
        json={
            **common,
            "idempotency_key": "synthetic-site-issue-confirm-cross-project",
            "reason": "跨项目确认必须失败",
        },
    )
    patch = client.patch(
        f"/api/maintenance/site-issues/{issue['issue_id']}",
        json={
            **common,
            "idempotency_key": "synthetic-site-issue-patch-cross-project",
            "receiver": "不得修改",
            "reason": "跨项目编辑必须失败",
        },
    )
    void = client.post(
        f"/api/maintenance/site-issues/{issue['issue_id']}/void",
        json={
            **common,
            "idempotency_key": "synthetic-site-issue-void-cross-project",
            "reason": "跨项目作废必须失败",
        },
    )

    assert [preview.status_code, confirm.status_code, patch.status_code, void.status_code] == [
        403,
        403,
        403,
        403,
    ]
    db.expire_all()
    assert db.get(MaintenanceSiteIssue, issue["issue_id"]).normalized_status == "draft"


def test_manager_row_scope_covers_site_issue_search_and_entity_routes(db):
    own = _project(db, project_id="project-site-issue-v2-manager-own")
    other = _project(db, project_id="project-site-issue-v2-manager-other")
    delivery = _delivery_source(
        db,
        project=other,
        delivery_line_id="synthetic-delivery-line-manager-scope",
    )
    admin = _client(db, username="site_issue_v2_manager_scope_admin")
    issue = _create_draft(
        admin,
        project_id=other.project_id,
        delivery_line_id=delivery.delivery_line_id,
        quantity="1",
        key="synthetic-site-issue-manager-scope-draft",
    )
    manager = SysUser(
        username="site_issue_v2_scoped_manager",
        role="purchaser",
        display_name="合成现场领用项目经理",
        password_hash=hash_password("synthetic-password-123"),
        permissions={
            "page_maintenance": True,
            "data_purchase_cost": True,
            "action_maintenance_site_issue_manage": True,
        },
    )
    db.add(manager)
    db.flush()
    db.add(
        MaintenanceProjectUserAssignment(
            assignment_id="site-issue-scope-assignment",
            project_id=own.project_id,
            responsibility_type="primary_manager",
            user_id=manager.id,
            assigned_by="site_issue_v2_manager_scope_admin",
            assignment_reason="限定为本人负责项目",
        )
    )
    db.commit()
    client = _client_for_existing_user(db, username=manager.username)

    candidate_search = client.post(
        f"/api/maintenance/site-issues/projects/{other.project_id}/candidates/search",
        json={"page": 1, "page_size": 20},
    )
    issue_search = client.post(
        "/api/maintenance/site-issues/search",
        json={"project_id": other.project_id, "page": 1, "page_size": 20},
    )
    create = client.post(
        f"/api/maintenance/site-issues/projects/{other.project_id}",
        json={
            "idempotency_key": "synthetic-site-issue-manager-scope-create",
            "issue_date": "2026-08-09",
            "receiver": "不得创建",
            "issued_by": "不得创建",
            "site_location": "他人项目",
            "lines": [{"delivery_line_id": delivery.delivery_line_id, "quantity": "1"}],
            "reason": "越权创建必须失败",
        },
    )
    common = {"project_id": other.project_id, "version": issue["version"]}
    preview = client.post(
        f"/api/maintenance/site-issues/{issue['issue_id']}/preview",
        json=common,
    )
    confirm = client.post(
        f"/api/maintenance/site-issues/{issue['issue_id']}/confirm",
        json={
            **common,
            "idempotency_key": "synthetic-site-issue-manager-scope-confirm",
            "reason": "越权确认必须失败",
        },
    )
    patch = client.patch(
        f"/api/maintenance/site-issues/{issue['issue_id']}",
        json={
            **common,
            "idempotency_key": "synthetic-site-issue-manager-scope-patch",
            "receiver": "不得修改",
            "reason": "越权编辑必须失败",
        },
    )
    void = client.post(
        f"/api/maintenance/site-issues/{issue['issue_id']}/void",
        json={
            **common,
            "idempotency_key": "synthetic-site-issue-manager-scope-void",
            "reason": "越权作废必须失败",
        },
    )

    assert [
        candidate_search.status_code,
        issue_search.status_code,
        create.status_code,
        preview.status_code,
        confirm.status_code,
        patch.status_code,
        void.status_code,
    ] == [403] * 7
    db.expire_all()
    assert db.get(MaintenanceSiteIssue, issue["issue_id"]).normalized_status == "draft"


def test_command_receipts_are_append_only_and_return_events_only_allow_one_downstream_registration(db):
    project = _project(db, project_id="project-site-issue-v2-immutable")
    delivery = _delivery_source(
        db,
        project=project,
        delivery_line_id="synthetic-delivery-line-immutable",
    )
    client = _client(db, username="site_issue_v2_immutable_admin")
    draft = _create_draft(
        client,
        project_id=project.project_id,
        delivery_line_id=delivery.delivery_line_id,
        quantity="1",
        key="synthetic-site-issue-create-immutable",
    )
    confirmed = client.post(
        f"/api/maintenance/site-issues/{draft['issue_id']}/confirm",
        json={
            "project_id": project.project_id,
            "version": draft["version"],
            "idempotency_key": "synthetic-site-issue-confirm-immutable",
            "reason": "生成不可篡改事实",
        },
    )
    assert confirmed.status_code == 200, confirmed.text
    event = (
        db.query(MaintenanceSiteIssueReturnEvent)
        .filter_by(issue_id=draft["issue_id"])
        .one()
    )
    assert event.downstream_reference.startswith("maintenance-return-obligations:")
    assert event.consumed_at is not None

    with pytest.raises(DBAPIError, match="append-only"):
        with db.begin_nested():
            db.execute(
                text(
                    "UPDATE maintenance_site_issue_return_event "
                    "SET payload = CAST(:payload AS jsonb) WHERE event_id = :event_id"
                ),
                {"payload": '{"tampered": true}', "event_id": event.event_id},
            )
    with pytest.raises(DBAPIError, match="append-only"):
        with db.begin_nested():
            db.execute(
                text(
                    "UPDATE maintenance_site_issue_command "
                    "SET response_json = CAST(:payload AS jsonb) "
                    "WHERE idempotency_key = :key"
                ),
                {
                    "payload": '{"tampered": true}',
                    "key": "synthetic-site-issue-confirm-immutable",
                },
            )
    assert (
        db.query(MaintenanceSiteIssueCommand)
        .filter_by(idempotency_key="synthetic-site-issue-confirm-immutable")
        .count()
        == 1
    )


def test_workbook_issue_void_is_not_blocked_by_unrelated_pending_return_event(db):
    """Codex P2（#319）：作废从未发过返还事件的单，不应替它清空整个项目的待投影事件。

    同项目另一张单留下一条畸形的待投影事件（消费方会抛契约错误）时，工作簿来源的单
    仍能作废；那条无关事件原样保留（不被本单的作废顺手消费或撞成 409）。
    """
    project = _project(db, project_id="project-si-wb-void-unrelated")
    issue, line = _workbook_issue(db, project=project, issue_no="CKD-WB-0101")
    other, _ = _workbook_issue(db, project=project, issue_no="CKD-WB-0102")
    stray = MaintenanceSiteIssueReturnEvent(
        event_id=str(uuid4()),
        project_id=project.project_id,
        issue_id=other.issue_id,
        event_type="return_obligation_created",
        issue_version=1,
        payload={"schema_version": 1, "project_id": project.project_id, "issue_id": other.issue_id},
    )
    db.add(stray)
    db.commit()

    client = _client(
        db,
        username="site_issue_workbook_void_unrelated_manager",
        role="boss",
        permissions={
            "page_maintenance": True,
            "data_purchase_cost": True,
            "action_maintenance_site_issue_manage": True,
        },
    )
    voided = client.post(
        f"/api/maintenance/site-issues/{issue.issue_id}/void",
        json={
            "project_id": project.project_id,
            "version": issue.version,
            "idempotency_key": "synthetic-workbook-void-unrelated-command",
            "reason": "同项目另一张单的畸形待投影事件不该挡住本单作废",
        },
    )
    assert voided.status_code == 200, voided.text
    assert voided.json()["return_obligation_event"] is None

    db.expire_all()
    assert db.get(MaintenanceSiteIssue, issue.issue_id).normalized_status == "void"
    assert db.get(MaintenanceSiteIssueLine, line.issue_line_id).is_active is False
    kept = db.get(MaintenanceSiteIssueReturnEvent, stray.event_id)
    assert kept.downstream_reference is None and kept.consumed_at is None


def test_workbook_sourced_issue_can_be_voided_through_the_panel_route(db):
    """生产缺陷：06 表建的领用单（source=workbook）在面板点「作废」一律 400。

    作废后按 D-01 退出全部读模型：单头 void、明细软作废、06 表导出不含该行、
    工作台与展示板的已领用成本归零；审计与幂等回执照写，重放 / 再作废与 v2 同口径。
    """
    project = _project(db, project_id="project-site-issue-workbook-void")
    issue, line = _workbook_issue(db, project=project, issue_no="CKD-WB-0001")
    assert line.issue_line_id in _site_sheet_entity_ids(db, project.project_id)
    assert _workspace_requisition_cost(db, project.project_id) == "113.00"
    assert _board_requisition_cost(db, project.project_id) == Decimal("113.00")

    client = _client(
        db,
        username="site_issue_workbook_void_manager",
        role="boss",
        permissions={
            "page_maintenance": True,
            "data_purchase_cost": True,
            "action_maintenance_site_issue_manage": True,
        },
    )
    path = f"/api/maintenance/site-issues/{issue.issue_id}/void"
    body = {
        "project_id": project.project_id,
        "version": issue.version,
        "idempotency_key": "synthetic-workbook-void-command",
        "reason": "06 表录错的领用单整单作废",
    }
    voided = client.post(path, json=body)
    assert voided.status_code == 200, voided.text
    payload = voided.json()
    assert payload["workflow_status"] == "void"
    assert payload["source"] == "workbook"
    assert payload["return_obligation_event"] is None
    assert payload["inventory_effect"] == "none"
    assert payload["idempotent_replay"] is False

    db.expire_all()
    stored = db.get(MaintenanceSiteIssue, issue.issue_id)
    assert stored.normalized_status == "void"
    assert stored.voided_at is not None
    assert stored.version == 2
    assert db.get(MaintenanceSiteIssueLine, line.issue_line_id).is_active is False
    audit = (
        db.query(MaintenanceProjectOperationAudit)
        .filter_by(entity_type="site_issue", entity_id=issue.issue_id, action="void")
        .one()
    )
    assert audit.reason == body["reason"]
    assert audit.before_json["workflow_status"] == "confirmed"
    assert audit.after_json["workflow_status"] == "void"
    # 从未进入返还义务接口的单：不发事件、不投影义务，也不能因此报错
    assert (
        db.query(MaintenanceSiteIssueReturnEvent).filter_by(issue_id=issue.issue_id).count()
        == 0
    )
    assert (
        db.query(MaintenanceReturnObligation).filter_by(issue_id=issue.issue_id).count()
        == 0
    )

    # D-01：作废行不再进入任何统计与导出
    assert line.issue_line_id not in _site_sheet_entity_ids(db, project.project_id)
    assert _workspace_requisition_cost(db, project.project_id) == "0.00"
    assert _board_requisition_cost(db, project.project_id) == Decimal("0")
    listed = client.post(
        "/api/maintenance/site-issues/search",
        json={"project_id": project.project_id},
    )
    assert listed.status_code == 200, listed.text
    listed_issue = next(
        row for row in listed.json()["rows"] if row["issue_id"] == issue.issue_id
    )
    assert listed_issue["workflow_status"] == "void"
    # 面板按明细铺行，明细软作废后作废单不再出现在「领用与返还」列表
    assert listed_issue["lines"] == []

    # 幂等：同键重放返回同一回执；换键再作废按既有口径 409，不写第二张回执
    replayed = client.post(path, json=body)
    assert replayed.status_code == 200, replayed.text
    assert replayed.json()["idempotent_replay"] is True
    again = client.post(
        path,
        json={
            **body,
            "version": 2,
            "idempotency_key": "synthetic-workbook-void-command-again",
        },
    )
    assert again.status_code == 409, again.text
    assert "已经作废" in again.json()["detail"]
    assert (
        db.query(MaintenanceSiteIssueCommand).filter_by(issue_id=issue.issue_id).count()
        == 1
    )
