"""T19 四项评审的页面、工作簿及权限回归。"""
from datetime import date
from io import BytesIO

from openpyxl import load_workbook
import pytest
from sqlalchemy import select

from app.models.fields_audience import FieldValue
from app.services import field_pool
from app.services.presales import common, inquiries, inquiry_workbook, quote_workbook, offers, visibility
from app.services.presales.workbook_registration import register_workbooks
from app.services.workbook import operations
from tests.test_presales_api_visibility import _client, users
from tests.test_presales_inquiry_workbook import _edit, _set, _preview, _apply, _ctx, presales_users, presales_parts
from tests.test_presales_quote_workbook import quote_data, _submitted, _published, _edit as _edit_quote, _apply as _apply_quote


def test_page_delivery_exports_and_excel_delivery_reads_back(db, users):
    register_workbooks()
    client = _client(users.seller)
    created = client.post('/api/presales/inquiries', json={"lines": [{"pn_raw": "T19", "expected_delivery": "2026-12-01"}]})
    assert created.status_code == 201, created.text
    row = created.json()
    inquiry = common.get_inquiry(db, row['id'])
    content = inquiry_workbook.export_inquiry(db, inquiry.id, _ctx(db, users.seller))
    sheet = load_workbook(BytesIO(content))['询价明细']
    cols = {cell.value: cell.column for cell in sheet[1]}
    assert sheet.cell(2, cols['期望交期']).value.date() == date(2026, 12, 1)
    changed = _edit(content, lambda sheet, cols: _set(sheet, cols, 2, **{'期望交期': '2026-12-02'}))
    _apply(db, users.seller, _preview(db, users.seller, changed, inquiry))
    db.commit()
    assert client.get(f'/api/presales/inquiries/{inquiry.id}').json()['lines'][0]['expected_delivery'] == '2026-12-02'


def test_legacy_delivery_reads_without_write_and_migrates_on_save(db, presales_users, presales_parts):
    user = presales_users.sales
    ctx = _ctx(db, user)
    inquiry = inquiries.create_draft(db, user.id, lines=[{'part_id': presales_parts[0].id}])
    line = inquiries.lines_for(db, inquiry.id)[0]
    subject = field_pool.Subject(inquiries.SUBJECT_TYPE, f'line:{line.id}', 'line')
    field_pool.save_values(db, subject, [{'new_name': '期望交期', 'value': '2026-12-03'}], ctx)
    definition = field_pool.list_defs(db, inquiries.SUBJECT_TYPE, 'line', ctx)[0]
    view = visibility.inquiry_view(db, inquiry, ctx)['lines'][0]
    assert view['expected_delivery'] == date(2026, 12, 3)
    assert view['field_values'] == [] and view['field_template'] == []
    assert line.expected_delivery is None
    from app.api.field_pool import template_for
    assert template_for(subject_type=inquiries.SUBJECT_TYPE, level='line', anchor=None, db=db, viewer=ctx) == []
    inquiries.update_draft(db, inquiry.id, user.id, inquiry.version, lines=[{'line_key': line.line_key, 'qty': 2}])
    assert line.expected_delivery == date(2026, 12, 3)
    assert list(db.scalars(select(FieldValue))) == []
    assert field_pool.list_defs(db, inquiries.SUBJECT_TYPE, 'line', ctx)[0].id == definition.id


def test_custom_export_schema_selection_visibility_and_roundtrip(db, presales_users, presales_parts):
    register_workbooks()
    user = presales_users.sales
    ctx = _ctx(db, user)
    inquiry = inquiries.create_draft(db, user.id, lines=[{'part_id': presales_parts[0].id,
        'field_values': [{'new_name': '质保', 'value': '三年'}, {'new_name': '包装', 'value': '原厂'}]}])
    schema = inquiry_workbook.export_schema(db, ctx, inquiry.id)
    custom = [column for column in schema['sheets'][0]['columns'] if column.get('custom')]
    assert {column['key'] for column in custom} == {'+质保', '+包装'}
    selected = {'lines': [column['key'] for column in schema['sheets'][0]['columns']]}
    content = inquiry_workbook.export_inquiry(db, inquiry.id, ctx, selected)
    sheet = load_workbook(BytesIO(content))['询价明细']
    cols = {cell.value: cell.column for cell in sheet[1]}
    assert sheet.cell(2, cols['+质保']).value == '三年'
    changed = _edit(content, lambda sheet, cols: _set(sheet, cols, 2, **{'数量': 2}))
    _apply(db, user, _preview(db, user, changed, inquiry))
    line = visibility.inquiry_view(db, inquiry, ctx)['lines'][0]
    assert {field['name']: field['value'] for field in line['field_values']} == {'质保': '三年', '包装': '原厂'}
    selected['lines'].remove('+包装')
    excluded = inquiry_workbook.export_inquiry(db, inquiry.id, ctx, selected)
    assert '+包装' not in [cell.value for cell in load_workbook(BytesIO(excluded))['询价明细'][1]]
    legacy = inquiry_workbook.export_inquiry(db, inquiry.id, ctx)
    assert {'+包装', '+质保'}.issubset({cell.value for cell in load_workbook(BytesIO(legacy))['询价明细'][1]})
    # 用字段作者之外的账号写入私密值，单据所有人也不可看到字段名。
    line_model = inquiries.lines_for(db, inquiry.id)[0]
    other_ctx = _ctx(db, presales_users.boss)
    hidden = field_pool.define_field(db, subject_type=inquiries.SUBJECT_TYPE, level='line',
        name='隐藏底价', owner_scope='team', viewer=other_ctx)
    field_pool.save_values(db, field_pool.Subject(inquiries.SUBJECT_TYPE, f'line:{line_model.id}', 'line'),
        [{'field_def_id': hidden.id, 'value': '99', 'rule': {'mode': 'private'}}], other_ctx)
    assert '+隐藏底价' not in {column['key'] for column in inquiry_workbook.export_schema(db, ctx, inquiry.id)['sheets'][0]['columns']}
    assert '隐藏底价' not in str(inquiry_workbook.export_schema(db, ctx, inquiry.id))


def test_shared_scope_excludes_own_before_pagination(db, users):
    shared = []
    for index in range(6):
        own = inquiries.create_draft(db, users.seller.id, lines=[{'pn_raw': f'OWN-{index}'}])
        inquiries.submit(db, own.id, users.seller.id, own.version)
        row = inquiries.create_draft(db, users.other.id, lines=[{'pn_raw': f'SHARED-{index}'}])
        inquiries.submit(db, row.id, users.other.id, row.version)
        shared.append(row.id)
    db.commit()
    client = _client(users.seller)
    seen = []
    for page in range(1, 4):
        response = client.get('/api/presales/inquiries', params={'scope': 'shared', 'page_size': 2, 'page': page})
        assert response.status_code == 200, response.text
        data = response.json()
        assert data['total'] == 6 and len(data['items']) == 2
        assert all(not row['is_owner'] for row in data['items'])
        seen.extend(row['id'] for row in data['items'])
    assert len(set(seen)) == 6 and set(seen) == set(shared)


@pytest.mark.parametrize('mode', ['skip', 'reuse'])
def test_excel_draft_response_hidden_until_reply(db, quote_data, mode):
    data = quote_data
    original = _published(db, data) if mode == 'reuse' else None
    inquiry = inquiries.create_draft(db, data.sales.id, assignee_ids=[data.buyer.id, data.peer.id],
        lines=[{'part_id': data.parts[0].id}])
    inquiries.submit(db, inquiry.id, data.sales.id, inquiry.version)
    data.other.template_perms = {**data.other.template_perms, 'action_presales_supervise': True}
    db.flush()
    line = inquiries.lines_for(db, inquiry.id)[0]
    changes = {'处理方式': '不报', '不报原因': '草稿供应商无货'} if mode == 'skip' else {'处理方式': '沿用', '沿用报价编号': original.id}
    content = quote_workbook.export_quote_sheet(db, [inquiry.id], data.buyer.id)
    _apply_quote(db, data, _edit_quote(content, {2: changes}))
    for viewer in (data.sales, data.peer):
        view = visibility.inquiry_view(db, inquiry, _ctx(db, viewer))
        assert view['line_responses'] == []
        assert view['progress']['quoted_lines'] == 0
        if original:
            assert original.id not in {offer['id'] for offer in view['offers']}
    supervisor = visibility.inquiry_view(db, inquiry, _ctx(db, data.other))
    assert supervisor['line_responses'][0]['is_draft'] is True
    assert supervisor['progress']['quoted_lines'] == 0
    own = visibility.inquiry_view(db, inquiry, _ctx(db, data.buyer))
    assert own['line_responses'][0]['is_draft'] is True
    assert own['progress']['quoted_lines'] == 0
    if original:
        assert own['line_responses'][0]['offer']['unit_price'] == original.unit_price
        assert original.id not in {offer['id'] for offer in own['offers']}
    choice = {'line_id': line.id, 'mode': mode, **({'reason': '正式无货'} if mode == 'skip' else {'ref_offer_id': original.id})}
    offers.reply(db, inquiry.id, data.buyer.id, expected_version=inquiry.version, line_responses=[choice])
    for viewer in (data.sales, data.buyer):
        view = visibility.inquiry_view(db, inquiry, _ctx(db, viewer))
        assert len(view['line_responses']) == 1 and not view['line_responses'][0].get('is_draft')
        assert view['progress']['quoted_lines'] == 1


@pytest.mark.parametrize('legacy', [False, True])
def test_quote_workbook_reads_the_same_delivery_date(db, quote_data, legacy):
    data = quote_data
    inquiry = _submitted(db, data)
    line = inquiries.lines_for(db, inquiry.id)[0]
    if legacy:
        field_pool.save_values(db, field_pool.Subject(inquiries.SUBJECT_TYPE, f'line:{line.id}', 'line'),
            [{'new_name': '期望交期', 'value': '2026-12-07'}], _ctx(db, data.sales))
    else:
        line.expected_delivery = date(2026, 12, 7)
        db.flush()
    content = quote_workbook.export_quote_sheet(db, [inquiry.id], data.buyer.id)
    sheet = load_workbook(BytesIO(content))['采购报价表']
    cols = {cell.value: cell.column for cell in sheet[1]}
    assert '期望交期 2026-12-07' in sheet.cell(2, cols['客户要求']).value


def test_schema_api_uses_current_document_and_user(db, users):
    inquiry = inquiries.create_draft(db, users.seller.id, lines=[{'pn_raw': 'CUSTOM-T19',
        'field_values': [{'new_name': '质保', 'value': '三年'}]}])
    db.commit()
    client = _client(users.seller)
    result = client.get('/api/presales/workbooks/inquiry-schema', params={'inquiry_id': inquiry.id})
    assert result.status_code == 200, result.text
    assert any(column['key'] == '+质保' and column['custom'] for column in result.json()['sheets'][0]['columns'])
    assert _client(users.other).get('/api/presales/workbooks/inquiry-schema', params={'inquiry_id': inquiry.id}).status_code == 403


@pytest.mark.parametrize('mode', ['private', 'include', 'exclude'])
def test_schema_hides_team_field_with_invisible_override(db, presales_users, presales_parts, mode):
    boss = _ctx(db, presales_users.boss)
    owner = _ctx(db, presales_users.sales)
    field = field_pool.define_field(db, subject_type=inquiries.SUBJECT_TYPE, level='line',
        name='受控列', owner_scope='team', viewer=boss)
    inquiry = inquiries.create_draft(db, presales_users.sales.id, lines=[{'part_id': presales_parts[0].id}])
    line = inquiries.lines_for(db, inquiry.id)[0]
    user_ids = [presales_users.other_sales.username] if mode == 'include' else [presales_users.sales.username] if mode == 'exclude' else []
    field_pool.save_values(db, field_pool.Subject(inquiries.SUBJECT_TYPE, f'line:{line.id}', 'line'),
        [{'field_def_id': field.id, 'value': '隐藏', 'rule': {'mode': mode, 'user_ids': user_ids}}], boss)
    assert '+受控列' not in {column['key'] for column in inquiry_workbook.export_schema(db, owner, inquiry.id)['sheets'][0]['columns']}
    assert any(column['key'] == '+受控列' for column in inquiry_workbook.export_schema(db, boss, inquiry.id)['sheets'][0]['columns'])


@pytest.mark.parametrize('expected', [None, '2026-12-09'])
def test_explicit_delivery_overrides_legacy_and_can_be_cleared(db, presales_users, expected):
    user = presales_users.sales
    ctx = _ctx(db, user)
    inquiry = inquiries.create_draft(db, user.id, lines=[{'pn_raw': 'LEGACY'}])
    line = inquiries.lines_for(db, inquiry.id)[0]
    field_pool.save_values(db, field_pool.Subject(inquiries.SUBJECT_TYPE, f'line:{line.id}', 'line'),
        [{'new_name': '期望交期', 'value': '2026-12-08'}], ctx)
    inquiries.update_draft(db, inquiry.id, user.id, inquiry.version,
        lines=[{'line_key': line.line_key, 'expected_delivery': expected}])
    assert line.expected_delivery == (date.fromisoformat(expected) if expected else None)
    assert list(db.scalars(select(FieldValue))) == []


def test_legacy_delivery_export_rechecks_m8_after_rule_changes(db, presales_users):
    register_workbooks()
    owner = presales_users.sales
    boss_ctx = _ctx(db, presales_users.boss)
    field = field_pool.define_field(db, subject_type=inquiries.SUBJECT_TYPE, level='line',
        name='期望交期', owner_scope='team', viewer=boss_ctx)
    inquiry = inquiries.create_draft(db, owner.id, lines=[{'pn_raw': 'DELIVERY-SCOPE'}])
    line = inquiries.lines_for(db, inquiry.id)[0]
    subject = field_pool.Subject(inquiries.SUBJECT_TYPE, f'line:{line.id}', 'line')
    field_pool.save_values(db, subject, [{'field_def_id': field.id, 'value': '2026-12-10'}], boss_ctx)
    content = inquiry_workbook.export_inquiry(db, inquiry.id, _ctx(db, owner))
    field_pool.save_values(db, subject, [{'field_def_id': field.id, 'value': '2026-12-10',
        'rule': {'mode': 'private'}}], boss_ctx)
    from app.services.workbook.registry import Forbidden
    with pytest.raises(Forbidden):
        _preview(db, owner, content, inquiry)


def test_quote_legacy_delivery_export_rechecks_m8_after_rule_changes(db, quote_data):
    data = quote_data
    inquiry = _submitted(db, data)
    line = inquiries.lines_for(db, inquiry.id)[0]
    owner_ctx = _ctx(db, data.sales)
    subject = field_pool.Subject(inquiries.SUBJECT_TYPE, f'line:{line.id}', 'line')
    field_pool.save_values(db, subject, [{'new_name': '期望交期', 'value': '2026-12-10'}], owner_ctx)
    field = field_pool.list_defs(db, inquiries.SUBJECT_TYPE, 'line', owner_ctx)[0]
    content = quote_workbook.export_quote_sheet(db, [inquiry.id], data.buyer.id)
    field_pool.save_values(db, subject, [{'field_def_id': field.id, 'value': '2026-12-10',
        'rule': {'mode': 'private'}}], owner_ctx)
    from app.services.workbook.registry import Forbidden
    with pytest.raises(Forbidden):
        operations.preview(db, quote_workbook.QUOTE_SCHEMA.name, content, ctx=_ctx(db, data.buyer))
