"""M4 的文件边界契约：通过真实 XLSX 覆盖类型、身份和办公软件往返。"""
from __future__ import annotations

import importlib
import io
import json
import os
import shutil
import subprocess
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import replace
from datetime import date, datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from app.services.workbook.export import COLORS, export_workbook
from app.services.workbook.parse import parse_workbook
from app.services.workbook.protocol import read_metadata, write_metadata
from app.services.workbook.schema import (
    ColumnSpec, SheetSpec, WorkbookContext, WorkbookError, WorkbookSchema,
)

CTX = WorkbookContext(field_groups={"data_supplier"}, actor="tester")
VALUES = {
    "pn": "00231", "code": "00001234567890123456", "text": "原始备注",
    "qty": 2, "price": Decimal("1234.50"), "day": date(2026, 1, 2),
    "state": "ok", "enabled": True, "reference": "参考", "supplier": "秘密供应商",
}


def _schema(*, allow_insert=True, allow_delete=True, custom_columns=True):
    columns = [
        ColumnSpec("pn", "PN", "code", "locked", "编码按文本保存", width=27),
        ColumnSpec("code", "用户编码", "code", "optional", "保留前导零"),
        ColumnSpec("text", "备注", "text", "optional", "最多20字", max_len=20),
        ColumnSpec("qty", "数量", "int", "required", "必须为整数"),
        ColumnSpec("price", "金额", "decimal", "optional", "最多2位小数", decimals=2),
        ColumnSpec("day", "日期", "date", "optional", "YYYY-MM-DD"),
        ColumnSpec("state", "状态", "enum", "optional", "请选择", choices={"ok": "正常", "bad": "异常"}),
        ColumnSpec("enabled", "启用", "bool", "optional", "是或否"),
        ColumnSpec("reference", "参考", "text", "reference", "仅供参考"),
        ColumnSpec("supplier", "供应商", "text", "optional", "受字段组权限控制", field_group="data_supplier"),
    ]
    return WorkbookSchema("demo", 1, [SheetSpec("items", "明细", columns,
        allow_insert=allow_insert, allow_delete=allow_delete, custom_columns=custom_columns)], ["仅修改允许填写的字段"])


def _export(schema=None, *, rows=None, export_id="export-one", ctx=CTX, **kwargs):
    schema = schema or _schema()
    if rows is None:
        rows = [{"entity_id": 101, "base_version": 3, "values": VALUES.copy(), "custom_fields": {"自定义": "旧值"}}]
    return export_workbook(schema, {"items": rows}, export_id=export_id, actor="tester", ctx=ctx, **kwargs)


def _save(wb):
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def _edit(payload, changes, *, row=2):
    wb = load_workbook(io.BytesIO(payload))
    ws = wb["明细"]
    columns = {cell.value: cell.column for cell in ws[1]}
    for header, value in changes.items():
        cell = ws.cell(row, columns[header])
        cell.value = value
        if isinstance(value, str) and not value.startswith("="):
            cell.data_type = "s"
    return _save(wb)


def _parsed(changes, *, schema=None):
    schema = schema or _schema()
    return parse_workbook(schema, _edit(_export(schema), changes), ctx=CTX)


def _rewrite(payload, transform=None, *, extra=None):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(payload)) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry)
            target.writestr(entry, transform(entry.filename, data) if transform else data)
        for name, data in (extra or {}).items():
            target.writestr(name, data)
    return output.getvalue()


def test_schema_serialization_and_export_column_permissions():
    schema = _schema()
    public = WorkbookContext()
    encoded = json.loads(json.dumps(schema.to_dict(ctx=public), ensure_ascii=False))
    columns = encoded["sheets"][0]["columns"]
    assert "supplier" not in {c["key"] for c in columns}
    assert next(c for c in columns if c["key"] == "pn")["mandatory"] is True
    payload = _export(ctx=public, visible_columns={"items": ["text"]})
    wb = load_workbook(io.BytesIO(payload))
    assert [c.value for c in wb["明细"][1]] == ["_row_token", "PN", "备注"]
    _, manifest = read_metadata(wb)
    assert "秘密供应商" not in json.dumps(manifest, ensure_ascii=False)
    parsed = parse_workbook(schema, payload, ctx=public)
    assert not parsed.errors
    assert "qty" not in parsed.rows[0].cells


@pytest.mark.parametrize("header", ["编码 ", " 编码", "编\u200b码", "编\u00a0码"])
def test_schema_rejects_headers_that_would_change_on_upload(header):
    with pytest.raises(ValueError, match="表头"):
        ColumnSpec("pn", header, "code", "locked", "文本")


@pytest.mark.parametrize("kind,value", [("int", 123456789012345678), ("decimal", Decimal("123456789012345678.12")), ("decimal", Decimal("1E+20"))])
def test_large_numbers_roundtrip_without_excel_precision_loss(kind, value):
    schema = WorkbookSchema("large", 1, [SheetSpec("items", "明细", [
        ColumnSpec("number", "数字", kind, "optional", "无损数字", decimals=2 if kind == "decimal" else None),
    ])])
    payload = _export(schema, rows=[{"entity_id": 1, "base_version": 1, "values": {"number": value}}])
    assert load_workbook(io.BytesIO(payload))["明细"]["B2"].data_type == "s"
    parsed = parse_workbook(schema, payload, ctx=CTX)
    assert not parsed.errors
    assert parsed.rows[0].cells["number"].value == value


def test_export_styles_and_compatibility_constraints():
    payload = _export(custom_fields={"items": ["自定义"]})
    wb = load_workbook(io.BytesIO(payload))
    ws = wb["明细"]
    headers = {cell.value: cell.column for cell in ws[1]}
    assert wb.sheetnames[0] == "填写说明"
    assert wb["_meta"].sheet_state == "hidden"
    assert ws.column_dimensions["A"].hidden
    assert ws.freeze_panes.endswith("2")
    assert ws.protection.sheet and not ws.protection.password
    assert ws.column_dimensions["B"].width == 27
    assert not ws.merged_cells.ranges and not ws.tables and not ws.conditional_formatting
    for column in _schema().sheets[0].columns:
        cell = ws.cell(2, headers[column.header])
        assert cell.fill.fgColor.rgb[-6:] == COLORS[column.edit]
        assert cell.protection.locked is not column.editable
        assert ws.cell(1, headers[column.header]).comment is not None
        if column.type in {"text", "code"}:
            assert cell.number_format == "@" and cell.data_type == "s"
    assert ws.cell(2, headers["+自定义"]).fill.fgColor.rgb[-6:] == COLORS["custom"]
    assert ws.cell(2, headers["+自定义"]).protection.locked is False
    validations = list(ws.data_validations.dataValidation)
    assert len(validations) == 2
    assert all(len(dv.formula1) <= 255 for dv in validations)
    instructions = " ".join(str(c.value or "") for row in wb["填写说明"] for c in row)
    assert all(value in instructions for value in ["灰底", "黄底", "绿底", "蓝底", "白底", "必须为整数", "新增", "删行", "仅修改允许填写的字段"])
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        assert not any(name.endswith("vbaProject.bin") or name.startswith("xl/tables/") for name in archive.namelist())


def test_long_enum_options_are_instructions_only_and_text_formula_stays_literal():
    choices = ["选项" + str(n) + "x" * 50 for n in range(8)]
    schema = WorkbookSchema("long", 1, [SheetSpec("items", "明细", [
        ColumnSpec("choice", "状态", "enum", "optional", "从说明选择", choices=choices),
        ColumnSpec("text", "备注", "text", "optional", "文本"),
    ])])
    payload = _export(schema, rows=[{"entity_id": 1, "base_version": 1, "values": {"choice": choices[0], "text": "=HYPERLINK(\"https://example.invalid\")"}}])
    wb = load_workbook(io.BytesIO(payload))
    assert not wb["明细"].data_validations.dataValidation
    assert wb["明细"]["C2"].data_type == "s"
    assert choices[-1] in " ".join(str(c.value or "") for row in wb["填写说明"] for c in row)
    assert not parse_workbook(schema, payload, ctx=CTX).errors


@pytest.mark.parametrize("raw", [231, 2.31e7, 12345678901234567890, 2310000.0])
@pytest.mark.parametrize("header,key", [("PN", "pn"), ("用户编码", "code")])
def test_numeric_codes_reject_row_without_guessing(raw, header, key):
    parsed = _parsed({header: raw})
    cell = parsed.rows[0].cells[key]
    assert "请将该列设为文本后重新粘贴" in cell.error
    assert any(e["row_number"] == 2 and e["column"] == header for e in parsed.errors)
    if key == "code":
        assert cell.value is None


@pytest.mark.parametrize("raw,expected", [(" ￥１，２３４．５０ ", Decimal("1234.50")), ("¥ 1,234.50", Decimal("1234.50")), (" 12.00 ", Decimal("12.00")), ("", None)])
def test_decimal_normalization(raw, expected):
    parsed = _parsed({"金额": raw})
    assert not parsed.errors
    assert parsed.rows[0].cells["price"].value == expected


@pytest.mark.parametrize("raw", ["1.234", "NaN", "Infinity", "1,23", True])
def test_decimal_errors_are_located_and_never_rounded(raw):
    parsed = _parsed({"金额": raw})
    assert parsed.rows[0].cells["price"].error
    assert parsed.rows[0].cells["price"].value is None
    assert parsed.errors[0]["column"] == "金额"


@pytest.mark.parametrize("raw,expected", [("￥１，２３４", 1234), ("¥ 1,234", 1234), (" １２ ", 12)])
def test_integer_normalization(raw, expected):
    parsed = _parsed({"数量": raw})
    assert not parsed.errors
    assert parsed.rows[0].cells["qty"].value == expected


@pytest.mark.parametrize("raw", ["1.1", "", None])
def test_required_integer_rejects_fraction_and_empty(raw):
    parsed = _parsed({"数量": raw})
    assert parsed.rows[0].cells["qty"].error


@pytest.mark.parametrize("raw", [date(2026, 1, 2), datetime(2026, 1, 2, 12), "2026-01-02", "2026/01/02", 46024])
def test_date_formats_and_1900_serial(raw):
    parsed = _parsed({"日期": raw})
    assert not parsed.errors
    assert parsed.rows[0].cells["day"].value == date(2026, 1, 2)


@pytest.mark.parametrize("raw", ["2026.01.02", "2026/02/30", "02/01/2026"])
def test_invalid_date_formats(raw):
    assert _parsed({"日期": raw}).rows[0].cells["day"].error


def test_date_1904_serial_zero_and_signed_metadata():
    schema = _schema()
    wb = load_workbook(io.BytesIO(_edit(_export(), {"日期": 0})))
    metadata, manifest = read_metadata(wb)
    metadata["date_system"] = "1904"
    del wb["_meta"]
    wb.epoch = CALENDAR_MAC_1904
    wb["明细"]["G2"].number_format = "General"
    write_metadata(wb, metadata, manifest)
    parsed = parse_workbook(schema, _save(wb), ctx=CTX)
    assert not parsed.errors
    assert parsed.rows[0].cells["day"].value == date(1904, 1, 1)


@pytest.mark.parametrize("raw,expected", [("是", True), ("否", False), ("Y", True), ("N", False), (1, True), (0, False), ("TRUE", True), ("false", False), ("对", True), ("错", False)])
def test_bool_aliases(raw, expected):
    parsed = _parsed({"启用": raw})
    assert not parsed.errors
    assert parsed.rows[0].cells["enabled"].value is expected


@pytest.mark.parametrize("raw,expected", [("OK", "ok"), ("正常", "ok"), ("bAd", "bad"), ("异常", "bad")])
def test_enum_codes_and_chinese_labels(raw, expected):
    parsed = _parsed({"状态": raw})
    assert not parsed.errors
    assert parsed.rows[0].cells["state"].value == expected


def test_boolean_enum_aliases():
    schema = WorkbookSchema("bool_enum", 1, [SheetSpec("items", "明细", [ColumnSpec("v", "是否", "enum", "optional", "是否", choices={"Y": "是", "N": "否"})])])
    payload = _export(schema, rows=[{"entity_id": 1, "base_version": 1, "values": {"v": "Y"}}])
    parsed = parse_workbook(schema, _edit(payload, {"是否": "错"}), ctx=CTX)
    assert not parsed.errors and parsed.rows[0].cells["v"].value == "N"


def test_text_cleaning_max_length_and_control_character_rejection():
    parsed = _parsed({"备注": " \ufeff\u200bA\u00a0B\u200d\u2060 "})
    assert parsed.rows[0].cells["text"].value == "AB"
    assert _parsed({"备注": "x" * 21}).rows[0].cells["text"].error
    with pytest.raises(WorkbookError, match="控制字符"):
        _parsed({"备注": "bad_x0001_text"})


def test_formula_requires_cache_and_accepts_only_cached_value():
    payload = _edit(_export(), {"金额": "=1+2"})
    parsed = parse_workbook(_schema(), payload, ctx=CTX)
    assert "缓存" in parsed.rows[0].cells["price"].error
    def cache_formula(name, data):
        if name == "xl/worksheets/sheet2.xml":
            root = ET.fromstring(data)
            namespace = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            cell = root.find(".//m:c[@r='F2']/m:v", namespace)
            assert cell is not None
            cell.text = "3"
            return ET.tostring(root)
        return data
    cached = parse_workbook(_schema(), _rewrite(payload, cache_formula), ctx=CTX)
    assert not cached.errors
    assert cached.rows[0].cells["price"].value == Decimal(3)


def test_external_formula_and_external_link_package_rejected():
    with pytest.raises(WorkbookError, match="外部链接"):
        _parsed({"金额": "='[other.xlsx]Sheet1'!A1"})
    with pytest.raises(WorkbookError, match="外部链接"):
        parse_workbook(_schema(), _rewrite(_export(), extra={"xl/externalLinks/externalLink1.xml": b"<externalLink/>"}), ctx=CTX)


def test_empty_rows_reordering_and_inserted_row_are_not_position_identity():
    rows = [{"entity_id": n, "base_version": n, "values": {**VALUES, "text": str(n)}} for n in (101, 102)]
    wb = load_workbook(io.BytesIO(_export(rows=rows)))
    ws = wb["明细"]
    first, second = [[c.value for c in ws[r]] for r in (2, 3)]
    for col, value in enumerate(second, 1):
        ws.cell(2, col).value = value
    for col, value in enumerate(first, 1):
        ws.cell(3, col).value = value
    ws.insert_rows(3)
    ws.cell(6, 5).value = 9
    parsed = parse_workbook(_schema(), _save(wb), ctx=CTX)
    assert not parsed.errors
    assert [row.entity_id for row in parsed.rows if row.kind == "exported"] == [102, 101]
    assert [row.kind for row in parsed.rows] == ["exported", "exported", "added"]
    assert parsed.rows[-1].row_number == 6


def test_insert_delete_flags_and_custom_fields():
    schema = _schema(allow_insert=False, allow_delete=False)
    wb = load_workbook(io.BytesIO(_export(schema)))
    ws = wb["明细"]
    ws.delete_rows(2)
    ws.cell(2, 5).value = 5
    parsed = parse_workbook(schema, _save(wb), ctx=CTX)
    assert {r.kind for r in parsed.rows} == {"added", "deleted"}
    assert any("不允许新增" in e["message"] for e in parsed.errors)
    assert any("不允许删除" in e["message"] for e in parsed.errors)
    payload = _export(custom_fields={"items": ["自定义"]})
    parsed = parse_workbook(_schema(), _edit(payload, {"+自定义": " \u200b新值 "}), ctx=CTX)
    assert parsed.rows[0].custom_fields == {"自定义": "新值"}


@pytest.mark.parametrize("mode", ["duplicate", "forged", "cross-export"])
def test_invalid_tokens_reject_whole_file(mode):
    payload = _export()
    wb = load_workbook(io.BytesIO(payload))
    ws = wb["明细"]
    if mode == "duplicate":
        for index, cell in enumerate(list(ws[2]), 1):
            ws.cell(3, index).value = cell.value
    elif mode == "forged":
        ws["A2"] = "forged-token"
    else:
        other = load_workbook(io.BytesIO(_export(export_id="other-export")))
        ws["A2"] = other["明细"]["A2"].value
    with pytest.raises(WorkbookError, match="令牌"):
        parse_workbook(_schema(), _save(wb), ctx=CTX)


def test_locked_and_reference_modifications_are_warnings_and_keep_baseline():
    parsed = _parsed({"PN": "changed", "参考": "changed", "备注": "新备注"})
    assert not parsed.errors
    assert parsed.rows[0].cells["pn"].value == VALUES["pn"]
    assert parsed.rows[0].cells["reference"].value == VALUES["reference"]
    assert len(parsed.warnings) == 2
    assert parsed.rows[0].cells["text"].value == "新备注"


def test_revoked_field_permission_ignores_uploaded_column_and_baseline():
    parsed = parse_workbook(_schema(), _export(), ctx=WorkbookContext())
    assert not parsed.errors
    assert "supplier" not in parsed.rows[0].cells
    assert "supplier" not in parsed.rows[0].baseline
    assert any(e["column"] == "供应商" for e in parsed.warnings)


@pytest.mark.parametrize("mode", ["missing", "tampered", "wrong-schema", "wrong-version"])
def test_metadata_and_schema_identity_are_required(mode):
    schema = _schema()
    wb = load_workbook(io.BytesIO(_export()))
    if mode == "missing":
        del wb["_meta"]
    elif mode == "tampered":
        wb["_meta"]["B1"] = "tampered"
    elif mode == "wrong-schema":
        schema = replace(schema, name="other")
    else:
        schema = replace(schema, version=2)
    with pytest.raises(WorkbookError, match="模板"):
        parse_workbook(schema, _save(wb), ctx=CTX)


def test_merged_region_and_absent_required_identity_columns_rejected():
    wb = load_workbook(io.BytesIO(_export()))
    wb["明细"].merge_cells("C2:D2")
    with pytest.raises(WorkbookError, match="C2:D2"):
        parse_workbook(_schema(), _save(wb), ctx=CTX)
    wb = load_workbook(io.BytesIO(_export()))
    wb["明细"].delete_cols(2)
    with pytest.raises(WorkbookError, match="身份编码"):
        parse_workbook(_schema(), _save(wb), ctx=CTX)


def test_zip_size_entry_limits_and_invalid_content(monkeypatch):
    parser = importlib.import_module("app.services.workbook.parse")
    payload = _export()
    monkeypatch.setattr(parser, "MAX_FILE_BYTES", len(payload) - 1)
    with pytest.raises(WorkbookError, match="10 MB"):
        parse_workbook(_schema(), payload, ctx=CTX)
    monkeypatch.setattr(parser, "MAX_FILE_BYTES", 10 * 1024 * 1024)
    monkeypatch.setattr(parser, "MAX_UNCOMPRESSED_BYTES", 100)
    with pytest.raises(WorkbookError, match="解压"):
        parse_workbook(_schema(), payload, ctx=CTX)
    monkeypatch.setattr(parser, "MAX_UNCOMPRESSED_BYTES", 64 * 1024 * 1024)
    monkeypatch.setattr(parser, "MAX_ZIP_ENTRIES", 1)
    with pytest.raises(WorkbookError, match="条目"):
        parse_workbook(_schema(), payload, ctx=CTX)
    with pytest.raises(WorkbookError, match="有效"):
        parse_workbook(_schema(), b"not a zip", ctx=CTX)


@pytest.mark.skipif(os.environ.get("RUN_SOFFICE_ROUNDTRIP") != "1" or not shutil.which("soffice"), reason="仅 RUN_SOFFICE_ROUNDTRIP=1 且安装 LibreOffice 时执行")
def test_libreoffice_real_roundtrip_preserves_values_tokens_and_types():
    schema = _schema()
    payload = _export(custom_fields={"items": ["自定义"]})
    with tempfile.TemporaryDirectory(prefix="t15-lo-", dir="/tmp") as directory:
        from pathlib import Path
        root = Path(directory)
        source = root / "demo.xlsx"
        output = root / "resaved"
        output.mkdir()
        source.write_bytes(payload)
        result = subprocess.run(["soffice", "-env:UserInstallation=file:///tmp/t15-lo-profile", "--headless", "--convert-to", "xlsx", "--outdir", str(output), str(source)], capture_output=True, text=True, timeout=60)
        target = output / source.name
        assert result.returncode == 0 and target.exists(), result.stdout + result.stderr
        parsed = parse_workbook(schema, target.read_bytes(), ctx=CTX)
        assert not parsed.errors
        assert len(parsed.rows) == 1 and parsed.rows[0].kind == "exported"
        row = parsed.rows[0]
        assert row.entity_id == 101 and row.base_version == 3
        assert {key: c.value for key, c in row.cells.items()} == VALUES
        assert row.custom_fields == {"自定义": "旧值"}
        wb = load_workbook(target)
        original = load_workbook(io.BytesIO(payload))
        ws = wb["明细"]
        assert ws["A2"].value == original["明细"]["A2"].value
        assert ws["B2"].data_type == "s" and ws["B2"].number_format == "@"
        assert ws["F2"].value == 1234.5 and ws["G2"].is_date
        assert wb["_meta"].sheet_state == "hidden"
        assert ws.protection.sheet
        assert ws["B1"].comment and ws["B2"].fill.fgColor.rgb[-6:] == COLORS["locked"]
        assert len(ws.data_validations.dataValidation) == 2


def test_allowed_deletion_preserves_identity_baseline_and_new_custom_only_row():
    wb = load_workbook(io.BytesIO(_export()))
    ws = wb["明细"]
    ws.delete_rows(2)
    ws.cell(1, 12).value = "+新增备注"
    ws.cell(4, 12).value = "自定义内容"
    ws.cell(4, 5).value = 1
    parsed = parse_workbook(_schema(), _save(wb), ctx=CTX)
    assert not parsed.errors
    deletion = next(row for row in parsed.rows if row.kind == "deleted")
    addition = next(row for row in parsed.rows if row.kind == "added")
    assert deletion.entity_id == 101 and deletion.base_version == 3
    assert deletion.baseline["pn"] == "00231"
    assert addition.custom_fields == {"新增备注": "自定义内容"}


def test_readonly_only_row_is_ignored_and_new_row_required_value_is_checked():
    payload = _edit(_export(rows=[]), {"PN": "readonly-only"})
    assert not parse_workbook(_schema(), payload, ctx=CTX).rows
    parsed = parse_workbook(_schema(), _edit(payload, {"备注": "有内容"}), ctx=CTX)
    assert len(parsed.rows) == 1 and parsed.rows[0].kind == "added"
    assert any(e["column"] == "数量" and "必填" in e["message"] for e in parsed.errors)


@pytest.mark.parametrize("column,value", [("启用", "maybe"), ("状态", "unknown")])
def test_invalid_bool_and_enum_are_located(column, value):
    parsed = _parsed({column: value})
    assert any(error["column"] == column and error["row_number"] == 2 for error in parsed.errors)


def test_new_unauthorized_column_does_not_bypass_export_scope():
    schema = _schema()
    payload = _export(visible_columns={"items": ["qty"]}, ctx=WorkbookContext())
    wb = load_workbook(io.BytesIO(payload))
    ws = wb["明细"]
    ws.cell(1, 4).value = "供应商"
    ws.cell(2, 4).value = "无权新值"
    parsed = parse_workbook(schema, _save(wb), ctx=WorkbookContext())
    assert "supplier" not in parsed.rows[0].cells
    assert any(e["column"] == "供应商" for e in parsed.warnings)


def test_signature_covers_manifest_baseline_and_control_text_in_unknown_sheet():
    wb = load_workbook(io.BytesIO(_export()))
    meta = wb["_meta"]
    for row in meta:
        if str(row[0].value).startswith("payload_"):
            row[1].value = row[1].value.replace("00231", "fake-value")
            break
    with pytest.raises(WorkbookError, match="签名"):
        parse_workbook(_schema(), _save(wb), ctx=CTX)
    wb = load_workbook(io.BytesIO(_export()))
    wb.create_sheet("用户附表")["A1"] = "_x0008_"
    with pytest.raises(WorkbookError, match="用户附表 A1.*控制字符"):
        parse_workbook(_schema(), _save(wb), ctx=CTX)


def test_sparse_worksheet_far_cell_does_not_become_a_new_row():
    wb = load_workbook(io.BytesIO(_export()))
    wb["明细"]["IV20000"] = "无表头的附加内容"
    parsed = parse_workbook(_schema(), _save(wb), ctx=CTX)
    assert not parsed.errors
    assert len(parsed.rows) == 1
    assert parsed.rows[0].entity_id == 101


def test_formula_without_cache_outside_business_columns_still_errors():
    wb = load_workbook(io.BytesIO(_export()))
    wb.create_sheet("用户附表")["A1"] = "=1+2"
    parsed = parse_workbook(_schema(), _save(wb), ctx=CTX)
    assert any(error["sheet"] == "用户附表" and "缓存" in error["message"]
               for error in parsed.errors)


@pytest.mark.parametrize("title", ["_meta", "_META", "_MeTa", "填写说明"])
def test_schema_rejects_case_insensitive_reserved_sheet_titles(title):
    with pytest.raises(ValueError, match="系统工作表"):
        SheetSpec("items", title, [])
