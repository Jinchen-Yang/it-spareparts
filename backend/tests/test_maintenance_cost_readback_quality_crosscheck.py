"""Independent post-normal cost evidence checks, using real PostgreSQL readers.

Expected values come from D-25 and the authoritative recompute, not the writer.
Only synthetic facts are created; conftest owns and removes the run database.
"""
import io
import os
from decimal import Decimal
from pathlib import Path
import subprocess
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.api import maintenance_project_master_workbook as workbook_api
from app.models.maintenance import FMaintenanceLine, MaintenanceManualCostOverride
from app.services import maintenance_cost as authority
from app.services import maintenance_project_master_workbook as master
from app.services import maintenance_workbook_cost_writes as writer
from tests.test_maintenance_expense_collection_workbook import uploader
from tests.test_maintenance_project_master_v2_editable import _make_project_with_line


@pytest.fixture(autouse=True)
def _optional_mutation(monkeypatch):
    mode = os.environ.get('COST_QUALITY_CROSSCHECK_MUTANT')
    if not mode:
        return
    if mode == 'baseline':
        source = subprocess.check_output([
            'git', 'show', '3d63d0d703140da8fa585af39631c33ca63845d7:'
            'backend/app/services/maintenance_workbook_cost_writes.py'], text=True)
    elif mode == 'launder-source':
        source = Path(writer.__file__).read_text().replace(
            'if line.cost_source in (None, "none", "manual"):', 'if True:')
    elif mode == 'proposed-qty-guard':
        source = Path(writer.__file__).read_text().replace(
            'if (line.cost_source == "manual" and line.cost_tax_basis',
            'if (line.qty is not None and line.cost_source == "manual" and line.cost_tax_basis')
    else:
        raise AssertionError(mode)
    namespace = dict(writer.__dict__)
    exec(compile(source, f'<{mode}-writer>', 'exec'), namespace)
    for name in ('_merge_manual_cost_to_line', '_recompute_line_amounts'):
        monkeypatch.setattr(writer, name, namespace[name])
        monkeypatch.setattr(master, name, namespace[name])


def _evidence(db, line_id, side):
    """Use the exact aggregate SQL expression, then compare workbook Python."""
    amount, actual, estimated, missing = authority._normalized_cost_expressions(
        FMaintenanceLine, basis=side)
    row = db.execute(select(amount, actual, estimated, missing).select_from(
        FMaintenanceLine).outerjoin(MaintenanceManualCostOverride,
        MaintenanceManualCostOverride.line_id == FMaintenanceLine.id).where(
        FMaintenanceLine.id == line_id)).one()
    db.expire_all()
    line = db.get(FMaintenanceLine, line_id)
    override = db.scalar(select(MaintenanceManualCostOverride).where(
        MaintenanceManualCostOverride.line_id == line_id))
    fact = master._line_cost_evidence(line, override, basis=side)
    tier = 'actual' if row[1] else 'estimated' if row[2] else 'missing'
    assert row[3] == (tier == 'missing')
    assert (fact['amount'], fact['tier']) == (row[0], tier)
    return row[0], tier


def _fill_http(db, project_id, line_id, price='7.13'):
    client = uploader(db, username=f'quality-{uuid4().hex[:8]}')
    response = client.get(f'/api/maintenance/projects/stable/{project_id}/master-workbook.xlsx',
                          params={'sheets': master.V2_SHEET_PARTS})
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content))
    sheet = book[master.V2_SHEET_PARTS]
    columns = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, columns['人工未税单位成本'], float(price))
    sheet.cell(2, columns['人工成本原因'], '独立质量核验')
    content = io.BytesIO()
    book.save(content)
    db.commit()
    response = client.post(f'/api/maintenance/projects/stable/{project_id}/master-workbook/apply',
        files={'file': ('quality.xlsx', content.getvalue(),
                       'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert response.status_code == 200, response.text
    db.expire_all()
    return db.get(FMaintenanceLine, line_id)


@pytest.mark.parametrize('side', ['inc', 'ex'])
def test_price_only_missing_quantity_stays_unknown_like_authoritative_recompute(db, monkeypatch, side):
    monkeypatch.setattr(workbook_api.get_settings(), 'maintenance_project_master_v2_enabled', True)
    project, _, _, line = _make_project_with_line(db, unit_cost=None, cost_source='none')
    line.qty = None
    line.anomaly_flags = ['missing_qty', 'no_cost']
    db.commit()
    project_id, line_id = project.project_id, line.id
    _fill_http(db, project_id, line_id)
    after_fill = _evidence(db, line_id, side)
    authority.recompute(db, line_ids={line_id})
    after_authority = _evidence(db, line_id, side)
    assert after_authority == (None, 'missing')
    assert after_fill == after_authority, (after_fill, after_authority)


@pytest.mark.parametrize('flag,side', [('inc_tax_estimated', 'inc'), ('ex_tax_estimated', 'ex')])
def test_historical_tax_estimate_retention_is_preexisting_not_new_promotion(db, flag, side):
    # Existing complete manual evidence already carries the stale flag. Both
    # baseline and candidate keep it; only authoritative recompute removes it.
    _, _, _, line = _make_project_with_line(db, unit_cost=Decimal('3'), cost_source='manual')
    line.anomaly_flags = [flag, 'tax_rate_estimated', 'no_cost']
    db.commit()
    line_id = line.id
    writer._merge_manual_cost_to_line(db, master.CostRefill(
        line_id, Decimal('7.13'), Decimal('8.06'), '替换失效证据'), operated_by='quality')
    db.commit()
    after_fill = _evidence(db, line_id, side)
    authority.recompute(db, line_ids={line_id})
    after_authority = _evidence(db, line_id, side)
    assert after_authority == (Decimal('16.12' if side == 'inc' else '14.26'), 'actual')
    assert after_fill == (after_authority[0], 'estimated')


@pytest.mark.parametrize('legacy_present', [False, True])
def test_missing_quantity_quantity_only_edit_does_not_create_new_legacy_evidence(db, legacy_present):
    _, _, _, line = _make_project_with_line(db, unit_cost=Decimal('7.13'), cost_source='manual')
    line.qty = None
    line.cost_amount = Decimal('0') if legacy_present else None
    line.cost_amount_ex_tax = line.cost_amount_inc_tax = Decimal('0')
    db.commit()
    line_id = line.id
    before = _evidence(db, line_id, 'inc')
    assert before == ((Decimal('0'), 'actual') if legacy_present else (None, 'missing'))
    writer._recompute_line_amounts(db.get(FMaintenanceLine, line_id))
    db.commit()
    assert _evidence(db, line_id, 'inc') == before


@pytest.mark.parametrize('source,basis,expected', [
    ('sales_ref', 'inc', 'estimated'),
    ('trace_avg', 'ex', 'estimated'),
    ('sales_history', 'inc', 'estimated'),
    ('direct', 'bad', 'missing'),
    ('future-source', 'inc', 'missing'),
])
def test_zero_manual_override_cannot_launder_or_replace_automatic_evidence(db, source, basis, expected):
    _, _, _, line = _make_project_with_line(db, qty=Decimal('0.5'),
        unit_cost=Decimal('2'), cost_source=source)
    line.cost_tax_basis = basis
    line.anomaly_flags = ['inc_tax_estimated']
    db.commit()
    line_id = line.id
    before = {side: _evidence(db, line_id, side) for side in ('ex', 'inc')}
    writer._merge_manual_cost_to_line(db, master.CostRefill(
        line_id, Decimal('0'), Decimal('0'), '人工零价证据'), operated_by='quality')
    db.commit()
    for side in ('ex', 'inc'):
        assert _evidence(db, line_id, side) == before[side]
        assert before[side][1] == expected
    override = db.scalar(select(MaintenanceManualCostOverride).where(
        MaintenanceManualCostOverride.line_id == line_id))
    assert override.active and override.unit_cost_inc_tax == Decimal('0')


def test_fractional_quantity_half_cent_rounding_matches_authority_on_both_sides(db):
    _, _, _, line = _make_project_with_line(db, qty=Decimal('0.5'),
        unit_cost=None, cost_source='none')
    line_id = line.id
    writer._merge_manual_cost_to_line(db, master.CostRefill(
        line_id, Decimal('0.01'), Decimal('0.01'), '合法小数'), operated_by='quality')
    db.commit()
    written = [_evidence(db, line_id, side) for side in ('ex', 'inc')]
    authority.recompute(db, line_ids={line_id})
    recomputed = [_evidence(db, line_id, side) for side in ('ex', 'inc')]
    assert written == recomputed == [(Decimal('0.01'), 'actual')] * 2
