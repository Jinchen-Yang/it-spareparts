"""Independent post-normal checks for manual cost transaction boundaries.

Real workbook HTTP commands and isolated PostgreSQL, with fresh counterexamples
for a previously materialized manual fact. No production network or data.
"""
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from datetime import date
from decimal import Decimal
import io
from threading import Event
from time import monotonic, sleep
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import event, select, text

from app.api import maintenance_project_master_workbook as workbook_api
from app.db import SessionLocal
from app.etl import loader
from app.main import app
from app.models.dimensions import DimPart
from app.models.maintenance import FMaintenanceLine, MaintenanceManualCostOverride
from app.models.maintenance_project_operations import (
    MaintenanceProjectOperationAudit, MaintenanceProjectWorkbookOperation,
    MaintenanceProjectWorkbookState,
)
from app.services import maintenance_cost
from app.services import maintenance_project_master_workbook as master
from tests import factories as f
from tests.test_maintenance_expense_collection_workbook import uploader
from tests.test_maintenance_project_master_v2_editable import _batch, _make_project_with_line

_XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
_TABLES = (FMaintenanceLine, MaintenanceManualCostOverride,
           MaintenanceProjectOperationAudit, MaintenanceProjectWorkbookOperation,
           MaintenanceProjectWorkbookState)


@pytest.fixture(autouse=True)
def _v2_with_bounded_waits(monkeypatch):
    monkeypatch.setattr(workbook_api.get_settings(), 'maintenance_project_master_v2_enabled', True)

    def bounded(_session, _transaction, connection):
        connection.execute(text("SET LOCAL lock_timeout = '8s'"))
        connection.execute(text("SET LOCAL statement_timeout = '15s'"))

    event.listen(SessionLocal, 'after_begin', bounded)
    try:
        yield
    finally:
        event.remove(SessionLocal, 'after_begin', bounded)


def _snapshot(db):
    return {model.__tablename__: deepcopy(list(db.execute(
        select(model.__table__).order_by(*model.__table__.primary_key.columns)
    ).mappings())) for model in _TABLES}


def _bytes(book):
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _export(client, project_id):
    response = client.get(
        f'/api/maintenance/projects/stable/{project_id}/master-workbook.xlsx',
        params={'sheets': master.V2_SHEET_PARTS})
    assert response.status_code == 200, response.text
    return response.content


def _edit(content, changes):
    book = load_workbook(io.BytesIO(content))
    sheet = book[master.V2_SHEET_PARTS]
    columns = {cell.value: cell.column for cell in sheet[1]}
    for name, value in changes.items():
        sheet.cell(2, columns[name], value)
    return _bytes(book)


def _post(client, project_id, content):
    return client.post(f'/api/maintenance/projects/stable/{project_id}/master-workbook/apply',
        files={'file': ('cost-boundary.xlsx', content, _XLSX)})


def _succeeded(response):
    assert response.status_code == 200, response.text
    return response.json()


def _row(client, project_id, line_id):
    response = client.get(f'/api/maintenance/projects/stable/{project_id}/master-workbook/rows',
                          params={'sheet': master.V2_SHEET_PARTS})
    assert response.status_code == 200, response.text
    return next(row for row in response.json()['rows'] if row['line_id'] == line_id)


def _seed(db):
    project, part, order, line = _make_project_with_line(
        db, unit_cost=None, cost_source='none')
    project_id, line_id = project.project_id, line.id
    client = uploader(db, username=f'cost-txn-{uuid4().hex[:8]}')
    _succeeded(_post(client, project_id, _edit(_export(client, project_id), {
        '人工未税单位成本': 7.13, '人工成本原因': '先前已确认凭据',
    })))
    db.expire_all()
    override = db.scalar(select(MaintenanceManualCostOverride).where(
        MaintenanceManualCostOverride.line_id == line_id))
    assert override.version == 1 and override.active
    db.commit()
    return project, part, order, line, override, client


def _assert_materialized(db, client, project_id, line_id, *, price, amount_ex, amount_inc):
    db.expire_all()
    line = db.get(FMaintenanceLine, line_id)
    assert (line.cost_source, line.cost_tax_basis, line.unit_cost, line.cost_amount,
            line.cost_amount_ex_tax, line.cost_amount_inc_tax) == (
        'manual', 'ex', Decimal(price), Decimal(amount_ex),
        Decimal(amount_ex), Decimal(amount_inc))
    db.commit()
    row = _row(client, project_id, line_id)
    assert row['missing_kind'] is None
    assert Decimal(str(row['cost_amount_inc_tax'])) == Decimal(amount_inc)


def test_existing_evidence_update_flush_failure_restores_price_fields_versions_and_protection(db):
    project, _part, _order, line, override, client = _seed(db)
    project_id, line_id = project.project_id, line.id
    content = _edit(_export(client, project_id), {
        '人工未税单位成本': 9.17, '人工成本原因': '更正发票',
        '需求数量': 4, '退货数量': 1, '备注': '与价格同事务',
    })
    before = _snapshot(db)
    db.commit()
    seen = []

    def fail_after_ledger_update(_mapper, connection, target):
        if target.line_id == line_id:
            # SQL UPDATE of the existing evidence has already happened. Check
            # real persisted-in-this-transaction values before aborting it.
            actual = connection.execute(select(MaintenanceManualCostOverride.__table__).where(
                MaintenanceManualCostOverride.line_id == line_id)).mappings().one()
            assert actual['version'] == 2 and actual['unit_cost_ex_tax'] == Decimal('9.17')
            seen.append(target.id)
            raise master.WorkbookError('synthetic_cost_update_failure', '合成账本更新后失败')

    event.listen(MaintenanceManualCostOverride, 'after_update', fail_after_ledger_update)
    try:
        response = _post(client, project_id, content)
    finally:
        event.remove(MaintenanceManualCostOverride, 'after_update', fail_after_ledger_update)
    assert seen == [override.id]
    assert response.status_code == 422, response.text
    assert response.json()['detail']['code'] == 'synthetic_cost_update_failure'
    assert _snapshot(db) == before
    db.commit()

    first = _succeeded(_post(client, project_id, content))
    assert not first['replayed']
    _assert_materialized(db, client, project_id, line_id,
                         price='9.17', amount_ex='27.51', amount_inc='31.08')
    db.refresh(override)
    db.refresh(line)
    assert override.version == 2
    assert set(line.manual_override) == {'qty', 'return_qty', 'line_note'}
    after = _snapshot(db)
    db.commit()
    assert _succeeded(_post(client, project_id, content)) == {**first, 'replayed': True}
    assert _snapshot(db) == after


def test_stale_different_price_upload_cannot_overwrite_new_evidence_or_quantities(db):
    project, _part, _order, line, override, client = _seed(db)
    project_id, line_id = project.project_id, line.id
    exported = _export(client, project_id)
    first_file = _edit(exported, {'人工未税单位成本': 9.17, '需求数量': 3})
    stale_file = _edit(exported, {'人工未税单位成本': 11.11, '需求数量': 8})
    _succeeded(_post(client, project_id, first_file))
    before = _snapshot(db)
    db.commit()
    rejected = _post(client, project_id, stale_file)
    assert rejected.status_code == 409, rejected.text
    assert _snapshot(db) == before
    _assert_materialized(db, client, project_id, line_id,
                         price='9.17', amount_ex='27.51', amount_inc='31.08')
    db.refresh(override)
    assert override.version == 2


def test_same_upload_waits_for_uncommitted_cost_update_then_replays_one_revision(db):
    project, _part, _order, line, override, client = _seed(db)
    project_id, line_id = project.project_id, line.id
    content = _edit(_export(client, project_id), {
        '人工未税单位成本': 9.17, '需求数量': 4, '退货数量': 1,
    })
    state = db.get(MaintenanceProjectWorkbookState, project_id)
    revision = state.revision
    audit_count = len(db.scalars(select(MaintenanceProjectOperationAudit)).all())
    db.commit()
    at_flush, release = Event(), Event()
    first_pid = []

    def pause_after_update(_mapper, connection, target):
        if target.line_id == line_id and not first_pid:
            first_pid.append(connection.scalar(text('select pg_backend_pid()')))
            at_flush.set()
            assert release.wait(12), 'test did not release the first uncommitted upload'

    other = TestClient(app)
    other.headers.update(client.headers)
    event.listen(MaintenanceManualCostOverride, 'after_update', pause_after_update)
    try:
        with ThreadPoolExecutor(max_workers=2) as pool:
            first_future = pool.submit(_post, client, project_id, content)
            try:
                assert at_flush.wait(8), 'first upload never reached evidence UPDATE'
                second_future = pool.submit(_post, other, project_id, content)
                deadline = monotonic() + 5
                blocked = False
                while monotonic() < deadline:
                    blocked = bool(db.scalar(text(
                        'SELECT EXISTS (SELECT 1 FROM pg_stat_activity '
                        'WHERE :blocker = ANY(pg_blocking_pids(pid)))'),
                        {'blocker': first_pid[0]}))
                    db.commit()
                    if blocked:
                        break
                    sleep(0.025)
                assert blocked, 'second HTTP upload must really wait on the first PostgreSQL transaction'
                assert not second_future.done()
            finally:
                release.set()
            first = _succeeded(first_future.result(timeout=12))
            replay = _succeeded(second_future.result(timeout=12))
    finally:
        release.set()
        event.remove(MaintenanceManualCostOverride, 'after_update', pause_after_update)
        other.close()
    assert replay == {**first, 'replayed': True}
    _assert_materialized(db, client, project_id, line_id,
                         price='9.17', amount_ex='27.51', amount_inc='31.08')
    db.refresh(override)
    db.refresh(state)
    assert override.version == 2 and state.revision == revision + 1
    assert len(db.scalars(select(MaintenanceProjectOperationAudit)).all()) == audit_count + 2


def test_price_only_http_fill_keeps_source_quantity_live_and_readable_after_source_upsert(db):
    project, part, order, line, override, client = _seed(db)
    project_id, line_id = project.project_id, line.id
    override_id, version = override.id, override.version
    # Match the actual import transaction: loader owns source fields, followed
    # by authoritative cost recompute before commit (imports.py / import_wbdd).
    source = f.maintenance_line(order.raw_order_id, line.raw_line_id, part.pn_std,
                                qty='5', return_qty='0')
    source['line_note'] = '原件仍可更新'
    head = f.maintenance_head(order.raw_order_id, order_no=order.order_no,
                              sales_order=order.linked_sales_order_no)
    loader.load(db, f.maintenance_result({order.raw_order_id: head}, [source]),
                _batch(db), date(2026, 10, 23), mode='upsert')
    maintenance_cost.recompute(db, commit=False)
    db.commit()
    db.refresh(line)
    db.refresh(override)
    assert line.qty == Decimal('5') and line.line_note == '原件仍可更新'
    assert line.manual_override == {} and line.edited_source == 'wbdd'
    assert override.id == override_id and override.version == version and override.active
    db.commit()
    # Retained manual evidence must remain usable after the completed import
    # transaction and the next HTTP quantity edit.
    row = _row(client, project_id, line_id)
    assert Decimal(str(row['cost_amount_inc_tax'])) == Decimal('40.30')
    content = _edit(_export(client, project_id), {'退货数量': 2})
    _succeeded(_post(client, project_id, content))
    db.expire_all()
    assert db.get(FMaintenanceLine, line_id).return_qty == Decimal('2')
    db.commit()
    row = _row(client, project_id, line_id)
    assert Decimal(str(row['cost_amount_inc_tax'])) == Decimal('24.18')
    db.refresh(override)
    assert override.version == version and override.active


@pytest.mark.parametrize('new_price', [None, 9.17])
def test_identity_change_retires_old_price_and_only_explicit_new_price_becomes_readable(db, new_price):
    project, _part, _order, line, override, client = _seed(db)
    project_id, line_id = project.project_id, line.id
    replacement = DimPart(pn_std=f'COST-NEW-PN-{uuid4().hex[:8].upper()}', status='active')
    db.add(replacement)
    db.commit()
    changes = {'PN': replacement.pn_std, '需求数量': 3}
    if new_price is not None:
        changes.update({'人工未税单位成本': new_price, '人工成本原因': '仅适用新 PN'})
    _succeeded(_post(client, project_id, _edit(_export(client, project_id), changes)))
    db.refresh(line)
    db.refresh(override)
    assert line.part_id == replacement.id
    assert set(line.manual_override) == {'pn_raw', 'pn_std', 'part_id', 'qty'}
    db.commit()
    row = _row(client, project_id, line_id)
    if new_price is None:
        assert override.active is False and override.version == 2
        assert row['cost_amount_inc_tax'] is None
    else:
        assert override.active is True and override.version == 3
        assert override.reason == '仅适用新 PN'
        _assert_materialized(db, client, project_id, line_id,
                             price='9.17', amount_ex='27.51', amount_inc='31.08')
