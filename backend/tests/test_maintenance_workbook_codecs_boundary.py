"""Compatibility and dependency contracts for configured workbook codecs."""
import ast
import hashlib
import hmac
import io
from importlib.util import resolve_name
import json
import subprocess
import sys
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_EVEN, localcontext
from pathlib import Path
from types import SimpleNamespace
from typing import get_type_hints

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import MAC_EPOCH, WINDOWS_EPOCH
from pydantic import SecretStr

from app.config import get_settings
from app.models.maintenance import FMaintenanceLine, MaintenanceManualCostOverride, FProjectExpense
from app.models.maintenance_project_operations import MaintenanceSiteIssue, MaintenanceSiteIssueLine
from app.services import maintenance_expense_collection_workbook as ec
from app.services import maintenance_project_master_workbook as master
from app.services import maintenance_workbook_excel as excel
from app.services import maintenance_workbook_rows as rows
from app.services import maintenance_workbook_signatures as signatures
from app.services import maintenance_workbook_values as values
from app.services.maintenance_workbook_protocol import V2_SHEET_META

_EXPORTS = {
    signatures: ('_global_cost_base_token', '_parse_global_cost_base_token',
                 '_v2_row_base_token', '_parse_v2_row_base_token', '_v2_sign_meta_rows'),
    excel: ('_decimal', '_v2_date', '_v2_decimal', '_v2_meta', '_v2_readonly_sheet_hash'),
    rows: ('_v2_part_row_values', '_v2_expense_row_values', '_v2_receipt_row_values',
           '_v2_site_row_values', '_v2_site_identity_values'),
}
_OLD_SECRET = 'w1b-old-test-secret-only-12345678901234567890'
_NEW_SECRET = 'w1b-new-test-secret-only-12345678901234567890'


@pytest.fixture
def keyring(monkeypatch):
    settings = get_settings()
    monkeypatch.setattr(settings, 'maintenance_manifest_active_key_id', 'w1b-old')
    monkeypatch.setattr(settings, 'maintenance_manifest_active_hmac_key', SecretStr(_OLD_SECRET))
    monkeypatch.setattr(settings, 'maintenance_manifest_previous_hmac_keys_json', SecretStr('none'))
    return settings


def test_codec_imports_do_not_load_adapters_queries_or_writers():
    backend = Path(__file__).resolve().parents[1]
    script = f'''
import importlib, json, sys
from sqlalchemy.engine import Engine
def unexpected_connection(*args, **kwargs):
    raise AssertionError('codec import attempted a database connection')
Engine.connect = unexpected_connection
for name in {[module.__name__ for module in _EXPORTS]!r}:
    importlib.import_module(name)
print(json.dumps(sorted(name for name in sys.modules if name.startswith('app.'))))
'''
    loaded = json.loads(subprocess.run([sys.executable, '-c', script], cwd=backend,
                       check=True, capture_output=True, text=True).stdout)
    forbidden = ('app.api', 'app.main',
        'app.services.maintenance_project_master_workbook',
        'app.services.maintenance_expense_collection_workbook',
        'app.services.maintenance_boss_board', 'app.services.maintenance_project_operations',
        'app.services.maintenance_site_issue_manual')
    assert not [name for name in loaded if name.startswith(forbidden)]
    allowed_services = {*(module.__name__ for module in _EXPORTS),
        'app.services.maintenance_field_protection', 'app.services.maintenance_workbook_plans',
        'app.services.maintenance_workbook_protocol', 'app.services.maintenance_workbook_values'}
    assert {name for name in loaded if name.startswith('app.services.')} <= allowed_services
    # Inspect the full local closure, including package initializers and imports
    # inside functions. ORM Base imports app.db; no adapter dependency is allowed.
    pending = [module.__name__ for module in _EXPORTS]
    visited = set()
    while pending:
        name = pending.pop()
        if name in visited:
            continue
        path = backend / (name.replace('.', '/') + '.py')
        package = backend / name.replace('.', '/') / '__init__.py'
        path = package if package.exists() else path
        if not path.exists():
            continue
        visited.add(name)
        assert not name.startswith(forbidden)
        if name.startswith('app.services.'):
            assert name in allowed_services
        for length in range(1, len(name.split('.'))):
            pending.append('.'.join(name.split('.')[:length]))
        for node in ast.walk(ast.parse(path.read_text())):
            if isinstance(node, ast.Import):
                imports = [alias.name for alias in node.names]
            elif isinstance(node, ast.ImportFrom):
                imported = node.module or ''
                if node.level:
                    parent = name if path.name == '__init__.py' else name.rsplit('.', 1)[0]
                    imported = resolve_name('.' * node.level + imported, parent)
                imports = [imported, *(imported + '.' + alias.name for alias in node.names)]
            else:
                continue
            pending.extend(item for item in imports if item == 'app' or item.startswith('app.'))



@pytest.mark.parametrize(('module', 'name'), [
    (module, name) for module, names in _EXPORTS.items() for name in names])
def test_legacy_exports_are_same_function_objects(module, name):
    assert getattr(master, name) is getattr(module, name)


def test_signature_annotations_keep_model_inputs_and_shared_exception():
    hints = get_type_hints(signatures._global_cost_base_token)
    assert hints['line'] is FMaintenanceLine
    assert hints['override'] == MaintenanceManualCostOverride | None
    assert excel.WorkbookError is signatures.WorkbookError is ec.WorkbookError


def _row_token():
    return signatures._v2_row_base_token(sheet='06_领用返还', entity_id='行-1',
        values={'PN': ' PN ', '日期': datetime(2026, 10, 23, 20, 30),
                '数量': Decimal('1.000'), '免返': False, '空': None})


def _global_token():
    return signatures._global_cost_base_token(line=FMaintenanceLine(id=7),
        override=MaintenanceManualCostOverride(version=2, active=False),
        display_unit_cost_ex_tax=Decimal('1.005'),
        display_unit_cost_inc_tax=Decimal('1.135'), visible_reason='人工来源')


def _parse_row(token):
    return signatures._parse_v2_row_base_token(token, sheet='06_领用返还', entity_id='行-1', row_no=8)


def _parse_global(token):
    return signatures._parse_global_cost_base_token(token, line_id=7, row_no=8)


def test_signed_bytes_match_vectors_captured_before_extraction(keyring):
    assert _row_token() == (
        '{"key_id":"w1b-old","payload":{"base":{"PN":" PN ","免返":"False",'
        '"数量":"1","日期":"2026-10-23","空":""},"entity":"行-1",'
        '"sheet":"06_领用返还"},"signature":'
        '"1c4129f4cc9de75262c7229a3cacdd181a3e6d94772b8e75489241b953256682"}')
    with localcontext() as context:
        context.rounding = ROUND_HALF_EVEN
        assert _global_token() == (
            '{"key_id":"w1b-old","payload":{"line_id":7,"override_active":false,'
            '"override_version":2,"reason":"人工来源","unit_cost_ex_tax":"1.00",'
            '"unit_cost_inc_tax":"1.14"},"signature":'
            '"890d8c93d9470f5c0b69cd567e0dae1ad5c15634ad370db237c2f049df534387"}')
    signed = signatures._v2_sign_meta_rows([
        ('project_id', '项目-1'), ('included_sheets', '03_备件明细'), ('parts_row_ids', '1,2')])
    assert signed[-1] == ('metadata_hmac', 'c4ed9fc117d6e44b9d4c887f797ce8b2049d855140e1b9698378bb8f9a13df21')


@pytest.mark.parametrize(('make', 'parse'), [(_row_token, _parse_row), (_global_token, _parse_global)])
def test_key_rotation_and_revocation_are_observed_after_import(keyring, make, parse):
    old = make()
    original_payload = parse(old)
    keyring.maintenance_manifest_active_key_id = 'w1b-new'
    keyring.maintenance_manifest_active_hmac_key = SecretStr(_NEW_SECRET)
    keyring.maintenance_manifest_previous_hmac_keys_json = SecretStr(json.dumps({'w1b-old': _OLD_SECRET}))
    new = make()
    assert json.loads(new)['key_id'] == 'w1b-new'
    assert new != old and parse(new) == parse(old) == original_payload
    keyring.maintenance_manifest_previous_hmac_keys_json = SecretStr('none')
    with pytest.raises(ec.WorkbookError) as caught:
        parse(old)
    assert caught.value.code == 'invalid_concurrency_token'
    assert parse(new) == original_payload


def test_identity_domain_and_malformed_token_failures_keep_adapter_error(keyring):
    token = _row_token()
    for kwargs in ({'sheet': '04_费用报销', 'entity_id': '行-1'}, {'sheet': '06_领用返还', 'entity_id': '行-2'}):
        with pytest.raises(ec.WorkbookError, match='第 9 行'):
            signatures._parse_v2_row_base_token(token, row_no=9, **kwargs)
    for make, parse in ((_row_token, _parse_row), (_global_token, _parse_global)):
        envelope = json.loads(make())
        body = json.dumps(envelope['payload'], ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode()
        envelope['signature'] = hmac.new(_OLD_SECRET.encode(), b'wrong-domain' + body, hashlib.sha256).hexdigest()
        for invalid in (json.dumps(envelope), 'not-json', 'null', '[]', '{}'):
            with pytest.raises(ec.WorkbookError) as caught:
                parse(invalid)
            assert caught.value.code == 'invalid_concurrency_token'


def test_global_absent_override_and_visible_zero_remain_distinct(keyring):
    token = signatures._global_cost_base_token(line=FMaintenanceLine(id=7), override=None,
        display_unit_cost_ex_tax=None, display_unit_cost_inc_tax=Decimal('0'), visible_reason='')
    assert _parse_global(token) == {'override_version': None, 'override_active': None,
        'unit_cost_ex_tax': None, 'unit_cost_inc_tax': Decimal('0.00'), 'reason': None}


@pytest.mark.parametrize(('source', 'epoch', 'expected'), [
    (None, WINDOWS_EPOCH, None), ('', WINDOWS_EPOCH, None),
    (date(2026, 10, 23), WINDOWS_EPOCH, date(2026, 10, 23)),
    (datetime(2026, 10, 23, 23, 59), WINDOWS_EPOCH, date(2026, 10, 23)),
    (' 2026/10/23 ', WINDOWS_EPOCH, date(2026, 10, 23)), ('2026-10', WINDOWS_EPOCH, date(2026, 10, 1)),
    (1, WINDOWS_EPOCH, date(1900, 1, 1)), (60, WINDOWS_EPOCH, date(1900, 2, 28)),
    (61, WINDOWS_EPOCH, date(1900, 3, 1)), (Decimal('61.75'), WINDOWS_EPOCH, date(1900, 3, 1)),
    (1, MAC_EPOCH, date(1904, 1, 2))])
def test_excel_dates_keep_workbook_epoch_and_calendar_semantics(source, epoch, expected):
    assert excel._v2_date(source, row_no=7, label='日期', epoch=epoch) == expected


@pytest.mark.parametrize('value', [True, False, 0, -1, float('nan'), float('inf'), 'not-a-date'])
def test_invalid_excel_dates_do_not_become_numeric_or_unix_dates(value):
    with pytest.raises(ec.WorkbookError) as caught:
        excel._v2_date(value, row_no=7, label='日期')
    assert caught.value.code == 'invalid_date'


def test_decimal_codecs_keep_independent_money_and_quantity_rounding():
    with localcontext() as context:
        context.rounding = ROUND_HALF_EVEN
        assert excel._decimal(' 1.005 ', row_no=2, label='成本') == Decimal('1.00')
        assert excel._v2_decimal('1.005', row_no=2, label='成本') == Decimal('1.01')
        assert excel._v2_decimal('1.2345', row_no=2, label='数量', demand_quantity=True) == Decimal('1.234')
    assert excel._v2_decimal(' 1,234.50 ', row_no=2, label='成本') == Decimal('1234.50')
    assert excel._v2_decimal(None, row_no=2, label='成本') is None
    assert excel._v2_decimal('-0.125', row_no=2, label='成本', allow_negative=True) == Decimal('-0.13')
    for value, kwargs, code in [('', {'required': True}, 'missing_amount'), (True, {}, 'invalid_amount'),
                                ('-1', {}, 'invalid_amount'), ('NaN', {}, 'invalid_amount'),
                                ('1000000000000', {}, 'invalid_amount')]:
        with pytest.raises(ec.WorkbookError) as caught:
            excel._v2_decimal(value, row_no=2, label='成本', **kwargs)
        assert caught.value.code == code


def test_metadata_chunk_round_trip_rebuilds_logical_signature_and_uses_current_key(keyring):
    long_value = ''.join(chr(0x4e00 + i) * 30000 for i in range(10)) + '末'
    logical = [('project_id', '项目-1'), ('parts_base_hashes', long_value)]
    for expected_key, secret in [('w1b-old', _OLD_SECRET), ('w1b-new', _NEW_SECRET)]:
        keyring.maintenance_manifest_active_key_id = expected_key
        keyring.maintenance_manifest_active_hmac_key = SecretStr(secret)
        signed = signatures._v2_sign_meta_rows(logical)
        assert all(len(value) <= 30000 for _, value in signed)
        wb = Workbook()
        ws = wb.create_sheet(V2_SHEET_META)
        for row in reversed(signed):
            ws.append(row)
        payload = io.BytesIO()
        wb.save(payload)
        meta = excel._v2_meta(load_workbook(io.BytesIO(payload.getvalue())))
        assert meta['parts_base_hashes'] == long_value
        assert meta['metadata_hmac_key_id'] == expected_key
        assert meta['metadata_hmac'] == values._v2_metadata_signature(meta, secret.encode())
        assert values._v2_metadata_signature({**meta, 'parts_base_hashes': long_value + '!'}, secret.encode()) != meta['metadata_hmac']


def test_metadata_parser_keeps_existing_overwrite_strip_and_sparse_chunk_tolerance():
    wb = Workbook()
    with pytest.raises(ec.WorkbookError) as caught:
        excel._v2_meta(wb)
    assert caught.value.code == 'template_version_mismatch'
    ws = wb.create_sheet(V2_SHEET_META)
    for row in [(' key ', ' first '), ('key', ' last '), ('key@10', ' tenth '),
                ('key@2', ' second '), ('orphan@3', 'third'), ('plain@x', 'literal'),
                ('zero', 0), (None, 'ignored')]:
        ws.append(row)
    assert excel._v2_meta(wb) == {'key': 'lastsecondtenth', 'orphan': 'third', 'plain@x': 'literal', 'zero': ''}


def test_readonly_hash_ignores_excel_padding_but_keeps_internal_blanks_and_booleans():
    ws = Workbook().active
    ws.append(['中文', None, False, Decimal('1.00'), datetime(2026, 10, 23, 23, 59)])
    original = excel._v2_readonly_sheet_hash(ws)
    ws.cell(8, 12, '')
    assert excel._v2_readonly_sheet_hash(ws) == original
    ws.cell(1, 3, 0)
    assert excel._v2_readonly_sheet_hash(ws) != original
    ws.cell(1, 3, False)
    ws.cell(1, 2, ' ')
    assert excel._v2_readonly_sheet_hash(ws) != original


def test_site_projection_separates_display_from_identity_without_mutating_inputs():
    issue = MaintenanceSiteIssue(issue_no='LY-AUTO-identity', display_no='LY-20261023-0001',
                                 issue_date=date(2026, 10, 23), site_location='现场')
    line = MaintenanceSiteIssueLine(issue_line_id='line-1', pn='PN', serial_number=None,
                                    quantity=Decimal('2.125'), no_return=False, remark=None, demand_order_no='WBDD-1')
    shown = rows._v2_site_row_values(line, issue)
    identity = rows._v2_site_identity_values(line, issue)
    assert shown['领用单号'] == 'LY-20261023-0001'
    assert identity == {**shown, '领用单号': 'LY-AUTO-identity'}
    assert shown['应返数量'] == Decimal('2.125')
    issue.display_no = 'LY-20261023-9999'
    assert rows._v2_site_identity_values(line, issue) == identity
    assert shown['领用单号'] == 'LY-20261023-0001'
    issue.issue_no = 'MANUAL-1'
    assert rows._v2_site_row_values(line, issue)['领用单号'] == 'MANUAL-1'
    assert line.no_return is False and line.quantity == Decimal('2.125')
    line.no_return = None
    assert rows._v2_site_row_values(line, issue)['应返数量'] == '—'
    details = {'return_requirement': {'requirement_status': 'exempt', 'required_quantity': '0'}}
    assert rows._v2_site_row_values(line, issue, details)['应返数量'] == Decimal('0')
    assert details['return_requirement']['required_quantity'] == '0'


def test_part_projection_keeps_zero_unknown_and_manual_evidence_distinct():
    order = SimpleNamespace(order_no='WBDD', order_date=None, linked_sales_order_no='XSDD',
                            demand_type=None, warehouse=None, salesperson=None, business_type=None)
    line = FMaintenanceLine(id=1, qty=Decimal('1.125'), return_qty=Decimal('0'), cost_source=None,
                            unit_cost_ex_tax=None, unit_cost_inc_tax=None, edited_source='page_manual', pn_raw='raw', pn_std=None)
    unknown = rows._v2_part_row_values(line, order, None)
    assert unknown['成本缺失类型'] == 'out_of_scope' and unknown['可补价'] == '是'
    assert unknown['来源'] == '页面人工登记' and unknown['需求数量'] == Decimal('1.125')
    line.unit_cost_ex_tax = Decimal('0')
    zero = rows._v2_part_row_values(line, order, MaintenanceManualCostOverride(unit_cost_ex_tax=Decimal('5.50'), reason='人工证据'))
    assert zero['系统未税单位成本'] == Decimal('0') and zero['人工未税单位成本'] == Decimal('5.50')
    assert zero['成本缺失类型'] == '' and zero['可补价'] == '否'
    assert line.unit_cost_ex_tax == Decimal('0')


def test_receipt_and_expense_projection_keep_supplied_context_and_source_amounts():
    snapshot = SimpleNamespace(report_month=date(2026, 10, 1), cumulative_amount=Decimal('0'),
                               status='confirmed', receipt_reference=None, remark=None, collection_id='c')
    assert rows._v2_receipt_row_values(snapshot, 'provided-XSDD') == {
        '合同编号': 'provided-XSDD', '报告月份': '2026-10', '累计实收金额（含税）': Decimal('0'),
        '状态': 'confirmed', '回款凭证号': '', '备注': '', '实体ID': 'c'}
    snapshot.report_month = '2026-10'
    assert rows._v2_receipt_row_values(snapshot, '')['报告月份'] == '2026-10'
    expense = FProjectExpense(raw_line_id='expense', line_no=1, amount=Decimal('113'),
                              amount_ex_tax=Decimal('100'), amount_inc_tax=Decimal('113'), tax_basis='inc', linked_sales_order_no='XSDD')
    projected = rows._v2_expense_row_values(expense, 'supplied project display')
    assert projected['项目名称'] == 'supplied project display'
    assert projected['原始报销金额'] == projected['含税金额（系统）'] == Decimal('113')
    assert projected['未税金额'] == Decimal('100') and projected['金额口径'] == 'inc'
