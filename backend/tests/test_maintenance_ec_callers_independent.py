"""Independent E1 HTTP contracts, executed unchanged on source and integration.

Real signed workbooks, authenticated requests and PostgreSQL facts are used.
Pass-through probes reject legacy dispatch; the sole fault injection is after
facts have been flushed, immediately before the transaction owner commits.
"""
import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import event, select
from sqlalchemy.orm import Session

from app.api import maintenance_project_master_workbook as master_api
from app.db import SessionLocal
from app.models.maintenance import FProjectExpense
from app.models.maintenance_project import MaintenanceProjectContract
from app.models.maintenance_project_operations import (
    MaintenanceCollectionSnapshot,
    MaintenanceProjectExpenseAttribution,
    MaintenanceProjectOperationAudit,
    MaintenanceProjectWorkbookOperation,
    MaintenanceProjectWorkbookState,
)
from app.services import maintenance_expense_collection_workbook as ec
from app.services import maintenance_project_master_workbook as master
from app.services import maintenance_project_operations as operations
from tests.test_maintenance_expense_collection_workbook import (
    _expense, project, uploader,  # noqa: F401
)
from tests.test_maintenance_master_v2_receipts_shared_path import _seed_month

_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_BASE = "/api/maintenance/projects/stable"


def _save(book):
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def _prepare(db, project, monkeypatch, route):
    monkeypatch.setattr(master_api.get_settings(), "maintenance_project_master_v2_enabled", route == "v2")
    _expense(db)
    operations.get_or_create_workbook_state(db, project_id=project.project_id)
    db.commit()
    client = uploader(db, username=f"e1-callers-{route}")
    suffix = "expense-collection-workbook" if route == "standalone" else "master-workbook"
    url = f"{_BASE}/{project.project_id}/{suffix}"
    params = None if route == "standalone" else {"sheets": master.SHEET_EXPENSE}
    response = client.get(url + ".xlsx", params=params)
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content))
    sheet = book[master.V2_SHEET_EXPENSE if route == "v2" else ec.SHEET_EXPENSE]
    columns = {cell.value: cell.column for cell in sheet[1]}
    amount_column = columns["未税金额"] if route == "v2" else 8
    sheet.cell(2, amount_column, 200)
    return client, url, _save(book)


def _post(client, url, data, *, force=False):
    return client.post(url + "/apply", files={"file": ("e1-callers.xlsx", data, _XLSX)},
                       data={"force_takeover": "true"} if force else {})


def _snapshot(bind, project_id):
    # A separate session makes an accidental nested commit observable.
    with Session(bind) as observer:
        expense = observer.scalar(select(FProjectExpense).where(FProjectExpense.raw_line_id == "BXD-1#1"))
        attribution = observer.get(MaintenanceProjectExpenseAttribution, "bxd:BXD-1#1")
        state = observer.get(MaintenanceProjectWorkbookState, project_id)
        return {
            "expense": (expense.amount_ex_tax, expense.amount_inc_tax),
            "attribution": (attribution.amount_ex_tax, attribution.amount_inc_tax, attribution.version),
            "revision": state.revision,
            "audits": tuple(observer.scalars(select(MaintenanceProjectOperationAudit.id).where(
                MaintenanceProjectOperationAudit.project_id == project_id).order_by(MaintenanceProjectOperationAudit.id))),
            "operations": tuple(observer.scalars(select(MaintenanceProjectWorkbookOperation.id).where(
                MaintenanceProjectWorkbookOperation.project_id == project_id).order_by(MaintenanceProjectWorkbookOperation.id))),
            "collections": tuple(observer.execute(select(
                MaintenanceCollectionSnapshot.collection_id,
                MaintenanceCollectionSnapshot.status,
                MaintenanceCollectionSnapshot.cumulative_amount,
                MaintenanceCollectionSnapshot.version,
            ).where(MaintenanceCollectionSnapshot.project_id == project_id).order_by(MaintenanceCollectionSnapshot.collection_id)).all()),
        }


@pytest.mark.parametrize("route", ["standalone", "v1", "v2"])
def test_real_http_uses_explicit_owner_and_preserves_public_receipt(db, project, monkeypatch, route):
    client, url, data = _prepare(db, project, monkeypatch, route)
    project_id, bind = project.project_id, db.get_bind()
    before = _snapshot(bind, project_id)
    entry = "apply_and_commit" if route == "standalone" else "apply_in_transaction"
    real = getattr(ec, entry)
    observed = []

    def forbidden_legacy(*args, **kwargs):
        pytest.fail("HTTP callers must not redispatch through legacy boolean-mode apply")

    def probe(session, plan, **kwargs):
        assert not {"commit", "bump_revision", "track_change"} & kwargs.keys()
        result = real(session, plan, **kwargs)
        session.flush()
        visible = _snapshot(bind, project_id)
        if route == "standalone":
            assert visible["expense"] == (Decimal("200.00"), Decimal("226.00"))
        else:
            assert visible == before
            assert result.operating_fact_changed is True
            assert kwargs["workbook_state"].project_id == project_id
        observed.append(result)
        return result

    monkeypatch.setattr(ec, "apply", forbidden_legacy)
    monkeypatch.setattr(ec, entry, probe)
    response = _post(client, url, data)
    assert response.status_code == 200, response.text
    assert len(observed) == 1
    assert "_operating_fact_changed" not in response.json()
    assert response.json()["applied_by"] == f"e1-callers-{route}"
    after = _snapshot(bind, project_id)
    assert after["expense"] == (Decimal("200.00"), Decimal("226.00"))
    assert after["attribution"][:2] == after["expense"]
    assert after["revision"] == before["revision"] + 1
    # Existing expense amount updates do not add a project operation audit;
    # V2 additionally persists exactly one command-completion receipt audit.
    assert len(after["audits"]) - len(before["audits"]) == int(route == "v2")
    assert after["attribution"][2] == before["attribution"][2] + 1


@pytest.mark.parametrize("route", ["standalone", "v1", "v2"])
def test_commit_stage_failure_rolls_back_real_http_facts_and_same_file_can_retry(db, project, monkeypatch, route):
    client, url, data = _prepare(db, project, monkeypatch, route)
    project_id, bind = project.project_id, db.get_bind()
    before = _snapshot(bind, project_id)
    seen = []

    def fail_after_facts(session):
        if session is db:
            return
        session.flush()
        expense = session.scalar(select(FProjectExpense).where(FProjectExpense.raw_line_id == "BXD-1#1"))
        if expense is not None and expense.amount_ex_tax == Decimal("200.00"):
            seen.append((expense.amount_ex_tax, _snapshot(bind, project_id)))
            raise ec.WorkbookError("cost_recompute_busy", "independent commit-stage fault")

    event.listen(SessionLocal, "before_commit", fail_after_facts)
    try:
        response = _post(client, url, data)
    finally:
        event.remove(SessionLocal, "before_commit", fail_after_facts)
    assert response.status_code == (422 if route == "standalone" else 409), response.text
    assert response.json()["detail"]["code"] == "cost_recompute_busy"
    if route != "standalone":
        assert response.headers["Retry-After"] == "5"
    assert seen == [(Decimal("200.00"), before)]
    assert _snapshot(bind, project_id) == before
    retry = _post(client, url, data)
    assert retry.status_code == 200, retry.text
    if route == "v2":
        assert retry.json()["replayed"] is False
    after = _snapshot(bind, project_id)
    assert after["expense"] == (Decimal("200.00"), Decimal("226.00"))
    assert after["revision"] == before["revision"] + 1


def test_v2_mixed_real_expense_and_skipped_void_month_replay_exact_first_receipt(db, project, monkeypatch):
    client, url, _data = _prepare(db, project, monkeypatch, "v2")
    contract = db.scalar(select(MaintenanceProjectContract).where(MaintenanceProjectContract.project_id == project.project_id))
    voided = _seed_month(db, project, contract, date(2026, 7, 1), "90000", status="void", voided_by="independent-voider")
    voided_id, project_id, bind = voided.collection_id, project.project_id, db.get_bind()
    download = client.get(url + ".xlsx", params={"sheets": f"{master.V2_SHEET_EXPENSE},{master.V2_SHEET_RECEIPTS}"})
    assert download.status_code == 200, download.text
    book = load_workbook(io.BytesIO(download.content))
    sheet = book[master.V2_SHEET_EXPENSE]
    columns = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, columns["未税金额"], 200)
    sheet = book[master.V2_SHEET_RECEIPTS]
    columns = {cell.value: cell.column for cell in sheet[1]}
    row = sheet.max_row + 1
    for key, value in {"合同编号": contract.contract_no, "报告月份": "2026-07", "累计实收金额（含税）": 85000}.items():
        sheet.cell(row, columns[key], value)
    data = _save(book)
    db.commit()
    first = _post(client, url, data)
    assert first.status_code == 200, first.text
    receipt = first.json()
    assert receipt["replayed"] is False
    assert len(receipt["voided_rows"]) == 1
    notice = receipt["voided_rows"][0]
    assert (notice["entity_id"], notice["sheet"], notice["code"], notice["voided_by"]) == (
        voided_id, master.V2_SHEET_RECEIPTS, "collection_voided", "independent-voider")
    assert notice["message"] in receipt["warnings"]
    after = _snapshot(bind, project_id)
    assert after["expense"] == (Decimal("200.00"), Decimal("226.00"))
    assert after["collections"] == ((voided_id, "void", Decimal("90000.00"), 3),)
    assert len(after["operations"]) == 1
    replay = _post(client, url, data)
    assert replay.status_code == 200, replay.text
    assert replay.json() == {**receipt, "replayed": True}
    assert _snapshot(bind, project_id) == after


def test_v2_expense_conflict_preserves_http_shape_and_takeover_is_explicit(db, project, monkeypatch):
    client, url, data = _prepare(db, project, monkeypatch, "v2")
    project_id, bind = project.project_id, db.get_bind()
    expense = db.scalar(select(FProjectExpense).where(FProjectExpense.raw_line_id == "BXD-1#1"))
    attribution = db.get(MaintenanceProjectExpenseAttribution, "bxd:BXD-1#1")
    expense.amount = expense.amount_ex_tax = attribution.amount_ex_tax = Decimal("150.00")
    expense.amount_inc_tax = attribution.amount_inc_tax = Decimal("169.50")
    attribution.version += 1
    operations.bump_workbook_revision(db, project_id=project_id)
    db.commit()
    before = _snapshot(bind, project_id)
    rejected = _post(client, url, data)
    assert rejected.status_code == 409, rejected.text
    detail = rejected.json()["detail"]
    assert detail["code"] == "row_conflicts"
    assert detail["conflicts"]
    assert _snapshot(bind, project_id) == before
    accepted = _post(client, url, data, force=True)
    assert accepted.status_code == 200, accepted.text
    assert accepted.json()["replayed"] is False
    after = _snapshot(bind, project_id)
    assert after["expense"] == (Decimal("200.00"), Decimal("226.00"))
    assert after["revision"] == before["revision"] + 1
