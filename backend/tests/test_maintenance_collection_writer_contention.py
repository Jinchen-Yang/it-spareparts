"""Actual PostgreSQL contention for all five collection snapshot writers.

Current implementation compatibility, not new business policy. Events position
transactions; pg_blocking_pids, not elapsed time, proves an actual lock wait.
"""
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from decimal import Decimal
import io
from threading import Event
from time import monotonic
from uuid import uuid4

from openpyxl import load_workbook
import pytest
from sqlalchemy import select, text

from app.db import SessionLocal
from app.models.maintenance_project import MaintenanceProject
from app.models.maintenance_project_operations import (
    MaintenanceCollectionSnapshot,
    MaintenanceProjectWorkbookOperation,
    MaintenanceProjectWorkbookState,
)
from app.services import maintenance_expense_collection_workbook as ec
from app.services import maintenance_project_identity as identity
from app.services import maintenance_project_master_workbook as master
from app.services import maintenance_project_operations as operations
from tests.test_maintenance_project_master_v2_editable import _project_with_receipt, _save
from tests.test_maintenance_project_xsdd_identity import _seed_mergeable_contract_conflict


def _connection_pid(session):
    # A broken lock graph must fail instead of hanging pytest cleanup.
    session.execute(text("SET LOCAL lock_timeout = '8s'"))
    session.execute(text("SET LOCAL statement_timeout = '15s'"))
    return session.scalar(text("SELECT pg_backend_pid()"))


def _assert_blocked(observer, *, waiter, holder, future, query_fragment):
    deadline, pause, latest = monotonic() + 5, Event(), None
    while monotonic() < deadline:
        observer.execute(text("SELECT pg_stat_clear_snapshot()"))
        latest = observer.execute(text(
            "SELECT pg_blocking_pids(pid), wait_event_type, query "
            "FROM pg_stat_activity WHERE pid = :pid"
        ), {"pid": waiter}).one()
        if holder in latest[0]:
            assert latest[1] == "Lock", latest
            assert query_fragment in latest[2], latest
            assert not future.done()
            return
        if future.done():
            future.result()  # Surface real service errors rather than a timeout.
            pytest.fail("contender completed without waiting for the held lock")
        pause.wait(0.01)  # Poll cadence only; not evidence of a lock wait.
    pytest.fail(f"backend {waiter} did not wait for {holder}; last activity={latest}")


def _note_plan(db, project_id):
    june = db.scalar(select(MaintenanceCollectionSnapshot).where(
        MaintenanceCollectionSnapshot.project_id == project_id,
        MaintenanceCollectionSnapshot.report_month == date(2026, 6, 1),
    ))
    may_id = str(uuid4())
    db.add(MaintenanceCollectionSnapshot(
        collection_id=may_id, project_id=project_id,
        project_contract_id=june.project_contract_id,
        report_month=date(2026, 5, 1), cumulative_amount=Decimal("50000"),
        status="confirmed", source="direct_api", version=1,
    ))
    db.commit()
    workbook = load_workbook(io.BytesIO(master.build_project_master_v2(
        db, project_id=project_id,
        sheets=(master.V2_SHEET_OVERVIEW, master.V2_SHEET_RECEIPTS),
    )))
    note = next(row[1] for row in workbook[master.V2_SHEET_OVERVIEW].iter_rows()
                if row[0].value == "项目备注")
    note.value = "V2 committed while collection writer waited"
    sheet = workbook[master.V2_SHEET_RECEIPTS]
    columns = {cell.value: cell.column for cell in sheet[1]}
    row = next(row for row in range(2, sheet.max_row + 1)
               if sheet.cell(row, columns["实体ID"]).value == may_id)
    sheet.cell(row, columns["累计实收金额（含税）"], 60000)
    plan = master.validate_project_master_v2(db, project_id=project_id, data=_save(workbook))
    assert plan.project_note_change is not None
    assert any(guard.entity_id == may_id for guard in plan.row_guards)
    assert not plan.conflicts
    db.commit()  # Export creates state: publish it to both independent sessions.
    return plan


def _receipt_plan(db, project_id, collection_id):
    workbook = load_workbook(io.BytesIO(master.build_project_master_v2(
        db, project_id=project_id, sheets=(master.V2_SHEET_RECEIPTS,),
    )))
    sheet = workbook[master.V2_SHEET_RECEIPTS]
    columns = {cell.value: cell.column for cell in sheet[1]}
    row = next(row for row in range(2, sheet.max_row + 1)
               if sheet.cell(row, columns["实体ID"]).value == collection_id)
    sheet.cell(row, columns["累计实收金额（含税）"], 90000)
    plan = master.validate_project_master_v2(db, project_id=project_id, data=_save(workbook))
    assert not plan.conflicts
    assert any(guard.entity_id == collection_id for guard in plan.row_guards)
    db.commit()
    return plan


def _race_after_v2_guard(db, monkeypatch, *, plan, writer, lock_query):
    guarded, release, writer_started = Event(), Event(), Event()
    pids = {}
    original_guard = master._v2_recheck_row_guards

    def pause_after_guard(session, current_plan):
        original_guard(session, current_plan)
        pids["v2"] = session.scalar(text("SELECT pg_backend_pid()"))
        guarded.set()
        assert release.wait(12), "test did not release the V2 guard barrier"

    monkeypatch.setattr(master, "_v2_recheck_row_guards", pause_after_guard)

    def apply_v2():
        with SessionLocal() as session:
            _connection_pid(session)
            return master.apply_project_master_v2(
                session, plan, operated_by="v2-contention", import_batch_id=str(uuid4()),
            )

    def apply_writer():
        with SessionLocal() as session:
            pids["writer"] = _connection_pid(session)
            writer_started.set()
            result = writer(session)
            session.commit()
            return result

    with ThreadPoolExecutor(max_workers=2) as pool:
        v2_future = pool.submit(apply_v2)
        try:
            assert guarded.wait(5), "V2 never reached its real post-lock guard"
            writer_future = pool.submit(apply_writer)
            assert writer_started.wait(5)
            _assert_blocked(db, waiter=pids["writer"], holder=pids["v2"],
                            future=writer_future, query_fragment=lock_query)
        finally:
            release.set()
        v2_result = v2_future.result(timeout=15)
        result = writer_future.result(timeout=15)
    assert v2_result["replayed"] is False
    db.expire_all()
    assert db.get(MaintenanceProject, plan.project_id).project_note == (
        "V2 committed while collection writer waited"
    )
    may = db.scalar(select(MaintenanceCollectionSnapshot).where(
        MaintenanceCollectionSnapshot.project_id == plan.project_id,
        MaintenanceCollectionSnapshot.report_month == date(2026, 5, 1),
    ))
    assert may.cumulative_amount == Decimal("60000") and may.version == 2
    return result


def _create(session, project_id, contract_id):
    return operations.create_collection(
        session, project_id=project_id, project_contract_id=contract_id,
        report_month=date(2026, 7, 1), cumulative_amount=Decimal("90000"),
        status="confirmed", receipt_reference="contention-create", remark=None,
        reason="contention regression", operated_by="collection-writer",
    )


@pytest.mark.parametrize("writer_kind", ["create", "update", "expense_collection_apply"])
def test_v2_holds_state_until_commit_against_real_collection_writers(db, monkeypatch, writer_kind):
    project, contract, snapshot = _project_with_receipt(db)
    project_id, contract_id, snapshot_id = project.project_id, contract.project_contract_id, snapshot.collection_id
    contract_no = contract.contract_no
    plan = _note_plan(db, project_id)

    def writer(session):
        if writer_kind == "create":
            return _create(session, project_id, contract_id)
        if writer_kind == "update":
            return operations.update_collection(
                session, collection_id=snapshot_id, version=1,
                updates={"cumulative_amount": Decimal("90000")},
                reason="contention update", operated_by="collection-writer",
            )
        return ec.apply(
            session, ec.WorkbookPlan(project_id=project_id, expense_updates=(), collection_ops=(
                ec.CollectionOp(operation="UPDATE", project_contract_id=contract_id,
                                contract_no=contract_no, report_month=date(2026, 6, 1),
                                cumulative_amount=Decimal("90000"),
                                receipt_reference="contention-ec", remark=None),
            )), operated_by="collection-writer", import_batch_id=str(uuid4()),
        )

    result = _race_after_v2_guard(db, monkeypatch, plan=plan, writer=writer,
                                  lock_query="maintenance_project_workbook_state")
    if writer_kind == "create":
        assert result is not None
        row = db.get(MaintenanceCollectionSnapshot, result["collection_id"])
        assert row.collection_id != snapshot_id
        assert row.report_month == date(2026, 7, 1)
        assert row.version == 1
    else:
        row = db.get(MaintenanceCollectionSnapshot, snapshot_id)
        assert row.version == 2
        if writer_kind == "expense_collection_apply":
            assert row.receipt_reference == "contention-ec"
    assert row.cumulative_amount == Decimal("90000.00")


def _identity_fixture(db):
    canonical, source, current, historical, survivor, duplicate, xsdd = (
        _seed_mergeable_contract_conflict(db, suffix="7191")
    )
    # Also require merge's dynamic project_id UPDATE, beyond exact deletion.
    extra_id = str(uuid4())
    db.add(MaintenanceCollectionSnapshot(
        collection_id=extra_id, project_id=source.project_id,
        project_contract_id=historical.project_contract_id,
        report_month=date(2026, 7, 1), cumulative_amount=Decimal("90000"),
        status="confirmed", source="direct_api", version=1,
    ))
    # Keep the original deterministic canonical selection when adding a source
    # fact: equal fact counts fall back to the earlier-created project.
    db.add(MaintenanceCollectionSnapshot(
        collection_id=str(uuid4()), project_id=canonical.project_id,
        project_contract_id=current.project_contract_id,
        report_month=date(2026, 8, 1), cumulative_amount=Decimal("100000"),
        status="confirmed", source="direct_api", version=1,
    ))
    operations.get_or_create_workbook_state(db, project_id=canonical.project_id)
    op = MaintenanceProjectWorkbookOperation(
        project_id=source.project_id, export_id=None, file_sha256="a" * 64,
        operation_key=f"contention-{uuid4()}", payload_hash="b" * 64,
        operation_type="collection_create", entity_id=duplicate.collection_id,
        operated_by="legacy-writer",
    )
    db.add(op)
    db.commit()
    data = dict(canonical=canonical.project_id, source=source.project_id,
                current=current.project_contract_id, historical=historical.project_contract_id,
                survivor=survivor.collection_id, duplicate=duplicate.collection_id,
                xsdd=xsdd, extra=extra_id, operation_id=op.id)
    conflict = next(row for row in identity.preview_historical_conflicts(db)["conflicts"]
                    if row["xsdd_norm"] == xsdd)
    assert conflict["canonical_project_id"] == canonical.project_id
    data["manifest_hash"] = conflict["manifest_hash"]
    db.commit()
    return data


def _merge(session, fixture):
    return identity.apply_historical_project_merge(
        session, xsdd=fixture["xsdd"], expected_manifest_hash=fixture["manifest_hash"],
        expected_canonical_project_id=fixture["canonical"],
        expected_member_project_ids=[fixture["canonical"], fixture["source"]],
        operated_by="reviewed-maintenance-operator",
        contract_resolution={
            "current_project_contract_id": fixture["current"],
            "archive_contracts": [{"project_contract_id": fixture["historical"],
                                   "effective_to": "2026-09-01"}],
            "collection_contract_repoints": [{"collection_id": fixture["survivor"],
                                              "target_project_contract_id": fixture["current"]}],
        },
    )


@pytest.mark.parametrize("writer_kind", ["merge", "exact_dedupe"])
def test_v2_global_lock_blocks_real_identity_snapshot_writers(db, monkeypatch, writer_kind):
    project, _, _ = _project_with_receipt(db)
    plan = _note_plan(db, project.project_id)
    fixture = _identity_fixture(db)
    # Global exclusion also covers unrelated projects: V2 must not invalidate
    # this reviewed merge manifest, so the waiting merge really moves facts.
    def writer(session):
        if writer_kind == "merge":
            return _merge(session, fixture)
        return identity.apply_exact_collection_dedupe(
            session, xsdd=fixture["xsdd"], operated_by="reviewed-maintenance-operator",
        )

    result = _race_after_v2_guard(db, monkeypatch, plan=plan, writer=writer,
                                  lock_query="pg_advisory_xact_lock")
    assert db.get(MaintenanceCollectionSnapshot, fixture["duplicate"]) is None
    assert db.get(MaintenanceCollectionSnapshot, fixture["survivor"]) is not None
    assert db.get(MaintenanceProjectWorkbookOperation, fixture["operation_id"]).entity_id == fixture["survivor"]
    if writer_kind == "merge":
        assert result["deleted_exact_collections"]
        assert db.get(MaintenanceProject, fixture["source"]).is_active is False
        assert db.get(MaintenanceCollectionSnapshot, fixture["extra"]).project_id == fixture["canonical"]
    else:
        assert result["deleted_collection_ids"] == [fixture["duplicate"]]
        assert result["repointed_operation_count"] == 1
        assert db.get(MaintenanceProject, fixture["source"]).is_active is True


def test_v2_waits_for_real_update_then_rechecks_touched_receipt(db):
    project, _, snapshot = _project_with_receipt(db)
    project_id, snapshot_id = project.project_id, snapshot.collection_id
    plan = _receipt_plan(db, project_id, snapshot_id)
    started, pids = Event(), {}

    def apply_v2():
        with SessionLocal() as session:
            pids["v2"] = _connection_pid(session)
            started.set()
            return master.apply_project_master_v2(
                session, plan, operated_by="v2-contention", import_batch_id=str(uuid4()),
            )

    with SessionLocal() as writer:
        writer_pid = _connection_pid(writer)
        result = operations.update_collection(
            writer, collection_id=snapshot_id, version=1,
            updates={"cumulative_amount": Decimal("123456.78")},
            reason="writer wins before V2 locks", operated_by="collection-writer",
        )
        assert result is not None
        with ThreadPoolExecutor(max_workers=1) as pool:
            future = pool.submit(apply_v2)
            try:
                assert started.wait(5)
                _assert_blocked(db, waiter=pids["v2"], holder=writer_pid,
                                future=future, query_fragment="maintenance_project_workbook_state")
            finally:
                writer.commit()
            with pytest.raises(master.WorkbookError) as raised:
                future.result(timeout=15)
        assert raised.value.code == "row_conflicts"
    db.expire_all()
    row = db.get(MaintenanceCollectionSnapshot, snapshot_id)
    assert row.cumulative_amount == Decimal("123456.78")
    assert row.version == 2
    assert db.scalar(select(MaintenanceProjectWorkbookOperation.id).where(
        MaintenanceProjectWorkbookOperation.operation_type == "file_apply"
    )) is None


def test_exact_dedupe_project_lock_blocks_real_create_without_state_lock(db):
    fixture = _identity_fixture(db)
    started, pids = Event(), {}

    def create():
        with SessionLocal() as session:
            pids["create"] = _connection_pid(session)
            started.set()
            result = _create(session, fixture["canonical"], fixture["current"])
            session.commit()
            return result

    with SessionLocal() as deduper:
        deduper_pid = _connection_pid(deduper)
        deleted = identity.apply_exact_collection_dedupe(
            deduper, xsdd=fixture["xsdd"], operated_by="reviewed-maintenance-operator",
        )
        assert deleted["deleted_collection_ids"] == [fixture["duplicate"]]
        with ThreadPoolExecutor(max_workers=1) as pool:
            future = pool.submit(create)
            try:
                assert started.wait(5)
                _assert_blocked(db, waiter=pids["create"], holder=deduper_pid,
                                future=future, query_fragment="maintenance_project.project_id")
            finally:
                deduper.commit()
            created = future.result(timeout=15)
    db.expire_all()
    assert created is not None
    assert db.get(MaintenanceCollectionSnapshot, created["collection_id"]).cumulative_amount == Decimal("90000")
    assert db.get(MaintenanceCollectionSnapshot, fixture["duplicate"]) is None
    assert db.get(MaintenanceProjectWorkbookOperation, fixture["operation_id"]).entity_id == fixture["survivor"]


def test_v2_waits_for_exact_dedupe_then_refuses_deleted_receipt_without_revision_bump(db):
    fixture = _identity_fixture(db)
    # Dedupe has neither a workbook-state lock nor a revision bump. Freshness
    # must still come from the global lock followed by the actual row guard.
    plan = _receipt_plan(db, fixture["source"], fixture["duplicate"])
    prior_revision = db.get(MaintenanceProjectWorkbookState, fixture["source"]).revision
    db.commit()
    started, pids = Event(), {}

    def apply_v2():
        with SessionLocal() as session:
            pids["v2"] = _connection_pid(session)
            started.set()
            return master.apply_project_master_v2(
                session, plan, operated_by="v2-contention", import_batch_id=str(uuid4()),
            )

    with SessionLocal() as deduper:
        deduper_pid = _connection_pid(deduper)
        deleted = identity.apply_exact_collection_dedupe(
            deduper, xsdd=fixture["xsdd"], operated_by="reviewed-maintenance-operator",
        )
        assert deleted["deleted_collection_ids"] == [fixture["duplicate"]]
        with ThreadPoolExecutor(max_workers=1) as pool:
            future = pool.submit(apply_v2)
            try:
                assert started.wait(5)
                _assert_blocked(db, waiter=pids["v2"], holder=deduper_pid,
                                future=future, query_fragment="pg_advisory_xact_lock")
            finally:
                deduper.commit()
            with pytest.raises(master.WorkbookError) as raised:
                future.result(timeout=15)
        assert raised.value.code == "row_conflicts"
    db.expire_all()
    assert db.get(MaintenanceCollectionSnapshot, fixture["duplicate"]) is None
    assert db.get(MaintenanceCollectionSnapshot, fixture["survivor"]).cumulative_amount == Decimal("82325.40")
    assert db.get(MaintenanceProjectWorkbookState, fixture["source"]).revision == prior_revision
    assert db.scalar(select(MaintenanceProjectWorkbookOperation.id).where(
        MaintenanceProjectWorkbookOperation.operation_type == "file_apply"
    )) is None
