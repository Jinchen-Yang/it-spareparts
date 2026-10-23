"""A manual price write must remain a usable fact across HTTP/SQL readers.

These regressions target the cost evidence writer, independently of workbook
module extraction. The old reader's dirty-evidence checks are left intact.
"""
import io
from decimal import Decimal
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.api import maintenance_project_master_workbook as workbook_api
from app.models.maintenance import FMaintenanceLine, MaintenanceManualCostOverride
from app.models.maintenance_project_operations import MaintenanceProjectWorkbookState
from app.services import maintenance_cost_quality as quality
from app.services import maintenance_demand_manual as demand_manual
from app.services import maintenance_project_master_workbook as master
from app.services import maintenance_project_operations as operations
from app.services import maintenance_workbook_cost_writes as costs
from tests.test_maintenance_expense_collection_workbook import uploader
from tests.test_maintenance_project_master_v2_editable import _make_project_with_line

_XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def _save(book):
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def _apply(client, route, content):
    response = client.post(route, files={'file': ('cost.xlsx', content, _XLSX)})
    assert response.status_code == 200, response.text
    return response.json()


def _v2_content(client, project_id, changes):
    response = client.get(
        f'/api/maintenance/projects/stable/{project_id}/master-workbook.xlsx',
        params={'sheets': master.V2_SHEET_PARTS})
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content))
    sheet = book[master.V2_SHEET_PARTS]
    columns = {cell.value: cell.column for cell in sheet[1]}
    for name, value in changes.items():
        sheet.cell(2, columns[name], value)
    return _save(book)


def _assert_fact(db, client, project_id, line_id, *, unit, amount_ex, amount_inc):
    db.expire_all()
    line = db.get(FMaintenanceLine, line_id)
    assert (line.cost_source, line.cost_tax_basis, line.unit_cost, line.cost_amount) == (
        'manual', 'ex', Decimal(unit), Decimal(amount_ex))
    assert (line.cost_amount_ex_tax, line.cost_amount_inc_tax) == (
        Decimal(amount_ex), Decimal(amount_inc))
    # The generated SQL bucket also requires the legacy evidence to be valid.
    assert line.cost_bucket == quality.COST_BUCKET_ACTUAL_EX
    db.commit()
    sheet = (master.V2_SHEET_PARTS
             if workbook_api.get_settings().maintenance_project_master_v2_enabled
             else master.SHEET_PARTS)
    response = client.get(
        f'/api/maintenance/projects/stable/{project_id}/master-workbook/rows',
        params={'sheet': sheet})
    assert response.status_code == 200, response.text
    row = response.json()['rows'][0]
    assert Decimal(str(row['cost_amount_inc_tax'])) == Decimal(amount_inc)
    assert Decimal(str(row['unit_cost_ex_tax'])) == Decimal(unit)
    if sheet == master.V2_SHEET_PARTS:
        assert row['missing_kind'] is None


@pytest.mark.parametrize('entry', ['v1', 'global', 'v2'])
@pytest.mark.parametrize('price,returned,amount_ex,amount_inc', [
    ('7.13', '1', '28.52', '32.24'),
    ('0', '1', '0', '0'),
    ('7.13', '5', '0', '0'),
])
def test_real_workbook_entry_writes_complete_manual_cost_evidence(
    db, monkeypatch, entry, price, returned, amount_ex, amount_inc,
):
    project, _part, _order, line = _make_project_with_line(
        db, qty=Decimal('5'), return_qty=Decimal(returned),
        unit_cost=None, cost_source='none')
    project_id, line_id = project.project_id, line.id
    line.cost_tax_basis = None
    db.commit()
    client = uploader(db, username=f'manual-cost-{uuid4().hex[:8]}')
    monkeypatch.setattr(workbook_api.get_settings(), 'maintenance_project_master_v2_enabled', entry == 'v2')
    route = f'/api/maintenance/projects/stable/{project_id}/master-workbook/apply'
    if entry == 'v2':
        content = _v2_content(client, project_id, {'人工未税单位成本': float(price)})
    else:
        if entry == 'global':
            response = client.get('/api/maintenance/spare-part-lines.xlsx',
                                  params={'range': 'custom', 'from': '2026-01-01', 'to': '2026-12-31'})
            sheet_name, column = master.GLOBAL_SHEET, 10
            route = '/api/maintenance/spare-part-lines/apply'
        else:
            response = client.get(
                f'/api/maintenance/projects/stable/{project_id}/master-workbook.xlsx',
                params={'sheets': master.SHEET_PARTS})
            sheet_name, column = master.SHEET_PARTS, 12
        assert response.status_code == 200, response.text
        book = load_workbook(io.BytesIO(response.content))
        book[sheet_name].cell(2, column, float(price))
        content = _save(book)
    _apply(client, route, content)
    _assert_fact(db, client, project_id, line_id,
                 unit=price, amount_ex=amount_ex, amount_inc=amount_inc)


def test_v1_inclusive_only_input_keeps_canonical_ex_tax_evidence(db, monkeypatch):
    project, _part, _order, line = _make_project_with_line(db, unit_cost=None, cost_source='none')
    project_id, line_id = project.project_id, line.id
    client = uploader(db, username='manual-inc-only')
    monkeypatch.setattr(workbook_api.get_settings(), 'maintenance_project_master_v2_enabled', False)
    response = client.get(
        f'/api/maintenance/projects/stable/{project_id}/master-workbook.xlsx',
        params={'sheets': master.SHEET_PARTS})
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content))
    book[master.SHEET_PARTS].cell(2, 13, 8.06)
    _apply(client, f'/api/maintenance/projects/stable/{project_id}/master-workbook/apply', _save(book))
    _assert_fact(db, client, project_id, line_id, unit='7.13', amount_ex='14.26', amount_inc='16.12')


def test_manual_correction_quantity_only_and_replay_keep_one_evidence_version(db, monkeypatch):
    project, _part, _order, line = _make_project_with_line(db, unit_cost=None, cost_source='none')
    project_id, line_id = project.project_id, line.id
    client = uploader(db, username='manual-edit-replay')
    monkeypatch.setattr(workbook_api.get_settings(), 'maintenance_project_master_v2_enabled', True)
    route = f'/api/maintenance/projects/stable/{project_id}/master-workbook/apply'
    for price in (3, 7.13):
        _apply(client, route, _v2_content(client, project_id, {'人工未税单位成本': price}))
    quantity_file = _v2_content(client, project_id, {'需求数量': 5, '退货数量': 1})
    first = _apply(client, route, quantity_file)
    _assert_fact(db, client, project_id, line_id, unit='7.13', amount_ex='28.52', amount_inc='32.24')
    override = db.scalar(select(MaintenanceManualCostOverride).where(
        MaintenanceManualCostOverride.line_id == line_id))
    state = db.get(MaintenanceProjectWorkbookState, project_id)
    version, revision, updated_at = override.version, state.revision, override.updated_at
    assert version == 2
    db.commit()
    assert _apply(client, route, quantity_file) == {**first, 'replayed': True}
    # A newly exported, unchanged workbook also performs no operating write.
    _apply(client, route, _v2_content(client, project_id, {}))
    db.expire_all()
    assert (override.version, state.revision, override.updated_at) == (version, revision, updated_at)


def test_legacy_only_write_repair_counts_as_change_but_next_equal_fill_does_not(db):
    project, _part, _order, line = _make_project_with_line(db, unit_cost=Decimal('7.13'), cost_source='manual')
    project_id, line_id = project.project_id, line.id
    line.unit_cost = line.cost_amount = line.cost_tax_basis = None
    db.add(MaintenanceManualCostOverride(line_id=line_id, unit_cost_ex_tax=Decimal('7.13'),
        unit_cost_inc_tax=Decimal('8.06'), reason='原有证据', active=True, version=4))
    state = operations.get_or_create_workbook_state(db, project_id=project_id)
    db.commit()
    revision = state.revision
    plan = master.MasterPlan(project_id=project_id, cost_refills=(master.CostRefill(
        line_id, Decimal('7.13'), Decimal('8.06'), '原有证据'),), sheets=(master.SHEET_PARTS,))
    master.apply(db, plan, operated_by='manual-legacy', import_batch_id=str(uuid4()))
    db.expire_all()
    assert state.revision == revision + 1
    assert (line.unit_cost, line.cost_amount, line.cost_tax_basis) == (
        Decimal('7.13'), Decimal('14.26'), 'ex')
    override = db.scalar(select(MaintenanceManualCostOverride).where(
        MaintenanceManualCostOverride.line_id == line_id))
    assert override.version == 4  # price/reason evidence itself has not changed
    db.commit()
    master.apply(db, plan, operated_by='manual-legacy', import_batch_id=str(uuid4()))
    db.expire_all()
    assert state.revision == revision + 1
    assert override.version == 4


def test_page_quantity_correction_retains_same_canonical_manual_fact(db, monkeypatch):
    project, _part, _order, line = _make_project_with_line(db, unit_cost=None, cost_source='none')
    project_id, line_id, raw_line_id = project.project_id, line.id, line.raw_line_id
    client = uploader(db, username='manual-page-quantity')
    monkeypatch.setattr(workbook_api.get_settings(), 'maintenance_project_master_v2_enabled', True)
    _apply(client, f'/api/maintenance/projects/stable/{project_id}/master-workbook/apply',
           _v2_content(client, project_id, {'人工未税单位成本': 7.13}))
    demand_manual.patch_demand_line(db, raw_line_id=raw_line_id, updates={'qty': 5, 'return_qty': 1},
                                    reason='数量更正', operated_by='manual-page-quantity')
    db.commit()
    _assert_fact(db, client, project_id, line_id, unit='7.13', amount_ex='28.52', amount_inc='32.24')


@pytest.mark.parametrize('basis', ['ex', 'inc'])
def test_existing_manual_quantity_edit_preserves_declared_legacy_tax_side(basis):
    line = FMaintenanceLine(qty=Decimal('5'), return_qty=Decimal('1'), cost_source='manual',
        cost_tax_basis=basis, unit_cost_ex_tax=Decimal('7.13'), unit_cost_inc_tax=Decimal('8.06'))
    costs._recompute_line_amounts(line)
    assert (line.unit_cost, line.cost_amount, line.cost_tax_basis) == (
        (Decimal('7.13'), Decimal('28.52'), 'ex') if basis == 'ex'
        else (Decimal('8.06'), Decimal('32.24'), 'inc'))


@pytest.mark.parametrize('source', ['direct', 'window', 'month_avg', 'pool_purchase', 'unknown-dirty'])
def test_manual_price_does_not_overwrite_automatic_or_dirty_source(db, source):
    _project, _part, _order, line = _make_project_with_line(db, cost_source=source)
    names = ('cost_source', 'cost_tax_basis', 'unit_cost', 'cost_amount', 'unit_cost_ex_tax',
             'unit_cost_inc_tax', 'cost_amount_ex_tax', 'cost_amount_inc_tax')
    before = tuple(getattr(line, name) for name in names)
    assert costs._merge_manual_cost_to_line(db, master.CostRefill(
        line.id, Decimal('7.13'), Decimal('8.06'), '仅记录人工证据'), operated_by='precedence')
    db.flush()
    assert tuple(getattr(line, name) for name in names) == before
    override = db.scalar(select(MaintenanceManualCostOverride).where(
        MaintenanceManualCostOverride.line_id == line.id))
    assert override.unit_cost_ex_tax == Decimal('7.13')


@pytest.mark.parametrize('source', ['direct', 'pool_purchase', 'none', 'unknown-dirty'])
def test_nonmanual_quantity_recompute_retains_existing_legacy_behavior(source):
    line = FMaintenanceLine(qty=Decimal('5'), return_qty=Decimal('1'), cost_source=source,
        cost_tax_basis='ex', unit_cost=Decimal('100'), cost_amount=Decimal('200'),
        unit_cost_ex_tax=Decimal('100'), unit_cost_inc_tax=Decimal('113'))
    costs._recompute_line_amounts(line)
    assert (line.cost_amount_ex_tax, line.cost_amount_inc_tax) == (Decimal('400'), Decimal('452'))
    assert (line.unit_cost, line.cost_amount, line.cost_tax_basis) == (Decimal('100'), Decimal('200'), 'ex')


def test_partial_tax_missing_quantity_and_dirty_basis_remain_missing():
    line = FMaintenanceLine(qty=Decimal('5'), return_qty=Decimal('1'), cost_source='manual',
        cost_tax_basis='ex', unit_cost_ex_tax=Decimal('7.13'))
    costs._recompute_line_amounts(line)
    assert line.cost_amount_inc_tax is None  # no synthetic opposite tax side
    assert quality.normalized_tax_tier(source=line.cost_source, tax_basis=line.cost_tax_basis,
        legacy_amount=line.cost_amount, normalized_amount=line.cost_amount_inc_tax,
        normalized_basis='inc', anomaly_flags=()) == 'missing'
    assert quality.normalized_line_cost(source='none', tax_basis=None, legacy_amount=None,
        normalized_amount=None, normalized_basis='inc', anomaly_flags=(), qty=None,
        return_qty=Decimal('0'), manual_unit_cost=Decimal('8.06'), manual_active=True)['amount'] is None
    assert quality.normalized_tax_tier(source='manual', tax_basis='dirty',
        legacy_amount=Decimal('28.52'), normalized_amount=Decimal('32.24'),
        normalized_basis='inc', anomaly_flags=()) == 'missing'


@pytest.mark.parametrize('basis', ['ex', 'inc'])
@pytest.mark.parametrize('missing_side', ['ex', 'inc'])
def test_quantity_edit_cannot_promote_residual_amount_without_complete_manual_prices(basis, missing_side):
    line = FMaintenanceLine(qty=Decimal('4'), return_qty=Decimal('0'),
        cost_source='manual', cost_tax_basis=basis, unit_cost=None, cost_amount=None,
        unit_cost_ex_tax=Decimal('7.13'), unit_cost_inc_tax=Decimal('8.06'),
        cost_amount_ex_tax=Decimal('14.26'), cost_amount_inc_tax=Decimal('16.12'))
    setattr(line, f'unit_cost_{missing_side}_tax', None)
    costs._recompute_line_amounts(line)
    assert line.unit_cost is None and line.cost_amount is None
    for side in ('ex', 'inc'):
        assert quality.normalized_tax_tier(source=line.cost_source, tax_basis=line.cost_tax_basis,
            legacy_amount=line.cost_amount, normalized_amount=getattr(line, f'cost_amount_{side}_tax'),
            normalized_basis=side, anomaly_flags=()) == 'missing'
