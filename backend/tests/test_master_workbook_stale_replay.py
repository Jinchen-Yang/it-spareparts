"""空实体ID的领用重放冲突报告；全部使用合成数据，身份及合并口径不变。"""

import hashlib
import io
import json
import uuid
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models.maintenance_project import MaintenanceProject
from app.models.maintenance_project_operations import (
    MaintenanceSiteIssue,
    MaintenanceSiteIssueLine,
)
from app.services import maintenance_project_master_workbook as master
from app.services import maintenance_project_operations as operations
from tests.test_maintenance_project_master_v2_editable import (
    _make_project_with_line,
    _save,
)


FIX_HINT = (
    "已入库记录请保留实体ID在原行修改；真正的新领用请将领用单号留空自动编号"
    "（或使用尚未占用的单号）"
)


def _seed_legacy_rows(db, count):
    project, part, *_ = _make_project_with_line(db)
    rows = []
    for index in range(1, count + 1):
        issue_no = f"WBDD-20261019-{index:04d}"
        serial_number = f"SYNTHETIC-SN-{index:04d}"
        identity = "|".join([project.project_id, issue_no, part.pn_std, serial_number])
        issue = MaintenanceSiteIssue(
            issue_id=str(uuid.uuid4()), project_id=project.project_id,
            issue_no=issue_no, issue_date=date(2026, 10, 18),
            raw_status="已确认", normalized_status="confirmed",
            status_mapping_state="mapped", status_mapping_version="synthetic-test",
            site_location=f"合成领用地址-{index}", source="legacy",
        )
        line = MaintenanceSiteIssueLine(
            issue_line_id="manual-site:" + hashlib.sha1(identity.encode()).hexdigest(),
            issue_id=issue.issue_id, line_no=1, part_id=part.id,
            pn=part.pn_std, serial_number=serial_number,
            quantity=Decimal("2"), no_return=False, remark=f"原备注-{index}",
            algorithm_version="synthetic-test", tax_rate_used=Decimal("0.13"),
            is_active=True,
        )
        db.add(issue)
        rows.append((issue, line))
    db.flush()
    db.add_all([line for _, line in rows])
    db.commit()
    return project, part, rows


def _book(db, project):
    return load_workbook(io.BytesIO(master.build_project_master_v2(
        db, project_id=project.project_id, sheets=(master.V2_SHEET_SITE,))))


def _append_replay(book, issue, line, **changes):
    sheet = book[master.V2_SHEET_SITE]
    values = {
        "领用单号": issue.issue_no,
        "领用日期": issue.issue_date,
        "PN": line.pn,
        "SN": line.serial_number or "",
        "领用数量": line.quantity,
        "是否应返还": "" if line.no_return is None else "否" if line.no_return else "是",
        "备注": line.remark or "",
        **changes,
    }
    sheet.append([values.get(cell.value) for cell in sheet[1]])
    return sheet.max_row


def _validate(db, project, book):
    return master.validate_project_master_v2(
        db, project_id=project.project_id, data=_save(book))


def _apply(db, project, book):
    plan = _validate(db, project, book)
    result = master.apply_project_master_v2(
        db, plan, operated_by="stale-replay-contract-test",
        import_batch_id=str(uuid.uuid4()))
    return plan, result


def _entity_row(book, entity_id):
    sheet = book[master.V2_SHEET_SITE]
    columns = {cell.value: cell.column for cell in sheet[1]}
    row_no = next(row_no for row_no in range(2, sheet.max_row + 1)
                  if sheet.cell(row_no, columns["实体ID"]).value == entity_id)
    return sheet, columns, row_no


def test_all_legacy_collisions_report_exact_rows_and_only_changed_fields(db):
    project, part, stored = _seed_legacy_rows(db, 4)
    book = _book(db, project)
    modifications = [
        {"领用日期": date(2026, 10, 19)},
        {"是否应返还": "否"},
        {"领用日期": date(2026, 10, 17), "领用数量": 3, "备注": "上传备注"},
        {"领用地址": "合成新地址"},
    ]
    expected_diffs = [
        [{"field": "领用日期", "old": "2026-10-18", "new": "2026-10-19"}],
        [{"field": "是否应返还", "old": "是", "new": "否"}],
        [
            {"field": "领用日期", "old": "2026-10-18", "new": "2026-10-17"},
            {"field": "领用数量", "old": "2", "new": "3"},
            {"field": "备注", "old": "原备注-3", "new": "上传备注"},
        ],
        [{"field": "领用地址", "old": "合成领用地址-4", "new": "合成新地址"}],
    ]
    expected_rows = []
    for (issue, line), changes, differences in zip(stored, modifications, expected_diffs):
        row_no = _append_replay(book, issue, line, **changes)
        expected_rows.append({
            "row_no": row_no, "entity_id": line.issue_line_id,
            "issue_no": issue.issue_no, "pn": part.pn_std,
            "sn": line.serial_number, "differences": differences,
        })

    with pytest.raises(master.WorkbookError) as raised:
        _validate(db, project, book)

    error = raised.value
    assert error.code == "stale_site_replay"
    assert "共 4 行" in error.message
    assert f"第 {expected_rows[0]['row_no']} 行" in error.message
    assert error.details == {
        "rows": expected_rows, "total_count": 4, "fix_hint": FIX_HINT,
    }
    # 日期、Decimal 等只可作为归一化字符串返回，不能依赖默认 JSON serializer。
    assert json.loads(json.dumps(error.details, ensure_ascii=False)) == error.details
    for index, (issue, line) in enumerate(stored, 1):
        db.refresh(issue)
        db.refresh(line)
        assert issue.issue_date == date(2026, 10, 18)
        assert issue.site_location == f"合成领用地址-{index}"
        assert line.quantity == Decimal("2")
        assert line.no_return is False
        assert line.remark == f"原备注-{index}"


def test_collision_details_are_capped_but_total_counts_every_row(db):
    project, _, stored = _seed_legacy_rows(db, 203)
    book = _book(db, project)
    row_numbers = [
        _append_replay(book, issue, line, **{"领用日期": date(2026, 10, 19)})
        for issue, line in stored
    ]

    with pytest.raises(master.WorkbookError) as raised:
        _validate(db, project, book)

    error = raised.value
    assert error.code == "stale_site_replay"
    assert "共 203 行" in error.message
    assert error.details["total_count"] == 203
    assert len(error.details["rows"]) == 200
    assert [row["row_no"] for row in error.details["rows"]] == row_numbers[:200]
    assert [row["entity_id"] for row in error.details["rows"]] == [
        line.issue_line_id for _, line in stored[:200]
    ]


def test_identical_legacy_replay_keeps_identity_and_creates_nothing(db):
    project, _, stored = _seed_legacy_rows(db, 3)
    book = _book(db, project)
    identities = [line.issue_line_id for _, line in stored]
    for issue, line in stored:
        _append_replay(book, issue, line, **{"领用地址": issue.site_location})

    plan, result = _apply(db, project, book)

    assert not plan.site_flags
    assert not plan.field_changes
    assert result["site_creates"] == 0
    assert result["site_updates"] == 0
    assert result["changes"] == []
    assert set(db.scalars(select(MaintenanceSiteIssueLine.issue_line_id))) == set(identities)


def test_generated_replay_reports_stored_diffs_without_using_display_identity(db):
    project, part, _ = _seed_legacy_rows(db, 0)
    book = _book(db, project)
    sheet = book[master.V2_SHEET_SITE]
    sheet.append([None, date(2026, 10, 18), part.pn_std, "SYNTHETIC-AUTO", 2, "是", "原备注"])
    row_no = sheet.max_row
    _apply(db, project, book)
    issue = db.scalar(select(MaintenanceSiteIssue).where(
        MaintenanceSiteIssue.project_id == project.project_id))
    line = db.scalar(select(MaintenanceSiteIssueLine).where(
        MaintenanceSiteIssueLine.issue_id == issue.issue_id))
    identity = line.issue_line_id
    assert issue.issue_no.startswith("LY-AUTO-")
    assert issue.display_no.startswith("LY-20261018-")

    # D-0006：可读展示号已分配后的原样重放仍按 hash 身份成功。
    replay_plan = _validate(db, project, book)
    assert not replay_plan.site_flags
    edited = _book(db, project)
    edited_sheet, columns, edited_row = _entity_row(edited, identity)
    edited_sheet.cell(edited_row, columns["领用数量"], 3)
    _apply(db, project, edited)
    line.demand_order_no = "WBDD-20261019-9999"
    db.commit()

    with pytest.raises(master.WorkbookError) as raised:
        _validate(db, project, book)

    assert raised.value.code == "stale_site_replay"
    assert raised.value.details["rows"] == [{
        "row_no": row_no, "entity_id": identity, "issue_no": issue.display_no_effective,
        "pn": part.pn_std, "sn": "SYNTHETIC-AUTO",
        "differences": [
            {"field": "领用数量", "old": "3", "new": "2"},
            {"field": "关联需求单号", "old": "WBDD-20261019-9999", "new": ""},
        ],
    }]
    db.refresh(line)
    assert line.issue_line_id == identity
    assert line.quantity == Decimal("3")


def test_explicit_clear_address_is_reported_as_empty_uploaded_value(db):
    project, _, [(issue, line)] = _seed_legacy_rows(db, 1)
    book = _book(db, project)
    _append_replay(book, issue, line, **{"地址操作": "清空"})

    with pytest.raises(master.WorkbookError) as raised:
        _validate(db, project, book)

    assert raised.value.details["rows"][0]["differences"] == [
        {"field": "领用地址", "old": "合成领用地址-1", "new": ""},
    ]
    db.refresh(issue)
    assert issue.site_location == "合成领用地址-1"


def test_signed_entity_row_still_rebases_untouched_server_fields(db):
    project, _, [(issue, line)] = _seed_legacy_rows(db, 1)
    book = _book(db, project)
    sheet, columns, row_no = _entity_row(book, line.issue_line_id)
    line.quantity = Decimal("5")
    issue.issue_date = date(2026, 10, 17)
    operations.bump_workbook_revision(db, project_id=project.project_id)
    db.commit()
    sheet.cell(row_no, columns["备注"], "用户只改备注")

    plan, _ = _apply(db, project, book)

    assert not plan.conflicts
    assert {change["field"] for change in plan.field_changes} == {"备注"}
    db.refresh(issue)
    db.refresh(line)
    assert issue.issue_date == date(2026, 10, 17)
    assert line.quantity == Decimal("5")
    assert line.remark == "用户只改备注"


def test_signed_entity_row_keeps_three_way_conflicts_and_zero_write(db):
    project, _, [(_, line)] = _seed_legacy_rows(db, 1)
    book = _book(db, project)
    sheet, columns, row_no = _entity_row(book, line.issue_line_id)
    line.quantity = Decimal("5")
    operations.bump_workbook_revision(db, project_id=project.project_id)
    db.commit()
    sheet.cell(row_no, columns["领用数量"], 7)

    plan = _validate(db, project, book)

    assert len(plan.conflicts) == 1
    assert plan.conflicts[0]["field"] == "领用数量"
    assert plan.conflicts[0]["entity_id"] == line.issue_line_id
    assert plan.conflicts[0]["base"] == "2"
    assert plan.conflicts[0]["old"] == "5"
    assert plan.conflicts[0]["new"] == "7"
    with pytest.raises(master.WorkbookError) as raised:
        master.apply_project_master_v2(
            db, plan, operated_by="stale-replay-contract-test",
            import_batch_id=str(uuid.uuid4()))
    assert raised.value.code == "row_conflicts"
    db.rollback()
    db.refresh(line)
    assert line.quantity == Decimal("5")


def test_duplicate_identity_within_file_keeps_existing_error(db):
    project, _, [(issue, line)] = _seed_legacy_rows(db, 1)
    book = _book(db, project)
    first = _append_replay(book, issue, line, **{"领用日期": date(2026, 10, 19)})
    second = _append_replay(book, issue, line, **{"领用日期": date(2026, 10, 20)})

    with pytest.raises(master.WorkbookError) as raised:
        _validate(db, project, book)

    assert raised.value.code == "duplicate_site_row"
    assert f"第 {first} 行" in raised.value.message
    assert f"第 {second} 行" in raised.value.message


def test_collision_does_not_expose_values_from_another_project(db):
    project, _, [(issue, line)] = _seed_legacy_rows(db, 1)
    book = _book(db, project)
    sheet, _, row_no = _entity_row(book, line.issue_line_id)
    sheet.delete_rows(row_no)
    _append_replay(book, issue, line, **{"领用日期": date(2026, 10, 19)})
    other = MaintenanceProject(
        project_id=str(uuid.uuid4()), project_code="SYNTHETIC-OTHER-SCOPE",
        display_name="另一个不可见项目", lifecycle_status="ongoing")
    db.add(other)
    db.flush()
    # 即使库内存在与本项目推导ID相同、却挂在其它项目的脏身份，也先拒绝越界。
    issue.project_id = other.project_id
    issue.site_location = "其它项目不可见地址"
    line.remark = "其它项目不可见备注"
    db.commit()

    with pytest.raises(master.WorkbookError) as raised:
        _validate(db, project, book)

    assert raised.value.code == "project_mismatch"
    assert not raised.value.details
    assert "其它项目不可见" not in raised.value.message


@pytest.mark.parametrize("sheet_name", [master.V2_SHEET_USAGE, master.V2_SHEET_DICTIONARY])
def test_guidance_explains_existing_identity_and_new_issue_numbering(db, sheet_name):
    project, _, _ = _seed_legacy_rows(db, 0)
    book = _book(db, project)
    text = "\n".join(str(value) for row in book[sheet_name].iter_rows(values_only=True)
                     for value in row if value is not None)

    for instruction in ("WBDD+PN+SN", "保留实体ID", "原行", "领用单号留空", "自动编号"):
        assert instruction in text
    assert any(instruction in text for instruction in ("不要追加", "不得追加", "请勿追加", "切勿追加"))
