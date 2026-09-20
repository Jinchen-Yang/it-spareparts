"""Receipt display resolves real names without changing historical identity/SN facts."""

import io
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.models.maintenance_doc_import import MaintenanceRkdReturnLine
from app.models.system import SysUser
from app.services import maintenance_return_receipts as ledger
from tests.test_maintenance_return_receipts_api import _audit_rows, _client, _project
from tests.test_return_receipt_import import apply, preview, project, workbook


def test_list_resolves_current_name_without_rewriting_creator_or_audit(db):
    project = _project(db, project_id=str(uuid4()))
    username = "13800000001"
    client = _client(db, username=username)
    url = f"/api/maintenance/projects/stable/{project.project_id}/return-receipts"
    response = client.post(url, json={"pn": "DISPLAY-PN", "qty": 2,
                                    "serial_numbers": ["000007", "000008"],
                                    "evidence_ref": "ORIGINAL-VOUCHER"})
    assert response.status_code == 201, response.text
    receipt_id = response.json()["receipt_id"]
    listed = client.get(url)
    assert listed.status_code == 200
    item = listed.json()["items"][0]
    assert item["created_by"] == username
    assert item["created_by_name"] == "合成返还台账操作人"
    assert item["serial_numbers"] == ["000007", "000008"]
    before = _audit_rows(db, receipt_id)[0].after_json.copy()
    actor = db.query(SysUser).filter_by(username=username).one()
    actor.display_name = "  更正后的合成姓名  "
    actor.is_active = False
    db.commit()
    # Disabled accounts still identify historical registrants.
    item = ledger.search_receipts(db, project_id=project.project_id)["items"][0]
    assert item["created_by_name"] == "更正后的合成姓名"
    stored = db.get(MaintenanceRkdReturnLine, receipt_id)
    assert stored.created_by == username
    assert stored.evidence_ref == "ORIGINAL-VOUCHER"
    assert stored.version == 1
    assert _audit_rows(db, receipt_id)[0].after_json == before
    actor.display_name = " "
    db.commit()
    assert ledger.search_receipts(db, project_id=project.project_id)["items"][0]["created_by_name"] is None


@pytest.mark.parametrize("machine", [False, True])
def test_imported_sn_keeps_leading_zeroes_without_becoming_quantity_bound_serials(db, project, machine):
    data = load_workbook(io.BytesIO(workbook(machine=machine)))
    sheet = data.active
    column = [c.value for c in sheet[2]].index("整机SN" if machine else "备件明细.备件SN") + 1
    sheet.cell(3, column, "000009")
    output = io.BytesIO()
    data.save(output)
    apply(db, preview(db, output.getvalue()))
    row = ledger.search_receipts(db, project_id=project.project_id)["items"][0]
    assert row["source_serial_number"] == "000009"
    assert row["serial_numbers"] == []
    assert row["qty"] == ("1.000" if machine else "2.000")
    assert row["created_by_name"] is None
