"""Independent W1b probes: signed field boundaries and composed XLSX codecs.

These tests use synthetic, loaded objects and never apply a business mutation.
The existing boundary suite owns scalar/date/key-rotation compatibility vectors.
"""
import copy
import io
import json
from datetime import date
from decimal import Decimal
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from pydantic import SecretStr

from app.config import get_settings
from app.models.maintenance import FMaintenanceLine, FProjectExpense, MaintenanceManualCostOverride
from app.models.maintenance_project_operations import MaintenanceSiteIssue, MaintenanceSiteIssueLine
from app.services import maintenance_project_master_workbook as master
from app.services import maintenance_workbook_excel as excel
from app.services import maintenance_workbook_protocol as protocol
from app.services import maintenance_workbook_rows as rows
from app.services import maintenance_workbook_signatures as signatures
from app.services import maintenance_workbook_values as values
from app.services.maintenance_workbook_plans import WorkbookError


@pytest.fixture
def signing_key(monkeypatch):
    settings = get_settings()
    monkeypatch.setattr(settings, 'maintenance_manifest_active_key_id', 'independent-codecs')
    monkeypatch.setattr(settings, 'maintenance_manifest_active_hmac_key',
                        SecretStr('independent-w1b-synthetic-secret-1234567890'))
    monkeypatch.setattr(settings, 'maintenance_manifest_previous_hmac_keys_json', SecretStr('none'))


def _xlsx(wb):
    stream = io.BytesIO()
    wb.save(stream)
    return load_workbook(io.BytesIO(stream.getvalue()), data_only=True)


def _row_token():
    return signatures._v2_row_base_token(
        sheet=protocol.V2_SHEET_PARTS, entity_id=41,
        values={'PN': '甲', 'SN': '乙', '需求数量': Decimal('3.125'), '备注': '保留'})


def _global_token():
    return signatures._global_cost_base_token(
        line=FMaintenanceLine(id=41),
        override=MaintenanceManualCostOverride(version=9, active=True),
        display_unit_cost_ex_tax=Decimal('12.50'),
        display_unit_cost_inc_tax=Decimal('14.13'), visible_reason='独立凭据')


def _parse_row(token):
    return signatures._parse_v2_row_base_token(
        token, sheet=protocol.V2_SHEET_PARTS, entity_id=41, row_no=23)


def _parse_global(token):
    return signatures._parse_global_cost_base_token(token, line_id=41, row_no=23)


@pytest.mark.parametrize('edit', ['replace', 'swap', 'remove', 'insert', 'splice'])
def test_row_signature_binds_field_names_and_complete_base(signing_key, edit):
    envelope = json.loads(_row_token())
    base = envelope['payload']['base']
    if edit == 'replace':
        base['需求数量'] = '31.25'
    elif edit == 'swap':
        base['PN'], base['SN'] = base['SN'], base['PN']
    elif edit == 'remove':
        del base['备注']
    elif edit == 'insert':
        base['人工成本原因'] = '新增伪证据'
    else:
        another = signatures._v2_row_base_token(
            sheet=protocol.V2_SHEET_PARTS, entity_id=41,
            values={'PN': '另一份合法文件', 'SN': '乙', '需求数量': Decimal('3.125'), '备注': '保留'})
        envelope['payload']['base'] = json.loads(another)['payload']['base']
    with pytest.raises(WorkbookError, match='第 23 行') as caught:
        _parse_row(json.dumps(envelope, ensure_ascii=False))
    assert caught.value.code == 'invalid_concurrency_token'


@pytest.mark.parametrize(('field', 'replacement'), [
    ('line_id', 42), ('override_version', 10), ('override_active', False),
    ('unit_cost_ex_tax', '125.00'), ('unit_cost_inc_tax', '141.30'), ('reason', '替换凭据'),
])
def test_global_signature_binds_each_concurrency_and_cost_field(signing_key, field, replacement):
    envelope = json.loads(_global_token())
    envelope['payload'][field] = replacement
    with pytest.raises(WorkbookError, match='第 23 行') as caught:
        _parse_global(json.dumps(envelope, ensure_ascii=False))
    assert caught.value.code == 'invalid_concurrency_token'


@pytest.mark.parametrize(('make', 'parse'), [(_row_token, _parse_row), (_global_token, _parse_global)])
def test_token_json_presentation_is_not_part_of_signed_payload(signing_key, make, parse):
    token = make()
    envelope = json.loads(token)
    envelope['payload'] = dict(reversed(list(envelope['payload'].items())))
    # A real XLSX cell stores the reformatted JSON, including escaped Chinese.
    reformatted = json.dumps(dict(reversed(list(envelope.items()))), indent=2, ensure_ascii=True)
    wb = Workbook()
    wb.active['A1'] = reformatted
    assert parse(_xlsx(wb).active['A1'].value) == parse(token)


def _projected_row(kind):
    if kind == 'parts':
        order = SimpleNamespace(order_no='WBDD-41', order_date=date(2026, 7, 3),
                                linked_sales_order_no='XSDD-41', demand_type='维保',
                                warehouse='甲库', salesperson='人员甲', business_type='服务')
        line = FMaintenanceLine(id=41, part_id=42, pn_std='标准PN', pn_raw='原PN',
                                description='描述', qty=Decimal('3.125'), return_qty=Decimal('1.125'),
                                serial_numbers='SN-1', unit_cost_ex_tax=Decimal('12.50'),
                                unit_cost_inc_tax=Decimal('14.13'), cost_source='purchase',
                                confidence='high', edited_source='wbdd', line_note='备注')
        override = MaintenanceManualCostOverride(unit_cost_ex_tax=Decimal('15.00'), reason='人工凭据')
        return protocol.V2_SHEET_PARTS, rows._v2_part_row_values(line, order, override), protocol.V2_PART_BASE_FIELDS
    if kind == 'expenses':
        expense = FProjectExpense(raw_line_id='expense-41', bxd_no='BXD-41', line_no=4,
                                  expense_date=date(2026, 7, 3), person='人员乙', expense_type='交通',
                                  fee_category='出差', reason='服务', linked_sales_order_no='XSDD-41',
                                  data_status='approved', amount=Decimal('226.00'), tax_basis='inc',
                                  amount_ex_tax=Decimal('200.00'), amount_inc_tax=Decimal('226.00'), remark='备注')
        return protocol.V2_SHEET_EXPENSE, rows._v2_expense_row_values(expense, '指定项目'), protocol.V2_EXPENSE_BASE_FIELDS
    if kind == 'receipts':
        snapshot = SimpleNamespace(report_month=date(2026, 7, 1), cumulative_amount=Decimal('1200.25'),
                                   status='confirmed', receipt_reference='到账凭据', remark='备注', collection_id='receipt-41')
        return protocol.V2_SHEET_RECEIPTS, rows._v2_receipt_row_values(snapshot, '指定合同'), protocol.V2_RECEIPT_BASE_FIELDS
    issue = MaintenanceSiteIssue(issue_no='LY-AUTO-stable', display_no='LY-20260703-0001',
                                 issue_date=date(2026, 7, 3), site_location='甲地址')
    line = MaintenanceSiteIssueLine(issue_line_id='site-41', pn='PN-41', serial_number='SN-41',
                                    quantity=Decimal('3.125'), no_return=False, remark='备注', demand_order_no='WBDD-41')
    return protocol.V2_SHEET_SITE, rows._v2_site_row_values(line, issue), protocol.V2_SITE_BASE_FIELDS


@pytest.mark.parametrize('kind', ['parts', 'expenses', 'receipts', 'site'])
def test_projected_rows_survive_actual_xlsx_in_both_signed_base_and_sheet_hash(signing_key, kind):
    sheet, projected, base_fields = _projected_row(kind)
    original = copy.deepcopy(projected)
    entity_id = projected['实体ID']
    token = signatures._v2_row_base_token(
        sheet=sheet, entity_id=entity_id, values={name: projected[name] for name in base_fields})
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(list(projected))
    ws.append(list(projected.values()))
    hash_before = excel._v2_readonly_sheet_hash(ws)
    loaded = _xlsx(wb).active
    observed = dict(zip((cell.value for cell in loaded[1]), (cell.value for cell in loaded[2])))
    assert excel._v2_readonly_sheet_hash(loaded) == hash_before
    assert signatures._parse_v2_row_base_token(token, sheet=sheet, entity_id=entity_id, row_no=2) == {
        name: values._v2_hash_value(observed[name]) for name in base_fields}
    assert projected == original


@pytest.mark.parametrize('no_return', [True, False, None])
@pytest.mark.parametrize(('state', 'quantity', 'expected', 'message'), [
    ('required', '1.375', Decimal('1.375'), '应返；收货见项目返还台账'),
    ('exempt', '0', Decimal('0'), '免返；收货记录保留'),
    ('pending_category', 'not-yet-known', '—', '按规则判断（品类待确认）'),
])
def test_return_requirement_summary_wins_without_rewriting_manual_flag(no_return, state, quantity, expected, message):
    issue = MaintenanceSiteIssue(issue_no='LY-AUTO-stable', display_no='LY-20260703-0001',
                                 issue_date=date(2026, 7, 3), site_location='甲地址')
    line = MaintenanceSiteIssueLine(issue_line_id='site-41', pn='PN', quantity=Decimal('9.125'),
                                    no_return=no_return)
    details = {'return_requirement': {'requirement_status': state, 'required_quantity': quantity}}
    before = copy.deepcopy(details)
    projected = rows._v2_site_row_values(line, issue, details)
    assert projected['应返数量'] == expected
    assert projected['返还状态'] == message
    assert projected['是否应返还'] == {True: '否', False: '是', None: ''}[no_return]
    assert projected['领用数量'] == Decimal('9.125')
    identity = rows._v2_site_identity_values(line, issue, details)
    assert identity == {**projected, '领用单号': 'LY-AUTO-stable'}
    assert details == before and line.no_return is no_return


@pytest.mark.parametrize(('source', 'edited_source', 'missing', 'refill', 'provenance'), [
    (None, 'workbook_manual', 'out_of_scope', '是', '工作簿'),
    ('none', 'page_manual', 'none', '是', '页面人工登记'),
    ('purchase', 'wbdd', 'none', '否', 'WBDD'),
    ('sales_fallback', None, 'none', '否', 'WBDD'),
])
def test_partial_tax_evidence_does_not_invent_net_cost_or_replace_source(source, edited_source, missing, refill, provenance):
    order = SimpleNamespace(order_no='WBDD', order_date=None, linked_sales_order_no=None,
                            demand_type=None, warehouse=None, salesperson=None, business_type=None)
    line = FMaintenanceLine(id=41, qty=Decimal('2.125'), return_qty=Decimal('0.25'),
                            unit_cost_ex_tax=None, unit_cost_inc_tax=Decimal('11.30'),
                            cost_source=source, edited_source=edited_source)
    override = MaintenanceManualCostOverride(unit_cost_ex_tax=Decimal('50.00'), reason='人工证据')
    projected = rows._v2_part_row_values(line, order, override)
    assert (projected['成本缺失类型'], projected['可补价'], projected['来源']) == (missing, refill, provenance)
    assert projected['系统未税单位成本'] is None
    assert projected['系统含税单位成本'] == Decimal('11.30')
    assert projected['人工未税单位成本'] == Decimal('50.00')
    assert projected['需求数量'] == Decimal('2.125')
    assert projected['退货数量'] == Decimal('0.25')


class _NoDatabaseAccess:
    def get(self, *args, **kwargs):
        raise AssertionError('signature/readonly rejection must precede project lookup')


def _signed_workbook(marker):
    wb = Workbook()
    wb.active.title = protocol.V2_SHEET_PARTS
    ledger = wb.create_sheet(protocol.V2_SHEET_RETURN_LEDGER)
    ledger.append(['需求单', '实收数量', '备注'])
    ledger.append(['WBDD-A', 2, '甲'])
    ledger.append(['WBDD-B', 3, '乙'])
    meta = {
        'protocol_id': protocol.V2_PROTOCOL_ID, 'template_version': protocol.V2_TEMPLATE_VERSION,
        'project_id': 'project-' + marker, 'export_id': 'export-' + marker,
        'included_sheets': protocol.V2_SHEET_PARTS, 'workbook_revision': '4',
        'workbook_data_version': ('a' if marker == 'A' else 'b') * 64,
        'parts_row_ids': '41,42' if marker == 'A' else '43,44',
        # Exercise metadata signing plus numeric chunk merging through the actual verifier.
        'parts_base_hashes': (marker + 'c' * 63 + ',') * 470,
        'receipt_ledger_hash': excel._v2_readonly_sheet_hash(ledger),
    }
    ws = wb.create_sheet(protocol.V2_SHEET_META)
    for item in signatures._v2_sign_meta_rows(list(meta.items())):
        ws.append(item)
    return _xlsx(wb)


@pytest.mark.parametrize('field', [
    'project_id', 'export_id', 'workbook_data_version', 'parts_row_ids',
    'parts_base_hashes@2', 'metadata_hmac',
])
def test_splicing_from_other_authentic_workbook_is_rejected_before_context_lookup(signing_key, field):
    original, donor = _signed_workbook('A'), _signed_workbook('B')
    assert master._v2_verify_meta(_NoDatabaseAccess(), original, 'project-A')['export_id'] == 'export-A'
    assert master._v2_verify_meta(_NoDatabaseAccess(), donor, 'project-B')['export_id'] == 'export-B'
    replacement = next(row[1].value for row in donor[protocol.V2_SHEET_META] if row[0].value == field)
    target = next(row[1] for row in original[protocol.V2_SHEET_META] if row[0].value == field)
    assert target.value != replacement
    target.value = replacement
    with pytest.raises(WorkbookError) as caught:
        master._v2_verify_meta(_NoDatabaseAccess(), _xlsx(original), 'project-A')
    assert caught.value.code == 'template_signature_invalid'


@pytest.mark.parametrize('edit', ['change', 'swap_rows', 'insert_blank', 'remove_sheet'])
def test_signed_readonly_snapshot_ignores_style_but_detects_structural_or_value_edits(signing_key, edit):
    wb = _signed_workbook('A')
    ws = wb[protocol.V2_SHEET_RETURN_LEDGER]
    ws['B2'].number_format = '0.000'
    ws['A1'].font = Font(bold=True, color='FF0000')
    ws.row_dimensions[2].height = 28
    ws.column_dimensions['A'].width = 35
    wb = _xlsx(wb)
    assert master._v2_verify_meta(_NoDatabaseAccess(), wb, 'project-A')['project_id'] == 'project-A'
    ws = wb[protocol.V2_SHEET_RETURN_LEDGER]
    if edit == 'change':
        ws['B2'] = 5
    elif edit == 'swap_rows':
        first, second = list(ws.values)[1:3]
        for column, value in enumerate(second, 1):
            ws.cell(2, column, value)
        for column, value in enumerate(first, 1):
            ws.cell(3, column, value)
    elif edit == 'insert_blank':
        ws.insert_rows(3)
    else:
        del wb[protocol.V2_SHEET_RETURN_LEDGER]
    with pytest.raises(WorkbookError) as caught:
        master._v2_verify_meta(_NoDatabaseAccess(), _xlsx(wb), 'project-A')
    assert caught.value.code == 'readonly_field_changed'
