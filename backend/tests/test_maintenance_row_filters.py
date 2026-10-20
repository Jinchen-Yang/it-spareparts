"""D-33：通过真实 HTTP 比较筛选网页、完整导出与独立期望集。"""

import csv
import io
import json
from datetime import date
from decimal import Decimal
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.models.maintenance import FMaintenanceLine, FMaintenanceOrder, MaintenanceManualCostOverride
from app.models.maintenance_project import MaintenanceProjectUserAssignment
from app.models.system import SysUser
from tests.boss_board_helpers import PASSWORD, client_for, make_project
from tests.test_maintenance_demand_export_ordering import export_data  # noqa: F401


BASE = "/api/maintenance/spare-part-lines"
WINDOW = {"range": "custom", "from": "2026-07-01", "to": "2026-07-31"}


def test_global_multi_field_range_count_and_complete_exports(db, export_data):
    client = client_for(db, username="row-filter-admin", role="admin")
    filters = json.dumps([
        {"field": "pn_std", "op": "prefix", "value": "PN-same-"},
        {"field": "qty", "op": "range", "min": "1", "max": "1"},
        {"field": "order_date", "op": "range",
         "min": "2026-07-10", "max": "2026-07-10"},
    ])
    params = {**WINDOW, "filters": filters}
    response = client.get(BASE + "/rows", params={**params, "page": 2, "page_size": 4})
    assert response.status_code == 200, response.text
    assert response.json()["total"] == 15
    assert len(response.json()["rows"]) == 4
    actual = db.scalar(select(func.count()).select_from(FMaintenanceLine).join(
        FMaintenanceOrder, FMaintenanceOrder.id == FMaintenanceLine.order_id,
    ).where(
        FMaintenanceLine.pn_std.startswith("PN-same-"),
        FMaintenanceLine.qty == 1,
        FMaintenanceOrder.order_date == date(2026, 7, 10),
    ))
    assert actual == response.json()["total"]
    expected = [
        f"PN-{order}-{line}"
        for order in ["same-a", "same-a-copy", "same-b"]
        for line in ["one-a", "one-b", "two", "null-a", "null-b"]
    ]
    downloaded = client.get(BASE + "/export", params={**params, "format": "csv"})
    assert downloaded.status_code == 200, downloaded.text
    rows = list(csv.DictReader(io.StringIO(downloaded.content.decode("utf-8-sig"))))
    assert [row["PN"] for row in rows] == expected
    downloaded = client.get(BASE + "/export", params={**params, "format": "xlsx"})
    assert downloaded.status_code == 200, downloaded.text
    workbook = load_workbook(io.BytesIO(downloaded.content), data_only=True)
    try:
        values = list(workbook.active.values)
        pns = [row[values[0].index("PN")] for row in values[1:]]
        assert pns == expected
        assert "实体ID" not in values[0]
        assert not any("令牌" in str(name) for name in values[0])
    finally:
        workbook.close()


def _project_path(project):
    return f"/api/maintenance/projects/stable/{project.project_id}/master-workbook"


def test_registry_covers_business_union_and_separates_system_effective_manual(db, export_data):
    client = client_for(db, username="registry-admin", role="admin")
    response = client.get(BASE + "/fields")
    assert response.status_code == 200, response.text
    fields = {item["field"]: item for item in response.json()["fields"]}
    expected = {
        "project_name", "sales_order_no", "order_no", "order_date", "pn_std",
        "description", "qty", "return_qty", "serial_numbers", "warehouse",
        "demand_type", "business_type", "salesperson", "line_no", "line_note", "source",
        "returned_qty", "pending_return_qty", "purchase_qty", "purchased_qty",
        "pending_purchase_qty", "direct_ship_qty", "warehouse_need_qty",
        "warehouse_shipped_qty", "supplied_qty", "pending_supply_qty", "consumed_qty",
        "demand_pending_return_qty", "system_unit_cost_ex_tax", "system_unit_cost_inc_tax",
        "effective_unit_cost_ex_tax", "effective_unit_cost_inc_tax",
        "effective_cost_amount_inc_tax", "effective_cost_amount_ex_tax",
        "manual_unit_cost_ex_tax", "manual_reason", "cost_source", "confidence",
        "effective_cost_tier",
        "missing_kind", "can_refill", "system_cost_source", "system_confidence",
        "system_missing_kind", "system_can_refill",
    }
    assert expected <= fields.keys()
    assert fields["qty"]["operators"] == ["eq", "range", "is_empty"]
    assert fields["demand_type"]["operators"] == ["in", "is_empty"]
    assert not {"line_id", "part_id", "raw_line_id", "baseline_token", "readonly_hash"} & fields.keys()


def test_project_rows_filter_in_database_and_export_same_complete_set(db, export_data):
    project, _ = export_data
    client = client_for(db, username="project-filter-admin", role="admin")
    params = {"filters": json.dumps([{"field": "pn_std", "op": "prefix", "value": "PN-same-"}])}
    base = _project_path(project)
    response = client.get(base + "/rows", params={
        **params, "sheet": "03_备件订单", "page": 2, "page_size": 4,
    })
    assert response.status_code == 200, response.text
    assert response.json()["total"] == 15
    assert len(response.json()["rows"]) == 4
    for format in ("csv", "xlsx"):
        response = client.get(base + "/rows/export", params={**params, "format": format})
        assert response.status_code == 200, response.text
        if format == "csv":
            rows = list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))
            assert len(rows) == 15
        else:
            wb = load_workbook(io.BytesIO(response.content))
            assert wb.active.max_row == 16
            wb.close()


def _assign_user(db, username, project, responsibility):
    user = db.scalar(select(SysUser).where(SysUser.username == username))
    db.add(MaintenanceProjectUserAssignment(
        assignment_id=str(uuid4()), project_id=project.project_id,
        responsibility_type=responsibility, user_id=user.id,
        assigned_by="fixture", assignment_reason="合成权限证据",
    ))
    db.commit()


def test_amount_filters_require_explicit_authorized_project_before_count(db, export_data):
    project, _ = export_data
    client = client_for(db, username="row-manager", overrides={
        "page_maintenance": True, "data_profit": False, "data_purchase_cost": False,
    })
    _assign_user(db, "row-manager", project, "primary_manager")
    money = json.dumps([{"field": "system_unit_cost_ex_tax", "op": "is_empty"}])
    # 即使目前可见集合只有本人项目，也不能用匹配结果暗中确定金额范围。
    denied = client.get(BASE + "/rows", params={**WINDOW, "filters": money})
    assert denied.status_code == 403, denied.text
    assert "total" not in denied.json()
    registry = client.get(BASE + "/fields").json()
    assert not registry["amount_filter_allowed"]
    assert "system_unit_cost_ex_tax" not in {f["field"] for f in registry["fields"]}
    for path, params in [
        (BASE + "/rows", {**WINDOW, "project_id": project.project_id}),
        (_project_path(project) + "/rows", {"sheet": "03_备件订单"}),
    ]:
        allowed = client.get(path, params={**params, "filters": money})
        assert allowed.status_code == 200, allowed.text
        assert allowed.json()["total"] in (25, 30)
    viewer = client_for(db, username="row-viewer", overrides={
        "page_maintenance": True, "data_profit": False, "data_purchase_cost": False,
    })
    _assign_user(db, "row-viewer", project, "viewer")
    for suffix in ("/rows", "/rows/export"):
        params = {"sheet": "03_备件订单", "filters": money}
        denied = viewer.get(_project_path(project) + suffix, params=params)
        assert denied.status_code == 403, denied.text
        assert "total" not in denied.json()
    facts = viewer.get(_project_path(project) + "/rows", params={"sheet": "03_备件订单"})
    assert facts.status_code == 200, facts.text
    assert facts.json()["total"] == 30
    assert all(row.get("unit_cost_ex_tax") is None for row in facts.json()["rows"])


def test_legacy_effective_filter_matches_its_existing_manual_projection(db, export_data):
    project, _ = export_data
    line = db.scalar(select(FMaintenanceLine).where(
        FMaintenanceLine.pn_std == "PN-early-one-a",
    ))
    line.cost_source = "manual"
    line.cost_tax_basis = "ex"
    line.unit_cost = line.unit_cost_ex_tax = line.cost_amount = line.cost_amount_ex_tax = Decimal("10")
    line.unit_cost_inc_tax = line.cost_amount_inc_tax = Decimal("11.30")
    db.add(MaintenanceManualCostOverride(
        line_id=line.id, unit_cost_ex_tax=Decimal("20"), unit_cost_inc_tax=Decimal("22.60"),
        reason="旧行投影反例", active=True,
    ))
    db.commit()
    client = client_for(db, username="legacy-manual-projection", role="admin")
    params = {"project": project.display_name, "filters": json.dumps([
        {"field": "effective_unit_cost_ex_tax", "op": "eq", "value": "10"},
    ])}
    result = client.get("/api/maintenance/lines", params=params)
    assert result.status_code == 200, result.text
    assert result.json()["total"] == 1
    assert result.json()["rows"][0]["unit_cost_ex_tax"] == 10
    csv_result = client.get("/api/maintenance/lines/export", params=params)
    assert csv_result.status_code == 200, csv_result.text
    rows = list(csv.DictReader(io.StringIO(csv_result.content.decode("utf-8-sig"))))
    assert Decimal(rows[0]["未税单位成本"]) == 10


def test_half_cent_cost_filter_matches_project_python_evidence(db, export_data):
    project, _ = export_data
    line = db.scalar(select(FMaintenanceLine).where(
        FMaintenanceLine.pn_std == "PN-early-one-a",
    ))
    line.qty = Decimal("0.5")
    db.add(MaintenanceManualCostOverride(
        line_id=line.id, unit_cost_ex_tax=Decimal("0.01"), unit_cost_inc_tax=Decimal("0.01"),
        reason="半分舍入反例", active=True,
    ))
    db.commit()
    client = client_for(db, username="half-cent-admin", role="admin")
    filters = json.dumps([
        {"field": "pn_std", "op": "eq", "value": "PN-early-one-a"},
        {"field": "effective_cost_amount_inc_tax", "op": "eq", "value": "0.00"},
    ])
    for endpoint, params in (
        (BASE + "/rows", WINDOW),
        (_project_path(project) + "/rows", {"sheet": "03_备件订单"}),
    ):
        response = client.get(endpoint, params={**params, "filters": filters})
        assert response.status_code == 200, response.text
        assert response.json()["total"] == 1
        row = response.json()["rows"][0]
        amount = row.get("effective_cost_amount_inc_tax", row.get("cost_amount_inc_tax"))
        assert Decimal(str(amount)) == Decimal("0.00")


def test_mixed_project_amount_scope_and_foreign_project_are_rejected(db, export_data):
    project, _ = export_data
    viewer_project = make_project(db, "仅可看事实的第二项目")
    foreign = make_project(db, "完全不可见第三项目")
    client = client_for(db, username="mixed-row-manager", overrides={
        "page_maintenance": True, "data_profit": False, "data_purchase_cost": False,
    })
    _assign_user(db, "mixed-row-manager", project, "primary_manager")
    _assign_user(db, "mixed-row-manager", viewer_project, "viewer")
    money = {"field": "system_unit_cost_ex_tax", "op": "is_empty"}
    # 第二项目即使没有任何命中行，也不能据筛选后的结果放行混合金额查询。
    for filters in ([money], [
        {"field": "project_name", "op": "eq", "value": project.display_name}, money,
    ]):
        for suffix in ("/rows", "/export"):
            response = client.get(BASE + suffix, params={**WINDOW, "filters": json.dumps(filters)})
            assert response.status_code == 403
            assert not {"total", "rows"} & response.json().keys()
    for project_id in (viewer_project.project_id, foreign.project_id):
        response = client.get(BASE + "/rows", params={
            **WINDOW, "project_id": project_id, "filters": json.dumps([money]),
        })
        assert response.status_code == 403
    for suffix in ("/fields", "/rows", "/export"):
        response = client.get(BASE + suffix, params={**WINDOW, "project_id": foreign.project_id})
        assert response.status_code == 403
    for suffix in ("/row-fields", "/rows", "/rows/export"):
        response = client.get(_project_path(foreign) + suffix, params={"sheet": "03_备件订单"})
        assert response.status_code == 403


def test_canonical_sales_direct_project_preserves_scope_distinct_from_global_directory(db, export_data):
    project, orders = export_data
    project.salesperson = "直达项目销售"
    orders["early"].demand_type = "仅本人直达项目需求"
    client = client_for(db, username="direct-project-sales", overrides={
        "page_maintenance": True, "data_profit": True,
        "own_maintenance_projects_only": False,
    })
    user = db.scalar(select(SysUser).where(SysUser.username == "direct-project-sales"))
    user.salesperson_name = project.salesperson
    db.commit()
    login = client.post("/api/auth/login", json={
        "username": user.username, "password": PASSWORD,
    })
    assert login.status_code == 200
    client.headers["Authorization"] = f"Bearer {login.json()['token']}"
    base = _project_path(project)
    params = {"sheet": "03_备件订单", "page": 1, "page_size": 4}
    response = client.get(base + "/rows", params=params)
    assert response.status_code == 200, response.text
    assert response.json()["total"] == 30
    catalog = client.get(base + "/row-fields")
    assert catalog.status_code == 200, catalog.text
    field = next(f for f in catalog.json()["fields"] if f["field"] == "demand_type")
    assert "仅本人直达项目需求" in {o["value"] for o in field["options"]}
    filters = json.dumps([{"field": "pn_std", "op": "prefix", "value": "PN-same-"}])
    response = client.get(base + "/rows", params={**params, "filters": filters})
    assert response.json()["total"] == 15
    for format in ("csv", "xlsx"):
        response = client.get(base + "/rows/export", params={"format": format, "filters": filters})
        assert response.status_code == 200, response.text
        if format == "csv":
            assert len(list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))) == 15
        else:
            wb = load_workbook(io.BytesIO(response.content))
            assert wb.active.max_row == 16
            wb.close()
    # 全局仍沿原目录集合，不能因为直达项目范围修复而扩大全局导出。
    for query_params in (WINDOW, {**WINDOW, "project_id": project.project_id}):
        response = client.get(BASE + "/rows", params=query_params)
        assert response.status_code == 200, response.text
        assert response.json()["total"] == 0
        response = client.get(BASE + "/export", params={**query_params, "format": "csv"})
        assert response.status_code == 200, response.text
        assert list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig")))) == []


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_readonly_export_does_not_enter_roundtrip_protocol(db, export_data, format):
    project, _ = export_data
    client = client_for(db, username="no-roundtrip-admin", role="admin")
    response = client.get(_project_path(project) + "/rows/export", params={
        "format": format, "filters": json.dumps([{"field": "qty", "op": "eq", "value": 1}]),
    })
    assert response.status_code == 200
    rejected = client.post(
        _project_path(project) + "/validate",
        files={"file": (f"核对.{format}", response.content,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert rejected.status_code == (415 if format == "csv" else 422), rejected.text


def test_nested_filter_json_is_validation_error(db, export_data):
    client = client_for(db, username="nested-filter-admin", role="admin")
    response = client.get(BASE + "/rows", params={
        **WINDOW, "filters": "[" * 1100 + "]" * 1100,
    })
    assert response.status_code == 422


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_export_hard_limit_rejects_instead_of_truncating(db, export_data, monkeypatch, format):
    from app.services import maintenance_row_workbench
    monkeypatch.setattr(maintenance_row_workbench, "MAX_EXPORT_ROWS", 10, raising=False)
    client = client_for(db, username="export-limit-admin", role="admin")
    response = client.get(BASE + "/export", params={**WINDOW, "format": format})
    assert response.status_code == 413
    assert "10" in str(response.json()["detail"])


def test_xlsx_cell_limit_rejects_and_formula_text_stays_literal(db, export_data):
    line = db.scalar(select(FMaintenanceLine).where(FMaintenanceLine.pn_std == "PN-early-one-a"))
    line.line_note = "=1+1"
    db.commit()
    client = client_for(db, username="cell-limit-admin", role="admin")
    params = {**WINDOW, "format": "xlsx", "filters": json.dumps([
        {"field": "pn_std", "op": "eq", "value": line.pn_std},
    ])}
    response = client.get(BASE + "/export", params=params)
    wb = load_workbook(io.BytesIO(response.content), data_only=False)
    note_col = [c.value for c in wb.active[1]].index("备注") + 1
    assert wb.active.cell(2, note_col).value == "=1+1"
    assert wb.active.cell(2, note_col).data_type == "s"
    wb.close()
    line.line_note = "中" * 32768
    db.commit()
    response = client.get(BASE + "/export", params=params)
    assert response.status_code == 413


def test_xlsx_uses_existing_control_character_sanitization(db, export_data):
    line = db.scalar(select(FMaintenanceLine).where(FMaintenanceLine.pn_std == "PN-early-one-a"))
    line.line_note = "前\x01后"
    db.commit()
    client = client_for(db, username="xlsx-controls-admin", role="admin")
    response = client.get(BASE + "/export", params={**WINDOW, "format": "xlsx"})
    assert response.status_code == 200, response.text
    wb = load_workbook(io.BytesIO(response.content))
    header = [c.value for c in wb.active[1]]
    rows = list(wb.active.values)[1:]
    row = next(r for r in rows if r[header.index("PN")] == "PN-early-one-a")
    assert row[header.index("备注")] == "前后"
    wb.close()


def test_v1_cost_source_filter_and_export_match_existing_web_projection(db, export_data, monkeypatch):
    from app.api import maintenance_project_master_workbook as master_api
    monkeypatch.setattr(master_api.get_settings(), "maintenance_project_master_v2_enabled", False)
    project, _ = export_data
    client = client_for(db, username="v1-source-admin", role="admin")
    params = {"filters": json.dumps([{"field": "cost_source", "op": "is_empty"}])}
    response = client.get(_project_path(project) + "/rows", params={
        **params, "sheet": "03_备件订单", "page": 1, "page_size": 4,
    })
    assert response.status_code == 200, response.text
    assert response.json()["total"] == 30
    assert {row["cost_source"] for row in response.json()["rows"]} == {""}
    response = client.get(_project_path(project) + "/rows/export", params={**params, "format": "csv"})
    assert response.status_code == 200, response.text
    rows = list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert len(rows) == 30
    assert {row["有效成本来源"] for row in rows} == {""}


def test_enum_options_include_values_from_authorized_scope(db, export_data):
    project, orders = export_data
    orders["early"].demand_type = "新增需求类型"
    db.commit()
    client = client_for(db, username="enum-scope-viewer", overrides={
        "page_maintenance": True, "data_profit": False, "data_purchase_cost": False,
    })
    _assign_user(db, "enum-scope-viewer", project, "viewer")
    response = client.get(BASE + "/fields", params={"project_id": project.project_id})
    assert response.status_code == 200, response.text
    fields = {f["field"]: f for f in response.json()["fields"]}
    assert "新增需求类型" in {option["value"] for option in fields["demand_type"]["options"]}
    assert "confidence" not in fields


@pytest.mark.parametrize("value", ["1e999999999", "1e-999999999", "99e131071"])
def test_extreme_numeric_exponents_are_validation_errors(db, export_data, value):
    client = client_for(db, username="extreme-number-admin", role="admin")
    response = client.get(BASE + "/rows", params={**WINDOW, "filters": json.dumps([
        {"field": "qty", "op": "eq", "value": value},
    ])})
    assert response.status_code == 422


def test_numeric_range_bound_is_not_a_business_amount_write(db, export_data):
    client = client_for(db, username="numeric-bound-admin", role="admin")
    filters = json.dumps([
        {"field": "pn_std", "op": "prefix", "value": "PN-same-"},
        {"field": "qty", "op": "range", "min": "-1e12", "max": "1e12"},
    ])
    response = client.get(BASE + "/rows", params={**WINDOW, "filters": filters})
    assert response.status_code == 200, response.text
    assert response.json()["total"] == 15
    response = client.get(BASE + "/export", params={
        **WINDOW, "filters": filters, "format": "csv",
    })
    assert response.status_code == 200, response.text
    assert len(list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))) == 15


@pytest.mark.parametrize("value", ["a\x00b", "\ud800"])
def test_unrepresentable_database_text_is_validation_error(db, export_data, value):
    client = client_for(db, username="invalid-text-admin", role="admin")
    response = client.get(BASE + "/rows", params={**WINDOW, "filters": json.dumps([
        {"field": "description", "op": "contains", "value": value},
    ])})
    assert response.status_code == 422
