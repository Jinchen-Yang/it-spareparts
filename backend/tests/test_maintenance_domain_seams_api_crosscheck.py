"""Independent post-suite R1/R2 counterexamples at real HTTP/Excel boundaries.

A subtotal is not a total; a workbook owner exemption is not an HTTP entitlement.
No production reader, permission function, or reference result is mocked.
"""
import io
from contextlib import contextmanager
from datetime import UTC, datetime
from decimal import Decimal
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries
from sqlalchemy import event, select

from app.config import get_settings
from app.db import engine
from app.models.maintenance_project import MaintenanceProjectUserAssignment
from app.models.maintenance_project_operations import (
    MaintenanceProjectOperationAudit, MaintenanceSiteIssueCommand, MaintenanceSiteIssueReturnEvent,
)
from app.models.system import SysUser
from app.services import maintenance_project_master_workbook as master
from app.services import maintenance_project_operations as operations
from app.services.maintenance_contract_reference import current_contract_references
from app.services.maintenance_manager_workbook_v3 import OVERVIEW_SHEET, OVERVIEW_TABLE
from app.services.maintenance_site_issue_representation import site_issue_line_dict
from tests.boss_board_helpers import assign, client_for, import_wbdd, make_project
from tests.test_maintenance_contract_reference import _contract
from tests.test_maintenance_site_issue_facts import _demand_prices, _facts
from tests.test_site_issue_line_correction_api import _body, _url


@pytest.fixture(autouse=True)
def _flags(monkeypatch):
    monkeypatch.setattr(get_settings(), "maintenance_boss_dashboard_enabled", True)
    monkeypatch.setattr(get_settings(), "maintenance_project_master_v2_enabled", False)


def _assign_manager(db, project, username, *, responsibility="primary_manager"):
    user = db.scalar(select(SysUser).where(SysUser.username == username))
    db.add(MaintenanceProjectUserAssignment(
        assignment_id=str(uuid4()), project_id=project.project_id,
        responsibility_type=responsibility, user_id=user.id,
        assigned_at=datetime.now(UTC), assigned_by="independent-crosscheck",
        assignment_reason="合成接口权限反例",
    ))
    db.commit()


@contextmanager
def _book(response):
    assert response.status_code == 200, response.text
    assert "spreadsheetml" in response.headers["content-type"]
    book = load_workbook(io.BytesIO(response.content), data_only=False)
    try:
        yield book
    finally:
        book.close()


def _metric(book, sheet_name, label):
    rows = [row for row in book[sheet_name].iter_rows(values_only=True) if row[0] == label]
    assert len(rows) == 1, (sheet_name, label)
    return rows[0][1]


def _manager_rows(book):
    sheet = book[OVERVIEW_SHEET]
    left, top, right, bottom = range_boundaries(sheet.tables[OVERVIEW_TABLE].ref)
    headers = [sheet.cell(top, col).value for col in range(left, right + 1)]
    return [dict(zip(headers, [sheet.cell(row, col).value for col in range(left, right + 1)]))
            for row in range(top + 1, bottom + 1)]


@pytest.mark.parametrize("case", ["partial", "shared_off_scope", "complete_zero"])
def test_same_contract_facts_keep_five_consumer_projections(db, tmp_path, case):
    """The card may publish partial(value); totals elsewhere remain blank."""
    project = make_project(db, "独立合同消费者")
    other = make_project(db, "不可混入经理月表的项目")
    assign(db, import_wbdd(db, tmp_path, project=project.display_name)[0], project)
    amount = "0.00" if case == "complete_zero" else "12345.67"
    _contract(db, project.project_id, amount=amount, contract_id="stable-contract")
    if case == "partial":
        _contract(db, project.project_id, amount=None, contract_no="XSDD-MISSING")
    elif case == "shared_off_scope":
        _contract(db, other.project_id, amount=amount, contract_id="stable-contract")
    db.commit()
    client = client_for(db, username="crosscheck-admin", role="admin")
    _assign_manager(db, project, "crosscheck-admin")
    raw = current_contract_references(db, [project.project_id])[project.project_id]
    assert raw["amount_inc_tax"] == (None if case == "shared_off_scope" else Decimal(amount))
    assert raw["contract_incomplete"] is (case != "complete_zero")

    response = client.get("/api/maintenance/boss-board/projects")
    assert response.status_code == 200, response.text
    card = next(row for row in response.json()["rows"] if row["project_id"] == project.project_id)
    assert card["contract_shared"] is (case == "shared_off_scope")
    assert card["contract_amount_inc_tax"]["state"] == ("ready" if case == "complete_zero" else "partial")
    assert card["contract_amount_inc_tax"]["value"] == (None if raw["amount_inc_tax"] is None else amount)
    response = client.get("/api/maintenance/projects")
    assert response.status_code == 200, response.text
    row = next(row for row in response.json()["rows"] if row["project"] == project.display_name)
    expected_total = 0 if case == "complete_zero" else None
    assert row["contract_amount"] == row["contract_amount_inc_tax"] == expected_total
    assert row["contract_incomplete"] is (case != "complete_zero")
    base = f"/api/maintenance/projects/stable/{project.project_id}"
    with _book(client.get(base + "/master-workbook.xlsx", params={"sheets": master.SHEET_OVERVIEW})) as book:
        assert _metric(book, master.SHEET_OVERVIEW, "合同总额(含税)") == expected_total
    with _book(client.get(base + "/workbook-v3.xlsx")) as book:
        assert _metric(book, "02_概览数据", "合同总额(含税)") == expected_total
    with _book(client.get("/api/maintenance/project-manager/workbooks/v3", params={"report_month": "2026-10"})) as book:
        rows = _manager_rows(book)
        assert [row["项目ID"] for row in rows] == [project.project_id]
        assert rows[0]["全部合同额（含税）"] == expected_total
        assert other.project_id not in str(list(book[OVERVIEW_SHEET].values))


@pytest.mark.parametrize("responsibility", ["primary_manager", "viewer"])
def test_internal_snapshot_does_not_define_http_or_excel_entitlement(db, responsibility):
    """D-03 owner exemption belongs to master, not every fact consumer."""
    project, issue, (line, _), _ = _facts(db)
    line.manual_unit_cost = Decimal("731.29")
    line.manual_unit_cost_inc_tax = Decimal("826.36")
    line.unit_cost = line.unit_cost_ex_tax = Decimal("731.29")
    line.unit_cost_inc_tax = Decimal("826.36")
    line.cost_amount = line.cost_amount_ex_tax = Decimal("1462.58")
    line.cost_amount_inc_tax = Decimal("1652.72")
    line.cost_source = "manual"
    line.manual_evidence = "PRIVATE-INVOICE-CROSSCHECK"
    line.reference_sample_ids = ["PRIVATE-SAMPLE-CROSSCHECK"]
    line.reference_samples = [{"document_no": "PRIVATE-INVOICE-CROSSCHECK", "unit_price": "731.29"}]
    line.reference_sample_count = 1
    db.commit()
    before = site_issue_line_dict(line)
    _contract(db, project.project_id, amount="64532.10")
    db.commit()
    restricted = client_for(db, username="crosscheck-restricted", overrides={
        "page_maintenance": True, "data_purchase_cost": False, "data_profit": False,
    })
    _assign_manager(db, project, "crosscheck-restricted", responsibility=responsibility)
    admin = client_for(db, username="crosscheck-reader", role="admin")
    path = "/api/maintenance/site-issues/search"
    allowed = admin.post(path, json={"project_id": project.project_id})
    assert allowed.status_code == 200, allowed.text
    row = next(row for row in allowed.json()["rows"] if row["issue_id"] == issue.issue_id)["lines"][0]
    for key in ("manual_unit_cost", "manual_evidence", "reference_samples", "reference_sample_ids"):
        assert row[key] == before[key]
    responses = [
        restricted.get(f"/api/maintenance/projects/stable/{project.project_id}/workbook-v3.xlsx"),
        restricted.post(path, json={"project_id": project.project_id}),
        restricted.get("/api/maintenance/project-manager/workbooks/v3", params={"report_month": "2026-10"}),
    ]
    for response in responses:
        assert response.status_code == 403, response.text
        assert "PRIVATE-" not in response.text and "731.29" not in response.text
    response = restricted.get(f"/api/maintenance/projects/stable/{project.project_id}/master-workbook.xlsx",
                              params={"sheets": f"{master.SHEET_SITE},{master.SHEET_OVERVIEW}"})
    assert response.status_code == (200 if responsibility == "primary_manager" else 403), response.text
    if responsibility == "primary_manager":
        with _book(response) as book:
            assert line.pn in str(list(book[master.SHEET_SITE].values))
            assert _metric(book, master.SHEET_OVERVIEW, "合同总额(含税)") == 64532.1
    db.refresh(line)
    assert site_issue_line_dict(line) == before


@pytest.mark.parametrize("missing", ["data_profit", "data_purchase_cost"])
@pytest.mark.parametrize("selector", [
    {"sort": "cost_ratio"}, {"card_status": "warning"}, {"lifecycle": "payment_complete"},
])
def test_money_selector_rejected_before_reference_count_and_pagination(db, missing, selector):
    """A 422 after executing an unauthorized aggregate is not sufficient."""
    project = make_project(db, "independent-money-filter")
    _contract(db, project.project_id, amount="87654.32")
    db.commit()
    client = client_for(db, username="crosscheck-no-money", overrides={
        "page_maintenance": True, "data_profit": missing != "data_profit",
        "data_purchase_cost": missing != "data_purchase_cost",
    })
    statements = []

    def capture(_conn, _cursor, statement, _parameters, _context, _executemany):
        statements.append(" ".join(statement.lower().split()))

    event.listen(engine, "before_cursor_execute", capture)
    try:
        responses = [
            client.get("/api/maintenance/boss-board/projects", params={**selector, "page": 2, "page_size": 1}),
            client.post("/api/maintenance/boss-board/projects/search", json={**selector, "q": "independent", "page": 2, "page_size": 1}),
        ]
    finally:
        event.remove(engine, "before_cursor_execute", capture)
    for response in responses:
        assert response.status_code == 422, response.text
        assert response.json()["detail"]["code"] == "cost_contract_permission_required"
        assert "total" not in response.json() and "rows" not in response.json()
    assert not [sql for sql in statements if sql.startswith("select") and (
        "from maintenance_project " in sql or "from maintenance_project_contract " in sql
    )], statements


def test_same_login_rechecks_cross_project_scope_and_revocation(db):
    project, issue, _, _ = _facts(db)
    other, foreign_issue, _, _ = _facts(db)
    _contract(db, other.project_id, amount="99887.66")
    db.commit()
    client = client_for(db, username="crosscheck-scoped", overrides={
        "page_maintenance": True, "data_profit": True, "data_purchase_cost": True,
    })
    _assign_manager(db, project, "crosscheck-scoped")
    allowed = client.post("/api/maintenance/site-issues/search", json={"project_id": project.project_id})
    assert allowed.status_code == 200, allowed.text
    assert [row["issue_id"] for row in allowed.json()["rows"]] == [issue.issue_id]
    for suffix in ("master-workbook.xlsx", "workbook-v3.xlsx"):
        denied = client.get(f"/api/maintenance/projects/stable/{other.project_id}/{suffix}")
        assert denied.status_code == 403, denied.text
        assert "99887.66" not in denied.text
    denied = client.post("/api/maintenance/site-issues/search", json={"project_id": other.project_id})
    assert denied.status_code == 403, denied.text
    assert foreign_issue.issue_id not in denied.text
    assignment = db.scalar(select(MaintenanceProjectUserAssignment).where(
        MaintenanceProjectUserAssignment.project_id == project.project_id))
    assignment.archived_at = datetime.now(UTC)
    assignment.archived_by = "crosscheck"
    assignment.archive_reason = "撤销本人项目的合成权限反例"
    assignment.version += 1
    db.commit()
    # The same bearer token must not retain the now-revoked scope.
    assert client.post("/api/maintenance/site-issues/search", json={"project_id": project.project_id}).status_code == 403


def test_http_demand_mismatch_is_atomic_then_same_key_can_succeed(db):
    """Fact-layer rejection must not reserve a replay key or create audit."""
    project, issue, (line, _), _ = _facts(db)
    foreign, _, _, foreign_part = _facts(db)
    _demand_prices(db, project=foreign, part=foreign_part, prices=(("FOREIGN-CROSSCHECK", Decimal("29.00")),))
    client = client_for(db, username="crosscheck-editor", role="admin")
    before = site_issue_line_dict(line)
    key = f"crosscheck-error-{uuid4()}"
    rejected = client.patch(_url(issue, line), json=_body(
        key=key, changes={"remark": "must be atomic", "demand_order_no": "FOREIGN-CROSSCHECK"},
    ))
    assert rejected.status_code == 409, rejected.text
    assert rejected.json()["detail"] == {
        "code": "site_issue_edit_demand_project_mismatch",
        "message": "关联需求单不属于当前项目，请核对单号或留空",
    }
    db.refresh(line)
    db.refresh(issue)
    assert site_issue_line_dict(line) == before and issue.version == 1
    for model in (MaintenanceProjectOperationAudit, MaintenanceSiteIssueCommand, MaintenanceSiteIssueReturnEvent):
        assert db.query(model).count() == 0
    assert operations.get_or_create_workbook_state(db, project_id=project.project_id).revision == 0
    accepted = client.patch(_url(issue, line), json=_body(key=key, changes={"remark": "corrected after rejection"}))
    assert accepted.status_code == 200, accepted.text
    assert accepted.json()["changed"] is True and accepted.json()["idempotent_replay"] is False
