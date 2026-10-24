"""Real module grants, stable owners and complete replenishment transactions."""
import json
from io import BytesIO

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.api import replenishment as api
from app.authz import runtime
from app.config import get_settings
from app.db import get_db
from app.models.authorization import SysAuthorizationArtifact, SysAuthorizationGrant
from app.models.dimensions import DimPart
from app.models.replenishment import ReplenishmentApplication as Application, ReplenishmentApplicationVersion as Version
from tests.test_authz_s3_checklist import setup, cerbos_url  # noqa: F401

RESOURCES = ['replenishment_capabilities', 'replenishment_catalog', 'replenishment_project_directory',
             'replenishment_project_commands', 'replenishment_request']


@pytest.fixture
def replenishment(db, setup, monkeypatch):
    uid, identity, pid, _, headers = setup
    settings = get_settings()
    monkeypatch.setattr(settings, 'replenishment_beta_enabled', True)
    monkeypatch.setattr(settings, 'replenishment_auto_review_enabled', False)
    monkeypatch.setattr(settings, 'authz_enabled_resources', RESOURCES)
    runtime.activate_configured_resources(db.get_bind())
    part = DimPart(pn_std='AUTHZ-REPLENISHMENT-PART', description='synthetic', status='active')
    db.add(part)
    db.commit()
    app = FastAPI()
    app.include_router(api.router, prefix='/api')
    app.dependency_overrides[get_db] = lambda: db
    client = TestClient(app)
    client.headers.update(headers)
    return uid, pid, part.id, client


def grant(db, uid, resource, actions, scope=None):
    if scope is None:
        scope = 'all_projects' if resource in ('replenishment_project_commands', 'replenishment_project_directory') else 'all'
    row = db.scalar(select(SysAuthorizationGrant).where(SysAuthorizationGrant.user_id == uid, SysAuthorizationGrant.resource == resource))
    if row is None:
        row = SysAuthorizationGrant(user_id=uid, resource=resource, actions=actions, scope=scope,
            field_permissions=[] if resource == 'replenishment_capabilities' else ['data_pool_price_governance'], version=1)
        db.add(row)
    else:
        row.actions = actions
        row.scope = scope
        row.version += 1
    db.commit()
    return row


def submit(client, pid, part_id, key='replenishment-command-1'):
    return client.post('/api/replenishment-beta/applications', json={'client_request_id': key,
        'project_id': pid, 'request_note': 'synthetic', 'lines': [{'part_id': part_id, 'quantity': 2}]})


def test_capabilities_is_separate_and_does_not_grant_business_reads(db, replenishment):
    uid, pid, part, client = replenishment
    grant(db, uid, 'replenishment_capabilities', ['read'])
    response = client.get('/api/replenishment-beta/capabilities')
    assert response.status_code == 200, response.text
    assert response.json()['enabled'] is True
    assert response.json()['can_create'] is False
    assert response.headers['cache-control'] == 'no-store'
    for path in ('catalog', 'projects', 'applications', f'cart-drafts/{pid}'):
        response = client.get('/api/replenishment-beta/' + path)
        assert response.status_code == 403, (path, response.text)
    assert submit(client, pid, part).status_code == 403
    assert db.scalar(select(func.count()).select_from(Application)) == 0


def test_submit_only_stable_version_receipt_and_independent_read_export(db, replenishment):
    uid, pid, part, client = replenishment
    grant(db, uid, 'replenishment_project_commands', ['submit'])
    response = submit(client, pid, part)
    assert response.status_code == 201, response.text
    aid = response.json()['application_id']
    assert db.get(Application, aid).owner_user_id == uid
    version = db.scalar(select(Version).where(Version.application_id == aid))
    assert version.owner_user_id == uid
    assert db.get(SysAuthorizationArtifact, ('replenishment_version', version.version_id)).owner_user_id == uid
    assert submit(client, pid, part).json()['idempotent'] is True
    url = '/api/replenishment-beta/applications/' + aid
    assert client.get(url).status_code == 403
    grant(db, uid, 'replenishment_request', ['export_screening'], scope='own_applications')
    assert client.get(url).status_code == 403
    response = client.get(url + '/exports/system-screening.xlsx')
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert workbook.active.max_row >= 2
    grant(db, uid, 'replenishment_request', ['read'], scope='own_applications')
    response = client.get(url)
    assert response.status_code == 200, response.text
    assert response.json()['versions'][0]['lines'][0]['pn_std'] == 'AUTHZ-REPLENISHMENT-PART'
    assert client.get(url + '/exports/system-screening.xlsx').status_code == 403
    grant(db, uid, 'replenishment_project_commands', [])
    assert submit(client, pid, part).status_code == 403
    assert db.scalar(select(func.count()).select_from(Application)) == 1


def test_cart_edit_and_submit_are_independent_and_use_actual_saved_version(db, replenishment):
    uid, pid, part, client = replenishment
    grant(db, uid, 'replenishment_project_commands', ['cart_read', 'cart_update'])
    path = f'/api/replenishment-beta/cart-drafts/{pid}'
    response = client.get(path)
    assert response.status_code == 200, response.text
    assert response.json()['draft'] is None
    response = client.put(path, json={'expected_version': None, 'lines': [{'part_id': part, 'quantity': 3}]})
    assert response.status_code == 200, response.text
    draft = response.json()['draft']
    assert client.post(path + '/submit', json={'expected_version': draft['version'], 'client_request_id': draft['client_request_id']}).status_code == 403
    grant(db, uid, 'replenishment_project_commands', ['submit'])
    payload = {'expected_version': draft['version'], 'client_request_id': draft['client_request_id']}
    response = client.post(path + '/submit', json=payload)
    assert response.status_code == 201, response.text
    retry = client.post(path + '/submit', json=payload)
    assert retry.status_code == 201, retry.text
    assert retry.json()['idempotent'] is True
    assert client.get(path).status_code == 403
    assert client.delete(path).status_code == 403


def test_catalog_view_binds_actual_query_and_readonly_data(db, replenishment):
    uid, pid, part, client = replenishment
    grant(db, uid, 'replenishment_catalog', ['read'])
    path = '/api/replenishment-beta/catalog'
    response = client.get(path, params={'q': 'AUTHZ', 'view': 'authorization'})
    assert response.status_code == 200, response.text
    ctx = response.json()['authorization_context']
    pre = {key: ctx[key] for key in ('catalog_version', 'authorization_epoch', 'context')}
    response = client.get(path, params={'q': 'OTHER', 'authorization_precondition': json.dumps(pre)})
    assert response.status_code == 409, response.text
    response = client.get(path, params={'q': 'AUTHZ', 'authorization_precondition': json.dumps(pre)})
    assert response.status_code == 200, response.text
    assert response.json()['total'] == 1


@pytest.mark.parametrize('mode', ['lines', 'resolutions'])
def test_both_revision_modes_keep_stable_version_owner_and_replay_current_permissions(db, replenishment, monkeypatch, mode):
    uid, pid, part, client = replenishment
    monkeypatch.setattr(get_settings(), 'replenishment_auto_review_enabled', True)
    grant(db, uid, 'replenishment_project_commands', ['submit', 'revise'])
    created = submit(client, pid, part)
    assert created.status_code == 201, created.text
    assert created.json()['status'] == 'needs_revision'
    aid = created.json()['application_id']
    body = {'expected_application_version': created.json()['version'], 'client_request_id': 'synthetic-revision-request'}
    body[mode] = [{'part_id': part, 'quantity': 4}] if mode == 'lines' else [{
        'request_line_id': created.json()['versions'][0]['lines'][0]['request_line_id'],
        'action': 'replace', 'part_id': part, 'quantity': 4}]
    path = f'/api/replenishment-beta/applications/{aid}/revisions'
    assert client.post(path, json=body).status_code == 403  # A project action never implies application ownership.
    grant(db, uid, 'replenishment_request', ['revise'], scope='own_applications')
    response = client.post(path, json=body)
    assert response.status_code == 200, response.text
    assert response.json()['latest_version_no'] == 2
    replay = client.post(path, json=body)
    assert replay.status_code == 200, replay.text
    assert replay.json()['idempotent'] is True
    assert all(version.owner_user_id == uid for version in db.scalars(select(Version)))
    assert db.scalar(select(func.count()).select_from(SysAuthorizationArtifact)) == 2
    grant(db, uid, 'replenishment_request', [], scope='own_applications')
    assert client.post(path, json=body).status_code == 403
    assert db.scalar(select(func.count()).select_from(Version)) == 2


def test_real_part_lock_wait_then_epoch_revocation_has_no_application_version_or_receipt(db, replenishment):
    from concurrent.futures import ThreadPoolExecutor
    from types import MappingProxyType
    import time
    from sqlalchemy import text
    from sqlalchemy.orm import Session
    from fastapi import HTTPException
    from app.authz.principal import VerifiedIdentity
    from app.models.system import SysUser
    from app.security import UserContext
    from app.services import replenishment_authorization as protected
    uid, pid, part, _ = replenishment
    grant(db, uid, 'replenishment_project_commands', ['submit'])
    owner = db.get(SysUser, uid)
    identity = VerifiedIdentity(owner.username, owner.token_version, False, 'sys_user', MappingProxyType({}), uid)
    legacy = UserContext(user_id=owner.username, role=owner.role, salesperson_name=None, permissions={}, is_authenticated=True)
    db.rollback()
    def create():
        with Session(db.get_bind()) as worker:
            try:
                protected.command(worker, identity, entry_id=protected.CREATE,
                    params={'project_id': pid, 'client_request_id': 'actual-lock-request', 'lines': [{'part_id': part, 'quantity': 2}]},
                    legacy_context=legacy)
                return 201
            except HTTPException as exc:
                return exc.status_code
    with Session(db.get_bind()) as blocker:
        blocker.scalar(select(DimPart).where(DimPart.id == part).with_for_update())
        blocker_pid = blocker.scalar(text('SELECT pg_backend_pid()'))
        with ThreadPoolExecutor(max_workers=1) as pool:
            future = pool.submit(create)
            blocked = False
            deadline = time.monotonic() + 10
            try:
                while time.monotonic() < deadline and not future.done():
                    with db.get_bind().connect() as observer:
                        blocked = observer.scalar(text('SELECT EXISTS(SELECT 1 FROM pg_stat_activity WHERE :pid = ANY(pg_blocking_pids(pid)))'), {'pid': blocker_pid})
                    if blocked: break
                    time.sleep(.02)
                assert blocked, 'Did not reach actual part row lock'
                with Session(db.get_bind()) as administrator:
                    user = administrator.get(SysUser, uid)
                    user.token_version += 1
                    administrator.commit()
            finally:
                blocker.rollback()
            assert future.result(timeout=10) == 401
    assert db.scalar(select(func.count()).select_from(Application)) == 0
    assert db.scalar(select(func.count()).select_from(Version)) == 0
    assert db.scalar(select(func.count()).select_from(SysAuthorizationArtifact)) == 0


def test_own_list_sql_scope_and_all_reader_do_not_depend_on_project_write_scope(db, replenishment):
    from app.auth import _make_token
    from app.models.system import SysUser
    uid, pid, part, client = replenishment
    grant(db, uid, 'replenishment_project_commands', ['submit'])
    own = submit(client, pid, part).json()['application_id']
    other = SysUser(username='other-replenishment-owner', role='admin', password_hash='synthetic')
    db.add(other); db.commit()
    grant(db, other.id, 'replenishment_project_commands', ['submit'])
    token, _ = _make_token(other.role, other.username, None, token_version=other.token_version, authn='sys_user')
    original = client.headers['Authorization']
    client.headers['Authorization'] = 'Bearer ' + token
    response = submit(client, pid, part, key='other-owner-application')
    assert response.status_code == 201, response.text
    foreign = response.json()['application_id']
    client.headers['Authorization'] = original
    grant(db, uid, 'replenishment_request', ['read'], scope='own_applications')
    response = client.get('/api/replenishment-beta/applications', params={'page_size': 1})
    assert response.status_code == 200, response.text
    assert response.json()['total'] == 1
    assert response.json()['items'][0]['application_id'] == own
    assert client.get('/api/replenishment-beta/applications/' + foreign).status_code == 403
    grant(db, uid, 'replenishment_project_commands', [])
    grant(db, uid, 'replenishment_request', ['read'])
    assert client.get('/api/replenishment-beta/applications/' + foreign).status_code == 200
    assert client.get('/api/replenishment-beta/applications').json()['total'] == 2


def test_project_directory_scoped_sql_and_cart_delete_preconditions(db, replenishment):
    from app.models.maintenance_project import MaintenanceProjectUserAssignment
    from tests.boss_board_helpers import make_project
    uid, pid, part, client = replenishment
    make_project(db, 'unrelated-replenishment-project')
    db.add(MaintenanceProjectUserAssignment(assignment_id='replenishment-assignment', project_id=pid, user_id=uid,
        responsibility_type='primary_manager', assigned_by='synthetic', assignment_reason='synthetic scope'))
    db.commit()
    grant(db, uid, 'replenishment_project_directory', ['read'], scope='assigned_or_sales_projects')
    response = client.get('/api/replenishment-beta/projects')
    assert response.status_code == 200, response.text
    assert [row['project_id'] for row in response.json()['items']] == [pid]
    grant(db, uid, 'replenishment_project_commands', ['cart_update', 'cart_delete'])
    path = f'/api/replenishment-beta/cart-drafts/{pid}'
    response = client.put(path, json={'lines': [{'part_id': part, 'quantity': 2}]})
    assert response.status_code == 200, response.text
    draft = response.json()['draft']
    response = client.request('DELETE', path, params={'view': 'authorization', 'expected_version': draft['version']})
    assert response.status_code == 200, response.text
    ctx = response.json()['authorization_context']
    pre = {key: ctx[key] for key in ('catalog_version', 'authorization_epoch', 'context')}
    pre['context']['authorization_revision'] = 'stale'
    response = client.request('DELETE', path, params={'expected_version': draft['version']}, json={'authorization_precondition': pre})
    assert response.status_code == 409, response.text
    response = client.delete(path, params={'expected_version': draft['version']})
    assert response.status_code == 200, response.text
    assert response.json()['deleted'] is True


def test_all_scope_reviser_has_own_version_identity_distinct_from_application_owner(db, replenishment, monkeypatch):
    from app.auth import _make_token
    from app.models.system import SysUser
    uid, pid, part, client = replenishment
    monkeypatch.setattr(get_settings(), 'replenishment_auto_review_enabled', True)
    grant(db, uid, 'replenishment_project_commands', ['submit'])
    response = submit(client, pid, part)
    assert response.status_code == 201, response.text
    original = response.json()
    reviewer = SysUser(username='explicit-all-reviser', role='readonly', password_hash='synthetic')
    db.add(reviewer); db.commit()
    grant(db, reviewer.id, 'replenishment_project_commands', ['revise'])
    grant(db, reviewer.id, 'replenishment_request', ['revise'])
    token, _ = _make_token(reviewer.role, reviewer.username, None, token_version=reviewer.token_version, authn='sys_user')
    client.headers['Authorization'] = 'Bearer ' + token
    path = '/api/replenishment-beta/applications/' + original['application_id'] + '/revisions'
    response = client.post(path, json={'client_request_id': 'different-version-producer', 'expected_application_version': original['version'],
        'lines': [{'part_id': part, 'quantity': 5}]})
    assert response.status_code == 200, response.text
    app = db.get(Application, original['application_id'])
    versions = list(db.scalars(select(Version).where(Version.application_id == app.application_id).order_by(Version.version_no)))
    assert app.owner_user_id == uid
    assert [version.owner_user_id for version in versions] == [uid, reviewer.id]
    assert db.get(SysAuthorizationArtifact, ('replenishment_version', versions[1].version_id)).owner_user_id == reviewer.id


def test_full_application_read_and_export_are_in_the_same_readonly_snapshot(db, replenishment, monkeypatch):
    from sqlalchemy import text
    from app.services import replenishment as business
    uid, pid, part, client = replenishment
    grant(db, uid, 'replenishment_project_commands', ['submit'])
    aid = submit(client, pid, part).json()['application_id']
    grant(db, uid, 'replenishment_request', ['read', 'export_screening'])
    observed = []
    original_read, original_export = business.get_application, business.system_screening_workbook
    def observe(snapshot):
        observed.append((snapshot.scalar(text('SHOW transaction_isolation')), snapshot.scalar(text('SHOW transaction_read_only'))))
    def read(snapshot, *args, **kwargs):
        observe(snapshot)
        return original_read(snapshot, *args, **kwargs)
    def export(snapshot, *args, **kwargs):
        observe(snapshot)
        return original_export(snapshot, *args, **kwargs)
    monkeypatch.setattr(business, 'get_application', read)
    monkeypatch.setattr(business, 'system_screening_workbook', export)
    path = '/api/replenishment-beta/applications/' + aid
    assert client.get(path).status_code == 200
    assert client.get(path + '/exports/system-screening.xlsx').status_code == 200
    assert observed == [('repeatable read', 'on')] * 2
