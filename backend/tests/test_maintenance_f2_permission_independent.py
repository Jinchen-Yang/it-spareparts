"""Independent F2 HTTP counterexamples; real login, no authorization mocks."""
from datetime import UTC, datetime
import io
import uuid

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.auth import hash_password
from app.config import get_settings
from app.main import app
from app.models.maintenance_project import MaintenanceProjectUserAssignment
from app.models.system import SysUser
from app.models.maintenance import MaintenanceManualCostOverride
from app.services import maintenance_project_master_workbook as master
from tests.test_maintenance_project_master_v2_editable import _make_project_with_line

ACTION = "action_maintenance_expense_collection_upload"
SHEETS = [None, "03_备件订单", "04_报销订单", "05_项目经理回款单", "06_现场领用与返还"]


def account(db, project, *, role="readonly", profit=False, cost=False,
            action=False, page=True, assignment="viewer", sales=False):
    name = f"f2-independent-{uuid.uuid4().hex[:10]}"
    raw = {"page_maintenance": page, "data_profit": profit,
           "data_purchase_cost": cost, ACTION: action,
           "own_maintenance_projects_only": True}
    user = SysUser(username=name, role=role, display_name=name,
                   salesperson_name="F2 Sales" if sales else None,
                   password_hash=hash_password("F2-password-123"), is_active=True,
                   template_code=role, template_version=1,
                   template_perms={} if role == "admin" else raw,
                   perm_overrides={}, permissions={})
    if sales:
        project.salesperson = "F2 Sales"
    db.add(user)
    db.flush()
    link = None
    if assignment:
        link = MaintenanceProjectUserAssignment(
            assignment_id=str(uuid.uuid4()), project_id=project.project_id,
            responsibility_type=assignment, user_id=user.id, version=1,
            assigned_by="independent-review", assignment_reason="F2 independent HTTP")
        db.add(link)
    db.commit()
    client = TestClient(app)
    login = client.post("/api/auth/login", json={"username": name, "password": "F2-password-123"})
    assert login.status_code == 200, login.text
    client.headers["Authorization"] = "Bearer " + login.json()["token"]
    return client, user, link


def post_file(client, path, content):
    return client.post(path, files={"file": ("review.xlsx", content,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


@pytest.mark.parametrize("kwargs,global_download,global_upload,project_download,project_upload", [
    ({}, False, False, False, False),
    ({"cost": True, "action": True}, False, False, False, False),
    ({"profit": True, "cost": False, "action": True, "sales": True}, False, False, False, False),
    ({"profit": True, "cost": True}, True, False, True, False),
    ({"profit": True, "cost": True, "action": True}, True, True, True, True),
    ({"assignment": "primary_manager"}, False, False, True, True),
    ({"sales": True, "profit": True, "cost": True, "assignment": None}, True, False, True, True),
    ({"role": "admin", "assignment": None}, True, True, True, True),
    ({"page": False, "assignment": "primary_manager", "profit": True, "cost": True, "action": True}, False, False, False, False),
])
def test_real_http_six_entry_contract(db, monkeypatch, kwargs, global_download,
                                     global_upload, project_download, project_upload):
    monkeypatch.setattr(get_settings(), "maintenance_project_master_v2_enabled", True)
    project, _, _, line = _make_project_with_line(db)
    client, _, _ = account(db, project, **kwargs)
    global_base = "/api/maintenance/spare-part-lines"
    response = client.get(global_base + ".xlsx", params={"range": "custom", "from": "2026-08-01", "to": "2026-08-31"})
    assert response.status_code == (200 if global_download else 403), response.text[:500]
    content = response.content if global_download else b"not-an-xlsx"
    for verb in ("validate", "apply"):
        result = post_file(client, global_base + "/" + verb, content)
        assert result.status_code == (200 if global_upload else 403), result.text

    base = f"/api/maintenance/projects/stable/{project.project_id}"
    if kwargs.get("page", True) or kwargs.get("role") == "admin":
        detail = client.get(base)
        card = client.get(f"/api/maintenance/boss-board/projects/{project.project_id}")
        assert detail.status_code == card.status_code == 200
        assert detail.json()["project"]["can_edit_master_workbook"] is project_upload
        assert card.json()["can_edit_master_workbook"] is project_upload
    for sheet in SHEETS:
        download = client.get(base + "/master-workbook.xlsx", params={"sheets": sheet} if sheet else {})
        assert download.status_code == (200 if project_download else 403), (sheet, download.text[:500])
        content = download.content if project_download else b"not-an-xlsx"
        for verb in ("validate", "apply"):
            result = post_file(client, base + "/master-workbook/" + verb, content)
            assert result.status_code == (200 if project_upload else 403), (sheet, verb, result.text)
    db.refresh(line)
    assert line.qty == 2


def test_assignment_revocation_after_download_blocks_both_upload_phases(db, monkeypatch):
    monkeypatch.setattr(get_settings(), "maintenance_project_master_v2_enabled", True)
    project, _, _, line = _make_project_with_line(db)
    client, _, link = account(db, project, assignment="primary_manager")
    base = f"/api/maintenance/projects/stable/{project.project_id}/master-workbook"
    downloaded = client.get(base + ".xlsx")
    assert downloaded.status_code == 200
    link.archived_at = datetime.now(UTC)
    link.archived_by = "independent-review"
    link.archive_reason = "revoke manager before upload"
    link.version += 1
    db.commit()
    for verb in ("validate", "apply"):
        denied = post_file(client, base + "/" + verb, downloaded.content)
        assert denied.status_code == 403, denied.text
    assert client.get(base + ".xlsx").status_code == 403
    db.refresh(line)
    assert line.qty == 2


def test_manager_workbook_cannot_transfer_to_other_account_or_project(db, monkeypatch):
    monkeypatch.setattr(get_settings(), "maintenance_project_master_v2_enabled", True)
    project, _, _, _ = _make_project_with_line(db)
    other, _, _, _ = _make_project_with_line(db)
    client, _, _ = account(db, project, assignment="primary_manager")
    viewer, _, _ = account(db, project)
    base = f"/api/maintenance/projects/stable/{project.project_id}/master-workbook"
    other_base = f"/api/maintenance/projects/stable/{other.project_id}/master-workbook"
    downloaded = client.get(base + ".xlsx")
    assert downloaded.status_code == 200
    assert client.get(other_base + ".xlsx").status_code == 403
    for verb in ("validate", "apply"):
        assert post_file(viewer, base + "/" + verb, downloaded.content).status_code == 403
        assert post_file(client, other_base + "/" + verb, downloaded.content).status_code == 403


def test_token_revocation_prevents_previously_authorized_download_and_apply(db, monkeypatch):
    monkeypatch.setattr(get_settings(), "maintenance_project_master_v2_enabled", True)
    project, _, _, _ = _make_project_with_line(db)
    client, user, _ = account(db, project, profit=True, cost=True, action=True)
    base = f"/api/maintenance/projects/stable/{project.project_id}/master-workbook"
    downloaded = client.get(base + ".xlsx")
    assert downloaded.status_code == 200
    user.token_version = (user.token_version or 0) + 1
    db.commit()
    assert client.get(base + ".xlsx").status_code == 401
    for verb in ("validate", "apply"):
        assert post_file(client, base + "/" + verb, downloaded.content).status_code == 401


def test_global_scope_shrinks_without_boolean_change_and_old_file_is_rejected(db):
    project, part, _, _ = _make_project_with_line(db)
    _, other_part, _, _ = _make_project_with_line(db)
    client, _, link = account(db, project, profit=True, cost=True, action=True)
    base = "/api/maintenance/spare-part-lines"
    params = {"range": "custom", "from": "2026-08-01", "to": "2026-08-31"}
    downloaded = client.get(base + ".xlsx", params=params)
    assert downloaded.status_code == 200
    wb = load_workbook(io.BytesIO(downloaded.content))
    sheet = wb[master.GLOBAL_SHEET]
    assert sheet.max_row == 2
    values = [cell.value for row in sheet for cell in row]
    assert part.pn_std in values
    assert other_part.pn_std not in values
    sheet.cell(row=2, column=10, value=88)
    payload = io.BytesIO()
    wb.save(payload)
    assert post_file(client, base + "/validate", payload.getvalue()).status_code == 200
    link.archived_at = datetime.now(UTC)
    link.archived_by = "independent-review"
    link.archive_reason = "scope shrinks after preflight"
    link.version += 1
    db.commit()
    empty = client.get(base + ".xlsx", params=params)
    assert empty.status_code == 200  # unchanged money/action booleans do not grant rows
    assert load_workbook(io.BytesIO(empty.content))[master.GLOBAL_SHEET].max_row == 1
    for verb in ("validate", "apply"):
        denied = post_file(client, base + "/" + verb, payload.getvalue())
        assert denied.status_code == 422, denied.text
        assert denied.json()["detail"]["code"] == "project_scope_denied"
    assert db.scalars(select(MaintenanceManualCostOverride)).all() == []
