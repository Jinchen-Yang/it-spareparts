"""Independent W2a checks through real workbook HTTP transactions and readers.

Failure windows combine moved helpers instead of testing them in isolation.
Synthetic data only; the normal test harness owns a fresh PostgreSQL database.
"""
from copy import deepcopy
from datetime import date
from decimal import Decimal
import io
import json
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import event, select, text

from app.api import maintenance_project_master_workbook as workbook_api
from app.db import SessionLocal
from app.models.maintenance import (
    FMaintenanceLine, FProjectExpense, MaintenanceDemandDeleteEvent,
    MaintenanceDemandTombstone, MaintenanceManualCostOverride,
)
from app.models.maintenance_project import (
    MaintenanceProject, MaintenanceProjectContract, MaintenanceProjectXsdd,
)
from app.models.maintenance_project_operations import (
    MaintenanceProjectExpenseAttribution, MaintenanceProjectOperationAudit,
    MaintenanceProjectWorkbookOperation, MaintenanceProjectWorkbookState,
    MaintenanceSiteIssue, MaintenanceSiteIssueDisplayCounter, MaintenanceSiteIssueLine,
)
from app.models.maintenance_source_assignment import MaintenanceSourceOrderAssignment
from app.models.maintenance_manager import MaintenanceCollectionMilestone
from app.models.sales import FSalesOrder
from app.models.system import SysImportBatch
from app.services import maintenance_expense_collection_workbook as ec
from app.services import maintenance_project_master_workbook as master
from app.services import maintenance_workbook_fact_writes as facts
from tests.test_maintenance_expense_collection_workbook import uploader
from tests.test_maintenance_project_master_v2_editable import _make_project_with_line

_XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
_DAY = date(2026, 10, 23)
_TABLES = (
    FMaintenanceLine, FProjectExpense, MaintenanceManualCostOverride,
    MaintenanceProjectXsdd, MaintenanceProjectContract,
    MaintenanceProjectExpenseAttribution, MaintenanceProjectOperationAudit,
    MaintenanceProjectWorkbookOperation, MaintenanceProjectWorkbookState,
    MaintenanceSiteIssue, MaintenanceSiteIssueLine, MaintenanceSiteIssueDisplayCounter,
    MaintenanceSourceOrderAssignment, MaintenanceDemandTombstone,
    MaintenanceDemandDeleteEvent, MaintenanceCollectionMilestone,
)


@pytest.fixture(autouse=True)
def _real_v2_with_bounded_database_waits(monkeypatch):
    monkeypatch.setattr(workbook_api.get_settings(), 'maintenance_project_master_v2_enabled', True)

    def bound(_session, _transaction, connection):
        connection.execute(text("SET LOCAL lock_timeout = '10s'"))

    event.listen(SessionLocal, 'after_begin', bound)
    try:
        yield
    finally:
        event.remove(SessionLocal, 'after_begin', bound)


def _snapshot(db):
    # Read SQL table values, not cached ORM attributes. Access/auth logs are
    # deliberately outside the business-transaction invariant.
    return {
        model.__tablename__: deepcopy([dict(row) for row in db.execute(
            select(model.__table__).order_by(*model.__table__.primary_key.columns)
        ).mappings()])
        for model in _TABLES
    }


def _save(book):
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def _post(db, client, project_id, data, action='apply'):
    # Export may have created a workbook state row. Release its lock before
    # invoking the independent real request SessionLocal.
    db.commit()
    return client.post(
        f'/api/maintenance/projects/stable/{project_id}/master-workbook/{action}',
        files={'file': ('cross-boundary.xlsx', data, _XLSX)},
    )


def _mixed_file(db, *, automatic=False):
    project, part, order, line = _make_project_with_line(
        db, unit_cost=Decimal('100') if automatic else None,
        cost_source='direct' if automatic else 'none',
    )
    project.period_from = date(2026, 1, 1)
    contract = db.scalar(select(MaintenanceProjectContract).where(
        MaintenanceProjectContract.project_id == project.project_id))
    # 04 requires a current relation; 02 separately exercises deferred XSDD
    # materialization. Do not make the test invent expense eligibility.
    contract.contract_no = "XSDD-CROSS-EXPENSE"
    batch = SysImportBatch(filename='cross-boundary-sales.xlsx', file_type='sales',
                           file_hash=uuid4().hex, status='success')
    db.add(batch)
    db.flush()
    db.add(FSalesOrder(raw_order_id=f'cross-sales-{uuid4()}',
        order_no=order.linked_sales_order_no, amount_ex_tax=Decimal('100.01'),
        tax_rate=Decimal('0.13'), data_status='已生效', import_batch_id=batch.id))
    # Exercise rollback of ON CONFLICT UPDATE, not only rollback of INSERT.
    db.add(MaintenanceSiteIssueDisplayCounter(issue_date=_DAY, last_sequence=7))
    db.commit()
    client = uploader(db, username=f'cross-boundary-{uuid4().hex[:8]}')
    book = load_workbook(io.BytesIO(master.build_project_master_v2(
        db, project_id=project.project_id,
        sheets=(master.V2_SHEET_PLAN, master.V2_SHEET_PARTS,
                master.V2_SHEET_EXPENSE, master.V2_SHEET_SITE))))
    sheet = book[master.V2_SHEET_PARTS]
    columns = {cell.value: cell.column for cell in sheet[1]}
    for name, value in {'需求数量': 5, '退货数量': 1,
                        '人工未税单位成本': 7.13, '人工成本原因': '独立事务证据'}.items():
        sheet.cell(2, columns[name], value)
    book[master.V2_SHEET_PLAN].append([
        'CREATE', order.linked_sales_order_no, 1, '2026-11-23', 'day', 50,
        None, None, None, None, None, None, None, None,
    ])
    book[master.V2_SHEET_EXPENSE].append([
        '', 'BXD-CROSS-001', 1, '2026-10-23 10:00:00', '独立核验',
        '维保费用', '交通费', '跨模块写入', contract.contract_no,
        project.display_name, '合成销售', '已生效', 17, 'ex', 17, 19.21,
        '与合同物化同事务', '',
    ])
    for quantity in (1, 2):
        book[master.V2_SHEET_SITE].append([
            None, _DAY.isoformat(), part.pn_std, None, quantity, '是', '自动号事务',
        ])
    data = _save(book)
    db.commit()
    return project.project_id, line.id, client, data


def _assert_mixed_success(db, project_id, line_id, *, automatic=False):
    db.expire_all()
    line = db.get(FMaintenanceLine, line_id)
    assert (line.qty, line.return_qty) == (Decimal('5'), Decimal('1'))
    assert (line.unit_cost_ex_tax, line.unit_cost_inc_tax,
            line.cost_amount_ex_tax, line.cost_amount_inc_tax, line.cost_source) == (
        (Decimal('100'), Decimal('113'), Decimal('400'), Decimal('452'), 'direct')
        if automatic else
        (Decimal('7.13'), Decimal('8.06'), Decimal('28.52'), Decimal('32.24'), 'manual')
    )
    evidence = db.scalar(select(MaintenanceManualCostOverride).where(
        MaintenanceManualCostOverride.line_id == line_id))
    assert (evidence.unit_cost_ex_tax, evidence.unit_cost_inc_tax,
            evidence.reason, evidence.version) == (
        Decimal('7.13'), Decimal('8.06'), '独立事务证据', 1)
    contracts = db.scalars(select(MaintenanceProjectContract).where(
        MaintenanceProjectContract.project_id == project_id)).all()
    assert len(contracts) == 2
    fallback = next(c for c in contracts if c.source == 'sales_fallback')
    expense_contract = next(c for c in contracts if c.contract_no == 'XSDD-CROSS-EXPENSE')
    assert (fallback.contract_amount, fallback.amount_inc_tax) == (Decimal('100.01'), Decimal('113.01'))
    milestone = db.scalar(select(MaintenanceCollectionMilestone))
    assert milestone.project_contract_id == fallback.project_contract_id
    assert milestone.planned_amount == Decimal('50')
    expenses = ec.project_expenses(db, project_id)
    assert len(expenses) == 1
    assert (expenses[0].amount_ex_tax, expenses[0].amount_inc_tax) == (Decimal('17'), Decimal('19.21'))
    attribution = db.scalar(select(MaintenanceProjectExpenseAttribution))
    assert attribution.project_contract_id == expense_contract.project_contract_id
    issues = db.scalars(select(MaintenanceSiteIssue).where(
        MaintenanceSiteIssue.project_id == project_id)).all()
    assert sorted(issue.display_no for issue in issues) == ['LY-20261023-0008', 'LY-20261023-0009']
    assert db.get(MaintenanceSiteIssueDisplayCounter, _DAY).last_sequence == 9
    assert len(db.scalars(select(MaintenanceSiteIssueLine)).all()) == 2
    audits = db.scalars(select(MaintenanceProjectOperationAudit)).all()
    assert len([a for a in audits if a.entity_type == 'project_contract' and a.action == 'CREATE']) == 1
    assert len([a for a in audits if a.entity_type == 'maintenance_line' and a.action == 'UPDATE']) == 1
    assert len([a for a in audits if a.entity_type == 'master_workbook_operation' and a.action == 'apply_completed']) == 1
    assert len(db.scalars(select(MaintenanceProjectWorkbookOperation).where(
        MaintenanceProjectWorkbookOperation.operation_type == 'file_apply')).all()) == 1


@pytest.mark.parametrize('window', ['contract_audit', 'after_expense', 'completion_receipt'])
def test_http_multisheet_failure_rolls_back_all_boundaries_then_retries_once(db, monkeypatch, window):
    project_id, line_id, client, data = _mixed_file(db)
    before = _snapshot(db)
    observed = []
    real_audit = facts._write_audit
    real_ec_apply = ec.apply

    def fail_contract_audit(session, **kwargs):
        # The relation has been flushed; include its audit in this transaction
        # before throwing, so an extracted helper with its own commit is caught.
        real_audit(session, **kwargs)
        session.flush()
        observed.append('contract_audit')
        raise master.WorkbookError('synthetic_boundary_failure', '合成合同审计失败')

    def fail_after_expense(session, *args, **kwargs):
        real_ec_apply(session, *args, **kwargs)
        session.flush()
        assert len(session.scalars(select(FProjectExpense)).all()) == 1
        assert session.get(MaintenanceSiteIssueDisplayCounter, _DAY).last_sequence == 9
        observed.append('after_expense')
        raise master.WorkbookError('synthetic_boundary_failure', '合成费用后失败')

    def fail_completion(_mapper, _connection, target):
        if target.entity_type == 'master_workbook_operation' and target.action == 'apply_completed':
            observed.append('completion_receipt')
            raise master.WorkbookError('synthetic_boundary_failure', '合成首次回执失败')

    with monkeypatch.context() as inject:
        if window == 'contract_audit':
            inject.setattr(facts, '_write_audit', fail_contract_audit)
        elif window == 'after_expense':
            inject.setattr(ec, 'apply', fail_after_expense)
        else:
            event.listen(MaintenanceProjectOperationAudit, 'before_insert', fail_completion)
        try:
            response = _post(db, client, project_id, data)
        finally:
            if window == 'completion_receipt':
                event.remove(MaintenanceProjectOperationAudit, 'before_insert', fail_completion)
    assert observed == [window], response.text
    assert response.status_code == 422, response.text
    assert response.json()['detail']['code'] == 'synthetic_boundary_failure'
    assert _snapshot(db) == before

    retry = _post(db, client, project_id, data)
    assert retry.status_code == 200, retry.text
    assert retry.json()['replayed'] is False
    _assert_mixed_success(db, project_id, line_id)
    assert db.get(MaintenanceProjectWorkbookState, project_id).revision == (
        before['maintenance_project_workbook_state'][0]['revision'] + 1)
    first = _snapshot(db)
    replay = _post(db, client, project_id, data)
    assert replay.status_code == 200, replay.text
    assert replay.json() == {**retry.json(), 'replayed': True}
    assert _snapshot(db) == first


@pytest.mark.parametrize('automatic', [False, True])
def test_http_quantity_and_manual_evidence_remain_consistent_in_page_readback(db, automatic):
    project_id, line_id, client, data = _mixed_file(db, automatic=automatic)
    response = _post(db, client, project_id, data)
    assert response.status_code == 200, response.text
    _assert_mixed_success(db, project_id, line_id, automatic=automatic)
    db.commit()
    page = client.get(f'/api/maintenance/projects/stable/{project_id}/master-workbook/rows',
                      params={'sheet': master.V2_SHEET_PARTS})
    assert page.status_code == 200, page.text
    row = next(row for row in page.json()['rows'] if row['line_id'] == line_id)
    assert Decimal(str(row['qty'])) == Decimal('5')
    assert Decimal(str(row['return_qty'])) == Decimal('1')
    assert Decimal(str(row['manual_unit_cost_ex_tax'])) == Decimal('7.13')
    assert row['cost_amount_inc_tax'] is not None, row
    assert Decimal(str(row['cost_amount_inc_tax'])) == (Decimal('452') if automatic else Decimal('32.24'))


def test_http_void_cascade_audit_failure_preserves_foreign_and_inactive_rows_then_retries(db):
    project, part, order, line = _make_project_with_line(db, source_line_id='w2a-shared-raw-id')
    foreign = MaintenanceProject(project_id=str(uuid4()), project_code='W2A-FOREIGN',
                                 display_name='其他项目', lifecycle_status='ongoing')
    db.add(foreign)
    db.flush()
    states = [(project.project_id, True), (project.project_id, False), (foreign.project_id, True)]
    sites = []
    for number, (owner, active) in enumerate(states, 1):
        issue = MaintenanceSiteIssue(issue_id=str(uuid4()), project_id=owner,
            issue_no=f'W2A-CASCADE-{number}', issue_date=_DAY, raw_status='已确认',
            status_mapping_state='mapped', normalized_status='confirmed',
            status_mapping_version='v1', source='legacy')
        db.add(issue)
        db.flush()
        site = MaintenanceSiteIssueLine(issue_line_id=str(uuid4()), issue_id=issue.issue_id,
            line_no=1, part_id=part.id, pn=part.pn_std, quantity=Decimal('1'),
            source_order_id=order.raw_order_id, source_line_id=line.raw_line_id,
            algorithm_version='v1', is_active=active)
        db.add(site)
        sites.append(site)
    db.commit()
    client = uploader(db, username='w2a-void-crosscheck')
    book = load_workbook(io.BytesIO(master.build_project_master_v2(
        db, project_id=project.project_id, sheets=(master.V2_SHEET_PARTS,))))
    book[master.V2_SHEET_PARTS].cell(2, 1, 'VOID')
    data = _save(book)
    db.commit()
    before = _snapshot(db)
    observed = []

    def fail_void_audit(_mapper, _connection, audit):
        if audit.entity_type == 'maintenance_line' and audit.action == 'VOID':
            observed.append(audit.after_json['cascaded_site_lines'])
            raise master.WorkbookError('synthetic_void_audit', '合成作废审计失败')

    event.listen(MaintenanceProjectOperationAudit, 'before_insert', fail_void_audit)
    try:
        failed = _post(db, client, project.project_id, data)
    finally:
        event.remove(MaintenanceProjectOperationAudit, 'before_insert', fail_void_audit)
    assert observed == [1]
    assert failed.status_code == 422, failed.text
    assert failed.json()['detail']['code'] == 'synthetic_void_audit'
    assert _snapshot(db) == before
    success = _post(db, client, project.project_id, data)
    assert success.status_code == 200, success.text
    db.expire_all()
    assert not db.get(FMaintenanceLine, line.id).is_active
    assert [db.get(MaintenanceSiteIssueLine, site.issue_line_id).is_active for site in sites] == [False, False, True]
    audits = db.scalars(select(MaintenanceProjectOperationAudit).where(
        MaintenanceProjectOperationAudit.entity_type == 'maintenance_line',
        MaintenanceProjectOperationAudit.action == 'VOID')).all()
    assert len(audits) == 1 and audits[0].after_json['cascaded_site_lines'] == 1
    after = _snapshot(db)
    replay = _post(db, client, project.project_id, data)
    assert replay.status_code == 200, replay.text
    assert replay.json() == {**success.json(), 'replayed': True}
    assert _snapshot(db) == after


def test_http_manual_cost_only_becomes_visible_after_successful_apply(db):
    """Minimal baseline bug: only fill manual price on one unpriced demand."""
    project, _part, _order, line = _make_project_with_line(db, unit_cost=None, cost_source='none')
    project_id, line_id = project.project_id, line.id
    client = uploader(db, username='w2a-minimal-manual-price')
    route = f'/api/maintenance/projects/stable/{project_id}/master-workbook/rows'
    before = client.get(route, params={'sheet': master.V2_SHEET_PARTS})
    assert before.status_code == 200, before.text
    book = load_workbook(io.BytesIO(master.build_project_master_v2(
        db, project_id=project_id, sheets=(master.V2_SHEET_PARTS,))))
    sheet = book[master.V2_SHEET_PARTS]
    columns = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, columns['人工未税单位成本'], 7.13)
    sheet.cell(2, columns['人工成本原因'], '仅补价的最小复现')
    applied = _post(db, client, project_id, _save(book))
    assert applied.status_code == 200, applied.text
    after = client.get(route, params={'sheet': master.V2_SHEET_PARTS})
    assert after.status_code == 200, after.text
    db.expire_all()
    stored = db.get(FMaintenanceLine, line_id)
    fields = ('cost_source', 'cost_amount_inc_tax', 'unit_cost_ex_tax',
              'unit_cost_inc_tax', 'manual_unit_cost_ex_tax', 'missing_kind', 'can_refill')
    evidence = {
        'before': {key: before.json()['rows'][0][key] for key in fields},
        'after': {key: after.json()['rows'][0][key] for key in fields},
        'stored': {key: str(getattr(stored, key)) for key in (
            'qty', 'return_qty', 'cost_source', 'cost_tax_basis', 'unit_cost',
            'cost_amount', 'unit_cost_ex_tax', 'unit_cost_inc_tax',
            'cost_amount_ex_tax', 'cost_amount_inc_tax')},
    }
    assert after.json()['rows'][0]['cost_amount_inc_tax'] == '16.12', json.dumps(evidence, ensure_ascii=False)
