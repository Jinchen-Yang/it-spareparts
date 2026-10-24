"""Master workbook alternatives are complete grants over one real project."""
from datetime import datetime, timezone
from io import BytesIO

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import event, select

from app.authz import runtime
from app.config import get_settings
from app.db import get_db
from app.models.authorization import SysAuthorizationGrant
from app.models.maintenance_project import MaintenanceProjectUserAssignment
from app.models.maintenance_project_operations import MaintenanceProjectWorkbookState as MaintenanceWorkbookState
from tests.test_authz_s3_checklist import setup, cerbos_url  # noqa: F401


@pytest.fixture
def master(db, setup, monkeypatch, request):
    from app.api import maintenance_project_master_workbook as api
    uid, identity, pid, _, headers = setup
    enabled = ['maintenance_acceptance', 'maintenance_general_workbook']
    if getattr(request, 'param', None) != 'general_only':
        enabled.append('maintenance_managed_workbook')
    monkeypatch.setattr(get_settings(), 'authz_enabled_resources', enabled)
    monkeypatch.setattr(get_settings(), 'maintenance_project_master_v2_enabled', True)
    runtime.activate_configured_resources(db.get_bind())
    app = FastAPI()
    app.include_router(api.router, prefix='/api')
    app.dependency_overrides[get_db] = lambda: db
    client = TestClient(app)
    client.headers.update(headers)
    return uid, pid, client, f'/api/maintenance/projects/stable/{pid}/master-workbook.xlsx'


def grant(db, uid, resource, fields, scope):
    db.add(SysAuthorizationGrant(user_id=uid, resource=resource, actions=['download'],
        scope=scope, field_permissions=fields, version=1))
    db.commit()


@pytest.mark.parametrize("master", ["both", "general_only"], indirect=True)
def test_general_whole_download_is_read_only_without_old_role_or_page(db, master):
    uid, pid, client, url = master
    grant(db, uid, 'maintenance_general_workbook', ['data_purchase_cost', 'data_profit', 'data_customer'], 'all_projects')
    assert db.scalar(select(MaintenanceWorkbookState).where(MaintenanceWorkbookState.project_id == pid)) is None
    resolution = client.get(url + '?view=authorization')
    assert resolution.status_code == 200, resolution.text
    assert set(resolution.json()) == {'locator', 'authorization_context'}
    assert resolution.json()['locator'] == {'project_id': pid}
    statements = []
    def record(conn, cursor, statement, parameters, context, executemany):
        statements.append(statement)
    event.listen(db.get_bind(), 'before_cursor_execute', record)
    try:
        response = client.get(url)
    finally:
        event.remove(db.get_bind(), 'before_cursor_execute', record)
    assert response.status_code == 200, response.text
    assert not any('FOR UPDATE' in sql or 'FOR SHARE' in sql or sql.lstrip().upper().startswith(('INSERT', 'UPDATE', 'DELETE')) for sql in statements)
    book = load_workbook(BytesIO(response.content))
    assert len(book.sheetnames) >= 6
    book.close()
    assert db.scalar(select(MaintenanceWorkbookState).where(MaintenanceWorkbookState.project_id == pid)) is None


def test_managed_complete_contract_needs_real_current_assignment(db, master):
    uid, pid, client, url = master
    grant(db, uid, 'maintenance_managed_workbook', ['data_purchase_cost', 'data_profit', 'data_customer'], 'primary_managed_projects')
    assert client.get(url).status_code == 403
    row = MaintenanceProjectUserAssignment(assignment_id='master-download-manager', project_id=pid,
        user_id=uid, responsibility_type='primary_manager', assigned_at=datetime.now(timezone.utc),
        assigned_by='synthetic', assignment_reason='test actual relation')
    db.add(row)
    db.commit()
    assert client.get(url).status_code == 200
    row.archived_at = datetime.now(timezone.utc)
    row.archived_by = "synthetic"
    row.archive_reason = "test relation removal"
    row.version += 1
    db.commit()
    assert client.get(url).status_code == 403


def test_scope_and_fields_from_different_alternatives_never_union(db, master):
    uid, pid, client, url = master
    # General has the broad scope but cannot render the complete cost/profit file.
    grant(db, uid, 'maintenance_general_workbook', ['data_purchase_cost'], 'all_projects')
    # Managed has all fields but no relationship to this real project.
    grant(db, uid, 'maintenance_managed_workbook', ['data_purchase_cost', 'data_profit', 'data_customer'], 'primary_managed_projects')
    assert client.get(url).status_code == 403


def test_cutover_suspended_alternative_cannot_be_treated_as_never_migrated(db, master, monkeypatch):
    uid, pid, client, url = master
    grant(db, uid, 'maintenance_general_workbook', ['data_purchase_cost', 'data_profit', 'data_customer'], 'all_projects')
    monkeypatch.setattr(get_settings(), 'authz_enabled_resources', ['maintenance_acceptance', 'maintenance_general_workbook'])
    assert client.get(url).status_code == 503


@pytest.mark.parametrize('resource,action,suffix', [
    ('maintenance_collection', 'read_plan', 'master-workbook/collection-plan'),
    ('maintenance_expense', 'read_rows', 'expense-rows'),
])
def test_project_finance_rows_use_own_action_and_profit_without_cost(db, master, monkeypatch, resource, action, suffix):
    uid, pid, client, _ = master
    settings = get_settings()
    monkeypatch.setattr(settings, 'authz_enabled_resources', settings.authz_enabled_resources + [resource])
    runtime.activate_configured_resources(db.get_bind())
    row = SysAuthorizationGrant(user_id=uid, resource=resource, actions=[action],
        scope='all_projects', field_permissions=['data_profit'], version=1)
    db.add(row)
    db.commit()
    url = f'/api/maintenance/projects/stable/{pid}/{suffix}'
    view = client.get(url + '?view=authorization')
    assert view.status_code == 200, view.text
    assert set(view.json()) == {'locator', 'authorization_context'}
    response = client.get(url)
    assert response.status_code == 200, response.text
    assert response.json()['rows'] == []
    decision = response.json()['authorization_context']['resources'][resource]
    assert decision['field_permissions'] == ['data_profit']
    row.actions = []
    row.version += 1
    db.commit()
    assert client.get(url).status_code == 403


def test_master_validate_parses_actual_plan_then_requires_note_action(db, master, monkeypatch):
    from app.services import maintenance_project_master_workbook as workbook
    from tests.test_maintenance_project_note_workbook import _note_cell
    uid, pid, client, url = master
    settings = get_settings()
    monkeypatch.setattr(settings, 'authz_enabled_resources', settings.authz_enabled_resources +
        ['maintenance_general_workbook_write', 'maintenance_project'])
    runtime.activate_configured_resources(db.get_bind())
    grant(db, uid, 'maintenance_general_workbook', ['data_purchase_cost', 'data_profit', 'data_customer'], 'all_projects')
    write = SysAuthorizationGrant(user_id=uid, resource='maintenance_general_workbook_write', actions=['validate'],
        scope='all_projects', field_permissions=['data_purchase_cost', 'data_profit', 'data_customer'], version=1)
    db.add(write)
    db.commit()
    downloaded = client.get(url)
    assert downloaded.status_code == 200, downloaded.text
    base = url.removesuffix('.xlsx') + '/validate'
    plain = client.post(base, files={'file': ('master.xlsx', downloaded.content)})
    assert plain.status_code == 200, plain.text
    book = load_workbook(BytesIO(downloaded.content))
    _note_cell(book).value = 'requires independent project update'
    output = BytesIO()
    book.save(output)
    book.close()
    response = client.post(base, files={'file': ('master.xlsx', output.getvalue())})
    assert response.status_code == 403, response.text
    view = client.post(base + '?view=authorization', files={'file': ('master.xlsx', output.getvalue())})
    assert view.status_code == 200, view.text
    assert set(view.json()) == {'locator', 'authorization_context'}
    assert view.json()['authorization_context']['resources']['maintenance_project']['decisions']['update']['effect'] == 'deny'
    db.add(SysAuthorizationGrant(user_id=uid, resource='maintenance_project', actions=['update'],
        scope='all_projects', field_permissions=[], version=1))
    db.commit()
    response = client.post(base, files={'file': ('master.xlsx', output.getvalue())})
    assert response.status_code == 200, response.text
    assert any(change['field'] == '项目备注' for change in response.json()['changes'])
    from app.models.maintenance_project import MaintenanceProject
    db.expire_all()
    assert db.get(MaintenanceProject, pid).project_note is None


def test_master_download_checks_formal_get_precondition(db, master):
    import json
    from app.models.maintenance_project import MaintenanceProject
    uid, pid, client, url = master
    grant(db, uid, 'maintenance_general_workbook', ['data_purchase_cost', 'data_profit', 'data_customer'], 'all_projects')
    view = client.get(url + '?view=authorization')
    context = view.json()['authorization_context']
    pre = {key: context[key] for key in ('catalog_version', 'authorization_epoch', 'context')}
    project = db.get(MaintenanceProject, pid)
    project.version += 1
    db.commit()
    assert client.get(url, params={'authorization_precondition': json.dumps(pre)}).status_code == 409
    assert client.get(url, params={'view': 'authorization', 'authorization_precondition': json.dumps(pre)}).status_code == 200


def test_master_source_reassignment_is_authorized_before_loading_foreign_costs(db, master, monkeypatch):
    from app.services import maintenance_project_master_workbook as workbook
    from tests.test_maintenance_project_master_v2_editable import _make_project_with_line
    uid, pid, client, url = master
    settings = get_settings()
    monkeypatch.setattr(settings, 'authz_enabled_resources', settings.authz_enabled_resources + ['maintenance_general_workbook_write'])
    runtime.activate_configured_resources(db.get_bind())
    _, part, order, line = _make_project_with_line(db)
    db.add(MaintenanceProjectUserAssignment(assignment_id='source-validate-manager', project_id=pid,
        user_id=uid, responsibility_type='primary_manager', assigned_at=datetime.now(timezone.utc),
        assigned_by='synthetic', assignment_reason='test actual relation'))
    db.add(SysAuthorizationGrant(user_id=uid, resource='maintenance_managed_workbook', actions=['download', 'validate'],
        scope='primary_managed_projects', field_permissions=['data_purchase_cost', 'data_profit', 'data_customer'], version=1))
    db.commit()
    downloaded = client.get(url, params={'sheets': workbook.V2_SHEET_PARTS})
    assert downloaded.status_code == 200, downloaded.text
    book = load_workbook(BytesIO(downloaded.content))
    sheet = book[workbook.V2_SHEET_PARTS]
    columns = {cell.value: cell.column for cell in sheet[1]}
    rowno = sheet.max_row + 1
    for name, value in {'维保单号': order.order_no, 'PN': part.pn_std, '描述': line.description,
                        '需求数量': float(line.qty), '退货数量': float(line.return_qty or 0)}.items():
        sheet.cell(rowno, columns[name], value)
    output = BytesIO()
    book.save(output)
    book.close()
    sql = []
    def capture(conn, cursor, statement, parameters, context, executemany):
        if statement.lstrip().lower().startswith('select'):
            sql.append(statement.lower())
    event.listen(db.get_bind(), 'before_cursor_execute', capture)
    try:
        denied = client.post(url.removesuffix('.xlsx') + '/validate',
            files={'file': ('reassign.xlsx', output.getvalue())})
    finally:
        event.remove(db.get_bind(), 'before_cursor_execute', capture)
    assert denied.status_code == 403, denied.text
    assert not any('f_maintenance_line.unit_cost' in statement for statement in sql)
    db.add(SysAuthorizationGrant(user_id=uid, resource='maintenance_general_workbook_write',
        actions=['reassign_source_orders'], scope='all_projects',
        field_permissions=['data_purchase_cost', 'data_profit', 'data_customer'], version=1))
    db.commit()
    response = client.post(url.removesuffix('.xlsx') + '/validate',
        files={'file': ('reassign.xlsx', output.getvalue())})
    assert response.status_code == 200, response.text
    assert response.json()['summary']['order_reassignments'] == 1


def test_master_cost_rows_uses_exact_grant_before_query_and_checks_context(db, master, monkeypatch):
    import json
    from app.models.maintenance_project import MaintenanceProject
    from app.services import maintenance_project_master_workbook as workbook
    from tests.test_maintenance_project_master_v2_editable import _make_project_with_line
    uid, _, client, _ = master
    project, _, order, line = _make_project_with_line(db)
    settings = get_settings()
    monkeypatch.setattr(settings, 'authz_enabled_resources', settings.authz_enabled_resources + ['maintenance_row_workbench'])
    runtime.activate_configured_resources(db.get_bind())
    row = SysAuthorizationGrant(user_id=uid, resource='maintenance_row_workbench',
        actions=['read_fields'], scope='all_projects', field_permissions=['data_purchase_cost'], version=1)
    db.add(row)
    db.commit()
    url = f'/api/maintenance/projects/stable/{project.project_id}/master-workbook/rows'
    params = {'sheet': workbook.V2_SHEET_PARTS, 'page': 1, 'page_size': 1}
    statements = []
    def record(conn, cursor, sql, parameters, context, executemany):
        statements.append(sql.lower())
    event.listen(db.get_bind(), 'before_cursor_execute', record)
    try:
        denied = client.get(url, params=params)
    finally:
        event.remove(db.get_bind(), 'before_cursor_execute', record)
    assert denied.status_code == 403, denied.text
    assert not any('f_maintenance_line.unit_cost' in sql for sql in statements)
    row.actions = ['read_cost_rows']
    row.version += 1
    db.commit()
    view = client.get(url, params={**params, 'view': 'authorization'})
    assert view.status_code == 200, view.text
    assert set(view.json()) == {'locator', 'authorization_context'}
    context = view.json()['authorization_context']
    pre = {key: context[key] for key in ('catalog_version', 'authorization_epoch', 'context')}
    result = client.get(url, params={**params, 'authorization_precondition': json.dumps(pre)})
    assert result.status_code == 200, result.text
    value = result.json()
    assert value['total'] == 1 and value['page_size'] == 1
    assert value['rows'][0]['line_id'] == line.id
    assert value['rows'][0]['order_no'] == order.order_no
    assert value['rows'][0]['unit_cost_ex_tax'] is not None
    db.get(MaintenanceProject, project.project_id).version += 1
    db.commit()
    assert client.get(url, params={**params, 'authorization_precondition': json.dumps(pre)}).status_code == 409
    row.field_permissions = []
    row.version += 1
    db.commit()
    assert client.get(url, params=params).status_code == 403


def test_global_roundtrip_filters_actual_project_scope_before_whole_file_render(db, master, monkeypatch):
    import json
    from app.services import maintenance_project_master_workbook as workbook
    from tests.test_maintenance_project_master_v2_editable import _make_project_with_line
    uid, _, client, _ = master
    project, part, _, _ = _make_project_with_line(db)
    settings = get_settings()
    monkeypatch.setattr(settings, 'authz_enabled_resources', settings.authz_enabled_resources + ['maintenance_global_workbook'])
    runtime.activate_configured_resources(db.get_bind())
    row = SysAuthorizationGrant(user_id=uid, resource='maintenance_global_workbook', actions=['download'],
        scope='assigned_projects', field_permissions=['data_purchase_cost', 'data_profit', 'data_customer'], version=1)
    db.add(row)
    db.commit()
    url = '/api/maintenance/spare-part-lines.xlsx'
    params = {'range': 'custom', 'from': '2026-08-01', 'to': '2026-08-31'}
    def visible_pns(response):
        assert response.status_code == 200, response.text
        book = load_workbook(BytesIO(response.content))
        try:
            return [cells[4] for cells in book[workbook.GLOBAL_SHEET].iter_rows(min_row=2, values_only=True)]
        finally:
            book.close()
    assert part.pn_std not in visible_pns(client.get(url, params=params))
    db.add(MaintenanceProjectUserAssignment(assignment_id='global-workbook-manager', project_id=project.project_id,
        user_id=uid, responsibility_type='primary_manager', assigned_at=datetime.now(timezone.utc),
        assigned_by='synthetic', assignment_reason='actual scoped workbook relation'))
    db.commit()
    assert part.pn_std in visible_pns(client.get(url, params=params))
    context = client.get(url, params={**params, 'view': 'authorization'}).json()['authorization_context']
    pre = {key: context[key] for key in ('catalog_version', 'authorization_epoch', 'context')}
    stale = client.get(url, params={**params, 'from': '2026-08-02', 'authorization_precondition': json.dumps(pre)})
    assert stale.status_code == 409, stale.text
    row.field_permissions = ['data_purchase_cost', 'data_profit']
    row.version += 1
    db.commit()
    assert client.get(url, params=params).status_code == 403


def test_master_apply_stable_receipt_and_current_complete_replay_contract(db, master, monkeypatch):
    from app.auth import _make_token
    from app.models.system import SysUser
    from app.models.authorization import SysAuthorizationArtifact
    from app.models.maintenance_project_operations import MaintenanceProjectWorkbookOperation
    uid, pid, client, url = master
    settings = get_settings()
    monkeypatch.setattr(settings, 'authz_enabled_resources', settings.authz_enabled_resources + ['maintenance_general_workbook_write'])
    runtime.activate_configured_resources(db.get_bind())
    grant(db, uid, 'maintenance_general_workbook', ['data_purchase_cost', 'data_profit', 'data_customer'], 'all_projects')
    write = SysAuthorizationGrant(user_id=uid, resource='maintenance_general_workbook_write', actions=['apply'],
        scope='all_projects', field_permissions=['data_purchase_cost', 'data_profit', 'data_customer'], version=1)
    db.add(write)
    db.commit()
    download = client.get(url)
    assert download.status_code == 200, download.text
    apply_url = url.removesuffix('.xlsx') + '/apply'
    applied = client.post(apply_url, files={'file': ('master.xlsx', download.content)})
    assert applied.status_code == 200, applied.text
    assert applied.json()['replayed'] is False
    operation = db.scalar(select(MaintenanceProjectWorkbookOperation).where(
        MaintenanceProjectWorkbookOperation.owner_user_id == uid,
        MaintenanceProjectWorkbookOperation.operation_type == 'file_apply'))
    assert operation is not None
    binding = db.get(SysAuthorizationArtifact, ('master_workbook_operation', str(operation.id)))
    assert binding.owner_user_id == uid and binding.artifact_manifests['apply_receipt']['complete'] is True
    user = db.get(SysUser, uid)
    user.username = 'master-renamed-owner'
    db.commit()
    token, _ = _make_token('readonly', user.username, None, token_version=user.token_version, authn='sys_user')
    client.headers['Authorization'] = 'Bearer ' + token
    replay = client.post(apply_url, files={'file': ('master.xlsx', download.content)})
    assert replay.status_code == 200, replay.text
    assert replay.json()['replayed'] is True
    assert replay.json()['import_batch_id'] == applied.json()['import_batch_id']
    write.actions = []
    write.version += 1
    db.commit()
    denied = client.post(apply_url, files={'file': ('master.xlsx', download.content)})
    assert denied.status_code == 403, denied.text


@pytest.mark.parametrize('revoke_before_final_check', [False, True])
def test_master_apply_actual_mutation_final_revocation_rolls_back_and_receipt_is_immutable(db, master, monkeypatch, revoke_before_final_check):
    from sqlalchemy import text
    from sqlalchemy.orm import Session
    from app.models.maintenance_project import MaintenanceProject
    from app.models.maintenance_project_operations import MaintenanceProjectWorkbookOperation
    from app.models.system import SysUser
    from app.services.maintenance_master_apply_access import MasterWriteAuthorization
    from tests.test_maintenance_project_note_workbook import _note_cell
    uid, pid, client, url = master
    settings = get_settings()
    monkeypatch.setattr(settings, 'authz_enabled_resources', settings.authz_enabled_resources +
        ['maintenance_general_workbook_write', 'maintenance_project'])
    runtime.activate_configured_resources(db.get_bind())
    grant(db, uid, 'maintenance_general_workbook', ['data_purchase_cost', 'data_profit', 'data_customer'], 'all_projects')
    db.add(SysAuthorizationGrant(user_id=uid, resource='maintenance_general_workbook_write', actions=['apply'],
        scope='all_projects', field_permissions=['data_purchase_cost', 'data_profit', 'data_customer'], version=1))
    db.add(SysAuthorizationGrant(user_id=uid, resource='maintenance_project', actions=['update'],
        scope='all_projects', field_permissions=[], version=1))
    db.commit()
    downloaded = client.get(url)
    assert downloaded.status_code == 200, downloaded.text
    book = load_workbook(BytesIO(downloaded.content))
    _note_cell(book).value = 'canonical mutation before final gate'
    output = BytesIO()
    book.save(output)
    book.close()
    original = MasterWriteAuthorization.authorize_after
    observed = []
    def intercept(self, session):
        observed.append(session.get(MaintenanceProject, pid).project_note)
        if revoke_before_final_check:
            with Session(session.get_bind()) as revoker, revoker.begin():
                revoker.execute(text("SET LOCAL lock_timeout='1500ms'"))
                revoker.get(SysUser, uid).token_version += 1
        return original(self, session)
    monkeypatch.setattr(MasterWriteAuthorization, 'authorize_after', intercept)
    endpoint = url.removesuffix('.xlsx') + '/apply'
    result = client.post(endpoint, files={'file': ('note.xlsx', output.getvalue())})
    assert observed == ['canonical mutation before final gate']
    db.expire_all()
    project = db.get(MaintenanceProject, pid)
    if revoke_before_final_check:
        assert result.status_code == 401, result.text
        assert project.project_note is None
        assert db.scalar(select(MaintenanceProjectWorkbookOperation.id).where(
            MaintenanceProjectWorkbookOperation.owner_user_id == uid)) is None
    else:
        assert result.status_code == 200, result.text
        assert project.project_note == 'canonical mutation before final gate'
        project.project_note = 'later unrelated note'
        project.version += 1
        db.commit()
        replay = client.post(endpoint, files={'file': ('note.xlsx', output.getvalue())})
        assert replay.status_code == 200, replay.text
        assert replay.json()['replayed'] is True
        assert replay.json()['import_batch_id'] == result.json()['import_batch_id']
        db.expire_all()
        assert db.get(MaintenanceProject, pid).project_note == 'later unrelated note'


def test_master_reassignment_write_and_receipt_keep_all_original_parents(db, master, monkeypatch):
    from app.models.maintenance_source_assignment import MaintenanceSourceOrderAssignment
    from app.models.maintenance_project import MaintenanceProject
    from app.services import maintenance_project_master_workbook as workbook
    from tests.test_maintenance_project_master_v2_editable import _make_project_with_line
    uid, pid, client, url = master
    settings = get_settings()
    monkeypatch.setattr(settings, 'authz_enabled_resources', settings.authz_enabled_resources + ['maintenance_general_workbook_write'])
    runtime.activate_configured_resources(db.get_bind())
    old, part, order, line = _make_project_with_line(db)
    order.linked_sales_order_no = None
    db.add(MaintenanceProjectUserAssignment(assignment_id='source-apply-manager', project_id=pid,
        user_id=uid, responsibility_type='primary_manager', assigned_at=datetime.now(timezone.utc),
        assigned_by='synthetic', assignment_reason='actual target responsibility'))
    db.add(SysAuthorizationGrant(user_id=uid, resource='maintenance_managed_workbook', actions=['download', 'apply'],
        scope='primary_managed_projects', field_permissions=['data_purchase_cost', 'data_profit', 'data_customer'], version=1))
    db.add(SysAuthorizationGrant(user_id=uid, resource='maintenance_general_workbook_write', actions=['reassign_source_orders'],
        scope='all_projects', field_permissions=['data_purchase_cost', 'data_profit', 'data_customer'], version=1))
    db.commit()
    response = client.get(url, params={'sheets': workbook.V2_SHEET_PARTS})
    assert response.status_code == 200, response.text
    book = load_workbook(BytesIO(response.content))
    sheet = book[workbook.V2_SHEET_PARTS]
    columns = {cell.value: cell.column for cell in sheet[1]}
    rowno = sheet.max_row + 1
    for name, value in {'维保单号': order.order_no, 'PN': part.pn_std, '描述': line.description,
                        '需求数量': float(line.qty), '退货数量': float(line.return_qty or 0)}.items():
        sheet.cell(rowno, columns[name], value)
    output = BytesIO()
    book.save(output)
    book.close()
    endpoint = url.removesuffix('.xlsx') + '/apply'
    applied = client.post(endpoint, files={'file': ('reassign.xlsx', output.getvalue())})
    assert applied.status_code == 200, applied.text
    assert applied.json()['order_reassignments'] == 1
    current = db.scalar(select(MaintenanceSourceOrderAssignment).where(
        MaintenanceSourceOrderAssignment.source_order_id == order.raw_order_id,
        MaintenanceSourceOrderAssignment.is_active.is_(True)))
    assert current.project_id == pid
    replay = client.post(endpoint, files={'file': ('reassign.xlsx', output.getvalue())})
    assert replay.status_code == 200 and replay.json()['replayed'] is True, replay.text
    outside = MaintenanceProject(project_id='post-receipt-outside', project_code='OUTSIDE-CLOSURE',
        display_name='outside receipt closure', lifecycle_status='ongoing')
    db.add(outside)
    db.flush()
    current.is_active = False
    current.archived_by = "synthetic"
    current.archived_at = datetime.now(timezone.utc)
    current.version += 1
    db.flush()
    db.add(MaintenanceSourceOrderAssignment(assignment_id="outside-receipt-generation",
        source_order_id=order.raw_order_id, project_id=outside.project_id,
        created_by="synthetic"))
    db.commit()
    assert client.post(endpoint, files={'file': ('reassign.xlsx', output.getvalue())}).status_code == 403


@pytest.mark.parametrize('revoke_at_final_gate', [False, True])
def test_master_v1_actual_cost_write_receipt_and_final_revocation(db, master, monkeypatch, revoke_at_final_gate):
    from decimal import Decimal
    from sqlalchemy.orm import Session
    from app.models.system import SysUser
    from app.models.maintenance import MaintenanceManualCostOverride as MaintenanceCostOverride
    from app.models.maintenance_project_operations import MaintenanceProjectWorkbookOperation
    from app.services import maintenance_project_master_workbook as workbook
    from app.services.maintenance_master_apply_access import MasterWriteAuthorization
    from tests.test_maintenance_project_master_v2_editable import _make_project_with_line
    uid, _, client, _ = master
    settings = get_settings()
    monkeypatch.setattr(settings, 'maintenance_project_master_v2_enabled', False)
    monkeypatch.setattr(settings, 'authz_enabled_resources', settings.authz_enabled_resources + ['maintenance_general_workbook_write'])
    runtime.activate_configured_resources(db.get_bind())
    project, _, _, line = _make_project_with_line(db)
    pid, line_id = project.project_id, line.id
    fields = ['data_purchase_cost', 'data_profit', 'data_customer']
    grant(db, uid, 'maintenance_general_workbook', fields, 'all_projects')
    db.add(SysAuthorizationGrant(user_id=uid, resource='maintenance_general_workbook_write',
        actions=['apply'], scope='all_projects', field_permissions=fields, version=1))
    db.commit()
    url = f'/api/maintenance/projects/stable/{pid}/master-workbook.xlsx'
    downloaded = client.get(url, params={'sheets': workbook.SHEET_PARTS})
    assert downloaded.status_code == 200, downloaded.text
    book = load_workbook(BytesIO(downloaded.content))
    sheet = book[workbook.SHEET_PARTS]
    sheet.cell(2, 12, 123)
    sheet.cell(2, 13).value = None
    sheet.cell(2, 14, 'actual V1 canonical cost')
    output = BytesIO()
    book.save(output)
    book.close()
    real_authorize = MasterWriteAuthorization.authorize_after
    def final_authorize(self, current_db):
        actual = current_db.scalar(select(MaintenanceCostOverride).where(MaintenanceCostOverride.line_id == line_id))
        assert actual is not None and actual.unit_cost_ex_tax == Decimal('123')
        if revoke_at_final_gate:
            with Session(current_db.get_bind()) as other:
                user = other.get(SysUser, uid)
                user.token_version += 1
                other.commit()
        return real_authorize(self, current_db)
    monkeypatch.setattr(MasterWriteAuthorization, 'authorize_after', final_authorize)
    endpoint = url.removesuffix('.xlsx') + '/apply'
    applied = client.post(endpoint, files={'file': ('v1.xlsx', output.getvalue())})
    assert applied.status_code == (401 if revoke_at_final_gate else 200), applied.text
    db.expire_all()
    actual = db.scalar(select(MaintenanceCostOverride).where(MaintenanceCostOverride.line_id == line_id))
    if revoke_at_final_gate:
        assert actual is None
        assert db.scalar(select(MaintenanceProjectWorkbookOperation.id).where(
            MaintenanceProjectWorkbookOperation.owner_user_id == uid)) is None
    else:
        assert actual.unit_cost_ex_tax == Decimal('123')
        replay = client.post(endpoint, files={'file': ('v1.xlsx', output.getvalue())})
        assert replay.status_code == 200 and replay.json()['replayed'] is True, replay.text
        assert replay.json()['import_batch_id'] == applied.json()['import_batch_id']
        write = db.scalar(select(SysAuthorizationGrant).where(SysAuthorizationGrant.user_id == uid,
            SysAuthorizationGrant.resource == 'maintenance_general_workbook_write'))
        write.actions = []
        write.version += 1
        db.commit()
        assert client.post(endpoint, files={'file': ('v1.xlsx', output.getvalue())}).status_code == 403


def test_master_apply_precondition_binds_current_command_and_unknown_view_never_writes(db, master, monkeypatch):
    import json
    from app.models.maintenance_project import MaintenanceProject
    from app.models.maintenance_project_operations import MaintenanceProjectWorkbookOperation
    uid, pid, client, url = master
    settings = get_settings()
    monkeypatch.setattr(settings, 'authz_enabled_resources', settings.authz_enabled_resources + ['maintenance_general_workbook_write'])
    runtime.activate_configured_resources(db.get_bind())
    fields = ['data_purchase_cost', 'data_profit', 'data_customer']
    grant(db, uid, 'maintenance_general_workbook', fields, 'all_projects')
    db.add(SysAuthorizationGrant(user_id=uid, resource='maintenance_general_workbook_write',
        actions=['apply'], scope='all_projects', field_permissions=fields, version=1))
    db.commit()
    content = client.get(url).content
    endpoint = url.removesuffix('.xlsx') + '/apply'
    assert client.post(endpoint + '?view=typo', files={'file': ('master.xlsx', content)}).status_code == 422
    assert db.scalar(select(MaintenanceProjectWorkbookOperation.id).where(
        MaintenanceProjectWorkbookOperation.owner_user_id == uid)) is None
    view = client.post(endpoint + '?view=authorization', files={'file': ('master.xlsx', content)})
    assert view.status_code == 200, view.text
    context = view.json()['authorization_context']
    precondition = {key: context[key] for key in ('catalog_version', 'authorization_epoch', 'context')}
    project = db.get(MaintenanceProject, pid)
    project.version += 1
    db.commit()
    stale = client.post(endpoint, files={'file': ('master.xlsx', content)},
        data={'authorization_precondition': json.dumps(precondition)})
    assert stale.status_code == 409, stale.text
    assert db.scalar(select(MaintenanceProjectWorkbookOperation.id).where(
        MaintenanceProjectWorkbookOperation.owner_user_id == uid)) is None
    fresh = client.post(endpoint + '?view=authorization', files={'file': ('master.xlsx', content)})
    assert fresh.status_code == 200, fresh.text
    context = fresh.json()['authorization_context']
    precondition = {key: context[key] for key in ('catalog_version', 'authorization_epoch', 'context')}
    applied = client.post(endpoint, files={'file': ('master.xlsx', content)},
        data={'authorization_precondition': json.dumps(precondition)})
    assert applied.status_code == 200, applied.text
