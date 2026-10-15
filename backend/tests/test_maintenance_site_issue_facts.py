"""Contract for the shared locked fact editor and existing workbook adapters."""

from datetime import date
from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal
import io
import threading
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import event, func, select, text

from app.db import SessionLocal

from app.models.dimensions import DimPart
from app.models.maintenance import FMaintenanceLine, FMaintenanceOrder
from app.models.maintenance_project import MaintenanceProject
from app.models.maintenance_source_assignment import MaintenanceSourceOrderAssignment
from app.models.system import SysImportBatch
from app.models.maintenance_project_operations import (
    MaintenanceProjectOperationAudit,
    MaintenanceSiteIssue,
    MaintenanceSiteIssueDeliverySource,
    MaintenanceSiteIssueLine,
    MaintenanceSiteIssueReturnEvent,
)
from app.services import maintenance_manual_site_issue as manual
from app.services import maintenance_project_master_workbook as master
from app.services import maintenance_project_operations as operations


def _facts(db, *, source="page_manual"):
    project = MaintenanceProject(
        project_id=str(uuid4()), project_code=f"FACT-{uuid4().hex[:8]}",
        display_name="领用事实契约项目", lifecycle_status="ongoing",
    )
    part = DimPart(pn_std=f"FACT-{uuid4().hex[:8]}")
    db.add_all([project, part])
    db.flush()
    issue = MaintenanceSiteIssue(
        issue_id=str(uuid4()), project_id=project.project_id,
        issue_no=f"FACT-ISSUE-{uuid4().hex[:8]}", issue_date=date(2026, 9, 20),
        receiver="接收人", issued_by="发出人", site_location="现场",
        raw_status="confirmed", status_mapping_state="mapped",
        normalized_status="confirmed", status_mapping_version="contract-v1",
        source=source, import_batch_id=str(uuid4()) if source == "workbook" else None,
        version=1,
    )
    db.add(issue)
    db.flush()
    lines = [MaintenanceSiteIssueLine(
        issue_line_id=str(uuid4()), issue_id=issue.issue_id, line_no=i,
        part_id=part.id, pn=part.pn_std, quantity=Decimal("2.000"),
        algorithm_version="contract-v1", is_active=True, version=1,
    ) for i in (1, 2)]
    db.add_all(lines)
    operations.get_or_create_workbook_state(db, project_id=project.project_id)
    db.commit()
    return project, issue, lines, part


def _site_workbook(db, project):
    book = load_workbook(io.BytesIO(master.build_project_master_v2(
        db, project_id=project.project_id, sheets=(master.V2_SHEET_SITE,),
    )))
    sheet = book[master.V2_SHEET_SITE]
    columns = {cell.value: cell.column for cell in sheet[1] if cell.value}
    rows = {
        sheet.cell(row, columns["实体ID"]).value: row
        for row in range(2, sheet.max_row + 1)
        if sheet.cell(row, columns["实体ID"]).value
    }
    return book, sheet, columns, rows


def _apply_site_workbook(db, project, book):
    content = io.BytesIO()
    book.save(content)
    plan = master.validate_project_master_v2(
        db, project_id=project.project_id, data=content.getvalue(),
    )
    assert not plan.conflicts
    return master.apply_project_master_v2(
        db, plan, operated_by="tester", import_batch_id=str(uuid4()),
    )


def _demand_prices(db, *, project, part, prices):
    batch = SysImportBatch(filename="synthetic-facts.xlsx", file_type="maintenance",
                           file_hash=uuid4().hex, status="success")
    db.add(batch)
    db.flush()
    source_lines = []
    for number, unit_cost in prices:
        order = FMaintenanceOrder(
            raw_order_id=str(uuid4()), order_no=number, order_date=date(2026, 9, 20),
            project_raw=project.display_name, data_status="已生效", import_batch_id=batch.id,
        )
        db.add(order)
        db.flush()
        line = FMaintenanceLine(
            raw_line_id=str(uuid4()), order_id=order.id, line_no=1,
            part_id=part.id, pn_raw=part.pn_std, pn_std=part.pn_std,
            qty=Decimal("2"), return_qty=Decimal("0"),
            unit_cost=unit_cost, unit_cost_ex_tax=unit_cost,
            unit_cost_inc_tax=unit_cost * Decimal("1.13"),
            cost_amount=unit_cost * 2, cost_amount_ex_tax=unit_cost * 2,
            cost_amount_inc_tax=unit_cost * Decimal("2.26"),
            cost_source="direct", cost_tax_basis="ex", confidence="high",
            import_batch_id=batch.id,
        )
        db.add(line)
        source_lines.append(line)
        db.add(MaintenanceSourceOrderAssignment(
            assignment_id=str(uuid4()), project_id=project.project_id,
            source_order_id=order.raw_order_id, is_active=True, created_by="tester",
        ))
    db.commit()
    return source_lines


@pytest.mark.parametrize(("patch", "line_version", "issue_version", "events"), [
    ({"quantity": "3"}, 3, 3, 1),
    ({"remark": "核对后更正"}, 2, 2, 0),
])
def test_legacy_manual_keeps_versions_and_audit(db, patch, line_version, issue_version, events):
    project, issue, lines, part = _facts(db)
    target, untouched = lines
    result = manual.patch_manual_site_issue(
        db, issue_id=issue.issue_id, project_id=project.project_id, version=1,
        idempotency_key=f"facts-legacy-{uuid4()}", reason="契约冻结", operated_by="tester",
        lines=[{
            "issue_line_id": line.issue_line_id, "part_id": part.id,
            "quantity": patch.get("quantity", "2") if line is target else "2",
            "remark": patch.get("remark") if line is target else None,
        } for line in lines],
    )
    db.commit()
    assert result["version"] == issue_version
    assert target.version == line_version
    assert untouched.version == 1
    assert db.scalar(select(func.count()).select_from(MaintenanceSiteIssueReturnEvent)) == events
    actions = list(db.scalars(select(MaintenanceProjectOperationAudit.action)))
    assert actions.count("correct") == 1
    assert actions.count("return_obligation_corrected") == (1 if events else 0)


def test_legacy_v1_remark_preserves_cas_and_does_not_reprice(db, monkeypatch):
    project, issue, lines, _part = _facts(db, source="legacy")
    def unexpected_reprice(*args, **kwargs):
        pytest.fail("V1 historically does not reprice a remark edit")
    monkeypatch.setattr(manual.maintenance_consumption_cost, "resolve_lines", unexpected_reprice)
    master.apply(
        db, master.MasterPlan(project_id=project.project_id, site_flags=(
            master.SiteReturnFlag(issue_line_id=lines[0].issue_line_id,
                                  no_return=None, remark="V1 备注"),
        )), operated_by="tester", import_batch_id=str(uuid4()),
    )
    assert lines[0].remark == "V1 备注"
    assert (issue.version, lines[0].version) == (1, 1)
    assert db.scalar(select(func.count()).select_from(MaintenanceProjectOperationAudit)) == 0


@pytest.mark.parametrize("entrypoint", ["v1", "v2"])
def test_legacy_workbook_two_line_edits_keep_one_issue_correction(db, entrypoint):
    """The workbook transaction emits one issue event for both final line facts."""
    project, issue, lines, _part = _facts(db, source="workbook")
    if entrypoint == "v1":
        master.apply(
            db, master.MasterPlan(project_id=project.project_id, site_flags=(
                master.SiteReturnFlag(issue_line_id=lines[0].issue_line_id,
                                      no_return=False, quantity=Decimal("3")),
                master.SiteReturnFlag(issue_line_id=lines[1].issue_line_id,
                                      no_return=True, quantity=Decimal("4")),
            )), operated_by="tester", import_batch_id=str(uuid4()),
        )
    else:
        book, sheet, columns, rows = _site_workbook(db, project)
        for line, quantity, flag in ((lines[0], 3, "是"), (lines[1], 4, "否")):
            row = rows[line.issue_line_id]
            sheet.cell(row, columns["领用数量"], quantity)
            sheet.cell(row, columns["是否应返还"], flag)
        _apply_site_workbook(db, project, book)
    assert [(line.quantity, line.no_return, line.version) for line in lines] == [
        (Decimal("3"), False, 2), (Decimal("4"), True, 2),
    ]
    assert issue.version == 2
    events = list(db.scalars(select(MaintenanceSiteIssueReturnEvent)))
    assert len(events) == 1 and events[0].issue_version == 2
    assert {change["after"]["issue_line_id"] for change in events[0].payload["changes"]} == {
        line.issue_line_id for line in lines
    }
    actions = list(db.scalars(select(MaintenanceProjectOperationAudit.action)))
    assert actions == ["return_obligation_corrected", "return_obligation_corrected"]
    assert operations.get_or_create_workbook_state(db, project_id=project.project_id).revision == 1


def test_legacy_v2_void_and_edit_share_one_final_correction(db):
    project, issue, lines, _part = _facts(db, source="workbook")
    removed, kept = lines
    book, sheet, columns, rows = _site_workbook(db, project)
    sheet.cell(rows[kept.issue_line_id], columns["领用数量"], 3)
    sheet.cell(rows[kept.issue_line_id], columns["是否应返还"], "是")
    sheet.delete_rows(rows[removed.issue_line_id])
    _apply_site_workbook(db, project, book)
    assert (removed.is_active, removed.version, removed.quantity) == (False, 2, Decimal("2"))
    assert (kept.is_active, kept.version, kept.quantity, kept.no_return) == (True, 2, Decimal("3"), False)
    assert (issue.normalized_status, issue.version) == ("confirmed", 2)
    events = list(db.scalars(select(MaintenanceSiteIssueReturnEvent)))
    assert len(events) == 1 and events[0].issue_version == 2
    corrected = {change["after"]["issue_line_id"]: change["after"]
                 for change in events[0].payload["changes"]}
    assert corrected[removed.issue_line_id]["is_active"] is False
    assert Decimal(corrected[kept.issue_line_id]["quantity"]) == Decimal("3")
    actions = list(db.scalars(select(MaintenanceProjectOperationAudit.action)))
    assert actions.count("VOID") == 1
    assert actions.count("return_obligation_corrected") == 2
    assert len(actions) == 3
    assert operations.get_or_create_workbook_state(db, project_id=project.project_id).revision == 1


def test_legacy_manual_rename_with_lines_keeps_old_number_price_order(db):
    """The complete manual command historically prices before assigning its header."""
    project, issue, lines, part = _facts(db)
    old_no, new_no = "WBDD-20260920-0101", "WBDD-20260920-0102"
    issue.issue_no = old_no
    _demand_prices(db, project=project, part=part,
                   prices=((old_no, Decimal("10")), (new_no, Decimal("20"))))
    manual.maintenance_consumption_cost.resolve_lines(db, lines=[(issue.issue_date, line) for line in lines])
    db.commit()
    result = manual.patch_manual_site_issue(
        db, issue_id=issue.issue_id, project_id=project.project_id, version=1,
        idempotency_key=f"facts-manual-rename-{uuid4()}", issue_no=new_no,
        lines=[{"issue_line_id": line.issue_line_id, "part_id": part.id,
                "quantity": "3" if line is lines[0] else "2"} for line in lines],
        operated_by="tester", reason="冻结人工整单改号取价顺序",
    )
    db.commit()
    assert issue.issue_no == new_no and result["version"] == 3
    assert [line.unit_cost_ex_tax for line in lines] == [Decimal("10"), Decimal("10")]
    assert [line.cost_amount_ex_tax for line in lines] == [Decimal("30"), Decimal("20")]
    assert [line.version for line in lines] == [3, 1]
    assert all(line.reference_samples[0]["document_no"] == old_no for line in lines)
    assert db.scalar(select(func.count()).select_from(MaintenanceSiteIssueReturnEvent)) == 1


def test_legacy_manual_complete_lines_reprice_unchanged_line(db):
    project, issue, lines, part = _facts(db)
    source_lines = _demand_prices(db, project=project, part=part,
                                  prices=((issue.issue_no, Decimal("10")),))
    manual.maintenance_consumption_cost.resolve_lines(db, lines=[(issue.issue_date, line) for line in lines])
    db.commit()
    source = source_lines[0]
    source.unit_cost = source.unit_cost_ex_tax = Decimal("20")
    source.unit_cost_inc_tax = Decimal("22.60")
    source.cost_amount = source.cost_amount_ex_tax = Decimal("40")
    source.cost_amount_inc_tax = Decimal("45.20")
    db.commit()
    manual.patch_manual_site_issue(
        db, issue_id=issue.issue_id, project_id=project.project_id, version=1,
        idempotency_key=f"facts-manual-complete-{uuid4()}",
        lines=[{"issue_line_id": line.issue_line_id, "part_id": part.id,
                "quantity": "2", "remark": "核对首行" if line is lines[0] else None} for line in lines],
        operated_by="tester", reason="完整整单提交仍逐行取价",
    )
    db.commit()
    assert [line.unit_cost_ex_tax for line in lines] == [Decimal("20"), Decimal("20")]
    assert [line.cost_amount_ex_tax for line in lines] == [Decimal("40"), Decimal("40")]
    assert [line.version for line in lines] == [3, 2]
    assert (issue.version, lines[1].remark, lines[1].quantity) == (2, None, Decimal("2"))
    assert db.scalar(select(func.count()).select_from(MaintenanceSiteIssueReturnEvent)) == 0
    assert list(db.scalars(select(MaintenanceProjectOperationAudit.action))) == ["correct"]


def test_locked_fact_patch_keeps_omitted_lines_and_invalidates_cas_with_audit(db):
    from app.services.maintenance_site_issue_facts import apply_locked_site_issue_edits

    project, issue, lines, _part = _facts(db)
    locked_project = operations._lock_project_for_fact_write(db, project.project_id)
    db.refresh(issue, with_for_update=True)
    for line in lines:
        db.refresh(line, with_for_update=True)
    result = apply_locked_site_issue_edits(
        db, project=locked_project, issue=issue,
        line_changes={lines[0].issue_line_id: {"remark": "只改这一行", "serial_number": "SN-1"}},
        header_changes=None, operated_by="tester", reason="核实领用事实",
    )
    assert result.changed_line_ids == (lines[0].issue_line_id,)
    assert result.issue_version == 2
    assert result.line_versions == {lines[0].issue_line_id: 2}
    assert result.before["lines"][lines[0].issue_line_id]["remark"] is None
    assert result.after["lines"][lines[0].issue_line_id]["remark"] == "只改这一行"
    assert (lines[1].version, lines[1].is_active, lines[1].quantity) == (1, True, Decimal("2"))
    assert db.scalar(select(func.count()).select_from(MaintenanceSiteIssueReturnEvent)) == 0
    audit = db.scalar(select(MaintenanceProjectOperationAudit).where(
        MaintenanceProjectOperationAudit.action == "site_issue_facts_edited"))
    assert audit.operated_by == "tester" and audit.reason == "核实领用事实"
    state = operations.get_or_create_workbook_state(db, project_id=project.project_id)
    assert state.revision == 0  # caller owns the revision and final commit
    db.rollback()
    assert lines[0].remark is None and lines[0].version == 1


def test_locked_quantity_records_one_correction_without_double_bump(db):
    from app.services.maintenance_site_issue_facts import apply_locked_site_issue_edits

    project, issue, lines, _part = _facts(db)
    line = lines[0]
    line.manual_unit_cost = Decimal("10.00")
    line.manual_unit_cost_inc_tax = Decimal("11.30")
    line.manual_evidence = "人工核实价"
    db.commit()
    result = apply_locked_site_issue_edits(
        db, project=project, issue=issue, line_changes={line.issue_line_id: {"quantity": Decimal("3")}},
        header_changes=None, operated_by="tester", reason="数量核对",
    )
    assert (issue.version, line.version) == (2, 2)
    assert result.after["lines"][line.issue_line_id]["cost_amount_ex_tax"] == "30.00"
    assert result.after["lines"][line.issue_line_id]["cost_amount_inc_tax"] == "33.90"
    correction = db.scalar(select(MaintenanceSiteIssueReturnEvent))
    assert correction.issue_version == 2
    assert len(correction.payload["changes"]) == 1
    state = operations.get_or_create_workbook_state(db, project_id=project.project_id, lock=True)
    operations.bump_locked_workbook_revision(db, state=state)
    operations.bump_locked_workbook_revision(db, state=state)
    db.commit()
    assert state.revision == 1


@pytest.mark.parametrize("patch", [{"is_active": False}, {"quantity": Decimal("0")}, {"part_id": 999999}, {"no_return": "是"}])
def test_locked_fact_editor_rejects_invalid_intent_before_any_assignment(db, patch):
    from app.services.maintenance_site_issue_facts import apply_locked_site_issue_edits, SiteIssueEditError

    project, issue, lines, _part = _facts(db)
    with pytest.raises(SiteIssueEditError):
        apply_locked_site_issue_edits(
            db, project=project, issue=issue,
            line_changes={lines[0].issue_line_id: {"remark": "不得先写"}, lines[1].issue_line_id: patch},
            header_changes={"receiver": "不得先写"}, operated_by="tester", reason="非法意图",
        )
    assert issue.receiver == "接收人" and all(line.remark is None for line in lines)
    assert (issue.version, lines[0].version, lines[1].version) == (1, 1, 1)
    assert db.scalar(select(func.count()).select_from(MaintenanceProjectOperationAudit)) == 0


def test_locked_fact_editor_does_not_restore_or_edit_foreign_lines(db):
    from app.services.maintenance_site_issue_facts import apply_locked_site_issue_edits, SiteIssueEditError

    project, issue, lines, _part = _facts(db)
    other_project, _other_issue, other_lines, _other_part = _facts(db)
    for target_project, target_line in ((project, other_lines[0]), (other_project, lines[0])):
        with pytest.raises(SiteIssueEditError):
            apply_locked_site_issue_edits(
                db, project=target_project, issue=issue,
                line_changes={target_line.issue_line_id: {"remark": "越界"}},
                header_changes=None, operated_by="tester", reason="越界",
            )
    lines[0].is_active = False
    db.commit()
    with pytest.raises(SiteIssueEditError):
        apply_locked_site_issue_edits(
            db, project=project, issue=issue,
            line_changes={lines[0].issue_line_id: {"remark": "不得复活"}},
            header_changes=None, operated_by="tester", reason="作废优先",
        )
    assert lines[0].is_active is False and lines[0].remark is None


@pytest.mark.parametrize("source", ["site_issue_v2", "direct_api"])
def test_locked_fact_editor_preserves_native_source_workflow_guards(db, source):
    from app.services.maintenance_site_issue_facts import apply_locked_site_issue_edits, SiteIssueEditError

    project, issue, lines, _part = _facts(db, source=source)
    with pytest.raises(SiteIssueEditError):
        apply_locked_site_issue_edits(
            db, project=project, issue=issue, line_changes={lines[0].issue_line_id: {"remark": "越过原生守卫"}},
            header_changes=None, operated_by="tester", reason="来源守卫",
        )
    assert lines[0].remark is None and issue.version == 1


@pytest.mark.parametrize("header_change", [False, True])
def test_locked_fact_editor_rejects_delivery_backed_legacy_line(db, header_change):
    from app.services.maintenance_site_issue_facts import apply_locked_site_issue_edits, SiteIssueEditError

    project, issue, lines, part = _facts(db, source="legacy")
    delivery = MaintenanceSiteIssueDeliverySource(
        delivery_line_id=str(uuid4()), adapter_key="synthetic_delivery_v1",
        project_id=project.project_id, source_order_id=f"SOURCE-{uuid4()}",
        source_line_id=f"LINE-{uuid4()}", delivery_no="CKD-SYNTHETIC",
        delivery_date=date(2026, 9, 20), part_id=part.id, pn=part.pn_std,
        delivered_quantity=Decimal("2"), mapping_state="ready", mapping_version="contract-v1",
    )
    db.add(delivery)
    db.flush()
    lines[1].delivery_line_id = delivery.delivery_line_id
    db.commit()
    with pytest.raises(SiteIssueEditError, match="发货关联"):
        apply_locked_site_issue_edits(
            db, project=project, issue=issue,
            line_changes=({lines[0].issue_line_id: {"remark": "不能先写"}} if header_change else
                          {lines[1].issue_line_id: {"remark": "不能先写"}}),
            header_changes={"issue_date": date(2026, 9, 21)} if header_change else None,
            operated_by="tester", reason="发货来源守卫",
        )
    assert issue.issue_date == date(2026, 9, 20) and all(line.remark is None for line in lines)


def test_locked_noop_does_not_change_version_audit_event_or_revision(db):
    from app.services.maintenance_site_issue_facts import apply_locked_site_issue_edits

    project, issue, lines, _part = _facts(db)
    result = apply_locked_site_issue_edits(
        db, project=project, issue=issue, line_changes={lines[0].issue_line_id: {"quantity": Decimal("2"), "remark": None}},
        header_changes={"issue_date": issue.issue_date}, operated_by="tester", reason="无变化核对",
    )
    assert result.changed_line_ids == () and result.line_versions == {lines[0].issue_line_id: 1}
    assert result.before == result.after
    assert (issue.version, lines[0].version) == (1, 1)
    assert db.scalar(select(func.count()).select_from(MaintenanceProjectOperationAudit)) == 0
    assert db.scalar(select(func.count()).select_from(MaintenanceSiteIssueReturnEvent)) == 0


def test_locked_part_change_retires_manual_price_evidence(db):
    from app.services.maintenance_site_issue_facts import apply_locked_site_issue_edits

    project, issue, lines, _part = _facts(db)
    part = DimPart(pn_std=f"NEW-FACT-{uuid4().hex[:8]}")
    db.add(part)
    line = lines[0]
    line.manual_unit_cost = Decimal("10")
    line.manual_unit_cost_inc_tax = Decimal("11.30")
    line.manual_evidence = "属于旧型号的凭据"
    db.commit()
    apply_locked_site_issue_edits(
        db, project=project, issue=issue, line_changes={line.issue_line_id: {"part_id": part.id, "pn": part.pn_std}},
        header_changes=None, operated_by="tester", reason="更正型号",
    )
    assert line.part_id == part.id and line.pn == part.pn_std
    assert (line.manual_unit_cost, line.manual_unit_cost_inc_tax, line.manual_evidence) == (None, None, None)
    assert line.unit_cost_ex_tax is None and line.cost_amount_ex_tax is None
    assert (issue.version, line.version) == (2, 2)


def test_locked_header_changes_reprice_all_lines_without_double_obligation_event(db):
    from app.services.maintenance_site_issue_facts import apply_locked_site_issue_edits

    project, issue, lines, _part = _facts(db)
    result = apply_locked_site_issue_edits(
        db, project=project, issue=issue, line_changes={},
        header_changes={"issue_date": date(2026, 9, 21), "site_location": None},
        operated_by="tester", reason="日期核对",
    )
    assert issue.version == 2 and issue.issue_date == date(2026, 9, 21) and issue.site_location is None
    assert result.changed_line_ids == tuple(sorted(line.issue_line_id for line in lines))
    assert all(line.version == 2 for line in lines)
    assert db.scalar(select(func.count()).select_from(MaintenanceSiteIssueReturnEvent)) == 0


def test_locked_rename_reprices_against_new_number_and_keeps_amount_checks_valid(db):
    """Renaming A→B must use B's same-PN cost, including a quantity edit."""
    from app.services.maintenance_site_issue_facts import apply_locked_site_issue_edits

    project, issue, lines, part = _facts(db)
    before_no, after_no = "WBDD-20260920-0001", "WBDD-20260920-0002"
    issue.issue_no = before_no
    _demand_prices(db, project=project, part=part,
                   prices=((before_no, Decimal("10.00")), (after_no, Decimal("20.00"))))
    manual.maintenance_consumption_cost.resolve_lines(db, lines=[(issue.issue_date, line) for line in lines])
    db.commit()
    assert all(line.unit_cost_ex_tax == Decimal("10.00") for line in lines)
    result = apply_locked_site_issue_edits(
        db, project=project, issue=issue,
        line_changes={lines[0].issue_line_id: {"quantity": Decimal("3.000")}},
        header_changes={"issue_no": after_no}, operated_by="tester", reason="改号后重取需求成本",
    )
    assert issue.issue_no == after_no
    assert [line.unit_cost_ex_tax for line in lines] == [Decimal("20.00"), Decimal("20.00")]
    assert [line.cost_amount_ex_tax for line in lines] == [Decimal("60.00"), Decimal("40.00")]
    assert result.after["lines"][lines[0].issue_line_id]["reference_samples"][0]["document_no"] == after_no


def test_locked_correction_failure_rolls_back_cost_fact_event_audit_and_revision(db, monkeypatch):
    from app.services import maintenance_site_issue_facts as facts

    project, issue, lines, _part = _facts(db)
    original = facts.requirements.record_corrections
    def fail_after_correction(*args, **kwargs):
        original(*args, **kwargs)
        raise RuntimeError("模拟更正后失败")
    monkeypatch.setattr(facts.requirements, "record_corrections", fail_after_correction)
    with pytest.raises(RuntimeError, match="模拟更正后失败"):
        facts.apply_locked_site_issue_edits(
            db, project=project, issue=issue, line_changes={lines[0].issue_line_id: {"quantity": Decimal("3")}},
            header_changes={"site_location": "新地址"}, operated_by="tester", reason="回滚契约",
        )
    db.rollback()
    assert lines[0].quantity == Decimal("2") and lines[0].algorithm_version == "contract-v1"
    assert (issue.site_location, issue.version, lines[0].version) == ("现场", 1, 1)
    assert db.scalar(select(func.count()).select_from(MaintenanceSiteIssueReturnEvent)) == 0
    assert db.scalar(select(func.count()).select_from(MaintenanceProjectOperationAudit)) == 0
    state = operations.get_or_create_workbook_state(db, project_id=project.project_id)
    assert state.revision == 0


def test_locked_editor_and_manual_command_serialize_then_reject_stale_cas(db):
    """Two real PG sessions meet at workbook state before either fact lock."""
    from app.services.maintenance_site_issue_facts import apply_locked_site_issue_edits

    project, issue, lines, _part = _facts(db)
    project_id, issue_id, line_id = project.project_id, issue.issue_id, lines[0].issue_line_id
    db.rollback()
    held = threading.Event()
    release = threading.Event()
    waiter_reached_state = threading.Event()
    lock_trace = []

    def canonical_writer():
        with SessionLocal() as session:
            connection = session.connection()
            connection.execute(text("SET LOCAL lock_timeout = '8s'"))
            connection.execute(text("SET LOCAL statement_timeout = '15s'"))
            locked_project = operations._lock_project_for_fact_write(session, project_id)
            locked_issue = session.scalar(select(MaintenanceSiteIssue).where(
                MaintenanceSiteIssue.issue_id == issue_id).with_for_update().execution_options(populate_existing=True))
            session.scalars(select(MaintenanceSiteIssueLine).where(
                MaintenanceSiteIssueLine.issue_id == issue_id).order_by(
                    MaintenanceSiteIssueLine.issue_line_id).with_for_update().execution_options(populate_existing=True)).all()
            held.set()
            assert release.wait(10)
            apply_locked_site_issue_edits(
                session, project=locked_project, issue=locked_issue,
                line_changes={line_id: {"remark": "先提交的事实"}},
                header_changes=None, operated_by="writer", reason="并发契约",
            )
            operations.bump_workbook_revision(session, project_id=project_id)
            session.commit()

    def manual_waiter():
        with SessionLocal() as session:
            connection = session.connection()
            connection.execute(text("SET LOCAL lock_timeout = '8s'"))
            connection.execute(text("SET LOCAL statement_timeout = '15s'"))
            def observe(_conn, _cursor, statement, _parameters, _context, _executemany):
                sql = " ".join(statement.lower().split())
                if "for update" in sql:
                    lock_trace.append(sql)
                    if "from maintenance_project_workbook_state " in sql:
                        waiter_reached_state.set()
            event.listen(connection, "before_cursor_execute", observe)
            try:
                with pytest.raises(operations.MaintenanceOperationConflict, match="版本已变化"):
                    manual.patch_manual_site_issue(
                        session, issue_id=issue_id, project_id=project_id, version=1,
                        idempotency_key=f"facts-wait-{uuid4()}", receiver="过期覆盖",
                        reason="并发契约", operated_by="waiter",
                    )
                session.rollback()
            finally:
                event.remove(connection, "before_cursor_execute", observe)

    with ThreadPoolExecutor(max_workers=2) as pool:
        writer = pool.submit(canonical_writer)
        try:
            assert held.wait(10)
            waiter = pool.submit(manual_waiter)
            assert waiter_reached_state.wait(10)
            assert not waiter.done()
        finally:
            release.set()
        writer.result(timeout=15)
        waiter.result(timeout=15)
    assert "from maintenance_project_workbook_state " in lock_trace[0]
    project_lock = next(i for i, sql in enumerate(lock_trace) if "from maintenance_project " in sql)
    issue_lock = next(i for i, sql in enumerate(lock_trace) if "from maintenance_site_issue " in sql)
    assert 0 < project_lock < issue_lock
    db.expire_all()
    assert db.get(MaintenanceSiteIssueLine, line_id).remark == "先提交的事实"
    assert db.get(MaintenanceSiteIssue, issue_id).receiver == "接收人"
    state = operations.get_or_create_workbook_state(db, project_id=project_id)
    assert state.revision == 1
