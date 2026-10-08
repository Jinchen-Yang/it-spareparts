"""回收清单导入测试（板块 D 第一批，D-1）+ SN 台账唯一性（D-4）。

覆盖：布局不符 422 零写入、行级错误阻断、警告打标放行、聚合去重、
SHA-256 文件幂等、权限失败关闭、SN 唯一约束。
"""
from __future__ import annotations

import io
from decimal import Decimal

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy.exc import IntegrityError

from app import auth
from app.api import circulation
from app.auth import hash_password
from app.models.circulation import CirculationSnItem
from app.models.dimensions import DimPart
from app.models.system import SysUser
from app.services import recycle_import as recycler

COLUMNS = ("备件号", "物料描述", "库房代码", "货位编码", "2WR",
           "中文品名", "备件状态", "公司信息", "单价", "总价")


def _row(pn, desc="2.5in SAS SSD", wh="IT_BJDC", bin_="A-01", qty=2,
         cond="好件", entity="厂商甲", price="100.00", total=None):
    return [pn, desc, wh, bin_, qty, "硬盘", cond, entity,
            Decimal(price), Decimal(total if total is not None else str(Decimal(price) * int(qty)))]


def build_xlsx(rows: list[list], *, header: bool = True, tail_empty_rows: int = 0) -> bytes:
    wb = Workbook()
    ws = wb.active
    if header:
        ws.append(list(COLUMNS))
    for r in rows:
        ws.append(r)
    for _ in range(tail_empty_rows):
        ws.append([None] * len(COLUMNS))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _seed_part(db, pn="PNSYN-0001", part_id=9901):
    db.add(DimPart(id=part_id, pn_std=pn, description="测试件"))
    db.commit()
    return part_id


def _parse(db, rows: list[list], **kw):
    return recycler.parse_recycle_workbook(build_xlsx(rows, **kw), db)


# ── 解析与校验 ──


def test_parse_aggregates_same_pn_warehouse_bin_condition(db):
    report = _parse(db, [
        _row("PNSYN-0001", qty=2, bin_="A-01"),
        _row("PNSYN-0001", qty=3, bin_="A-01"),          # 同键聚合
        _row("PNSYN-0001", qty=1, bin_="A-02"),          # 不同货位 → 独立行
        _row("PNSYN-0002", qty=1, cond="坏件", wh="IT_HZGZDC"),
    ])
    assert report.ok
    assert report.rows_total == 4
    assert len(report.lines) == 3
    by_key = {(line.pn_raw, line.bin_code): line for line in report.lines}
    assert by_key[("PNSYN-0001", "A-01")].qty == Decimal("5")   # 2+3 聚合（A-01）
    assert by_key[("PNSYN-0001", "A-02")].qty == Decimal("1")   # 不同货位独立
    assert by_key[("PNSYN-0001", "A-01")].needs_review is True  # 未收录 PN
    assert report.total_amount == Decimal("700.00")


def test_parse_part_id_matched_when_dim_part_exists(db):
    _seed_part(db, "PNSYN-0001", part_id=9902)
    report = _parse(db, [_row("PNSYN-0001", qty=1)])
    assert report.lines[0].part_id == 9902
    assert report.lines[0].needs_review is False


def test_parse_warnings_and_errors(db):
    report = _parse(db, [
        _row("PNSYN-0001", qty=2, price="100.00", total="150.00"),   # 总价不一致 → 警告
        _row("PNSYN-0002", qty=1, cond="待修"),                       # 非法状态 → 错误
        _row("PNSYN-0003", qty=1, wh="XX_未知"),                      # 未知库区 → 错误
        _row("PNSYN-0004", qty=0),                                    # 数量非正 → 错误
        _row(None),                                                   # 备件号缺失 → 错误
    ])
    assert not report.ok
    assert len(report.warnings) == 1
    fields = {e.field for e in report.errors}
    assert fields == {"备件状态", "库房代码", "2WR", "备件号"}
    assert report.lines and report.lines[0].price_mismatch is True  # 警告行仍聚合


def test_parse_bad_layout_raises(db):
    wb = Workbook()
    ws = wb.active
    ws.append(["备件号", "随便"])
    ws.append(["PNSYN-0001", 1])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(recycler.RecycleImportError):
        recycler.parse_recycle_workbook(buf.getvalue(), db)


def test_parse_trims_long_empty_tail(db):
    rows = [_row("PNSYN-0001", qty=1)]
    report = _parse(db, rows, tail_empty_rows=300)  # 超过 EOF 空行阈值
    assert report.ok and report.rows_total == 1


def test_apply_idempotent_by_sha256(db):
    rows = [_row("PNSYN-0001", qty=2)]
    report = _parse(db, rows)
    batch, duplicate = recycler.apply_recycle_import(
        db, report, source_filename="a.xlsx", imported_by="tester")
    assert duplicate is False
    again_report = _parse(db, rows)
    batch2, duplicate2 = recycler.apply_recycle_import(
        db, again_report, source_filename="副本a.xlsx", imported_by="tester")
    assert duplicate2 is True
    assert batch2.batch_id == batch.batch_id


def test_apply_rejects_error_rows(db):
    report = _parse(db, [_row(None)])
    with pytest.raises(recycler.RecycleImportError):
        recycler.apply_recycle_import(db, report, source_filename="bad.xlsx", imported_by="t")


# ── SN 台账（D-4）──


def test_sn_item_unique_per_sn(db):
    db.add(CirculationSnItem(sn="SN-1", pn_std="PNSYN-0001"))
    db.commit()
    db.add(CirculationSnItem(sn="SN-1", pn_std="PNSYN-0002"))
    with pytest.raises(IntegrityError):
        db.commit()
    db.rollback()


# ── API ──


def _client(db, *, username="recycle_admin", with_action=True, role="admin") -> TestClient:
    permissions = {"page_maintenance": True}
    if with_action:
        permissions["action_recycle_manage"] = True
    db.add(SysUser(
        username=username, role=role, display_name="循环管理员",
        password_hash=hash_password("recycle-password-123"),
        permissions=permissions, is_active=True,
    ))
    db.commit()
    app = FastAPI()
    app.include_router(auth.router, prefix="/api")
    app.include_router(circulation.router, prefix="/api")
    client = TestClient(app)
    login = client.post("/api/auth/login",
                        json={"username": username, "password": "recycle-password-123"})
    assert login.status_code == 200, login.text
    client.headers["Authorization"] = f"Bearer {login.json()['token']}"
    return client


def test_api_apply_then_duplicate(db):
    client = _client(db)
    data = build_xlsx([_row("PNSYN-0001", qty=2)])
    resp = client.post("/api/circulation/recycle-imports",
                       files={"file": ("第十批.xlsx", data)},
                       data={"batch_label": "第十批", "source_doc_no": "305306"})
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["duplicate"] is False
    assert body["line_count"] == 1
    # 同文件重传 → duplicate=true
    resp2 = client.post("/api/circulation/recycle-imports",
                        files={"file": ("副本第十批.xlsx", data)})
    assert resp2.status_code == 200
    assert resp2.json()["duplicate"] is True
    # 批次列表与详情
    listing = client.get("/api/circulation/recycle-imports/batches").json()
    assert len(listing["batches"]) == 1
    detail = client.get(f"/api/circulation/recycle-imports/batches/{body['batch_id']}").json()
    assert detail["lines"][0]["pn_raw"] == "PNSYN-0001"


def test_api_rejects_invalid_rows_422_zero_write(db):
    client = _client(db)
    data = build_xlsx([_row("PNSYN-0001", qty=1), _row("PNSYN-0009", qty=1, cond="未知")])
    resp = client.post("/api/circulation/recycle-imports", files={"file": ("bad.xlsx", data)})
    assert resp.status_code == 422
    assert resp.json()["detail"]["code"] == "rows_invalid"
    listing = client.get("/api/circulation/recycle-imports/batches").json()
    assert listing["batches"] == []  # 零写入


def test_api_rejects_bad_layout_422(db):
    client = _client(db)
    wb = Workbook()
    wb.active.append(["hello", "world"])
    buf = io.BytesIO()
    wb.save(buf)
    resp = client.post("/api/circulation/recycle-imports/preview",
                       files={"file": ("x.xlsx", buf.getvalue())})
    assert resp.status_code == 422
    assert resp.json()["detail"]["code"] == "bad_layout"


def test_api_preview_zero_write(db):
    client = _client(db)
    data = build_xlsx([_row("PNSYN-0001", qty=1)])
    resp = client.post("/api/circulation/recycle-imports/preview",
                       files={"file": ("p.xlsx", data)})
    assert resp.status_code == 200 and resp.json()["ok"] is True
    listing = client.get("/api/circulation/recycle-imports/batches").json()
    assert listing["batches"] == []


def test_api_forbidden_without_action_key(db):
    # 注意：admin 角色动作恒全开（permissions.py 铁律），403 必须用非 admin 角色验证
    client = _client(db, username="recycle_noperm", with_action=False, role="readonly")
    data = build_xlsx([_row("PNSYN-0001", qty=1)])
    resp = client.post("/api/circulation/recycle-imports", files={"file": ("x.xlsx", data)})
    assert resp.status_code == 403
