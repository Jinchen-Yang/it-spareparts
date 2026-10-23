"""Independent consumer contracts for W1a; no new workbook business rules."""
import hashlib
import hmac
import io
import json
from dataclasses import asdict, replace
from datetime import date
from decimal import Decimal
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from app.services import maintenance_expense_collection_workbook as ec
from app.services import maintenance_project_master_workbook as master
from app.services import maintenance_workbook_plans as plans
from app.services import maintenance_workbook_protocol as protocol
from app.services import maintenance_workbook_values as values


@pytest.fixture
def rotating_keys(monkeypatch):
    state = {'active': 'old', 'keys': {'old': b'independent-old', 'new': b'independent-new'}}
    settings = SimpleNamespace(
        maintenance_manifest_signing_material=lambda: (state['active'], state['keys'][state['active']]),
        maintenance_manifest_verification_keys=lambda: state['keys'],
    )
    monkeypatch.setattr(master, 'get_settings', lambda: settings)
    return state


def _metadata_rows():
    return [
        ('protocol_id', protocol.V2_PROTOCOL_ID), ('template_version', protocol.V2_TEMPLATE_VERSION),
        ('project_id', 'independent-project'), ('included_sheets', protocol.V2_SHEET_PARTS),
        ('workbook_revision', '0'), ('workbook_data_version', 'a' * 64),
        ('parts_row_ids', 'line-1,line-2'),
    ]


def _xlsx_roundtrip(rows):
    workbook = Workbook()
    workbook.active.title = protocol.V2_SHEET_PARTS
    sheet = workbook.create_sheet(protocol.V2_SHEET_META)
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return load_workbook(io.BytesIO(output.getvalue()))


def _external_signature(metadata, key):
    # Independent wire consumer, with literal domain rather than production constants.
    payload = {name: str(value) for name, value in metadata.items() if name != 'metadata_hmac'}
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode()
    return hmac.new(key, b'ITDATA_MAINT_PROJECT_MASTER_META_V1\x00' + encoded, hashlib.sha256).hexdigest()


def test_xlsx_external_signature_reassembles_distinct_twelve_chunks(rotating_keys):
    # A uniform long value cannot expose lexical @10-before-@2 joining.
    long_value = ''.join(chr(0x4e00 + index) * 30000 for index in range(12))
    rows = _metadata_rows() + [('parts_base_hashes', long_value), ('future_scope', '范围/甲')]
    metadata = dict(rows) | {'metadata_hmac_algorithm': 'HMAC-SHA256', 'metadata_hmac_key_id': 'old'}
    metadata['metadata_hmac'] = _external_signature(metadata, rotating_keys['keys']['old'])
    physical = []
    for name, value in metadata.items():
        physical.extend(values._chunk_meta_value(name, value))
    workbook = _xlsx_roundtrip(sorted(physical, key=lambda row: row[0]))
    parsed = master._v2_verify_meta(None, workbook, 'independent-project')
    assert parsed == metadata
    assert parsed['parts_base_hashes'] == long_value


def test_adapter_signer_uses_current_key_and_old_files_survive_rotation(rotating_keys):
    old = _xlsx_roundtrip(master._v2_sign_meta_rows(_metadata_rows()))
    rotating_keys['active'] = 'new'
    new = _xlsx_roundtrip(master._v2_sign_meta_rows(_metadata_rows()))
    for workbook, key_id in ((old, 'old'), (new, 'new')):
        parsed = master._v2_verify_meta(None, workbook, 'independent-project')
        assert parsed['metadata_hmac_key_id'] == key_id
        assert parsed['metadata_hmac'] == _external_signature(parsed, rotating_keys['keys'][key_id])
    del rotating_keys['keys']['old']
    with pytest.raises(ec.WorkbookError) as caught:
        master._v2_verify_meta(None, old, 'independent-project')
    assert type(caught.value) is plans.WorkbookError
    assert caught.value.code == 'template_signature_invalid'
    assert master._v2_verify_meta(None, new, 'independent-project')['metadata_hmac_key_id'] == 'new'


@pytest.mark.parametrize('field,replacement', [
    ('parts_row_ids', 'line-1'), ('included_sheets', protocol.V2_SHEET_SITE),
    ('workbook_revision', '1'), ('workbook_data_version', 'b' * 64),
    ('metadata_hmac_key_id', 'new'), ('future_scope', 'changed'),
])
def test_real_meta_verifier_rejects_post_signing_scope_edits(rotating_keys, field, replacement):
    rows = master._v2_sign_meta_rows(_metadata_rows() + [('future_scope', 'original')])
    workbook = _xlsx_roundtrip([(key, replacement if key == field else value) for key, value in rows])
    with pytest.raises(ec.WorkbookError) as caught:
        master._v2_verify_meta(None, workbook, 'independent-project')
    assert caught.value.code == 'template_signature_invalid'


def test_row_token_and_meta_signatures_are_not_interchangeable(rotating_keys):
    normalized = {'日期': '2026-01-02', '数量': '2', '标记': 'False', '空值': ''}
    token = master._v2_row_base_token(
        sheet=protocol.V2_SHEET_SITE, entity_id=71,
        values={'日期': date(2026, 1, 2), '数量': Decimal('2.00'), '标记': False, '空值': None},
    )
    envelope = json.loads(token)
    payload = json.dumps(envelope['payload'], ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode()
    expected = hmac.new(rotating_keys['keys']['old'], b'ITDATA_MAINT_MASTER_ROW_BASE_V1\x00' + payload, hashlib.sha256).hexdigest()
    assert envelope['signature'] == expected
    assert master._parse_v2_row_base_token(token, sheet=protocol.V2_SHEET_SITE, entity_id='71', row_no=9) == normalized
    envelope['signature'] = hmac.new(rotating_keys['keys']['old'], b'ITDATA_MAINT_PROJECT_MASTER_META_V1\x00' + payload, hashlib.sha256).hexdigest()
    with pytest.raises(ec.WorkbookError) as caught:
        master._parse_v2_row_base_token(json.dumps(envelope), sheet=protocol.V2_SHEET_SITE, entity_id=71, row_no=9)
    assert type(caught.value) is plans.WorkbookError
    assert caught.value.code == 'invalid_concurrency_token'
    assert isinstance(caught.value.__cause__, ValueError)


def test_neutral_row_hash_keeps_selected_field_order_and_address_guard_separate():
    row = dict(zip(protocol.V2_SITE_BASE_FIELDS, ('LY-1', date(2026, 1, 2), 'PN-1', None, Decimal('2.00'), False, '备注')))
    row['领用地址'] = '甲地'
    baseline = values._v2_row_base_hash(row, protocol.V2_SITE_BASE_FIELDS)
    guard = master.V2RowGuard(protocol.V2_SHEET_SITE, 'line-1', 'row-2', baseline, protocol.V2_SITE_BASE_FIELDS)
    changed = {**row, '领用地址': '乙地', 'unrelated': 'ignored'}
    assert values._v2_row_base_hash(changed, guard.fields) == baseline
    assert values._v2_row_base_hash(changed, protocol.V2_SITE_ADDRESS_GUARD_FIELDS) != values._v2_row_base_hash(row, protocol.V2_SITE_ADDRESS_GUARD_FIELDS)
    assert values._v2_row_base_hash(row, tuple(reversed(guard.fields))) != baseline
    assert values._v2_row_base_hash({**row, '领用数量': 2.0}, guard.fields) == baseline
    assert values._v2_row_base_hash({**row, '是否应返还': 0}, guard.fields) != baseline


def test_mixed_v2_plan_summary_preserves_each_consumers_counting_meaning():
    cost = plans.CostRefill(1, None, None, None)
    receipt = ec.CollectionOp('CREATE', 'contract-id', 'XSDD1', date(2026, 1, 1), Decimal('0'), None, None)
    milestone = plans.V2MilestoneChange('CREATE', 'XSDD1', 1, None, 'month', Decimal('0'), None, None)
    plan = master.MasterV2Plan(
        'project', (protocol.V2_SHEET_PARTS,),
        cost_refills=(replace(cost, qty=Decimal('0')), replace(cost, reason=''),
                      replace(cost, unit_cost_inc_tax=Decimal('1.13')),
                      replace(cost, is_create=True, qty=Decimal('1')),
                      replace(cost, operation='VOID', qty=Decimal('2'))),
        site_flags=(plans.SiteReturnFlag('u', False), plans.SiteReturnFlag('c', None, is_create=True),
                    plans.SiteReturnFlag('v', False, is_void=True)),
        expense_updates=(ec.ExpenseUpdate('u'), plans.ExpenseUpdate('c', is_create=True)),
        receipt_ops=tuple(replace(receipt, operation=operation) for operation in ('CREATE', 'UPDATE', 'VOID')),
        milestone_changes=tuple(replace(milestone, operation=operation) for operation in ('CREATE', 'UPDATE', 'VOID')),
        assignment_changes=(plans.SourceOrderAssignmentChange('order', 'WBDD1', None, None, None, None),),
        contract_amount_change=plans.ContractAmountChange('contract-id', 1, None, Decimal('0')),
        project_note_change=plans.ProjectNoteChange(1, 'old', None),
        expense_voids=('void-1',), site_issue_merges=({'target': 'site-1'},),
    )
    assert plan.summary == {
        'line_updates': 3, 'qty_updates': 1, 'cost_overrides': 1, 'line_creates': 1,
        'line_voids': 1, 'order_reassignments': 1, 'contract_updates': 1,
        'project_note_updates': 1, 'expense_creates': 1, 'expense_updates': 1,
        'expense_voids': 1, 'plan_creates': 1, 'plan_updates': 1, 'plan_voids': 1,
        'collection_updates': 2, 'collection_voids': 1, 'site_creates': 1,
        'site_voids': 1, 'site_updates': 1, 'site_issue_merges': 1,
    }
    inner = ec.WorkbookPlan('project', plan.expense_updates, plan.receipt_ops)
    assert inner.summary['collection_updates'] == 1
    assert plans.MasterPlan('project', inner=inner).summary['collection_creates'] == 1
    snapshot = asdict(plan)
    assert snapshot['expense_updates'][1]['raw_line_id'] == 'c'
    assert snapshot['contract_amount_change']['amount_inc_tax'] == Decimal('0')
    assert isinstance(replace(plan, warnings=('observed',)), master.MasterV2Plan)


def test_empty_master_summary_does_not_acquire_inner_only_keys():
    empty = plans.MasterPlan('project')
    assert empty.summary == {'cost_refills': 0, 'site_return_flags': 0,
                             'expense_updates': 0, 'collection_creates': 0, 'collection_voids': 0}
    with_inner = master.MasterPlan('project', inner=ec.WorkbookPlan('project', (), ()))
    assert with_inner.summary['expense_creates'] == 0
    assert with_inner.summary['collection_updates'] == 0
    changed = with_inner.summary
    changed['expense_updates'] = 99
    assert with_inner.summary['expense_updates'] == 0


def test_error_empty_containers_remain_detached():
    empty_issues, empty_details = [], {}
    error = plans.WorkbookError('code', 'message', empty_issues, details=empty_details)
    empty_issues.append({'row': 1})
    empty_details['actor'] = 'new'
    assert error.issues == [] and error.details == {}
    assert error.args == ('message',)
    with pytest.raises(TypeError):
        ec.WorkbookError('code', 'message', [], {'actor': 'positional-details'})
