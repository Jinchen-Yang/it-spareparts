"""维保数据看板 API：权限矩阵、项目范围、参数校验与导出（R33 / D-0022 +
R34 / D-0024 试算比例、M01/M03 新端点）。

关键口径：
- 数量端点对 maintenance_manager / purchaser 开放，但响应与导出不得出现
  任何金额字段（键都不出现，不是 0/null）——递归扫描断言。
- 金额端点（valuation）= 登录 + page_maintenance + 项目范围 + 利润门；
- readonly 兜底模板有 data_profit 但无 page_maintenance → 仍 403（利润门单独不够）。
- 受限账号手填未授权项目 ID → 空交集 → 空结果，绝不回退全量。
- R34：value_retention_ratio 试算参数 [0,1]，**缺省 = 50%**（Sunny 01:20 修正），
  20%/30% 仅可选；显式 0 合法（不得吞成默认）；非法/越界/非有限 422。
"""

import csv
import io
import uuid
from datetime import date, datetime, timezone
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from app import permissions
from app.auth import hash_password
from app.main import app
from app.models.dimensions import DimPart
from app.models.maintenance_project import (
    MaintenanceProject,
    MaintenanceProjectUserAssignment,
)
from app.models.purchase import FPurchaseLine, FPurchaseOrder
from app.models.sales import FSalesLine, FSalesOrder
from app.models.system import SysImportBatch, SysUser
from app.services import maintenance_dashboard as dash
from app.services.maintenance_dashboard_valuation_policy import POLICY_VERSION
from tests.test_maintenance_return_metrics import _issue, _receipt


def _part(db, pn, *, unit=None, description=""):
    part = DimPart(pn_std=pn, unit=unit, description=description)
    db.add(part)
    db.flush()
    return part


def _project(db, code):
    project = MaintenanceProject(
        project_id=str(uuid.uuid4()), project_code=code, display_name=code,
        lifecycle_status="ongoing",
    )
    db.add(project)
    db.flush()
    return project

BASE = "/api/maintenance/analytics/maintenance-dashboard"
PASSWORD = "mdd-api-password-1"
TODAY = date(2026, 10, 21)
UTC = timezone.utc
# 无金额权响应里递归不允许出现的键片段（金额域键集）。
FORBIDDEN_KEY_PARTS = ("valuation", "value_inc", "value_ex", "amount", "price",
                       "factor", "reference", "coverage", "policy")


@pytest.fixture(autouse=True)
def _freeze_today(monkeypatch):
    monkeypatch.setattr(dash, "business_today", lambda: TODAY)


def client_for(db, *, username, role="readonly", overrides=None) -> TestClient:
    base = permissions.effective(role, None)
    effective = permissions.effective_from_snapshot(base, overrides or {})
    db.add(SysUser(
        username=username, role=role, display_name=username,
        password_hash=hash_password(PASSWORD), is_active=True,
        template_code=role, template_version=1, template_perms=base,
        perm_overrides=overrides or {}, permissions=effective,
    ))
    db.commit()
    client = TestClient(app)
    login = client.post("/api/auth/login",
                        json={"username": username, "password": PASSWORD})
    assert login.status_code == 200, login.text
    client.headers["Authorization"] = f"Bearer {login.json()['token']}"
    return client


def _assign_manager(db, client_username: str, project: MaintenanceProject):
    user = db.scalar(select(SysUser).where(SysUser.username == client_username))
    db.add(MaintenanceProjectUserAssignment(
        assignment_id=str(uuid.uuid4()), project_id=project.project_id,
        responsibility_type="primary_manager", user_id=user.id,
        source_manager_text=client_username, version=1, assigned_by="t",
        assignment_reason="test",
    ))
    db.commit()


def _purchase_ref(db, part, *, order_date=date(2026, 10, 18)):
    batch = SysImportBatch(
        filename="mdd-api-cg.xlsx", file_type="purchase",
        file_hash=uuid.uuid4().hex, status="success",
    )
    db.add(batch)
    db.flush()
    order = FPurchaseOrder(
        raw_order_id=f"mdd-api-cg-{uuid.uuid4().hex[:8]}", order_no="MDDA-CG-1",
        purchaser="t", data_status="已生效", order_date=order_date,
        tax_rate=Decimal("0.13"), is_tax_inclusive=True, import_batch_id=batch.id,
    )
    db.add(order)
    db.flush()
    db.add(FPurchaseLine(
        raw_line_id=f"mdd-api-cg-l-{uuid.uuid4().hex[:8]}", order_id=order.id,
        line_no=1, part_id=part.id, pn_std=part.pn_std, qty=Decimal(2),
        unit_price=Decimal("113.00"), unit="件", import_batch_id=batch.id,
    ))
    db.flush()
    return order


@pytest.fixture()
def seeded(db):
    """两个项目 + 坏件实收（期内）+ 领用 + 采购参考（113 含税 / 100 未税）。"""
    p1 = _project(db, "MDDA-P1")
    p2 = _project(db, "MDDA-P2")
    part = _part(db, "MDDA-PN-1", unit="件")
    _purchase_ref(db, part)
    _receipt(db, p1, part, "4", condition="坏品",
             occurred_at=datetime(2026, 10, 18, 8, tzinfo=UTC))
    _receipt(db, p2, part, "2", condition="废品",
             occurred_at=datetime(2026, 10, 18, 9, tzinfo=UTC))
    _issue(db, p1, part, "5", issue_date=date(2026, 10, 18))
    db.commit()
    return p1, p2, part


def _forbidden_keys(node, path="payload") -> set[str]:
    found: set[str] = set()
    if isinstance(node, dict):
        for key, value in node.items():
            if any(part in str(key) for part in FORBIDDEN_KEY_PARTS):
                found.add(f"{path}.{key}")
            found |= _forbidden_keys(value, f"{path}.{key}")
    elif isinstance(node, list):
        for index, item in enumerate(node):
            found |= _forbidden_keys(item, f"{path}[{index}]")
    return found


def test_unauthenticated_401(db):
    client = TestClient(app)
    for endpoint in ("/summary", "/projects", "/parts", "/valuation",
                     "/receipts", "/quantity-checks", "/export", "/options",
                     "/orders", "/consumption", "/consumption/issues"):
        response = client.get(f"{BASE}{endpoint}")
        assert response.status_code == 401, (endpoint, response.status_code)


def test_page_permission_required_even_with_profit(db, seeded):
    """readonly 有 data_profit 但无 page_maintenance → 数量/金额端点都 403。"""
    client = client_for(db, username="mdd-readonly", role="readonly")
    for endpoint in ("/summary", "/valuation", "/orders", "/consumption",
                     "/consumption/issues"):
        assert client.get(f"{BASE}{endpoint}").status_code == 403, endpoint


def test_admin_without_ratio_defaults_to_half(db, seeded):
    """Sunny 01:20 修正：缺省 value_retention_ratio → 按 50% 试算并回显默认。

    有利润权者首次查询即得 0.5 试算金额；政策载荷声明默认与可选项。
    """
    client = client_for(db, username="mdd-admin-aw", role="admin")
    summary = client.get(f"{BASE}/summary")
    assert summary.status_code == 200
    valuation = summary.json()["cards"]["valuation"]
    # 与显式 0.5 同额：坏品 4×113×0.5 + 废品 2×113×0.5 → 339.00/300.00。
    assert valuation["state"] == "ready"
    assert valuation["retention_ratio"] == "0.5"
    assert valuation["coverage_pct"] == 100.0
    assert valuation["bad_estimated_value"]["inc_tax"] == "339.00"
    assert valuation["awaiting_ratio_bad_qty"] == "0.000"  # HTTP 层不再有待试算态

    valuation_page = client.get(f"{BASE}/valuation")
    assert valuation_page.status_code == 200
    page = valuation_page.json()
    assert page["policy"]["retention_ratio"] == "0.5"
    assert page["policy"]["retention_ratio_default"] == "0.5"
    assert page["policy"]["retention_ratio_required"] is False
    assert page["policy"]["quick_options"] == ["0.5", "0.3", "0.2"]
    assert page["cards"]["bad_estimated_value"]["inc_tax"] == "339.00"
    assert page["cards"]["coverage_pct"] == 100.0
    parts = client.get(f"{BASE}/parts").json()["rows"]
    row = next(item for item in parts if item["pn"] == "MDDA-PN-1")
    assert row["valuation"]["state"] == "ready"
    assert row["valuation"]["bad_estimated_value_inc_tax"] == "339.00"
    assert row["valuation"]["coverage_pct"] == 100.0
    # 参考价仍并列展示。
    assert row["references"]["purchase"]["inc_tax"] == "113.00"

    # options 政策说明也回显默认（UI 首屏即可展示 50%）。
    options = client.get(f"{BASE}/options").json()
    assert options["valuation_policy"]["retention_ratio"] == "0.5"
    assert options["valuation_policy"]["retention_ratio_default"] == "0.5"


@pytest.mark.parametrize("filters", [
    {"q": "MDDA-NO-MATCH"},
    {"project": "00000000-0000-0000-0000-000000000000"},
    {"range": "custom", "date_from": "2020-01-01", "date_to": "2020-01-02"},
])
def test_valuation_empty_filter_returns_zero_quantities(db, seeded, filters):
    client = client_for(db, username="mdd-admin-empty", role="admin")
    response = client.get(f"{BASE}/valuation", params=filters)

    assert response.status_code == 200, response.text
    assert response.headers["cache-control"] == "no-store"
    page = response.json()
    assert page["rows"] == []
    assert page["pagination"] == {"page": 1, "page_size": 20, "total": 0}
    assert page["policy"]["retention_ratio"] == "0.5"
    assert page["policy"]["retention_ratio_default"] == "0.5"
    assert page["cards"] == {
        "estimated_bad_qty": "0.000",
        "pending_bad_qty": "0.000",
        "awaiting_ratio_bad_qty": "0.000",
        "coverage_pct": None,
        "unit": None,
        "unit_breakdown": [],
        "domains": {},
    }


def test_admin_sees_amounts_with_ratio_and_policy_version(db, seeded):
    p1, _p2, part = seeded
    client = client_for(db, username="mdd-admin", role="admin")
    summary = client.get(f"{BASE}/summary",
                         params={"value_retention_ratio": "0.5"})
    assert summary.status_code == 200
    assert summary.headers["cache-control"] == "no-store"
    body = summary.json()
    assert body["range"] == {"start": "2026-10-15", "end": "2026-10-21"}
    valuation = body["cards"]["valuation"]
    # R34 统一系数：坏品 4×113×0.5=226.00 + 废品 2×113×0.5=113.00 → 339.00（含税）；
    # 未税：200.00 + 100.00 → 300.00。
    assert valuation["state"] == "ready"
    assert valuation["bad_estimated_value"]["inc_tax"] == "339.00"
    assert valuation["bad_estimated_value"]["ex_tax"] == "300.00"
    assert valuation["retention_ratio"] == "0.5"
    assert valuation["policy_version"] == POLICY_VERSION
    assert valuation["coverage_pct"] == 100.0
    # 三域并列（B-F07）：采购域同额；无销售/成本参考的域仍报覆盖
    # （坏件分母照给、覆盖率 0、金额键缺失——未知不填 0）。
    assert valuation["domains"]["purchase"]["value_inc_tax"] == "339.00"
    for domain_name in ("sales", "cost"):
        domain = valuation["domains"][domain_name]
        assert "value_inc_tax" not in domain
        assert "value_ex_tax" not in domain
        assert domain["bad_qty"] == "6.000"
        assert domain["estimated_bad_qty"] == "0.000"
        assert domain["coverage_pct"] == 0.0
        assert domain["coverage_basis"] == "bad_receipts"

    valuation_page = client.get(
        f"{BASE}/valuation", params={"value_retention_ratio": "0.5"}
    )
    assert valuation_page.status_code == 200
    page = valuation_page.json()
    assert page["policy"]["policy_version"] == POLICY_VERSION
    assert page["policy"]["retention_ratio"] == "0.5"
    assert "condition_factors" not in page["policy"]  # R33 旧系数键已废除
    assert page["cards"]["bad_estimated_value"]["inc_tax"] == "339.00"

    parts = client.get(f"{BASE}/parts",
                       params={"value_retention_ratio": "0.5"}).json()["rows"]
    row = next(item for item in parts if item["pn"] == "MDDA-PN-1")
    assert row["valuation"]["bad_estimated_value_inc_tax"] == "339.00"
    assert row["valuation"]["domains"]["purchase"]["value_inc_tax"] == "339.00"
    assert row["references"]["chosen"] == "purchase"
    assert row["references"]["purchase"]["inc_tax"] == "113.00"
    # /parts 金额可见时附政策（比例回显随请求参数）。
    assert "policy" in client.get(
        f"{BASE}/parts", params={"value_retention_ratio": "0.3"}
    ).json()

    receipts = client.get(
        f"{BASE}/receipts", params={"value_retention_ratio": "0.5"}
    ).json()["rows"]
    bad = next(item for item in receipts if item["test_result"] == "坏品")
    assert bad["valuation"]["state"] == "ready"
    assert bad["valuation"]["factor"] == "0.5"
    assert bad["valuation"]["value_inc_tax"] == "226.00"
    assert bad["valuation"]["domains"]["cost"]["available"] is False


def test_boss_role_sees_amounts_by_role_template(db, seeded):
    """R33 权限矩阵：老板（boss）与管理员同看金额，不依赖 page_maintenance_boss。

    账号不带 overrides（按 ROLE_TEMPLATES 兜底）：boss 模板 page_maintenance=True、
    data_profit=True，而维保展示板的 page_maintenance_boss 本就显式 False——
    数据看板的门是 page_maintenance + 利润门，与展示板入口键无关。
    """
    base = permissions.effective("boss", None)
    assert base["page_maintenance"] is True
    assert base["data_profit"] is True
    assert base["page_maintenance_boss"] is False

    client = client_for(db, username="mdd-boss", role="boss")

    summary = client.get(f"{BASE}/summary",
                         params={"value_retention_ratio": "0.5"})
    assert summary.status_code == 200
    cards = summary.json()["cards"]
    # FULL_SCOPE_ROLES 含 boss：两个项目的实收（4+2）都在。
    assert cards["receipt_qty"]["total"] == "6.000"
    valuation = cards["valuation"]
    assert valuation["bad_estimated_value"]["inc_tax"] == "339.00"
    assert valuation["bad_estimated_value"]["ex_tax"] == "300.00"
    assert valuation["policy_version"] == POLICY_VERSION

    valuation_page = client.get(f"{BASE}/valuation")
    assert valuation_page.status_code == 200
    assert valuation_page.json()["policy"]["policy_version"] == POLICY_VERSION

    parts = client.get(f"{BASE}/parts",
                       params={"value_retention_ratio": "0.5"})
    assert parts.status_code == 200
    row = next(item for item in parts.json()["rows"] if item["pn"] == "MDDA-PN-1")
    assert row["valuation"]["bad_estimated_value_inc_tax"] == "339.00"
    assert row["references"]["chosen"] == "purchase"
    assert row["references"]["purchase"]["inc_tax"] == "113.00"

    export = client.get(
        f"{BASE}/export",
        params={"kind": "parts", "format": "csv",
                "value_retention_ratio": "0.5"},
    )
    assert export.status_code == 200
    header = next(csv.reader(io.StringIO(export.content.decode("utf-8"))))
    for amount_col in ("坏件试算价值-含税", "估价政策版本", "价值保留比例"):
        assert amount_col in header


def test_maintenance_manager_quantity_only_no_amount_leak(db, seeded):
    p1, _p2, _part = seeded
    username = "mdd-manager"
    client = client_for(db, username=username, role="maintenance_manager")
    _assign_manager(db, username, p1)

    options = client.get(f"{BASE}/options")
    assert options.status_code == 200
    options_body = options.json()
    assert options_body["amounts_visible"] is False
    assert "valuation_policy" not in options_body
    assert "bad_value_desc" not in options_body["sorts"]["projects"]
    assert [p["project_code"] for p in options_body["projects"]] == ["MDDA-P1"]

    summary = client.get(f"{BASE}/summary")
    assert summary.status_code == 200
    assert _forbidden_keys(summary.json()) == set()
    assert summary.json()["cards"]["receipt_qty"]["total"] == "4.000"

    projects = client.get(f"{BASE}/projects")
    assert projects.status_code == 200
    assert _forbidden_keys(projects.json()) == set()
    assert [r["project_code"] for r in projects.json()["rows"]] == ["MDDA-P1"]

    parts = client.get(f"{BASE}/parts")
    assert parts.status_code == 200
    assert _forbidden_keys(parts.json()) == set()

    receipts = client.get(f"{BASE}/receipts")
    assert receipts.status_code == 200
    assert _forbidden_keys(receipts.json()) == set()

    # 新端点（M01/M03）数量域干净；即使显式带比例参数也不泄漏金额键。
    orders = client.get(f"{BASE}/orders")
    assert orders.status_code == 200
    assert _forbidden_keys(orders.json()) == set()
    assert orders.json()["cards"]["order_qty"] == "0.000"  # 无期内需求单
    consumption = client.get(
        f"{BASE}/consumption", params={"value_retention_ratio": "0.5"},
    )
    assert consumption.status_code == 200
    assert _forbidden_keys(consumption.json()) == set()
    assert consumption.json()["cards"]["issued_qty"] == "5.000"
    issues = client.get(f"{BASE}/consumption/issues")
    assert issues.status_code == 200
    assert _forbidden_keys(issues.json()) == set()
    assert len(issues.json()["rows"]) == 1
    assert issues.json()["rows"][0]["qty"] == "5.000"

    # 金额端点 403；金额排序 422（不静默降级）。
    assert client.get(f"{BASE}/valuation").status_code == 403
    denied = client.get(f"{BASE}/projects",
                        params={"sort": "bad_value_desc"})
    assert denied.status_code == 422
    assert denied.json()["detail"]["code"] == "sort_requires_profit_permission"

    # 导出（数量列完整、无任何金额列/估价元数据）。
    export = client.get(f"{BASE}/export",
                        params={"kind": "parts", "format": "csv"})
    assert export.status_code == 200
    rows = list(csv.reader(io.StringIO(export.content.decode("utf-8"))))
    header = rows[0]
    assert "坏件实收" in header
    for forbidden in ("坏件试算价值-含税", "估价政策版本", "参考价-含税",
                      "覆盖率%", "价值保留比例", "成本来源", "已计成本-含税"):
        assert forbidden not in header


def test_purchaser_full_scope_quantity_only(db, seeded):
    client = client_for(db, username="mdd-purchaser", role="purchaser")
    projects = client.get(f"{BASE}/projects")
    assert projects.status_code == 200
    body = projects.json()
    assert _forbidden_keys(body) == set()
    assert {r["project_code"] for r in body["rows"]} == {"MDDA-P1", "MDDA-P2"}
    assert client.get(f"{BASE}/valuation").status_code == 403
    checks = client.get(f"{BASE}/quantity-checks")
    assert checks.status_code == 200
    assert _forbidden_keys(checks.json()) == set()


def test_scope_hand_filled_foreign_project_sees_nothing(db, seeded):
    p1, p2, _part = seeded
    username = "mdd-manager-foreign"
    client = client_for(db, username=username, role="maintenance_manager")
    _assign_manager(db, username, p1)
    # 手填未授权项目 ID：交集为空 → 空结果，绝不回退全量。
    response = client.get(f"{BASE}/projects",
                          params={"project": p2.project_id})
    assert response.status_code == 200
    assert response.json()["rows"] == []
    leak = client.get(f"{BASE}/summary", params={"project": p2.project_id})
    assert leak.json()["cards"]["receipt_qty"]["total"] == "0.000"
    # 选自己可见的项目正常返回。
    own = client.get(f"{BASE}/projects", params={"project": p1.project_id})
    assert [r["project_code"] for r in own.json()["rows"]] == ["MDDA-P1"]


def test_invalid_params_422_not_500(db, seeded):
    client = client_for(db, username="mdd-admin-422", role="admin")
    for params in (
        {"range": "ytd"},
        {"range": "custom"},
        {"range": "custom", "date_from": "2026-10-20", "date_to": "2026-10-01"},
        {"page_size": 0},
        {"page": -1},
        {"sort": "no_such_sort"},
    ):
        response = client.get(f"{BASE}/projects", params=params)
        assert response.status_code == 422, (params, response.status_code)
    # R34 比例参数：越界/非数值/NaN → 422 invalid_retention_ratio（所有角色一致）。
    for bad in ("1.5", "-0.1", "abc", "NaN"):
        for endpoint in ("/summary", "/parts", "/valuation", "/receipts",
                         "/export"):
            response = client.get(
                f"{BASE}{endpoint}",
                params={"kind": "parts", "value_retention_ratio": bad},
            )
            assert response.status_code == 422, (endpoint, bad)
            assert response.json()["detail"]["code"] == "invalid_retention_ratio"
    assert client.get(f"{BASE}/receipts",
                      params={"sort": "bad_value_desc"}).status_code == 422
    assert client.get(f"{BASE}/export",
                      params={"kind": "nope"}).status_code == 422
    assert client.get(f"{BASE}/export",
                      params={"kind": "parts", "format": "pdf"}).status_code == 422


def test_retention_ratio_zero_is_legal_blank_defaults_half(db, seeded):
    """显式 0 合法（全 0 试算，绝不被吞成默认）；空白等价缺省 → 0.5。"""
    client = client_for(db, username="mdd-admin-r0", role="admin")
    zero = client.get(f"{BASE}/valuation",
                      params={"value_retention_ratio": "0"})
    assert zero.status_code == 200
    body = zero.json()
    assert body["policy"]["retention_ratio"] == "0"
    assert body["cards"]["bad_estimated_value"]["inc_tax"] == "0.00"
    assert body["cards"]["estimated_bad_qty"] == "6.000"
    assert body["cards"]["coverage_pct"] == 100.0
    blank = client.get(f"{BASE}/valuation",
                       params={"value_retention_ratio": " "})
    assert blank.status_code == 200
    assert blank.json()["policy"]["retention_ratio"] == "0.5"
    assert blank.json()["cards"]["bad_estimated_value"]["inc_tax"] == "339.00"
    # 数量角色传非法比例同样 422（一致校验，不给静默忽略）；
    # 数量角色不提供比例也照常 200（默认不得要求其提供或泄露金额键）。
    manager = client_for(db, username="mdd-manager-r0", role="maintenance_manager")
    _assign_manager(db, "mdd-manager-r0", seeded[0])
    denied = manager.get(f"{BASE}/summary",
                         params={"value_retention_ratio": "9"})
    assert denied.status_code == 422
    assert denied.json()["detail"]["code"] == "invalid_retention_ratio"
    plain = manager.get(f"{BASE}/summary")
    assert plain.status_code == 200
    assert _forbidden_keys(plain.json()) == set()


def test_projects_accepts_retention_ratio_like_other_views(db, seeded):
    """/projects 行级试算与 summary/parts/导出同参同口径（root 裁决补传）。

    缺省 0.5、显式 0.2、显式 0 与 /valuation 卡同额；数量账号不因补参泄露。
    """
    p1, _p2, _part = seeded
    client = client_for(db, username="mdd-admin-pj", role="admin")

    def row_for(params=None):
        rows = client.get(f"{BASE}/projects", params=params).json()["rows"]
        return next(r for r in rows if r["project_id"] == p1.project_id)["valuation"]

    default = row_for()
    assert default["state"] == "ready"
    assert default["awaiting_ratio_bad_qty"] == "0.000"
    # P1 坏品 4：4×113×0.5 = 226.00 含税 / 4×100×0.5 = 200.00 未税。
    assert default["bad_estimated_value_inc_tax"] == "226.00"
    assert default["bad_estimated_value_ex_tax"] == "200.00"
    assert default["coverage_pct"] == 100.0
    # 显式 0.2 → 90.40；显式 0 → 0.00 且仍 ready（不吞成默认）。
    explicit = row_for({"value_retention_ratio": "0.2"})
    assert explicit["bad_estimated_value_inc_tax"] == "90.40"
    zero = row_for({"value_retention_ratio": "0"})
    assert zero["state"] == "ready"
    assert zero["bad_estimated_value_inc_tax"] == "0.00"
    # 与 /valuation 卡同额（同屏口径一致）。
    card = client.get(f"{BASE}/valuation",
                      params={"value_retention_ratio": "0.2"}
                      ).json()["cards"]["bad_estimated_value"]
    assert card["inc_tax"] == "135.60"  # 6×113×0.2（P1+P2 全量卡）
    assert explicit["bad_estimated_value_inc_tax"] == "90.40"  # 4×113×0.2（P1 行）
    # 数量账号：无参数/显式 0.2 都不出 valuation 键、不泄露金额。
    manager = client_for(db, username="mdd-manager-pj", role="maintenance_manager")
    _assign_manager(db, "mdd-manager-pj", p1)
    mrows = manager.get(f"{BASE}/projects").json()["rows"]
    assert "valuation" not in next(
        r for r in mrows if r["project_id"] == p1.project_id)
    with_ratio = manager.get(
        f"{BASE}/projects", params={"value_retention_ratio": "0.2"}
    ).json()
    assert _forbidden_keys(with_ratio) == set()


def test_range_all_and_literal_keyword_five_entrypoints(db, seeded):
    """B-F01/B-F02（独立首轮真实 HTTP 反例）：

    - range=all 曾在五个金额入口 500（occurred_at >= None → ArgumentError）；
    - 非空 q 曾 500（ilike(autoescape=True) 本 SQLAlchemy 不支持）。
    修复后：all 期间完整可查、不吞数据；关键词按**字面**匹配——
    %/_/反斜杠不得变成通配符（含未解析行原始 PN 兜底路径与导出）。
    """
    p1, _p2, _base = seeded
    # 字面通配字符种子 + 各自“若未转义则会被误中”的诱饵行。
    lit_pct = _part(db, "MDDA-LIT-%-PN", unit="件")
    lit_us = _part(db, "MDDA-US_A", unit="件")
    decoy_pct = _part(db, "MDDA-LIT-ABC-PN", unit="件")  # % 当通配会误中
    decoy_us = _part(db, "MDDA-USXA", unit="件")         # _ 当通配会误中
    at = datetime(2026, 10, 18, 10, tzinfo=UTC)
    _receipt(db, p1, lit_pct, "2", occurred_at=at)
    _receipt(db, p1, lit_us, "3", occurred_at=at)
    _receipt(db, p1, decoy_pct, "5", occurred_at=at)
    _receipt(db, p1, decoy_us, "7", occurred_at=at)
    # 未解析行（part_id=None）走原始 PN 文本兜底 ilike；反斜杠按字面。
    _receipt(db, p1, None, "11", pn="MDDA-RAW\\-LIT", occurred_at=at)
    _receipt(db, p1, None, "13", pn="MDDA-RAW-LIT", occurred_at=at)
    db.commit()
    client = client_for(db, username="mdd-admin-allq", role="admin")

    # B-F01：五入口 range=all 全部 200（不 500），且不吞数据：
    # 47 = 默认夹具 6（坏4+废2）+ 2+3+5+7+11+13。
    for endpoint in ("summary", "projects", "parts", "valuation", "receipts"):
        response = client.get(f"{BASE}/{endpoint}", params={"range": "all"})
        assert response.status_code == 200, (endpoint, response.text[:300])
    summary = client.get(f"{BASE}/summary", params={"range": "all"}).json()
    # B-F09：range=all 下「件」与未解析行（part_id=None → 待核桶）混单位，
    # 平面 total = null，unit_breakdown 呈现真实分桶（23 + 24，不猜单位）。
    receipt_card = summary["cards"]["receipt_qty"]
    assert receipt_card["total"] is None
    assert {
        entry["unit"]: entry["total"]
        for entry in receipt_card["unit_breakdown"]
    } == {"件": "23.000", None: "24.000"}

    # B-F02：非空 q 五入口 200；%/_ 按字面匹配，诱饵 PN 不得混入。
    for endpoint in ("summary", "projects", "parts", "valuation", "receipts"):
        response = client.get(
            f"{BASE}/{endpoint}", params={"range": "all", "q": "MDDA-LIT-%-PN"}
        )
        assert response.status_code == 200, (endpoint, response.text[:300])
    rows = {
        row["pn"]: row
        for row in client.get(
            f"{BASE}/parts", params={"range": "all", "q": "MDDA-LIT-%-PN"}
        ).json()["rows"]
    }
    assert set(rows) == {"MDDA-LIT-%-PN"}
    assert rows["MDDA-LIT-%-PN"]["bad_receipt_qty"] == "2.000"
    rows_us = {
        row["pn"]: row
        for row in client.get(
            f"{BASE}/parts", params={"range": "all", "q": "MDDA-US_A"}
        ).json()["rows"]
    }
    assert set(rows_us) == {"MDDA-US_A"}  # MDDA-USXA（_ 当单字通配）不得误中
    assert rows_us["MDDA-US_A"]["bad_receipt_qty"] == "3.000"
    # 反斜杠字面：原始未解析行命中带反斜杠者；无反斜杠诱饵不得命中。
    rows_bs = {
        row["pn"]: row
        for row in client.get(
            f"{BASE}/receipts", params={"range": "all", "q": "MDDA-RAW\\-LIT"}
        ).json()["rows"]
    }
    assert set(rows_bs) == {"MDDA-RAW\\-LIT"}
    assert rows_bs["MDDA-RAW\\-LIT"]["qty"] == "11.000"
    # 导出同参路径（B-F01/B-F02 服务层共用 seam）。
    export = client.get(
        f"{BASE}/export",
        params={"kind": "parts", "format": "csv", "range": "all",
                "q": "MDDA-LIT-%-PN"},
    )
    assert export.status_code == 200, export.text[:200]
    text = export.content.decode("utf-8-sig")
    assert "MDDA-LIT-%-PN" in text
    assert "MDDA-LIT-ABC-PN" not in text


def test_quantity_checks_endpoint_and_unassigned(db, seeded):
    p1, _p2, part = seeded
    # 未关联需求单的实收（全范围账号可见汇总）。
    client = client_for(db, username="mdd-admin-t4", role="admin")
    response = client.get(f"{BASE}/quantity-checks")
    assert response.status_code == 200
    body = response.json()
    assert body["cards"]["unassigned_receipt_qty"] == "6.000"
    assert body["cards"]["note"].startswith("当前累计")
    for row in body["rows"]:
        assert not ({"status", "return_status", "overdue"} & set(row))


def test_export_csv_formula_guard_and_policy_column(db, seeded):
    p1, _p2, _seed_part = seeded
    formula = _part(db, "=1+1", unit="件")
    negative = _part(db, "-5", unit="件")
    _receipt(db, p1, formula, "1", condition="坏品",
             occurred_at=datetime(2026, 10, 18, 10, tzinfo=UTC))
    _receipt(db, p1, negative, "1", condition="坏品",
             occurred_at=datetime(2026, 10, 18, 11, tzinfo=UTC))
    db.commit()
    client = client_for(db, username="mdd-admin-csv", role="admin")
    export = client.get(f"{BASE}/export",
                        params={"kind": "parts", "format": "csv",
                                "value_retention_ratio": "0.5"})
    assert export.status_code == 200
    assert export.headers["content-type"].startswith("text/csv")
    text = export.content.decode("utf-8")
    header = next(csv.reader(io.StringIO(text)))
    assert "坏件试算价值-含税" in header
    assert "估价政策版本" in header
    assert "价值保留比例" in header
    # 公式形态 PN 加撇号；纯负数 PN（如 -5）不误伤。
    assert "'=1+1" in text
    assert "'-5" not in text
    assert "-5" in text
    # 比例伴随导出：值列与政策版本列成对出现。
    rows = list(csv.reader(io.StringIO(text)))
    ratio_index = header.index("价值保留比例")
    assert {row[ratio_index] for row in rows[1:]} == {"0.5"}


def test_export_xlsx_numbers_and_amount_gating(db, seeded):
    from openpyxl import load_workbook

    client = client_for(db, username="mdd-admin-xlsx", role="admin")
    # 不带比例参数：缺省即按 0.5 试算并回显（Sunny 01:20）。
    export = client.get(f"{BASE}/export",
                        params={"kind": "projects", "format": "xlsx"})
    assert export.status_code == 200
    assert export.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    book = load_workbook(io.BytesIO(export.content))
    sheet = book.active
    header = [cell.value for cell in sheet[1]]
    assert "坏件试算价值-含税" in header
    assert "估价政策版本" in header
    row2 = [cell.value for cell in sheet[2]]
    value_cell = sheet[2][header.index("坏件试算价值-含税")]
    assert isinstance(value_cell.value, (int, float))  # 数值写数值单元格
    assert row2[header.index("估价政策版本")] == POLICY_VERSION
    assert row2[header.index("价值保留比例")] == "0.5"  # 缺省回显 0.5
    explicit = client.get(
        f"{BASE}/export",
        params={"kind": "projects", "format": "xlsx",
                "value_retention_ratio": "0.3"},
    )
    explicit_row = [cell.value for cell in load_workbook(
        io.BytesIO(explicit.content)).active[2]]
    assert explicit_row[header.index("价值保留比例")] == "0.3"  # 显式值不被默认覆盖

    # 无金额权：XLSX 同样不含金额列。
    username = "mdd-manager-xlsx"
    manager = client_for(db, username=username, role="maintenance_manager")
    _assign_manager(db, username, seeded[0])
    plain = manager.get(f"{BASE}/export",
                        params={"kind": "projects", "format": "xlsx"})
    assert plain.status_code == 200
    plain_book = load_workbook(io.BytesIO(plain.content))
    plain_header = [cell.value for cell in plain_book.active[1]]
    assert "坏件试算价值-含税" not in plain_header
    assert "估价政策版本" not in plain_header
    assert "价值保留比例" not in plain_header
    assert "返件实收合计" in plain_header


def test_receipts_test_result_filter(db, seeded):
    client = client_for(db, username="mdd-admin-rf", role="admin")
    response = client.get(f"{BASE}/receipts", params={"test_result": "坏品"})
    assert response.status_code == 200
    rows = response.json()["rows"]
    assert len(rows) == 1
    assert rows[0]["test_result"] == "坏品"


# ---------- M01/M03 新端点与导出（D-0024） ----------

def _costed_issue(db, project, part, qty, *, issue_date=date(2026, 10, 18),
                  unit_cost_ex=Decimal("25.00"), cost_source="direct_purchase"):
    line = _issue(db, project, part, qty, issue_date=issue_date)
    unit_inc = (unit_cost_ex * Decimal("1.13")).quantize(Decimal("0.01"))
    line.unit_cost = unit_cost_ex
    line.unit_cost_ex_tax = unit_cost_ex
    line.unit_cost_inc_tax = unit_inc
    line.cost_amount = (Decimal(qty) * unit_cost_ex).quantize(Decimal("0.01"))
    line.cost_amount_ex_tax = line.cost_amount
    line.cost_amount_inc_tax = (Decimal(qty) * unit_inc).quantize(Decimal("0.01"))
    line.cost_source = cost_source
    db.flush()
    return line


@pytest.fixture()
def consumption_seeded(db):
    """P1：期内成本领用 4×25 + 未计领用 3；P2：成本领用 2×10（manual）；
    期内需求单两张（P1，10 + 5，另一张带退货 2 不扣）。"""
    from tests.test_maintenance_analytics import _line

    p1 = _project(db, "MDDA-CP1")
    p2 = _project(db, "MDDA-CP2")
    part = _part(db, "MDDA-CPN", unit="件")
    _costed_issue(db, p1, part, "4")
    _issue(db, p1, part, "3", issue_date=date(2026, 10, 18))
    _costed_issue(db, p2, part, "2", issue_date=date(2026, 10, 19),
                  unit_cost_ex=Decimal("10.00"), cost_source="manual")
    _line(db, p1, part, order_no="WBDD-MDDA-O1", order_date=date(2026, 10, 16),
          qty=Decimal("10"), return_qty=Decimal("2"))
    _line(db, p1, part, order_no="WBDD-MDDA-O2", order_date=date(2026, 10, 17),
          qty=Decimal("5"))
    db.commit()
    return p1, p2, part


def test_orders_endpoint_admin_and_scope(db, consumption_seeded):
    p1, p2, _part = consumption_seeded
    client = client_for(db, username="mdd-admin-ord", role="admin")
    response = client.get(f"{BASE}/orders")
    assert response.status_code == 200
    body = response.json()
    # 原 qty 不扣退货：10 + 5 = 15；单数 2。单位分桶 + 无日期单独立报告
    # （B-F09/B-F12：单一单位平面兼容；undated 单列不冒充入期）。
    assert body["cards"] == {
        "order_qty": "15.000",
        "unit": "件",
        "unit_breakdown": [{"unit": "件", "order_qty": "15.000"}],
        "order_count": 2,
        "undated_orders": {
            "order_qty": "0.000",
            "order_count": 0,
            "unit_breakdown": [],
            "note": "无制单日期下单未计入期间（不猜日期）",
        },
    }
    (row,) = body["rows"]
    assert row["project_id"] == p1.project_id
    assert row["order_qty"] == "15.000"
    assert row["order_count"] == 2
    assert row["unit"] == "件"
    assert row["undated_order_qty"] == "0.000"
    assert row["undated_order_count"] == 0
    assert _forbidden_keys(body) == set()  # 数量域无金额键（任何角色）

    # 项目筛选与授权交集。
    only_p2 = client.get(f"{BASE}/orders", params={"project": p2.project_id})
    assert only_p2.json()["cards"] == {
        "order_qty": "0.000",
        "unit": None,
        "unit_breakdown": [],
        "order_count": 0,
        "undated_orders": {
            "order_qty": "0.000",
            "order_count": 0,
            "unit_breakdown": [],
            "note": "无制单日期下单未计入期间（不猜日期）",
        },
    }
    assert only_p2.json()["rows"] == []


def test_consumption_endpoints_amounts_gating(db, consumption_seeded):
    p1, p2, _part = consumption_seeded
    client = client_for(db, username="mdd-admin-cons", role="admin")
    daily = client.get(f"{BASE}/consumption")
    assert daily.status_code == 200
    body = daily.json()
    cards = body["cards"]
    assert cards["issued_qty"] == "9.000"
    # 已知小计：ex = 4×25 + 2×10 = 120.00；inc = 113.00 + 22.60 = 135.60。
    assert cards["cost_amount_ex_tax"] == "120.00"
    assert cards["cost_amount_inc_tax"] == "135.60"
    assert cards["uncosted_qty"] == "3.000"
    assert cards["coverage_pct"] == 66.7  # 6/9 已计，独立算得
    assert cards["sources"] == ["direct_purchase", "manual"]
    assert "只读" in cards["note"]
    by_date = {row["issue_date"]: row for row in body["rows"]}
    day18 = by_date["2026-10-18"]
    assert day18["issued_qty"] == "7.000"
    assert day18["uncosted_qty"] == "3.000"
    assert day18["sources"] == ["direct_purchase"]

    # 下钻：日 × 项目过滤；行级成本证据可核对。
    drill = client.get(f"{BASE}/consumption/issues",
                       params={"issue_date": "2026-10-19"})
    assert drill.status_code == 200
    (drill_row,) = drill.json()["rows"]
    assert drill_row["project_id"] == p2.project_id
    assert drill_row["cost_source"] == "manual"
    assert drill_row["unit_cost_ex_tax"] == "10.00"
    assert drill_row["cost_amount_inc_tax"] == "22.60"
    # 未计成本行：cost_source null 且无金额键（不填 0）。
    uncosted_day = client.get(f"{BASE}/consumption/issues",
                              params={"issue_date": "2026-10-18"}).json()
    uncosted = next(row for row in uncosted_day["rows"]
                    if row["qty"] == "3.000")
    assert uncosted["cost_source"] is None
    assert "cost_amount_ex_tax" not in uncosted
    assert "unit_cost_ex_tax" not in uncosted

    # 越权 project_id：与授权范围取交集 → 空（admin 全范围 + 外部 id 无单）。
    foreign = client.get(f"{BASE}/consumption/issues",
                         params={"project_id": "not-a-project"})
    assert foreign.json()["rows"] == []

    # 数量角色：数量键在、成本/覆盖/来源/note 键整个不出现。
    username = "mdd-manager-cons"
    manager = client_for(db, username=username, role="maintenance_manager")
    _assign_manager(db, username, p1)
    manager_daily = manager.get(f"{BASE}/consumption")
    assert manager_daily.status_code == 200
    assert _forbidden_keys(manager_daily.json()) == set()
    manager_cards = manager_daily.json()["cards"]
    assert manager_cards["issued_qty"] == "7.000"
    assert "cost_amount_ex_tax" not in manager_cards
    assert "coverage_pct" not in manager_cards
    assert "sources" not in manager_cards
    assert "note" not in manager_cards
    manager_rows = manager.get(f"{BASE}/consumption/issues").json()["rows"]
    assert {row["qty"] for row in manager_rows} == {"4.000", "3.000"}
    assert all("cost_source" not in row for row in manager_rows)


def test_export_new_kinds_orders_and_consumption(db, consumption_seeded):
    p1, _p2, _part = consumption_seeded
    client = client_for(db, username="mdd-admin-exp2", role="admin")

    orders = client.get(f"{BASE}/export",
                        params={"kind": "orders", "format": "csv"})
    assert orders.status_code == 200
    rows = list(csv.reader(io.StringIO(orders.content.decode("utf-8-sig"))))
    assert rows[0] == ["项目编码", "项目名称", "业务类型", "单位",
                       "期间下单量（原qty，不扣退货/免返）", "期间下单单数",
                       "无制单日期下单量（未计入期间）", "无制单日期单数"]
    assert rows[1][0] == "MDDA-CP1"
    assert rows[1][3] == "件"  # B-F09：单一单位照给
    assert rows[1][4] == "15.000"  # 数值列 Decimal 原样，不丢精度口径
    assert rows[1][5] == "2"
    assert rows[1][6] == "0.000"  # B-F12：无日期单单列，本夹具为 0
    assert rows[1][7] == "0"

    daily = client.get(f"{BASE}/export",
                       params={"kind": "consumption_daily", "format": "csv"})
    assert daily.status_code == 200
    header = next(csv.reader(io.StringIO(daily.content.decode("utf-8-sig"))))
    for column in ("业务日", "项目编码", "领用量", "已计成本-含税",
                   "未计成本量", "成本来源"):
        assert column in header

    issues = client.get(f"{BASE}/export",
                        params={"kind": "consumption_issues", "format": "csv"})
    assert issues.status_code == 200
    text = issues.content.decode("utf-8")
    header = next(csv.reader(io.StringIO(text.lstrip("﻿"))))
    for column in ("业务日", "领用单号", "PN", "数量", "成本来源",
                   "不含税成本", "含税成本"):
        assert column in header
    assert "113.00" in text  # 4×25×1.13 可核对

    # 数量角色：新导出同样无金额列。
    username = "mdd-manager-exp2"
    manager = client_for(db, username=username, role="maintenance_manager")
    _assign_manager(db, username, p1)
    for kind in ("orders", "consumption_daily", "consumption_issues"):
        plain = manager.get(f"{BASE}/export",
                            params={"kind": kind, "format": "csv"})
        assert plain.status_code == 200, kind
        plain_header = next(
            csv.reader(io.StringIO(plain.content.decode("utf-8-sig")))
        )
        for forbidden in ("已计成本-含税", "已计成本-不含税", "成本来源",
                          "含税成本", "不含税成本", "不含税单价", "含税单价",
                          "未计成本量", "已计成本覆盖率%", "参考侧"):
            assert forbidden not in plain_header, (kind, forbidden)


def test_cost_only_valuation_export_consistency(db):
    """P-B01 追加：市场缺样、成本域有效时——receipts/parts/summary/导出
    的成本域并列值互不矛盾（主展示仍待估，不填 0）。"""
    p1 = _project(db, "MDDA-CVO")
    part = _part(db, "MDDA-CVO-PN", unit="件")
    # 领用 qty 10 × ex 40（inc 45.20）→ 单位成本 45.20/40.00。
    _costed_issue(db, p1, part, "10", unit_cost_ex=Decimal("40.00"))
    _receipt(db, p1, part, "2", condition="坏品",
             occurred_at=datetime(2026, 10, 18, 8, tzinfo=UTC))
    db.commit()
    client = client_for(db, username="mdd-admin-cvo", role="admin")
    ratio = {"value_retention_ratio": "0.5"}

    summary = client.get(f"{BASE}/summary", params=ratio).json()
    card = summary["cards"]["valuation"]
    assert card["pending_bad_qty"] == "2.000"
    assert card["domains"]["cost"]["value_inc_tax"] == "45.20"

    parts = client.get(f"{BASE}/parts", params=ratio).json()["rows"]
    valuation = parts[0]["valuation"]
    assert valuation["state"] == "pending"
    assert "bad_estimated_value_inc_tax" not in valuation
    assert valuation["domains"]["cost"]["value_inc_tax"] == "45.20"

    receipts = client.get(f"{BASE}/receipts", params=ratio).json()["rows"]
    (line,) = receipts
    assert line["valuation"]["state"] == "pending"
    assert line["valuation"]["domains"]["cost"]["value_inc_tax"] == "45.20"

    export = client.get(f"{BASE}/export",
                        params={"kind": "parts", "format": "csv", **ratio})
    rows = list(csv.reader(io.StringIO(export.content.decode("utf-8"))))
    header = rows[0]
    cost_column = header.index("三域-领用已计成本-价值-含税")
    assert rows[1][cost_column] == "45.20"
    value_column = header.index("坏件试算价值-含税")
    assert rows[1][value_column] == ""  # 主展示待估 → 空单元格，不是 0


def test_api_module_does_not_import_giant_maintenance_api():
    """生产隔离红线：api.maintenance 无 _safe_formula_only，也不得把巨型
    API 拖进看板（会带入无关 #432 业务）；净化函数走独立导出模块。"""
    import ast
    from pathlib import Path

    api_file = Path(__file__).resolve().parent.parent / "app" / "api" / "maintenance_dashboard.py"
    source = api_file.read_text()
    tree = ast.parse(source)
    for node in ast.walk(tree):
        if isinstance(node, ast.ImportFrom) and node.module:
            assert node.module != "app.api.maintenance", (
                "禁止 import 生产巨型 api.maintenance（含 _safe_formula_only "
                "等私有 helper）"
            )
            assert not node.module.startswith("app.api.maintenance."), node.module
        if isinstance(node, ast.Import):
            for alias in node.names:
                assert alias.name != "app.api.maintenance"
    assert "from app.services.maintenance_dashboard_export_text import" in source
