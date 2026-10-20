"""#432 审查修复：导出数值类型、CSV 仅转义公式、金额遮蔽、需求数量空按 0、直达项目范围。"""

import csv
import io
import json
from decimal import Decimal

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select

from app import permissions
from app.models.maintenance import FMaintenanceLine, MaintenanceManualCostOverride
from app.security import UserContext
from tests.boss_board_helpers import client_for, make_project
from tests.test_maintenance_demand_export_ordering import export_data  # noqa: F401
from tests.test_maintenance_row_filters import BASE, WINDOW, _assign_user, _project_path


PN = "PN-early-one-a"
NO_AMOUNT = {"page_maintenance": True, "data_profit": False, "data_purchase_cost": False}


def _line(db, pn=PN):
    return db.scalar(select(FMaintenanceLine).where(FMaintenanceLine.pn_std == pn))


def _only(pn=PN, *extra):
    return json.dumps([{"field": "pn_std", "op": "eq", "value": pn}, *extra])


def _csv(response):
    assert response.status_code == 200, response.text
    text = response.content.decode("utf-8-sig")
    return list(csv.reader(io.StringIO(text)))


def _xlsx(response):
    assert response.status_code == 200, response.text
    wb = load_workbook(io.BytesIO(response.content), data_only=False)
    try:
        header = [c.value for c in wb.active[1]]
        rows = [list(row) for row in wb.active.iter_rows(min_row=2)]
    finally:
        wb.close()
    return header, rows


def _manual_cost(db, line, ex="12.34", inc="13.94"):
    db.add(MaintenanceManualCostOverride(
        line_id=line.id, unit_cost_ex_tax=Decimal(ex), unit_cost_inc_tax=Decimal(inc),
        reason="审查修复真实成本", active=True,
    ))


def _exports(client, project, filters):
    yield BASE + "/export", {**WINDOW, "filters": filters}
    yield _project_path(project) + "/rows/export", {"filters": filters}


def test_xlsx_writes_numbers_as_numeric_cells_and_text_as_text(db, export_data):
    project, _ = export_data
    line = _line(db)
    line.return_qty = Decimal("-1.500")
    line.description = "-1"
    line.line_note = "=1+1"
    _manual_cost(db, line)
    db.commit()
    client = client_for(db, username="xlsx-number-admin", role="admin")
    for path, params in _exports(client, project, _only()):
        header, rows = _xlsx(client.get(path, params={**params, "format": "xlsx"}))
        (row,) = rows
        cell = {label: row[index] for index, label in enumerate(header)}
        for label, expected in (
            ("需求数量", 1), ("退货数量", -1.5), ("人工未税单位成本", 12.34),
            ("人工含税单位成本", 13.94),
        ):
            assert cell[label].value is not None, (path, label)
            assert cell[label].data_type == "n", (path, label, cell[label].value)
            assert cell[label].value == pytest.approx(expected), (path, label)
        # 文本保持文本：公式样式不执行，数字样式文本不变成数值。
        assert (cell["备注"].data_type, cell["备注"].value) == ("s", "=1+1")
        assert (cell["描述"].data_type, cell["描述"].value) == ("s", "-1")
        assert cell["制单日期"].value == "2026-07-01"


def test_csv_keeps_ordinary_signed_values_and_escapes_only_formulas(db, export_data):
    project, _ = export_data
    line = _line(db)
    line.description = "-48V"
    line.line_note = "+5V"
    line.serial_numbers = "SN-1"
    line.return_qty = Decimal("-1.500")
    other = _line(db, "PN-early-one-b")
    other.description = "=1+1"
    other.line_note = "@SUM(A1)"
    other.serial_numbers = "-2+3+cmd|' /C calc'!A0"
    other.return_old_part = "  =A1+1"  # varchar(16)
    db.commit()
    client = client_for(db, username="csv-formula-admin", role="admin")
    filters = json.dumps([{"field": "pn_std", "op": "prefix", "value": "PN-early-one-"}])
    for path, params in _exports(client, project, filters):
        header, *rows = _csv(client.get(path, params={**params, "format": "csv"}))
        by_pn = {row[header.index("PN")]: dict(zip(header, row)) for row in rows}
        plain, formula = by_pn["PN-early-one-a"], by_pn["PN-early-one-b"]
        assert plain["描述"] == "-48V"
        assert plain["备注"] == "+5V"
        assert plain["SN"] == "SN-1"
        assert plain["退货数量"] == "-1.500"
        assert plain["需求数量"] == "1.000"
        assert formula["描述"] == "'=1+1"
        assert formula["备注"] == "'@SUM(A1)"
        assert formula["SN"] == "'-2+3+cmd|' /C calc'!A0"
        assert formula["退返旧件"] == "'  =A1+1"


@pytest.mark.parametrize("v2", [True, False])
def test_project_viewer_never_sees_real_costs_in_rows_or_exports(db, export_data, monkeypatch, v2):
    from app.api import maintenance_project_master_workbook as master_api
    monkeypatch.setattr(master_api.get_settings(), "maintenance_project_master_v2_enabled", v2)
    project, _ = export_data
    line = _line(db)
    line.cost_source = "direct"
    line.cost_tax_basis = "ex"
    line.unit_cost = line.unit_cost_ex_tax = line.cost_amount = line.cost_amount_ex_tax = Decimal("12.34")
    line.unit_cost_inc_tax = line.cost_amount_inc_tax = Decimal("13.94")
    line.confidence = "high"
    _manual_cost(db, _line(db, "PN-early-one-b"), ex="56.78", inc="64.16")
    db.commit()
    admin = client_for(db, username=f"mask-admin-{v2}", role="admin")
    viewer = client_for(db, username=f"mask-viewer-{v2}", overrides=NO_AMOUNT)
    _assign_user(db, f"mask-viewer-{v2}", project, "viewer")
    rows_params = {"sheet": "03_备件订单", "filters": json.dumps([
        {"field": "pn_std", "op": "prefix", "value": "PN-early-one-"},
    ])}
    costed = admin.get(_project_path(project) + "/rows", params=rows_params)
    assert costed.status_code == 200, costed.text
    by_pn = {row["pn_std"]: row for row in costed.json()["rows"]}
    # 夹具必须带真实成本，否则遮蔽断言恒真。
    for pn, ex, inc in (("PN-early-one-a", "12.34", "13.94"), ("PN-early-one-b", "56.78", "64.16")):
        assert Decimal(str(by_pn[pn]["unit_cost_ex_tax"])) == Decimal(ex)
        assert Decimal(str(by_pn[pn]["cost_amount_inc_tax"])) == Decimal(inc)
    masked = viewer.get(_project_path(project) + "/rows", params=rows_params)
    assert masked.status_code == 200, masked.text
    assert {row["pn_std"] for row in masked.json()["rows"]} == set(by_pn)
    for row in masked.json()["rows"]:
        for key in ("unit_cost_ex_tax", "unit_cost_inc_tax", "cost_amount_inc_tax"):
            assert row[key] is None, (row["pn_std"], key)
    secrets = ("12.34", "13.94", "56.78", "64.16")
    admin_csv = admin.get(_project_path(project) + "/rows/export", params={
        "filters": rows_params["filters"], "format": "csv",
    })
    assert any(s in admin_csv.content.decode("utf-8-sig") for s in secrets)
    for format in ("csv", "xlsx"):
        response = viewer.get(_project_path(project) + "/rows/export", params={
            "filters": rows_params["filters"], "format": format,
        })
        if format == "csv":
            header, *rows = _csv(response)
            values = [value for row in rows for value in row]
        else:
            header, rows = _xlsx(response)
            values = [str(cell.value) for row in rows for cell in row]
        assert len(rows) == 2
        assert not any("成本" in str(label) for label in header), header
        assert not any(s in value for s in secrets for value in values)


def test_null_demand_quantity_counts_as_zero_in_numeric_filters(db, export_data):
    project, _ = export_data
    _line(db, "PN-early-one-a").qty = None
    _line(db, "PN-early-one-b").qty = Decimal("0")
    _line(db, "PN-early-two").qty = Decimal("-2")
    db.commit()
    client = client_for(db, username="null-qty-admin", role="admin")
    scope = {"field": "pn_std", "op": "prefix", "value": "PN-early-"}
    cases = [
        ({"op": "eq", "value": "0"}, {"one-a", "one-b"}),
        ({"op": "range", "max": "0"}, {"one-a", "one-b", "two"}),
        ({"op": "range", "min": "0", "max": "0"}, {"one-a", "one-b"}),
        ({"op": "range", "min": "-5", "max": "0.5"}, {"one-a", "one-b", "two"}),
        ({"op": "range", "min": "0.5"}, {"null-a", "null-b"}),
        ({"op": "eq", "value": "1"}, {"null-a", "null-b"}),
        ({"op": "is_empty"}, {"one-a"}),
    ]
    for condition, expected in cases:
        filters = json.dumps([scope, {"field": "qty", **condition}])
        expected_pns = {f"PN-early-{suffix}" for suffix in expected}
        for path, params in (
            (BASE + "/rows", {**WINDOW, "page_size": 200}),
            (_project_path(project) + "/rows", {"sheet": "03_备件订单"}),
        ):
            response = client.get(path, params={**params, "filters": filters})
            assert response.status_code == 200, response.text
            assert {row["pn_std"] for row in response.json()["rows"]} == expected_pns, (path, condition)
        header, *rows = _csv(client.get(BASE + "/export", params={
            **WINDOW, "filters": filters, "format": "csv",
        }))
        assert {row[header.index("PN")] for row in rows} == expected_pns, condition
    # 已返数量等其他数值字段不套用该口径：空仍不等于 0。
    _line(db, "PN-early-one-a").returned_qty = None
    _line(db, "PN-early-one-b").returned_qty = Decimal("0")
    db.commit()
    response = client.get(BASE + "/rows", params={**WINDOW, "filters": json.dumps([
        scope, {"field": "returned_qty", "op": "eq", "value": "0"},
    ])})
    assert "PN-early-one-a" not in {row["pn_std"] for row in response.json()["rows"]}


def test_direct_project_query_enforces_project_scope(db, export_data):
    from app.services import maintenance_row_workbench as workbench
    project, _ = export_data
    own = make_project(db, "范围校验本人项目")
    client_for(db, username="scope-direct-viewer", overrides=NO_AMOUNT)
    _assign_user(db, "scope-direct-viewer", own, "viewer")
    ctx = UserContext(
        user_id="scope-direct-viewer", role="readonly", is_authenticated=True,
        permissions=permissions.effective_from_snapshot(
            permissions.effective("readonly", None), NO_AMOUNT,
        ),
    )
    stmt = workbench.query(db, ctx, project_id=own.project_id, direct_project=True)
    assert workbench.count(db, stmt) == 0
    # 直达查询不带可见项目集合，范围校验是唯一闸门；缺失时会返回该项目 30 行。
    with pytest.raises(HTTPException) as denied:
        workbench.query(db, ctx, project_id=project.project_id, direct_project=True)
    assert denied.value.status_code == 403
    with pytest.raises(HTTPException) as missing:
        workbench.query(db, ctx, project_id="no-such-project", direct_project=True)
    assert missing.value.status_code == 404
    admin = UserContext(user_id="scope-admin", role="admin", is_authenticated=True)
    stmt = workbench.query(db, admin, project_id=project.project_id, direct_project=True)
    assert workbench.count(db, stmt) == 30
