"""Global uploaded rows and milestone restore use actual whole-object contracts."""
import io,json
from datetime import date
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest
from sqlalchemy import select,func,update,text,event
from sqlalchemy.orm import Session
from app.api import maintenance_project_master_workbook as api
from app.db import get_db
from app.models.authorization import SysAuthorizationGrant
from app.models.maintenance import MaintenanceManualCostOverride
from app.models.maintenance_manager import MaintenanceCollectionMilestone
from app.models.maintenance_source_assignment import MaintenanceSourceOrderAssignment
from app.models.system import SysUser
from app.services import maintenance_project_master_workbook as business
from app.services import maintenance_global_lines_access as protected
from tests.test_authz_s3_projects import canonical
from tests.test_authz_s3_project_foundation import projects
from tests.authz_pdp_helpers import cerbos_url
from tests.test_maintenance_project_master_workbook import project_with_lines
from tests.test_maintenance_collection_milestone_restore import _seed,_payload
from tests.test_authz_s3_finance import _pre


@pytest.fixture
def client(db,canonical):
    app=FastAPI();app.include_router(api.router,prefix='/api');app.dependency_overrides[get_db]=lambda:db
    return TestClient(app),canonical[1]


def grant(db,actor,resource,actions,fields,scope='all_projects'):
    db.add(SysAuthorizationGrant(user_id=actor.id,resource=resource,actions=actions,scope=scope,field_permissions=fields,version=1));db.commit()


def workbook(db,project):
    content=business.build_global_lines(db,preset='custom',date_from=date(2026,1,1),date_to=date(2026,12,31),allowed_project_ids={project.project_id})
    wb=load_workbook(io.BytesIO(content));wb[business.GLOBAL_SHEET].cell(2,10,88);out=io.BytesIO();wb.save(out);db.commit();return out.getvalue()


def upload(client,headers,path,content,**extra):
    return client.post('/api/maintenance/spare-part-lines/'+path,headers=headers,
        files={'file':('global.xlsx',content,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},**extra)


def test_global_validate_and_apply_are_independent_complete_contracts(db,client,projects,project_with_lines):
    c,h=client;actor=projects[0];data=workbook(db,project_with_lines)
    assert upload(c,h,'validate',data).status_code==403
    grant(db,actor,protected.RESOURCE,['validate'],['data_purchase_cost','data_profit','data_customer'])
    view=upload(c,h,'validate?view=authorization',data);assert view.status_code==200,view.text
    assert set(view.json())=={'locator','authorization_context'}
    valid=upload(c,h,'validate',data,data={'authorization_precondition':json.dumps(_pre(view.json()))})
    assert valid.status_code==200,valid.text
    assert upload(c,h,'apply',data).status_code==403
    db.execute(update(SysAuthorizationGrant).where(SysAuthorizationGrant.user_id==actor.id,SysAuthorizationGrant.resource==protected.RESOURCE).values(actions=['apply'],version=2));db.commit()
    view=upload(c,h,'apply?view=authorization',data);assert view.status_code==200,view.text
    applied=upload(c,h,'apply',data,data={'authorization_precondition':json.dumps(_pre(view.json()))});assert applied.status_code==200,applied.text
    assert db.scalar(select(func.count()).select_from(MaintenanceManualCostOverride))==1
    replay=upload(c,h,'apply',data);assert replay.status_code==200,replay.text
    assert replay.json()['replayed_line_ids']


def test_global_whole_parent_scope_and_input_precondition(db,client,projects,project_with_lines):
    c,h=client;actor=projects[0];data=workbook(db,project_with_lines)
    grant(db,actor,protected.RESOURCE,['apply'],['data_purchase_cost','data_profit','data_customer'],'assigned_or_sales_projects')
    assert upload(c,h,'apply',data).status_code==403
    db.execute(update(SysAuthorizationGrant).where(SysAuthorizationGrant.user_id==actor.id,SysAuthorizationGrant.resource==protected.RESOURCE).values(scope='all_projects',version=2));db.commit()
    view=upload(c,h,'apply?view=authorization',data)
    wb=load_workbook(io.BytesIO(data));wb[business.GLOBAL_SHEET].cell(2,10,99);out=io.BytesIO();wb.save(out)
    stale=upload(c,h,'apply',out.getvalue(),data={'authorization_precondition':json.dumps(_pre(view.json()))})
    assert stale.status_code==409,stale.text
    assert db.scalar(select(func.count()).select_from(MaintenanceManualCostOverride))==0


def test_global_late_epoch_revocation_rolls_back_pending_override(db,client,projects,project_with_lines,monkeypatch):
    c,h=client;actor=projects[0];data=workbook(db,project_with_lines)
    grant(db,actor,protected.RESOURCE,['apply'],['data_purchase_cost','data_profit','data_customer'])
    real=protected.GlobalLinesAuthorization.authorize_after
    def late(self,session):
        session.flush();assert session.scalar(select(func.count()).select_from(MaintenanceManualCostOverride))==1
        with Session(session.get_bind()) as writer:
            writer.execute(text("SET LOCAL lock_timeout='1500ms'"));writer.execute(update(SysUser).where(SysUser.id==actor.id).values(token_version=actor.token_version+1));writer.commit()
        return real(self,session)
    monkeypatch.setattr(protected.GlobalLinesAuthorization,'authorize_after',late)
    response=upload(c,h,'apply',data);assert response.status_code==401,response.text
    assert db.scalar(select(func.count()).select_from(MaintenanceManualCostOverride))==0


def test_restore_f_only_named_action_zero_write_view_and_current_replay(db,client,projects):
    c,h=client;actor=projects[0];project,contract,milestones=_seed(db);body=_payload(contract,milestones)
    path=f'/api/maintenance/projects/stable/{project.project_id}/collection-milestones/restore'
    assert c.post(path,headers=h,json=body).status_code==403

    grant(db,actor,'maintenance_collection',['restore_plan'],['data_profit'])
    view=c.post(path+'?view=authorization',headers=h,json=body);assert view.status_code==200,view.text
    assert set(view.json())=={'locator','authorization_context'}
    assert db.scalar(select(func.count()).select_from(MaintenanceCollectionMilestone).where(MaintenanceCollectionMilestone.is_active.is_(True)))==0
    restored=c.post(path,headers=h,json={**body,'authorization_precondition':_pre(view.json())});assert restored.status_code==200,restored.text
    assert restored.json()['restored_count']==2
    replay=c.post(path,headers=h,json=body);assert replay.status_code==200,replay.text
    assert replay.json()['idempotent_replay_count']==2
    db.execute(update(SysAuthorizationGrant).where(SysAuthorizationGrant.user_id==actor.id,SysAuthorizationGrant.resource=='maintenance_collection').values(actions=[],version=2));db.commit()
    assert c.post(path,headers=h,json=body).status_code==403


def test_global_missing_complete_field_denies_before_business_facts(db,client,projects,project_with_lines):
    c,h=client;actor=projects[0];data=workbook(db,project_with_lines)
    grant(db,actor,protected.RESOURCE,['validate'],['data_purchase_cost','data_profit'])
    sql=[]
    def capture(conn,cursor,statement,parameters,context,executemany):sql.append(statement.lower())
    event.listen(db.get_bind(),'before_cursor_execute',capture)
    try:response=upload(c,h,'validate',data)
    finally:event.remove(db.get_bind(),'before_cursor_execute',capture)
    assert response.status_code==403,response.text
    assert not any('f_maintenance_line.pn' in q or 'f_maintenance_line.unit_cost' in q for q in sql)


def test_restore_invalid_child_whole_batch_is_zero_write(db,client,projects):
    c,h=client;actor=projects[0];project,contract,milestones=_seed(db);body=_payload(contract,milestones)
    grant(db,actor,'maintenance_collection',['restore_plan'],['data_profit'])
    body['items'][1]['planned_amount']='1.00'
    response=c.post(f'/api/maintenance/projects/stable/{project.project_id}/collection-milestones/restore',headers=h,json=body)
    assert response.status_code==409,response.text
    assert db.scalar(select(func.count()).select_from(MaintenanceCollectionMilestone).where(MaintenanceCollectionMilestone.is_active.is_(True)))==0
