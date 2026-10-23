"""Independent W1b consumer checks using real XLSX, HTTP and PostgreSQL.

Exercise codec composition across application boundaries, rather than repeating
the unit-level byte vectors or alias checks. No paused business policy is added.
"""
from copy import deepcopy
from datetime import date
from decimal import Decimal
from io import BytesIO
import json

import pytest
from openpyxl import load_workbook
from openpyxl.utils.datetime import MAC_EPOCH, WINDOWS_EPOCH, to_excel
from pydantic import SecretStr
from sqlalchemy import event, select, text

from app.api import maintenance_project_master_workbook as master_api
from app.config import get_settings
from app.db import SessionLocal
from app.models.maintenance import FMaintenanceLine
from app.models.maintenance_project_operations import (
    MaintenanceCollectionSnapshot,
    MaintenanceProjectOperationAudit,
    MaintenanceProjectWorkbookOperation,
    MaintenanceProjectWorkbookState,
    MaintenanceSiteIssue,
    MaintenanceSiteIssueLine,
)
from app.services import maintenance_project_master_workbook as master
from tests.test_maintenance_expense_collection_workbook import _XLSX, uploader
from tests.test_maintenance_project_master_v2_editable import _make_project_with_line


_OLD_SECRET = "w1b-consumer-old-synthetic-only-1234567890"
_NEW_SECRET = "w1b-consumer-new-synthetic-only-1234567890"


@pytest.fixture(autouse=True)
def workbook_environment(monkeypatch):
    settings = get_settings()
    monkeypatch.setattr(master_api.get_settings(), "maintenance_project_master_v2_enabled", True)
    # Mutate the shared Settings object: both the relocated signer and the
    # remaining master metadata verifier must see the same runtime keyring.
    monkeypatch.setattr(settings, "maintenance_manifest_active_key_id", "consumer-old")
    monkeypatch.setattr(settings, "maintenance_manifest_active_hmac_key", SecretStr(_OLD_SECRET))
    monkeypatch.setattr(settings, "maintenance_manifest_previous_hmac_keys_json", SecretStr("none"))

    def bound(_session, _transaction, connection):
        connection.execute(text("SET LOCAL lock_timeout = '10s'"))

    event.listen(SessionLocal, "after_begin", bound)
    try:
        yield settings
    finally:
        event.remove(SessionLocal, "after_begin", bound)


def _rotate(settings, *, keep_old):
    settings.maintenance_manifest_active_key_id = "consumer-new"
    settings.maintenance_manifest_active_hmac_key = SecretStr(_NEW_SECRET)
    settings.maintenance_manifest_previous_hmac_keys_json = SecretStr(
        json.dumps({"consumer-old": _OLD_SECRET}) if keep_old else "none"
    )


def _setup(db, *, qty=Decimal("2")):
    project, part, order, line = _make_project_with_line(db, qty=qty)
    client = uploader(db, username="w1b-codec-consumer")
    route = f"/api/maintenance/projects/stable/{project.project_id}/master-workbook"
    return project, part, order, line, client, route


def _download(client, route, sheets):
    response = client.get(route + ".xlsx", params={"sheets": ",".join(sheets)})
    assert response.status_code == 200, response.text
    return load_workbook(BytesIO(response.content))


def _columns(sheet):
    return {cell.value: cell.column for cell in sheet[1] if cell.value}


def _save(book):
    result = BytesIO()
    book.save(result)
    return result.getvalue()


def _post(client, route, data, action="apply"):
    return client.post(route + "/" + action, files={"file": ("codec-consumer.xlsx", data, _XLSX)})


def _snapshot(db, project_id, line_id):
    db.expire_all()
    line = db.get(FMaintenanceLine, line_id)
    state = db.get(MaintenanceProjectWorkbookState, project_id)
    audits = list(db.scalars(select(MaintenanceProjectOperationAudit).where(
        MaintenanceProjectOperationAudit.project_id == project_id,
    ).order_by(MaintenanceProjectOperationAudit.id)))
    operations = list(db.scalars(select(MaintenanceProjectWorkbookOperation).where(
        MaintenanceProjectWorkbookOperation.project_id == project_id,
    ).order_by(MaintenanceProjectWorkbookOperation.id)))
    result = {
        "line": (line.qty, line.cost_amount_ex_tax, line.cost_amount_inc_tax,
                 deepcopy(line.manual_override)),
        "revision": None if state is None else state.revision,
        "audits": [(item.id, item.action, deepcopy(item.after_json)) for item in audits],
        "operations": [(item.id, item.operation_key) for item in operations],
    }
    # An assertion-side read transaction must not keep a lock across HTTP.
    db.rollback()
    return result


def test_http_rotated_old_file_applies_then_revoked_signature_still_replays_saved_receipt(
    db, workbook_environment,
):
    project, _part, _order, line, client, route = _setup(db)
    project_id, line_id = project.project_id, line.id
    book = _download(client, route, (master.V2_SHEET_PARTS,))
    sheet = book[master.V2_SHEET_PARTS]
    columns = _columns(sheet)
    sheet.cell(2, columns["需求数量"], "3.125")
    old_data = _save(book)
    _rotate(workbook_environment, keep_old=True)
    preview = _post(client, route, old_data, "validate")
    assert preview.status_code == 200, preview.text
    applied = _post(client, route, old_data)
    assert applied.status_code == 200, applied.text
    assert applied.json()["replayed"] is False
    after = _snapshot(db, project_id, line_id)
    assert after["line"][0] == Decimal("3.125")
    assert len(after["operations"]) == 1

    _rotate(workbook_environment, keep_old=False)
    # D-38 resolves an already completed, currently authorized request from its
    # persisted receipt before parsing today's revoked signing material.
    replay = _post(client, route, old_data)
    assert replay.status_code == 200, replay.text
    assert replay.json() == {**applied.json(), "replayed": True}
    assert _snapshot(db, project_id, line_id) == after

    # Different bytes are a new request and cannot piggyback on that receipt.
    book.properties.subject = "different file after key revocation"
    rejected = _post(client, route, _save(book))
    assert rejected.status_code == 422, rejected.text
    assert rejected.json()["detail"]["code"] == "template_signature_invalid"
    assert _snapshot(db, project_id, line_id) == after
    fresh = _download(client, route, (master.V2_SHEET_PARTS,))
    assert master._v2_meta(fresh)["metadata_hmac_key_id"] == "consumer-new"
    token = fresh[master.V2_SHEET_PARTS].cell(2, columns["基线令牌"])
    assert json.loads(token.value)["key_id"] == "consumer-new"


@pytest.mark.parametrize("revoked_layer", ["metadata", "row"])
def test_http_rejects_each_revoked_signature_layer_without_writes_then_accepts_restored_key(
    db, workbook_environment, revoked_layer,
):
    project, _part, _order, line, client, route = _setup(db)
    project_id, line_id = project.project_id, line.id
    book = _download(client, route, (master.V2_SHEET_PARTS,))
    sheet = book[master.V2_SHEET_PARTS]
    sheet.cell(2, _columns(sheet)["需求数量"], "3.125")
    _rotate(workbook_environment, keep_old=False)
    if revoked_layer == "row":
        # Rotate only the outer envelope, deliberately leaving the actual
        # signed 03 row token from the retired key intact.
        metadata = master._v2_meta(book)
        logical = [(key, value) for key, value in metadata.items()
                   if key not in {"metadata_hmac", "metadata_hmac_key_id", "metadata_hmac_algorithm"}]
        meta_sheet = book[master.V2_SHEET_META]
        meta_sheet.delete_rows(1, meta_sheet.max_row)
        for row in master._v2_sign_meta_rows(logical):
            meta_sheet.append(row)
    data = _save(book)
    before = _snapshot(db, project_id, line_id)
    expected = "template_signature_invalid" if revoked_layer == "metadata" else "invalid_concurrency_token"
    for action in ("validate", "apply"):
        rejected = _post(client, route, data, action)
        assert rejected.status_code == 422, rejected.text
        assert rejected.json()["detail"]["code"] == expected
        assert _snapshot(db, project_id, line_id) == before
    _rotate(workbook_environment, keep_old=True)
    accepted = _post(client, route, data)
    assert accepted.status_code == 200, accepted.text
    assert _snapshot(db, project_id, line_id)["line"][0] == Decimal("3.125")


@pytest.mark.parametrize("epoch", [WINDOWS_EPOCH, MAC_EPOCH], ids=["1900", "1904"])
def test_http_mixed_precision_epoch_and_auto_display_identity_survive_database_readback(db, epoch):
    project, part, _order, line, client, route = _setup(db)
    project_id, line_id = project.project_id, line.id
    sheets = (master.V2_SHEET_PARTS, master.V2_SHEET_RECEIPTS, master.V2_SHEET_SITE)
    book = _download(client, route, sheets)
    book.epoch = epoch
    parts = book[master.V2_SHEET_PARTS]
    parts.cell(2, _columns(parts)["需求数量"], "1.2345")
    receipts = book[master.V2_SHEET_RECEIPTS]
    columns = _columns(receipts)
    row_no = receipts.max_row + 1
    for key, value in {"合同编号": "XSDD-EDIT-001", "报告月份": to_excel(date(2026, 10, 1), epoch=epoch),
                       "累计实收金额（含税）": "3.005", "回款凭证号": "W1B-半分边界"}.items():
        receipts.cell(row_no, columns[key], value)
    site = book[master.V2_SHEET_SITE]
    columns = _columns(site)
    row_no = site.max_row + 1
    for key, value in {"PN": part.pn_std, "领用日期": to_excel(date(2026, 10, 23), epoch=epoch),
                       "领用数量": "2.345", "是否应返还": "是"}.items():
        site.cell(row_no, columns[key], value)
    data = _save(book)
    preview = _post(client, route, data, "validate")
    assert preview.status_code == 200, preview.text
    applied = _post(client, route, data)
    assert applied.status_code == 200, applied.text
    db.expire_all()
    assert db.get(FMaintenanceLine, line_id).qty == Decimal("1.234")
    receipt = db.scalar(select(MaintenanceCollectionSnapshot).where(
        MaintenanceCollectionSnapshot.project_id == project_id))
    assert receipt.cumulative_amount == Decimal("3.01")
    assert receipt.report_month == date(2026, 10, 1)
    issue, issue_line = db.execute(select(MaintenanceSiteIssue, MaintenanceSiteIssueLine).join(
        MaintenanceSiteIssueLine, MaintenanceSiteIssueLine.issue_id == MaintenanceSiteIssue.issue_id,
    ).where(MaintenanceSiteIssue.project_id == project_id)).one()
    identity = (issue.issue_id, issue.issue_no, issue.display_no, issue_line.issue_line_id)
    assert issue.issue_no.startswith("LY-AUTO-")
    assert issue.display_no.startswith("LY-20261023-")
    assert issue.issue_date == date(2026, 10, 23)
    assert issue_line.quantity == Decimal("2.35")
    db.rollback()

    downloaded = _download(client, route, sheets)
    site = downloaded[master.V2_SHEET_SITE]
    columns = _columns(site)
    assert site.cell(2, columns["领用单号"]).value == identity[2]
    assert site.cell(2, columns["实体ID"]).value == identity[3]
    assert Decimal(str(site.cell(2, columns["领用数量"]).value)) == Decimal("2.35")
    # A newly signed display projection must still match the internal identity
    # used for the deleted-row hash/guard; an unchanged re-export is clean.
    clean = _post(client, route, _save(downloaded), "validate")
    assert clean.status_code == 200, clean.text
    assert all(value == 0 for value in clean.json()["summary"].values())
    before = _snapshot(db, project_id, line_id)
    replay = _post(client, route, data)
    assert replay.status_code == 200 and replay.json()["replayed"] is True
    assert _snapshot(db, project_id, line_id) == before
    db.expire_all()
    same = db.get(MaintenanceSiteIssue, identity[0])
    assert (same.issue_id, same.issue_no, same.display_no) == identity[:3]


def test_http_locked_guard_preserves_third_decimal_change_committed_after_preview(db, monkeypatch):
    project, _part, _order, line, client, route = _setup(db, qty=Decimal("1.234"))
    project_id, line_id = project.project_id, line.id
    book = _download(client, route, (master.V2_SHEET_PARTS,))
    sheet = book[master.V2_SHEET_PARTS]
    sheet.cell(2, _columns(sheet)["需求数量"], "1.236")
    data = _save(book)
    before = _snapshot(db, project_id, line_id)
    original_validate = master.validate_project_master_v2
    observed = []

    def commit_after_real_preview(session, **kwargs):
        plan = original_validate(session, **kwargs)
        assert plan.row_guards and not plan.conflicts
        with SessionLocal() as peer:
            peer.get(FMaintenanceLine, line_id).qty = Decimal("1.235")
            peer.commit()
        observed.append(True)
        return plan

    monkeypatch.setattr(master, "validate_project_master_v2", commit_after_real_preview)
    rejected = _post(client, route, data)
    assert observed == [True]
    assert rejected.status_code == 409, rejected.text
    assert rejected.json()["detail"]["code"] == "row_conflicts"
    after = _snapshot(db, project_id, line_id)
    assert after["line"][0] == Decimal("1.235")
    assert after["line"][1:] == before["line"][1:]
    assert {key: value for key, value in after.items() if key != "line"} == {
        key: value for key, value in before.items() if key != "line"}

