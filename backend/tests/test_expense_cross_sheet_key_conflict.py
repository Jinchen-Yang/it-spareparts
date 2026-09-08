"""跨报销页同键不同值：静默保留首条 = 金额由 Sheet 顺序决定（2026-09-08 审查 P1）。

`pipeline.transform_workbook` 把一个工作簿里的多张报销页合并成一次装载，然后按
`raw_line_id` 跨页去重（`app/etl/pipeline.py` 的「跨页同键……保留首次，防 upsert 撞键」）。
但它**只比较键、不比较行内容**——注释里说的「完全相同的行」并没有真的校验。

于是：同一笔报销在两张页上金额不同（100 / 900），系统静默保留第一张页的 100，
错误清单为空。最终入库金额取决于 Sheet 在工作簿里的先后顺序，而不是任何业务事实。

对照：**同一张页内**的同键重复是有守卫的——`transform` 里的 `composite_seen` 会报
`duplicate_key`「同一文件内 单号+序号 重复」。跨页缺的正是这一道。

口径：同键同值 = 同一笔费用，保留首次（原行为，不变）；同键**不同值** = 事实冲突，
必须显式报错让人处置，不能替用户选一个。
"""

from __future__ import annotations

import io

from openpyxl import Workbook

from app.etl import pipeline

_HEADER = [
    "数据ID(不可修改)", "数据标题", "流程状态", "报销人员", "报销类别",
    "支出事由", "维保销售订单", "报销日期",
    "报销明细.数据ID(不可修改)", "报销明细.序号",
    "报销明细.费用分类", "报销明细.报销金额",
]


def _row(*, line_id: str, amount, order_no: str = "XSDD-20260101-0001") -> list:
    return [
        "BXD-DATA-1", "BXD-0001 张三", "已通过", "张三", "差旅",
        "出差", order_no, "2026-07-01",
        line_id, 1, "交通", amount,
    ]


def _workbook(sheets: list[tuple[str, list[list]]], tmp_path) -> str:
    workbook = Workbook()
    for index, (name, rows) in enumerate(sheets):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = name
        for row in rows:
            sheet.append(row)
    path = str(tmp_path / "expense.xlsx")
    workbook.save(path)
    return path


def _transform(tmp_path, sheets):
    loaded = pipeline.load_workbook(_workbook(sheets, tmp_path))
    assert len(loaded.expense_sheets) == 2, "夹具必须真的产生两张报销页"
    return pipeline.transform_workbook(loaded)


def _amounts(transformed) -> list:
    return [line.get("amount_inc_tax") or line.get("amount")
            for line in transformed.result.lines]


def test_same_key_different_amount_across_sheets_is_a_conflict(tmp_path):
    """两页同一条报销明细，金额 100 / 900：不能替用户挑一个。"""

    transformed = _transform(tmp_path, [
        ("报销一", [_HEADER, _row(line_id="LINE-1", amount=100)]),
        ("报销二", [_HEADER, _row(line_id="LINE-1", amount=900)]),
    ])

    codes = {error.error_type for error in transformed.result.errors}
    assert "cross_sheet_conflict" in codes, (
        f"跨页同键不同金额被静默吞掉了；lines={_amounts(transformed)} errors={codes}"
    )
    # 冲突行不得带着任意一方的金额进入写入计划
    assert _amounts(transformed) == []


def test_conflict_message_names_both_sheets_and_both_values(tmp_path):
    """报错要能让人直接去改：说清是哪两张页、两个值分别是多少。"""

    transformed = _transform(tmp_path, [
        ("报销一", [_HEADER, _row(line_id="LINE-1", amount=100)]),
        ("报销二", [_HEADER, _row(line_id="LINE-1", amount=900)]),
    ])

    detail = " ".join(
        error.error_detail for error in transformed.result.errors
        if error.error_type == "cross_sheet_conflict"
    )
    assert "报销一" in detail and "报销二" in detail
    assert "100" in detail and "900" in detail


def test_same_key_same_values_across_sheets_still_dedupes_silently(tmp_path):
    """回归锁：同键同值本来就是同一笔（导出重复），保留首次、不报错。"""

    transformed = _transform(tmp_path, [
        ("报销一", [_HEADER, _row(line_id="LINE-1", amount=100)]),
        ("报销二", [_HEADER, _row(line_id="LINE-1", amount=100)]),
    ])

    codes = {error.error_type for error in transformed.result.errors}
    assert "cross_sheet_conflict" not in codes
    assert len(transformed.result.lines) == 1


def test_different_keys_across_sheets_are_both_kept(tmp_path):
    """回归锁：不同键就是两笔费用，两张页各自的行都要在。"""

    transformed = _transform(tmp_path, [
        ("报销一", [_HEADER, _row(line_id="LINE-1", amount=100)]),
        ("报销二", [_HEADER, _row(line_id="LINE-2", amount=900)]),
    ])

    assert len(transformed.result.lines) == 2
    codes = {error.error_type for error in transformed.result.errors}
    assert "cross_sheet_conflict" not in codes
