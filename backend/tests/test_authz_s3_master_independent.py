"""Independent receipt topology and publication failure counterexamples."""
from datetime import datetime, timezone
from io import BytesIO
import json

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.authz import runtime
from app.config import get_settings
from app.models.authorization import SysAuthorizationArtifact, SysAuthorizationGrant
from app.models.maintenance_project import MaintenanceProject, MaintenanceProjectUserAssignment
from app.models.maintenance_source_assignment import MaintenanceSourceOrderAssignment
from app.models.maintenance_project_operations import MaintenanceProjectWorkbookOperation, MaintenanceProjectOperationAudit
from tests.test_authz_s3_master_workbooks import master, grant  # noqa:F401
from tests.test_authz_s3_checklist import setup  # noqa:F401
from tests.authz_pdp_helpers import cerbos_url  # noqa:F401


def test_receipt_inside_closure_reparent_changes_precondition_but_keeps_history_readable(db,master,monkeypatch):
    from app.services import maintenance_project_master_workbook as workbook
    from tests.test_maintenance_project_master_v2_editable import _make_project_with_line
    uid,pid,client,url=master
    settings=get_settings()
    monkeypatch.setattr(settings,'authz_enabled_resources',settings.authz_enabled_resources+['maintenance_general_workbook_write'])
    runtime.activate_configured_resources(db.get_bind())
    old,part,order,line=_make_project_with_line(db)
    old_id=old.project_id;source_id=order.raw_order_id;order.linked_sales_order_no=None
    db.add(MaintenanceProjectUserAssignment(assignment_id='independent-receipt-manager',project_id=pid,
        user_id=uid,responsibility_type='primary_manager',assigned_by='independent',assignment_reason='test'))
    fields=['data_purchase_cost','data_profit','data_customer']
    db.add(SysAuthorizationGrant(user_id=uid,resource='maintenance_managed_workbook',actions=['download','apply'],
        scope='primary_managed_projects',field_permissions=fields,version=1))
    db.add(SysAuthorizationGrant(user_id=uid,resource='maintenance_general_workbook_write',actions=['reassign_source_orders'],
        scope='all_projects',field_permissions=fields,version=1));db.commit()
    response=client.get(url,params={'sheets':workbook.V2_SHEET_PARTS})
    assert response.status_code==200,response.text
    book=load_workbook(BytesIO(response.content));sheet=book[workbook.V2_SHEET_PARTS]
    columns={cell.value:cell.column for cell in sheet[1]};row=sheet.max_row+1
    for name,value in {'维保单号':order.order_no,'PN':part.pn_std,'描述':line.description,
                      '需求数量':float(line.qty),'退货数量':float(line.return_qty or 0)}.items():sheet.cell(row,columns[name],value)
    output=BytesIO();book.save(output);book.close();content=output.getvalue()
    endpoint=url.removesuffix('.xlsx')+'/apply';files={'file':('test.xlsx',content)}
    applied=client.post(endpoint,files=files);assert applied.status_code==200,applied.text
    view=client.post(endpoint+'?view=authorization',files=files)
    assert view.status_code==200,view.text
    context=view.json()['authorization_context'];pre={k:context[k] for k in ('catalog_version','authorization_epoch','context')}
    versions={p:db.get(MaintenanceProject,p).version for p in (pid,old_id)}
    current=db.scalar(select(MaintenanceSourceOrderAssignment).where(
        MaintenanceSourceOrderAssignment.source_order_id==source_id,MaintenanceSourceOrderAssignment.is_active.is_(True)))
    current.is_active=False;current.archived_at=datetime.now(timezone.utc);current.archived_by='independent';current.version+=1
    db.flush()
    db.add(MaintenanceSourceOrderAssignment(assignment_id='independent-inside-generation',
        source_order_id=source_id,project_id=old_id,created_by='independent'));db.commit()
    assert {p:db.get(MaintenanceProject,p).version for p in (pid,old_id)}==versions
    stale=client.post(endpoint,files=files,data={'authorization_precondition':json.dumps(pre)})
    assert stale.status_code==409,stale.text
    fresh=client.post(endpoint+'?view=authorization',files=files)
    assert fresh.status_code==200,fresh.text
    context=fresh.json()['authorization_context'];new_pre={k:context[k] for k in ('catalog_version','authorization_epoch','context')}
    assert new_pre['context']['authorization_revision']!=pre['context']['authorization_revision']
    replay=client.post(endpoint,files=files,data={'authorization_precondition':json.dumps(new_pre)})
    assert replay.status_code==200,replay.text
    assert replay.json()['replayed'] is True
    assert replay.json()['import_batch_id']==applied.json()['import_batch_id']


def test_publication_failure_after_real_sidecar_and_bytes_rolls_back_business_and_receipt(db,master,monkeypatch):
    from app.services import authorization_artifacts as artifacts
    from tests.test_maintenance_project_note_workbook import _note_cell
    uid,pid,client,url=master;settings=get_settings()
    monkeypatch.setattr(settings,'authz_enabled_resources',settings.authz_enabled_resources+['maintenance_general_workbook_write','maintenance_project'])
    runtime.activate_configured_resources(db.get_bind())
    grant(db,uid,'maintenance_general_workbook',['data_purchase_cost','data_profit','data_customer'],'all_projects')
    db.add(SysAuthorizationGrant(user_id=uid,resource='maintenance_general_workbook_write',actions=['apply'],scope='all_projects',
        field_permissions=['data_purchase_cost','data_profit','data_customer'],version=1))
    db.add(SysAuthorizationGrant(user_id=uid,resource='maintenance_project',actions=['update'],scope='all_projects',field_permissions=[],version=1));db.commit()
    download=client.get(url);assert download.status_code==200,download.text
    book=load_workbook(BytesIO(download.content));_note_cell(book).value='must roll back after publication'
    output=BytesIO();book.save(output);book.close()
    before={model:db.scalar(select(func.count()).select_from(model)) for model in
        (MaintenanceProjectWorkbookOperation,MaintenanceProjectOperationAudit,SysAuthorizationArtifact)}
    original=artifacts.publish_artifact;observed=[]
    def fault(session,*args,**kwargs):
        result=original(session,*args,**kwargs)
        if kwargs['purpose']=='apply_receipt':
            observed.append(result)
            assert session.get(MaintenanceProject,pid).project_note=='must roll back after publication'
            assert session.scalar(select(func.count()).select_from(MaintenanceProjectWorkbookOperation))==before[MaintenanceProjectWorkbookOperation]+1
            raise HTTPException(503,'independent post-publication fault')
        return result
    monkeypatch.setattr(artifacts,'publish_artifact',fault)
    response=client.post(url.removesuffix('.xlsx')+'/apply',files={'file':('note.xlsx',output.getvalue())})
    assert response.status_code==503,response.text
    assert len(observed)==1 and observed[0]['complete'] is True
    db.expire_all()
    assert db.get(MaintenanceProject,pid).project_note is None
    for model,count in before.items():assert db.scalar(select(func.count()).select_from(model))==count
