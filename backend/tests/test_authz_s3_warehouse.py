"""Warehouse HTTP contracts using actual PostgreSQL and Cerbos decisions."""
import json
from datetime import datetime, timezone
from types import MappingProxyType

import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from sqlalchemy import func, select, text
from sqlalchemy.orm import Session

from app.api import maintenance_warehouse as api
from app.auth import _make_token
from app.authz import runtime
from app.authz.principal import VerifiedIdentity
from app.config import get_settings
from app.db import get_db
from app.models.authorization import SysAuthorizationArtifact, SysAuthorizationGrant
from app.models.maintenance_project import MaintenanceProject, MaintenanceProjectUserAssignment
from app.models.maintenance_source_assignment import MaintenanceSourceOrderAssignment
from app.models.maintenance_warehouse import (
    MaintenanceWarehouseImportBatch as Batch, MaintenanceWarehouseDocument as Document,
    MaintenanceWarehouseAmbiguity as Ambiguity, MaintenanceWarehouseAuditEvent as Audit,
)
from app.models.system import SysUser
from app.services import maintenance_warehouse_authorization as protected
from app.services import maintenance_warehouse as business
from tests.authz_pdp_helpers import cerbos_url  # noqa: F401
from tests.test_maintenance_warehouse_api import _approve_synthetic_contract, _seed, _content
from tests.boss_board_helpers import make_project


@pytest.fixture
def warehouse(db, monkeypatch, cerbos_url):
    _approve_synthetic_contract(monkeypatch)
    _seed(db)
    settings = get_settings()
    for key, value in {'authorization_mode': 'unified_modules',
        'authz_enabled_resources': ['maintenance_warehouse', 'maintenance_warehouse_import'],
        'authz_suspended_resources': [], 'authz_sales_analysis_enabled': False,
        'authz_sales_analysis_suspended': False, 'cerbos_url': cerbos_url,
        'cerbos_timeout_seconds': 2.0}.items(): monkeypatch.setattr(settings, key, value)
    runtime.activate_configured_resources(db.get_bind())
    actor = SysUser(username='warehouse-canonical', role='readonly', password_hash='synthetic')
    db.add(actor); db.flush()
    db.add_all([SysAuthorizationGrant(user_id=actor.id, resource=protected.IMPORT,
        actions=['preview', 'apply', 'reconcile'], scope='all', field_permissions=[], version=1),
        SysAuthorizationGrant(user_id=actor.id, resource=protected.WAREHOUSE,
        actions=['read_documents', 'read_ambiguities', 'resolve_ambiguity'],
        scope='all_projects', field_permissions=[], version=1)])
    db.commit()
    app = FastAPI(); app.include_router(api.router, prefix='/api')
    app.dependency_overrides[get_db] = lambda: db
    client = TestClient(app)
    login(client, actor)
    return client, actor.id, _content()


def login(client, actor):
    token, _ = _make_token(actor.role, actor.username, None, token_version=actor.token_version, authn='sys_user')
    client.headers['Authorization'] = 'Bearer ' + token


def post_file(client, path, content, *, fields=None, view=False):
    return client.post('/api/maintenance/' + path + ('?view=authorization' if view else ''),
        files={'file': ('warehouse.xlsx', content)}, data=fields or {})


def precondition(response):
    assert response.status_code == 200, response.text
    ctx = response.json()['authorization_context']
    return {k: ctx[k] for k in ('catalog_version', 'authorization_epoch', 'context')}


def applied(client, content):
    preview = post_file(client, 'warehouse-imports/preview', content)
    assert preview.status_code == 200, preview.text
    data = preview.json()
    fields = {'reason': 'Synthetic verified import', 'preview_token': data['preview_token']}
    path = 'warehouse-imports/' + data['import_id'] + '/apply'
    view = post_file(client, path, content, fields=fields, view=True)
    fields['authorization_precondition'] = json.dumps(precondition(view))
    result = post_file(client, path, content, fields=fields)
    assert result.status_code == 200, result.text
    return data, result, fields


def grant(db, uid, resource):
    return db.scalar(select(SysAuthorizationGrant).where(SysAuthorizationGrant.user_id == uid,
                                                        SysAuthorizationGrant.resource == resource))


def set_actions(db, uid, resource, actions):
    row = grant(db, uid, resource); row.actions = actions; row.version += 1; db.commit()


def test_preview_and_all_authorization_views_are_zero_write(db, warehouse):
    client, uid, content = warehouse
    before = db.scalar(select(func.count()).select_from(Audit))
    view = post_file(client, 'warehouse-imports/preview', content, view=True)
    assert set(view.json()) == {'locator', 'authorization_context'}
    assert view.json()['authorization_context']['context']['object_type'] == 'operation_input'
    fields = {'authorization_precondition': json.dumps(precondition(view))}
    response = post_file(client, 'warehouse-imports/preview', content, fields=fields)
    assert response.status_code == 200, response.text
    assert response.json()['can_apply'] is True
    assert db.scalar(select(func.count()).select_from(Batch)) == 0
    assert db.scalar(select(func.count()).select_from(SysAuthorizationArtifact)) == 0
    assert db.scalar(select(func.count()).select_from(Audit)) == before
    set_actions(db, uid, protected.IMPORT, ['apply'])
    assert post_file(client, 'warehouse-imports/preview', content).status_code == 403


def test_apply_stable_receipt_replay_and_revocation(db, warehouse):
    client, uid, content = warehouse
    plan, result, fields = applied(client, content)
    batch = db.get(Batch, plan['import_id'])
    assert batch.owner_user_id == uid
    binding = db.get(SysAuthorizationArtifact, (protected.KIND, batch.import_id))
    assert binding.owner_user_id == uid and binding.artifact_manifests['apply_receipt']['complete'] is True
    path = 'warehouse-imports/' + batch.import_id + '/apply'
    fields.pop('authorization_precondition')
    again = post_file(client, path, content, fields=fields)
    assert again.status_code == 200, again.text
    assert again.json()['idempotent_replay'] is True
    assert again.json()['writes'] == {'documents': 0, 'lines': 0, 'links': 0, 'ambiguities': 0, 'audits': 0}
    set_actions(db, uid, protected.IMPORT, ['preview'])
    assert post_file(client, path, content, fields=fields).status_code == 403
    assert db.scalar(select(func.count()).select_from(Batch)) == 1


def test_search_and_resolve_are_independent_and_use_current_context(db, warehouse):
    client, uid, content = warehouse
    applied(client, content)
    url = '/api/maintenance/warehouse-documents/search'
    view = client.post(url + '?view=authorization', json={})
    response = client.post(url, json={'authorization_precondition': precondition(view)})
    assert response.status_code == 200, response.text
    assert response.json()['total'] == 1
    set_actions(db, uid, protected.WAREHOUSE, ['read_ambiguities'])
    assert client.post(url, json={}).status_code == 403
    url = '/api/maintenance/warehouse-ambiguities/search'
    response = client.post(url, json={})
    assert response.status_code == 200, response.text
    item = next(row for row in response.json()['items'] if row['ambiguity_type'] == 'controlled_attachment')
    resolve = '/api/maintenance/warehouse-ambiguities/' + item['ambiguity_id'] + '/resolve'
    body = {'version': item['version'], 'reason': 'Actual controlled evidence checked', 'decision': 'acknowledge'}
    assert client.post(resolve, json=body).status_code == 403
    set_actions(db, uid, protected.WAREHOUSE, ['resolve_ambiguity'])
    view = client.post(resolve + '?view=authorization', json=body)
    before = db.get(Ambiguity, item['ambiguity_id']).version
    result = client.post(resolve, json={**body, 'authorization_precondition': precondition(view)})
    assert result.status_code == 200, result.text
    db.expire_all()
    assert db.get(Ambiguity, item['ambiguity_id']).version == before + 1
    assert client.post(url, json={}).status_code == 403


def test_reconcile_is_separate_and_empty_plan_has_no_implicit_permission(db, warehouse):
    client, uid, _ = warehouse
    url = '/api/maintenance/warehouse-integrations/reconcile'
    body = {'reason': 'Explicit empty target reconciliation'}
    view = client.post(url + '?view=authorization', json=body)
    result = client.post(url, json={**body, 'authorization_precondition': precondition(view)})
    assert result.status_code == 200, result.text
    assert result.json()['documents'] == 0
    set_actions(db, uid, protected.IMPORT, ['preview', 'apply'])
    assert client.post(url, json=body).status_code == 403


def source_project(db, uid, suffix='a'):
    project = make_project(db, 'warehouse-project-' + suffix)
    db.add(MaintenanceSourceOrderAssignment(assignment_id='warehouse-source-' + suffix,
        source_order_id='SYN-API-WBDD', project_id=project.project_id, created_by='synthetic'))
    db.add(MaintenanceProjectUserAssignment(assignment_id='warehouse-manager-' + suffix,
        project_id=project.project_id, user_id=uid, responsibility_type='primary_manager',
        assigned_by='synthetic', assignment_reason='explicit project responsibility'))
    db.commit()
    return project.project_id



def move_source(db, target_project_id):
    row = db.get(MaintenanceSourceOrderAssignment, 'warehouse-source-a')
    row.is_active = False; row.archived_at = datetime.now(timezone.utc)
    row.archived_by = 'synthetic'; row.version += 1
    db.flush()
    db.add(MaintenanceSourceOrderAssignment(assignment_id='warehouse-source-next',
        source_order_id=row.source_order_id, project_id=target_project_id, created_by='synthetic'))


def test_scope_uses_current_full_parent_set_before_count_and_paging(db, warehouse):
    client, uid, content = warehouse
    pid = source_project(db, uid)
    applied(client, content)
    row = grant(db, uid, protected.WAREHOUSE)
    row.scope = 'primary_managed_projects'; row.version += 1; db.commit()
    url = '/api/maintenance/warehouse-documents/search'
    result = client.post(url, json={'page_size': 1})
    assert result.status_code == 200, result.text
    assert result.json()['total'] == 1
    other = make_project(db, 'warehouse-new-parent')
    move_source(db, other.project_id); db.commit()
    # The old warehouse project link cannot lend its scope to the new actual parent.
    result = client.post(url, json={'page_size': 1})
    assert result.status_code == 200, result.text
    assert result.json()['total'] == 0 and result.json()['items'] == []
    ambiguities = client.post('/api/maintenance/warehouse-ambiguities/search', json={})
    assert ambiguities.status_code == 200 and ambiguities.json()['total'] == 0


def test_receipt_survives_rename_but_not_username_reuse_or_unknown_owner(db, warehouse):
    client, uid, content = warehouse
    plan, _, fields = applied(client, content)
    fields.pop('authorization_precondition')
    url = 'warehouse-imports/' + plan['import_id'] + '/apply'
    actor = db.get(SysUser, uid); old_name = actor.username
    actor.username = 'warehouse-renamed'; db.commit(); login(client, actor)
    assert post_file(client, url, content, fields=fields).status_code == 200
    replacement = SysUser(username=old_name, role='admin', password_hash='synthetic')
    db.add(replacement); db.flush()
    db.add(SysAuthorizationGrant(user_id=replacement.id, resource=protected.IMPORT,
        actions=['apply'], scope='all', field_permissions=[], version=1)); db.commit(); login(client, replacement)
    assert post_file(client, url, content, fields=fields).status_code == 403
    assert db.get(Batch, plan['import_id']).owner_user_id == uid


def test_legacy_receipt_is_not_claimed_by_same_username_and_same_upload(db, warehouse):
    client, uid, content = warehouse
    secret = get_settings().secret_key.encode()
    plan = business.preview_import(content, filename='old.xlsx', hmac_key=secret)
    business.apply_import(db, content, filename='old.xlsx', import_id=plan['import_id'],
        preview_token=plan['preview_token'], reason='legacy original', operated_by='warehouse-canonical', hmac_key=secret)
    db.commit()
    response = post_file(client, 'warehouse-imports/' + plan['import_id'] + '/apply', content,
        fields={'reason': 'cannot claim', 'preview_token': plan['preview_token']})
    assert response.status_code == 403
    assert db.get(Batch, plan['import_id']).owner_user_id is None
    assert db.scalar(select(func.count()).select_from(SysAuthorizationArtifact)) == 0


def test_upload_bytes_and_stale_precondition_cannot_be_reused(db, warehouse):
    client, uid, content = warehouse
    preview = post_file(client, 'warehouse-imports/preview', content).json()
    path = 'warehouse-imports/' + preview['import_id'] + '/apply'
    fields = {'preview_token': preview['preview_token'], 'reason': 'verified input'}
    view = post_file(client, path, content, fields=fields, view=True)
    pre = precondition(view); pre['context']['authorization_revision'] = 'stale'
    denied = post_file(client, path, content, fields={**fields, 'authorization_precondition': json.dumps(pre)})
    assert denied.status_code == 409
    from openpyxl import load_workbook
    from io import BytesIO
    workbook = load_workbook(BytesIO(content)); workbook.active.cell(3, 11, 99)
    output = BytesIO(); workbook.save(output)
    assert post_file(client, path, output.getvalue(), fields=fields).status_code == 409
    assert db.scalar(select(func.count()).select_from(Batch)) == 0


def test_search_materializes_facts_in_same_read_only_repeatable_read(db, warehouse, monkeypatch):
    client, _, content = warehouse; applied(client, content)
    original = business.search_documents; observations = []
    def observed(session, **kwargs):
        observations.append((session.scalar(text('SHOW transaction_isolation')),
                             session.scalar(text('SHOW transaction_read_only'))))
        return original(session, **kwargs)
    monkeypatch.setattr(business, 'search_documents', observed)
    response = client.post('/api/maintenance/warehouse-documents/search', json={})
    assert response.status_code == 200 and response.json()['total'] == 1
    assert observations == [('repeatable read', 'on')]


@pytest.mark.parametrize('change', ['revoke', 'move_parent'])
def test_real_project_lock_wait_rechecks_actor_and_actual_parent(db, warehouse, change):
    from concurrent.futures import ThreadPoolExecutor
    import time
    client, uid, content = warehouse
    pid = source_project(db, uid)
    other = make_project(db, 'warehouse-lock-other')
    owner = db.get(SysUser, uid)
    identity = VerifiedIdentity(owner.username, owner.token_version, False, 'sys_user', MappingProxyType({}), uid)
    secret = get_settings().secret_key.encode()
    plan = business.preview_import(content, filename='waiting.xlsx', hmac_key=secret)
    params = {'import_id': plan['import_id'], 'preview_token': plan['preview_token'], 'reason': 'Lock wait adversary'}
    bind = db.get_bind(); db.rollback()
    with Session(bind) as view_session:
        response = protected.command(view_session, identity, entry_id=protected.APPLY, params=params,
            content=content, filename='waiting.xlsx', hmac_key=secret, view='authorization')
    context = response['authorization_context']; pre = {k: context[k] for k in ('catalog_version','authorization_epoch','context')}
    def apply():
        with Session(bind) as worker:
            worker.execute(text("SET application_name='warehouse-authz-lock-test'"))
            try:
                protected.command(worker, identity, entry_id=protected.APPLY, params=params,
                    content=content, filename='waiting.xlsx', hmac_key=secret, precondition=pre)
                return 200
            except HTTPException as exc:
                worker.rollback(); return exc.status_code
    with Session(bind) as blocker:
        blocker.scalar(select(MaintenanceProject).where(MaintenanceProject.project_id == pid).with_for_update())
        with ThreadPoolExecutor(max_workers=1) as pool:
            future = pool.submit(apply)
            waited = False
            for _ in range(100):
                with Session(bind) as observer:
                    waited = bool(observer.scalar(text("SELECT EXISTS(SELECT 1 FROM pg_stat_activity WHERE application_name='warehouse-authz-lock-test' AND wait_event_type='Lock')")))
                if waited: break
                time.sleep(.02)
            assert waited, 'worker must actually wait for the held domain lock'
            if change == 'revoke':
                set_actions(db, uid, protected.IMPORT, [])
            else:
                move_source(blocker, other.project_id)
            blocker.commit()
            assert future.result(timeout=10) in ({403} if change == 'revoke' else {409})
    assert db.scalar(select(func.count()).select_from(Batch)) == 0
    assert db.scalar(select(func.count()).select_from(Document)) == 0
    assert db.scalar(select(func.count()).select_from(SysAuthorizationArtifact)) == 0


def test_stable_warehouse_owner_blocks_lossy_schema_downgrade(db, monkeypatch):
    from tests.test_authz_s3_migration import migrate, S2, insert
    _approve_synthetic_contract(monkeypatch); _seed(db)
    uid = insert(db, 'warehouse-owner-guard', eligible=False)
    content = _content(); secret = get_settings().secret_key.encode()
    preview = business.preview_import(content, filename='owner.xlsx', hmac_key=secret)
    business.apply_import(db, content, filename='owner.xlsx', import_id=preview['import_id'],
        preview_token=preview['preview_token'], reason='Stable owner cannot be removed',
        operated_by='warehouse-owner-guard', hmac_key=secret, owner_user_id=uid)
    with pytest.raises(RuntimeError, match='stable command or file owner'):
        migrate(db, S2, down=True)
    db.rollback()
    assert db.scalar(text("SELECT EXISTS(SELECT 1 FROM information_schema.columns WHERE table_name='maintenance_warehouse_import_batch' AND column_name='owner_user_id')"))


def test_nonempty_reconcile_binds_every_document_and_replay_cannot_ignore_parent_move(db, warehouse):
    from app.models.maintenance_warehouse import MaintenanceWarehouseDocumentLink as Link
    client, uid, content = warehouse
    first = source_project(db, uid)
    plan, _, fields = applied(client, content)
    fields.pop('authorization_precondition')
    url = '/api/maintenance/warehouse-integrations/reconcile'
    body = {'reason': 'Repair exact current assignment edge'}
    old_view = client.post(url + '?view=authorization', json=body)
    second = make_project(db, 'warehouse-reconcile-target')
    move_source(db, second.project_id); db.commit()
    stale = client.post(url, json={**body, 'authorization_precondition': precondition(old_view)})
    assert stale.status_code == 409
    assert db.scalar(select(Link.target_id).where(Link.target_type == 'maintenance_project', Link.status == 'active')) == first
    view = client.post(url + '?view=authorization', json=body)
    result = client.post(url, json={**body, 'authorization_precondition': precondition(view)})
    assert result.status_code == 200, result.text
    assert result.json()['documents'] == 1 and result.json()['links_superseded'] == 1
    assert db.scalar(select(Link.target_id).where(Link.target_type == 'maintenance_project', Link.status == 'active')) == second.project_id
    replay = post_file(client, 'warehouse-imports/' + plan['import_id'] + '/apply', content, fields=fields)
    assert replay.status_code in (403, 409), replay.text
    assert db.scalar(select(func.count()).select_from(Batch)) == 1


def test_scoped_resolve_cannot_borrow_original_document_parent(db, warehouse):
    client, uid, content = warehouse
    source_project(db, uid); applied(client, content)
    item = db.scalar(select(Ambiguity).where(Ambiguity.ambiguity_type == 'controlled_attachment'))
    row = grant(db, uid, protected.WAREHOUSE)
    row.scope = 'primary_managed_projects'; row.actions = ['resolve_ambiguity']; row.version += 1
    db.commit()
    url = '/api/maintenance/warehouse-ambiguities/' + item.ambiguity_id + '/resolve'
    body = {'version': item.version, 'decision': 'acknowledge', 'reason': 'Only the actual owner may resolve'}
    old_view = client.post(url + '?view=authorization', json=body)
    other = make_project(db, 'warehouse-resolve-other')
    move_source(db, other.project_id); db.commit()
    response = client.post(url, json={**body, 'authorization_precondition': precondition(old_view)})
    assert response.status_code == 403, response.text
    db.expire_all()
    assert db.get(Ambiguity, item.ambiguity_id).status == 'open'


def test_receipt_sidecar_content_must_match_immutable_business_result(db, warehouse):
    client, uid, content = warehouse
    plan, _, fields = applied(client, content)
    fields.pop('authorization_precondition')
    binding = db.get(SysAuthorizationArtifact, (protected.KIND, plan['import_id']))
    manifests = json.loads(json.dumps(binding.artifact_manifests))
    manifests['apply_receipt']['content_sha256'] = '0' * 64
    binding.artifact_manifests = manifests; db.commit()
    response = post_file(client, 'warehouse-imports/' + plan['import_id'] + '/apply', content, fields=fields)
    assert response.status_code == 403, response.text
    assert db.scalar(select(func.count()).select_from(Batch)) == 1


def test_upstream_document_parent_closure_cannot_lend_scope_and_cycles_terminate(db, warehouse):
    from uuid import uuid4
    from app.models.maintenance_warehouse import MaintenanceWarehouseDocumentLink as Link
    client, uid, content = warehouse
    source_project(db, uid)
    plan, _, _ = applied(client, content)
    first = db.scalar(select(Document))
    second_project = make_project(db, 'warehouse-upstream-parent')
    second = Document(document_id=str(uuid4()), document_type='shipment', source_document_id='UPSTREAM-ID',
        document_no='UPSTREAM-NO', normalized_status='confirmed', raw_fields_json={},
        raw_fingerprint='b' * 64, first_import_id=plan['import_id'])
    db.add(second); db.flush()
    def link(document_id, kind, target_type, target_id):
        return Link(link_id=str(uuid4()), document_id=document_id, line_id=None,
            link_kind=kind, target_type=target_type, target_id=target_id,
            stable_key_kind='source_document_id', stable_key_hash='b' * 64,
            source='automatic', status='active', version=1, reason='exact upstream lineage', operated_by='synthetic')
    db.add_all([link(first.document_id, 'warehouse_document', 'warehouse_document', second.document_id),
                link(second.document_id, 'warehouse_document', 'warehouse_document', first.document_id),
                link(second.document_id, 'project', 'maintenance_project', second_project.project_id)])
    row = grant(db, uid, protected.WAREHOUSE); row.scope = 'primary_managed_projects'; row.version += 1
    db.commit()
    url = '/api/maintenance/warehouse-documents/search'
    response = client.post(url, json={})
    assert response.status_code == 200, response.text
    assert response.json()['total'] == 0 and response.json()['items'] == []
    db.add(MaintenanceProjectUserAssignment(assignment_id='warehouse-upstream-manager', project_id=second_project.project_id,
        user_id=uid, responsibility_type='primary_manager', assigned_by='synthetic', assignment_reason='complete parent scope'))
    db.commit()
    response = client.post(url, json={})
    assert response.status_code == 200, response.text
    assert response.json()['total'] == 2
