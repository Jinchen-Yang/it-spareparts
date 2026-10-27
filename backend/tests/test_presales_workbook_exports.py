"""售前下载入口：只读汇总、列选择、身份隔离和访问审计。"""

from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models.dimensions import DimCustomer
from app.models.system import SysAccessLog
from app.services.presales import inquiries, offers, sales_workbook
from app.services.workbook.schema import WorkbookError
from app.services.workflow.registry import Forbidden
from tests.test_presales_api_visibility import _client
from tests.test_presales_r6_existing import r6_actors


def _summary(content):
    workbook = load_workbook(BytesIO(content))
    sheet = workbook["销售询价汇总"]
    headers = [cell.value for cell in sheet[1]]
    return workbook, [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]


def test_sales_summary_aggregates_bound_and_raw_pn_without_roundtrip_identity(db, r6_actors):
    users = r6_actors
    first = inquiries.create_draft(db, users.sales.id, lines=[
        {"part_id": users.part.id, "qty": "2.100", "description": "第一处描述"},
        {"pn_raw": "待建档原文", "qty": 3},
    ])
    second = inquiries.create_draft(db, users.sales.id, lines=[
        {"part_id": users.part.id, "qty": "4.200", "require_exact_pn": True},
        {"pn_raw": "待建档原文", "qty": 5},
    ])
    content = sales_workbook.export_sales_summary(db, [first.id, second.id, first.id], users.sales.id)
    workbook, rows = _summary(content)
    assert workbook.sheetnames == ["填写说明", "销售询价汇总"]
    assert workbook["销售询价汇总"].protection.sheet
    assert len(rows) == 2
    bound = next(row for row in rows if row["PN"] == users.part.pn_std)
    assert bound["合计数量"] == "6.300"
    assert bound["要求 PN 一致"] == "是"
    assert bound["来源单号"] == first.inquiry_no + "\n" + second.inquiry_no
    assert next(row for row in rows if row["PN"] == "待建档原文")["合计数量"] == "8.000"
    assert "_meta" not in workbook.sheetnames
    assert all("_row_token" not in row for row in rows)


def test_summary_column_selection_keeps_pn_and_rejects_other_owners_even_for_boss(db, r6_actors):
    users = r6_actors
    inquiry = inquiries.create_draft(db, users.sales.id, lines=[{"part_id": users.part.id}])
    content = sales_workbook.export_sales_summary(db, [inquiry.id], users.sales.id,
                                                 {"summary": ["qty"]})
    _, rows = _summary(content)
    assert set(rows[0]) == {"PN", "合计数量"}
    for viewer in (users.other, users.boss):
        with pytest.raises(Forbidden, match="本人"):
            sales_workbook.export_sales_summary(db, [inquiry.id], viewer.id)
    with pytest.raises(WorkbookError):
        sales_workbook.export_sales_summary(db, [inquiry.id], users.sales.id, {"summary": "qty"})


def test_summary_prices_use_pinned_adoption_and_hide_private_reporter(db, r6_actors):
    users = r6_actors
    source = inquiries.create_draft(db, users.buyer.id, kind="proactive_offer",
                                    lines=[{"part_id": users.part.id}])
    offer = offers.save_draft(db, source.id, inquiries.lines_for(db, source.id)[0].id,
        users.buyer.id, {"unit_price": "12.34", "is_tax_inclusive": True,
                         "reporter_audience": {"mode": "private"}},
        expected_version=source.version)
    offers.publish_proactive(db, source.id, users.buyer.id, expected_version=source.version)
    inquiry = inquiries.create_draft(db, users.sales.id, lines=[{
        "part_id": users.part.id, "ref_offer_id": offer.id}])
    inquiries.submit(db, inquiry.id, users.sales.id, inquiry.version)
    inquiries.close(db, inquiry.id, users.sales.id, inquiry.version, outcome="won", won_offer_id=offer.id)
    offers.update_offer(db, offer.id, users.buyer.id, {"unit_price": "23.45"},
                        expected_version=offer.version)
    rows = sales_workbook.summary_rows(db, [inquiry.id], users.sales.id)
    assert "buyer" not in rows[0]["handler_min_prices"]
    assert "其他采购" in rows[0]["handler_min_prices"]
    assert "12.34" in rows[0]["adopted_prices"]
    assert "23.45" not in rows[0]["adopted_prices"]


def test_export_endpoints_download_and_write_literal_access_actions(db, r6_actors):
    users = r6_actors
    customer = DimCustomer(name_raw="绝不出现在采购文件的客户")
    db.add(customer)
    db.flush()
    inquiry = inquiries.create_draft(db, users.sales.id, customer_id=customer.id,
        lines=[{"part_id": users.part.id}], assignee_ids=[users.buyer.id])
    inquiry_id = inquiry.id
    db.commit()
    seller = _client(users.sales)
    template = seller.get("/api/presales/workbooks/inquiry-template")
    assert template.status_code == 200, template.text
    assert template.headers["cache-control"] == "no-store"
    exported = seller.get(f"/api/presales/workbooks/inquiry/{inquiry_id}")
    assert exported.status_code == 200, exported.text
    summary = seller.post("/api/presales/workbooks/sales-summary", json={
        "inquiry_ids": [inquiry_id], "visible_columns": {"summary": ["qty"]}})
    assert summary.status_code == 200, summary.text
    assert set(_summary(summary.content)[1][0]) == {"PN", "合计数量"}
    assert seller.get("/api/presales/workbooks/inquiry-template", params={
        "visible_columns": "not-json"}).status_code == 422
    assert _client(users.other).post("/api/presales/workbooks/sales-summary", json={
        "inquiry_ids": [inquiry_id]}).status_code == 403
    inquiries.submit(db, inquiry_id, users.sales.id, inquiry.version)
    db.commit()
    quote = _client(users.buyer).post("/api/presales/workbooks/quote-sheet", json={
        "inquiry_ids": [inquiry_id]})
    assert quote.status_code == 200, quote.text
    workbook = load_workbook(BytesIO(quote.content))
    assert customer.name_raw not in str([list(sheet.values) for sheet in workbook])
    actions = set(db.scalars(select(SysAccessLog.action)))
    assert {"presales_template_export", "presales_inquiry_export",
            "presales_summary_export", "presales_quote_sheet_export"} <= actions


@pytest.mark.parametrize("inquiry_ids", [[], [0], [True], ["1"]])
def test_export_invalid_identifiers_are_rejected(db, r6_actors, inquiry_ids):
    db.commit()
    response = _client(r6_actors.sales).post("/api/presales/workbooks/sales-summary",
                                            json={"inquiry_ids": inquiry_ids})
    assert response.status_code == 422


def test_summary_schema_supports_column_picker_but_export_cannot_be_imported(db, r6_actors):
    users = r6_actors
    inquiry = inquiries.create_draft(db, users.sales.id,
                                    lines=[{"part_id": users.part.id}])
    inquiry_id = inquiry.id
    db.commit()
    seller = _client(users.sales)
    schema = seller.get("/api/workbooks/presales_sales_summary/schema")
    assert schema.status_code == 200, schema.text
    columns = schema.json()["sheets"][0]["columns"]
    assert all(column["edit"] == "reference" for column in columns)
    assert {column["key"] for column in columns} == {
        "pn", "description", "qty", "inquiry_nos", "require_exact_pn",
        "handler_min_prices", "adopted_prices",
    }
    assert _client(users.buyer).get("/api/workbooks/presales_sales_summary/schema").status_code == 403
    exported = seller.post("/api/presales/workbooks/sales-summary",
                           json={"inquiry_ids": [inquiry_id]})
    assert exported.status_code == 200, exported.text
    preview = seller.post("/api/workbooks/presales_sales_summary/preview", files={
        "file": ("summary.xlsx", exported.content,
                 "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    })
    assert preview.status_code == 422
    with pytest.raises(WorkbookError, match="只读"):
        sales_workbook.reject_import(db, {}, users.sales.id)


def test_summary_keeps_untrusted_text_as_text_and_omits_customer_data(db, r6_actors):
    users = r6_actors
    customer = DimCustomer(name_raw="汇总不能带出的客户名称")
    db.add(customer)
    db.flush()
    inquiry = inquiries.create_draft(db, users.sales.id, customer_id=customer.id, lines=[{
        "pn_raw": "=HYPERLINK(\"https://example.invalid\")", "description": "=1+1", "qty": 2,
    }])
    content = sales_workbook.export_sales_summary(db, [inquiry.id], users.sales.id)
    workbook, rows = _summary(content)
    assert rows[0]["描述"] == "=1+1"
    assert rows[0]["PN"] == '=HYPERLINK("https://example.invalid")'
    sheet = workbook["销售询价汇总"]
    assert sheet.cell(2, 1).data_type == "s"
    assert sheet.cell(2, 2).data_type == "s"
    assert customer.name_raw not in str([list(sheet.values) for sheet in workbook])
