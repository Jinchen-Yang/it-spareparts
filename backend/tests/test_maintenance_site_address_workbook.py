"""D-26 领用地址往返契约；使用合成数据，不包含历史客户原件。"""
import io
import uuid
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models.maintenance_project_operations import (
    MaintenanceProjectOperationAudit, MaintenanceSiteIssue,
)
from app.services import maintenance_project_master_workbook as master
from app.services import maintenance_project_operations as operations
from tests.test_maintenance_project_master_v2_editable import (
    _append_site_line, _make_project_with_line, _site_issue,
)


def _download(db, project, *, monkeypatch=None, legacy=False):
    if legacy:
        with monkeypatch.context() as patch:
            patch.setattr(master, "V2_SITE_HEADERS", [name for name in master.V2_SITE_HEADERS
                if name not in {"领用地址", "地址操作", master.V2_SITE_DEMAND_ADDRESS_COLUMN}])
            patch.setattr(master, "V2_SITE_ADDRESS_GUARD_FIELDS", master.V2_SITE_BASE_FIELDS)
            data = master.build_project_master_v2(db, project_id=project.project_id,
                                                  sheets=(master.V2_SHEET_SITE,))
    else:
        data = master.build_project_master_v2(db, project_id=project.project_id,
                                              sheets=(master.V2_SHEET_SITE,))
    book = load_workbook(io.BytesIO(data))
    return book, book[master.V2_SHEET_SITE]


def _headers(sheet):
    return {cell.value: cell.column for cell in sheet[1] if cell.value}


def _row(sheet, line):
    column = _headers(sheet)["实体ID"]
    return next(row for row in range(2, sheet.max_row + 1)
                if sheet.cell(row, column).value == line.issue_line_id)


def _write(sheet, line, **values):
    row = _row(sheet, line)
    for name, value in values.items():
        sheet.cell(row, _headers(sheet)[name]).value = value


def _validate(db, project, book, *, takeover=False):
    output = io.BytesIO()
    book.save(output)
    return master.validate_project_master_v2(db, project_id=project.project_id,
                                            data=output.getvalue(), force_takeover=takeover)


def _apply(db, project, book, *, takeover=False):
    plan = _validate(db, project, book, takeover=takeover)
    result = master.apply_project_master_v2(db, plan, operated_by="address-test",
                                           import_batch_id=str(uuid.uuid4()))
    return plan, result


def _seed(db, *, second=False):
    project, part, _, _ = _make_project_with_line(db)
    issue, line = _site_issue(db, project, part)
    issue.site_location = "原领用地址"
    db.commit()
    return project, part, issue, line, _append_site_line(db, issue, part) if second else None


@pytest.mark.parametrize("alias", ["领用地址", "地址", "现场位置"])
@pytest.mark.parametrize("position", [1, 6, 17])
def test_address_alias_at_any_position_updates_once_without_repricing(db, monkeypatch, alias, position):
    project, _, issue, line, _ = _seed(db)
    book, sheet = _download(db, project)
    old_col = _headers(sheet)["领用地址"]
    values = [sheet.cell(row, old_col).value for row in range(1, sheet.max_row + 1)]
    sheet.delete_cols(old_col)
    sheet.insert_cols(position)
    for row, value in enumerate(values, 1):
        sheet.cell(row, position, value)
    sheet.cell(1, position, alias)
    _write(sheet, line, **{alias: "新领用地址\n二楼"})
    from app.services import maintenance_consumption_cost
    monkeypatch.setattr(maintenance_consumption_cost, "resolve_lines",
                        lambda *args, **kwargs: pytest.fail("地址不应触发重新取价"))
    version = issue.version
    plan, _ = _apply(db, project, book)
    db.refresh(issue)
    assert issue.site_location == "新领用地址\n二楼"
    assert issue.version == version + 1
    assert len([item for item in plan.field_changes if item["field"] == "领用地址"]) == 1
    audits = db.scalars(select(MaintenanceProjectOperationAudit).where(
        MaintenanceProjectOperationAudit.entity_id == issue.issue_id,
        MaintenanceProjectOperationAudit.action == "UPDATE")).all()
    assert len(audits) == 1
    assert audits[0].before_json["site_location"] == "原领用地址"
    assert audits[0].after_json["site_location"] == "新领用地址\n二楼"


def test_signed_unchanged_rows_do_not_conflict_with_one_changed_address(db):
    project, _, issue, first, second = _seed(db, second=True)
    book, sheet = _download(db, project)
    _write(sheet, second, **{"领用地址": "单头新地址"})
    plan, _ = _apply(db, project, book)
    db.refresh(issue)
    assert issue.site_location == "单头新地址"
    assert not plan.warnings
    assert len(plan.site_flags) == 1
    book2, sheet2 = _download(db, project)
    assert sheet2.cell(_row(sheet2, first), _headers(sheet2)["领用地址"]).value == "单头新地址"
    assert sheet2.cell(_row(sheet2, second), _headers(sheet2)["领用地址"]).value == "单头新地址"


@pytest.mark.parametrize("mode", ["blank", "missing", "legacy"])
def test_blank_missing_and_old_signed_template_preserve_current_address(db, monkeypatch, mode):
    project, _, issue, line, _ = _seed(db)
    book, sheet = _download(db, project, monkeypatch=monkeypatch, legacy=mode == "legacy")
    if mode == "blank":
        _write(sheet, line, **{"领用地址": None})
    elif mode == "missing":
        sheet.delete_cols(_headers(sheet)["领用地址"])
    issue.site_location = "同事已改地址"
    operations.bump_workbook_revision(db, project_id=project.project_id)
    db.commit()
    _write(sheet, line, **{"备注": "只改备注"})
    plan, _ = _apply(db, project, book)
    db.refresh(issue)
    db.refresh(line)
    assert issue.site_location == "同事已改地址"
    assert line.remark == "只改备注"
    assert not plan.conflicts


@pytest.mark.parametrize("operation", ["清空", "保留", "设置"])
def test_explicit_address_operation(db, operation):
    project, _, issue, line, _ = _seed(db)
    book, sheet = _download(db, project)
    _write(sheet, line, **{"地址操作": operation, "领用地址": "新地址"})
    _apply(db, project, book)
    db.refresh(issue)
    assert issue.site_location == {"清空": None, "保留": "原领用地址", "设置": "新地址"}[operation]


def test_literal_clear_is_an_address_and_empty_operation_retains_it(db):
    project, _, issue, line, _ = _seed(db)
    book, sheet = _download(db, project)
    _write(sheet, line, **{"领用地址": "清空"})
    _apply(db, project, book)
    db.refresh(issue)
    assert issue.site_location == "清空"


@pytest.mark.parametrize("conflict", ["addresses", "operations", "invalid_operation"])
def test_conflicting_addresses_preserve_that_header_but_apply_other_facts_and_order(db, conflict):
    project, part, issue, first, second = _seed(db, second=True)
    other, other_line = _site_issue(db, project, part, issue_no="ISS-OTHER")
    other.site_location = "另一单旧地址"
    db.commit()
    book, sheet = _download(db, project)
    _write(sheet, first, **{"领用地址": "地址甲", "备注": "合法备注"})
    _write(sheet, second, **{"领用地址": "地址乙", "领用数量": 3})
    if conflict == "operations":
        _write(sheet, second, **{"领用地址": "地址甲", "地址操作": "清空"})
        _write(sheet, first, **{"地址操作": "设置"})
    elif conflict == "invalid_operation":
        _write(sheet, second, **{"领用地址": "地址甲", "地址操作": "删除"})
    _write(sheet, other_line, **{"领用地址": "另一单新地址"})
    plan, _ = _apply(db, project, book)
    for record in (issue, first, second, other):
        db.refresh(record)
    assert issue.site_location == "原领用地址"
    assert first.remark == "合法备注"
    assert second.quantity == Decimal("3")
    assert other.site_location == "另一单新地址"
    assert any(issue.issue_no in warning and str(_row(sheet, first)) in warning
               and str(_row(sheet, second)) in warning for warning in plan.warnings)


def test_address_only_edit_rebases_current_quantity_and_notes(db):
    project, _, issue, line, _ = _seed(db)
    book, sheet = _download(db, project)
    line.quantity = Decimal("7")
    line.remark = "同事备注"
    operations.bump_workbook_revision(db, project_id=project.project_id)
    db.commit()
    _write(sheet, line, **{"领用地址": "我的新地址"})
    plan, _ = _apply(db, project, book)
    db.refresh(line)
    assert line.quantity == Decimal("7")
    assert line.remark == "同事备注"
    assert not plan.conflicts


def test_address_concurrency_uses_existing_conflict_and_takeover_contract(db):
    project, _, issue, line, _ = _seed(db)
    book, sheet = _download(db, project)
    issue.site_location = "同事的新地址"
    issue.version += 1
    operations.bump_workbook_revision(db, project_id=project.project_id)
    db.commit()
    _write(sheet, line, **{"领用地址": "我填写的地址"})
    plan = _validate(db, project, book)
    assert plan.conflicts[0]["field"] == "领用地址"
    with pytest.raises(master.WorkbookError, match="他人更新"):
        master.apply_project_master_v2(db, plan, operated_by="address-test", import_batch_id=str(uuid.uuid4()))
    db.rollback()
    plan, _ = _apply(db, project, book, takeover=True)
    db.refresh(issue)
    assert issue.site_location == "我填写的地址"
    assert plan.overridden[0]["base"] == "原领用地址"


def test_address_change_between_validate_and_apply_is_guarded(db):
    project, _, issue, line, _ = _seed(db)
    book, sheet = _download(db, project)
    _write(sheet, line, **{"领用地址": "预检地址"})
    plan = _validate(db, project, book)
    issue.site_location = "预检后被改"
    db.commit()
    with pytest.raises(master.WorkbookError, match="校验与写入之间"):
        master.apply_project_master_v2(db, plan, operated_by="address-test", import_batch_id=str(uuid.uuid4()))
    db.rollback()
    db.refresh(issue)
    assert issue.site_location == "预检后被改"


def test_old_template_missing_row_hash_keeps_existing_deletion_contract(db, monkeypatch):
    project, _, issue, first, second = _seed(db, second=True)
    book, sheet = _download(db, project, monkeypatch=monkeypatch, legacy=True)
    issue.site_location = "后来补的地址"
    db.commit()
    sheet.delete_rows(_row(sheet, first))
    plan = _validate(db, project, book)
    assert not plan.conflicts
    assert [flag.issue_line_id for flag in plan.site_flags if flag.is_void] == [first.issue_line_id]


def test_new_multi_line_issue_uses_address_even_if_only_later_row_filled(db):
    project, part, _, _, _ = _seed(db)
    book, sheet = _download(db, project)
    for sn, address in (("NEW-SN-1", None), ("NEW-SN-2", "新单实际地址")):
        values = {"领用单号": "ISS-NEW", "领用日期": date(2026, 10, 12),
                  "PN": part.pn_std, "SN": sn, "领用数量": 1,
                  "领用地址": address}
        sheet.append([values.get(name) for name in [cell.value for cell in sheet[1]]])
    _apply(db, project, book)
    issue = db.scalar(select(MaintenanceSiteIssue).where(MaintenanceSiteIssue.issue_no == "ISS-NEW"))
    assert issue.site_location == "新单实际地址"


def test_address_length_rejected_without_truncation(db):
    project, _, _, line, _ = _seed(db)
    book, sheet = _download(db, project)
    _write(sheet, line, **{"领用地址": "址" * 257})
    with pytest.raises(master.WorkbookError, match="256"):
        _validate(db, project, book)


def test_unchanged_roundtrip_does_not_write_address(db):
    project, _, issue, _, _ = _seed(db)
    book, _ = _download(db, project)
    version = issue.version
    plan, _ = _apply(db, project, book)
    db.refresh(issue)
    assert plan.site_flags == ()
    assert issue.version == version


def _append_new(sheet, part, *, issue_no, address, sn="NEW-SN", operation=None):
    values = {"领用单号": issue_no, "领用日期": date(2026, 8, 10),
              "PN": part.pn_std, "SN": sn, "领用数量": 1,
              "领用地址": address, "地址操作": operation}
    sheet.append([values.get(cell.value) for cell in sheet[1]])


def test_explicit_set_can_restore_baseline_after_concurrent_address_change(db):
    project, _, issue, line, _ = _seed(db)
    book, sheet = _download(db, project)
    issue.site_location = "他人修改"
    db.commit()
    _write(sheet, line, **{"地址操作": "设置"})
    plan, _ = _apply(db, project, book, takeover=True)
    db.refresh(issue)
    assert issue.site_location == "原领用地址"
    assert plan.overridden[0]["old"] == "他人修改"


def test_new_line_under_existing_issue_updates_address_and_is_guarded(db):
    project, part, issue, _, _ = _seed(db)
    book, sheet = _download(db, project)
    _append_new(sheet, part, issue_no=issue.issue_no, address="新增明细所填地址")
    plan = _validate(db, project, book)
    assert any(guard.entity_type == "site_issue" for guard in plan.row_guards)
    master.apply_project_master_v2(db, plan, operated_by="address-test", import_batch_id=str(uuid.uuid4()))
    db.refresh(issue)
    assert issue.site_location == "新增明细所填地址"


def test_new_line_address_guard_rejects_concurrent_change(db):
    project, part, issue, _, _ = _seed(db)
    book, sheet = _download(db, project)
    _append_new(sheet, part, issue_no=issue.issue_no, address="新增明细地址")
    plan = _validate(db, project, book)
    issue.site_location = "预检后修改"
    db.commit()
    with pytest.raises(master.WorkbookError, match="校验与写入之间"):
        master.apply_project_master_v2(db, plan, operated_by="address-test", import_batch_id=str(uuid.uuid4()))


def test_new_line_uses_untouched_signed_header_baseline_for_concurrency(db):
    project, part, issue, _, _ = _seed(db)
    book, sheet = _download(db, project)
    issue.site_location = "同事修改的地址"
    db.commit()
    _append_new(sheet, part, issue_no=issue.issue_no, address="新增明细地址")
    plan = _validate(db, project, book)
    conflict = next(item for item in plan.conflicts if item["field"] == "领用地址")
    assert conflict["base"] == "原领用地址"
    assert conflict["old"] == "同事修改的地址"
    assert conflict["new"] == "新增明细地址"


def test_renamed_issue_takes_address_from_new_line_with_new_number(db):
    project, part, issue, first, _ = _seed(db)
    book, sheet = _download(db, project)
    _write(sheet, first, **{"领用单号": "ISS-RENAMED"})
    _append_new(sheet, part, issue_no="ISS-RENAMED", address="新增明细地址")
    version = issue.version
    plan, _ = _apply(db, project, book)
    db.refresh(issue)
    assert issue.issue_no == "ISS-RENAMED"
    assert issue.site_location == "新增明细地址"
    assert issue.version == version + 1
    assert not plan.conflicts


@pytest.mark.parametrize("void_address", ["作废行的矛盾地址", "址" * 257])
def test_voided_line_does_not_vote_on_active_header_address(db, void_address):
    project, _, issue, first, second = _seed(db, second=True)
    book, sheet = _download(db, project)
    first.is_active = False
    db.commit()
    _write(sheet, first, **{"领用地址": void_address})
    _write(sheet, second, **{"领用地址": "活跃行的新地址"})
    plan, _ = _apply(db, project, book)
    db.refresh(issue)
    db.refresh(first)
    assert not first.is_active
    assert issue.site_location == "活跃行的新地址"
    assert len(plan.voided_rows) == 1
    assert not any("地址操作冲突" in text for text in plan.warnings)


def test_voided_row_baseline_does_not_contribute_to_new_line_address(db):
    project, part, issue, first, _ = _seed(db)
    book, sheet = _download(db, project)
    first.is_active = False
    issue.site_location = "当前地址"
    db.commit()
    _write(sheet, first, **{"领用地址": "作废行地址"})
    _append_new(sheet, part, issue_no=issue.issue_no, address="新增行地址")
    plan, _ = _apply(db, project, book)
    db.refresh(issue)
    assert not plan.conflicts
    assert issue.site_location == "新增行地址"


def test_old_signature_can_accept_new_address_column(db, monkeypatch):
    project, _, issue, line, _ = _seed(db)
    book, sheet = _download(db, project, monkeypatch=monkeypatch, legacy=True)
    column = sheet.max_column + 1
    sheet.cell(1, column, "地址")
    sheet.cell(_row(sheet, line), column, "旧模板新填地址")
    _apply(db, project, book)
    db.refresh(issue)
    assert issue.site_location == "旧模板新填地址"


def test_same_row_alias_conflict_preserves_address_and_keeps_notes(db):
    project, _, issue, line, _ = _seed(db)
    book, sheet = _download(db, project)
    column = sheet.max_column + 1
    sheet.cell(1, column, "地址")
    sheet.cell(_row(sheet, line), column, "另一个地址")
    _write(sheet, line, **{"备注": "合法备注"})
    plan, _ = _apply(db, project, book)
    db.refresh(issue)
    db.refresh(line)
    assert issue.site_location == "原领用地址"
    assert line.remark == "合法备注"
    assert plan.warnings


def test_exported_address_is_literal_and_max_length_roundtrips(db):
    project, _, issue, _, _ = _seed(db)
    issue.site_location = "=" + "地址" * 127 + "址"
    db.commit()
    assert len(issue.site_location) == 256
    book, sheet = _download(db, project)
    column = _headers(sheet)["领用地址"]
    assert sheet.cell(2, column).data_type == "s"
    assert sheet.cell(2, column).value == issue.site_location
    plan = _validate(db, project, book)
    assert plan.site_flags == ()


def test_address_does_not_change_generated_issue_identity(db):
    project, part, _, _, _ = _seed(db)
    book, sheet = _download(db, project)
    _append_new(sheet, part, issue_no=None, address="地址甲")
    first_plan = _validate(db, project, book)
    sheet.cell(sheet.max_row, _headers(sheet)["领用地址"], "地址乙")
    second_plan = _validate(db, project, book)
    first = next(flag for flag in first_plan.site_flags if flag.is_create)
    second = next(flag for flag in second_plan.site_flags if flag.is_create)
    assert (first.issue_line_id, first.issue_no) == (second.issue_line_id, second.issue_no)


def test_auto_numbered_rows_with_same_wbdd_keep_independent_addresses(db):
    project, part, _, _, _ = _seed(db)
    historic, _ = _site_issue(db, project, part, issue_no="WBDD-20261012-0001")
    historic.site_location = "历史WBDD单头地址"
    db.commit()
    book, sheet = _download(db, project)
    _append_new(sheet, part, issue_no="WBDD-20261012-0001", address="甲地址", sn="NEW-A")
    _append_new(sheet, part, issue_no="WBDD-20261012-0001", address="乙地址", sn="NEW-B")
    plan, _ = _apply(db, project, book)
    created = [flag for flag in plan.site_flags if flag.is_create]
    assert len({flag.issue_no for flag in created}) == 2
    assert {row.site_location for row in db.scalars(select(MaintenanceSiteIssue).where(
        MaintenanceSiteIssue.issue_no.in_([flag.issue_no for flag in created])))} == {"甲地址", "乙地址"}
    db.refresh(historic)
    assert historic.site_location == "历史WBDD单头地址"


def test_conflicting_new_issue_addresses_leave_new_address_empty(db):
    project, part, _, _, _ = _seed(db)
    book, sheet = _download(db, project)
    _append_new(sheet, part, issue_no="ISS-NEW-CONFLICT", address="甲地址", sn="NEW-A")
    _append_new(sheet, part, issue_no="ISS-NEW-CONFLICT", address="乙地址", sn="NEW-B")
    plan, _ = _apply(db, project, book)
    issue = db.scalar(select(MaintenanceSiteIssue).where(MaintenanceSiteIssue.issue_no == "ISS-NEW-CONFLICT"))
    assert issue.site_location is None
    assert plan.warnings
