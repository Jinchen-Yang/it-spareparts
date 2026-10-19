"""06 单号批内唯一性、合法互换及事务失败的公开 HTTP / 服务契约。"""

import io
import uuid
from dataclasses import replace
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, insert, select
from sqlalchemy.orm import Session

from app.api import maintenance_project_master_workbook as workbook_api
from app.models.maintenance_project_operations import (
    MaintenanceProjectOperationAudit,
    MaintenanceProjectWorkbookOperation,
    MaintenanceSiteIssue,
    MaintenanceSiteIssueLine,
    MaintenanceSiteIssueReturnEvent,
)
from app.services import maintenance_project_master_workbook as master
from app.services import maintenance_project_operations as operations
from tests.test_maintenance_expense_collection_workbook import _XLSX, uploader
from tests.test_maintenance_project_master_v2_editable import (
    _make_project_with_line,
    _save,
    _site_issue,
)


@pytest.fixture(autouse=True)
def _enable_v2(monkeypatch):
    monkeypatch.setattr(
        workbook_api.get_settings(), "maintenance_project_master_v2_enabled", True,
    )


def _seed(db, numbers=("ISS-A", "ISS-B")):
    project, part, *_ = _make_project_with_line(db)
    issues = [_site_issue(db, project, part, issue_no=number) for number in numbers]
    book = load_workbook(io.BytesIO(master.build_project_master_v2(
        db, project_id=project.project_id, sheets=(master.V2_SHEET_SITE,),
    )))
    sheet = book[master.V2_SHEET_SITE]
    columns = {cell.value: cell.column for cell in sheet[1]}
    rows = {
        sheet.cell(row, columns["实体ID"]).value: row
        for row in range(2, sheet.max_row + 1)
        if sheet.cell(row, columns["实体ID"]).value
    }
    db.commit()
    return project, part, issues, book, sheet, columns, rows


def _post(client, project, book, action):
    return client.post(
        f"/api/maintenance/projects/stable/{project.project_id}/master-workbook/{action}",
        files={"file": ("synthetic-issue-numbers.xlsx", _save(book), _XLSX)},
    )


def _validate(db, project, book):
    return master.validate_project_master_v2(
        db, project_id=project.project_id, data=_save(book),
    )


def _apply(db, plan):
    return master.apply_project_master_v2(
        db, plan, operated_by="synthetic-number-editor",
        import_batch_id=str(uuid.uuid4()),
    )


def _append_new(sheet, part, *, number=None, serial="SYNTHETIC-NEW", address=None,
                issue_date=date(2026, 8, 10)):
    values = {
        "领用单号": number, "领用日期": issue_date,
        "PN": part.pn_std, "SN": serial, "领用数量": 1,
        "是否应返还": "是", "领用地址": address,
    }
    sheet.append([values.get(cell.value) for cell in sheet[1]])
    return sheet.max_row


def _snapshot(db, project_id):
    db.expire_all()
    return {
        "headers": list(db.execute(select(
            MaintenanceSiteIssue.issue_id, MaintenanceSiteIssue.issue_no,
            MaintenanceSiteIssue.issue_date, MaintenanceSiteIssue.site_location,
            MaintenanceSiteIssue.version,
        ).where(MaintenanceSiteIssue.project_id == project_id).order_by(
            MaintenanceSiteIssue.issue_id,
        ))),
        "lines": list(db.execute(select(
            MaintenanceSiteIssueLine.issue_line_id, MaintenanceSiteIssueLine.issue_id,
            MaintenanceSiteIssueLine.quantity, MaintenanceSiteIssueLine.version,
        ).join(MaintenanceSiteIssue).where(
            MaintenanceSiteIssue.project_id == project_id,
        ).order_by(MaintenanceSiteIssueLine.issue_line_id))),
        "revision": operations.get_or_create_workbook_state(
            db, project_id=project_id,
        ).revision,
        "counts": tuple(db.scalar(select(func.count()).select_from(model)) for model in (
            MaintenanceProjectOperationAudit, MaintenanceProjectWorkbookOperation,
            MaintenanceSiteIssueReturnEvent,
        )),
    }


def _assert_duplicate(detail, target, row_numbers):
    assert detail["code"] == "duplicate_issue_no"
    assert target in detail["message"]
    assert "06_领用返还" in detail["message"]
    assert all(str(row) in detail["message"] for row in row_numbers)
    assert sorted(row["row_no"] for row in detail["rows"]) == sorted(row_numbers)


def test_prod_two_signed_issues_renamed_to_unused_number_are_rejected(db):
    """ae32737: validate 200 后 apply 在真实 PG 唯一键 flush 抛 IntegrityError。"""
    project, _, issues, book, sheet, columns, rows = _seed(db)
    target = "WBDD-20261013-0005"
    row_numbers = sorted(rows[line.issue_line_id] for _, line in issues)
    for row in row_numbers:
        sheet.cell(row, columns["领用单号"], target)
    client = uploader(db, username="synthetic-duplicate-numbers")

    validated = _post(client, project, book, "validate")
    # Keeping the apply attempt here makes the old bug reproduce the actual
    # validate-200 -> IntegrityError path, rather than only a loose status diff.
    applied = _post(client, project, book, "apply")

    for response in (validated, applied):
        assert response.status_code == 422, response.text
        _assert_duplicate(response.json()["detail"], target, row_numbers)
    for issue, _ in issues:
        db.refresh(issue)
    assert [issue.issue_no for issue, _ in issues] == ["ISS-A", "ISS-B"]


@pytest.mark.parametrize("entrypoint", ["service", "http"])
def test_apply_rechecks_duplicate_final_numbers_in_replayed_plan(db, monkeypatch, entrypoint):
    project, _, issues, book, sheet, columns, rows = _seed(db)
    for (issue, line), number in zip(issues, ("ISS-FINAL-A", "ISS-FINAL-B")):
        sheet.cell(rows[line.issue_line_id], columns["领用单号"], number)
    plan = _validate(db, project, book)
    forged = replace(plan, site_flags=tuple(
        replace(flag, issue_no="ISS-FINAL-A") for flag in plan.site_flags
    ))
    before = _snapshot(db, project.project_id)
    db.commit()

    if entrypoint == "service":
        with pytest.raises(master.WorkbookError) as raised:
            _apply(db, forged)
        error = raised.value
        _assert_duplicate({"code": error.code, "message": error.message, **error.details},
                          "ISS-FINAL-A", list(rows.values()))
        db.rollback()
    else:
        client = uploader(db, username="synthetic-replayed-numbers")
        monkeypatch.setattr(master, "validate_project_master_v2", lambda *_a, **_kw: forged)
        response = _post(client, project, book, "apply")
        assert response.status_code == 422, response.text
        _assert_duplicate(response.json()["detail"], "ISS-FINAL-A", list(rows.values()))
    assert _snapshot(db, project.project_id) == before


@pytest.mark.parametrize("entrypoint", ["validate", "apply"])
def test_generated_new_header_cannot_share_a_renamed_header_number(db, entrypoint):
    project, part, [(issue, line)], book, sheet, columns, rows = _seed(db, ("ISS-A",))
    new_row = _append_new(sheet, part)
    plan = _validate(db, project, book)
    generated = next(flag for flag in plan.site_flags if flag.is_create)
    assert generated.issue_no.startswith("LY-AUTO-")
    existing_row = rows[line.issue_line_id]
    sheet.cell(existing_row, columns["领用单号"], generated.issue_no)
    if entrypoint == "validate":
        client = uploader(db, username="synthetic-auto-header-collision")
        response = _post(client, project, book, "validate")
        assert response.status_code == 422, response.text
        _assert_duplicate(response.json()["detail"], generated.issue_no,
                          [existing_row, new_row])
    else:
        # Valid plans collected independently can be replayed together; the
        # create without a bound issue_id is a separate header, never an append.
        sheet.delete_rows(new_row)
        renamed = _validate(db, project, book)
        forged = replace(renamed, site_flags=(*renamed.site_flags, generated))
        before = _snapshot(db, project.project_id)
        db.commit()
        with pytest.raises(master.WorkbookError) as raised:
            _apply(db, forged)
        _assert_duplicate({"code": raised.value.code, "message": raised.value.message,
                           **raised.value.details}, generated.issue_no, [existing_row, new_row])
        db.rollback()
        assert _snapshot(db, project.project_id) == before


@pytest.mark.parametrize("targets", [("ISS-B", "ISS-A"), ("ISS-B", "ISS-C")], ids=["swap", "chain"])
@pytest.mark.parametrize("change_quantity", [False, True], ids=["headers", "headers-and-facts"])
def test_unique_final_numbers_succeed_with_original_to_final_audits(db, targets, change_quantity):
    project, _, issues, book, sheet, columns, rows = _seed(db)
    before_revision = operations.get_or_create_workbook_state(db, project_id=project.project_id).revision
    versions = {issue.issue_id: issue.version for issue, _ in issues}
    line_versions = {line.issue_line_id: line.version for _, line in issues}
    original = {issue.issue_id: issue.issue_no for issue, _ in issues}
    for (_, line), target in zip(issues, targets):
        row = rows[line.issue_line_id]
        sheet.cell(row, columns["领用单号"], target)
        if change_quantity:
            sheet.cell(row, columns["领用数量"], 3)
    client = uploader(db, username="synthetic-unique-final-numbers")

    validated = _post(client, project, book, "validate")
    assert validated.status_code == 200, validated.text
    applied = _post(client, project, book, "apply")
    assert applied.status_code == 200, applied.text

    db.expire_all()
    assert [issue.issue_no for issue, _ in issues] == list(targets)
    assert set(db.scalars(select(MaintenanceSiteIssue.issue_no).where(
        MaintenanceSiteIssue.project_id == project.project_id,
    ))) == set(targets)
    for (issue, line), target in zip(issues, targets):
        assert issue.version == versions[issue.issue_id] + 1 + int(change_quantity)
        assert line.version == line_versions[line.issue_line_id] + int(change_quantity)
        assert line.quantity == Decimal("3" if change_quantity else "2")
        audits = list(db.scalars(select(MaintenanceProjectOperationAudit).where(
            MaintenanceProjectOperationAudit.entity_type == "site_issue",
            MaintenanceProjectOperationAudit.entity_id == issue.issue_id,
            MaintenanceProjectOperationAudit.action == "UPDATE",
        )))
        assert len(audits) == 1
        assert audits[0].before_json["issue_no"] == original[issue.issue_id]
        assert audits[0].after_json["issue_no"] == target
    assert operations.get_or_create_workbook_state(db, project_id=project.project_id).revision == before_revision + 1
    assert db.scalar(select(func.count()).select_from(MaintenanceSiteIssueReturnEvent)) == 2 * int(change_quantity)


@pytest.mark.parametrize("number", ["ISS-RENAMED", "WBDD-20261013-0005"])
def test_new_details_preserve_explicit_header_append_and_wbdd_auto_numbering(db, number):
    project, part, [(issue, line)], book, sheet, columns, rows = _seed(db, ("ISS-A",))
    sheet.cell(rows[line.issue_line_id], columns["领用单号"], number)
    _append_new(sheet, part, number=number, address="合成新增明细地址")
    plan = _validate(db, project, book)
    added = next(flag for flag in plan.site_flags if flag.is_create)
    _apply(db, plan)
    db.expire_all()
    new_line = db.get(MaintenanceSiteIssueLine, added.issue_line_id)
    assert issue.issue_no == number
    if number.startswith("WBDD-"):
        assert new_line.issue_id != issue.issue_id
        assert db.get(MaintenanceSiteIssue, new_line.issue_id).issue_no.startswith("LY-AUTO-")
        assert new_line.demand_order_no == number
    else:
        assert new_line.issue_id == issue.issue_id
        assert issue.site_location == "合成新增明细地址"
        assert new_line.line_no == 2
        assert db.scalar(select(func.count()).select_from(MaintenanceSiteIssue)) == 1


def test_number_held_by_unrenamed_issue_is_still_rejected(db):
    project, _, issues, book, sheet, columns, rows = _seed(db)
    sheet.cell(rows[issues[0][1].issue_line_id], columns["领用单号"], "ISS-B")
    client = uploader(db, username="synthetic-occupied-number")
    before = _snapshot(db, project.project_id)
    db.commit()
    for action in ("validate", "apply"):
        response = _post(client, project, book, action)
        assert response.status_code == 422, response.text
        assert response.json()["detail"]["code"] == "duplicate_issue_no"
        assert "ISS-B" in response.json()["detail"]["message"]
    assert _snapshot(db, project.project_id) == before


def test_apply_rechecks_number_claimed_after_validate(db):
    project, part, [(issue, line)], book, sheet, columns, rows = _seed(db, ("ISS-A",))
    sheet.cell(rows[line.issue_line_id], columns["领用单号"], "ISS-CLAIMED")
    plan = _validate(db, project, book)
    _site_issue(db, project, part, issue_no="ISS-CLAIMED")
    before = _snapshot(db, project.project_id)
    db.commit()
    with pytest.raises(master.WorkbookError) as raised:
        _apply(db, plan)
    assert raised.value.code == "duplicate_issue_no"
    assert "ISS-CLAIMED" in raised.value.message
    db.rollback()
    assert _snapshot(db, project.project_id) == before


@pytest.mark.parametrize("second_target", ["ISS-A", "ISS-C"], ids=["swap", "chain"])
def test_swap_and_chain_new_detail_attaches_to_final_owner_with_date_and_address(db, second_target):
    project, part, issues, book, sheet, columns, rows = _seed(db)
    for (_, line), number, issue_date in zip(
        issues, ("ISS-B", second_target), (date(2026, 8, 11), date(2026, 8, 12)),
    ):
        sheet.cell(rows[line.issue_line_id], columns["领用单号"], number)
        sheet.cell(rows[line.issue_line_id], columns["领用日期"], issue_date)
    _append_new(sheet, part, number="ISS-B", issue_date=date(2026, 8, 11), address="合成最终 B 地址")
    plan = _validate(db, project, book)
    added = next(flag for flag in plan.site_flags if flag.is_create)
    _apply(db, plan)
    db.expire_all()
    first, second = (issue for issue, _ in issues)
    added_line = db.get(MaintenanceSiteIssueLine, added.issue_line_id)
    assert (first.issue_no, first.issue_date, first.site_location) == (
        "ISS-B", date(2026, 8, 11), "合成最终 B 地址",
    )
    assert (second.issue_no, second.issue_date, second.site_location) == (
        second_target, date(2026, 8, 12), None,
    )
    assert (added_line.issue_id, added_line.line_no) == (first.issue_id, 2)


def test_new_header_reuses_vacated_number_without_old_address_or_line_number(db):
    project, part, [(issue, line)], book, sheet, columns, rows = _seed(db, ("ISS-A",))
    issue.site_location = "合成原领用地址"
    db.commit()
    book = load_workbook(io.BytesIO(master.build_project_master_v2(
        db, project_id=project.project_id, sheets=(master.V2_SHEET_SITE,),
    )))
    sheet = book[master.V2_SHEET_SITE]
    sheet.cell(rows[line.issue_line_id], columns["领用单号"], "ISS-B")
    sheet.cell(rows[line.issue_line_id], columns["领用日期"], date(2026, 8, 11))
    _append_new(sheet, part, number="ISS-A", issue_date=date(2026, 8, 12), address="合成全新 A 地址")
    plan = _validate(db, project, book)
    added = next(flag for flag in plan.site_flags if flag.is_create)
    _apply(db, plan)
    db.expire_all()
    new_line = db.get(MaintenanceSiteIssueLine, added.issue_line_id)
    new_header = db.get(MaintenanceSiteIssue, new_line.issue_id)
    assert (issue.issue_no, issue.issue_date, issue.site_location) == (
        "ISS-B", date(2026, 8, 11), "合成原领用地址",
    )
    assert new_header.issue_id != issue.issue_id
    assert (new_header.issue_no, new_header.issue_date, new_header.site_location) == (
        "ISS-A", date(2026, 8, 12), "合成全新 A 地址",
    )
    assert new_line.line_no == 1


@pytest.mark.parametrize("entrypoint", ["service", "http"])
@pytest.mark.parametrize("phase", ["rename", "create"])
@pytest.mark.parametrize("constraint", ["issue_number", "unknown"])
def test_real_flush_integrity_errors_are_4xx_and_rollback_every_write(
    db, monkeypatch, caplog, entrypoint, phase, constraint,
):
    project, part, [(issue, line)], book, sheet, columns, rows = _seed(db, ("ISS-A",))
    project_id, original_issue_id = project.project_id, issue.issue_id
    target = "ISS-FLUSH-CONFLICT"
    if phase == "rename":
        sheet.cell(rows[line.issue_line_id], columns["领用单号"], target)
    else:
        _append_new(sheet, part, number=target)
    plan = _validate(db, project, book)
    before = _snapshot(db, project_id)
    db.commit()
    client = uploader(db, username="synthetic-flush-conflict") if entrypoint == "http" else None
    # Alembic's fileConfig disables already-imported application loggers while
    # building this test database. Restore this logger for the warning contract.
    monkeypatch.setattr(master.logger, "disabled", False)
    caplog.set_level("WARNING", logger=master.__name__)
    original_flush = Session.flush
    injected = []

    def inject_constraint_failure(session, *args, **kwargs):
        candidates = [obj for obj in (*session.new, *session.dirty)
                      if isinstance(obj, MaintenanceSiteIssue) and obj.issue_no == target]
        if candidates and not injected:
            injected.append(session)
            # Insert through Core to reproduce the actual PG constraint error
            # on the next ORM flush, after the apply-side preflight completed.
            # The competing row belongs to this transaction and rolls back too.
            session.execute(insert(MaintenanceSiteIssue).values(
                issue_id=str(uuid.uuid4()) if constraint == "issue_number" else original_issue_id,
                project_id=project_id, issue_no=target,
                issue_date=date(2026, 8, 10), raw_status="已确认",
                normalized_status="confirmed", status_mapping_state="mapped",
                status_mapping_version="synthetic-test", source="legacy", version=1,
            ))
        return original_flush(session, *args, **kwargs)

    monkeypatch.setattr(Session, "flush", inject_constraint_failure)
    expected_code = "duplicate_issue_no" if constraint == "issue_number" else "apply_conflict"
    if entrypoint == "service":
        with pytest.raises(master.WorkbookError) as raised:
            _apply(db, plan)
        assert raised.value.code == expected_code
        assert db.is_active
        assert db.scalar(select(1)) == 1  # No caller rollback needed after DB failure.
        message = raised.value.message
    else:
        response = _post(client, project, book, "apply")
        assert response.status_code == 422, response.text
        assert response.json()["detail"]["code"] == expected_code
        message = response.json()["detail"]["message"]
    assert injected, "The test must exercise a real write/flush collision"
    if constraint == "issue_number":
        assert target in message
    assert any(record.levelname == "WARNING" and record.exc_info for record in caplog.records)
    assert _snapshot(db, project_id) == before
