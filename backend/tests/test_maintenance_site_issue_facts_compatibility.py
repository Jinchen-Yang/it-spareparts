"""Preserve the existing workbook/manual protocols around shared fact edits."""

from collections import Counter
from datetime import date
from decimal import Decimal
import io
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.models.maintenance_project_operations import (
    MaintenanceProjectOperationAudit,
    MaintenanceProjectWorkbookOperation,
    MaintenanceSiteIssue,
    MaintenanceSiteIssueCommand,
    MaintenanceSiteIssueLine,
    MaintenanceSiteIssueReturnEvent,
)
from app.services import maintenance_manual_site_issue as manual
from app.services import maintenance_project_master_workbook as master
from app.services import maintenance_project_operations as operations
from tests.test_maintenance_site_issue_facts import (
    _apply_site_workbook,
    _facts,
    _site_workbook,
)


def _another_issue(db, *, project, part):
    """A second independent two-line document inside the same workbook scope."""
    issue = MaintenanceSiteIssue(
        issue_id=str(uuid4()), project_id=project.project_id,
        issue_no=f"FACT-SECOND-{uuid4().hex[:8]}", issue_date=date(2026, 9, 20),
        receiver="接收人", issued_by="发出人", site_location="现场",
        raw_status="confirmed", status_mapping_state="mapped",
        normalized_status="confirmed", status_mapping_version="contract-v1",
        source="workbook", import_batch_id=str(uuid4()), version=1,
    )
    db.add(issue)
    db.flush()
    lines = [MaintenanceSiteIssueLine(
        issue_line_id=str(uuid4()), issue_id=issue.issue_id, line_no=line_no,
        part_id=part.id, pn=part.pn_std, quantity=Decimal("2.000"),
        algorithm_version="contract-v1", is_active=True, version=1,
    ) for line_no in (1, 2)]
    db.add_all(lines)
    db.commit()
    return issue, lines


def _revision(db, project):
    return operations.get_or_create_workbook_state(
        db, project_id=project.project_id,
    ).revision


def _actions(db):
    return Counter(db.scalars(select(MaintenanceProjectOperationAudit.action)))


def _side_effects(db, *, project, issue, lines):
    """Receipt replay must leave every persisted write family unchanged."""
    return {
        "audits": db.scalar(select(func.count()).select_from(
            MaintenanceProjectOperationAudit)),
        "events": db.scalar(select(func.count()).select_from(
            MaintenanceSiteIssueReturnEvent)),
        "manual_commands": db.scalar(select(func.count()).select_from(
            MaintenanceSiteIssueCommand)),
        "workbook_operations": db.scalar(select(func.count()).select_from(
            MaintenanceProjectWorkbookOperation)),
        "revision": _revision(db, project),
        "facts": operations.site_issue_dict(issue, lines),
    }


@pytest.mark.parametrize("entrypoint", ["v1", "v2"])
def test_workbook_multiple_documents_emit_one_final_event_each(db, entrypoint):
    project, first_issue, first_lines, part = _facts(db, source="workbook")
    second_issue, second_lines = _another_issue(db, project=project, part=part)
    issues = (first_issue, second_issue)
    all_lines = first_lines + second_lines
    if entrypoint == "v1":
        master.apply(
            db, master.MasterPlan(
                project_id=project.project_id,
                site_flags=tuple(master.SiteReturnFlag(
                    issue_line_id=line.issue_line_id,
                    quantity=Decimal(3 if line.line_no == 1 else 4),
                    no_return=line.line_no == 2,
                ) for line in all_lines),
            ), operated_by="tester", import_batch_id=str(uuid4()),
        )
    else:
        book, sheet, columns, rows = _site_workbook(db, project)
        for line in all_lines:
            row = rows[line.issue_line_id]
            sheet.cell(row, columns["领用数量"], 3 if line.line_no == 1 else 4)
            sheet.cell(row, columns["是否应返还"], "是" if line.line_no == 1 else "否")
        _apply_site_workbook(db, project, book)

    db.expire_all()
    assert all(issue.version == 2 for issue in issues)
    assert [(line.quantity, line.no_return, line.version) for line in all_lines] == [
        (Decimal("3"), False, 2), (Decimal("4"), True, 2),
        (Decimal("3"), False, 2), (Decimal("4"), True, 2),
    ]
    events = list(db.scalars(select(MaintenanceSiteIssueReturnEvent)))
    assert len(events) == 2
    assert Counter(event.issue_id for event in events) == {
        issue.issue_id: 1 for issue in issues
    }
    for event in events:
        assert event.issue_version == 2
        assert {change["after"]["issue_line_id"] for change in event.payload["changes"]} == {
            line.issue_line_id for line in all_lines if line.issue_id == event.issue_id
        }
    assert _actions(db) == {
        "return_obligation_corrected": 4,
        **({"apply_completed": 1} if entrypoint == "v2" else {}),
    }
    assert _revision(db, project) == 1


@pytest.mark.parametrize("entrypoint", ["v1", "v2"])
def test_header_edit_internal_v1_apply_and_v2_upload_preserve_legacy_protocol(db, entrypoint):
    """V1's internal plan seam is broader than its unsigned upload parser."""
    project, issue, lines, _part = _facts(db, source="workbook")
    new_number = f"FACT-RENAMED-{uuid4().hex[:8]}"
    new_date = date(2026, 9, 21)
    if entrypoint == "v1":
        master.apply(
            db, master.MasterPlan(
                project_id=project.project_id,
                site_flags=tuple(master.SiteReturnFlag(
                    issue_line_id=line.issue_line_id, no_return=None,
                    issue_no=new_number, issue_date=new_date,
                ) for line in lines),
            ), operated_by="tester", import_batch_id=str(uuid4()),
        )
    else:
        book, sheet, columns, rows = _site_workbook(db, project)
        # One changed repeated header propagates to the whole document once.
        row = rows[lines[0].issue_line_id]
        sheet.cell(row, columns["领用单号"], new_number)
        sheet.cell(row, columns["领用日期"], new_date)
        _apply_site_workbook(db, project, book)

    db.expire_all()
    assert (issue.issue_no, issue.issue_date) == (new_number, new_date)
    assert issue.version == (1 if entrypoint == "v1" else 2)
    assert [line.version for line in lines] == [1, 1]
    assert [line.quantity for line in lines] == [Decimal("2"), Decimal("2")]
    assert all(line.is_active and line.no_return is None for line in lines)
    assert db.scalar(select(func.count()).select_from(MaintenanceSiteIssueReturnEvent)) == 0
    assert _actions(db) == ({} if entrypoint == "v1" else {"UPDATE": 1, "apply_completed": 1})
    if entrypoint == "v2":
        audit = db.scalar(select(MaintenanceProjectOperationAudit).where(
            MaintenanceProjectOperationAudit.entity_type == "site_issue",
            MaintenanceProjectOperationAudit.entity_id == issue.issue_id,
            MaintenanceProjectOperationAudit.action == "UPDATE",
        ))
        assert (audit.entity_type, audit.entity_id, audit.operated_by) == (
            "site_issue", issue.issue_id, "tester",
        )
        assert audit.before_json["issue_date"] == "2026-09-20"
        assert audit.after_json["issue_date"] == "2026-09-21"
        assert audit.after_json["issue_no"] == new_number
    assert _revision(db, project) == 1


@pytest.mark.parametrize("field", ["现场领用单号", "领用日期"])
def test_v1_uploaded_header_edit_is_rejected_without_any_writes(db, field):
    project, issue, lines, _part = _facts(db, source="workbook")
    before = _side_effects(db, project=project, issue=issue, lines=lines)
    book = load_workbook(io.BytesIO(master.build_project_master(
        db, project_id=project.project_id, sheets=(master.SHEET_SITE,),
    )))
    sheet = book[master.SHEET_SITE]
    columns = {cell.value: cell.column for cell in sheet[1] if cell.value}
    sheet.cell(2, columns[field],
               "FACT-V1-FORBIDDEN" if field == "现场领用单号" else "2026-09-21")
    content = io.BytesIO()
    book.save(content)
    with pytest.raises(master.WorkbookError) as raised:
        master.validate(db, project_id=project.project_id, data=content.getvalue())
    assert raised.value.code == "legacy_site_identity_readonly"
    # Even an unrelated commit after a caught validation error must be harmless.
    db.commit()
    db.expire_all()
    assert _side_effects(db, project=project, issue=issue, lines=lines) == before
    assert (issue.version, *(line.version for line in lines)) == (1, 1, 1)
    assert _actions(db) == {}
    assert _revision(db, project) == 0


def _manual_two_line_edit(db, *, project, issue, lines, part, key):
    return manual.patch_manual_site_issue(
        db, project_id=project.project_id, issue_id=issue.issue_id, version=1,
        idempotency_key=key, operated_by="tester", reason="人工多行核对",
        lines=[{
            "issue_line_id": line.issue_line_id, "part_id": part.id,
            "quantity": "3" if line.line_no == 1 else "4",
            "no_return": line.line_no == 2,
            "remark": f"核对第 {line.line_no} 行",
        } for line in lines],
    )


def test_manual_multiple_lines_keep_one_document_event_and_one_correct_audit(db):
    project, issue, lines, part = _facts(db)
    result = _manual_two_line_edit(
        db, project=project, issue=issue, lines=lines, part=part,
        key=f"facts-manual-multi-{uuid4()}",
    )
    db.commit()
    assert result["version"] == issue.version == 3
    assert [line.version for line in lines] == [3, 3]
    assert [line.quantity for line in lines] == [Decimal("3"), Decimal("4")]
    assert [line.no_return for line in lines] == [False, True]
    events = list(db.scalars(select(MaintenanceSiteIssueReturnEvent)))
    assert len(events) == 1
    assert (events[0].issue_id, events[0].issue_version) == (issue.issue_id, 2)
    assert {change["after"]["issue_line_id"] for change in events[0].payload["changes"]} == {
        line.issue_line_id for line in lines
    }
    assert _actions(db) == {"return_obligation_corrected": 2, "correct": 1}
    assert db.scalar(select(func.count()).select_from(MaintenanceSiteIssueCommand)) == 1
    assert _revision(db, project) == 1


def test_manual_replay_returns_persisted_receipt_without_any_new_writes(db):
    project, issue, lines, part = _facts(db)
    key = f"facts-manual-replay-{uuid4()}"
    first = _manual_two_line_edit(
        db, project=project, issue=issue, lines=lines, part=part, key=key,
    )
    db.commit()
    receipt = db.scalar(select(MaintenanceSiteIssueCommand).where(
        MaintenanceSiteIssueCommand.idempotency_key == key,
    ))
    assert receipt.response_json == first
    before_replay = _side_effects(db, project=project, issue=issue, lines=lines)
    replay = _manual_two_line_edit(
        db, project=project, issue=issue, lines=lines, part=part, key=key,
    )
    db.commit()
    assert first["idempotent_replay"] is False
    assert replay == {**first, "idempotent_replay": True}
    db.expire_all()
    assert _side_effects(db, project=project, issue=issue, lines=lines) == before_replay
    assert _actions(db) == {"return_obligation_corrected": 2, "correct": 1}


def test_v2_same_plan_replay_returns_original_result_without_any_new_writes(db):
    """The apply seam reuses one validated plan; HTTP revalidation is separate."""
    project, issue, lines, _part = _facts(db, source="workbook")
    book, sheet, columns, rows = _site_workbook(db, project)
    for line in lines:
        row = rows[line.issue_line_id]
        sheet.cell(row, columns["领用数量"], 3 if line.line_no == 1 else 4)
        sheet.cell(row, columns["是否应返还"], "是" if line.line_no == 1 else "否")
    content = io.BytesIO()
    book.save(content)
    plan = master.validate_project_master_v2(
        db, project_id=project.project_id, data=content.getvalue(),
    )
    assert not plan.conflicts
    first = master.apply_project_master_v2(
        db, plan, operated_by="tester", import_batch_id=str(uuid4()),
    )
    before_replay = _side_effects(db, project=project, issue=issue, lines=lines)
    replay = master.apply_project_master_v2(
        db, plan, operated_by="tester", import_batch_id=str(uuid4()),
    )
    assert first["replayed"] is False
    assert replay == {**first, "replayed": True}
    db.expire_all()
    assert _side_effects(db, project=project, issue=issue, lines=lines) == before_replay
    assert before_replay["workbook_operations"] == 1
    assert before_replay["events"] == 1
    assert before_replay["revision"] == 1
    assert _actions(db) == {"return_obligation_corrected": 2, "apply_completed": 1}


def test_v2_revalidated_upload_replays_original_receipt_without_writes(db):
    """Actual repeated uploads must keep the entire first successful response."""
    project, issue, lines, _part = _facts(db, source="workbook")
    book, sheet, columns, rows = _site_workbook(db, project)
    for line in lines:
        row = rows[line.issue_line_id]
        sheet.cell(row, columns["领用数量"], 3 if line.line_no == 1 else 4)
        sheet.cell(row, columns["是否应返还"], "是" if line.line_no == 1 else "否")
    content = io.BytesIO()
    book.save(content)
    data = content.getvalue()
    first_plan = master.validate_project_master_v2(
        db, project_id=project.project_id, data=data,
    )
    assert not first_plan.conflicts
    first = master.apply_project_master_v2(
        db, first_plan, operated_by="tester", import_batch_id=str(uuid4()),
    )
    assert first["replayed"] is False
    assert first["changes"]
    before_replay = _side_effects(db, project=project, issue=issue, lines=lines)

    # Actual duplicate uploads run validation again against the changed facts.
    replay_plan = master.validate_project_master_v2(
        db, project_id=project.project_id, data=data,
    )
    assert not replay_plan.conflicts
    assert replay_plan.file_sha256 == first_plan.file_sha256
    replay = master.apply_project_master_v2(
        db, replay_plan, operated_by="tester", import_batch_id=str(uuid4()),
    )
    assert replay["replayed"] is True
    assert replay["import_batch_id"] == first["import_batch_id"]
    assert replay == {**first, "replayed": True}
    db.expire_all()
    assert _side_effects(db, project=project, issue=issue, lines=lines) == before_replay
    assert before_replay["events"] == 1
    assert before_replay["revision"] == 1
    assert _actions(db) == {"return_obligation_corrected": 2, "apply_completed": 1}
