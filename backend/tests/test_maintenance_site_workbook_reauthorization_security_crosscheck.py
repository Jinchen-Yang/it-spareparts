"""06 锁后授权反例：保持 token_version 不变，验证当前身份、权限快照和字段资格。

复用准备数据的夹具，使用独立等待编排与真实已签名 HTTP 请求。
"""

from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
from threading import Event
from time import monotonic

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, text

from app.db import SessionLocal, engine
from app.models.maintenance_project import MaintenanceProjectContract, MaintenanceProjectUserAssignment
from app.models.maintenance_project_operations import MaintenanceSiteIssueLine
from app.models.system import SysUser
from app.services import maintenance_project_master_workbook as master
from tests.test_maintenance_expense_collection_workbook import _PASSWORD, _XLSX
from tests.test_maintenance_project_master_v2_editable import _make_project_with_line
from tests.test_maintenance_site_workbook_reauthorization import _case, _hold, _login, _post, _snapshot


def _edit_account(case, mutate):
    with SessionLocal() as db:
        user = db.scalar(select(SysUser).where(SysUser.username == case.username))
        version = user.token_version
        mutate(user)
        db.commit()
        db.refresh(user)
        assert user.token_version == version, "counterexample must bypass JWT-version rejection"


def _put_right_in_template(case, key, value):
    def mutate(user):
        # 刻意保持旧 permissions 列及 JWT 不变。
        user.template_perms = {**user.template_perms, key: value}
        user.perm_overrides = {k: v for k, v in user.perm_overrides.items() if k != key}
    _edit_account(case, mutate)


def _book_with_overview(case, field, *, mixed=False):
    sheets = [master.V2_SHEET_SITE, master.V2_SHEET_OVERVIEW]
    if mixed:
        sheets.append(master.V2_SHEET_PLAN)
    exported = case.client.get(case.route + ".xlsx", params={"sheets": ",".join(sheets)})
    assert exported.status_code == 200, exported.text
    wb = load_workbook(BytesIO(exported.content))
    site = wb[master.V2_SHEET_SITE]
    cols = {cell.value: cell.column for cell in site[1] if cell.value}
    site.cell(2, cols["领用数量"], "3.25")
    label, value = {
        "amount": ("合同总额（含税）", "12345.67"),
        "note": ("项目备注", "独立复核：现场维护记录"),
    }[field]
    next(row[1] for row in wb[master.V2_SHEET_OVERVIEW] if row[0].value == label).value = value
    if mixed:
        row = next(row for row in wb[master.V2_SHEET_PLAN].iter_rows(min_row=2)
                   if row[12].value == case.milestone_id)
        row[5].value = "725.50"
    output = BytesIO()
    wb.save(output)
    wb.close()
    case.data = output.getvalue()
    preview = case.client.post(case.route + "/validate", files={"file": ("crosscheck.xlsx", case.data, _XLSX)})
    assert preview.status_code == 200, preview.text
    assert preview.json()["contract_updates" if field == "amount" else "project_note_updates"] == 1


def _reject_after_wait(case, change, *, lock_kind="global_advisory", status=403, message=None):
    before = _snapshot()
    with SessionLocal() as blocker, engine.connect().execution_options(isolation_level="AUTOCOMMIT") as observer:
        blocker_pid = blocker.scalar(text("SELECT pg_backend_pid()"))
        expected_query = _hold(blocker, case, lock_kind)
        with ThreadPoolExecutor(max_workers=1) as pool:
            future = pool.submit(_post, case)
            try:
                deadline = monotonic() + 15
                waiter_pid = None
                while monotonic() < deadline:
                    row = observer.execute(text("""
                        SELECT pid, query, wait_event_type FROM pg_stat_activity
                        WHERE datname=current_database()
                          AND :blocker = ANY(pg_blocking_pids(pid))
                    """), {"blocker": blocker_pid}).mappings().one_or_none()
                    # Activity sampling and lock-manager inspection can briefly disagree.
                    # Require both observations before revoking rights; keep the deadline.
                    if row and row["wait_event_type"] == "Lock":
                        assert row["wait_event_type"] == "Lock"
                        assert expected_query in row["query"], row
                        waiter_pid = row["pid"]
                        break
                    assert not future.done(), future.result().text if future.done() else ""
                    Event().wait(0.01)
                assert waiter_pid is not None, "request never reached the intended real PG lock"
                change()
                assert not future.done()
                assert blocker_pid in observer.scalar(text("SELECT pg_blocking_pids(:pid)"), {"pid": waiter_pid})
                assert _snapshot() == before
            finally:
                blocker.rollback()
            response = future.result(timeout=30)
    after = _snapshot()
    changed_tables = [name for name in before if before[name] != after[name]]
    assert response.status_code == status, f"{response.text}; committed_changed_tables={changed_tables}"
    if message:
        assert message in response.text
    assert after == before, "failed authorization must preserve all 16 business fact/audit tables"


@pytest.mark.parametrize("key,message", [
    ("page_maintenance", "无权访问该页面"),
    ("action_maintenance_expense_collection_upload", "项目总表编辑"),
    ("data_profit", "项目总表编辑"),
])
def test_template_snapshot_revocation_without_token_bump(db, monkeypatch, key, message):
    with _case(db, monkeypatch) as case:
        _put_right_in_template(case, key, True)
        _reject_after_wait(case, lambda: _put_right_in_template(case, key, False), message=message)


def test_account_deactivation_without_token_bump_during_line_wait(db, monkeypatch):
    with _case(db, monkeypatch) as case:
        _reject_after_wait(case, lambda: _edit_account(case, lambda user: setattr(user, "is_active", False)),
                           lock_kind="site_line", status=401, message="账号已失效")


def test_salesperson_identity_change_without_token_bump(db, monkeypatch):
    with _case(db, monkeypatch, editor="sales") as case:
        _reject_after_wait(case, lambda: _edit_account(case, lambda user: setattr(user, "salesperson_name", "Other salesperson")),
                           lock_kind="site_header", message="项目总表编辑")


def test_manager_demotion_keeps_viewer_scope_but_loses_workbook_edit(db, monkeypatch):
    with _case(db, monkeypatch, editor="manager") as case:
        def demote():
            with SessionLocal() as current:
                assignment = current.scalar(select(MaintenanceProjectUserAssignment).join(
                    SysUser, SysUser.id == MaintenanceProjectUserAssignment.user_id,
                ).where(SysUser.username == case.username,
                        MaintenanceProjectUserAssignment.project_id == case.project_id,
                        MaintenanceProjectUserAssignment.responsibility_type == "primary_manager"))
                # _case 另有 viewer 挂靠，仍可读项目，但失去负责人编辑资格。
                assignment.archived_at = datetime.now(timezone.utc)
                assignment.archived_by = "site-rights-admin"
                assignment.archive_reason = "Independent manager qualification revocation"
                assignment.version += 1
                current.commit()
        _reject_after_wait(case, demote, lock_kind="site_line", message="项目总表编辑")
        assert case.client.get(case.route + ".xlsx", params={"sheets": master.V2_SHEET_SITE}).status_code == 200


@pytest.mark.parametrize("field", ["amount", "note"])
@pytest.mark.parametrize("lock_kind", ["global_advisory", "late_milestone"])
def test_mixed_overview_field_right_revocation_without_token_bump(db, monkeypatch, field, lock_kind):
    with _case(db, monkeypatch, mixed=True) as case:
        _put_right_in_template(case, "action_maintenance_project_manage", True)
        _book_with_overview(case, field, mixed=True)
        _reject_after_wait(case, lambda: _put_right_in_template(case, "action_maintenance_project_manage", False),
                           lock_kind=lock_kind, message="合同总额" if field == "amount" else "无此操作权限")


@pytest.mark.parametrize("with_amount", [False, True])
def test_manager_without_profit_upload_or_page_site_action_still_allowed(db, monkeypatch, with_amount):
    with _case(db, monkeypatch, editor="manager") as case:
        _put_right_in_template(case, "data_profit", False)
        _login(case.client, case.username, _PASSWORD)
        if with_amount:
            _book_with_overview(case, "amount")
        applied = _post(case)
        assert applied.status_code == 200, applied.text
        with SessionLocal() as current:
            assert current.get(MaintenanceSiteIssueLine, case.site_line_id).quantity == Decimal("3.250")
            if with_amount:
                contract = current.scalar(select(MaintenanceProjectContract).where(
                    MaintenanceProjectContract.project_id == case.project_id))
                assert contract.amount_inc_tax == Decimal("12345.67")


def test_salesperson_without_profit_is_not_manager_exception(db, monkeypatch):
    with _case(db, monkeypatch, editor="sales") as case:
        _put_right_in_template(case, "data_profit", True)
        _reject_after_wait(case, lambda: _put_right_in_template(case, "data_profit", False),
                           lock_kind="site_line", message="项目总表编辑")


@pytest.mark.parametrize("lock_kind", ["global_advisory", "late_milestone"])
def test_mixed_assignment_requires_full_scope_after_role_downgrade(db, monkeypatch, lock_kind):
    """viewer/uploader 可编辑 06，但不具备跨项目改派 WBDD 的资格。"""
    with _case(db, monkeypatch, mixed=True) as case:
        _edit_account(case, lambda user: setattr(user, "role", "boss"))
        _login(case.client, case.username, _PASSWORD)
        old_project, part, order, line = _make_project_with_line(db)
        order.order_no = "WBDD-INDEPENDENT-FOREIGN"
        order.linked_sales_order_no = "XSDD-INDEPENDENT-FOREIGN"
        contract = db.scalar(select(MaintenanceProjectContract).where(
            MaintenanceProjectContract.project_id == old_project.project_id))
        contract.contract_no = order.linked_sales_order_no
        contract.contract_id = "INDEPENDENT-FOREIGN-CONTRACT"
        order.project_raw = "可编辑测试"
        order.project_std = "可编辑测试"
        db.commit()
        exported = case.client.get(case.route + ".xlsx", params={"sheets": ",".join([
            master.V2_SHEET_SITE, master.V2_SHEET_PARTS, master.V2_SHEET_PLAN])})
        assert exported.status_code == 200, exported.text
        wb = load_workbook(BytesIO(exported.content))
        site = wb[master.V2_SHEET_SITE]
        cols = {cell.value: cell.column for cell in site[1] if cell.value}
        site.cell(2, cols["领用数量"], "3.25")
        parts = wb[master.V2_SHEET_PARTS]
        cols = {cell.value: cell.column for cell in parts[1] if cell.value}
        row_no = parts.max_row + 1
        for key, value in {"维保单号": order.order_no, "PN": part.pn_std,
                           "描述": line.description, "需求数量": float(line.qty), "退货数量": 0}.items():
            parts.cell(row_no, cols[key], value)
        row = next(row for row in wb[master.V2_SHEET_PLAN].iter_rows(min_row=2)
                   if row[12].value == case.milestone_id)
        row[5].value = "725.50"
        output = BytesIO()
        wb.save(output)
        wb.close()
        case.data = output.getvalue()
        preview = case.client.post(case.route + "/validate", files={"file": ("scope.xlsx", case.data, _XLSX)})
        assert preview.status_code == 200, preview.text
        assert preview.json()["order_reassignments"] == 1
        assert preview.json()["will_reassign_orders"][0]["from_project_id"] == old_project.project_id
        db.rollback()
        _reject_after_wait(case, lambda: _edit_account(case, lambda user: setattr(user, "role", "readonly")),
                           lock_kind=lock_kind, status=422, message="仅全量项目范围账号")
