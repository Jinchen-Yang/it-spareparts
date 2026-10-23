"""Transaction ownership and independent receipt access after V2 extraction."""
import ast
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import subprocess
import sys
from typing import get_type_hints

import pytest
from openpyxl import load_workbook
from sqlalchemy import event, select
from sqlalchemy.exc import IntegrityError

from app.models.maintenance import FMaintenanceLine
from app.models.maintenance_project_operations import (
    MaintenanceProjectOperationAudit, MaintenanceProjectWorkbookOperation,
    MaintenanceProjectWorkbookState,
)
from app.services import maintenance_project_master_workbook as master
from app.services import maintenance_workbook_coordinator as coordinator
from app.services import maintenance_workbook_receipts as receipts
from app.services.maintenance_workbook_plans import WorkbookError
from tests.test_maintenance_project_master_v2_editable import _make_project_with_line
from tests.test_maintenance_workbook_boundaries import _local_dependency_paths

_EXPORTS = {
    coordinator: ('apply_project_master_v2', '_apply_project_master_v2'),
    receipts: ('_v2_apply_operation_key', '_v2_apply_result',
               '_v2_stored_apply_result', 'get_project_master_v2_replay'),
}

@pytest.mark.parametrize(('module', 'name'), [
    (module, name) for module, names in _EXPORTS.items() for name in names
])
def test_old_entrypoints_use_the_same_implementation_and_resolvable_annotations(module, name):
    function = getattr(module, name)
    assert getattr(master, name) is function
    assert function.__globals__ is vars(module)
    get_type_hints(function)


def test_transaction_and_receipt_closures_cannot_reenter_master_or_http():
    for module in _EXPORTS:
        paths = _local_dependency_paths([module.__name__])
        assert not {name: path for name, path in paths.items() if name.startswith((
            'app.api', 'app.main', 'app.services.maintenance_project_master_workbook',
        ))}
    # W2 still reaches the actual maintenance_manual_site_issue via facts.
    # R2 removes that legacy edge in the combined candidate.
    # W2 is stacked independently from R1. R1 removes cost -> board via the
    # contract query seam; this legacy edge is not a route back to master.
    receipt_paths = _local_dependency_paths([receipts.__name__])
    assert 'app.services.maintenance_workbook_coordinator' not in receipt_paths
    assert 'app.services.maintenance_expense_collection_workbook' not in receipt_paths
    assert 'app.services.maintenance_project_operations' not in receipt_paths
    tree = ast.parse(Path(receipts.__file__).read_text())
    assert not [n for n in ast.walk(tree) if isinstance(n, ast.Call)
                and isinstance(n.func, ast.Attribute)
                and n.func.attr in {'commit', 'rollback', 'flush', 'add', 'delete'}]


def test_coordinator_import_does_not_load_master_or_connect_to_database():
    script = '''
import json, sys
from sqlalchemy.engine import Engine
def fail(*args, **kwargs):
    raise AssertionError('import must not open a database connection')
Engine.connect = fail
from app.services import maintenance_workbook_coordinator
print(json.dumps(sorted(sys.modules)))
'''
    loaded = json.loads(subprocess.run([sys.executable, '-c', script],
        cwd=Path(__file__).resolve().parents[1], check=True,
        capture_output=True, text=True).stdout)
    assert 'app.services.maintenance_project_master_workbook' not in loaded
    assert not any(name.startswith('app.api.') for name in loaded)


def _plan(db):
    project, _part, _order, line = _make_project_with_line(db)
    data = master.build_project_master_v2(db, project_id=project.project_id,
                                         sheets=(master.V2_SHEET_PARTS,))
    book = load_workbook(BytesIO(data))
    sheet = book[master.V2_SHEET_PARTS]
    columns = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, columns['需求数量'], 7)
    stream = BytesIO()
    book.save(stream)
    db.commit()
    plan = master.validate_project_master_v2(db, project_id=project.project_id, data=stream.getvalue())
    return project.project_id, line.id, plan


def _snapshot(db, project_id, line_id):
    db.expire_all()
    line = db.get(FMaintenanceLine, line_id)
    state = db.get(MaintenanceProjectWorkbookState, project_id)
    return (line.qty, line.cost_amount_ex_tax, state.revision, state.data_version,
            [(row.id, deepcopy(row.after_json)) for row in db.scalars(
                select(MaintenanceProjectOperationAudit).order_by(MaintenanceProjectOperationAudit.id))],
            list(db.scalars(select(MaintenanceProjectWorkbookOperation.id))))


def test_completion_flush_failure_rolls_back_facts_revision_and_receipt_then_same_plan_retries(db):
    project_id, line_id, plan = _plan(db)
    before = _snapshot(db, project_id, line_id)
    commits = []
    def committed(_session):
        commits.append(True)
    def fail_completion(session, *_):
        if any(isinstance(row, MaintenanceProjectOperationAudit) and row.action == 'apply_completed'
               for row in session.new):
            raise IntegrityError('synthetic completion flush', {}, RuntimeError('synthetic constraint'))
    event.listen(db, 'after_commit', committed)
    event.listen(db, 'before_flush', fail_completion)
    try:
        with pytest.raises(WorkbookError) as caught:
            coordinator.apply_project_master_v2(db, plan, operated_by='coordinator-test', import_batch_id='first')
        assert caught.value.code == 'apply_conflict'
        assert db.is_active and not commits
        assert _snapshot(db, project_id, line_id) == before
    finally:
        event.remove(db, 'before_flush', fail_completion)
        event.remove(db, 'after_commit', committed)
    result = coordinator.apply_project_master_v2(db, plan, operated_by='coordinator-test', import_batch_id='retry')
    assert result['replayed'] is False
    after = _snapshot(db, project_id, line_id)
    assert after[0] == 7 and after[2] == before[2] + 1
    assert len(after[5]) == 1
    assert sum(1 for _, envelope in after[4] if isinstance(envelope, dict)
               and envelope.get('schema') == 'master_v2_apply_receipt_v1') == 1


def test_receipt_read_never_commits_unrelated_pending_changes_and_returns_detached_response(db):
    project_id, line_id, plan = _plan(db)
    original = coordinator.apply_project_master_v2(db, plan, operated_by='coordinator-test', import_batch_id='first')
    state = db.get(MaintenanceProjectWorkbookState, project_id)
    revision = state.revision
    state.revision += 100
    commits = []
    def committed(_session):
        commits.append(True)
    event.listen(db, 'after_commit', committed)
    try:
        replay, needs_amount = receipts.get_project_master_v2_replay(
            db, project_id=project_id, file_sha256=plan.file_sha256, operated_by='coordinator-test')
        assert replay == {**original, 'replayed': True} and needs_amount is False
        replay['summary'].clear()
        replay['changes'].append({'synthetic': 'caller mutation'})
        second, _ = receipts.get_project_master_v2_replay(
            db, project_id=project_id, file_sha256=plan.file_sha256, operated_by='coordinator-test')
        assert second == {**original, 'replayed': True}
        assert receipts.get_project_master_v2_replay(db, project_id=project_id,
            file_sha256=plan.file_sha256, operated_by='different-operator') is None
        assert not commits
    finally:
        event.remove(db, 'after_commit', committed)
        db.rollback()
    assert db.get(MaintenanceProjectWorkbookState, project_id).revision == revision
    assert db.get(FMaintenanceLine, line_id).qty == 7
