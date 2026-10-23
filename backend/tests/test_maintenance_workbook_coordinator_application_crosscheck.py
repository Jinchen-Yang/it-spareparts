"""Independent W2b application probes against real HTTP and PostgreSQL.

The fixtures combine 03 demand quantity, 04 canonical expense and 05 collections.
They exercise the actual extracted coordinator closure, not fake repository data.
"""
from copy import deepcopy
from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import event, select, text
from sqlalchemy.exc import IntegrityError

from app.api import maintenance_project_master_workbook as master_api
from app.db import SessionLocal
from app.models.maintenance import FMaintenanceLine, FProjectExpense
from app.models.maintenance_project import MaintenanceProject, MaintenanceProjectContract
from app.models.maintenance_project_operations import (
    MaintenanceCollectionSnapshot,
    MaintenanceProjectExpenseAttribution,
    MaintenanceProjectOperationAudit,
    MaintenanceProjectWorkbookOperation,
    MaintenanceProjectWorkbookState,
)
from app.models.system import SysImportBatch
from app.services import maintenance_project_master_workbook as master
from app.services import maintenance_workbook_coordinator as coordinator
from app.services import maintenance_project_operations as operations
from app.services.maintenance_workbook_plans import WorkbookError
from tests.test_maintenance_expense_collection_workbook import _BASE, _XLSX, uploader
from tests.test_maintenance_project_master_v2_editable import _make_project_with_expense


@pytest.fixture(autouse=True)
def _v2_and_bounded_waits(monkeypatch):
    monkeypatch.setattr(master_api.get_settings(), "maintenance_project_master_v2_enabled", True)

    def bound(_session, _transaction, connection):
        connection.execute(text("SET LOCAL lock_timeout = '5s'"))

    event.listen(SessionLocal, "after_begin", bound)
    try:
        yield
    finally:
        event.remove(SessionLocal, "after_begin", bound)


def _read_state(ids, session=None):
    if session is None:
        with SessionLocal() as observer:
            return _read_state(ids, observer)
    project_id, line_id, expense_id, collection_id = ids
    return {
        "line": tuple(session.execute(select(
            FMaintenanceLine.qty, FMaintenanceLine.cost_amount_ex_tax,
            FMaintenanceLine.cost_amount_inc_tax,
        ).where(FMaintenanceLine.id == line_id)).one()),
        "expense": tuple(session.execute(select(
            FProjectExpense.amount, FProjectExpense.amount_ex_tax,
            FProjectExpense.amount_inc_tax, FProjectExpense.data_status,
        ).where(FProjectExpense.raw_line_id == expense_id)).one()),
        "attribution": tuple(session.execute(select(
            MaintenanceProjectExpenseAttribution.amount_ex_tax,
            MaintenanceProjectExpenseAttribution.amount_inc_tax,
            MaintenanceProjectExpenseAttribution.version,
        ).where(MaintenanceProjectExpenseAttribution.expense_id == f"bxd:{expense_id}")).one()),
        "collection": tuple(session.execute(select(
            MaintenanceCollectionSnapshot.cumulative_amount,
            MaintenanceCollectionSnapshot.version,
            MaintenanceCollectionSnapshot.status,
        ).where(MaintenanceCollectionSnapshot.collection_id == collection_id)).one()),
        "project": session.scalar(select(MaintenanceProject.version).where(
            MaintenanceProject.project_id == project_id)),
        "state": tuple(session.execute(select(
            MaintenanceProjectWorkbookState.revision,
            MaintenanceProjectWorkbookState.data_version,
        ).where(MaintenanceProjectWorkbookState.project_id == project_id)).one()),
        "audits": [tuple(row) for row in session.execute(select(
            MaintenanceProjectOperationAudit.id,
            MaintenanceProjectOperationAudit.entity_type,
            MaintenanceProjectOperationAudit.action,
            MaintenanceProjectOperationAudit.after_json,
        ).where(MaintenanceProjectOperationAudit.project_id == project_id)
          .order_by(MaintenanceProjectOperationAudit.id))],
        "operations": list(session.scalars(select(MaintenanceProjectWorkbookOperation.id).where(
            MaintenanceProjectWorkbookOperation.project_id == project_id))),
        "batches": list(session.scalars(select(SysImportBatch.id).order_by(SysImportBatch.id))),
    }


def _mixed_upload(db):
    project, _order, line, expense = _make_project_with_expense(db, data_status="已生效")
    contract = db.scalar(select(MaintenanceProjectContract).where(
        MaintenanceProjectContract.project_id == project.project_id))
    collection = MaintenanceCollectionSnapshot(
        collection_id=str(uuid4()), project_id=project.project_id,
        project_contract_id=contract.project_contract_id,
        report_month=date(2026, 6, 1), cumulative_amount=Decimal("1200.00"),
        status="confirmed", source="direct_api", version=1,
    )
    db.add(collection)
    operations.get_or_create_workbook_state(db, project_id=project.project_id)
    db.commit()
    ids = (project.project_id, line.id, expense.raw_line_id, collection.collection_id)
    client = uploader(db, username="coordinator-application-crosscheck")
    response = client.get(f"{_BASE}/{project.project_id}/master-workbook.xlsx", params={
        "sheets": ",".join((master.V2_SHEET_PARTS, master.V2_SHEET_EXPENSE, master.V2_SHEET_RECEIPTS)),
    })
    assert response.status_code == 200, response.text
    book = load_workbook(BytesIO(response.content))
    for sheet_name, entity_id, field, value in (
        (master.V2_SHEET_PARTS, str(line.id), "需求数量", 3.125),
        (master.V2_SHEET_EXPENSE, expense.raw_line_id, "未税金额", 610.25),
        (master.V2_SHEET_RECEIPTS, collection.collection_id, "累计实收金额（含税）", 1500.50),
    ):
        sheet = book[sheet_name]
        columns = {cell.value: cell.column for cell in sheet[1]}
        row = next(row for row in range(2, sheet.max_row + 1)
                   if str(sheet.cell(row, columns["实体ID"]).value) == entity_id)
        sheet.cell(row, columns[field], value)
    output = BytesIO()
    book.save(output)
    db.commit()
    return client, ids, output.getvalue()


def _post(client, ids, data):
    return client.post(f"{_BASE}/{ids[0]}/master-workbook/apply",
                       files={"file": ("independent-mixed.xlsx", data, _XLSX)})


def _assert_success(before, after):
    assert after["line"] == (Decimal("3.125"), Decimal("312.50"), Decimal("353.13"))
    assert after["expense"][:3] == (Decimal("610.25"), Decimal("610.25"), Decimal("689.58"))
    assert after["attribution"] == (Decimal("610.25"), Decimal("689.58"), before["attribution"][2] + 1)
    assert after["collection"] == (Decimal("1500.50"), before["collection"][1] + 1, "confirmed")
    assert after["state"][0] == before["state"][0] + 1
    assert len(after["operations"]) == 1
    assert sum(row[1:3] == ("master_workbook_operation", "apply_completed")
               for row in after["audits"]) == 1


@pytest.mark.parametrize("entry", ["legacy", "coordinator"])
def test_both_entrypoints_keep_all_three_sheets_in_one_transaction(db, monkeypatch, entry):
    client, ids, data = _mixed_upload(db)
    before = _read_state(ids)
    plan = master.validate_project_master_v2(db, project_id=ids[0], data=data)
    db.commit()
    calls = []
    commits = []
    real_apply = coordinator.ec.apply

    def observed_nested(session, inner_plan, **kwargs):
        result = real_apply(session, inner_plan, **kwargs)
        session.flush()
        inside = _read_state(ids, session)
        assert inside["line"][0] == Decimal("3.125")
        assert inside["expense"][1] == Decimal("610.25")
        assert inside["collection"][0] == Decimal("1500.50")
        assert _read_state(ids) == before  # Separate connection sees no partial writes.
        assert session is db
        assert kwargs["commit"] is False and kwargs["bump_revision"] is False
        assert kwargs["track_change"] is True
        calls.append(session)
        return result

    def committed(_session):
        commits.append(True)

    monkeypatch.setattr(coordinator.ec, "apply", observed_nested)
    event.listen(db, "after_commit", committed)
    try:
        apply = master.apply_project_master_v2 if entry == "legacy" else coordinator.apply_project_master_v2
        result = apply(db, plan, operated_by="coordinator-application-crosscheck", import_batch_id="independent")
    finally:
        event.remove(db, "after_commit", committed)
        client.close()
    assert len(calls) == len(commits) == 1
    assert result["replayed"] is False
    _assert_success(before, _read_state(ids))


@pytest.mark.parametrize("failure", ["domain", "integrity", "duplicate_issue"])
def test_http_error_after_real_nested_writes_rolls_back_then_exact_bytes_retry(db, monkeypatch, failure):
    client, ids, data = _mixed_upload(db)
    before = _read_state(ids)
    real_apply = coordinator.ec.apply
    reached = []

    def fail_after_nested(session, plan, **kwargs):
        real_apply(session, plan, **kwargs)
        session.flush()
        inside = _read_state(ids, session)
        assert inside["line"][0] == Decimal("3.125")
        assert inside["attribution"][0] == Decimal("610.25")
        assert inside["collection"][0] == Decimal("1500.50")
        assert _read_state(ids) == before
        reached.append(True)
        if failure == "domain":
            raise WorkbookError("cost_recompute_busy", "独立测试：稍后重试")
        error = RuntimeError("synthetic internal SQL/credential must not reach HTTP")
        error.diag = SimpleNamespace(
            constraint_name="uq_maintenance_site_issue_project_no" if failure == "duplicate_issue" else "synthetic_constraint",
            message_detail="synthetic private database detail",
        )
        raise IntegrityError("SYNTHETIC PRIVATE SQL", {"private": "synthetic-private"}, error)

    with monkeypatch.context() as faults:
        faults.setattr(coordinator.ec, "apply", fail_after_nested)
        failed = _post(client, ids, data)
    expected = {"domain": (409, "cost_recompute_busy"),
                "integrity": (422, "apply_conflict"),
                "duplicate_issue": (422, "duplicate_issue_no")}[failure]
    assert (failed.status_code, failed.json()["detail"]["code"]) == expected, failed.text
    assert "SYNTHETIC PRIVATE" not in failed.text and "synthetic-private" not in failed.text
    if failure == "domain":
        assert failed.headers["Retry-After"] == "5"
    assert reached == [True]
    assert _read_state(ids) == before
    retried = _post(client, ids, data)
    assert retried.status_code == 200, retried.text
    assert retried.json()["replayed"] is False
    _assert_success(before, _read_state(ids))
    client.close()


def test_http_locked_guard_rejects_peer_change_after_validate_without_writing_other_sheets(db, monkeypatch):
    client, ids, data = _mixed_upload(db)
    before = _read_state(ids)
    real_validate = master.validate_project_master_v2
    real_guard = coordinator._v2_recheck_row_guards
    validated_sessions = []
    checked_sessions = []

    def validate_then_peer_commit(session, **kwargs):
        plan = real_validate(session, **kwargs)
        with SessionLocal() as peer:
            peer.get(FMaintenanceLine, ids[1]).qty = Decimal("9.125")
            peer.commit()
        validated_sessions.append(session)
        return plan

    def guard_after_locks(session, plan):
        assert session.scalar(text(
            "SELECT count(*) FROM pg_locks WHERE pid = pg_backend_pid() "
            "AND locktype = 'advisory' AND granted")) >= 1
        assert session.get(FMaintenanceLine, ids[1]).qty == Decimal("9.125")
        checked_sessions.append(session)
        return real_guard(session, plan)

    with monkeypatch.context() as probes:
        probes.setattr(master, "validate_project_master_v2", validate_then_peer_commit)
        probes.setattr(coordinator, "_v2_recheck_row_guards", guard_after_locks)
        response = _post(client, ids, data)
    assert response.status_code == 409, response.text
    assert response.json()["detail"]["code"] == "row_conflicts"
    assert response.json()["detail"]["conflicts"]
    assert len(validated_sessions) == len(checked_sessions) == 1
    assert checked_sessions[0] is validated_sessions[0]
    after = _read_state(ids)
    expected = deepcopy(before)
    expected["line"] = (Decimal("9.125"), *before["line"][1:])
    assert after == expected
    client.close()
