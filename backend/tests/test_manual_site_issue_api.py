"""页面人工登记领用（page_manual）——真实 service/API 全链路测试。

覆盖：prod 下可用、幂等回放/改 payload 拒绝、version CAS、跨项目拒绝、
缺价留空、数量/PN 变化重算、返还义务修正审计、作废退出统计、
06 导出可见、merged/excluded 主档拒绝、数量边界（0.0001/1e100/NaN）。
"""

from datetime import date
from decimal import Decimal
from uuid import uuid4

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app import auth
from app.api import maintenance_project_operations
from app.auth import hash_password
from app.models.dimensions import DimPart
from app.models.maintenance import FMaintenanceLine, FMaintenanceOrder
from app.models.maintenance_project import MaintenanceProject
from app.models.maintenance_project_operations import (
    MaintenanceProjectOperationAudit,
    MaintenanceSiteIssue,
    MaintenanceSiteIssueCommand,
    MaintenanceSiteIssueLine,
)
from app.models.maintenance_source_assignment import (
    MaintenanceSourceOrderAssignment,
)
from app.models.system import SysImportBatch, SysUser
from app.services import maintenance_manual_site_issue as manual_service


def _client(db, *, username: str = "manual_site_admin") -> TestClient:
    db.add(
        SysUser(
            username=username,
            role="admin",
            display_name="人工领用管理员",
            password_hash=hash_password("manual-password-123"),
            permissions={
                "page_maintenance": True,
                "action_maintenance_site_issue_manage": True,
                "data_purchase_cost": True,
            },
        )
    )
    db.commit()
    app = FastAPI()
    app.include_router(auth.router, prefix="/api")
    app.include_router(
        maintenance_project_operations.stable_site_issue_router, prefix="/api"
    )
    client = TestClient(app)
    login = client.post(
        "/api/auth/login",
        json={"username": username, "password": "manual-password-123"},
    )
    assert login.status_code == 200, login.text
    client.headers["Authorization"] = f"Bearer {login.json()['token']}"
    return client


def _project(db, project_id: str) -> MaintenanceProject:
    row = MaintenanceProject(
        project_id=project_id,
        project_code=f"MAN-{project_id[-6:]}",
        display_name="人工领用项目",
        lifecycle_status="ongoing",
    )
    db.add(row)
    db.commit()
    return row


def _part(db, pn: str, **kwargs) -> DimPart:
    row = DimPart(pn_std=pn, **kwargs)
    db.add(row)
    db.flush()
    return row


def _line_payload(part: DimPart, quantity: str = "2", **extra) -> dict:
    return {"part_id": part.id, "quantity": quantity, **extra}


def _create_body(parts, *, key: str, quantity: str = "2") -> dict:
    return {
        "idempotency_key": key,
        "issue_date": "2026-09-20",
        "receiver": "人工接收人",
        "issued_by": "人工发出人",
        "site_location": "人工现场",
        "lines": [_line_payload(part, quantity) for part in parts],
        "reason": "页面人工登记领用",
    }


def test_manual_create_preview_read_update_void_full_cycle(db, monkeypatch):
    """ENVIRONMENT=prod 下真实 service/API：新建→读回→修改→作废 + 成本/义务证据。"""
    project = _project(db, "project-manual-full-cycle")
    part = _part(db, "PN-MANUAL-001")
    other = _project(db, "project-manual-other")
    client = _client(db, username="manual_cycle_admin")

    # 生产闸门打开：人工通道必须仍可用（这是明确人工业务事实，不是合成发货）
    monkeypatch.setattr(
        maintenance_project_operations.operations,
        "_site_issue_is_production_blocked",
        lambda: True,
    )

    # 预览：不落库、无价格时成本留空
    preview = client.post(
        f"/api/maintenance/site-issues/manual/preview?project_id={project.project_id}",
        json={
            "issue_date": "2026-09-20",
            "receiver": "人工接收人",
            "issued_by": "人工发出人",
            "site_location": "人工现场",
            "lines": [_line_payload(part, "3")],
        },
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()["lines"][0]["cost_source"] is None
    assert preview.json()["lines"][0]["cost_amount_ex_tax"] is None
    assert preview.json()["inventory_effect"] == "none"
    assert db.query(MaintenanceSiteIssue).count() == 0  # 预览零写入

    # 新建
    key = f"manual-create-{uuid4().hex[:8]}-0001"
    created = client.post(
        f"/api/maintenance/site-issues/projects/{project.project_id}/manual",
        json=_create_body([part], key=key, quantity="3"),
    )
    assert created.status_code == 201, created.text
    issue = created.json()
    assert issue["source"] == "page_manual"
    assert issue["workflow_status"] == "confirmed"
    assert issue["issue_no"].startswith("LYR-")
    assert issue["lines"][0]["pn"] == "PN-MANUAL-001"
    assert issue["lines"][0]["quantity"] == "3.000"
    assert issue["idempotent_replay"] is False
    issue_id = issue["issue_id"]

    # 幂等回放：同 key 同 payload → 同结果 + idempotent_replay，不重复建单
    replayed = client.post(
        f"/api/maintenance/site-issues/projects/{project.project_id}/manual",
        json=_create_body([part], key=key, quantity="3"),
    )
    assert replayed.status_code == 201, replayed.text
    assert replayed.json()["issue_id"] == issue_id
    assert replayed.json()["idempotent_replay"] is True
    assert db.query(MaintenanceSiteIssue).count() == 1

    # 改 payload 复用 key → 409
    mutated = client.post(
        f"/api/maintenance/site-issues/projects/{project.project_id}/manual",
        json=_create_body([part], key=key, quantity="5"),
    )
    assert mutated.status_code == 409, mutated.text

    # 跨项目创建：admin 全局 scope 下 URL 项目即归属（建到 other 名下），
    # 受限角色（sales 无挂靠）会被项目 scope 拒绝——单独用例覆盖。

    # 读回：search 白名单包含 page_manual
    found = client.post(
        "/api/maintenance/site-issues/search",
        json={"project_id": project.project_id},
    )
    assert found.status_code == 200, found.text
    rows = found.json()["rows"]
    assert any(row["issue_id"] == issue_id for row in rows)

    # 修改：version CAS + 幂等
    db.expire_all()
    current = db.get(MaintenanceSiteIssue, issue_id)
    patch_key = f"manual-patch-{uuid4().hex[:8]}-0001"
    patched = client.patch(
        f"/api/maintenance/site-issues/manual/{issue_id}",
        json={
            "project_id": project.project_id,
            "version": current.version,
            "idempotency_key": patch_key,
            "issue_date": "2026-09-21",
            "lines": [{
                "issue_line_id": issue["lines"][0]["issue_line_id"],
                "part_id": part.id,
                "quantity": "4",
                "no_return": False,
                "remark": "更正",
            }],
            "reason": "数量更正",
        },
    )
    assert patched.status_code == 200, patched.text
    assert patched.json()["workflow_status"] == "corrected"
    assert patched.json()["lines"][0]["quantity"] == "4.000"
    assert patched.json()["lines"][0]["no_return"] is False

    # 旧 version 再改 → 409 CAS 冲突
    stale = client.patch(
        f"/api/maintenance/site-issues/manual/{issue_id}",
        json={
            "project_id": project.project_id,
            "version": current.version,
            "idempotency_key": f"manual-stale-{uuid4().hex[:8]}-01",
            "receiver": "新人",
            "reason": "过期版本",
        },
    )
    assert stale.status_code == 409, stale.text

    # patch 幂等回放
    replay_patch = client.patch(
        f"/api/maintenance/site-issues/manual/{issue_id}",
        json={
            "project_id": project.project_id,
            "version": current.version,
            "idempotency_key": patch_key,
            "issue_date": "2026-09-21",
            "lines": [{
                "issue_line_id": issue["lines"][0]["issue_line_id"],
                "part_id": part.id,
                "quantity": "4",
                "no_return": False,
                "remark": "更正",
            }],
            "reason": "数量更正",
        },
    )
    assert replay_patch.status_code == 200, replay_patch.text
    assert replay_patch.json()["idempotent_replay"] is True

    # 跨项目 patch：单据属 project，body 报 other → 服务层 403
    cross_patch = client.patch(
        f"/api/maintenance/site-issues/manual/{issue_id}",
        json={
            "project_id": other.project_id,
            "version": 999,
            "idempotency_key": f"manual-cp-{uuid4().hex[:8]}-001",
            "receiver": "越权",
            "reason": "跨项目",
        },
    )
    assert cross_patch.status_code == 403, cross_patch.text
    assert "不属于当前项目" in cross_patch.json()["detail"]

    # 作废（复用 stable void 端点）
    db.expire_all()
    fresh = db.get(MaintenanceSiteIssue, issue_id)
    voided = client.post(
        f"/api/maintenance/site-issues/{issue_id}/void",
        json={
            "project_id": project.project_id,
            "version": fresh.version,
            "idempotency_key": f"manual-void-{uuid4().hex[:8]}-01",
            "reason": "录错，整单作废",
        },
    )
    assert voided.status_code == 200, voided.text
    assert voided.json()["workflow_status"] == "void"

    # 作废后退出统计、不硬删
    db.expire_all()
    final = db.get(MaintenanceSiteIssue, issue_id)
    assert final.normalized_status == "void"
    lines = (
        db.query(MaintenanceSiteIssueLine)
        .filter_by(issue_id=issue_id)
        .all()
    )
    assert lines and all(not line.is_active for line in lines)
    audits = (
        db.query(MaintenanceProjectOperationAudit)
        .filter_by(entity_id=issue_id)
        .all()
    )
    assert any(a.action == "void" for a in audits)
    assert any(a.action == "create" for a in audits)
    assert any(a.action == "correct" for a in audits)


def test_manual_create_missing_price_keeps_cost_empty_never_fabricates(db):
    project = _project(db, "project-manual-no-price")
    part = _part(db, "PN-MANUAL-NO-PRICE")
    client = _client(db, username="manual_noprice_admin")
    created = client.post(
        f"/api/maintenance/site-issues/projects/{project.project_id}/manual",
        json=_create_body([part], key=f"manual-np-{uuid4().hex[:8]}-001"),
    )
    assert created.status_code == 201, created.text
    line = created.json()["lines"][0]
    assert line["cost_source"] is None
    assert line["unit_cost_ex_tax"] is None
    assert line["cost_amount_ex_tax"] is None
    assert line["cost_evidence_kind"] == "missing"


def test_manual_create_resolves_demand_price_and_recalculates_on_change(db):
    """需求单价格层生效；数量变化后金额按新数量重算（同需求单价）。"""
    project = _project(db, "project-manual-demand-price")
    part = _part(db, "PN-MANUAL-DEMAND")
    batch = SysImportBatch(
        filename="manual-demand.xlsx",
        file_type="maintenance",
        file_hash="manual-demand",
        status="success",
    )
    db.add(batch)
    db.flush()
    order = FMaintenanceOrder(
        raw_order_id="WBDD-raw-manual-001",
        order_no="WBDD-20260901-0001",
        order_date=date(2026, 9, 1),
        data_status="已生效",
        import_batch_id=batch.id,
    )
    db.add(order)
    db.flush()
    # 需求行成本 100 ex-tax / 净数量 2 → 单价 50
    db.add(
        FMaintenanceLine(
            raw_line_id="wbdd-line-manual-001",
            order_id=order.id,
            part_id=part.id,
            qty=Decimal("2"),
            unit_cost_ex_tax=Decimal("50"),
            cost_amount_ex_tax=Decimal("100.00"),
            cost_source="direct",
            import_batch_id=batch.id,
        )
    )
    db.add(
        MaintenanceSourceOrderAssignment(
            assignment_id=str(uuid4()),
            source_order_id=order.raw_order_id,
            project_id=project.project_id,
            is_active=True,
            version=1,
            created_by="tester",
        )
    )
    db.commit()
    client = _client(db, username="manual_demand_admin")

    created = client.post(
        f"/api/maintenance/site-issues/projects/{project.project_id}/manual",
        json=_create_body(
            [part],
            key=f"manual-dp-{uuid4().hex[:8]}-001",
            quantity="2",
        ),
    )
    assert created.status_code == 201, created.text
    line = created.json()["lines"][0]
    assert line["cost_source"] == "maint_demand"
    assert line["unit_cost_ex_tax"] == "50.00"
    assert line["cost_amount_ex_tax"] == "100.00"

    # 数量改 3 → 重算 150（同单价 × 新数量）
    patched = client.patch(
        f"/api/maintenance/site-issues/manual/{created.json()['issue_id']}",
        json={
            "project_id": project.project_id,
            "version": created.json()["version"],
            "idempotency_key": f"manual-dp2-{uuid4().hex[:8]}-01",
            "lines": [{
                "issue_line_id": created.json()["lines"][0]["issue_line_id"],
                "part_id": part.id,
                "quantity": "3",
            }],
            "reason": "数量更正",
        },
    )
    assert patched.status_code == 200, patched.text
    updated = patched.json()["lines"][0]
    assert updated["unit_cost_ex_tax"] == "50.00"
    assert updated["cost_amount_ex_tax"] == "150.00"


def test_manual_demand_order_must_belong_to_same_project(db):
    project = _project(db, "project-manual-demand-scope")
    stranger = _project(db, "project-manual-demand-stranger")
    part = _part(db, "PN-MANUAL-SCOPE")
    batch = SysImportBatch(
        filename="manual-scope.xlsx",
        file_type="maintenance",
        file_hash="manual-scope",
        status="success",
    )
    db.add(batch)
    db.flush()
    order = FMaintenanceOrder(
        raw_order_id="WBDD-raw-manual-scope",
        order_no="WBDD-20260901-0002",
        order_date=date(2026, 9, 1),
        data_status="已生效",
        import_batch_id=batch.id,
    )
    db.add(order)
    db.flush()
    db.add(
        FMaintenanceLine(
            raw_line_id="wbdd-line-manual-scope",
            order_id=order.id,
            part_id=part.id,
            qty=Decimal("1"),
            import_batch_id=batch.id,
        )
    )
    db.add(
        MaintenanceSourceOrderAssignment(
            assignment_id=str(uuid4()),
            source_order_id=order.raw_order_id,
            project_id=stranger.project_id,
            is_active=True,
            version=1,
            created_by="tester",
        )
    )
    db.commit()
    client = _client(db, username="manual_scope_admin")
    rejected = client.post(
        f"/api/maintenance/site-issues/projects/{project.project_id}/manual",
        json={
            "idempotency_key": f"manual-sc-{uuid4().hex[:8]}-001",
            "issue_date": "2026-09-20",
            "receiver": "接收",
            "issued_by": "发出",
            "site_location": "现场",
            "lines": [
                _line_payload(part, "1", demand_order_no="WBDD-20260901-0002")
            ],
            "reason": "跨项目需求单",
        },
    )
    assert rejected.status_code == 400, rejected.text
    assert "不属于当前项目" in rejected.json()["detail"]


def test_manual_part_identity_edges_merged_and_excluded_rejected(db):
    project = _project(db, "project-manual-part-edges")
    survivor = _part(db, "PN-MANUAL-SURVIVOR")
    merged = _part(db, "PN-MANUAL-OLD", status="merged", merged_into_id=survivor.id)
    excluded = _part(db, "PN-MANUAL-EXC", is_excluded=True)
    client = _client(db, username="manual_edges_admin")

    for bad_part in (merged, excluded):
        response = client.post(
            f"/api/maintenance/site-issues/projects/{project.project_id}/manual",
            json=_create_body(
                [bad_part], key=f"manual-edge-{uuid4().hex[:8]}-01"
            ),
        )
        assert response.status_code == 400, response.text


def test_manual_quantity_edges_rejected_without_500(db):
    project = _project(db, "project-manual-qty-edges")
    part = _part(db, "PN-MANUAL-QTY")
    client = _client(db, username="manual_qty_admin")
    for quantity in ("0.0001", "1e100", "NaN", "Infinity", "0", "-2"):
        response = client.post(
            f"/api/maintenance/site-issues/projects/{project.project_id}/manual",
            json=_create_body(
                [part], key=f"manual-qe-{uuid4().hex[:8]}-01", quantity=quantity
            ),
        )
        assert response.status_code == 422, (
            f"quantity={quantity} expected 422, got {response.status_code}"
        )
    # 合法 3 位小数仍可提交
    ok = client.post(
        f"/api/maintenance/site-issues/projects/{project.project_id}/manual",
        json=_create_body(
            [part], key=f"manual-qe-{uuid4().hex[:8]}-02", quantity="0.125"
        ),
    )
    assert ok.status_code == 201, ok.text


def test_manual_patch_not_page_manual_source_rejected(db):
    """workbook 来源的单不能走页面人工通道改（源语义必须诚实）。"""
    project = _project(db, "project-manual-source-guard")
    part = _part(db, "PN-MANUAL-WB")
    issue = MaintenanceSiteIssue(
        issue_id=str(uuid4()),
        project_id=project.project_id,
        issue_no="WB-MANUAL-0001",
        issue_date=date(2026, 9, 20),
        raw_status="已确认",
        status_mapping_state="mapped",
        normalized_status="confirmed",
        status_mapping_version="workbook-manual-v1",
        source="workbook",
        import_batch_id="batch-wb",
        version=1,
    )
    db.add(issue)
    db.flush()
    db.add(
        MaintenanceSiteIssueLine(
            issue_line_id=f"{issue.issue_id}:1",
            issue_id=issue.issue_id,
            line_no=1,
            part_id=part.id,
            pn=part.pn_std,
            quantity=Decimal("1"),
            is_active=True,
            algorithm_version="workbook-manual-v1",
        )
    )
    db.commit()
    client = _client(db, username="manual_guard_admin")
    response = client.patch(
        f"/api/maintenance/site-issues/manual/{issue.issue_id}",
        json={
            "project_id": project.project_id,
            "version": 1,
            "idempotency_key": f"manual-grd-{uuid4().hex[:8]}-01",
            "receiver": "新人",
            "reason": "误闯通道",
        },
    )
    assert response.status_code == 400, response.text
    assert "不是页面人工登记单" in response.json()["detail"]


def test_manual_voided_issue_cannot_be_edited_and_no_hard_delete(db):
    project = _project(db, "project-manual-void-rules")
    part = _part(db, "PN-MANUAL-VOID")
    client = _client(db, username="manual_void_admin")
    created = client.post(
        f"/api/maintenance/site-issues/projects/{project.project_id}/manual",
        json=_create_body([part], key=f"manual-vd-{uuid4().hex[:8]}-001"),
    )
    assert created.status_code == 201, created.text
    issue = created.json()
    voided = client.post(
        f"/api/maintenance/site-issues/{issue['issue_id']}/void",
        json={
            "project_id": project.project_id,
            "version": issue["version"],
            "idempotency_key": f"manual-vd2-{uuid4().hex[:8]}-01",
            "reason": "整单作废",
        },
    )
    assert voided.status_code == 200, voided.text
    edited = client.patch(
        f"/api/maintenance/site-issues/manual/{issue['issue_id']}",
        json={
            "project_id": project.project_id,
            "version": voided.json()["version"],
            "idempotency_key": f"manual-vd3-{uuid4().hex[:8]}-01",
            "receiver": "新人",
            "reason": "作废后修改",
        },
    )
    assert edited.status_code == 409, edited.text
    db.expire_all()
    assert db.get(MaintenanceSiteIssue, issue["issue_id"]) is not None
    assert (
        db.query(MaintenanceSiteIssueLine)
        .filter_by(issue_id=issue["issue_id"])
        .count()
        > 0
    )


def test_manual_create_receipt_persisted_with_create_action(db):
    """幂等回执真实落库（action='create'），回放走回执不重算。"""
    project = _project(db, "project-manual-receipt")
    part = _part(db, "PN-MANUAL-RECEIPT")
    client = _client(db, username="manual_receipt_admin")
    key = f"manual-rc-{uuid4().hex[:8]}-001"
    created = client.post(
        f"/api/maintenance/site-issues/projects/{project.project_id}/manual",
        json=_create_body([part], key=key),
    )
    assert created.status_code == 201, created.text
    db.expire_all()
    receipt = (
        db.query(MaintenanceSiteIssueCommand)
        .filter_by(idempotency_key=key)
        .one()
    )
    assert receipt.action == "create"
    assert receipt.issue_id == created.json()["issue_id"]


def test_manual_scoped_user_cannot_cross_project_create_or_void(db):
    """受限角色（sales，无挂靠无销售关系）：项目 scope 失败关闭，跨项目拒绝。"""
    project = _project(db, "project-manual-scoped")
    stranger = _project(db, "project-manual-scoped-stranger")
    part = _part(db, "PN-MANUAL-SCOPED")
    db.add(
        SysUser(
            username="manual_scoped_sales",
            role="sales",
            display_name="受限销售",
            password_hash=hash_password("manual-password-123"),
            permissions={
                "page_maintenance": True,
                "action_maintenance_site_issue_manage": True,
                "data_purchase_cost": True,
            },
        )
    )
    db.commit()
    app = FastAPI()
    app.include_router(auth.router, prefix="/api")
    app.include_router(
        maintenance_project_operations.stable_site_issue_router, prefix="/api"
    )
    client = TestClient(app)
    login = client.post(
        "/api/auth/login",
        json={
            "username": "manual_scoped_sales",
            "password": "manual-password-123",
        },
    )
    assert login.status_code == 200, login.text
    client.headers["Authorization"] = f"Bearer {login.json()['token']}"

    # 未挂靠项目 → create 拒绝（含自己项目的 URL）
    denied = client.post(
        f"/api/maintenance/site-issues/projects/{stranger.project_id}/manual",
        json=_create_body([part], key=f"manual-scp-{uuid4().hex[:8]}-01"),
    )
    assert denied.status_code == 403, denied.text
    assert db.query(MaintenanceSiteIssue).count() == 0

    # admin 先建一张，再让受限用户尝试 void（跨项目 + 无挂靠）
    admin_client = _client(db, username="manual_scoped_admin")
    created = admin_client.post(
        f"/api/maintenance/site-issues/projects/{project.project_id}/manual",
        json=_create_body([part], key=f"manual-scp2-{uuid4().hex[:8]}-01"),
    )
    assert created.status_code == 201, created.text
    void_denied = client.post(
        f"/api/maintenance/site-issues/{created.json()['issue_id']}/void",
        json={
            "project_id": project.project_id,
            "version": created.json()["version"],
            "idempotency_key": f"manual-scp3-{uuid4().hex[:8]}-01",
            "reason": "越权作废",
        },
    )
    assert void_denied.status_code == 403, void_denied.text


def test_manual_issue_visible_in_06_export_and_workbook_edit_conflicts(db):
    """page_manual 行进 06 导出；页面改过后，旧 Excel 的行 guard 冲突明确暴露。"""
    import io

    from openpyxl import load_workbook

    from app.services import (
        maintenance_project_master_workbook as master,
    )

    project = _project(db, "project-manual-06-export")
    part = _part(db, "PN-MANUAL-06")
    client = _client(db, username="manual_export_admin")
    created = client.post(
        f"/api/maintenance/site-issues/projects/{project.project_id}/manual",
        json={
            **_create_body([part], key=f"manual-ex-{uuid4().hex[:8]}-001"),
            "issue_no": "LYR-EXPORT-0001",
        },
    )
    assert created.status_code == 201, created.text
    issue_id = created.json()["issue_id"]
    line_id = created.json()["lines"][0]["issue_line_id"]

    # 06 导出包含该行（身份稳定 = issue_line_id）
    db.commit()  # 释放 TestClient 提交后的会话状态
    content = master.build_project_master_v2(
        db, project_id=project.project_id, sheets=(master.V2_SHEET_SITE,)
    )
    assert content is not None
    db.rollback()
    worksheet = load_workbook(io.BytesIO(content))[master.V2_SHEET_SITE]
    headers = {cell.value: cell.column for cell in worksheet[1]}
    exported_ids = {
        str(worksheet.cell(row, headers["实体ID"]).value or "")
        for row in range(2, worksheet.max_row + 1)
    }
    assert line_id in exported_ids

    # 页面先改掉该行数量；再把旧导出（基线=旧数量）中该行数量改成第三个值
    # 上传——用户触碰 + 服务端已变 + 与用户值不同 → 必须报 conflict，
    # 不能静默覆盖页面改动。
    db.expire_all()
    fresh = db.get(MaintenanceSiteIssue, issue_id)
    patched = client.patch(
        f"/api/maintenance/site-issues/manual/{issue_id}",
        json={
            "project_id": project.project_id,
            "version": fresh.version,
            "idempotency_key": f"manual-ex2-{uuid4().hex[:8]}-01",
            "lines": [
                {
                    "issue_line_id": line_id,
                    "part_id": part.id,
                    "quantity": "9",
                }
            ],
            "reason": "页面更正数量",
        },
    )
    assert patched.status_code == 200, patched.text

    edited = load_workbook(io.BytesIO(content))
    sheet = edited[master.V2_SHEET_SITE]
    headers = {cell.value: cell.column for cell in sheet[1]}
    for row in range(2, sheet.max_row + 1):
        if str(sheet.cell(row, headers["实体ID"]).value or "") == line_id:
            sheet.cell(row, headers["领用数量"]).value = 5  # 第三个值：≠基线≠服务端
    buffer = io.BytesIO()
    edited.save(buffer)
    db.rollback()
    plan = master.validate_project_master_v2(
        db, project_id=project.project_id, data=buffer.getvalue()
    )
    conflicts = [
        conflict
        for conflict in getattr(plan, "conflicts", ())
        if conflict.get("entity_id") == line_id
    ]
    assert conflicts, f"旧 Excel 触碰行未与页面改动冲突: {plan.conflicts}"


def test_manual_patch_bumps_line_version_for_cost_cas(db):
    """更正后 line.version 必须 +1：fill_manual_cost 按 line.version 做 CAS，
    不 bump 会接受更正前的过期补价凭据（Codex 复审 P1）。"""
    project = _project(db, "project-manual-line-version")
    part = _part(db, "PN-MANUAL-LV")
    client = _client(db, username="manual_lv_admin")
    created = client.post(
        f"/api/maintenance/site-issues/projects/{project.project_id}/manual",
        json=_create_body([part], key=f"manual-lv-{uuid4().hex[:8]}-001"),
    )
    assert created.status_code == 201, created.text
    line_id = created.json()["lines"][0]["issue_line_id"]
    old_version = created.json()["lines"][0]["version"]

    patched = client.patch(
        f"/api/maintenance/site-issues/manual/{created.json()['issue_id']}",
        json={
            "project_id": project.project_id,
            "version": created.json()["version"],
            "idempotency_key": f"manual-lv2-{uuid4().hex[:8]}-01",
            "lines": [
                {
                    "issue_line_id": line_id,
                    "part_id": part.id,
                    "quantity": "7",
                }
            ],
            "reason": "数量更正，行版本必须前进",
        },
    )
    assert patched.status_code == 200, patched.text
    updated_line = patched.json()["lines"][0]
    assert updated_line["issue_line_id"] == line_id
    # 行事实与成本结果都变 → version 至少前进 1（次数是实现细节，前进是契约）
    assert updated_line["version"] > old_version

    # 旧 line.version 提交补价 → 必须撞 CAS（409），不能接受过期凭据
    db.expire_all()
    from app.services import maintenance_project_operations as ops

    with pytest.raises(ops.MaintenanceOperationConflict):
        ops.fill_manual_cost(
            db,
            project_id=project.project_id,
            issue_line_id=line_id,
            version=old_version,
            manual_unit_cost=Decimal("10.00"),
            evidence="过期凭据",
            reason="应当被 CAS 拒绝",
            operated_by="tester",
        )
    # 新版本可读且仍是原值（未被过期凭据写入）
    line = db.get(MaintenanceSiteIssueLine, line_id)
    assert line.manual_unit_cost is None


def test_manual_patch_same_part_without_sn_keeps_both_line_identities(db):
    """同 PN 且无 SN 的两行：只编辑其中一行，另一行的身份/价格证据不变，
    不允许 (PN, SN) 猜测顶替（Codex 复审 P1）。"""
    project = _project(db, "project-manual-twin-lines")
    part = _part(db, "PN-MANUAL-TWIN")
    batch = SysImportBatch(
        filename="manual-twin.xlsx",
        file_type="maintenance",
        file_hash="manual-twin",
        status="success",
    )
    db.add(batch)
    db.flush()
    client = _client(db, username="manual_twin_admin")
    created = client.post(
        f"/api/maintenance/site-issues/projects/{project.project_id}/manual",
        json={
            "idempotency_key": f"manual-tw-{uuid4().hex[:8]}-001",
            "issue_date": "2026-09-20",
            "receiver": "接收",
            "issued_by": "发出",
            "site_location": "现场",
            "lines": [
                {"part_id": part.id, "quantity": "1"},
                {"part_id": part.id, "quantity": "2"},
            ],
            "reason": "同 PN 双行",
        },
    )
    assert created.status_code == 201, created.text
    body = created.json()
    line_a, line_b = body["lines"][0], body["lines"][1]
    assert line_a["issue_line_id"] != line_b["issue_line_id"]

    # 只更正第一行（显式 id + 新数量），第二行原样保留（显式 id + 原数量）
    patched = client.patch(
        f"/api/maintenance/site-issues/manual/{body['issue_id']}",
        json={
            "project_id": project.project_id,
            "version": body["version"],
            "idempotency_key": f"manual-tw2-{uuid4().hex[:8]}-01",
            "lines": [
                {
                    "issue_line_id": line_a["issue_line_id"],
                    "part_id": part.id,
                    "quantity": "5",
                },
                {
                    "issue_line_id": line_b["issue_line_id"],
                    "part_id": part.id,
                    "quantity": "2",
                },
            ],
            "reason": "只改第一行数量",
        },
    )
    assert patched.status_code == 200, patched.text
    lines = patched.json()["lines"]
    by_id = {line["issue_line_id"]: line for line in lines}
    assert by_id[line_a["issue_line_id"]]["quantity"] == "5.000"
    assert by_id[line_b["issue_line_id"]]["quantity"] == "2.000"
    assert by_id[line_b["issue_line_id"]]["version"] == line_b["version"]

    # 外来 id / 重复 id → 拒绝
    stale_patch = client.patch(
        f"/api/maintenance/site-issues/manual/{body['issue_id']}",
        json={
            "project_id": project.project_id,
            "version": patched.json()["version"],
            "idempotency_key": f"manual-tw3-{uuid4().hex[:8]}-01",
            "lines": [
                {
                    "issue_line_id": "foreign-line-id",
                    "part_id": part.id,
                    "quantity": "1",
                }
            ],
            "reason": "外来行 id",
        },
    )
    assert stale_patch.status_code == 400, stale_patch.text
    assert "不属于本单" in stale_patch.json()["detail"]


def test_manual_patch_add_after_removing_last_line_keeps_history(db):
    """连续更正删除末行再新增：新编号不能复用软作废历史行的唯一键。"""
    project = _project(db, "project-manual-history-lines")
    part = _part(db, "PN-MANUAL-HISTORY")
    client = _client(db, username="manual_history_admin")
    created = client.post(
        f"/api/maintenance/site-issues/projects/{project.project_id}/manual",
        json=_create_body([part, part], key=f"manual-history-{uuid4().hex}"),
    )
    assert created.status_code == 201, created.text
    issue = created.json()
    kept, removed = issue["lines"]
    kept_payload = {
        "issue_line_id": kept["issue_line_id"],
        "part_id": part.id,
        "quantity": kept["quantity"],
    }
    endpoint = f"/api/maintenance/site-issues/manual/{issue['issue_id']}"
    trimmed = client.patch(endpoint, json={
        "project_id": project.project_id,
        "version": issue["version"],
        "idempotency_key": f"manual-remove-{uuid4().hex}",
        "lines": [kept_payload],
        "reason": "删除末行，保留历史",
    })
    assert trimmed.status_code == 200, trimmed.text
    appended = client.patch(endpoint, json={
        "project_id": project.project_id,
        "version": trimmed.json()["version"],
        "idempotency_key": f"manual-append-{uuid4().hex}",
        "lines": [kept_payload, _line_payload(part, "3")],
        "reason": "再次更正新增一行",
    })
    assert appended.status_code == 200, appended.text
    active = appended.json()["lines"]
    assert active[0]["issue_line_id"] == kept["issue_line_id"]
    assert active[1]["issue_line_id"] not in {
        kept["issue_line_id"], removed["issue_line_id"],
    }
    db.expire_all()
    history = db.query(MaintenanceSiteIssueLine).filter_by(
        issue_id=issue["issue_id"],
    ).order_by(MaintenanceSiteIssueLine.line_no).all()
    assert [line.line_no for line in history] == [1, 2, 3]
    assert [line.is_active for line in history] == [True, False, True]
    assert history[1].issue_line_id == removed["issue_line_id"]
    assert history[1].version > removed["version"]
    assert history[2].quantity == Decimal("3.000")
    assert db.query(MaintenanceProjectOperationAudit).filter_by(
        entity_id=issue["issue_id"], action="correct",
    ).count() == 2


def test_manual_patch_rejects_archived_project(db):
    """归档项目只读：PATCH 与 create 同样拒绝（Codex 核对补充）。"""
    project = _project(db, "project-manual-archived")
    part = _part(db, "PN-MANUAL-ARCH")
    client = _client(db, username="manual_arch_admin")
    created = client.post(
        f"/api/maintenance/site-issues/projects/{project.project_id}/manual",
        json=_create_body([part], key=f"manual-ar-{uuid4().hex[:8]}-001"),
    )
    assert created.status_code == 201, created.text
    db.expire_all()
    row = db.get(MaintenanceProject, project.project_id)
    row.is_active = False
    db.commit()
    rejected = client.patch(
        f"/api/maintenance/site-issues/manual/{created.json()['issue_id']}",
        json={
            "project_id": project.project_id,
            "version": created.json()["version"],
            "idempotency_key": f"manual-ar2-{uuid4().hex[:8]}-01",
            "receiver": "新人",
            "reason": "归档后修改",
        },
    )
    assert rejected.status_code == 400, rejected.text
    assert "已归档" in rejected.json()["detail"]


def test_manual_concurrent_duplicate_issue_no_fails_closed(db):
    """同单号并发/重复：唯一约束 + 显式 409，不产生第二张单。"""
    project = _project(db, "project-manual-dup-no")
    part = _part(db, "PN-MANUAL-DUP")
    client = _client(db, username="manual_dup_admin")
    first = client.post(
        f"/api/maintenance/site-issues/projects/{project.project_id}/manual",
        json={
            **_create_body(
                [part], key=f"manual-dup-{uuid4().hex[:8]}-001"
            ),
            "issue_no": "LYR-DUP-0001",
        },
    )
    assert first.status_code == 201, first.text
    second = client.post(
        f"/api/maintenance/site-issues/projects/{project.project_id}/manual",
        json={
            **_create_body(
                [part], key=f"manual-dup-{uuid4().hex[:8]}-002"
            ),
            "issue_no": "LYR-DUP-0001",
        },
    )
    assert second.status_code == 409, second.text
    db.expire_all()
    assert (
        db.query(MaintenanceSiteIssue)
        .filter_by(issue_no="LYR-DUP-0001")
        .count()
        == 1
    )
