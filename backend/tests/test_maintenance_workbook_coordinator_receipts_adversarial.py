"""Independent W2b adversarial evidence: PostgreSQL receipts and commit ownership."""
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from dataclasses import replace
from io import BytesIO
from threading import Event
from time import monotonic, sleep

import pytest
from openpyxl import load_workbook
from sqlalchemy import event, func, insert, select, text
from sqlalchemy.exc import IntegrityError

from app.db import SessionLocal
from app.models.maintenance import FMaintenanceLine
from app.models.maintenance_project_operations import (
    MaintenanceProjectOperationAudit as Audit,
    MaintenanceProjectWorkbookOperation as Operation,
    MaintenanceProjectWorkbookState as State,
)
from app.services import maintenance_project_master_workbook as master
from app.services import maintenance_workbook_coordinator as coordinator
from app.services import maintenance_workbook_receipts as receipts
from app.services.maintenance_workbook_plans import MasterV2Plan, WorkbookError
from tests.test_maintenance_project_master_v2_editable import _make_project_with_line

ACTOR = 'independent-receipts'


def _fixture(db):
    project, _, _, line = _make_project_with_line(db)
    content = master.build_project_master_v2(db, project_id=project.project_id,
                                            sheets=(master.V2_SHEET_PARTS,))
    book = load_workbook(BytesIO(content))
    sheet = book[master.V2_SHEET_PARTS]
    header = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, header['需求数量'], 11)
    stream = BytesIO()
    book.save(stream)
    project_id, line_id = project.project_id, line.id
    db.commit()
    plan = master.validate_project_master_v2(db, project_id=project_id, data=stream.getvalue())
    db.rollback()
    return project_id, line_id, plan


def _operation(db, plan, **changes):
    values = dict(project_id=plan.project_id, export_id=plan.export_id,
                  file_sha256=plan.file_sha256, payload_hash=plan.file_sha256,
                  operation_key=receipts._v2_apply_operation_key(plan, ACTOR),
                  operated_by=ACTOR, operation_type='file_apply')
    values.update(changes)
    row = Operation(**values)
    db.add(row)
    db.flush()
    return row


def _envelope(operation):
    return dict(schema='master_v2_apply_receipt_v1', operation_key=operation.operation_key,
                requires_contract_amount_manage=True,
                response=dict(project_id=operation.project_id, applied_by=ACTOR,
                              replayed=False, import_batch_id='original-batch',
                              changes=[{'nested': {'rows': [1, {'value': 'original'}]}}]))


def _audit(db, operation, envelope, **changes):
    values = dict(project_id=operation.project_id, entity_type='master_workbook_operation',
                  entity_id=str(operation.id), action='apply_completed', after_json=envelope,
                  reason='independent synthetic evidence', operated_by=ACTOR,
                  operated_at=operation.created_at)
    values.update(changes)
    row = Audit(**values)
    db.add(row)
    db.flush()
    return row


def _read(db, plan):
    return receipts.get_project_master_v2_replay(db, project_id=plan.project_id,
                                               file_sha256=plan.file_sha256, operated_by=ACTOR)


def _snapshot(db, project_id, line_id):
    db.expire_all()
    state = db.get(State, project_id)
    return (db.get(FMaintenanceLine, line_id).qty, state.revision, state.data_version,
            db.scalar(select(func.count()).select_from(Operation)),
            db.scalar(select(func.count()).select_from(Audit)))


@pytest.mark.parametrize('field', ['project_id', 'export_id', 'file_sha256', 'operator'])
def test_exact_request_identity_includes_each_dimension(field):
    plan = MasterV2Plan(project_id='p', sheets=(), export_id='e', file_sha256='f')
    original = receipts._v2_apply_operation_key(plan, ACTOR)
    changed = (receipts._v2_apply_operation_key(plan, 'other') if field == 'operator'
               else receipts._v2_apply_operation_key(replace(plan, **{field: 'different'}), ACTOR))
    assert original != changed
    assert receipts._v2_apply_operation_key(replace(plan, export_id=''), ACTOR) is None
    assert receipts._v2_apply_operation_key(replace(plan, file_sha256=''), ACTOR) is None


@pytest.mark.parametrize('corruption', [
    'non_object', 'schema', 'operation_key', 'requires_int', 'requires_missing',
    'response_list', 'project', 'actor', 'replayed_int', 'audit_actor', 'duplicate',
])
def test_persisted_corrupt_evidence_is_not_historical_fallback(db, corruption):
    _, _, plan = _fixture(db)
    operation = _operation(db, plan)
    envelope = _envelope(operation)
    audit_changes = {}
    if corruption == 'non_object': envelope = ['not-an-object']
    elif corruption == 'schema': envelope['schema'] = 'unknown'
    elif corruption == 'operation_key': envelope['operation_key'] = 'other'
    elif corruption == 'requires_int': envelope['requires_contract_amount_manage'] = 1
    elif corruption == 'requires_missing': del envelope['requires_contract_amount_manage']
    elif corruption == 'response_list': envelope['response'] = []
    elif corruption == 'project': envelope['response']['project_id'] = 'other'
    elif corruption == 'actor': envelope['response']['applied_by'] = 'other'
    elif corruption == 'replayed_int': envelope['response']['replayed'] = 0
    elif corruption == 'audit_actor': audit_changes['operated_by'] = 'other'
    _audit(db, operation, envelope, **audit_changes)
    if corruption == 'duplicate': _audit(db, operation, deepcopy(envelope))
    db.commit()
    with pytest.raises(WorkbookError) as caught:
        _read(db, plan)
    assert caught.value.code == 'idempotency_conflict'
    # The locked service replay must reject the same poisoned evidence.
    with pytest.raises(WorkbookError) as locked:
        coordinator.apply_project_master_v2(db, plan, operated_by=ACTOR, import_batch_id='retry')
    assert locked.value.code == 'idempotency_conflict'
    db.rollback()


def test_nested_detachment_preserves_amount_metadata_and_unrelated_evidence(db):
    _, _, plan = _fixture(db)
    operation = _operation(db, plan)
    audit = _audit(db, operation, _envelope(operation))
    db.commit()
    stored = deepcopy(audit.after_json)
    result, needs_amount = _read(db, plan)
    assert needs_amount is True
    result['changes'][0]['nested']['rows'][1]['value'] = 'caller-corruption'
    assert audit.after_json == stored
    assert _read(db, plan)[0] == {**stored['response'], 'replayed': True}
    assert receipts.get_project_master_v2_replay(db, project_id='other-project',
               file_sha256=plan.file_sha256, operated_by=ACTOR) is None
    assert receipts.get_project_master_v2_replay(db, project_id=plan.project_id,
               file_sha256='other-sha', operated_by=ACTOR) is None


def test_genuine_historical_operation_does_not_write_or_backfill_receipt(db):
    project_id, line_id, plan = _fixture(db)
    _operation(db, plan)
    db.commit()
    before = _snapshot(db, project_id, line_id)
    assert _read(db, plan) is None
    result = coordinator.apply_project_master_v2(db, plan, operated_by=ACTOR, import_batch_id='new-attempt')
    assert result['replayed'] is True
    assert _snapshot(db, project_id, line_id) == before
    assert _read(db, plan) is None


@pytest.mark.parametrize('bad_identity', ['payload_hash', 'export_id', 'operated_by', 'operation_type'])
def test_locked_receipt_rejects_identity_mismatch_before_any_fact_write(db, bad_identity):
    project_id, line_id, plan = _fixture(db)
    value = 'file_export' if bad_identity == 'operation_type' else 'other'
    _operation(db, plan, **{bad_identity: value})
    db.commit()
    before = _snapshot(db, project_id, line_id)
    with pytest.raises(WorkbookError) as caught:
        coordinator.apply_project_master_v2(db, plan, operated_by=ACTOR, import_batch_id='attempt')
    assert caught.value.code == 'idempotency_conflict'
    db.rollback()
    assert _snapshot(db, project_id, line_id) == before


def test_read_rejects_ambiguous_same_upload_even_when_operation_keys_differ(db):
    _, _, plan = _fixture(db)
    first = _operation(db, plan)
    _audit(db, first, _envelope(first))
    _operation(db, plan, operation_key='other-key', export_id='other-export')
    db.commit()
    with pytest.raises(WorkbookError) as caught:
        _read(db, plan)
    assert caught.value.code == 'idempotency_conflict'


def test_real_postgres_constraint_at_commit_rolls_back_all_and_session_recovers(db):
    project_id, line_id, plan = _fixture(db)
    before = _snapshot(db, project_id, line_id)
    commits = []
    observed = []
    def record_commit(session): commits.append(True)
    def collide_at_commit(session):
        # Operation is already flushed; completion audit and facts share this tx.
        operation = session.scalar(select(Operation).where(Operation.operation_type == 'file_apply'))
        assert operation is not None
        try:
            session.execute(insert(Operation).values(project_id=project_id,
                export_id=plan.export_id, file_sha256=plan.file_sha256,
                payload_hash=plan.file_sha256, operation_key=operation.operation_key,
                operation_type='file_apply', operated_by=ACTOR))
        except IntegrityError as error:
            observed.append(error.orig.diag.constraint_name)
            raise
    event.listen(db, 'before_commit', collide_at_commit)
    event.listen(db, 'after_commit', record_commit)
    try:
        with pytest.raises(WorkbookError) as caught:
            coordinator.apply_project_master_v2(db, plan, operated_by=ACTOR, import_batch_id='first')
        assert caught.value.code == 'apply_conflict'
        assert observed == ['uq_maintenance_project_workbook_operation_key']
        assert not commits and db.is_active
        assert _snapshot(db, project_id, line_id) == before
        assert 'INSERT' not in str(caught.value) and 'spareparts' not in str(caught.value)
    finally:
        event.remove(db, 'before_commit', collide_at_commit)
        event.remove(db, 'after_commit', record_commit)
    result = coordinator.apply_project_master_v2(db, plan, operated_by=ACTOR, import_batch_id='retry')
    assert result['replayed'] is False
    assert db.get(FMaintenanceLine, line_id).qty == 11


def test_concurrent_duplicate_waits_then_replays_exact_committed_receipt(db):
    project_id, line_id, plan = _fixture(db)
    before = _snapshot(db, project_id, line_id)
    db.rollback()
    first_at_commit, release_first, second_lock_attempt = Event(), Event(), Event()
    commit_counts = {'first': 0, 'second': 0}
    backend_pids = {}
    def run(label):
        with SessionLocal() as session:
            session.execute(text("SET LOCAL lock_timeout = '12s'"))
            backend_pids[label] = session.scalar(select(func.pg_backend_pid()))
            def before_commit(_):
                if label == 'first':
                    first_at_commit.set()
                    assert release_first.wait(10), 'test failed to release first writer'
            def after_commit(_): commit_counts[label] += 1
            def before_execute(state):
                if label == 'second' and 'pg_advisory_xact_lock' in str(state.statement):
                    second_lock_attempt.set()
            event.listen(session, 'before_commit', before_commit)
            event.listen(session, 'after_commit', after_commit)
            event.listen(session, 'do_orm_execute', before_execute)
            return coordinator.apply_project_master_v2(session, plan, operated_by=ACTOR,
                                                        import_batch_id=label)
    with ThreadPoolExecutor(max_workers=2) as pool:
        first = pool.submit(run, 'first')
        try:
            assert first_at_commit.wait(10)
            with SessionLocal() as observer:
                assert _read(observer, plan) is None  # uncommitted completion is invisible
            second = pool.submit(run, 'second')
            assert second_lock_attempt.wait(10)
            assert not second.done()
            # Observe PostgreSQL's actual blocking graph, not only thread timing.
            deadline = monotonic() + 5
            blockers = []
            with SessionLocal() as observer:
                while monotonic() < deadline:
                    blockers = observer.scalar(select(func.pg_blocking_pids(backend_pids['second'])))
                    if backend_pids['first'] in blockers:
                        break
                    sleep(0.01)
            assert backend_pids['first'] in blockers
        finally:
            release_first.set()
        initial, replay = first.result(timeout=20), second.result(timeout=20)
    assert initial['replayed'] is False and replay == {**initial, 'replayed': True}
    assert commit_counts == {'first': 1, 'second': 1}
    after = _snapshot(db, project_id, line_id)
    assert after[0] == 11 and after[1] == before[1] + 1
    assert after[3] == before[3] + 1
    assert db.scalar(select(func.count()).select_from(Audit).where(Audit.action == 'apply_completed')) == 1
