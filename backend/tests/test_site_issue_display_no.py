"""LY-YYYYMMDD-NNNN 展示号（决策 0006）：分配、replay 复用、并发唯一性与展示口径。

内部身份仍是 LY-AUTO-<sha256[:20]>-<sha256[:24]>（去重/replay/唯一约束原样），
展示号只是首次落库时按日分配的可读别名；迁移链与回填确定性的验收在
``test_site_issue_display_no_migration.py``。
"""
import io
import re
from datetime import date
from threading import Thread

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError

from app.db import SessionLocal
from app.models.maintenance_project_operations import (
    MaintenanceSiteIssue,
    MaintenanceSiteIssueLine,
)
from app.services import maintenance_project_master_workbook as master
from app.services.maintenance_project_operations import (
    search_site_issues,
    site_issue_dict,
)
from tests.test_maintenance_project_master_v2_editable import (
    _make_project_with_line,
    _save,
)
from tests.test_maintenance_master_v2_row_merge import _apply

DISPLAY_RE = re.compile(r"^LY-\d{8}-\d{4}$")


def _book(db, project_id):
    return load_workbook(io.BytesIO(master.build_project_master_v2(
        db, project_id=project_id, sheets=(master.V2_SHEET_SITE,))))


def _append(wb, pn, *, no=None, when="2025-04-21", qty=1):
    wb[master.V2_SHEET_SITE].append([no, when, pn, None, qty])
    return wb[master.V2_SHEET_SITE].max_row


def _issues(db, project_id):
    return list(db.scalars(select(MaintenanceSiteIssue).where(
        MaintenanceSiteIssue.project_id == project_id).order_by(
        MaintenanceSiteIssue.created_at, MaintenanceSiteIssue.issue_id)))


def _lines(db, issues):
    return list(db.scalars(select(MaintenanceSiteIssueLine).where(
        MaintenanceSiteIssueLine.issue_id.in_([i.issue_id for i in issues]))))


def test_auto_rows_get_sequential_display_no_and_user_rows_do_not(db):
    project, part, *_ = _make_project_with_line(db)
    wb = _book(db, project.project_id)
    _append(wb, part.pn_std, when="2025-04-21")
    _append(wb, part.pn_std, when="2025-04-21")
    _append(wb, part.pn_std, when="2025-04-22")
    _append(wb, part.pn_std, no="USER-0009", when="2025-04-21")
    _, result = _apply(db, project.project_id, wb)
    assert result["site_creates"] == 4
    issues = _issues(db, project.project_id)
    autos = [i for i in issues if i.issue_no.startswith("LY-AUTO-")]
    manual = [i for i in issues if i.issue_no == "USER-0009"]
    assert len(autos) == 3 and len(manual) == 1
    by_date: dict[date, list[str]] = {}
    for issue in autos:
        # 身份不变：hash 单号仍是内部主键口径，且比展示号长得多。
        assert issue.issue_no.startswith("LY-AUTO-")
        assert len(issue.issue_no) > len(issue.display_no)
        assert DISPLAY_RE.fullmatch(issue.display_no)
        assert issue.display_no.startswith(f"LY-{issue.issue_date:%Y%m%d}-")
        assert issue.display_no_effective == issue.display_no
        by_date.setdefault(issue.issue_date, []).append(issue.display_no)
    assert sorted(by_date[date(2025, 4, 21)]) == ["LY-20250421-0001", "LY-20250421-0002"]
    assert by_date[date(2025, 4, 22)] == ["LY-20250422-0001"]
    user_issue = manual[0]
    assert user_issue.display_no is None and user_issue.display_no_effective == "USER-0009"


def test_downloaded_sheet_shows_display_no_instead_of_hash(db):
    project, part, *_ = _make_project_with_line(db)
    wb = _book(db, project.project_id)
    _append(wb, part.pn_std)
    _apply(db, project.project_id, wb)
    issue = _issues(db, project.project_id)[0]
    downloaded = _book(db, project.project_id)
    cell = downloaded[master.V2_SHEET_SITE].cell(2, 1).value
    assert cell == issue.display_no
    assert "LY-AUTO-" not in str(cell)


def test_replay_reuses_same_display_no_without_duplicates(db):
    project, part, *_ = _make_project_with_line(db)
    original = _book(db, project.project_id)
    _append(original, part.pn_std)
    _append(original, part.pn_std)
    _apply(db, project.project_id, original)
    issues = _issues(db, project.project_id)
    before = {(i.issue_no, i.display_no) for i in issues}
    assert len(before) == 2
    # 原始文件重放：不产生重复行、不重新分配展示号。
    plan, result = _apply(db, project.project_id, original)
    assert not plan.site_flags and result["site_creates"] == 0
    # 下载件（单号列已是可读号）原样回传：零回写、无改名、无重复。
    plan, result = _apply(db, project.project_id, _book(db, project.project_id))
    assert not plan.site_flags and result["site_creates"] == 0
    issues = _issues(db, project.project_id)
    assert {(i.issue_no, i.display_no) for i in issues} == before
    assert len(_lines(db, issues)) == 2


def test_downloaded_display_cell_is_not_treated_as_rename(db):
    project, part, *_ = _make_project_with_line(db)
    wb = _book(db, project.project_id)
    _append(wb, part.pn_std)
    _apply(db, project.project_id, wb)
    issue = _issues(db, project.project_id)[0]
    downloaded = _book(db, project.project_id)
    downloaded[master.V2_SHEET_SITE].cell(2, 5, 3)  # 只改数量
    _, result = _apply(db, project.project_id, downloaded)
    assert result["site_creates"] == 0
    refreshed = _issues(db, project.project_id)[0]
    # 展示号列的既有值不是改号：hash 身份与展示号都保持不变。
    assert refreshed.issue_no == issue.issue_no
    assert refreshed.display_no == issue.display_no
    (line,) = _lines(db, [refreshed])
    assert line.quantity == 3


def test_deleting_downloaded_row_voids_and_replay_receipt_shows_display(db):
    project, part, *_ = _make_project_with_line(db)
    wb = _book(db, project.project_id)
    _append(wb, part.pn_std)
    _apply(db, project.project_id, wb)
    issue = _issues(db, project.project_id)[0]
    downloaded = _book(db, project.project_id)
    downloaded[master.V2_SHEET_SITE].delete_rows(2)
    plan, _ = _apply(db, project.project_id, downloaded)
    assert [flag.issue_line_id for flag in plan.site_flags if flag.is_void]
    (line,) = _lines(db, [issue])
    assert not line.is_active
    # D-02 作废优先回执面向用户：展示可读号，不再露出 hash。
    plan, _ = _apply(db, project.project_id, wb)
    assert plan.voided_rows
    assert plan.voided_rows[0]["row"] == issue.display_no
    assert "LY-AUTO-" not in plan.voided_rows[0]["row"]


def test_old_era_hash_export_delete_row_still_voids(db, monkeypatch):
    """部署切换兼容：旧版导出（单号列=hash、meta 按身份口径）删行仍按作废处理。"""
    project, part, *_ = _make_project_with_line(db)
    wb = _book(db, project.project_id)
    _append(wb, part.pn_std)
    _apply(db, project.project_id, wb)
    # 模拟旧版导出：展示口径退回 issue_no（hash），签名基线与 meta 随之。
    monkeypatch.setattr(
        MaintenanceSiteIssue, "display_no_effective",
        property(lambda self: self.issue_no))
    old_style = _book(db, project.project_id)
    monkeypatch.undo()
    assert "LY-AUTO-" in str(old_style[master.V2_SHEET_SITE].cell(2, 1).value)
    old_style[master.V2_SHEET_SITE].delete_rows(2)
    plan, _ = _apply(db, project.project_id, old_style)
    assert not plan.conflicts
    (line,) = _lines(db, _issues(db, project.project_id))
    assert not line.is_active


def test_v1_master_roundtrip_accepts_display_cells(db):
    project, part, *_ = _make_project_with_line(db)
    wb = _book(db, project.project_id)
    _append(wb, part.pn_std)
    _apply(db, project.project_id, wb)
    issue = _issues(db, project.project_id)[0]
    v1 = master.build_project_master(
        db, project_id=project.project_id, sheets=(master.SHEET_SITE,))
    assert v1
    wb1 = load_workbook(io.BytesIO(v1))
    assert wb1[master.SHEET_SITE].cell(2, 1).value == issue.display_no
    # 单号列的可读展示值不是身份修改，V1 只读校验不得拒绝。
    plan = master.validate(db, project_id=project.project_id, data=_save(wb1))
    assert not plan.site_flags


def test_dict_and_search_expose_display_no(db):
    project, part, *_ = _make_project_with_line(db)
    wb = _book(db, project.project_id)
    _append(wb, part.pn_std)
    _apply(db, project.project_id, wb)
    issue = _issues(db, project.project_id)[0]
    payload = site_issue_dict(issue, [])
    assert payload["issue_no"] == issue.issue_no  # 身份字段原样保留
    assert payload["display_no"] == issue.display_no
    by_display = search_site_issues(
        db, project_id=project.project_id, q_text=issue.display_no,
        workflow_statuses=["confirmed"], page=1, page_size=10)
    assert [row["issue_id"] for row in by_display["rows"]] == [issue.issue_id]
    by_hash = search_site_issues(
        db, project_id=project.project_id, q_text=issue.issue_no[:14],
        workflow_statuses=["confirmed"], page=1, page_size=10)
    assert [row["issue_id"] for row in by_hash["rows"]] == [issue.issue_id]


def _allocate_in_thread(when, results):
    session = SessionLocal()
    try:
        results.append(master._v2_allocate_site_display_no(session, when))
        session.commit()
    finally:
        session.close()


def test_concurrent_allocations_are_unique(db):
    when = date(2026, 10, 14)
    results: list[str] = []
    threads = [Thread(target=_allocate_in_thread, args=(when, results))
               for _ in range(6)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()
    assert sorted(results) == [f"LY-{when:%Y%m%d}-{i:04d}" for i in range(1, 7)]


def test_duplicate_display_no_rejected_by_unique_index(db):
    project, *_ = _make_project_with_line(db)
    shared = "LY-20260421-0099"
    for suffix in ("one", "two"):
        db.add(MaintenanceSiteIssue(
            issue_id=f"disp-uniq-{suffix}", project_id=project.project_id,
            issue_no=f"LY-AUTO-uniq-{suffix}", display_no=shared,
            issue_date=date(2025, 4, 21), raw_status="已确认",
            status_mapping_state="mapped", normalized_status="confirmed",
            status_mapping_version="workbook-manual-v1", source="workbook",
            import_batch_id=f"disp-uniq-batch-{suffix}", version=1))
    with pytest.raises(IntegrityError):
        db.flush()
    db.rollback()
