"""Independent canonical amount-write branch evidence, real HTTP/PG/Cerbos."""
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.authz import runtime
from app.config import get_settings
from app.models.authorization import SysAuthorizationArtifact, SysAuthorizationGrant
from app.models.maintenance_project import MaintenanceProjectAuditLog, MaintenanceProjectContract
from app.models.maintenance_project_operations import MaintenanceProjectWorkbookOperation
from app.models.system import SysImportBatch
from app.services import maintenance_project_master_workbook as workbook
from tests.test_authz_s3_master_workbooks import master, grant, setup, cerbos_url  # noqa: F401
from tests.test_maintenance_project_master_v2_editable import _make_project_with_line

AMOUNT = 'maintenance_contract_amount_write'


@pytest.fixture
def amount_case(db, master, monkeypatch):
    uid, _, client, _ = master
    settings = get_settings()
    monkeypatch.setattr(settings, 'authz_enabled_resources', settings.authz_enabled_resources +
                        ['maintenance_general_workbook_write', AMOUNT])
    runtime.activate_configured_resources(db.get_bind())
    project, _, _, _ = _make_project_with_line(db)
    pid = project.project_id
    grant(db, uid, 'maintenance_general_workbook',
          ['data_purchase_cost', 'data_profit', 'data_customer'], 'all_projects')
    db.add(SysAuthorizationGrant(user_id=uid, resource='maintenance_general_workbook_write',
        actions=['validate', 'apply'], scope='all_projects',
        field_permissions=['data_purchase_cost', 'data_profit', 'data_customer'], version=1))
    db.commit()
    url = f'/api/maintenance/projects/stable/{pid}/master-workbook'
    response = client.get(url + '.xlsx', params={'sheets': workbook.V2_SHEET_OVERVIEW})
    assert response.status_code == 200, response.text
    book = load_workbook(BytesIO(response.content))
    sheet = book[workbook.V2_SHEET_OVERVIEW]
    cell = next(row[1] for row in sheet.iter_rows(min_col=1, max_col=2)
                if row[0].value == '合同总额（含税）')
    assert Decimal(str(cell.value)) == Decimal('10000.00')
    cell.value = Decimal('12345.67')
    output = BytesIO()
    book.save(output)
    book.close()
    return uid, pid, client, url, output.getvalue()


def _state(db, pid):
    db.expire_all()
    row = db.scalar(select(MaintenanceProjectContract).where(MaintenanceProjectContract.project_id == pid))
    return (row.amount_inc_tax, row.version,
            *[db.scalar(select(func.count()).select_from(model)) for model in
              (MaintenanceProjectAuditLog, MaintenanceProjectWorkbookOperation,
               SysAuthorizationArtifact, SysImportBatch)])


def _post(client, url, suffix, content):
    return client.post(url + suffix, files={'file': ('amount.xlsx', content)})


@pytest.mark.parametrize('permission', ['missing', 'validate_only', 'foreign_scope', 'missing_profit'])
def test_amount_apply_cannot_borrow_general_write_action_fields_or_scope(db, amount_case, permission):
    uid, pid, client, url, content = amount_case
    if permission != 'missing':
        db.add(SysAuthorizationGrant(user_id=uid, resource=AMOUNT,
            actions=['validate'] if permission == 'validate_only' else ['validate', 'apply'],
            scope='assigned_or_sales_projects' if permission == 'foreign_scope' else 'all_projects',
            field_permissions=[] if permission == 'missing_profit' else ['data_profit'], version=1))
        db.commit()
    before = _state(db, pid)
    view = _post(client, url, '/apply?view=authorization', content)
    if permission == 'missing_profit':
        # The frozen action requires its own field. Corrupt persisted Grant is
        # diagnosed before returning a view, never repaired from another Grant.
        assert view.status_code == 403, view.text
        assert view.json()['detail']['code'] == 'AUTHZ_INVALID_CONFIGURATION'
    else:
        assert view.status_code == 200, view.text
        decision = view.json()['authorization_context']['resources'][AMOUNT]['decisions']['apply']
        assert decision['effect'] == 'deny'
    denied = _post(client, url, '/apply', content)
    assert denied.status_code == 403, denied.text
    assert _state(db, pid) == before
    validate = _post(client, url, '/validate', content)
    assert validate.status_code == (200 if permission == 'validate_only' else 403), validate.text
    assert _state(db, pid) == before


def test_amount_exact_validate_apply_and_replay_require_current_independent_grant(db, amount_case):
    uid, pid, client, url, content = amount_case
    amount = SysAuthorizationGrant(user_id=uid, resource=AMOUNT, actions=['validate', 'apply'],
        scope='all_projects', field_permissions=['data_profit'], version=1)
    db.add(amount)
    db.commit()
    before = _state(db, pid)
    validate = _post(client, url, '/validate', content)
    assert validate.status_code == 200, validate.text
    assert _state(db, pid) == before
    applied = _post(client, url, '/apply', content)
    assert applied.status_code == 200, applied.text
    assert applied.json()['replayed'] is False
    after = _state(db, pid)
    assert after[0] == Decimal('12345.67') and after[1] == before[1] + 1
    audit = db.scalar(select(MaintenanceProjectAuditLog).where(
        MaintenanceProjectAuditLog.project_id == pid,
        MaintenanceProjectAuditLog.action == 'workbook_update_amount'))
    assert audit is not None
    assert audit.before_json['amount_inc_tax'] == '10000.00'
    assert audit.after_json['amount_inc_tax'] == '12345.67'
    replay = _post(client, url, '/apply', content)
    assert replay.status_code == 200 and replay.json()['replayed'] is True, replay.text
    assert _state(db, pid) == after
    amount.actions = ['validate']
    amount.version += 1
    db.commit()
    denied = _post(client, url, '/apply', content)
    assert denied.status_code == 403, denied.text
    assert _state(db, pid) == after
