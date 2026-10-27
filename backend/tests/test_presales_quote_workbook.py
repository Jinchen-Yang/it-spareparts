"""采购往返模板的合并分发、草稿确认与可见性契约。"""

from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
import os
import shutil
import subprocess
from types import SimpleNamespace

from openpyxl import load_workbook
import pytest
from sqlalchemy import select, text
from sqlalchemy.exc import DBAPIError

from app import permissions
from app.db import SessionLocal
from app.models.dimensions import DimPart, DimSupplier
from app.models.presales import PresalesLineResponse, PresalesOffer, PresalesWorkbookExport
from app.models.system import SysUser
from app.services.presales import common, inquiries, offers, quote_workbook, visibility
from app.services.presales.workbook_registration import register_workbooks
from app.services.workbook import operations
from app.services.workbook.registry import ApplyFailed, Conflict, Forbidden
from app.services.workbook.protocol import read_metadata
from app.services.workflow import work_items
from app.services.workflow.registry import InvalidTransition


@pytest.fixture()
def quote_data(db, monkeypatch):
    register_workbooks()
    monkeypatch.setattr(common, "now", lambda: datetime(2026, 10, 25, 7, tzinfo=timezone.utc))
    users = {}
    for name, role in (("sales", "sales"), ("other", "sales"), ("buyer", "purchaser"), ("peer", "purchaser")):
        user = SysUser(username=f"workbook-{name}", role=role, display_name=name,
            password_hash="测试不登录", template_perms=permissions.effective(role, None), perm_overrides={})
        db.add(user)
        users[name] = user
    parts = [DimPart(pn_std=f"QUOTE-{i}", description=f"型号{i}") for i in range(4)]
    supplier = DimSupplier(name_raw="精确供应商")
    db.add_all([*parts, supplier])
    db.flush()
    return SimpleNamespace(**users, parts=parts, supplier=supplier)


def _submitted(db, data, *, owner=None, lines=None, exact=False, assigned=None):
    inquiry = inquiries.create_draft(db, (owner or data.sales).id,
        assignee_ids=[(assigned or data.buyer).id],
        lines=lines or [{"part_id": data.parts[0].id, "qty": 2, "require_exact_pn": exact}])
    inquiries.submit(db, inquiry.id, (owner or data.sales).id, inquiry.version)
    return inquiry


def _edit(content, changes):
    book = load_workbook(BytesIO(content))
    sheet = book["采购报价表"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    for row_number, values in changes.items():
        for key, value in values.items():
            sheet.cell(row_number, columns[key], value)
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def _apply(db, data, content):
    preview = operations.preview(db, quote_workbook.QUOTE_SCHEMA.name, content,
        ctx=common.resolve_actor(db, data.buyer.id)[1])
    return operations.apply(db, preview["operation_id"], data.buyer.id, preview["plan_digest"])


def _published(db, data, *, private=False):
    inquiry = inquiries.create_draft(db, data.buyer.id, kind="proactive_offer",
        offer_audience={"mode": "private"} if private else None,
        lines=[{"part_id": data.parts[0].id, "qty": 1}])
    line = inquiries.lines_for(db, inquiry.id)[0]
    offers.save_draft(db, inquiry.id, line.id, data.buyer.id,
        {"unit_price": "19.00", "validity_days": 3}, expected_version=inquiry.version)
    return offers.publish_proactive(db, inquiry.id, data.buyer.id,
        expected_version=inquiry.version, on_existing="new")[0]


def test_same_pn_merges_and_distributes_without_reply(db, quote_data):
    data = quote_data
    first = _submitted(db, data)
    second = _submitted(db, data, owner=data.other)
    content = quote_workbook.export_quote_sheet(db, [first.id, second.id], data.buyer.id)
    sheet = load_workbook(BytesIO(content))["采购报价表"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    assert sheet.max_row == 2
    assert sheet.cell(2, headers["合计数量"]).value == 4
    assert first.inquiry_no in sheet.cell(2, headers["来源单号"]).value
    assert second.inquiry_no in sheet.cell(2, headers["来源单号"]).value
    changed = _edit(content, {2: {"处理方式": "报价", "单价": "１２.５０", "含税": "是", "供应商": data.supplier.name_raw}})
    result = _apply(db, data, changed)
    rows = list(db.scalars(select(PresalesOffer).order_by(PresalesOffer.id)))
    assert len(rows) == 2 and {row.inquiry_id for row in rows} == {first.id, second.id}
    assert all(row.status == "draft" and row.unit_price == Decimal("12.50") for row in rows)
    assert all(row.supplier_id == data.supplier.id for row in rows)
    assert result["response"]["replied"] is False
    assert all(item.status == "pending" for inquiry in (first, second)
        for item in work_items.subject_items(db, inquiries.subject(inquiry), data.buyer.id))
    replay = operations.preview(db, quote_workbook.QUOTE_SCHEMA.name, changed,
        ctx=common.resolve_actor(db, data.buyer.id)[1])
    assert replay["replayed"] and len(list(db.scalars(select(PresalesOffer)))) == 2


def test_exact_requirement_splits_and_unknown_pn_can_skip(db, quote_data):
    data = quote_data
    first = _submitted(db, data)
    second = _submitted(db, data, exact=True)
    third = _submitted(db, data, lines=[{"pn_raw": "新型号", "qty": 1}])
    content = quote_workbook.export_quote_sheet(db, [first.id, second.id, third.id], data.buyer.id)
    sheet = load_workbook(BytesIO(content))["采购报价表"]
    assert sheet.max_row == 4
    result = _apply(db, data, _edit(content, {n: {"处理方式": "不报", "不报原因": "待确认"} for n in range(2, 5)}))
    assert len(result["response"]["saved"]) == 3
    responses = list(db.scalars(select(PresalesLineResponse)))
    assert len(responses) == 3 and all(row.mode == "skip" and row.is_draft for row in responses)
    assert visibility.inquiry_view(db, first, common.resolve_actor(db, data.sales.id)[1])["progress"]["quoted_lines"] == 0


def test_reuse_pins_version_only_when_every_owner_can_view(db, quote_data):
    data = quote_data
    offer = _published(db, data)
    first = _submitted(db, data)
    second = _submitted(db, data, owner=data.other)
    content = quote_workbook.export_quote_sheet(db, [first.id, second.id], data.buyer.id)
    result = _apply(db, data, _edit(content, {2: {"处理方式": "沿用", "沿用报价编号": offer.id}}))
    responses = list(db.scalars(select(PresalesLineResponse)))
    assert len(responses) == 2
    assert all(row.ref_offer_id == offer.id and row.ref_offer_revision_no == offer.revision_no for row in responses)
    assert result["response"]["replied"] is False
    for inquiry in (first, second):
        offers.reply(db, inquiry.id, data.buyer.id, expected_version=inquiry.version)
    assert all(not row.is_draft for row in responses)


def test_hidden_reuse_fails_and_rolls_back(db, quote_data):
    data = quote_data
    offer = _published(db, data, private=True)
    first = _submitted(db, data)
    content = quote_workbook.export_quote_sheet(db, [first.id], data.buyer.id)
    assert f"#{offer.id}：" not in str(list(load_workbook(BytesIO(content))["采购报价表"].values))
    with pytest.raises(ApplyFailed) as error:
        _apply(db, data, _edit(content, {2: {"处理方式": "沿用", "沿用报价编号": offer.id}}))
    assert "不可见" in str(error.value.__cause__)
    assert not list(db.scalars(select(PresalesLineResponse)))


@pytest.mark.parametrize("choice", ["更新", "另报"])
def test_existing_quote_requires_explicit_choice_and_saves_for_reply(db, quote_data, choice):
    data = quote_data
    original = _published(db, data)
    inquiry = _submitted(db, data)
    content = quote_workbook.export_quote_sheet(db, [inquiry.id], data.buyer.id)
    result = _apply(db, data, _edit(content, {2: {"处理方式": "报价", "单价": 22,
        "已有报价时的选择": choice, "目标报价编号": original.id if choice == "更新" else None}}))
    draft = db.scalar(select(PresalesOffer).where(PresalesOffer.inquiry_id == inquiry.id))
    assert draft.status == "draft" and original.unit_price == Decimal("19.00")
    assert result["response"]["replied"] is False
    published = offers.reply(db, inquiry.id, data.buyer.id, expected_version=inquiry.version)
    assert len(published) == 1
    assert (published[0].id == original.id) == (choice == "更新")
    assert published[0].unit_price == Decimal("22.00")


@pytest.mark.parametrize("split", [False, True])
def test_batch_cannot_save_multiple_updates_to_same_target(db, quote_data, split):
    data = quote_data
    original = _published(db, data)
    first = _submitted(db, data)
    second = _submitted(db, data, exact=split)
    content = quote_workbook.export_quote_sheet(db, [first.id, second.id], data.buyer.id)
    rows = [2, 3] if split else [2]
    changed = _edit(content, {row: {"处理方式": "报价", "单价": 25,
        "已有报价时的选择": "更新", "目标报价编号": original.id} for row in rows})
    with pytest.raises(ApplyFailed) as error:
        _apply(db, data, changed)
    assert "多次更新同一有效报价" in str(error.value.__cause__)
    assert not list(db.scalars(select(PresalesOffer).where(PresalesOffer.status == "draft")))
    assert original.unit_price == Decimal("19.00")


def test_missing_existing_choice_fails_without_persisting_draft(db, quote_data):
    data = quote_data
    _published(db, data)
    inquiry = _submitted(db, data)
    content = quote_workbook.export_quote_sheet(db, [inquiry.id], data.buyer.id)
    with pytest.raises(ApplyFailed):
        _apply(db, data, _edit(content, {2: {"处理方式": "报价", "单价": 22}}))
    assert not list(db.scalars(select(PresalesOffer).where(PresalesOffer.inquiry_id == inquiry.id)))


def test_missing_supplier_fails_without_persisting_draft(db, quote_data):
    data = quote_data
    inquiry = _submitted(db, data)
    content = quote_workbook.export_quote_sheet(db, [inquiry.id], data.buyer.id)
    with pytest.raises(ApplyFailed) as error:
        _apply(db, data, _edit(content, {2: {"处理方式": "报价", "单价": 22, "供应商": "不存在"}}))
    assert "精确匹配唯一" in str(error.value.__cause__)
    assert not list(db.scalars(select(PresalesOffer)))


def test_ambiguous_supplier_never_chooses_a_candidate():
    # 当前数据库 name_raw 有唯一约束，模拟兼容旧档案或未来适配器返回重复候选。
    db = SimpleNamespace(scalars=lambda statement: [SimpleNamespace(id=1), SimpleNamespace(id=2)])
    with pytest.raises(InvalidTransition, match="精确匹配唯一"):
        quote_workbook._supplier_id(db, "重复供应商")


def test_late_source_failure_rolls_back_earlier_distribution(db, quote_data, monkeypatch):
    data = quote_data
    first = _submitted(db, data)
    second = _submitted(db, data)
    content = quote_workbook.export_quote_sheet(db, [first.id, second.id], data.buyer.id)
    actual = offers.save_line_decision
    def reject_second(db, inquiry_id, *args, **kwargs):
        if inquiry_id == second.id:
            raise ValueError("第二张来源单校验失败")
        return actual(db, inquiry_id, *args, **kwargs)
    monkeypatch.setattr(offers, "save_line_decision", reject_second)
    with pytest.raises(ApplyFailed):
        _apply(db, data, _edit(content, {2: {"处理方式": "报价", "单价": 22}}))
    assert not list(db.scalars(select(PresalesOffer)))


def test_only_own_active_task_exports_and_source_version_conflicts(db, quote_data):
    data = quote_data
    other = _submitted(db, data, assigned=data.peer)
    with pytest.raises(Exception, match="待办"):
        quote_workbook.export_quote_sheet(db, [other.id], data.buyer.id)
    inquiry = _submitted(db, data)
    content = quote_workbook.export_quote_sheet(db, [inquiry.id], data.buyer.id)
    changed = _edit(content, {2: {"处理方式": "不报", "不报原因": "无货"}})
    preview = operations.preview(db, quote_workbook.QUOTE_SCHEMA.name, changed,
        ctx=common.resolve_actor(db, data.buyer.id)[1])
    line = inquiries.lines_for(db, inquiry.id)[0]
    line.version += 1
    db.flush()
    with pytest.raises(Conflict):
        operations.apply(db, preview["operation_id"], data.buyer.id, preview["plan_digest"])


def test_reuse_offer_updated_after_preview_requires_fresh_export(db, quote_data):
    data = quote_data
    original = _published(db, data)
    inquiry = _submitted(db, data)
    content = _edit(quote_workbook.export_quote_sheet(db, [inquiry.id], data.buyer.id),
        {2: {"处理方式": "沿用", "沿用报价编号": original.id}})
    preview = operations.preview(db, quote_workbook.QUOTE_SCHEMA.name, content,
        ctx=common.resolve_actor(db, data.buyer.id)[1])
    offers.update_offer(db, original.id, data.buyer.id, {"unit_price": "100.00"},
        expected_version=original.version)
    with pytest.raises(Conflict):
        operations.apply(db, preview["operation_id"], data.buyer.id, preview["plan_digest"])
    assert not list(db.scalars(select(PresalesLineResponse)))


def test_reference_lock_is_nowait_and_survives_until_outer_commit(db, quote_data):
    data = quote_data
    original = _published(db, data)
    inquiry = _submitted(db, data)
    content = quote_workbook.export_quote_sheet(db, [inquiry.id], data.buyer.id)
    _, manifest = read_metadata(load_workbook(BytesIO(content)))
    entity = manifest["sheets"]["lines"]["rows"][0]["entity_id"]
    original_id = original.id
    ctx = common.resolve_actor(db, data.buyer.id)[1]
    db.commit()
    with SessionLocal() as other:
        other.scalar(select(PresalesOffer).where(PresalesOffer.id == original_id).with_for_update())
        with pytest.raises(Conflict, match="有效报价正在更新"):
            quote_workbook.load_current(db, "lines", [entity], ctx)
        assert db.scalar(text("SELECT 1")) == 1
        other.rollback()
        assert entity in quote_workbook.load_current(db, "lines", [entity], ctx)
        with pytest.raises(DBAPIError):
            other.scalar(select(PresalesOffer).where(PresalesOffer.id == original_id).with_for_update(nowait=True))
        other.rollback()
    db.commit()


def test_identity_columns_survive_selection_and_rows_cannot_be_added_or_deleted(db, quote_data):
    data = quote_data
    inquiry = _submitted(db, data)
    content = quote_workbook.export_quote_sheet(db, [inquiry.id], data.buyer.id,
        visible_columns={"lines": ["mode", "reason"]})
    sheet = load_workbook(BytesIO(content))["采购报价表"]
    assert [cell.value for cell in sheet[1]] == ["_row_token", "来源单号", "PN", "处理方式", "不报原因"]
    book = load_workbook(BytesIO(content))
    book["采购报价表"].delete_rows(2)
    output = BytesIO()
    book.save(output)
    preview = operations.preview(db, quote_workbook.QUOTE_SCHEMA.name, output.getvalue(),
        ctx=common.resolve_actor(db, data.buyer.id)[1])
    assert preview["plan"]["errors"]
    added = _edit(content, {3: {"处理方式": "不报", "不报原因": "无货"}})
    preview = operations.preview(db, quote_workbook.QUOTE_SCHEMA.name, added,
        ctx=common.resolve_actor(db, data.buyer.id)[1])
    assert preview["plan"]["errors"]
    corrected = operations.preview(db, quote_workbook.QUOTE_SCHEMA.name, content,
        ctx=common.resolve_actor(db, data.buyer.id)[1])
    assert corrected["plan"]["counts"]["unchanged"] == 1


def test_price_only_patch_preserves_different_source_drafts_and_hidden_columns(db, quote_data):
    data = quote_data
    sources = [_submitted(db, data), _submitted(db, data)]
    drafts = []
    for index, inquiry in enumerate(sources):
        line = inquiries.lines_for(db, inquiry.id)[0]
        drafts.append(offers.save_draft(db, inquiry.id, line.id, data.buyer.id,
            {"unit_price": str(10 + index), "supplier_id": data.supplier.id,
             "notes": f"来源备注{index}", "stock_location": f"仓库{index}", "validity_days": 3},
            expected_version=inquiry.version))
    content = quote_workbook.export_quote_sheet(db, [row.id for row in sources], data.buyer.id,
        visible_columns={"lines": ["unit_price"]})
    _apply(db, data, _edit(content, {2: {"单价": "25.00"}}))
    assert [row.unit_price for row in drafts] == [Decimal("25.00"), Decimal("25.00")]
    assert [row.notes for row in drafts] == ["来源备注0", "来源备注1"]
    assert [row.stock_location for row in drafts] == ["仓库0", "仓库1"]
    assert all(row.supplier_id == data.supplier.id and row.validity_days == 3 for row in drafts)


def test_historical_scope_replay_rechecks_task_ownership(db, quote_data):
    data = quote_data
    inquiry = _submitted(db, data)
    content = _edit(quote_workbook.export_quote_sheet(db, [inquiry.id], data.buyer.id),
        {2: {"处理方式": "不报", "不报原因": "无货"}})
    _apply(db, data, content)
    task = work_items.subject_items(db, inquiries.subject(inquiry), data.buyer.id)[0]
    task.assignee_id = data.peer.id
    db.flush()
    with pytest.raises(Forbidden):
        operations.preview(db, quote_workbook.QUOTE_SCHEMA.name, content,
            ctx=common.resolve_actor(db, data.buyer.id)[1])


@pytest.mark.parametrize("level", ["header", "line"])
@pytest.mark.parametrize("preview_before_change", [False, True])
def test_deleted_error_plan_cannot_replay_newly_private_requirements(db, quote_data, level, preview_before_change):
    data = quote_data
    inquiry = _submitted(db, data, lines=[{"part_id": data.parts[0].id, "qty": 2,
        "notes": "原公开行备注"}])
    inquiry.notes = "原公开单头备注"
    db.flush()
    content = quote_workbook.export_quote_sheet(db, [inquiry.id], data.buyer.id)
    book = load_workbook(BytesIO(content))
    book["采购报价表"].delete_rows(2)
    output = BytesIO()
    book.save(output)
    content = output.getvalue()
    ctx = common.resolve_actor(db, data.buyer.id)[1]
    preview = None
    if preview_before_change:
        preview = operations.preview(db, quote_workbook.QUOTE_SCHEMA.name, content, ctx=ctx)
        assert preview["plan"]["errors"]
        assert "原公开" in preview["plan"]["rows"][0]["values"]["customer_requirements"]
    target = inquiry if level == "header" else inquiries.lines_for(db, inquiry.id)[0]
    target.notes_audience_rule_id = common.rule_id(db, {"mode": "private"},
        common.resolve_actor(db, data.sales.id)[1])
    db.flush()
    with pytest.raises(Forbidden):
        operations.preview(db, quote_workbook.QUOTE_SCHEMA.name, content, ctx=ctx)
    if preview is not None:
        with pytest.raises(Forbidden):
            operations.apply(db, preview["operation_id"], data.buyer.id, preview["plan_digest"])


@pytest.mark.parametrize("preview_before_change", [False, True])
def test_deleted_error_plan_cannot_replay_newly_hidden_offer_summary(db, quote_data, preview_before_change):
    data = quote_data
    original = _published(db, data)
    inquiry = _submitted(db, data)
    content = quote_workbook.export_quote_sheet(db, [inquiry.id], data.buyer.id)
    book = load_workbook(BytesIO(content))
    book["采购报价表"].delete_rows(2)
    output = BytesIO()
    book.save(output)
    content = output.getvalue()
    ctx = common.resolve_actor(db, data.buyer.id)[1]
    if preview_before_change:
        preview = operations.preview(db, quote_workbook.QUOTE_SCHEMA.name, content, ctx=ctx)
        assert f"#{original.id}：" in preview["plan"]["rows"][0]["values"]["active_offers"]
    origin = common.get_inquiry(db, original.inquiry_id)
    origin.offer_audience_rule_id = common.rule_id(db, {"mode": "private"}, ctx)
    db.flush()
    with pytest.raises(Forbidden):
        operations.preview(db, quote_workbook.QUOTE_SCHEMA.name, content, ctx=ctx)


def test_export_scope_only_tracks_selected_visible_fields_and_contains_no_values(db, quote_data):
    data = quote_data
    inquiry = _submitted(db, data)
    inquiry.notes = "受众控制的客户要求"
    db.flush()
    content = quote_workbook.export_quote_sheet(db, [inquiry.id], data.buyer.id,
        visible_columns={"lines": ["mode", "reason"]})
    scope = db.scalar(select(PresalesWorkbookExport)).scope
    assert scope["rows"][0]["notes"] == []
    assert scope["rows"][0]["active_offer_ids"] == []
    assert "受众控制的客户要求" not in str(scope)
    inquiry.notes_audience_rule_id = common.rule_id(db, {"mode": "private"},
        common.resolve_actor(db, data.sales.id)[1])
    db.flush()
    preview = operations.preview(db, quote_workbook.QUOTE_SCHEMA.name, content,
        ctx=common.resolve_actor(db, data.buyer.id)[1])
    assert preview["plan"]["counts"]["unchanged"] == 1


@pytest.mark.parametrize("preview_before_change", [False, True])
def test_forwarded_supervisor_export_cannot_leak_private_requirements(db, quote_data, preview_before_change):
    data = quote_data
    data.peer.perm_overrides = {"action_presales_supervise": True}
    db.flush()
    inquiry = inquiries.create_draft(db, data.sales.id, assignee_ids=[data.buyer.id, data.peer.id],
        lines=[{"part_id": data.parts[0].id, "notes": "主管导出的受控客户要求"}])
    inquiries.submit(db, inquiry.id, data.sales.id, inquiry.version)
    line = inquiries.lines_for(db, inquiry.id)[0]
    private = common.rule_id(db, {"mode": "private"}, common.resolve_actor(db, data.sales.id)[1])
    if not preview_before_change:
        line.notes_audience_rule_id = private
        db.flush()
    content = quote_workbook.export_quote_sheet(db, [inquiry.id], data.peer.id)
    book = load_workbook(BytesIO(content))
    assert "主管导出的受控客户要求" in str(list(book["采购报价表"].values))
    book["采购报价表"].delete_rows(2)
    output = BytesIO()
    book.save(output)
    content = output.getvalue()
    ctx = common.resolve_actor(db, data.buyer.id)[1]
    if preview_before_change:
        previous = operations.preview(db, quote_workbook.QUOTE_SCHEMA.name, content, ctx=ctx)
        assert previous["plan"]["errors"]
        line.notes_audience_rule_id = private
        db.flush()
    with pytest.raises(Forbidden):
        operations.preview(db, quote_workbook.QUOTE_SCHEMA.name, content, ctx=ctx)
    if preview_before_change:
        with pytest.raises(Forbidden):
            operations.apply(db, previous["operation_id"], data.buyer.id, previous["plan_digest"])


def test_other_rows_in_another_export_do_not_expand_authorization_scope(db, quote_data):
    data = quote_data
    data.peer.perm_overrides = {"action_presales_supervise": True}
    db.flush()
    first = inquiries.create_draft(db, data.sales.id, assignee_ids=[data.buyer.id, data.peer.id],
        lines=[{"part_id": data.parts[0].id}])
    inquiries.submit(db, first.id, data.sales.id, first.version)
    second = _submitted(db, data, assigned=data.peer,
        lines=[{"part_id": data.parts[1].id, "notes": "无关来源的私密要求",
                "notes_audience": {"mode": "private"}}])
    quote_workbook.export_quote_sheet(db, [first.id, second.id], data.peer.id)
    content = quote_workbook.export_quote_sheet(db, [first.id], data.buyer.id,
        visible_columns={"lines": ["mode", "reason"]})
    preview = operations.preview(db, quote_workbook.QUOTE_SCHEMA.name, content,
        ctx=common.resolve_actor(db, data.buyer.id)[1])
    assert preview["plan"]["counts"]["unchanged"] == 1


def test_forwarded_buyer_draft_cannot_leak_private_purchase_notes(db, quote_data):
    data = quote_data
    inquiry = inquiries.create_draft(db, data.sales.id, assignee_ids=[data.buyer.id, data.peer.id],
        lines=[{"part_id": data.parts[0].id}])
    inquiries.submit(db, inquiry.id, data.sales.id, inquiry.version)
    line = inquiries.lines_for(db, inquiry.id)[0]
    offers.save_draft(db, inquiry.id, line.id, data.peer.id,
        {"unit_price": "12.00", "notes": "另一采购的私密备注", "notes_audience": {"mode": "private"}},
        expected_version=inquiry.version)
    content = quote_workbook.export_quote_sheet(db, [inquiry.id], data.peer.id)
    with pytest.raises(Forbidden):
        operations.preview(db, quote_workbook.QUOTE_SCHEMA.name, content,
            ctx=common.resolve_actor(db, data.buyer.id)[1])


def test_deleted_own_draft_does_not_block_first_receipt_replay(db, quote_data):
    data = quote_data
    inquiry = _submitted(db, data)
    line = inquiries.lines_for(db, inquiry.id)[0]
    draft = offers.save_draft(db, inquiry.id, line.id, data.buyer.id,
        {"unit_price": "12.00", "notes": "本人草稿备注", "notes_audience": {"mode": "private"}},
        expected_version=inquiry.version)
    content = _edit(quote_workbook.export_quote_sheet(db, [inquiry.id], data.buyer.id),
        {2: {"处理方式": "不报", "不报原因": "本轮不报"}})
    result = _apply(db, data, content)
    assert db.get(PresalesOffer, draft.id) is None
    replay = operations.preview(db, quote_workbook.QUOTE_SCHEMA.name, content,
        ctx=common.resolve_actor(db, data.buyer.id)[1])
    assert replay["replayed"] and replay["response"] == result["response"]


def test_optional_source_can_be_cleared_without_clearing_supplier(db, quote_data):
    data = quote_data
    inquiry = _submitted(db, data)
    line = inquiries.lines_for(db, inquiry.id)[0]
    draft = offers.save_draft(db, inquiry.id, line.id, data.buyer.id,
        {"unit_price": "12.00", "source_key": "仓库A", "supplier_id": data.supplier.id},
        expected_version=inquiry.version)
    content = quote_workbook.export_quote_sheet(db, [inquiry.id], data.buyer.id)
    book = load_workbook(BytesIO(content))
    sheet = book["采购报价表"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, columns["货源"]).value = None
    output = BytesIO()
    book.save(output)
    _apply(db, data, output.getvalue())
    assert draft.source_key == "" and draft.supplier_id == data.supplier.id


@pytest.mark.skipif(os.getenv("RUN_SOFFICE_ROUNDTRIP") != "1", reason="显式开启 LibreOffice 往返")
def test_quote_libreoffice_roundtrip(db, quote_data, tmp_path):
    binary = shutil.which("soffice")
    assert binary, "RUN_SOFFICE_ROUNDTRIP=1 需要 soffice"
    inquiry = _submitted(db, quote_data)
    source = tmp_path / "quote.xlsx"
    source.write_bytes(_edit(quote_workbook.export_quote_sheet(db, [inquiry.id], quote_data.buyer.id),
        {2: {"处理方式": "报价", "单价": "12.50", "含税": "是"}}))
    destination = tmp_path / "converted"
    destination.mkdir()
    result = subprocess.run([binary, "-env:UserInstallation=file:///tmp/t16-lo-profile", "--headless",
        "--convert-to", "xlsx", "--outdir", str(destination), str(source)],
        capture_output=True, text=True, timeout=60)
    assert result.returncode == 0, result.stdout + result.stderr
    output = destination / source.name
    assert output.exists(), result.stdout + result.stderr
    result = _apply(db, quote_data, output.read_bytes())
    assert result["response"]["saved"][0]["mode"] == "quote"
