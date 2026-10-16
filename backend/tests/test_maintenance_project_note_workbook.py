"""Signed 01 round trips the independent project note through the master command."""

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.config import get_settings
from app.main import app
from app.models.maintenance_project import MaintenanceProject, MaintenanceProjectAuditLog
from app.models.maintenance_project_operations import MaintenanceProjectWorkbookState
from app.services import maintenance_project_master_workbook as master, maintenance_project_catalog as catalog
from app.db import SessionLocal
from tests.test_maintenance_project_management_api import _admin_client
from tests.test_maintenance_project_master_v2_editable import _make_project_with_line, _save
from tests.test_maintenance_master_v2_editor_permissions import _manager_client, _upload


@pytest.fixture(autouse=True)
def enable_v2(monkeypatch):
    monkeypatch.setattr(get_settings(), "maintenance_project_master_v2_enabled", True)


def _book(client, pid, *, full=False):
    response = client.get(f"/api/maintenance/projects/stable/{pid}/master-workbook.xlsx",
                          params={} if full else {"sheets": master.V2_SHEET_OVERVIEW})
    assert response.status_code == 200, response.text
    return load_workbook(io.BytesIO(response.content))


def _note_cell(wb):
    return next(row[1] for row in wb[master.V2_SHEET_OVERVIEW] if row[0].value == "项目备注")


def _send(client, pid, data, *, action="apply", force=False):
    return client.post(f"/api/maintenance/projects/stable/{pid}/master-workbook/{action}",
        files={"file": ("note.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"force_takeover": "true"} if force else None)


def _patch(client, pid, version, **updates):
    return client.patch(f"/api/maintenance/projects/stable/{pid}", json={
        "version": version, "reason": "维护项目备注", **updates})


def test_signed_note_roundtrip_uses_master_audit_and_replays_once(db):
    client = _admin_client(db, "workbook_note_admin")
    project, _part, _order, _line = _make_project_with_line(db)
    pid = project.project_id
    wb = _book(client, pid)
    _note_cell(wb).value = "现场有过漏水\n下次携带防水箱"
    data = _save(wb)
    validated = _send(client, pid, data, action="validate")
    assert validated.status_code == 200, validated.text
    assert validated.json()["project_note_updates"] == 1
    response = _send(client, pid, data)
    assert response.status_code == 200, response.text
    panel = client.get(f"/api/maintenance/projects/stable/{pid}").json()["project"]
    assert panel["project_note"] == "现场有过漏水\n下次携带防水箱"
    assert panel["no_return_note"] is None
    assert panel["version"] == 2
    audit = db.scalar(select(MaintenanceProjectAuditLog).where(
        MaintenanceProjectAuditLog.action == "update"))
    assert audit.operated_by == "workbook_note_admin"
    assert audit.before_json["project_note"] is None
    assert audit.after_json["project_note"] == panel["project_note"]
    state = db.scalar(select(MaintenanceProjectWorkbookState).where(
        MaintenanceProjectWorkbookState.project_id == pid))
    assert state.revision == 1
    replay = _send(client, pid, data)
    assert replay.status_code == 200, replay.text
    assert replay.json()["replayed"] is True
    db.refresh(state)
    assert state.revision == 1
    assert db.scalar(select(func.count()).select_from(MaintenanceProjectAuditLog)) == 1
    assert _patch(client, pid, 2, project_note="后来的备注").status_code == 200
    assert _send(client, pid, data).json()["replayed"] is True
    assert client.get(f"/api/maintenance/projects/stable/{pid}").json()["project"]["project_note"] == "后来的备注"
    other = _admin_client(db, "other_note_admin")
    assert _send(other, pid, data).status_code == 409


@pytest.mark.parametrize("value", ["汉" * 500, "😀" * 500, "=不执行公式\n@现场标记"])
def test_page_note_exports_as_text_and_excel_can_clear(db, value):
    client = _admin_client(db)
    project, *_ = _make_project_with_line(db)
    pid = project.project_id
    assert _patch(client, pid, 1, project_note=value, no_return_note="硬盘免返").status_code == 200
    wb = _book(client, pid)
    cell = _note_cell(wb)
    assert cell.value == value
    assert cell.data_type == "s"
    assert _upload(client, pid, wb).status_code == 200
    assert client.get(f"/api/maintenance/projects/stable/{pid}").json()["project"]["version"] == 2
    cell.value = None
    assert _upload(client, pid, wb).status_code == 200
    panel = client.get(f"/api/maintenance/projects/stable/{pid}").json()["project"]
    assert panel["project_note"] is None
    assert panel["no_return_note"] == "硬盘免返"


@pytest.mark.parametrize("value", ["汉" * 501, "😀" * 501, " " * 501, "=1+1", 123])
@pytest.mark.parametrize("action", ["validate", "apply"])
def test_note_invalid_excel_is_readable_and_never_clears(db, value, action):
    client = _admin_client(db)
    project, *_ = _make_project_with_line(db)
    pid = project.project_id
    assert _patch(client, pid, 1, project_note="原始备注").status_code == 200
    wb = _book(client, pid)
    _note_cell(wb).value = value
    response = _upload(client, pid, wb, action=action)
    assert response.status_code == 422, response.text
    assert "项目备注" in response.json()["detail"]["message"]
    assert client.get(f"/api/maintenance/projects/stable/{pid}").json()["project"]["project_note"] == "原始备注"


@pytest.mark.parametrize("force", [False, True])
def test_note_version_conflict_does_not_write_other_sheet_or_allow_takeover(db, force):
    client = _admin_client(db)
    project, _part, _order, line = _make_project_with_line(db)
    pid = project.project_id
    wb = _book(client, pid, full=True)
    _note_cell(wb).value = "Excel草稿"
    ws = wb[master.V2_SHEET_PARTS]
    headers = {cell.value: cell.column for cell in ws[1]}
    row = next(r for r in range(2, ws.max_row + 1)
               if str(ws.cell(r, headers["实体ID"]).value or "") == str(line.id))
    ws.cell(row, headers["需求数量"], 3)
    data = _save(wb)
    assert _send(client, pid, data, action="validate").status_code == 200
    assert _patch(client, pid, 1, no_return_note="另一主档字段更新").status_code == 200
    response = _send(client, pid, data, force=force)
    assert response.status_code == 409, response.text
    assert response.json()["detail"]["code"] == "stale_project_note"
    assert "conflicts" not in response.json()["detail"]
    db.refresh(line)
    db.refresh(project)
    assert line.qty == 2
    assert project.project_note is None
    assert project.no_return_note == "另一主档字段更新"


def test_note_apply_rechecks_preloaded_project_after_validation(db):
    client = _admin_client(db)
    project, *_ = _make_project_with_line(db)
    pid = project.project_id
    wb = _book(client, pid)
    _note_cell(wb).value = "Excel草稿"
    with SessionLocal() as stale:
        stale.get(MaintenanceProject, pid)
        plan = master.validate_project_master_v2(stale, project_id=pid, data=_save(wb))
        assert not stale.new and not stale.dirty and not stale.deleted
        assert _patch(client, pid, 1, project_note="并发网页备注").status_code == 200
        with pytest.raises(master.WorkbookError, match="项目备注") as error:
            master.apply_project_master_v2(stale, plan, operated_by="project_master_admin", import_batch_id="note-race")
        assert error.value.code == "stale_project_note"
        stale.rollback()
    assert client.get(f"/api/maintenance/projects/stable/{pid}").json()["project"]["project_note"] == "并发网页备注"


def test_note_and_line_share_one_revision_and_audit_failure_rolls_back(db, monkeypatch):
    client = _admin_client(db)
    project, _part, _order, line = _make_project_with_line(db)
    pid = project.project_id
    wb = _book(client, pid, full=True)
    _note_cell(wb).value = "组合修改"
    ws = wb[master.V2_SHEET_PARTS]
    headers = {cell.value: cell.column for cell in ws[1]}
    row = next(r for r in range(2, ws.max_row + 1) if str(ws.cell(r, headers["实体ID"]).value or "") == str(line.id))
    ws.cell(row, headers["需求数量"], 3)
    data = _save(wb)
    original_audit = catalog._audit
    def fail_audit(*args, **kwargs):
        raise RuntimeError("synthetic audit unavailable")
    monkeypatch.setattr(catalog, "_audit", fail_audit)
    broken = TestClient(app, raise_server_exceptions=False)
    broken.headers.update(client.headers)
    assert _send(broken, pid, data).status_code == 500
    db.refresh(project)
    db.refresh(line)
    assert project.project_note is None and project.version == 1
    assert line.qty == 2
    monkeypatch.setattr(catalog, "_audit", original_audit)
    assert _send(client, pid, data).status_code == 200
    db.refresh(project)
    db.refresh(line)
    assert project.project_note == "组合修改" and line.qty == 3
    assert db.scalar(select(MaintenanceProjectWorkbookState.revision).where(MaintenanceProjectWorkbookState.project_id == pid)) == 1


@pytest.mark.parametrize("action", ["validate", "apply"])
def test_note_upload_requires_login_and_master_rights_beyond_workbook_editor(db, action):
    project, *_ = _make_project_with_line(db)
    manager = _manager_client(db, project)
    wb = _book(manager, project.project_id)
    # A manager may still use an untouched workbook under its existing gate.
    assert _upload(manager, project.project_id, wb, action=action).status_code == 200
    _note_cell(wb).value = "越权备注"
    assert _upload(TestClient(app), project.project_id, wb, action=action).status_code == 401
    response = _upload(manager, project.project_id, wb, action=action)
    assert response.status_code == 403, response.text


@pytest.mark.parametrize("mutation", ["missing", "duplicate", "signature", "cross_project", "legacy"])
def test_note_signed_envelope_and_old_template_compatibility(db, mutation):
    client = _admin_client(db)
    project, *_ = _make_project_with_line(db)
    pid = project.project_id
    assert _patch(client, pid, 1, project_note="保留原备注").status_code == 200
    wb = _book(client, pid)
    ws = wb[master.V2_SHEET_OVERVIEW]
    if mutation == "missing":
        ws.delete_rows(_note_cell(wb).row)
    elif mutation == "duplicate":
        ws.append(["项目备注", "重复"])
    elif mutation == "signature":
        for row in wb[master.V2_SHEET_META]:
            if row[0].value == "project_note_version":
                row[1].value = "999"
    elif mutation == "cross_project":
        other = client.post("/api/maintenance/projects/stable", json={
            "project_code": "OTHER-PROJECT", "display_name": "其他项目", "reason": "合成项目"}).json()
        response = _upload(client, other["project_id"], wb)
        assert response.status_code == 422 and response.json()["detail"]["code"] == "project_mismatch"
        return
    else:
        # Re-sign an authentic previous-format fixture without the new keys.
        meta_rows = [(r[0].value, r[1].value) for r in wb[master.V2_SHEET_META]
                     if r[0].value not in {"project_note_version", "project_note_baseline", "metadata_hmac", "metadata_hmac_key_id", "metadata_hmac_algorithm"}]
        wb.remove(wb[master.V2_SHEET_META])
        meta = wb.create_sheet(master.V2_SHEET_META)
        for row in master._v2_sign_meta_rows(meta_rows):
            meta.append(row)
        ws.delete_rows(_note_cell(wb).row)
    response = _upload(client, pid, wb)
    assert response.status_code == (200 if mutation == "legacy" else 422), response.text
    db.refresh(project)
    assert project.project_note == "保留原备注"


def test_note_changes_do_not_enable_other_overview_writeback(db):
    client = _admin_client(db)
    project, *_ = _make_project_with_line(db)
    name = project.display_name
    wb = _book(client, project.project_id)
    _note_cell(wb).value = "只写项目备注"
    for row in wb[master.V2_SHEET_OVERVIEW]:
        if row[0].value in {"项目名称", "销售人员", "项目经理（负责人）", "服务期"}:
            row[1].value = "不应回写"
    assert _upload(client, project.project_id, wb).status_code == 200
    db.refresh(project)
    assert project.display_name == name
    assert project.project_note == "只写项目备注"


def test_note_and_existing_contract_amount_replay_reaches_the_same_actor_receipt(db):
    client = _admin_client(db)
    project, *_ = _make_project_with_line(db)
    pid = project.project_id
    wb = _book(client, pid)
    _note_cell(wb).value = "合同核对后的项目备注"
    for row in wb[master.V2_SHEET_OVERVIEW]:
        if row[0].value == "合同总额（含税）":
            row[1].value = 12000
    data = _save(wb)
    assert _send(client, pid, data).status_code == 200
    assert _send(client, pid, data, action="validate").status_code == 200
    replay = _send(client, pid, data)
    assert replay.status_code == 200, replay.text
    assert replay.json()["replayed"] is True
    assert _patch(client, pid, 2, project_note="后来另行补充的备注").status_code == 200
    assert _send(client, pid, data).json()["replayed"] is True
    db.refresh(project)
    assert project.project_note == "后来另行补充的备注"
    another = _admin_client(db, "other_contract_note_admin")
    assert _send(another, pid, data).status_code in (409, 422)
    assert db.scalar(select(MaintenanceProjectWorkbookState.revision).where(
        MaintenanceProjectWorkbookState.project_id == pid)) == 2


@pytest.mark.parametrize("bad", ["坏_x000B_字符", "_x0001_", "尾_x001F_", "_xFFFE_"])
@pytest.mark.parametrize("action", ["validate", "apply"])
def test_note_excel_with_control_character_is_rejected(db, action, bad):
    client = _admin_client(db)
    project, *_ = _make_project_with_line(db)
    pid = project.project_id
    assert _patch(client, pid, 1, project_note="原始备注").status_code == 200
    wb = _book(client, pid)
    # Excel 把单元格里的 U+000B 等存成 _x000B_ 转义；openpyxl 读回的就是这串字面文本。
    _note_cell(wb).value = bad
    response = _upload(client, pid, wb, action=action)
    assert response.status_code == 422, response.text
    assert "不可见控制字符" in response.json()["detail"]["message"]
    assert client.get(f"/api/maintenance/projects/stable/{pid}").json()["project"]["project_note"] == "原始备注"


def test_export_survives_legacy_control_characters_in_workbook_text_fields(db):
    client = _admin_client(db)
    project, *_ = _make_project_with_line(db)
    pid = project.project_id
    # 绕过入口校验，模拟入口加固前已落库的脏数据。
    row = db.get(MaintenanceProject, pid)
    row.project_note = "旧\x0b备注\ufffe"
    row.display_name = "旧\x01名称"
    row.salesperson = "销\x1f售"
    db.commit()
    wb = _book(client, pid, full=True)
    assert _note_cell(wb).value == "旧备注"
    values = [cell.value for ws in wb.worksheets for r in ws.iter_rows() for cell in r]
    assert "旧名称" in " ".join(str(v) for v in values if v is not None)
    assert not any(isinstance(v, str) and any(ch in v for ch in "\x0b\x01\x1f\ufffe") for v in values)


@pytest.mark.parametrize("action", ["validate", "apply"])
def test_note_upload_is_422_when_v2_switch_is_off(db, monkeypatch, action):
    client = _admin_client(db)
    project, *_ = _make_project_with_line(db)
    pid = project.project_id
    assert _patch(client, pid, 1, project_note="原始备注").status_code == 200
    wb = _book(client, pid)
    _note_cell(wb).value = "开关关闭时不得写入"
    monkeypatch.setattr(get_settings(), "maintenance_project_master_v2_enabled", False)
    response = _upload(client, pid, wb, action=action)
    assert response.status_code == 422, response.text
    assert client.get(f"/api/maintenance/projects/stable/{pid}").json()["project"]["project_note"] == "原始备注"


def test_upload_text_check_covers_raw_and_escaped_forms_but_not_ordinary_text():
    from app.services.excel_text_safety import has_illegal_excel_upload_text
    assert has_illegal_excel_upload_text("a\x0bb")
    assert has_illegal_excel_upload_text("a\uffffb")
    assert has_illegal_excel_upload_text("a_x000c_b")
    assert not has_illegal_excel_upload_text("制表\t换行\n回车\r")
    assert not has_illegal_excel_upload_text("_x0009_ _x000A_ _x000D_ _x0020_ _x4E2D_")
    assert not has_illegal_excel_upload_text(None)
