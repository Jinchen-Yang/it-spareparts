"""Actual relations with SQL pagination, export and field authorization."""
import csv
import io
import json
from datetime import datetime, timezone

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import event, select

from app.api import maintenance_row_workbench as api
from app.authz import runtime
from app.config import get_settings
from app.db import get_db
from app.models.authorization import SysAuthorizationGrant
from app.models.maintenance_project import MaintenanceProjectUserAssignment
from app.models.maintenance_source_assignment import MaintenanceSourceOrderAssignment
from tests.test_authz_s3_checklist import setup, cerbos_url  # noqa: F401
from tests.test_maintenance_demand_export_ordering import export_data  # noqa: F401
from tests.boss_board_helpers import make_project

BASE = '/api/maintenance/spare-part-lines'
WINDOW = {'range': 'custom', 'from': '2026-07-01', 'to': '2026-07-31'}


@pytest.fixture
def rows(db, setup, export_data, monkeypatch):
    uid, identity, _, _, headers = setup
    monkeypatch.setattr(get_settings(), 'authz_enabled_resources', ['maintenance_acceptance', *api.canonical.RESOURCES])
    runtime.activate_configured_resources(db.get_bind())
    app = FastAPI()
    app.include_router(api.router, prefix='/api')
    app.dependency_overrides[get_db] = lambda: db
    client = TestClient(app)
    client.headers.update(headers)
    project, orders = export_data
    return uid, project, orders, client


def grant(db, uid, resource, actions, *, fields=(), scope='all_projects'):
    row = SysAuthorizationGrant(user_id=uid, resource=resource, actions=list(actions),
        scope=scope, field_permissions=list(fields), version=1)
    db.add(row)
    db.commit()
    return row


def assign(db, uid, pid):
    row = MaintenanceProjectUserAssignment(assignment_id='canonical-row-manager', project_id=pid,
        user_id=uid, responsibility_type='primary_manager', assigned_at=datetime.now(timezone.utc),
        assigned_by='synthetic', assignment_reason='actual scope')
    db.add(row)
    db.commit()
    return row


def test_sql_scope_precedes_count_page_and_both_exports(db, rows):
    uid, project, orders, client = rows
    grant(db, uid, 'maintenance_global_rows', ['read', 'export', 'read_fields'], scope='assigned_projects')
    assign(db, uid, project.project_id)
    other = make_project(db, 'outside-row-project')
    moved = db.scalar(select(MaintenanceSourceOrderAssignment).where(
        MaintenanceSourceOrderAssignment.source_order_id == orders['late'].raw_order_id))
    orders['late'].linked_sales_order_no = None
    db.flush()
    moved.is_active = False
    moved.archived_at = datetime.now(timezone.utc)
    moved.archived_by = 'synthetic'
    moved.version += 1
    db.flush()
    db.add(MaintenanceSourceOrderAssignment(assignment_id='moved-row-order',
        source_order_id=moved.source_order_id, project_id=other.project_id, created_by='synthetic'))
    db.commit()
    response = client.get(BASE + '/rows', params={**WINDOW, 'page': 2, 'page_size': 3})
    assert response.status_code == 200, response.text
    data = response.json()
    assert data['total'] == 20 and len(data['rows']) == 3
    assert all('PN-late-' not in row['pn_std'] for row in data['rows'])
    assert all('effective_cost_amount_ex_tax' not in row for row in data['rows'])
    response = client.get(BASE + '/export', params={**WINDOW, 'format': 'csv'})
    assert response.status_code == 200, response.text
    exported = list(csv.DictReader(io.StringIO(response.content.decode('utf-8-sig'))))
    assert len(exported) == 20 and all('PN-late-' not in row['PN'] for row in exported)
    from openpyxl import load_workbook
    response = client.get(BASE + '/export', params={**WINDOW, 'format': 'xlsx'})
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    assert book.active.max_row == 21
    book.close()
    denied = client.get(BASE + '/rows', params={**WINDOW, 'project_id': other.project_id})
    assert denied.status_code == 200 and denied.json()['total'] == 0


def test_managed_alternative_uses_own_fields_and_real_assignment(db, rows):
    uid, project, _, client = rows
    grant(db, uid, 'maintenance_managed_rows', ['read', 'read_fields', 'export'],
          fields=['data_purchase_cost'], scope='primary_managed_projects')
    relation = assign(db, uid, project.project_id)
    response = client.get(BASE + '/fields', params={'view': 'authorization'})
    assert response.status_code == 200, response.text
    assert set(response.json()) == {'locator', 'authorization_context'}
    assert set(response.json()['authorization_context']['resources']) == {'maintenance_managed_rows'}
    response = client.get(BASE + '/fields')
    assert response.status_code == 200, response.text
    assert response.json()['amount_filter_allowed'] is True
    base = f'/api/maintenance/projects/stable/{project.project_id}/master-workbook'
    assert client.get(base + '/row-fields').status_code == 200
    assert client.get(base + '/rows/export', params={'format': 'csv'}).status_code == 200
    relation.archived_at = datetime.now(timezone.utc)
    relation.archived_by = 'synthetic'
    relation.archive_reason = 'actual removal'
    relation.version += 1
    db.commit()
    assert client.get(base + '/rows/export', params={'format': 'csv'}).status_code == 403
    response = client.get(BASE + '/rows', params=WINDOW)
    assert response.status_code == 200 and response.json()['total'] == 0


def test_broad_scope_cannot_borrow_managed_cost_fields(db, rows):
    uid, project, _, client = rows
    grant(db, uid, 'maintenance_global_rows', ['read', 'read_fields'], fields=[])
    grant(db, uid, 'maintenance_managed_rows', ['read', 'read_fields'],
          fields=['data_purchase_cost'], scope='primary_managed_projects')
    assign(db, uid, project.project_id)
    response = client.get(BASE + '/fields')
    assert response.status_code == 200 and response.json()['amount_filter_allowed'] is False
    filters = json.dumps([{'field': 'effective_cost_amount_ex_tax', 'op': 'range', 'min': '1', 'max': '2'}])
    response = client.get(BASE + '/rows', params={**WINDOW, 'filters': filters})
    assert response.status_code == 403, response.text


def test_read_and_fields_never_grant_export_or_other_resource(db, rows):
    uid, project, _, client = rows
    grant(db, uid, 'maintenance_global_rows', ['read', 'read_fields'])
    assert client.get(BASE + '/rows', params=WINDOW).status_code == 200
    assert client.get(BASE + '/export', params={**WINDOW, 'format': 'csv'}).status_code == 403
    assert client.get(f'/api/maintenance/projects/stable/{project.project_id}/master-workbook/row-fields').status_code == 403


def test_current_relation_and_query_precondition_without_business_writes(db, rows):
    uid, project, _, client = rows
    grant(db, uid, 'maintenance_global_rows', ['read'], scope='assigned_projects')
    relation = assign(db, uid, project.project_id)
    view = client.get(BASE + '/rows', params={**WINDOW, 'view': 'authorization'})
    assert view.status_code == 200
    ctx = view.json()['authorization_context']
    pre = json.dumps({key: ctx[key] for key in ('catalog_version', 'authorization_epoch', 'context')})
    statements = []
    def record(conn, cursor, statement, parameters, context, executemany): statements.append(statement)
    event.listen(db.get_bind(), 'before_cursor_execute', record)
    try:
        result = client.get(BASE + '/rows', params={**WINDOW, 'authorization_precondition': pre})
        assert result.status_code == 200, result.text
    finally:
        event.remove(db.get_bind(), 'before_cursor_execute', record)
    assert not any(sql.lstrip().upper().startswith(('INSERT ', 'UPDATE ', 'DELETE ')) or 'FOR UPDATE' in sql or 'FOR SHARE' in sql for sql in statements)
    assert client.get(BASE + '/rows', params={**WINDOW, 'project_id': project.project_id, 'authorization_precondition': pre}).status_code == 409
    relation.archived_at = datetime.now(timezone.utc)
    relation.archived_by = 'synthetic'
    relation.archive_reason = 'actual removal'
    relation.version += 1
    db.commit()
    assert client.get(BASE + '/rows', params={**WINDOW, 'authorization_precondition': pre}).status_code == 409
    assert client.get(BASE + '/rows', params=WINDOW).json()['total'] == 0


def test_suspended_alternative_cannot_fall_back(db, rows, monkeypatch):
    uid, _, _, client = rows
    grant(db, uid, 'maintenance_global_rows', ['read'])
    monkeypatch.setattr(get_settings(), 'authz_enabled_resources', ['maintenance_acceptance', 'maintenance_global_rows'])
    assert client.get(BASE + '/rows', params=WINDOW).status_code == 503
