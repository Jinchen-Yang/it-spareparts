"""Streaming values obey the existing business-export parsing contract."""

from datetime import datetime
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook

from app.services import maintenance_bulk_import as bulk
from app.services import maintenance_import_streaming as streaming

_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_HEADERS = ["销售订单", "实收金额", "收款单号", "收款日期", "数据状态", "不相关内容"]
_SYSTEM = ["F0000008", "F0000062", "SeqNo", "F0000001", "Status", "Unused"]


def _workbook(*, double_header=False, merged=False, blank=False, second=False):
    workbook = Workbook()
    sheet = workbook.active
    if double_header:
        sheet.append(_SYSTEM)
    sheet.append(_HEADERS)
    first = sheet.max_row + 1
    sheet.append(["XSDD-1", 12, "SKD-1", datetime(2026, 10, 12), "已生效", "discard"])
    if merged:
        sheet.append([None] * 6 if blank else ["XSDD-2", 7, None, None, None, "discard"])
        for column in ("C", "D", "E"):
            sheet.merge_cells(f"{column}{first}:{column}{first + 1}")
    if second:
        other = workbook.create_sheet("Second")
        for row in sheet.iter_rows(values_only=True):
            other.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _rewrite(data, change):
    output = BytesIO()
    with ZipFile(BytesIO(data)) as original, ZipFile(output, "w", ZIP_DEFLATED) as rewritten:
        for item in original.infolist():
            rewritten.writestr(item.filename, change(item.filename, original.read(item.filename)))
    return output.getvalue()


def _detect(data, **limits):
    return streaming.detect_streaming(data, bulk._ADAPTERS, **limits)[1]


@pytest.mark.parametrize("double_header", [False, True])
def test_single_and_system_headers_keep_existing_detection_and_values(double_header):
    data = _workbook(double_header=double_header)
    adapter, ordinary = bulk._detect(data)
    actual_adapter, result = streaming.detect_streaming(data, bulk._ADAPTERS)
    assert actual_adapter.key == adapter.key == "receipt_collection"
    assert result["headers"] == ordinary.headers
    assert result["system_headers"] == ordinary.system_headers
    assert result["field_indexes"] == ordinary.field_indexes
    assert result["field_matches"] == ordinary.field_matches
    assert result["header_rows"] == ordinary.header_rows
    assert result["rows"][0][1][:5] == ordinary.rows[0][1][:5]
    assert result["rows"][0][1][5] is None
    assert isinstance(result["rows"][0][1][3], datetime)


def test_bad_dimension_and_reversed_cell_order_do_not_drop_money():
    def change(name, data):
        if name != "xl/worksheets/sheet1.xml":
            return data
        root = ET.fromstring(data)
        root.find(_NS + "dimension").set("ref", "A1")
        for row in root.findall(_NS + "sheetData/" + _NS + "row"):
            row[:] = list(reversed(row[:]))
        return ET.tostring(root)
    result = _detect(_rewrite(_workbook(double_header=True), change))
    assert result["field_indexes"]["actual_amount"] == 1
    assert result["rows"][0][1][1] == 12
    assert result["rows"][0][1][0] == "XSDD-1"


def test_merge_expansion_matches_normal_reader_and_does_not_keep_unused_values():
    data = _workbook(double_header=True, merged=True)
    _, ordinary = bulk._detect(data)
    result = _detect(data)
    assert result["merged_cells_expanded"] == ordinary.merged_cells_expanded == 3
    assert len(result["rows"]) == 2
    for (_n, got), (_m, expected) in zip(result["rows"], ordinary.rows, strict=True):
        assert got[:5] == expected[:5]
        assert got[5] is None


def test_raw_empty_row_inside_a_merge_is_not_created_as_a_receipt():
    result = _detect(_workbook(merged=True, blank=True))
    assert len(result["rows"]) == 1
    assert result["merged_cells_expanded"] == 3


def test_merged_amount_is_rejected_with_original_range():
    def change(name, data):
        if name != "xl/worksheets/sheet1.xml":
            return data
        root = ET.fromstring(data)
        ET.SubElement(root.find(_NS + "mergeCells"), _NS + "mergeCell", {"ref": "B2:B3"})
        return ET.tostring(root)
    with pytest.raises(streaming.StreamingWorkbookError, match="逐行金额") as error:
        _detect(_rewrite(_workbook(merged=True), change))
    assert error.value.issues[0]["cell_range"] == "B2:B3"
    assert error.value.issues[0]["canonical_field"] == "actual_amount"


@pytest.mark.parametrize("state", ["hidden", "veryHidden", "name_prefix"])
def test_hidden_sheet_is_ignored(state):
    def change(name, data):
        if name != "xl/workbook.xml":
            return data
        root = ET.fromstring(data)
        first = root.find(_NS + "sheets/" + _NS + "sheet")
        first.set("name", "hidden_copy") if state == "name_prefix" else first.set("state", state)
        return ET.tostring(root)
    result = _detect(_rewrite(_workbook(second=True), change))
    assert result["name"] == "Second"


def test_equally_matching_visible_sheets_are_rejected():
    with pytest.raises(streaming.StreamingWorkbookError, match="多个同等匹配"):
        _detect(_workbook(second=True))


@pytest.mark.parametrize("limits, message", [
    ({"max_source_rows": 1}, "源文件行数"),
    ({"max_columns": 5}, "源文件列数"),
    ({"max_cells": 11}, "源文件单元格数"),
])
def test_actual_coordinates_enforce_limits_even_when_dimension_lies(limits, message):
    def change(name, data):
        if name != "xl/worksheets/sheet1.xml":
            return data
        root = ET.fromstring(data)
        root.find(_NS + "dimension").set("ref", "A1")
        return ET.tostring(root)
    with pytest.raises(streaming.StreamingWorkbookError, match=message):
        _detect(_rewrite(_workbook(), change), **limits)


def test_unsorted_row_coordinates_are_rejected_instead_of_silently_skipped():
    def change(name, data):
        if name != "xl/worksheets/sheet1.xml":
            return data
        root = ET.fromstring(data)
        rows = root.find(_NS + "sheetData")
        rows[:] = list(reversed(rows[:]))
        return ET.tostring(root)
    with pytest.raises(streaming.StreamingWorkbookError, match="行号重复或顺序错误"):
        _detect(_rewrite(_workbook(), change))


def test_reader_rejects_oversized_uncompressed_archive(monkeypatch):
    monkeypatch.setattr(streaming, "MAX_UNCOMPRESSED_BYTES", 1)
    with pytest.raises(streaming.StreamingWorkbookError, match="解压总量"):
        _detect(_workbook())


def test_shared_strings_are_resolved_without_losing_rich_text():
    data = _workbook()
    shared = ET.Element(_NS + "sst")
    def change(name, content):
        if name != "xl/worksheets/sheet1.xml":
            return content
        root = ET.fromstring(content)
        for cell in root.findall(".//" + _NS + "c"):
            if cell.get("t") != "inlineStr":
                continue
            value = "".join(node.text or "" for node in cell.iter(_NS + "t"))
            index = len(shared)
            si = ET.SubElement(shared, _NS + "si")
            # Multiple rich-text runs must concatenate to the same source value.
            for part in (value[:2], value[2:]):
                run = ET.SubElement(si, _NS + "r")
                ET.SubElement(run, _NS + "t").text = part
            for child in list(cell):
                cell.remove(child)
            cell.set("t", "s")
            ET.SubElement(cell, _NS + "v").text = str(index)
        return ET.tostring(root)
    converted = _rewrite(data, change)
    output = BytesIO()
    with ZipFile(BytesIO(converted)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for name in source.namelist():
            content = source.read(name)
            if name == "[Content_Types].xml":
                types = ET.fromstring(content)
                ET.SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override", {
                    "PartName": "/xl/sharedStrings.xml",
                    "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml",
                })
                content = ET.tostring(types)
            target.writestr(name, content)
        target.writestr("xl/sharedStrings.xml", ET.tostring(shared))
    result = _detect(output.getvalue())
    assert result["rows"][0][1][0] == "XSDD-1"
    assert result["rows"][0][1][2] == "SKD-1"


def test_receipt_source_rows_have_the_same_limit_on_both_sides_of_compressed_threshold(monkeypatch):
    data = _workbook(double_header=True, merged=True)
    monkeypatch.setattr(bulk, "MAX_PREVIEW_ROWS", 1)
    monkeypatch.setattr(bulk, "MAX_SOURCE_ROWS", 100)
    monkeypatch.setattr(bulk, "STREAM_THRESHOLD_BYTES", len(data) + 1)
    small_adapter, small = bulk._detect(data)
    monkeypatch.setattr(bulk, "STREAM_THRESHOLD_BYTES", len(data) - 1)
    large_adapter, large = bulk._detect(data)
    assert small_adapter.key == large_adapter.key == "receipt_collection"
    assert small == large
    assert len(small.rows) == 2


def test_small_receipt_uses_read_only_loading_and_keeps_malformed_cell_order(monkeypatch):
    def malformed(name, content):
        if name != "xl/worksheets/sheet1.xml":
            return content
        root = ET.fromstring(content)
        root.find(_NS + "dimension").set("ref", "A1")
        for row in root.findall(_NS + "sheetData/" + _NS + "row"):
            row[:] = list(reversed(row[:]))
        return ET.tostring(root)
    data = _rewrite(_workbook(double_header=True, merged=True), malformed)
    original = bulk.load_workbook
    read_modes = []
    def observe_load(*args, **kwargs):
        read_modes.append(kwargs.get("read_only"))
        return original(*args, **kwargs)
    monkeypatch.setattr(bulk, "load_workbook", observe_load)
    monkeypatch.setattr(bulk, "STREAM_THRESHOLD_BYTES", len(data) + 1)
    adapter, result = bulk._detect(data)
    assert adapter.key == "receipt_collection"
    assert result.rows[0][1][0] == "XSDD-1"
    assert result.rows[1][1][2] == "SKD-1"
    assert read_modes == [True]


@pytest.mark.parametrize("large", [False, True])
def test_sales_still_uses_preview_row_limit(large, monkeypatch):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["销售单号", "合同总额(含税)", "不含税金额"])
    sheet.append(["XSDD-1", 100, 100])
    sheet.append(["XSDD-2", 200, 200])
    output = BytesIO()
    workbook.save(output)
    data = output.getvalue()
    monkeypatch.setattr(bulk, "MAX_PREVIEW_ROWS", 1)
    monkeypatch.setattr(bulk, "MAX_SOURCE_ROWS", 100)
    monkeypatch.setattr(bulk, "STREAM_THRESHOLD_BYTES", len(data) - 1 if large else len(data) + 1)
    with pytest.raises(bulk.BulkImportInvalid, match="数据行超过安全上限 1"):
        bulk._detect(data)


def test_small_receipt_still_rejects_two_equally_matching_visible_sheets(monkeypatch):
    data = _workbook(second=True)
    monkeypatch.setattr(bulk, "STREAM_THRESHOLD_BYTES", len(data) + 1)
    with pytest.raises(bulk.BulkImportInvalid, match="多个同等匹配"):
        bulk._detect(data)
