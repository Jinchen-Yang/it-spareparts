"""Independent W1a checks through real Excel, HTTP, and PostgreSQL consumers.

These preserve the current V1/V2/EC adapter differences; they do not introduce
new workbook policy or implement the paused field-protection/restore work.
"""
from copy import deepcopy
from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import event, select, text

from app.api import maintenance_project_master_workbook as master_api
from app.db import SessionLocal
from app.models.maintenance import FMaintenanceLine, FProjectExpense, MaintenanceManualCostOverride
from app.models.maintenance_project_operations import (
    MaintenanceCollectionSnapshot,
    MaintenanceProjectExpenseAttribution,
    MaintenanceProjectOperationAudit,
    MaintenanceProjectWorkbookOperation,
    MaintenanceProjectWorkbookState,
)
from app.services import maintenance_expense_collection_workbook as ec
from app.services import maintenance_project_master_workbook as master
from app.services import maintenance_project_operations as operations
from app.services import maintenance_workbook_plans as plans
from tests.test_maintenance_expense_collection_workbook import _XLSX, _expense, uploader
from tests.test_maintenance_project_master_v2_editable import _make_project_with_line


@pytest.fixture(autouse=True)
def bounded_waits():
    def bound(_session, _transaction, connection):
        connection.execute(text("SET LOCAL lock_timeout = '10s'"))

    event.listen(SessionLocal, "after_begin", bound)
    try:
        yield
    finally:
        event.remove(SessionLocal, "after_begin", bound)


def _columns(sheet):
    return {cell.value: cell.column for cell in sheet[1]}


def _save(book):
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def _setup(db, monkeypatch, mode):
    monkeypatch.setattr(master_api.get_settings(), "maintenance_project_master_v2_enabled", mode == "v2")
    project, _part, _order, line = _make_project_with_line(db, unit_cost=None, cost_source=None)
    _expense(db, raw_line_id="W1A-EXPENSE#1", contract_no="XSDD-EDIT-001")
    state = operations.get_or_create_workbook_state(db, project_id=project.project_id)
    db.commit()
    client = uploader(db, username=f"w1a-workflows-{mode}")
    name = "expense-collection-workbook" if mode == "ec" else "master-workbook"
    route = f"/api/maintenance/projects/stable/{project.project_id}/{name}"
    response = client.get(route + ".xlsx")
    assert response.status_code == 200, response.text
    book = load_workbook(BytesIO(response.content))
    expense_name = master.V2_SHEET_EXPENSE if mode == "v2" else ec.SHEET_EXPENSE
    receipt_name = master.V2_SHEET_RECEIPTS if mode == "v2" else ec.SHEET_COLLECTION
    expenses, receipts = book[expense_name], book[receipt_name]
    expense_columns, receipt_columns = _columns(expenses), _columns(receipts)
    expenses.cell(2, expense_columns["未税金额"], 125.25)
    expenses.cell(2, expense_columns["备注"], "跨层核验／零丢失 α")
    receipt_row = receipts.max_row + 1
    receipts.cell(receipt_row, receipt_columns["合同编号"], "XSDD-EDIT-001")
    receipts.cell(receipt_row, receipt_columns["报告月份"], "2026-09")
    receipts.cell(receipt_row, receipt_columns["回款凭证号"], "W1A-回执-01")
    receipts.cell(receipt_row, receipt_columns["备注"], "跨表同次写入")
    if mode == "v2":
        receipts.cell(receipt_row, receipt_columns["累计实收金额（含税）"], 450.50)
    else:
        receipts.cell(receipt_row, receipt_columns["操作"], "CREATE")
        receipts.cell(receipt_row, receipt_columns["累计回款金额(含税)"], 450.50)
    if mode != "ec":
        parts_name = master.V2_SHEET_PARTS if mode == "v2" else master.SHEET_PARTS
        parts = book[parts_name]
        columns = _columns(parts)
        price = "人工未税单位成本" if mode == "v2" else "未税单位成本"
        reason = "人工成本原因" if mode == "v2" else "变更原因"
        parts.cell(2, columns[price], 17.35)
        parts.cell(2, columns[reason], "W1a 合成跨表核验")
        if mode == "v2":
            parts.cell(2, columns["需求数量"], 3)
    data = _save(book)
    return project, line, state, client, route, data


def _post(client, route, data, action="apply"):
    return client.post(route + "/" + action, files={"file": ("crosscheck.xlsx", data, _XLSX)})


def _snapshot(db, project, line):
    db.expire_all()
    # The commit hook uses a different request Session from the fixture.
    # Resolve the line in this Session instead of reading the fixture's cache.
    line = db.get(FMaintenanceLine, line.id)
    expense = db.scalar(select(FProjectExpense).where(FProjectExpense.raw_line_id == "W1A-EXPENSE#1"))
    attribution = db.get(MaintenanceProjectExpenseAttribution, "bxd:W1A-EXPENSE#1")
    state = db.get(MaintenanceProjectWorkbookState, project.project_id)
    override = db.scalar(select(MaintenanceManualCostOverride).where(MaintenanceManualCostOverride.line_id == line.id))
    snapshots = list(db.scalars(select(MaintenanceCollectionSnapshot).where(
        MaintenanceCollectionSnapshot.project_id == project.project_id,
    )))
    audits = list(db.scalars(select(MaintenanceProjectOperationAudit).where(
        MaintenanceProjectOperationAudit.project_id == project.project_id,
    ).order_by(MaintenanceProjectOperationAudit.id)))
    receipts = list(db.scalars(select(MaintenanceProjectWorkbookOperation).where(
        MaintenanceProjectWorkbookOperation.project_id == project.project_id,
    ).order_by(MaintenanceProjectWorkbookOperation.id)))
    return {
        "line": (line.qty, line.cost_amount_ex_tax, line.cost_amount_inc_tax),
        "line_protection": deepcopy(line.manual_override),
        "expense": (expense.amount, expense.amount_ex_tax, expense.amount_inc_tax, expense.remark),
        "attribution": (attribution.amount_ex_tax, attribution.amount_inc_tax, attribution.version),
        "override": None if override is None else (
            override.unit_cost_ex_tax, override.unit_cost_inc_tax, override.reason, override.version,
        ),
        "collections": [(row.collection_id, row.report_month, row.cumulative_amount, row.status, row.version)
                        for row in snapshots],
        "audits": [(row.id, row.action, deepcopy(row.before_json), deepcopy(row.after_json)) for row in audits],
        "operations": [(row.id, row.operation_key, row.file_sha256) for row in receipts],
        "revision": state.revision,
    }


@pytest.mark.parametrize("mode", ["v1", "v2", "ec"])
def test_real_mixed_roundtrip_uses_neutral_plans_and_preserves_each_adapter(db, monkeypatch, mode):
    project, line, _state, client, route, data = _setup(db, monkeypatch, mode)
    before = _snapshot(db, project, line)
    db.commit()
    seen = []
    original_apply = ec.apply

    def observe_inner(session, plan, **kwargs):
        assert type(plan) is plans.WorkbookPlan
        assert all(type(item) is plans.ExpenseUpdate for item in plan.expense_updates)
        assert all(type(item) is plans.CollectionOp for item in plan.collection_ops)
        seen.append((kwargs.get("commit", True), kwargs.get("bump_revision", True)))
        return original_apply(session, plan, **kwargs)

    monkeypatch.setattr(ec, "apply", observe_inner)
    preview = _post(client, route, data, "validate")
    assert preview.status_code == 200, preview.text
    assert preview.json()["expense_updates"] == 1
    assert seen == []
    assert _snapshot(db, project, line) == before
    db.commit()
    applied = _post(client, route, data)
    assert applied.status_code == 200, applied.text
    expected = [(True, True)] if mode == "ec" else [(False, False)]
    assert seen == expected
    after = _snapshot(db, project, line)
    assert after["expense"] == (Decimal("125.25"), Decimal("125.25"), Decimal("141.53"), "跨层核验／零丢失 α")
    assert after["attribution"][:2] == (Decimal("125.25"), Decimal("141.53"))
    assert len(after["collections"]) == 1
    assert after["collections"][0][1:4] == (date(2026, 9, 1), Decimal("450.50"), "confirmed")
    assert after["revision"] == before["revision"] + 1
    if mode == "ec":
        assert after["line"] == before["line"] and after["override"] is None
    else:
        assert after["override"][:2] == (Decimal("17.35"), Decimal("19.61"))
        assert after["line"][0] == Decimal("3" if mode == "v2" else "2")
    db.commit()
    if mode == "v2":
        # Receipt persistence is tested after a mixed 03/04/05 application,
        # including nested EC; neither plan reconstruction nor EC may run again.
        replay = _post(client, route, data)
        assert replay.status_code == 200, replay.text
        assert replay.json() == {**applied.json(), "replayed": True}
        assert seen == [(False, False)]
        assert _snapshot(db, project, line) == after
        completed = [row for row in after["audits"] if row[1] == "apply_completed"]
        assert len(completed) == len(after["operations"]) == 1
        assert completed[0][3]["response"] == applied.json()
        db.commit()
    exported = client.get(route + ".xlsx")
    assert exported.status_code == 200, exported.text
    book = load_workbook(BytesIO(exported.content))
    sheet = book[master.V2_SHEET_EXPENSE if mode == "v2" else ec.SHEET_EXPENSE]
    columns = _columns(sheet)
    assert Decimal(str(sheet.cell(2, columns["未税金额"]).value)) == Decimal("125.25")
    assert sheet.cell(2, columns["备注"]).value == "跨层核验／零丢失 α"
    book.close()


@pytest.mark.parametrize("mode", ["v1", "v2", "ec"])
@pytest.mark.parametrize("error_code", ["cost_recompute_busy", "row_conflicts"])
def test_neutral_error_after_flushed_mixed_writes_rolls_back_and_allows_same_file_retry(
    db, monkeypatch, mode, error_code,
):
    project, line, _state, client, route, data = _setup(db, monkeypatch, mode)
    before = _snapshot(db, project, line)
    db.commit()
    reached = []
    issues = [{"sheet": "04", "row": 2, "field": "未税金额", "server": "125.25"}]
    details = {"rows": [{"row_no": 2}], "total_count": 1, "fix_hint": "合成失败后重试", "internal": "do-not-expose"}

    def fail_after_flush(session):
        session.flush()
        expense = session.scalar(select(FProjectExpense).where(
            FProjectExpense.raw_line_id == "W1A-EXPENSE#1",
        ))
        if expense is None or expense.amount_ex_tax != Decimal("125.25"):
            return
        snapshot = _snapshot(session, project, line)
        assert len(snapshot["collections"]) == 1
        assert snapshot["expense"][1] == Decimal("125.25")
        if mode != "ec":
            assert snapshot["override"][0] == Decimal("17.35")
        if mode == "v2":
            assert len(snapshot["operations"]) == 1
            assert any(row[1] == "apply_completed" for row in snapshot["audits"])
        reached.append(snapshot)
        raise plans.WorkbookError(error_code, "中立异常的合成提交失败", issues, details=details)

    event.listen(SessionLocal, "before_commit", fail_after_flush)
    try:
        rejected = _post(client, route, data)
    finally:
        event.remove(SessionLocal, "before_commit", fail_after_flush)
    assert len(reached) == 1
    assert rejected.status_code == (422 if mode == "ec" else 409), rejected.text
    body = rejected.json()["detail"]
    assert body["code"] == error_code
    assert body["message"] == "中立异常的合成提交失败"
    assert body["issues"] == issues
    assert "internal" not in body
    if mode != "ec":
        assert {key: body[key] for key in ("rows", "total_count", "fix_hint")} == {
            key: details[key] for key in ("rows", "total_count", "fix_hint")
        }
        if error_code == "cost_recompute_busy":
            assert rejected.headers["Retry-After"] == "5"
        else:
            assert body["conflicts"] == issues
    assert _snapshot(db, project, line) == before
    db.commit()
    retry = _post(client, route, data)
    assert retry.status_code == 200, retry.text
    after = _snapshot(db, project, line)
    assert after["revision"] == before["revision"] + 1
    assert after["expense"][1] == Decimal("125.25")
    assert len(after["collections"]) == 1
    if mode == "v2":
        assert retry.json()["replayed"] is False
        assert len(after["operations"]) == 1
