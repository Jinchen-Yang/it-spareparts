"""销售对账区（D-40）与含税三件套（D-41）导入行为契约：transform / loader / pipeline。

全部用测试生成的小 xlsx（双表头：第 0 行 F 码、第 1 行中文列名），无真实客户数据。
"""
from datetime import datetime, timezone
from decimal import Decimal

import openpyxl
import pytest
from sqlalchemy import func, select

from app.etl import loader as etl_loader, mapping, pipeline, reader, transform
from app.models.sales import FSalesOrder, FSalesReconciliationSnapshot
from app.models.system import SysImportError
from app.services.sales_reconciliation import (
    latest_snapshots_by_raw_order_id,
    snapshot_content_hash,
)

RECON_HEADERS = ["已对账金额", "待对账金额", "退货已对账金额", "修改时间(必填)"]
LINE_ID = "订单明细.数据ID(不可修改)"
LINE_NO = "订单明细.序号"


def _row(**over):
    """一行销售导出（头级字段随行重复，模拟 ffill 后形态）。列名带特殊字符，
    调用方用 ``_row(**{"列名": 值})`` 覆盖。"""
    row = {
        "订单编号(必填)": "XSDD-20260630-0001", "数据ID(不可修改)": "RAW-RECON-1",
        "订单日期": "2026-06-30", "销售人员": "测试销售", "客户名称": "测试客户",
        "业务类型#": "备件销售", "订单金额": Decimal("113.00"), "是否含税": "含税",
        "不含税金额": Decimal("100.00"), "税率": "13%", "税金": Decimal("13.00"),
        "数据状态": "已生效",
        LINE_ID: "LINE-1", LINE_NO: 1,
        "订单明细.产品名称": "TEST-PN-1", "订单明细.订单数量": 1,
        "订单明细.单价": Decimal("113.00"), "订单明细.金额": Decimal("113.00"),
        "已对账金额": Decimal("100.00"), "待对账金额": Decimal("13.00"),
        "退货已对账金额": Decimal("0.00"), "修改时间(必填)": "2026-07-01 09:30:00",
    }
    row.update(over)
    return row


def _second_line(**over):
    """同一订单的第二行（头级字段与首行一致）。"""
    row = _row(**{LINE_ID: "LINE-2", LINE_NO: 2})
    row.update(over)
    return row


def _write_xlsx(path, rows, *, extra_sheets=()):
    """双表头销售 xlsx；extra_sheets: [(sheet 名, 是否隐藏, 行内容)] 模拟多页导出。"""
    headers = list(rows[0])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "销售订单"
    ws.append([f"F{i:07d}" for i in range(1, len(headers) + 1)])
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h) for h in headers])
    for name, hidden, content in extra_sheets:
        extra = wb.create_sheet(title=name)
        extra.sheet_state = "hidden" if hidden else "visible"
        for line in content:
            extra.append(line)
    wb.save(path)
    return str(path)


def _transform_file(path):
    df, file_type = reader.read_excel(path)
    assert file_type == mapping.SALES
    return transform.transform(df, file_type)


def _snapshots(db):
    return db.scalars(
        select(FSalesReconciliationSnapshot).order_by(FSalesReconciliationSnapshot.id)
    ).all()


def _purchase_xlsx(path):
    headers = ["采购单号(必填)", "数据ID(不可修改)", "采购日期(必填)", "采购人员(必填)",
               "供应商(必填)", "批量采购(必填)", "不含税金额", "税率(必填)",
               "是否含税(必填)", "税金", "采购金额", "数据状态",
               "明细.数据ID(不可修改)", "明细.序号", "明细.产品名称(必填)",
               "明细.采购数量(必填)", "明细.单价(必填)", "明细.合计金额"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([f"F{i:07d}" for i in range(1, len(headers) + 1)])
    ws.append(headers)
    ws.append(["CGDD-1", "PUR-RAW-1", "2026-06-30", "测试采购", "测试供应商", "备货",
               Decimal("100.00"), "13%", "含税", Decimal("13.00"), Decimal("113.00"),
               "已生效", "PL-1", 1, "TEST-PN-1", 1, Decimal("113.00"), Decimal("113.00")])
    wb.save(path)
    return str(path)


# ---------- transform：列存在性 / 变体 / 旧格式 ----------

def test_transform_parses_reconciliation_columns(tmp_path):
    path = _write_xlsx(tmp_path / "s.xlsx", [_row(), _second_line()])
    result = _transform_file(path)
    assert result.has_reconciliation_columns is True
    head = result.orders["RAW-RECON-1"]
    assert head["reconciled_amount"] == Decimal("100.00")
    assert head["pending_amount"] == Decimal("13.00")
    assert head["return_reconciled_amount"] == Decimal("0.00")
    # 氚云北京时间 2026-07-01 09:30 → UTC 01:30
    assert head["source_modified_at"] == datetime(2026, 7, 1, 1, 30, tzinfo=timezone.utc)
    assert not result.errors


def test_transform_annotated_amounts_and_bare_modified_time(tmp_path):
    row = _row(**{"已对账金额(必填)": Decimal("88.00"), "待对账金额(必填)": Decimal("12.00"),
                  "退货已对账金额(必填)": Decimal("1.00"), "修改时间": "2026-07-02 08:00:00"})
    for h in RECON_HEADERS:
        row.pop(h, None)
    path = _write_xlsx(tmp_path / "s.xlsx", [row])
    result = _transform_file(path)
    assert result.has_reconciliation_columns is True
    head = result.orders["RAW-RECON-1"]
    assert head["reconciled_amount"] == Decimal("88.00")
    assert head["pending_amount"] == Decimal("12.00")
    assert head["return_reconciled_amount"] == Decimal("1.00")
    assert head["source_modified_at"] == datetime(2026, 7, 2, 0, 0, tzinfo=timezone.utc)


def test_transform_bare_amounts_and_annotated_modified_time(tmp_path):
    path = _write_xlsx(tmp_path / "s.xlsx", [
        _row(已对账金额=Decimal("60.00"), **{"修改时间(必填)": "2026-07-03 18:00:00"})])
    result = _transform_file(path)
    head = result.orders["RAW-RECON-1"]
    assert head["reconciled_amount"] == Decimal("60.00")
    assert head["source_modified_at"] == datetime(2026, 7, 3, 10, 0, tzinfo=timezone.utc)


def test_transform_old_format_without_reconciliation_columns(tmp_path):
    row = _row()
    for h in RECON_HEADERS:
        row.pop(h)
    path = _write_xlsx(tmp_path / "s.xlsx", [row])
    result = _transform_file(path)
    assert result.has_reconciliation_columns is False
    head = result.orders["RAW-RECON-1"]
    assert "reconciled_amount" not in head and "source_modified_at" not in head
    assert not result.errors


def test_transform_empty_cells_are_null_not_zero(tmp_path):
    path = _write_xlsx(tmp_path / "s.xlsx", [
        _row(已对账金额=None, 待对账金额="", 退货已对账金额=None, **{"修改时间(必填)": None})])
    result = _transform_file(path)
    head = result.orders["RAW-RECON-1"]
    assert head["reconciled_amount"] is None
    assert head["pending_amount"] is None
    assert head["return_reconciled_amount"] is None
    assert head["source_modified_at"] is None
    assert not result.errors


# ---------- transform：坏值与数值格式 ----------

def test_transform_bad_amount_and_bad_time_do_not_block(tmp_path):
    path = _write_xlsx(tmp_path / "s.xlsx", [
        _row(已对账金额="not-a-number", **{"修改时间(必填)": "乱码日期"}),
        _second_line()])
    result = _transform_file(path)
    head = result.orders["RAW-RECON-1"]
    # 首个非空值是坏值 → 锁定 NULL，不被第二行的好值顶掉（确定性规则）
    assert head["reconciled_amount"] is None
    assert head["source_modified_at"] is None
    assert len(result.lines) == 2                  # 行/单照常导入
    types = {e.error_type for e in result.errors}
    assert "bad_reconciliation_amount" in types
    assert "bad_source_modified_time" in types
    assert all(t not in transform.SOFT_ERROR_TYPES for t in
               ("bad_reconciliation_amount", "bad_source_modified_time"))


def test_transform_money_formats_negative_thousands_huge(tmp_path):
    path = _write_xlsx(tmp_path / "s.xlsx", [
        _row(已对账金额="-1,234.56", 待对账金额="¥9876.5",
             退货已对账金额=Decimal("999999999.99"))])
    result = _transform_file(path)
    head = result.orders["RAW-RECON-1"]
    assert head["reconciled_amount"] == Decimal("-1234.56")   # 负数原样保留
    assert head["pending_amount"] == Decimal("9876.50")       # 千分位/货币符
    assert head["return_reconciled_amount"] == Decimal("999999999.99")
    assert not result.errors


def test_transform_modified_time_format_variants(tmp_path):
    excel_serial = ((datetime(2026, 7, 4, 12, 0) - datetime(1899, 12, 30)).total_seconds()
                    / 86400)
    variants = [
        ("2026-07-04 12:00:00", datetime(2026, 7, 4, 4, 0, tzinfo=timezone.utc)),
        ("2026/7/4 09:05", datetime(2026, 7, 4, 1, 5, tzinfo=timezone.utc)),
        (excel_serial, datetime(2026, 7, 4, 4, 0, tzinfo=timezone.utc)),
        (datetime(2026, 7, 4, 20, 0, 0), datetime(2026, 7, 4, 12, 0, tzinfo=timezone.utc)),
        ("2026-07-04", datetime(2026, 7, 3, 16, 0, tzinfo=timezone.utc)),
    ]
    for i, (raw, expected) in enumerate(variants):
        path = _write_xlsx(tmp_path / f"s{i}.xlsx", [
            _row(**{"数据ID(不可修改)": f"RAW-T{i}", "修改时间(必填)": raw})])
        result = _transform_file(path)
        head = result.orders[f"RAW-T{i}"]
        # openpyxl 日期序列化往返会带微秒噪声，比较按秒精度
        actual = head["source_modified_at"].replace(microsecond=0)
        assert actual == expected, f"variant {raw!r}"
        assert not result.errors


def test_transform_out_of_range_serial_is_rejected(tmp_path):
    path = _write_xlsx(tmp_path / "s.xlsx", [_row(**{"修改时间(必填)": 200000})])
    result = _transform_file(path)
    assert result.orders["RAW-RECON-1"]["source_modified_at"] is None
    assert any(e.error_type == "bad_source_modified_time" for e in result.errors)


# ---------- transform：多行单头聚合 ----------

def test_transform_head_conflict_takes_first_nonempty_with_soft_warning(tmp_path):
    path = _write_xlsx(tmp_path / "s.xlsx", [
        _row(已对账金额=Decimal("100.00")),
        _second_line(已对账金额=Decimal("200.00"))])
    result = _transform_file(path)
    head = result.orders["RAW-RECON-1"]
    assert head["reconciled_amount"] == Decimal("100.00")
    conflicts = [e for e in result.errors if e.error_type == "reconciliation_head_conflict"]
    assert len(conflicts) == 1
    assert "reconciliation_head_conflict" in transform.SOFT_ERROR_TYPES


def test_transform_equal_values_across_rows_are_not_conflict(tmp_path):
    path = _write_xlsx(tmp_path / "s.xlsx", [
        _row(已对账金额="100", **{"修改时间(必填)": "2026-07-01 09:30:00"}),
        _second_line(已对账金额=Decimal("100.00"), **{"修改时间(必填)": "2026-07-01 09:30:00"})])
    result = _transform_file(path)
    assert not [e for e in result.errors if e.error_type == "reconciliation_head_conflict"]


def test_transform_first_row_empty_later_row_value_wins(tmp_path):
    path = _write_xlsx(tmp_path / "s.xlsx", [
        _row(已对账金额=None), _second_line(已对账金额=Decimal("250.00"))])
    result = _transform_file(path)
    assert result.orders["RAW-RECON-1"]["reconciled_amount"] == Decimal("250.00")


def test_transform_recon_collected_from_invalid_line_rows(tmp_path):
    """对账区收集点在明细级校验（PN/数量/单价）之前：校验失败行的单头值照常收集。

    原「闸门之后收集」口径（值只在坏行上 → 落 NULL）被 GitHub Codex 审查判定
    P1：单头对账值绝不因明细坏行被静默跳过。本用例第二段即原
    test_transform_recon_only_collected_from_surviving_rows 的反向断言，按新
    口径取反（非削弱：坏行的 bad_number 仍照常上报，收集到的值进入订单头）。
    """
    # 存活行带值、坏行重复同值 → 照常收集（坏行不阻断）
    path = _write_xlsx(tmp_path / "s1.xlsx", [
        _row(已对账金额=Decimal("70.00")),
        _second_line(**{"订单明细.订单数量": "abc"}, 已对账金额=Decimal("70.00"))])
    result = _transform_file(path)
    assert len(result.lines) == 1
    assert result.orders["RAW-RECON-1"]["reconciled_amount"] == Decimal("70.00")

    # 值只在坏行上（存活行该列为空，ffill 只向下不向上）→ 照常收集，不再落 NULL
    path = _write_xlsx(tmp_path / "s2.xlsx", [
        _row(已对账金额=None),
        _second_line(**{"订单明细.订单数量": "abc"}, 已对账金额=Decimal("70.00"))])
    result = _transform_file(path)
    assert len(result.lines) == 1
    assert any(e.error_type == "bad_number" for e in result.errors)   # 坏行仍上报
    assert result.orders["RAW-RECON-1"]["reconciled_amount"] == Decimal("70.00")


def test_transform_recon_failed_row_value_wins_with_conflict_warning(tmp_path):
    """坏行携带的对账值与存活行不同：取文件序首个非空（坏行的），差异软告警。"""
    path = _write_xlsx(tmp_path / "s.xlsx", [
        _row(**{"订单明细.订单数量": "abc"}, 已对账金额=Decimal("70.00")),
        _second_line(已对账金额=Decimal("200.00"))])
    result = _transform_file(path)
    assert len(result.lines) == 1
    assert result.orders["RAW-RECON-1"]["reconciled_amount"] == Decimal("70.00")
    conflicts = [e for e in result.errors if e.error_type == "reconciliation_head_conflict"]
    assert len(conflicts) == 1
    assert "200" in conflicts[0].error_detail


def test_transform_bad_recon_on_detail_invalid_row_still_reported(tmp_path):
    """明细校验失败行上的坏对账值：逐行报 bad_reconciliation_amount，不因行被拒而静默。"""
    path = _write_xlsx(tmp_path / "s.xlsx", [
        _row(**{"订单明细.订单数量": "abc"}, 已对账金额="not-a-number"),
        _second_line(已对账金额=Decimal("100.00"))])
    result = _transform_file(path)
    assert len(result.lines) == 1
    types = {e.error_type for e in result.errors}
    assert "bad_number" in types                       # 明细坏行本身照常上报
    assert "bad_reconciliation_amount" in types        # 坏行上的对账解析错误不再被吞
    # 首个非空是坏值 → 锁定 NULL，存活行的好值只换来一条不一致告警（既有口径）
    assert result.orders["RAW-RECON-1"]["reconciled_amount"] is None
    assert len([e for e in result.errors
                if e.error_type == "reconciliation_head_conflict"]) == 1


def test_import_all_lines_failed_no_order_no_snapshot(db, tmp_path):
    """订单全部明细被拒：订单不入库、不写快照（与订单头一致），行级错误如实可见。

    同一行再植入坏对账值/坏修改时间：对账解析错误不随订单被拒而消失（收集点在
    明细闸门之前，逐行上报）——审查意见 1「全部明细被拒 → 错误要报」的完整锁定。
    """
    path = _write_xlsx(tmp_path / "a.xlsx", [
        _row(**{"订单明细.订单数量": "abc"}, 已对账金额="not-a-number",
              **{"修改时间(必填)": "乱码日期"})])
    batch = pipeline.run_import(db, path, "a.xlsx")
    db.commit()

    assert batch.status == "success"                   # 坏行隔离，整批照常成功
    assert batch.rows_error == 3                       # bad_number + 坏对账金额 + 坏修改时间
    assert db.scalar(select(func.count()).select_from(FSalesOrder)) == 0
    assert _snapshots(db) == []
    error_types = db.scalars(select(SysImportError.error_type)
                             .where(SysImportError.batch_id == batch.id)).all()
    assert {"bad_number", "bad_reconciliation_amount",
            "bad_source_modified_time"} <= set(error_types)
    # 坏对账值错误不带业务身份（identity 只为修复模式作废豁免等真实消费方预留，
    # 无意义的调试残留 _identity 已删）——raw_row 留痕里不应出现 _identity 键
    bad_recon_rows = db.scalars(select(SysImportError.raw_row).where(
        SysImportError.batch_id == batch.id,
        SysImportError.error_type == "bad_reconciliation_amount")).all()
    assert bad_recon_rows
    assert all("_identity" not in (row or {}) for row in bad_recon_rows)


def test_transform_tax_trio_collected_from_invalid_line_rows(tmp_path):
    """含税三件套同类修复：值只在明细校验失败行上 → 按行聚合，不再随首个存活行丢失。"""
    path = _write_xlsx(tmp_path / "s.xlsx", [
        _row(订单金额=None, 是否含税=None, 税金=None),
        _second_line(**{"订单明细.订单数量": "abc"},
                     订单金额=Decimal("226.00"), 是否含税="含税", 税金=Decimal("26.00"))])
    result = _transform_file(path)
    assert len(result.lines) == 1
    head = result.orders["RAW-RECON-1"]
    assert head["amount_inc_tax"] == Decimal("226.00")
    assert head["tax_amount"] == Decimal("26.00")
    assert head["is_tax_inclusive"] is True


def test_import_tax_trio_from_invalid_line_row_persisted(db, tmp_path):
    """上一用例经 loader 落库：三件套与快照 order_amount 都取到坏行携带的值。"""
    path = _write_xlsx(tmp_path / "a.xlsx", [
        _row(订单金额=None, 是否含税=None, 税金=None),
        _second_line(**{"订单明细.订单数量": "abc"},
                     订单金额=Decimal("226.00"), 是否含税="含税", 税金=Decimal("26.00"))])
    pipeline.run_import(db, path, "a.xlsx")
    db.commit()

    order = db.scalar(select(FSalesOrder).where(FSalesOrder.raw_order_id == "RAW-RECON-1"))
    assert order.amount_inc_tax == Decimal("226.00")
    assert order.tax_amount == Decimal("26.00")
    assert order.is_tax_inclusive is True
    snap = db.scalar(select(FSalesReconciliationSnapshot))
    assert snap.order_amount == Decimal("226.00")


def test_transform_tax_trio_conflict_takes_first_nonempty(tmp_path):
    """三件套跨行不一致（坏行 vs 存活行）：文件序首个非空生效 + 软告警（与对账区同规则）。"""
    path = _write_xlsx(tmp_path / "s.xlsx", [
        _row(**{"订单明细.订单数量": "abc"}, 税金=Decimal("26.00"), 是否含税="不含税"),
        _second_line(税金=Decimal("13.00"), 是否含税="含税")])
    result = _transform_file(path)
    head = result.orders["RAW-RECON-1"]
    assert head["tax_amount"] == Decimal("26.00")
    assert head["is_tax_inclusive"] is False          # 坏行的不含税先到先得
    conflicts = [e for e in result.errors if e.error_type == "reconciliation_head_conflict"]
    assert len(conflicts) == 2                         # 税金 + 是否含税 各一条
    joined = "\n".join(e.error_detail for e in conflicts)
    assert "税金" in joined and "是否含税" in joined
    assert "不含税" in joined                           # 布尔按业务语义渲染


def test_transform_tax_trio_equal_values_across_rows_no_conflict(tmp_path):
    """三件套同值跨行（"113" vs 113.00）：解析值相等不算冲突。"""
    path = _write_xlsx(tmp_path / "s.xlsx", [
        _row(订单金额="113"),
        _second_line(订单金额=Decimal("113.00"))])
    result = _transform_file(path)
    assert not [e for e in result.errors if e.error_type == "reconciliation_head_conflict"]
    assert result.orders["RAW-RECON-1"]["amount_inc_tax"] == Decimal("113.00")


def test_transform_tax_trio_old_format_collected_from_invalid_line_rows(tmp_path):
    """旧格式（无对账区列）孪生场景：三件套聚合不依赖对账区列的存在。

    D-41 回填正是重放这类 PR-2 之前的旧件——若聚合回退成「仅新格式文件生效」，
    旧件的 file_value 会静默变 None、被误记 unresolved/漏报冲突，故必须锁死。
    """
    rows = [
        _row(订单金额=None, 是否含税=None, 税金=None),
        _second_line(**{"订单明细.订单数量": "abc"},
                     订单金额=Decimal("226.00"), 是否含税="含税", 税金=Decimal("26.00"))]
    for r in rows:
        for h in RECON_HEADERS:
            r.pop(h, None)
    path = _write_xlsx(tmp_path / "s.xlsx", rows)
    result = _transform_file(path)
    assert result.has_reconciliation_columns is False
    head = result.orders["RAW-RECON-1"]
    assert "reconciled_amount" not in head and "source_modified_at" not in head
    assert head["amount_inc_tax"] == Decimal("226.00")
    assert head["tax_amount"] == Decimal("26.00")
    assert head["is_tax_inclusive"] is True


def test_transform_tax_trio_bad_value_falls_back_to_later_good_row(tmp_path):
    """三件套坏值口径（docstring 承诺）：静默 None、不报错、不锁定，后续好行可补位。

    与对账区「坏值锁定 NULL + 逐行报错」刻意不同（沿用 _build_head 旧口径，
    不新增错误码）；锁定该分歧，防止未来「统一规则」的重构悄悄改变落库值。
    """
    path = _write_xlsx(tmp_path / "s.xlsx", [
        _row(订单金额="not-a-number", 税金="乱码", 是否含税="随便写的"),
        _second_line(订单金额=Decimal("113.00"), 税金=Decimal("13.00"))])
    result = _transform_file(path)
    head = result.orders["RAW-RECON-1"]
    assert head["amount_inc_tax"] == Decimal("113.00")  # 首行坏值不锁空，次行好值补位
    assert head["tax_amount"] == Decimal("13.00")
    assert head["is_tax_inclusive"] is True             # 次行 含税 补位（首行无法识别）
    assert not result.errors                           # 三件套坏值静默（既有口径）


def test_transform_multi_sheet_with_hidden_sheet(tmp_path):
    path = _write_xlsx(
        tmp_path / "s.xlsx", [_row()],
        extra_sheets=[
            ("备注页", True, [["内部备注"], ["客户对账周期按月"]]),
            ("副本", False, [["杂项"]]),
        ])
    result = _transform_file(path)
    assert result.has_reconciliation_columns is True
    assert result.orders["RAW-RECON-1"]["reconciled_amount"] == Decimal("100.00")


def test_transform_purchase_file_has_no_reconciliation_flag(tmp_path):
    df, file_type = reader.read_excel(_purchase_xlsx(tmp_path / "purchase.xlsx"))
    assert file_type == mapping.PURCHASE
    result = transform.transform(df, file_type)
    assert result.has_reconciliation_columns is False


# ---------- pipeline / loader：入库与统计 ----------

def test_skip_first_import_writes_tax_columns_and_snapshot(db, tmp_path):
    path = _write_xlsx(tmp_path / "a.xlsx", [_row()])
    batch = pipeline.run_import(db, path, "a.xlsx")
    db.commit()

    order = db.scalar(select(FSalesOrder).where(FSalesOrder.raw_order_id == "RAW-RECON-1"))
    assert order.amount_inc_tax == Decimal("113.00")
    assert order.tax_amount == Decimal("13.00")
    assert order.is_tax_inclusive is True

    snaps = _snapshots(db)
    assert len(snaps) == 1
    snap = snaps[0]
    assert snap.source == "import"
    assert snap.import_batch_id == batch.id
    assert snap.sales_order_id == order.id
    assert snap.reconciled_amount == Decimal("100.00")
    assert snap.pending_amount == Decimal("13.00")
    assert snap.source_modified_at == datetime(2026, 7, 1, 1, 30, tzinfo=timezone.utc)
    assert snap.content_hash == snapshot_content_hash(
        raw_order_id="RAW-RECON-1",
        source_modified_at=datetime(2026, 7, 1, 1, 30, tzinfo=timezone.utc),
        order_amount=Decimal("113.00"), reconciled_amount=Decimal("100.00"),
        pending_amount=Decimal("13.00"), return_reconciled_amount=Decimal("0.00"))

    report = batch.report_json
    assert report["sales_recon_snapshots_inserted"] == 1
    assert report["sales_recon_snapshots_skipped_same_batch"] == 0
    assert report["sales_tax_columns_backfilled"] == 0


def test_skip_reupload_same_content_appends_new_batch_observation(db, tmp_path):
    """内容相同的重传（不同批次）各记一条观测（D-40：内容不去重），最新取批次大者。"""
    path1 = _write_xlsx(tmp_path / "a.xlsx", [_row()])
    batch1 = pipeline.run_import(db, path1, "a.xlsx")
    db.commit()
    # 不同文件哈希（多一个新单）、旧单对账内容未变
    path2 = _write_xlsx(tmp_path / "b.xlsx", [
        _row(),
        _row(**{"数据ID(不可修改)": "RAW-RECON-2", "订单编号(必填)": "XSDD-20260630-0002",
                LINE_ID: "LINE-B1", "已对账金额": Decimal("50.00")})])
    batch2 = pipeline.run_import(db, path2, "b.xlsx")
    db.commit()

    assert batch2.report_json["sales_recon_snapshots_inserted"] == 2   # 旧单+新单各一条
    assert batch2.report_json["sales_recon_snapshots_skipped_same_batch"] == 0
    assert len(_snapshots(db)) == 3
    old_order_rows = [s for s in _snapshots(db) if s.raw_order_id == "RAW-RECON-1"]
    assert len(old_order_rows) == 2                                   # 两批次各一条
    assert {s.import_batch_id for s in old_order_rows} == {batch1.id, batch2.id}
    latest = latest_snapshots_by_raw_order_id(db, ["RAW-RECON-1"])
    assert latest["RAW-RECON-1"].import_batch_id == batch2.id         # 最新=批次较大者
    assert latest["RAW-RECON-1"].reconciled_amount == Decimal("100.00")


def test_import_aba_with_null_modified_times_latest_is_a(db, tmp_path):
    """A→B→A（修改时间全空）：三个批次各一条观测，最新 = 第三批的 A（不被内容去重吞）。

    三个文件字节不同（明细产品名不同，避免同哈希被 skip 模式拒绝），但
    RAW-RECON-1 的对账内容 A→B→A；旧口径下第三批与首批同哈希会被 ON CONFLICT
    丢弃、最新错误停在 B。
    """

    def _aba(name, reconciled, pn):
        return _write_xlsx(tmp_path / name, [
            _row(已对账金额=Decimal(reconciled), **{"修改时间(必填)": None,
                                                    "订单明细.产品名称": pn})])

    pipeline.run_import(db, _aba("aba1.xlsx", "100.00", "TEST-PN-A"), "aba1.xlsx")
    db.commit()
    pipeline.run_import(db, _aba("aba2.xlsx", "60.00", "TEST-PN-B"), "aba2.xlsx")
    db.commit()
    batch3 = pipeline.run_import(db, _aba("aba3.xlsx", "100.00", "TEST-PN-C"),
                                 "aba3.xlsx")
    db.commit()

    snaps = [s for s in _snapshots(db) if s.raw_order_id == "RAW-RECON-1"]
    assert [s.reconciled_amount for s in snaps] == [
        Decimal("100.00"), Decimal("60.00"), Decimal("100.00")]
    assert all(s.source_modified_at is None for s in snaps)     # 修改时间全空
    latest = latest_snapshots_by_raw_order_id(db, ["RAW-RECON-1"])
    assert latest["RAW-RECON-1"].import_batch_id == batch3.id
    assert latest["RAW-RECON-1"].reconciled_amount == Decimal("100.00")


def test_skip_reupload_fills_null_tax_keeps_existing(db, tmp_path):
    path1 = _write_xlsx(tmp_path / "a.xlsx", [_row(税金=None, 是否含税=None)])
    pipeline.run_import(db, path1, "a.xlsx")
    db.commit()
    order = db.scalar(select(FSalesOrder).where(FSalesOrder.raw_order_id == "RAW-RECON-1"))
    assert order.tax_amount is None and order.is_tax_inclusive is None

    # 第二份文件：非空列(订单金额/不含税)值不同也不覆盖；NULL 列(税金/是否含税)补齐
    path2 = _write_xlsx(tmp_path / "b.xlsx", [
        _row(订单金额=Decimal("999.00"), 不含税金额=Decimal("888.00"), 是否含税="不含税")])
    batch2 = pipeline.run_import(db, path2, "b.xlsx")
    db.commit()
    db.refresh(order)
    assert order.amount_inc_tax == Decimal("113.00")     # 非空不覆盖
    assert order.amount_ex_tax == Decimal("100.00")
    assert order.tax_amount == Decimal("13.00")          # NULL 补齐
    assert order.is_tax_inclusive is False               # NULL 按第二份文件值补齐
    assert batch2.report_json["sales_tax_columns_backfilled"] == 2
    # 新批次 → 本单再记一条观测（本例文件内容也变了；内容未变同样会追加）
    assert batch2.report_json["sales_recon_snapshots_inserted"] == 1
    assert len(_snapshots(db)) == 2


def test_upsert_reimport_overwrites_tax_columns(db, tmp_path):
    path1 = _write_xlsx(tmp_path / "a.xlsx", [_row(税金=None)])
    pipeline.run_import(db, path1, "a.xlsx")
    db.commit()
    path2 = _write_xlsx(tmp_path / "b.xlsx", [
        _row(订单金额=Decimal("226.00"), 不含税金额=Decimal("200.00"), 税金=Decimal("26.00"),
             已对账金额=Decimal("150.00"))])
    batch2 = pipeline.run_import(db, path2, "b.xlsx", mode="upsert")
    db.commit()
    order = db.scalar(select(FSalesOrder).where(FSalesOrder.raw_order_id == "RAW-RECON-1"))
    assert order.amount_inc_tax == Decimal("226.00")     # 覆盖模式如实覆盖
    assert order.tax_amount == Decimal("26.00")
    assert batch2.report_json["sales_recon_snapshots_inserted"] == 1
    snaps = _snapshots(db)
    assert len(snaps) == 2
    assert snaps[-1].reconciled_amount == Decimal("150.00")


def test_file_without_recon_columns_writes_no_snapshot(db, tmp_path):
    row = _row()
    for h in RECON_HEADERS:
        row.pop(h)
    path = _write_xlsx(tmp_path / "a.xlsx", [row])
    batch = pipeline.run_import(db, path, "a.xlsx")
    db.commit()
    assert batch.report_json["sales_recon_snapshots_missing_columns"] == 1
    assert "sales_recon_snapshots_inserted" not in batch.report_json
    assert _snapshots(db) == []
    order = db.scalar(select(FSalesOrder).where(FSalesOrder.raw_order_id == "RAW-RECON-1"))
    assert order.amount_inc_tax == Decimal("113.00")     # 三件套照常写


def test_pipeline_imports_file_with_bad_recon_value(db, tmp_path):
    # 坏值不阻断整批：run_import 照常成功，订单/快照落库，行级错误如实上报
    path = _write_xlsx(tmp_path / "a.xlsx", [_row(已对账金额="not-a-number")])
    batch = pipeline.run_import(db, path, "a.xlsx")
    db.commit()

    assert batch.status == "success"
    assert batch.rows_error == 1
    error_types = db.scalars(select(SysImportError.error_type)
                             .where(SysImportError.batch_id == batch.id)).all()
    assert "bad_reconciliation_amount" in error_types

    order = db.scalar(select(FSalesOrder).where(FSalesOrder.raw_order_id == "RAW-RECON-1"))
    assert order is not None
    snap = db.scalar(select(FSalesReconciliationSnapshot))
    assert snap.reconciled_amount is None            # 坏值 → NULL，快照仍在
    assert snap.pending_amount == Decimal("13.00")


def test_loader_same_batch_replay_skips_existing_observation(db, tmp_path):
    """同批次经 loader 本身重放：ON CONFLICT 吞掉，报告计数如实上报（幂等兜底）。"""
    path = _write_xlsx(tmp_path / "a.xlsx", [_row()])
    batch = pipeline.run_import(db, path, "a.xlsx")
    db.commit()
    assert len(_snapshots(db)) == 1

    # 同一 transform 结果、同一批次再走一次 loader 写入路径
    df, file_type = reader.read_excel(path)
    result = transform.transform(df, file_type)
    inserted, skipped = etl_loader._write_sales_reconciliation_snapshots(
        db, result, batch.id, {})
    db.commit()
    assert (inserted, skipped) == (0, 1)
    assert len(_snapshots(db)) == 1                     # 不新增重复观测


def test_import_rollback_discards_snapshots(db, tmp_path):
    path = _write_xlsx(tmp_path / "a.xlsx", [_row()])
    pipeline.run_import(db, path, "a.xlsx")
    # 不 commit 直接回滚（导入失败的等价事务边界）：快照与订单一并消失
    db.rollback()
    assert db.scalar(select(func.count()).select_from(FSalesOrder)) == 0
    assert db.scalar(select(func.count()).select_from(FSalesReconciliationSnapshot)) == 0


def test_purchase_import_leaves_snapshots_untouched(db, tmp_path):
    batch = pipeline.run_import(db, _purchase_xlsx(tmp_path / "purchase.xlsx"),
                                "purchase.xlsx")
    db.commit()
    assert batch.file_type == mapping.PURCHASE
    for key in ("sales_recon_snapshots_inserted",
                "sales_recon_snapshots_skipped_same_batch",
                "sales_recon_snapshots_missing_columns", "sales_tax_columns_backfilled"):
        assert key not in batch.report_json
    assert _snapshots(db) == []
