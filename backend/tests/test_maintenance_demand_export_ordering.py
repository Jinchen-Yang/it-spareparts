"""D-33：真实下载按业务日期升序，订单连续，稳定尾键与网页原序兼容。"""

import csv
import io
from datetime import date
from decimal import Decimal
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from app.models.dimensions import DimPart
from app.models.maintenance import FMaintenanceLine, FMaintenanceOrder
from app.models.system import SysImportBatch
from app.api import maintenance_project_master_workbook as master_api
from app.services import maintenance_project_master_workbook as master
from app.services import maintenance_cost
from tests.boss_board_helpers import assign, client_for, make_project


_CONTRACT = "XSDD-20260701-0001"
_PROJECT = "导出排序合成项目"
_ORDER_KEYS = ["early", "same-a", "same-a-copy", "same-b", "late", "undated"]
_LINE_KEYS = ["one-a", "one-b", "two", "null-a", "null-b"]


def _expected_pns(*, dated_only=False):
    orders = _ORDER_KEYS[:-1] if dated_only else _ORDER_KEYS
    return [f"PN-{order}-{line}" for order in orders for line in _LINE_KEYS]


@pytest.fixture
def export_data(db):
    project = make_project(db, _PROJECT)
    batch = SysImportBatch(
        filename="ascending-synthetic.xlsx", file_type="maintenance",
        file_hash="ascending-synthetic",
    )
    db.add(batch)
    db.flush()
    orders = {}
    # 插入序刻意与日期/单号排序不同；同日同号也必须按单头身份分组。
    for key, on, number in [
        ("late", date(2026, 7, 20), "WBDD-20260720-0001"),
        ("undated", None, "WBDD-20260730-0001"),
        ("same-b", date(2026, 7, 10), "WBDD-20260710-0002"),
        ("same-a", date(2026, 7, 10), "WBDD-20260710-0001"),
        ("same-a-copy", date(2026, 7, 10), "WBDD-20260710-0001"),
        ("early", date(2026, 7, 1), "WBDD-20260701-0001"),
    ]:
        order = FMaintenanceOrder(
            raw_order_id=f"asc-{key}", order_no=number, order_date=on,
            linked_sales_order_no=_CONTRACT, project_std=_PROJECT,
            project_raw=_PROJECT, data_status="已生效", import_batch_id=batch.id,
        )
        db.add(order)
        db.flush()
        orders[key] = order
    # 跨单交错插入，使单号打平时仅按行号/行 ID 的错误排序必定交织。
    for key, line_no in [
        ("two", 2), ("null-a", None), ("one-a", 1),
        ("one-b", 1), ("null-b", None),
    ]:
        for order_key, order in orders.items():
            pn = f"PN-{order_key}-{key}"
            part = DimPart(pn_std=pn)
            db.add(part)
            db.flush()
            db.add(FMaintenanceLine(
                raw_line_id=f"asc-{order_key}-{key}", order_id=order.id,
                line_no=line_no, part_id=part.id, pn_std=pn, pn_raw=pn,
                qty=Decimal("1"), return_qty=Decimal("0"),
                cost_source="none",
                import_batch_id=batch.id,
            ))
    db.commit()
    for order in orders.values():
        assign(db, order, project)
    return project, orders


def _xlsx_rows(payload, sheet):
    workbook = load_workbook(io.BytesIO(payload), data_only=True)
    try:
        rows = iter(workbook[sheet].values)
        headers = next(rows)
        return [dict(zip(headers, row)) for row in rows]
    finally:
        workbook.close()


def test_order_download_sorts_headers_and_lines_with_nulls_last(db, export_data):
    _project, orders = export_data
    client = client_for(db, username="asc-admin", role="admin")
    response = client.get("/api/maintenance/orders/export")
    assert response.status_code == 200, response.text
    headers = _xlsx_rows(response.content, "维保订单")
    assert [row["原始订单ID"] for row in headers] == [
        orders[key].raw_order_id for key in _ORDER_KEYS
    ]
    rows = _xlsx_rows(response.content, "订单明细")
    assert [row["PN"] for row in rows] == _expected_pns()


@pytest.mark.parametrize("v2", [False, True], ids=["v1", "v2"])
@pytest.mark.parametrize("single_sheet", [False, True], ids=["full", "03-only"])
def test_master_download_ascending_keeps_web_order(
    db, export_data, monkeypatch, v2, single_sheet,
):
    project, _orders = export_data
    monkeypatch.setattr(
        master_api.get_settings(), "maintenance_project_master_v2_enabled", v2,
    )
    client = client_for(db, username="asc-master-admin", role="admin")
    base = f"/api/maintenance/projects/stable/{project.project_id}/master-workbook"
    web_params = {"sheet": master.SHEET_PARTS, "page": 1, "page_size": 10}
    before = client.get(base + "/rows", params=web_params)
    assert before.status_code == 200, before.text
    # 既有网页：DESC 未指定 NULLS LAST，空日期仍排最前，其次最新日期。
    assert before.json()["rows"][0]["pn_std"] == "PN-undated-one-a"
    params = {"sheets": master.SHEET_PARTS} if single_sheet else {}
    response = client.get(base + ".xlsx", params=params)
    assert response.status_code == 200, response.text
    sheet = master.V2_SHEET_PARTS if v2 else master.SHEET_PARTS
    rows = _xlsx_rows(response.content, sheet)
    # V2 表尾有带示例 PN 的 CREATE 行，只核对带实体身份的已存在明细。
    exported = [row for row in rows if row.get("实体ID")] if v2 else rows
    assert [row["PN"] for row in exported] == _expected_pns()
    after = client.get(base + "/rows", params=web_params)
    assert after.json() == before.json()


def test_global_download_ascending_preserves_date_window(db, export_data):
    client = client_for(db, username="asc-global-admin", role="admin")
    response = client.get("/api/maintenance/spare-part-lines.xlsx", params={
        "range": "custom", "from": "2026-07-01", "to": "2026-07-31",
    })
    assert response.status_code == 200, response.text
    rows = _xlsx_rows(response.content, master.GLOBAL_SHEET)
    # 此入口必选日期区间，原有筛选不包含 NULL 日期，不新增 all 范围。
    assert [row["PN"] for row in rows] == _expected_pns(dated_only=True)


@pytest.mark.parametrize("endpoint", ["lines/export", "export-workbook", "export-workbooks"])
def test_legacy_financial_downloads_keep_orders_contiguous(db, export_data, endpoint):
    client = client_for(db, username="asc-financial-admin", role="admin")
    params = ({"project": _PROJECT} if endpoint == "lines/export" else
              {"contract": _CONTRACT} if endpoint == "export-workbook" else {})
    response = client.get(f"/api/maintenance/{endpoint}", params=params)
    assert response.status_code == 200, response.text
    if endpoint == "lines/export":
        rows = list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))
        pns = [row["PN"] for row in rows]
    else:
        payload = response.content
        if endpoint == "export-workbooks":
            with ZipFile(io.BytesIO(payload)) as archive:
                books = [name for name in archive.namelist() if name.endswith(".xlsx")]
                assert len(books) == 1
                payload = archive.read(books[0])
        rows = _xlsx_rows(payload, "备件明细-氚云")
        pns = [row["需供货产品"] for row in rows]
    # 成本起算日条件在排序修改前就排除 NULL 日期，保持原有集合。
    assert pns == _expected_pns(dated_only=True)


def test_v3_download_uses_complete_order_and_line_keys(db, export_data):
    project, _orders = export_data
    client = client_for(db, username="asc-v3-admin", role="admin")
    response = client.get(
        f"/api/maintenance/projects/stable/{project.project_id}/workbook-v3.xlsx",
    )
    assert response.status_code == 200, response.text
    rows = _xlsx_rows(response.content, "03_备件订单")
    assert [row["PN"] for row in rows if row["PN"]] == _expected_pns()


@pytest.mark.parametrize("bundle", [False, True], ids=["single", "zip"])
def test_roundtrip_download_ascending_including_manual_cost_sheet(db, export_data, bundle):
    _project, orders = export_data
    client = client_for(db, username="asc-roundtrip-admin", role="admin")
    endpoint = "roundtrip-templates" if bundle else "roundtrip-template"
    response = client.get(
        f"/api/maintenance/{endpoint}", params={} if bundle else {"contract": _CONTRACT},
    )
    assert response.status_code == 200, response.text
    payload = response.content
    if bundle:
        with ZipFile(io.BytesIO(payload)) as archive:
            books = [name for name in archive.namelist() if name.endswith(".xlsx")]
            assert len(books) == 1
            payload = archive.read(books[0])
    headers = _xlsx_rows(payload, "02_维保订单")
    assert [row["__entity_id"] for row in headers if row["__entity_id"]] == [
        orders[key].id for key in _ORDER_KEYS
    ]
    for sheet in ("03_订单明细", "05_人工成本回填"):
        rows = _xlsx_rows(payload, sheet)
        assert [row["PN"] for row in rows if row["PN"]] == _expected_pns()


def test_streaming_export_crosses_batches_without_changing_web_pagination(db, export_data):
    before = maintenance_cost.project_lines(db, _PROJECT, page=1, page_size=2)
    assert [row["pn_std"] for row in before["rows"]] == [
        "PN-late-null-b", "PN-late-one-b",
    ]
    rows = list(maintenance_cost.iter_project_lines(db, _PROJECT, yield_per=2))
    assert [row["pn_std"] for row in rows] == _expected_pns(dated_only=True)
    assert maintenance_cost.project_lines(db, _PROJECT, page=1, page_size=2) == before


@pytest.mark.parametrize("v2", [False, True], ids=["v1", "v2"])
def test_sorted_and_reordered_master_files_validate_without_changes(db, export_data, v2):
    project, _orders = export_data
    if v2:
        payload = master.build_project_master_v2(
            db, project_id=project.project_id, sheets=(master.V2_SHEET_PARTS,),
        )
        sheet = master.V2_SHEET_PARTS
    else:
        payload = master.build_project_master(
            db, project_id=project.project_id, sheets=(master.SHEET_PARTS,),
        )
        sheet = master.SHEET_PARTS
    # 新升序与重新排列的在途文件都按实体和签名定位，不以 Excel 行位定位。
    for reordered in (False, True):
        workbook = load_workbook(io.BytesIO(payload))
        try:
            if reordered:
                ws = workbook[sheet]
                existing = list(ws.iter_rows(min_row=2, max_row=31, values_only=True))
                for row_no, values in enumerate(reversed(existing), 2):
                    for column, value in enumerate(values, 1):
                        ws.cell(row_no, column).value = value
            buffer = io.BytesIO()
            workbook.save(buffer)
        finally:
            workbook.close()
        if v2:
            plan = master.validate_project_master_v2(
                db, project_id=project.project_id, data=buffer.getvalue(),
            )
            assert plan.field_changes == ()
            assert plan.conflicts == ()
            assert plan.will_void_rows == ()
        else:
            plan = master.validate(db, project_id=project.project_id, data=buffer.getvalue())
        assert plan.cost_refills == ()


@pytest.mark.parametrize("bundle", [False, True], ids=["single", "zip"])
def test_contract_download_merges_headless_orders_by_business_date(db, export_data, bundle):
    _project, orders = export_data
    for number, on in [
        ("WBDD-20260715-0001", date(2026, 7, 15)),
        ("WBDD-20260705-0001", date(2026, 7, 5)),
    ]:
        db.add(FMaintenanceOrder(
            raw_order_id=f"asc-headless-{number}", order_no=number, order_date=on,
            linked_sales_order_no=_CONTRACT, project_std=_PROJECT,
            project_raw=_PROJECT, data_status="已生效",
            import_batch_id=orders["early"].import_batch_id,
        ))
    db.commit()
    client = client_for(db, username="asc-headless-admin", role="admin")
    endpoint = "export-workbooks" if bundle else "export-workbook"
    response = client.get(
        f"/api/maintenance/{endpoint}", params={} if bundle else {"contract": _CONTRACT},
    )
    assert response.status_code == 200, response.text
    payload = response.content
    if bundle:
        with ZipFile(io.BytesIO(payload)) as archive:
            books = [name for name in archive.namelist() if name.endswith(".xlsx")]
            assert len(books) == 1
            payload = archive.read(books[0])
    rows = _xlsx_rows(payload, "备件明细-氚云")
    assert [row["制单日期"] for row in rows] == (
        ["2026-07-01"] * 5 + ["2026-07-05"] + ["2026-07-10"] * 15
        + ["2026-07-15"] + ["2026-07-20"] * 5
    )
    assert [row["数据标题(WBDD单号)"] for row in rows
            if row["产品描述"] == "暂无配件明细"] == [
        "WBDD-20260705-0001", "WBDD-20260715-0001",
    ]
    assert [row["需供货产品"] for row in rows if row["需供货产品"]] == _expected_pns(
        dated_only=True,
    )
