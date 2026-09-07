"""05_实收回款经项目总表（V2）走共享写路径的契约：作废优先、重复月份、状态列。

复现自评审探针（V2P1 / V2P2 / V2P3）。V2 解析器（高风险模块）不判同合同同月重复、
不校验状态列、把落在已作废月份上的新行当 UPDATE 送来；守卫全部落在共享写路径
``maintenance_expense_collection_workbook.apply``，总表只把 05 的行级回执并进
``voided_rows``。这里只从总表 HTTP 入口证明，不测解析器内部。
"""
import io
import uuid
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook
from sqlalchemy import select

from app.api import maintenance_project_master_workbook as master_api
from app.models.maintenance_project_operations import (
    MaintenanceCollectionSnapshot,
    MaintenanceProjectOperationAudit,
)
from app.services import maintenance_project_master_workbook as master
from app.services import maintenance_project_operations as operations

from tests.test_maintenance_expense_collection_workbook import _BASE, _XLSX, _client
from tests.test_maintenance_project_master_v2_editable import (  # noqa: F401
    _project_with_receipt, _save,
)


def _v2_client(db, monkeypatch):
    """V2 开关打开 + 能改总表的实名账号；项目须先建好（客户端建立时挂可见范围）。"""
    monkeypatch.setattr(
        master_api.get_settings(), "maintenance_project_master_v2_enabled", True)
    return _client(db, username="v2-receipts-editor", overrides={
        "page_maintenance": True,
        "data_profit": True,
        "action_maintenance_expense_collection_upload": True,
        "action_maintenance_project_manage": True,
    })


def _receipt_sheet(client, project):
    resp = client.get(f"{_BASE}/{project.project_id}/master-workbook.xlsx",
                      params={"sheets": master.V2_SHEET_RECEIPTS})
    assert resp.status_code == 200, resp.text
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb[master.V2_SHEET_RECEIPTS]
    headers = {cell.value: cell.column for cell in ws[1]}
    return wb, ws, headers


def _add_row(ws, headers, contract_no, month_text, amount, *, status=None):
    row = ws.max_row + 1
    ws.cell(row, headers["合同编号"], contract_no)
    ws.cell(row, headers["报告月份"], month_text)
    ws.cell(row, headers["累计实收金额（含税）"], amount)
    if status is not None:
        ws.cell(row, headers["状态"], status)


def _apply(client, project, wb):
    return client.post(
        f"{_BASE}/{project.project_id}/master-workbook/apply",
        files={"file": ("m.xlsx", io.BytesIO(_save(wb)), _XLSX)})


def _seed_month(db, project, contract, month: date, amount, *, status,
                voided_by: str | None = None) -> MaintenanceCollectionSnapshot:
    row = MaintenanceCollectionSnapshot(
        collection_id=str(uuid.uuid4()), project_id=project.project_id,
        project_contract_id=contract.project_contract_id,
        report_month=month, cumulative_amount=Decimal(amount),
        status=status, source="direct_api", version=3)
    db.add(row)
    if voided_by is not None:
        operations._fact_audit(
            db, project_id=project.project_id, entity_type="collection",
            entity_id=row.collection_id, action="void",
            before={"status": "confirmed"}, after={"status": "void"},
            reason="面板作废", operated_by=voided_by)
    db.commit()
    return row


def _rows(db):
    return sorted(
        (r.report_month, r.status, r.cumulative_amount, r.version)
        for r in db.scalars(select(MaintenanceCollectionSnapshot)))


def _audits(db):
    return [(a.action, a.entity_id) for a in db.scalars(
        select(MaintenanceProjectOperationAudit)
        .where(MaintenanceProjectOperationAudit.entity_type == "collection")
        .order_by(MaintenanceProjectOperationAudit.id))]


def test_v2_fresh_row_for_a_voided_month_is_a_row_receipt_not_a_revival(db, monkeypatch):
    """红线：05 导出只带 confirmed 行，7 月已作废对上传者不可见；补一行 7 月以前被
    解析成 UPDATE 并把作废月复活成 confirmed（版本 3→4、审计「状态 void→confirmed」）。
    现在 7 月不动、回执点名作废人，其余（6 月未触碰）照常，整次上传 200。"""
    project, contract, june = _project_with_receipt(db)     # 6 月 82325.40 confirmed
    july = _seed_month(db, project, contract, date(2026, 7, 1), "90000.00",
                       status="void", voided_by="李四")
    client = _v2_client(db, monkeypatch)
    wb, ws, headers = _receipt_sheet(client, project)
    exported_months = {str(ws.cell(r, headers["报告月份"]).value or "")
                       for r in range(2, ws.max_row + 1)}
    assert "2026-06" in exported_months and "2026-07" not in exported_months   # 作废月不导出
    _add_row(ws, headers, contract.contract_no, "2026-07", 85000)

    resp = _apply(client, project, wb)

    assert resp.status_code == 200, resp.text
    body = resp.json()
    (receipt,) = body["voided_rows"]
    assert receipt["sheet"] == master.V2_SHEET_RECEIPTS
    assert receipt["entity_id"] == july.collection_id
    assert receipt["code"] == "collection_voided" and receipt["reason"] == "row_voided"
    assert receipt["voided_by"] == "李四"
    assert "XSDD-EDIT-001 2026-07" in receipt["message"]
    assert "已被 李四 于 " in receipt["message"] and "修改未生效" in receipt["message"]
    assert receipt["message"] in body["warnings"]
    db.expire_all()
    assert _rows(db) == [
        (date(2026, 6, 1), "confirmed", Decimal("82325.40"), 1),
        (date(2026, 7, 1), "void", Decimal("90000.00"), 3),
    ]
    assert _audits(db) == [("void", july.collection_id)]     # 只有种数据时那条


def test_v2_two_fresh_rows_for_the_same_month_are_a_422_receipt_not_a_500(db, monkeypatch):
    """红线：复制粘贴出两行 8 月，以前 validate 放行、apply 在终态 flush 撞
    uq_maintenance_collection_contract_month 变成 500，整份文件（含其他表）作废且无回执。"""
    project, contract, _june = _project_with_receipt(db)
    client = _v2_client(db, monkeypatch)
    wb, ws, headers = _receipt_sheet(client, project)
    _add_row(ws, headers, contract.contract_no, "2026-08", 90000)
    _add_row(ws, headers, contract.contract_no, "2026-08", 95000)

    resp = _apply(client, project, wb)

    assert resp.status_code == 422, resp.text
    detail = resp.json()["detail"]
    assert detail["code"] == "duplicate_month"
    assert [(i["contract_no"], i["report_month"], i["code"]) for i in detail["issues"]] == [
        ("XSDD-EDIT-001", "2026-08", "duplicate_month")]
    db.expire_all()
    assert _rows(db) == [(date(2026, 6, 1), "confirmed", Decimal("82325.40"), 1)]
    assert _audits(db) == []


def test_v2_two_changed_rows_with_the_same_entity_id_are_a_422_not_last_wins(db, monkeypatch):
    """同一实体ID出现两行且都改了：以前两条 UPDATE 静默后者覆盖前者、留两条审计。"""
    project, contract, june = _project_with_receipt(db)
    client = _v2_client(db, monkeypatch)
    wb, ws, headers = _receipt_sheet(client, project)
    source = next(r for r in range(2, ws.max_row + 1)
                  if ws.cell(r, headers["实体ID"]).value == june.collection_id)
    target = ws.max_row + 1
    for col in range(1, ws.max_column + 1):
        ws.cell(target, col, ws.cell(source, col).value)
    ws.cell(source, headers["累计实收金额（含税）"], 85000)
    ws.cell(target, headers["累计实收金额（含税）"], 88000)

    resp = _apply(client, project, wb)

    assert resp.status_code == 422, resp.text
    detail = resp.json()["detail"]
    assert detail["code"] == "duplicate_month"
    assert [i["report_month"] for i in detail["issues"]] == ["2026-06"]
    db.expire_all()
    assert _rows(db) == [(date(2026, 6, 1), "confirmed", Decimal("82325.40"), 1)]
    assert _audits(db) == []


def test_v2_free_text_status_is_a_422_receipt_not_a_500(db, monkeypatch):
    """红线：状态列写「已确认」以前撞 ck_maintenance_collection_status 变 500。"""
    project, contract, _june = _project_with_receipt(db)
    client = _v2_client(db, monkeypatch)
    wb, ws, headers = _receipt_sheet(client, project)
    _add_row(ws, headers, contract.contract_no, "2026-08", 90000, status="已确认")

    resp = _apply(client, project, wb)

    assert resp.status_code == 422, resp.text
    detail = resp.json()["detail"]
    assert detail["code"] == "invalid_status"
    assert [(i["report_month"], i["code"]) for i in detail["issues"]] == [
        ("2026-08", "invalid_status")]
    assert "confirmed、unconfirmed 或 void" in detail["message"]
    db.expire_all()
    assert _rows(db) == [(date(2026, 6, 1), "confirmed", Decimal("82325.40"), 1)]
    assert _audits(db) == []
