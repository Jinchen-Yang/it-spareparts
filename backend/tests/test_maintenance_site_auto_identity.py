"""Unnumbered consumption preserves occurrences, actual dates and replay identity."""
import io
import hashlib
import uuid
from datetime import date

import pytest
from openpyxl import load_workbook
from openpyxl.utils.datetime import MAC_EPOCH, WINDOWS_EPOCH, to_excel
from sqlalchemy import select

from app.models.maintenance_project_operations import MaintenanceSiteIssue, MaintenanceSiteIssueLine
from app.services import maintenance_project_master_workbook as master
from tests.test_maintenance_project_master_v2_editable import _make_project_with_line, _save, _site_issue
from tests.test_maintenance_master_v2_row_merge import _apply


def _book(db, project_id):
    return load_workbook(io.BytesIO(master.build_project_master_v2(
        db, project_id=project_id, sheets=(master.V2_SHEET_SITE,))))


def _append(wb, pn, *, no=None, when="2025-04-21", qty=1):
    wb[master.V2_SHEET_SITE].append([no, when, pn, None, qty])
    return wb[master.V2_SHEET_SITE].max_row


def _lines(db, project_id):
    return db.execute(select(MaintenanceSiteIssue, MaintenanceSiteIssueLine).join(
        MaintenanceSiteIssueLine, MaintenanceSiteIssueLine.issue_id == MaintenanceSiteIssue.issue_id
    ).where(MaintenanceSiteIssue.project_id == project_id)).all()


def test_auto_identity_preserves_identical_occurrences_and_sorted_replay(db):
    project, part, *_ = _make_project_with_line(db)
    wb = _book(db, project.project_id)
    for qty in (1, 1, 2):
        _append(wb, part.pn_std, qty=qty)
    _, result = _apply(db, project.project_id, wb)
    assert result["site_creates"] == 3
    before = {(i.issue_no, l.issue_line_id, l.quantity) for i, l in _lines(db, project.project_id)}
    assert len(before) == 3
    ws = wb[master.V2_SHEET_SITE]
    rows = list(ws.iter_rows(min_row=3, values_only=True))
    for row_no, values in enumerate(reversed(rows), 3):
        for col, value in enumerate(values, 1):
            ws.cell(row_no, col).value = value
    plan, result = _apply(db, project.project_id, wb)
    assert not plan.site_flags
    assert result["site_creates"] == 0
    assert {(i.issue_no, l.issue_line_id, l.quantity) for i, l in _lines(db, project.project_id)} == before


def test_wbdd_reference_keeps_actual_date_and_old_template_works(db):
    project, part, *_ = _make_project_with_line(db)
    wb = _book(db, project.project_id)
    # Existing customer files end at the baseline column.
    wb[master.V2_SHEET_SITE].delete_cols(13)
    for when in ("2025-04-21", "2025-04-22", "2025-04-22"):
        _append(wb, part.pn_std, no="WBDD-20250506-0010", when=when)
    _apply(db, project.project_id, wb)
    rows = _lines(db, project.project_id)
    assert len(rows) == 3
    assert sorted(i.issue_date for i, _ in rows) == [date(2025, 4, 21), date(2025, 4, 22), date(2025, 4, 22)]
    assert all(l.demand_order_no == "WBDD-20250506-0010" for _, l in rows)
    assert all(i.issue_no.startswith("LY-AUTO-") for i, _ in rows)
    downloaded = _book(db, project.project_id)
    headers = {c.value: c.column for c in downloaded[master.V2_SHEET_SITE][1]}
    assert not downloaded[master.V2_SHEET_SITE].column_dimensions["M"].hidden
    assert downloaded[master.V2_SHEET_SITE].cell(2, headers["关联需求单号"]).value == "WBDD-20250506-0010"
    plan, _ = _apply(db, project.project_id, downloaded)
    assert not plan.site_flags


def test_existing_number_does_not_override_explicit_date(db):
    project, part, *_ = _make_project_with_line(db)
    wb = _book(db, project.project_id)
    _append(wb, part.pn_std, no="LY-20250506-0001", when="2025-04-21")
    _apply(db, project.project_id, wb)
    assert _lines(db, project.project_id)[0][0].issue_date == date(2025, 4, 21)


@pytest.mark.parametrize("epoch", [WINDOWS_EPOCH, MAC_EPOCH])
def test_date_serial_without_cell_style_honors_workbook_epoch(db, epoch):
    project, part, *_ = _make_project_with_line(db)
    wb = _book(db, project.project_id)
    wb.epoch = epoch
    _append(wb, part.pn_std, when=to_excel(date(2025, 5, 1), epoch=epoch))
    _apply(db, project.project_id, wb)
    assert _lines(db, project.project_id)[0][0].issue_date == date(2025, 5, 1)


@pytest.mark.parametrize("value", [True, -1, 0.5, float("inf"), float("nan"), 99999999])
def test_invalid_date_serials_fail_closed(value):
    with pytest.raises(master.WorkbookError):
        master._v2_date(value, row_no=3, label="领用")


def test_preflight_reports_dates_quantities_and_unknown_pns_together(db):
    project, part, *_ = _make_project_with_line(db)
    wb = _book(db, project.project_id)
    first = _append(wb, "UNKNOWN-PN-A", when="4月21", qty=-1)
    second = _append(wb, "UNKNOWN-PN-B", when=None)
    third = _append(wb, part.pn_std, when="2025-04")
    with pytest.raises(master.WorkbookError) as exc:
        master.validate_project_master_v2(db, project_id=project.project_id, data=_save(wb))
    assert exc.value.code == "invalid_site_rows"
    for message in (f"第 {first} 行", f"第 {second} 行", f"第 {third} 行", "UNKNOWN-PN-A", "UNKNOWN-PN-B", "领用数量超出允许范围"):
        assert message in exc.value.message
    assert not _lines(db, project.project_id)


def test_downloaded_auto_row_can_be_edited_and_original_replay_cannot_undo_it(db):
    project, part, *_ = _make_project_with_line(db)
    original = _book(db, project.project_id)
    _append(original, part.pn_std)
    _apply(db, project.project_id, original)
    updated = _book(db, project.project_id)
    updated[master.V2_SHEET_SITE].cell(2, 5, 3)
    _apply(db, project.project_id, updated)
    assert len(_lines(db, project.project_id)) == 1
    assert _lines(db, project.project_id)[0][1].quantity == 3
    with pytest.raises(master.WorkbookError) as exc:
        _apply(db, project.project_id, original)
    assert exc.value.code == "stale_site_replay"
    assert _lines(db, project.project_id)[0][1].quantity == 3


def test_editing_original_unnumbered_row_requires_download_instead_of_duplicating(db):
    project, part, *_ = _make_project_with_line(db)
    original = _book(db, project.project_id)
    row_no = _append(original, part.pn_std)
    _apply(db, project.project_id, original)
    original[master.V2_SHEET_SITE].cell(row_no, 5, 9)
    with pytest.raises(master.WorkbookError, match="重新下载"):
        _apply(db, project.project_id, original)
    rows = _lines(db, project.project_id)
    assert len(rows) == 1 and rows[0][1].quantity == 1


def test_auto_scope_is_rechecked_between_validation_and_apply(db):
    project, part, *_ = _make_project_with_line(db)
    original = _book(db, project.project_id)
    row_no = _append(original, part.pn_std)
    first = master.validate_project_master_v2(db, project_id=project.project_id, data=_save(original))
    original[master.V2_SHEET_SITE].cell(row_no, 5, 9)
    second = master.validate_project_master_v2(db, project_id=project.project_id, data=_save(original))
    master.apply_project_master_v2(db, first, operated_by="test", import_batch_id=str(uuid.uuid4()))
    with pytest.raises(master.WorkbookError, match="重新下载"):
        master.apply_project_master_v2(db, second, operated_by="test", import_batch_id=str(uuid.uuid4()))
    rows = _lines(db, project.project_id)
    assert len(rows) == 1 and rows[0][1].quantity == 1


def test_auto_row_deleted_then_original_replayed_stays_void(db):
    project, part, *_ = _make_project_with_line(db)
    original = _book(db, project.project_id)
    _append(original, part.pn_std)
    _apply(db, project.project_id, original)
    edited = _book(db, project.project_id)
    edited[master.V2_SHEET_SITE].delete_rows(2)
    _apply(db, project.project_id, edited)
    plan, _ = _apply(db, project.project_id, original)
    assert plan.voided_rows
    rows = _lines(db, project.project_id)
    assert len(rows) == 1 and not rows[0][1].is_active


def test_concurrent_edit_rebases_untouched_auto_row_fields(db):
    project, part, *_ = _make_project_with_line(db)
    wb = _book(db, project.project_id)
    _append(wb, part.pn_std)
    _apply(db, project.project_id, wb)
    first = _book(db, project.project_id)
    second = _book(db, project.project_id)
    first[master.V2_SHEET_SITE].cell(2, 5, 4)
    _apply(db, project.project_id, first)
    second[master.V2_SHEET_SITE].cell(2, 7, "补充现场备注")
    _apply(db, project.project_id, second)
    line = _lines(db, project.project_id)[0][1]
    assert line.quantity == 4 and line.remark == "补充现场备注"


def test_legacy_wbdd_replay_does_not_duplicate_or_silently_rewrite_date(db):
    project, part, *_ = _make_project_with_line(db)
    wb = _book(db, project.project_id)
    no = "WBDD-20250506-0010"
    issue, line = _site_issue(db, project, part, qty="1", issue_no=no)
    identity = "|".join([project.project_id, no, part.pn_std, ""])
    line.issue_line_id = "manual-site:" + hashlib.sha1(identity.encode()).hexdigest()
    issue.issue_date = date(2025, 5, 6)
    db.commit()
    row_no = _append(wb, part.pn_std, no=no, when="2025-05-06")
    plan, _ = _apply(db, project.project_id, wb)
    assert not plan.site_flags and len(_lines(db, project.project_id)) == 1
    wb[master.V2_SHEET_SITE].cell(row_no, 2, "2025-04-21")
    with pytest.raises(master.WorkbookError, match="保留实体ID"):
        _apply(db, project.project_id, wb)
    assert issue.issue_date == date(2025, 5, 6)
