"""总表校验错误保持 HTTP 契约，且失败访问审计不受请求回滚影响。"""

import hashlib
import io
import json
from datetime import date

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select

from app.api import maintenance_project_master_workbook as workbook_api
from app.models.system import SysAccessLog
from app.services import maintenance_expense_collection_workbook as ec
from app.services import maintenance_project_master_workbook as master
from tests.test_maintenance_expense_collection_workbook import _XLSX, uploader
from tests.test_maintenance_project_master_v2_editable import (
    _make_project_with_line,
    _save,
    _site_issue,
)


@pytest.fixture(autouse=True)
def _enable_v2(monkeypatch):
    monkeypatch.setattr(
        workbook_api.get_settings(), "maintenance_project_master_v2_enabled", True,
    )
    monkeypatch.setattr(workbook_api.config, "ENABLE_ACCESS_LOG", True)


def _post(client, project_id, content):
    return client.post(
        f"/api/maintenance/projects/stable/{project_id}/master-workbook/validate",
        files={"file": ("synthetic-collisions.xlsx", content, _XLSX)},
    )


def _validate_logs(db):
    return list(db.scalars(select(SysAccessLog).where(
        SysAccessLog.action == "workbook_validate",
        SysAccessLog.resource == "maintenance_project_master_workbook",
    ).order_by(SysAccessLog.id)))


def _collisions(db):
    project, part, *_ = _make_project_with_line(db)
    entity_ids = []
    for number in range(1, 4):
        issue_no = f"WBDD-20260810-{number:04d}"
        issue, line = _site_issue(db, project, part, qty="1", issue_no=issue_no)
        identity = "|".join([project.project_id, issue_no, part.pn_std, ""])
        line.issue_line_id = "manual-site:" + hashlib.sha1(identity.encode()).hexdigest()
        issue.site_location = "合成领用地址"
        entity_ids.append(line.issue_line_id)
    db.commit()
    book = load_workbook(io.BytesIO(master.build_project_master_v2(
        db, project_id=project.project_id, sheets=(master.V2_SHEET_SITE,),
    )))
    sheet = book[master.V2_SHEET_SITE]
    columns = {cell.value: cell.column for cell in sheet[1]}
    rows = []
    for number, entity_id in enumerate(entity_ids):
        original = next(
            row for row in sheet.iter_rows(min_row=2, values_only=True)
            if row[columns["实体ID"] - 1] == entity_id
        )
        values = list(original)
        for field in ("实体ID", "基线令牌", "应返数量", "返还状态", "返还单号", "关联需求单号"):
            values[columns[field] - 1] = ""
        changes = (
            {"领用日期": date(2026, 9, 1)},
            {"是否应返还": "是"},
            {"领用数量": 3, "备注": "合成上传备注", "领用地址": "合成新地址"},
        )[number]
        for field, value in changes.items():
            values[columns[field] - 1] = value
        sheet.append(values)
        rows.append(sheet.max_row)
    content = _save(book)
    book.close()
    db.commit()
    return project, entity_ids, rows, content


def test_stale_replay_error_exposes_only_supported_details():
    rows = [{"row_no": 8, "entity_id": "synthetic-id", "differences": []}]
    error = ec.WorkbookError("stale_site_replay", "合成冲突", details={
        "rows": rows, "total_count": 201, "fix_hint": "保留实体ID",
        "code": "must-not-override", "private": "must-not-leak",
    })
    with pytest.raises(HTTPException) as raised:
        workbook_api._fail(error)
    assert raised.value.status_code == 422
    assert raised.value.detail == {
        "code": "stale_site_replay", "message": "合成冲突", "issues": [],
        "rows": rows, "total_count": 201, "fix_hint": "保留实体ID",
    }


def test_existing_plain_and_three_way_errors_keep_their_shape():
    with pytest.raises(HTTPException) as plain:
        workbook_api._fail(ec.WorkbookError("stale_site_replay", "旧式重放提示"))
    assert plain.value.status_code == 422
    assert plain.value.detail == {
        "code": "stale_site_replay", "message": "旧式重放提示", "issues": [],
    }
    issues = [{"entity_id": "synthetic-id", "field": "备注", "current": "新值"}]
    with pytest.raises(HTTPException) as conflict:
        workbook_api._fail(ec.WorkbookError("row_conflicts", "三路合并冲突", issues))
    assert conflict.value.status_code == 409
    assert conflict.value.detail["issues"] == conflict.value.detail["conflicts"] == issues


def test_http_multi_collision_returns_all_rows_and_persists_failure_audit(db):
    project, entity_ids, row_numbers, content = _collisions(db)
    client = uploader(db, username="synthetic-conflict-uploader")

    response = _post(client, project.project_id, content)

    assert response.status_code == 422, response.text
    detail = response.json()["detail"]
    assert detail["code"] == "stale_site_replay"
    assert "共 3 行" in detail["message"]
    assert detail["total_count"] == 3
    assert [row["row_no"] for row in detail["rows"]] == row_numbers
    assert [row["entity_id"] for row in detail["rows"]] == entity_ids
    assert detail["rows"][0]["differences"] == [
        {"field": "领用日期", "old": "2026-08-10", "new": "2026-09-01"},
    ]
    assert detail["rows"][1]["differences"] == [
        {"field": "是否应返还", "old": "", "new": "是"},
    ]
    assert {item["field"] for item in detail["rows"][2]["differences"]} == {
        "领用数量", "备注", "领用地址",
    }
    assert "保留实体ID" in detail["fix_hint"]
    logs = _validate_logs(db)
    assert len(logs) == 1
    assert logs[0].detail == {
        "project_id": project.project_id,
        "error_code": "stale_site_replay",
        "conflict_count": 3,
    }
    json.dumps(logs[0].detail)


def test_validation_failure_audit_uses_total_count_instead_of_capped_rows(db, monkeypatch):
    project, *_ = _make_project_with_line(db)
    client = uploader(db, username="synthetic-capped-error")
    content = master.build_project_master_v2(db, project_id=project.project_id)
    db.commit()

    def fail(*_args, **_kwargs):
        raise ec.WorkbookError("stale_site_replay", "共 203 行冲突", details={
            "rows": [{"row_no": row} for row in range(2, 202)],
            "total_count": 203, "fix_hint": "保留实体ID",
        })

    monkeypatch.setattr(master, "validate_project_master_v2", fail)
    response = _post(client, project.project_id, content)
    assert response.status_code == 422, response.text
    assert len(response.json()["detail"]["rows"]) == 200
    assert _validate_logs(db)[0].detail["conflict_count"] == 203


def test_audit_failure_never_replaces_workbook_validation_error(db, monkeypatch):
    project, _, _, content = _collisions(db)
    client = uploader(db, username="synthetic-audit-fault")

    def fail_audit(*_args, **_kwargs):
        raise RuntimeError("合成审计连接故障")

    monkeypatch.setattr(workbook_api, "record_access_log", fail_audit)
    response = _post(client, project.project_id, content)
    assert response.status_code == 422, response.text
    assert response.json()["detail"]["code"] == "stale_site_replay"
    assert response.json()["detail"]["total_count"] == 3


@pytest.mark.parametrize("code,status_code,conflict_count", [
    ("row_conflicts", 409, 1),
    ("invalid_site_rows", 422, 0),
])
def test_plain_validation_errors_also_have_failure_audits(
    db, monkeypatch, code, status_code, conflict_count,
):
    project, *_ = _make_project_with_line(db)
    client = uploader(db, username="synthetic-plain-error")
    content = master.build_project_master_v2(db, project_id=project.project_id)
    db.commit()

    def fail(*_args, **_kwargs):
        raise ec.WorkbookError(code, "合成旧式校验错误", [{"row": 2}])

    monkeypatch.setattr(master, "validate_project_master_v2", fail)
    response = _post(client, project.project_id, content)
    assert response.status_code == status_code, response.text
    assert [log.detail for log in _validate_logs(db)] == [{
        "project_id": project.project_id, "error_code": code,
        "conflict_count": conflict_count,
    }]


def test_successful_validation_access_log_is_unchanged(db):
    project, *_ = _make_project_with_line(db)
    client = uploader(db, username="synthetic-success-uploader")
    content = master.build_project_master_v2(db, project_id=project.project_id)
    db.commit()
    response = _post(client, project.project_id, content)
    assert response.status_code == 200, response.text
    assert [log.detail for log in _validate_logs(db)] == [{
        "project_id": project.project_id, "conflicts": 0, "changes": 0,
    }]


def test_out_of_scope_upload_cannot_read_collision_details(db, monkeypatch):
    project, _, _, content = _collisions(db)
    client = uploader(
        db, username="synthetic-unassigned-uploader", assign_existing_projects=False,
    )

    def forbidden_validation(*_args, **_kwargs):
        raise AssertionError("未授权项目不应开始校验或返回碰撞信息")

    monkeypatch.setattr(master, "validate_project_master_v2", forbidden_validation)
    response = _post(client, project.project_id, content)
    assert response.status_code == 403, response.text
    assert "synthetic-id" not in response.text
    assert _validate_logs(db) == []
