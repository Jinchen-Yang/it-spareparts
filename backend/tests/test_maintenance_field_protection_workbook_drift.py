"""D-34 Q3=A（2026-10-21 需求单）：旧工作簿保护基准/周期漂移阻断。

旧簿里用户编辑过的需求事实字段，若其比较基准（本轮首次快照）或保护周期
自导出后已变化（含期间被解除又重新保护、显示值恰好相同的情形），阻断本次
应用并给出字段级提示；不因无关行/未触碰字段的变化作废整本；同一保护周期
内仅最新来源变化不算漂移；应用时在锁内再验一次。不可被强制接管。
"""

import hashlib
import hmac
import json
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.config import get_settings
from app.models.maintenance import FMaintenanceLine
from app.services import maintenance_demand_manual as manual
from app.services import maintenance_project_master_workbook as master
from app.services.maintenance_expense_collection_workbook import WorkbookError
from tests.test_maintenance_field_protection_workbook import reimport
from tests.test_maintenance_project_master_v2_editable import (
    _make_project_with_line,
)


def _export_parts(db, project_id):
    content = master.build_project_master_v2(
        db, project_id=project_id, sheets=(master.V2_SHEET_PARTS,))
    return load_workbook(BytesIO(content))


def _edit_cell(wb, row_no, header, value):
    ws = wb[master.V2_SHEET_PARTS]
    headers = {cell.value: cell.column for cell in ws[1]}
    ws.cell(row_no, headers[header], value)


def _apply_workbook(db, project_id, wb, *, force_takeover=False):
    buf = BytesIO()
    wb.save(buf)
    plan = master.validate_project_master_v2(
        db, project_id=project_id, data=buf.getvalue(),
        force_takeover=force_takeover,
    )
    return master.apply_project_master_v2(
        db, plan, operated_by="drift-test", import_batch_id=str(uuid4()))


def _page_patch(db, line, updates, *, operated_by="other-editor"):
    result = manual.patch_demand_line(
        db, raw_line_id=line.raw_line_id, updates=updates,
        reason="他人在页面修改", operated_by=operated_by)
    db.commit()
    return result


def _add_second_line(db, order, part, *, qty=Decimal(2)):
    extra = FMaintenanceLine(
        raw_line_id=f"raw-line-{uuid4()}", order_id=order.id, line_no=2,
        part_id=part.id, pn_std=part.pn_std, pn_raw=part.pn_std,
        qty=qty, return_qty=Decimal(0), is_active=True,
        import_batch_id=order.import_batch_id,
    )
    db.add(extra)
    db.commit()
    return extra


def _strip_protect_from_token(wb, row_no):
    """把新令牌降级为 2.7.0 旧格式（无 protect 键），模拟升级前的旧簿。"""
    ws = wb[master.V2_SHEET_PARTS]
    headers = {cell.value: cell.column for cell in ws[1]}
    col = headers[master.V2_BASE_COLUMN]
    envelope = json.loads(ws.cell(row_no, col).value)
    payload = envelope["payload"]
    payload.pop("protect", None)
    payload_bytes = json.dumps(
        payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")
    key_id, signing_key = get_settings().maintenance_manifest_signing_material()
    signature = hmac.new(
        signing_key, master._V2_ROW_BASE_DOMAIN + payload_bytes, hashlib.sha256
    ).hexdigest()
    ws.cell(row_no, col, json.dumps(
        {"key_id": key_id, "payload": payload, "signature": signature},
        ensure_ascii=False, sort_keys=True, separators=(",", ":")))


def test_stale_workbook_edit_blocked_after_release_since_export(db):
    """场景1：导出后保护被解除——旧簿编辑该字段被阻断并给字段级提示。"""
    project, _part, _order, line = _make_project_with_line(db)
    _page_patch(db, line, {"qty": 5})  # 保护建立，首次快照 2
    stale = _export_parts(db, project.project_id)
    _page_patch(db, line, {"qty": 2})  # 导出后他人改回原值解除

    _edit_cell(stale, 2, "需求数量", 3)
    with pytest.raises(WorkbookError) as exc_info:
        _apply_workbook(db, project.project_id, stale)
    assert exc_info.value.code == "protection_baseline_drift"
    assert len(exc_info.value.issues) == 1
    issue = exc_info.value.issues[0]
    assert issue["sheet"] == "03_备件明细"
    assert issue["field"] == "需求数量"
    assert issue["reason"] == "protection_released_since_export"
    db.refresh(line)
    assert line.qty == Decimal("2.000")  # 未写入任何数据


def test_stale_workbook_edit_blocked_when_reprotected_with_same_display_value(db):
    """场景2：期间被解除又重新保护、显示值恰好相同——按保护周期变化阻断。"""
    project, _part, _order, line = _make_project_with_line(db)
    _page_patch(db, line, {"qty": 5})   # 保护，首次快照 2，显示 5
    stale = _export_parts(db, project.project_id)
    _page_patch(db, line, {"qty": 2})   # 解除
    _page_patch(db, line, {"qty": 5})   # 重新保护：新周期，首次快照仍 2，显示仍 5

    # 行值三路合并不冲突（server==baseline==5），只有周期校验能拦住。
    _edit_cell(stale, 2, "需求数量", 3)
    with pytest.raises(WorkbookError) as exc_info:
        _apply_workbook(db, project.project_id, stale)
    assert exc_info.value.code == "protection_baseline_drift"
    assert exc_info.value.issues[0]["reason"] == "protection_cycle_changed"

    # 同一份旧簿即使填首次快照 2（想解除）也阻断：周期已换，先重新核对。
    _edit_cell(stale, 2, "需求数量", 2)
    with pytest.raises(WorkbookError) as again:
        _apply_workbook(db, project.project_id, stale)
    assert again.value.code == "protection_baseline_drift"
    assert again.value.issues[0]["reason"] == "protection_cycle_changed"


def test_stale_workbook_edit_blocked_when_protection_started_after_export(db):
    """场景3：导出时不保护、之后他人建立保护——旧簿编辑被阻断（不可接管）。"""
    project, _part, _order, line = _make_project_with_line(db)
    stale = _export_parts(db, project.project_id)  # 导出时 qty=2 无保护
    _page_patch(db, line, {"qty": 5})               # 导出后建立保护

    _edit_cell(stale, 2, "需求数量", 3)
    with pytest.raises(WorkbookError) as exc_info:
        _apply_workbook(db, project.project_id, stale, force_takeover=True)
    assert exc_info.value.code == "protection_baseline_drift"
    assert exc_info.value.issues[0]["reason"] == "protection_started_since_export"
    db.refresh(line)
    assert line.qty == Decimal("5.000")


def test_same_cycle_source_refresh_is_not_drift_and_first_snapshot_release_applies(db):
    """场景4：同一保护周期内仅最新来源变化不算漂移；旧簿按首次快照解除。"""
    project, part, order, line = _make_project_with_line(db)
    _page_patch(db, line, {"qty": 5})   # 保护，首次快照 2
    stale = _export_parts(db, project.project_id)
    reimport(db, order, line, pn=part.pn_std, qty="3")  # 来源刷新：original 仍 2
    db.refresh(line)
    assert line.qty == Decimal("5.000")

    # 旧簿填首次快照 2：无漂移（同周期仅来源变化），按 Q1=A 解除。
    _edit_cell(stale, 2, "需求数量", 2)
    result = _apply_workbook(db, project.project_id, stale)
    db.refresh(line)
    assert result["line_updates"] == 1
    assert line.qty == Decimal("2.000")
    assert line.manual_override == {}
    # 解除后来源上传恢复生效
    reimport(db, order, line, pn=part.pn_std, qty="9")
    db.refresh(line)
    assert line.qty == Decimal("9.000")


def test_unrelated_row_or_untouched_field_changes_do_not_block_stale_workbook(db):
    """场景5：仅因无关行/未触碰字段的变化不拒绝；同簿无漂移编辑照常应用。"""
    project, part, order, line = _make_project_with_line(db)
    other = _add_second_line(db, order, part)
    _page_patch(db, line, {"qty": 5})           # 行 A（第2行）保护，快照 2
    _page_patch(db, other, {"qty": 4})          # 行 B（第3行）也保护，快照 2
    stale = _export_parts(db, project.project_id)
    # 导出后：行 B 被解除+重新保护（周期变化）；行 A 的 note 建立保护。
    _page_patch(db, other, {"qty": 2})
    _page_patch(db, other, {"qty": 4})
    _page_patch(db, line, {"line_note": "他人备注"})

    # 旧簿只把行 A 的 qty 改回首次快照 2：行 A qty 无漂移 → 放行并解除。
    _edit_cell(stale, 2, "需求数量", 2)
    result = _apply_workbook(db, project.project_id, stale)
    db.refresh(line)
    assert result["line_updates"] == 1
    assert line.qty == Decimal("2.000")
    db.expire(line)
    assert "qty" not in (line.manual_override or {})
    assert (line.manual_override or {}).get("line_note", {}).get("value") == "他人备注"
    db.refresh(other)
    assert other.qty == Decimal("4.000")  # 无关行未被触碰、未被改写


def test_drift_between_validate_and_apply_blocks_write(db):
    """场景6：校验完成→实际应用之间基准变化——锁内复核整本拒绝、零写入。

    用「解除后重新保护、显示值恰好相同」构造只动保护、不动行值的窗口变化：
    行指纹不变，只有保护基准复核能拦住。
    """
    project, _part, _order, line = _make_project_with_line(db)
    _page_patch(db, line, {"qty": 5})
    stale = _export_parts(db, project.project_id)
    _edit_cell(stale, 2, "需求数量", 3)

    buf = BytesIO()
    stale.save(buf)
    plan = master.validate_project_master_v2(
        db, project_id=project.project_id, data=buf.getvalue())
    # validate→apply 窗口内：他人解除又重新保护，行值回到 5 与校验时相同。
    _page_patch(db, line, {"qty": 2})
    _page_patch(db, line, {"qty": 5})

    with pytest.raises(WorkbookError) as exc_info:
        master.apply_project_master_v2(
            db, plan, operated_by="drift-test", import_batch_id=str(uuid4()))
    assert exc_info.value.code == "row_conflicts"
    assert exc_info.value.issues[0]["reason"] == "protection_changed_since_validate"
    db.rollback()
    db.refresh(line)
    assert line.qty == Decimal("5.000")


def test_legacy_token_without_protect_payload_blocked_for_protected_field(db):
    """场景7：2.7.0 旧格式令牌（无 protect 键）编辑当前保护字段——阻断重下。"""
    project, _part, _order, line = _make_project_with_line(db)
    _page_patch(db, line, {"qty": 5})
    stale = _export_parts(db, project.project_id)
    _strip_protect_from_token(stale, 2)

    _edit_cell(stale, 2, "需求数量", 3)
    with pytest.raises(WorkbookError) as exc_info:
        _apply_workbook(db, project.project_id, stale)
    assert exc_info.value.code == "protection_baseline_drift"
    assert exc_info.value.issues[0]["reason"] == "protection_evidence_missing"


def test_unprotected_field_edit_in_old_workbook_still_applies(db):
    """对照：旧簿编辑从未受保护的字段（导出前后都无保护）照常生效并建保护。"""
    project, _part, _order, line = _make_project_with_line(db)
    stale = _export_parts(db, project.project_id)
    _edit_cell(stale, 2, "需求数量", 3)
    result = _apply_workbook(db, project.project_id, stale)
    db.refresh(line)
    assert result["line_updates"] == 1
    assert line.qty == Decimal("3.000")
    assert set(line.manual_override) == {"qty"}


# ---------------------------------------------------------------------------
# 返工2（#430 第 2 轮独立验收必须改）：
#   B1/MUST-2 锁内复核只针对本次实际编辑的事实字段；
#   MUST-1 保护基准改定长指纹 + 隐藏令牌整体长度防护，长文本不再撑爆令牌。
# ---------------------------------------------------------------------------


def test_untouched_field_cycle_change_between_validate_and_apply_keeps_write(db):
    """返工2 B1/MUST-2：未编辑字段的保护周期变化不得作废整本应用。

    备注受保护后导出；工作簿只把需求数量 2→3 并 validate；窗口内网页把
    备注 A→空→A（显示值不变、仅备注保护周期变化）；apply 必须照常写入
    数量 3，备注及其保护保留（D-34/0012：不因未触碰字段作废整本）。
    """
    project, _part, _order, line = _make_project_with_line(db)
    _page_patch(db, line, {"line_note": "A"})  # 备注保护：首次快照=空
    wb = _export_parts(db, project.project_id)
    _edit_cell(wb, 2, "需求数量", 3)             # 工作簿只改数量

    buf = BytesIO()
    wb.save(buf)
    plan = master.validate_project_master_v2(
        db, project_id=project.project_id, data=buf.getvalue())
    # validate→apply 窗口内：备注解除（A→空）又重新保护（空→A，新周期）。
    _page_patch(db, line, {"line_note": None})
    _page_patch(db, line, {"line_note": "A"})

    result = master.apply_project_master_v2(
        db, plan, operated_by="drift-test", import_batch_id=str(uuid4()))
    db.refresh(line)
    assert result["line_updates"] == 1
    assert line.qty == Decimal("3.000")          # 数量编辑照常写入
    assert line.line_note == "A"                 # 备注显示值不变
    ledger = line.manual_override or {}
    assert "line_note" in ledger                 # 备注保护保留（新周期）
    assert "qty" in ledger                       # 数量编辑建保护，快照 2


def test_long_protected_description_roundtrip_fresh_workbook(db):
    """返工2 MUST-1：长文本保护不能把基线令牌撑爆 Excel 单元格上限。

    描述原值 17000 字、网页改成另一段 17000 字（建立保护）后导出的工作簿，
    原样上传 validate 必须成功（旧实现把首次快照原文放进签名令牌 →
    超 32767 被 openpyxl 静默截断 → invalid_concurrency_token），apply 零
    写入、描述与保护保持不变。
    """
    project, _part, _order, line = _make_project_with_line(db)
    original = "原" * 17000
    replaced = "改" * 17000
    line.description = original
    db.commit()
    _page_patch(db, line, {"description": replaced})

    wb = _export_parts(db, project.project_id)
    ws = wb[master.V2_SHEET_PARTS]
    headers = {cell.value: cell.column for cell in ws[1]}
    token = str(ws.cell(2, headers[master.V2_BASE_COLUMN]).value or "")
    assert len(token) <= 32767, (
        f"基线令牌长度 {len(token)} 超过 Excel 单元格上限，导出即损坏")

    result = _apply_workbook(db, project.project_id, wb)
    db.refresh(line)
    assert result["line_updates"] == 0            # 原样上传零写入
    assert line.description == replaced
    assert "description" in (line.manual_override or {})


def test_row_base_token_over_cell_limit_raises_instead_of_silent_truncation():
    """返工2 MUST-1：令牌整体长度防护——超 32767 显式报错，绝不静默截断。"""
    with pytest.raises(WorkbookError) as exc_info:
        master._v2_row_base_token(
            sheet="03_备件明细", entity_id="ent-long",
            values={"描述": "长" * 40000})
    assert exc_info.value.code == "concurrency_token_too_long"
    assert "32767" in exc_info.value.message
