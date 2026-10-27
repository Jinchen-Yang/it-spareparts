"""T18：业务预览定位、保存点零写入及应用时重检的回执契约。"""
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest
from sqlalchemy import select

from app import permissions
from app.auth import _make_token
from app.main import app
from app.models.fields_audience import AudienceRule, FieldDef, FieldValue, FieldAnchorUsage
from app.models.presales import PresalesInquiry, PresalesInquiryLine, PresalesOffer, PresalesLineResponse
from app.models.system import SysAccessLog
from app.models.workbook import WorkbookOperation
from app.models.workflow import SubjectEvent, WorkItem
from app.services.presales import common, inquiries, inquiry_workbook, quote_workbook
from app.services.presales.workbook_registration import register_workbooks
from app.services.workbook import operations, registry
from app.services.workbook.registry import ApplyFailed, Conflict, WorkbookSpec
from app.services.workbook.validation import RowError
from tests.test_presales_quote_workbook import quote_data, _submitted, _published, _edit
from tests.test_workbook_operations import demo, _ctx, _edit as _edit_demo


def _snapshot(db):
    models = (PresalesInquiry, PresalesInquiryLine, PresalesOffer, PresalesLineResponse,
              AudienceRule, FieldDef, FieldValue, FieldAnchorUsage, SubjectEvent, WorkItem)
    return {model.__tablename__: list(db.execute(select(model.__table__).order_by(*model.__table__.primary_key.columns)))
            for model in models}


def _preview_quote(db, data, content):
    return operations.preview(db, quote_workbook.QUOTE_SCHEMA.name, content,
                              ctx=common.resolve_actor(db, data.buyer.id)[1])


def test_preview_collects_supplier_errors_for_every_row_and_rolls_back(db, quote_data):
    data = quote_data
    inquiry = _submitted(db, data, lines=[{"part_id": part.id, "qty": 1} for part in data.parts[:2]])
    content = _edit(quote_workbook.export_quote_sheet(db, [inquiry.id], data.buyer.id), {
        2: {"处理方式": "报价", "单价": 12, "供应商": "不存在甲"},
        3: {"处理方式": "报价", "单价": 13, "供应商": "不存在乙"},
    })
    before = _snapshot(db)
    result = _preview_quote(db, data, content)
    assert _snapshot(db) == before
    errors = result["plan"]["errors"]
    assert {error["row_number"] for error in errors} == {2, 3}
    assert all(error["column"] == "supplier" and "精确匹配唯一" in error["message"] for error in errors)
    assert all(row["errors"] for row in result["plan"]["rows"])


def test_invisible_reuse_is_located_before_apply(db, quote_data):
    data = quote_data
    hidden = _published(db, data, private=True)
    inquiry = _submitted(db, data)
    content = _edit(quote_workbook.export_quote_sheet(db, [inquiry.id], data.buyer.id),
                    {2: {"处理方式": "沿用", "沿用报价编号": hidden.id}})
    before = _snapshot(db)
    result = _preview_quote(db, data, content)
    error = result["plan"]["rows"][0]["errors"][0]
    assert error["row_number"] == 2 and error["column"] == "ref_offer_id"
    assert "不可见" in error["message"]
    assert _snapshot(db) == before


def test_preview_version_conflict_is_a_located_row_error(db, quote_data):
    data = quote_data
    inquiry = _submitted(db, data)
    content = _edit(quote_workbook.export_quote_sheet(db, [inquiry.id], data.buyer.id),
                    {2: {"处理方式": "不报", "不报原因": "无货"}})
    inquiries.lines_for(db, inquiry.id)[0].version += 1
    db.flush()
    result = _preview_quote(db, data, content)
    row = result["plan"]["rows"][0]
    assert row["status"] == "conflict"
    assert row["errors"][0]["row_number"] == 2
    assert "已变更" in row["errors"][0]["message"]


def test_valid_quote_preview_does_not_keep_drafts_events_audit_or_versions(db, quote_data):
    data = quote_data
    inquiry = _submitted(db, data)
    content = _edit(quote_workbook.export_quote_sheet(db, [inquiry.id], data.buyer.id),
                    {2: {"处理方式": "报价", "单价": 12, "供应商": data.supplier.name_raw}})
    business_logs = select(SysAccessLog.id).where(SysAccessLog.action.like("presales.%"))
    before, logs = _snapshot(db), list(db.scalars(business_logs))
    result = _preview_quote(db, data, content)
    assert not result["plan"]["errors"]
    assert _snapshot(db) == before
    assert list(db.scalars(business_logs)) == logs


def test_sales_preview_rolls_back_created_draft_fields_and_bound_pn_errors(db, quote_data):
    data = quote_data
    register_workbooks()
    ctx = common.resolve_actor(db, data.sales.id)[1]
    content = inquiry_workbook.export_template(db, ctx)
    book = load_workbook(BytesIO(content))
    sheet = book["询价明细"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, columns["PN 原文"], data.parts[0].pn_std)
    sheet.cell(2, columns["数量"], 2)
    sheet.cell(1, sheet.max_column + 1, "+质保")
    sheet.cell(2, sheet.max_column, "三年")
    out = BytesIO()
    book.save(out)
    before = _snapshot(db)
    preview = operations.preview(db, inquiry_workbook.INQUIRY_SCHEMA.name, out.getvalue(), ctx=ctx)
    assert not preview["plan"]["errors"]
    assert _snapshot(db) == before
    inquiry = inquiries.create_draft(db, data.sales.id, lines=[{"part_id": data.parts[0].id}])
    book = load_workbook(BytesIO(inquiry_workbook.export_inquiry(db, inquiry.id, ctx)))
    sheet = book["询价明细"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, columns["PN 原文"], data.parts[1].pn_std)
    out = BytesIO()
    book.save(out)
    preview = operations.preview(db, inquiry_workbook.INQUIRY_SCHEMA.name, out.getvalue(), ctx=ctx,
                                 subject_type=inquiries.SUBJECT_TYPE, subject_id=inquiry.id)
    error = preview["plan"]["rows"][0]["errors"][0]
    assert error["row_number"] == 2 and "已绑定行不能修改 PN" in error["message"]


def test_apply_revalidates_supplier_and_http_failure_receipt_has_row_errors(db, quote_data):
    data = quote_data
    inquiry = _submitted(db, data)
    content = _edit(quote_workbook.export_quote_sheet(db, [inquiry.id], data.buyer.id),
                    {2: {"处理方式": "报价", "单价": 12, "供应商": data.supplier.name_raw}})
    preview = _preview_quote(db, data, content)
    assert not preview["plan"]["errors"]
    data.supplier.name_raw = "档案已更名"
    db.commit()
    token, _ = _make_token(data.buyer.role, data.buyer.username, data.buyer.display_name,
                           perms=permissions.effective_for_user(data.buyer), token_version=data.buyer.token_version)
    client = TestClient(app)
    result = client.post(f"/api/workbooks/operations/{preview['operation_id']}/apply",
        headers={"Authorization": f"Bearer {token}"}, json={"expected_plan_digest": preview["plan_digest"]})
    assert result.status_code == 422, result.text
    detail = result.json()["detail"]
    assert detail["status"] == "failed"
    assert detail["response"]["row_errors"][0]["column"] == "supplier"
    assert detail["plan"]["rows"][0]["errors"][0]["row_number"] == 2
    db.expire_all()
    assert db.get(WorkbookOperation, preview["operation_id"]).status == "failed"
    assert not list(db.scalars(select(PresalesOffer)))


def test_apply_version_drift_has_structured_409_failure(db, quote_data):
    data = quote_data
    inquiry = _submitted(db, data)
    content = _edit(quote_workbook.export_quote_sheet(db, [inquiry.id], data.buyer.id),
                    {2: {"处理方式": "不报", "不报原因": "无货"}})
    preview = _preview_quote(db, data, content)
    inquiries.lines_for(db, inquiry.id)[0].version += 1
    db.flush()
    with pytest.raises(Conflict) as exc:
        operations.apply(db, preview["operation_id"], data.buyer.id, preview["plan_digest"])
    assert isinstance(exc.value, ApplyFailed)
    assert exc.value.status_code == 409
    assert exc.value.result["response"]["row_errors"][0]["row_number"] == 2
    assert not list(db.scalars(select(PresalesLineResponse)))


def test_generic_validator_writes_are_rolled_back_and_errors_join_plan(db, demo):
    def validate(db, comparison, ctx):
        db.get(type(demo.a), demo.a.id).display_name = "校验临时写入"
        db.flush()
        row = comparison.rows[0]
        return [RowError(row["sheet"], "业务不允许此名称", row["key"], row["row_number"], "name")]
    registry.unregister("demo")
    registry.register("demo", WorkbookSpec(demo.schema, demo.spec.load_current, demo.spec.apply_fn,
                                            demo.spec.can_access, validate))
    content = _edit_demo(demo.export(), lambda sheet: setattr(sheet["C2"], "value", "修改"))
    result = operations.preview(db, "demo", content, ctx=_ctx(demo.owner))
    assert db.get(type(demo.a), demo.a.id).display_name == "原值 A"
    assert result["plan"]["errors"][0]["message"] == "业务不允许此名称"
    assert result["plan"]["rows"][0]["errors"][0]["column"] == "name"


def test_generic_validator_cannot_commit_business_writes(db, demo):
    def validate(db, comparison, ctx):
        db.get(type(demo.a), demo.a.id).display_name = "不能提交"
        db.flush()
        db.commit()
        return []
    registry.unregister("demo")
    registry.register("demo", WorkbookSpec(demo.schema, demo.spec.load_current, demo.spec.apply_fn,
                                            demo.spec.can_access, validate))
    with pytest.raises(Exception, match="不得自行提交"):
        operations.preview(db, "demo", demo.export(), ctx=_ctx(demo.owner))
    assert db.get(type(demo.a), demo.a.id).display_name == "原值 A"


def test_validator_supports_plan_session_contract_without_retrying_type_errors(db, demo):
    calls = []
    def validate(comparison, session):
        calls.append(comparison.schema_name)
        assert session is db
        return []
    registry.unregister("demo")
    registry.register("demo", WorkbookSpec(demo.schema, demo.spec.load_current, demo.spec.apply_fn,
                                            demo.spec.can_access, validate))
    operations.preview(db, "demo", demo.export(), ctx=_ctx(demo.owner))
    assert calls == ["demo"]
    def invalid(comparison, session):
        calls.append("invalid")
        raise TypeError("业务内部 TypeError")
    registry.unregister("demo")
    registry.register("demo", WorkbookSpec(demo.schema, demo.spec.load_current, demo.spec.apply_fn,
                                            demo.spec.can_access, invalid))
    with pytest.raises(TypeError, match="业务内部"):
        operations.preview(db, "demo", demo.export(export_id="different"), ctx=_ctx(demo.owner))
    assert calls == ["demo", "invalid"]
