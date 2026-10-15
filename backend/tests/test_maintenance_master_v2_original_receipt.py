"""A completed V2 upload replays its original response without editing facts."""

from copy import deepcopy
from dataclasses import replace
from decimal import Decimal
import hashlib
import io
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import event, func, select

from app import permissions
from app.api import maintenance_project_master_workbook as workbook_api
from app.db import SessionLocal
from app.models.maintenance_project_operations import (
    MaintenanceProjectOperationAudit,
    MaintenanceProjectWorkbookOperation,
)
from app.models.system import SysUser
from app.services import maintenance_project_master_workbook as master
from app.services import maintenance_project_operations as operations
from app.services.maintenance_site_issue_facts import apply_locked_site_issue_edits
from tests.test_maintenance_expense_collection_workbook import (
    _PASSWORD,
    _client as workbook_client,
    uploader,
)
from tests.test_maintenance_project_master_v2_editable import _make_project_with_line
from tests.test_maintenance_site_issue_facts import _facts, _site_workbook
from tests.test_maintenance_site_issue_facts_compatibility import _side_effects

_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_ENTITY = "master_workbook_operation"
_ACTION = "apply_completed"


@pytest.fixture(autouse=True)
def _enable_v2(monkeypatch):
    monkeypatch.setattr(
        workbook_api.get_settings(), "maintenance_project_master_v2_enabled", True,
    )


def _site_file(db, project, lines):
    book, sheet, columns, rows = _site_workbook(db, project)
    for line in lines:
        row = rows[line.issue_line_id]
        sheet.cell(row, columns["领用数量"], 3 if line.line_no == 1 else 4)
        sheet.cell(row, columns["是否应返还"], "是" if line.line_no == 1 else "否")
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def _post(client, project, data):
    return client.post(
        f"/api/maintenance/projects/stable/{project.project_id}/master-workbook/apply",
        files={"file": ("synthetic-receipt.xlsx", data, _XLSX)},
    )


def _completion_audits(db):
    return list(db.scalars(select(MaintenanceProjectOperationAudit).where(
        MaintenanceProjectOperationAudit.entity_type == _ENTITY,
        MaintenanceProjectOperationAudit.action == _ACTION,
    )))


def _assert_snapshot(db, *, project, first, requires_amount=False):
    operations_rows = list(db.scalars(select(MaintenanceProjectWorkbookOperation).where(
        MaintenanceProjectWorkbookOperation.project_id == project.project_id,
        MaintenanceProjectWorkbookOperation.operation_type == "file_apply",
    )))
    audits = _completion_audits(db)
    assert len(operations_rows) == len(audits) == 1
    operation, audit = operations_rows[0], audits[0]
    assert audit.project_id == project.project_id
    assert audit.entity_id == str(operation.id)
    assert audit.operated_by == operation.operated_by == first["applied_by"]
    assert audit.operated_at == operation.created_at
    assert audit.after_json == {
        "schema": "master_v2_apply_receipt_v1",
        "operation_key": operation.operation_key,
        "response": first,
        "requires_contract_amount_manage": requires_amount,
    }
    return audit


def _change_permissions(db, client, *, username, changes):
    user = db.scalar(select(SysUser).where(SysUser.username == username))
    overrides = {**user.perm_overrides, **changes}
    user.perm_overrides = overrides
    user.permissions = permissions.effective_from_snapshot(user.template_perms, overrides)
    user.token_version += 1
    db.commit()
    login = client.post("/api/auth/login", json={"username": username, "password": _PASSWORD})
    assert login.status_code == 200, login.text
    client.headers["Authorization"] = f"Bearer {login.json()['token']}"


@pytest.mark.parametrize("later_change", ["quantity", "void"])
def test_http_reupload_keeps_complete_receipt_after_later_fact_change(db, later_change):
    project, issue, lines, _part = _facts(db, source="workbook")
    client = uploader(db, username=f"receipt-{later_change}")
    data = _site_file(db, project, lines)
    # An unrelated write after export makes the first authoritative result
    # carry revision_drift=True. Replays must preserve it too.
    operations.bump_workbook_revision(db, project_id=project.project_id)
    db.commit()
    response = _post(client, project, data)
    assert response.status_code == 200, response.text
    first = response.json()
    assert first["replayed"] is False
    assert first["revision_drift"] is True
    assert first["changes"]
    db.expire_all()
    audit = _assert_snapshot(db, project=project, first=first)
    audit_snapshot = deepcopy(audit.after_json)
    before_immediate_replay = _side_effects(db, project=project, issue=issue, lines=lines)
    immediate = _post(client, project, data)
    assert immediate.status_code == 200, immediate.text
    assert immediate.json() == {**first, "replayed": True}
    db.expire_all()
    assert _side_effects(db, project=project, issue=issue, lines=lines) == before_immediate_replay

    if later_change == "quantity":
        locked_project = operations._lock_project_for_fact_write(db, project.project_id)
        db.refresh(issue, with_for_update=True)
        for line in lines:
            db.refresh(line, with_for_update=True)
        apply_locked_site_issue_edits(
            db, project=locked_project, issue=issue,
            line_changes={lines[0].issue_line_id: {"quantity": "8"}},
            header_changes=None, operated_by="later-editor", reason="后续独立领用更正",
        )
        operations.bump_workbook_revision(db, project_id=project.project_id)
    else:
        operations.void_site_issue(
            db, project_id=project.project_id, issue_id=issue.issue_id,
            version=issue.version, idempotency_key=f"receipt-void-{uuid4()}",
            reason="后续独立作废", operated_by="later-editor",
        )
    db.commit()
    before_late_replay = _side_effects(db, project=project, issue=issue, lines=lines)
    replay = _post(client, project, data)
    assert replay.status_code == 200, replay.text
    assert replay.json() == {**first, "replayed": True}
    db.expire_all()
    assert _side_effects(db, project=project, issue=issue, lines=lines) == before_late_replay
    assert _assert_snapshot(db, project=project, first=first).after_json == audit_snapshot
    if later_change == "quantity":
        assert lines[0].quantity == Decimal("8")
    else:
        assert issue.normalized_status == "void"
        assert all(not line.is_active for line in lines)


def test_completion_audit_failure_rolls_back_facts_events_revision_and_receipt(db):
    project, issue, lines, _part = _facts(db, source="workbook")
    client = uploader(db, username="receipt-audit-fault")
    data = _site_file(db, project, lines)
    before = _side_effects(db, project=project, issue=issue, lines=lines)

    def fail_completion(_mapper, _connection, target):
        if target.entity_type == _ENTITY and target.action == _ACTION:
            raise RuntimeError("合成完成回执审计写入失败")

    event.listen(MaintenanceProjectOperationAudit, "before_insert", fail_completion)
    try:
        with pytest.raises(RuntimeError, match="合成完成回执审计写入失败"):
            _post(client, project, data)
    finally:
        event.remove(MaintenanceProjectOperationAudit, "before_insert", fail_completion)
    db.expire_all()
    assert _side_effects(db, project=project, issue=issue, lines=lines) == before
    assert _completion_audits(db) == []
    retry = _post(client, project, data)
    assert retry.status_code == 200, retry.text
    assert retry.json()["replayed"] is False
    _assert_snapshot(db, project=project, first=retry.json())


@pytest.mark.parametrize("validation_outcome", ["stale_contract", "conflict_plan"])
def test_http_replay_rechecks_completion_when_another_writer_wins_during_validation(
    db, monkeypatch, validation_outcome,
):
    project, issue, lines, _part = _facts(db, source="workbook")
    operator = f"receipt-race-{validation_outcome}"
    client = uploader(db, username=operator)
    data = _site_file(db, project, lines)
    original_validate = master.validate_project_master_v2
    original_replay = master.get_project_master_v2_replay
    lookup_results = []
    first_results = []

    def observe_lookup(session, **kwargs):
        result = original_replay(session, **kwargs)
        lookup_results.append(result)
        return result

    def winner_commits_before_validation_finishes(session, **kwargs):
        # The initial lookup missed. A separate real PostgreSQL session now
        # commits that exact upload while this request is still validating.
        assert lookup_results == [None]
        with SessionLocal() as writer:
            writer_plan = original_validate(writer, **kwargs)
            first_results.append(master.apply_project_master_v2(
                writer, writer_plan, operated_by=operator,
                import_batch_id=str(uuid4()), user_ctx=kwargs.get("user_ctx"),
            ))
        if validation_outcome == "stale_contract":
            raise master.WorkbookError("stale_contract", "合成校验窗口内合同版本已变化")
        plan = original_validate(session, **kwargs)
        return replace(plan, conflicts=({
            "sheet": master.V2_SHEET_SITE,
            "entity_id": lines[0].issue_line_id,
            "field": "领用数量", "base": "2", "server": "3", "user": "3",
        },))

    monkeypatch.setattr(master, "get_project_master_v2_replay", observe_lookup)
    monkeypatch.setattr(master, "validate_project_master_v2", winner_commits_before_validation_finishes)
    response = _post(client, project, data)
    assert response.status_code == 200, response.text
    assert len(first_results) == 1
    first = first_results[0]
    assert first["replayed"] is False and first["changes"]
    assert response.json() == {**first, "replayed": True}
    assert lookup_results[0] is None
    assert any(result is not None for result in lookup_results[1:])
    db.expire_all()
    after = _side_effects(db, project=project, issue=issue, lines=lines)
    assert (after["audits"], after["events"], after["workbook_operations"], after["revision"]) == (3, 1, 1, 1)
    assert (issue.version, *(line.version for line in lines)) == (2, 2, 2)
    _assert_snapshot(db, project=project, first=first)


def test_legacy_receipt_without_snapshot_rebuilds_without_writes_or_history_backfill(db):
    project, issue, lines, _part = _facts(db, source="workbook")
    operator = "receipt-historical"
    client = uploader(db, username=operator)
    data = _site_file(db, project, lines)
    plan = master.validate_project_master_v2(db, project_id=project.project_id, data=data)
    for line in lines:
        line.quantity = Decimal(3 if line.line_no == 1 else 4)
        line.no_return = line.line_no == 2
        line.version = 2
    issue.version = 2
    state = operations.get_or_create_workbook_state(db, project_id=project.project_id)
    state.revision = 1
    # Seed the old persisted receipt directly. Append-only audit history is
    # never deleted to simulate a version that did not yet write snapshots.
    db.add(MaintenanceProjectWorkbookOperation(
        project_id=project.project_id, export_id=plan.export_id,
        file_sha256=plan.file_sha256, payload_hash=plan.file_sha256,
        operation_key=master._v2_apply_operation_key(plan, operator),
        operation_type="file_apply", operated_by=operator,
    ))
    db.commit()
    before = _side_effects(db, project=project, issue=issue, lines=lines)
    assert master.get_project_master_v2_replay(
        db, project_id=project.project_id, file_sha256=hashlib.sha256(data).hexdigest(),
        operated_by=operator,
    ) is None
    response = _post(client, project, data)
    assert response.status_code == 200, response.text
    assert response.json()["replayed"] is True
    assert response.json()["changes"] == []
    db.expire_all()
    assert _side_effects(db, project=project, issue=issue, lines=lines) == before
    assert _completion_audits(db) == []


def test_noop_upload_records_one_receipt_without_business_versions_or_events(db):
    project, issue, lines, _part = _facts(db, source="workbook")
    client = uploader(db, username="receipt-noop")
    data = master.build_project_master_v2(
        db, project_id=project.project_id, sheets=(master.V2_SHEET_SITE,),
    )
    response = _post(client, project, data)
    assert response.status_code == 200, response.text
    first = response.json()
    assert first["changes"] == [] and first["replayed"] is False
    db.expire_all()
    before_replay = _side_effects(db, project=project, issue=issue, lines=lines)
    assert before_replay["revision"] == before_replay["events"] == 0
    assert before_replay["audits"] == before_replay["workbook_operations"] == 1
    assert (issue.version, *(line.version for line in lines)) == (1, 1, 1)
    _assert_snapshot(db, project=project, first=first)
    replay = _post(client, project, data)
    assert replay.status_code == 200, replay.text
    assert replay.json() == {**first, "replayed": True}
    db.expire_all()
    assert _side_effects(db, project=project, issue=issue, lines=lines) == before_replay


def test_replay_still_requires_current_upload_permission(db):
    project, issue, lines, _part = _facts(db, source="workbook")
    username = "receipt-revoked"
    client = uploader(db, username=username)
    data = _site_file(db, project, lines)
    first = _post(client, project, data)
    assert first.status_code == 200, first.text
    _change_permissions(db, client, username=username, changes={
        "action_maintenance_expense_collection_upload": False,
    })
    before = _side_effects(db, project=project, issue=issue, lines=lines)
    denied = _post(client, project, data)
    assert denied.status_code == 403, denied.text
    assert "changes" not in denied.json()
    db.expire_all()
    assert _side_effects(db, project=project, issue=issue, lines=lines) == before


def test_replay_is_bound_to_original_operator_and_project(db):
    project, issue, lines, _part = _facts(db, source="workbook")
    username = "receipt-bound"
    client = uploader(db, username=username)
    data = _site_file(db, project, lines)
    first = _post(client, project, data)
    assert first.status_code == 200, first.text
    other_project, _other_issue, _other_lines, _other_part = _facts(db)
    file_hash = hashlib.sha256(data).hexdigest()
    for project_id, operator in (
        (project.project_id, "another-operator"),
        (other_project.project_id, username),
    ):
        assert master.get_project_master_v2_replay(
            db, project_id=project_id, file_sha256=file_hash, operated_by=operator,
        ) is None
    before = _side_effects(db, project=project, issue=issue, lines=lines)
    denied = _post(client, other_project, data)
    assert denied.status_code == 403, denied.text
    assert "changes" not in denied.json()
    db.expire_all()
    assert _side_effects(db, project=project, issue=issue, lines=lines) == before


def test_contract_amount_replay_rechecks_current_amount_permission(db):
    project, _part, _order, _line = _make_project_with_line(db)
    username = "receipt-amount"
    client = workbook_client(db, username=username, overrides={
        "page_maintenance": True, "data_purchase_cost": True, "data_profit": True,
        "action_maintenance_expense_collection_upload": True,
        "action_maintenance_project_manage": True,
    })
    book = load_workbook(io.BytesIO(master.build_project_master_v2(
        db, project_id=project.project_id, sheets=(master.V2_SHEET_OVERVIEW,),
    )))
    sheet = book[master.V2_SHEET_OVERVIEW]
    amount_cell = next(row[1] for row in sheet.iter_rows(min_col=1, max_col=2)
                       if row[0].value == "合同总额（含税）")
    amount_cell.value = Decimal("12345.67")
    output = io.BytesIO()
    book.save(output)
    data = output.getvalue()
    first = _post(client, project, data)
    assert first.status_code == 200, first.text
    assert first.json()["contract_updates"] == 1
    _assert_snapshot(db, project=project, first=first.json(), requires_amount=True)
    _change_permissions(db, client, username=username, changes={
        "action_maintenance_project_manage": False,
    })
    audits_before = db.scalar(select(func.count()).select_from(MaintenanceProjectOperationAudit))
    revision_before = operations.get_or_create_workbook_state(db, project_id=project.project_id).revision
    denied = _post(client, project, data)
    assert denied.status_code == 403, denied.text
    assert "changes" not in denied.json()
    assert db.scalar(select(func.count()).select_from(MaintenanceProjectOperationAudit)) == audits_before
    assert operations.get_or_create_workbook_state(db, project_id=project.project_id).revision == revision_before
