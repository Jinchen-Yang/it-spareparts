"""Independent HTTP counterexamples for the R3b state-consumer migration.

Every request uses real auth, routers, services and PostgreSQL. Failure injection
is at a downstream boundary, after facts/audits/state have actually been flushed.
"""

from decimal import Decimal
from io import BytesIO

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import func, select

from app import auth
from app.api import maintenance_projects, maintenance_source_assignments
from app.api import maintenance_project_workbook_v3
from app.auth import hash_password
from app.models.dimensions import DimPart
from app.models.maintenance import FMaintenanceLine, FMaintenanceOrder
from app.models.maintenance_project import (
    MaintenanceProject,
    MaintenanceProjectAuditLog,
    MaintenanceProjectUserAssignment,
)
from app.models.maintenance_project_operations import MaintenanceProjectWorkbookState
from app.models.maintenance_source_assignment import MaintenanceSourceOrderAssignment
from app.models.system import SysImportBatch, SysUser
from app.services import maintenance_project_assignments as viewers
from app.services import maintenance_project_operations as legacy
from app.services import maintenance_warehouse as warehouse
from app.services import maintenance_workbook_state as state


PROJECTS = "/api/maintenance/projects/stable"
ORDERS = "/api/maintenance/project-assignments/orders"
PASSWORD = "r3b-independent-synthetic-password"
STATE_NAMES = (
    "get_or_create_workbook_state", "lock_workbook_states",
    "bump_workbook_revision", "bump_locked_workbook_revision",
    "_workbook_data_version", "_lock_project_for_fact_write",
)


@pytest.fixture(autouse=True)
def retired_state_exports_are_unusable(monkeypatch):
    def forbidden(*_args, **_kwargs):
        pytest.fail("a real HTTP consumer called the retired operations state seam")

    for name in STATE_NAMES:
        monkeypatch.setattr(legacy, name, forbidden)


def _user(db, username, *, role="admin", permissions=None):
    user = SysUser(
        username=username, display_name=username, role=role,
        password_hash=hash_password(PASSWORD), permissions=permissions,
    )
    db.add(user)
    db.commit()
    return user


def _client(user):
    application = FastAPI()
    for router in (
        auth.router, maintenance_projects.router,
        maintenance_source_assignments.router, maintenance_project_workbook_v3.router,
    ):
        application.include_router(router, prefix="/api")
    client = TestClient(application, raise_server_exceptions=False)
    login = client.post(
        "/api/auth/login", json={"username": user.username, "password": PASSWORD},
    )
    assert login.status_code == 200, login.text
    client.headers["Authorization"] = f"Bearer {login.json()['token']}"
    return client


def _project(db, suffix):
    project = MaintenanceProject(
        project_id=f"r3b-api-{suffix}", project_code=f"R3B-API-{suffix}",
        display_name=f"独立接口核验{suffix}", lifecycle_status="ongoing",
    )
    db.add(project)
    db.commit()
    return project.project_id


def _order(db, suffix):
    batch = SysImportBatch(
        filename=f"{suffix}.xlsx", file_type="maintenance",
        file_hash=f"r3b-independent-{suffix}", status="success",
    )
    db.add(batch)
    db.flush()
    raw_id = f"WBDD-R3B-{suffix}"
    order = FMaintenanceOrder(
        raw_order_id=raw_id, order_no=raw_id,
        project_raw="原始线索不决定稳定归属", project_std="原始线索不决定稳定归属",
        data_status="已生效", import_batch_id=batch.id,
    )
    part = DimPart(pn_std=f"PN-R3B-{suffix}")
    db.add_all([order, part])
    db.flush()
    db.add(FMaintenanceLine(
        raw_line_id=f"{raw_id}-L1", order_id=order.id, line_no=1,
        part_id=part.id, pn_std=part.pn_std, qty=Decimal("3.125"),
        import_batch_id=batch.id,
    ))
    db.commit()
    return raw_id


def _count(db, model):
    db.expire_all()
    return db.scalar(select(func.count()).select_from(model))


def _state(db, project_id):
    db.expire_all()
    row = db.get(MaintenanceProjectWorkbookState, project_id)
    return None if row is None else (row.revision, row.data_version)


def _export_baseline(client, project_id, source_ids=None):
    response = client.get(f"{PROJECTS}/{project_id}/workbook-v3.xlsx")
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    try:
        metadata = dict(workbook["99_元数据"].iter_rows(values_only=True))
        assert metadata["项目ID"] == project_id
        if source_ids is not None:
            sheet = workbook["03_备件订单"]
            headers = next(sheet.iter_rows(values_only=True))
            rows = [dict(zip(headers, row)) for row in sheet.iter_rows(
                min_row=2, values_only=True,
            ) if row[0]]
            assert {row["维保单号"] for row in rows} == set(source_ids)
            assert all(row["需求数量"] == 3.125 for row in rows)
            assert all(row["PN"] == "PN-R3B-transfer" for row in rows)
        return metadata["基线版本"]
    finally:
        workbook.close()


def _assign(client, project_id, source_id, current=None):
    item = {"source_order_id": source_id}
    if current is not None:
        item.update(
            expected_assignment_id=current["assignment_id"],
            expected_version=current["version"],
        )
    return client.post(f"{ORDERS}/assign", json={
        "project_id": project_id, "items": [item],
        "reason": "独立接口核验确认来源归属",
    })


def test_source_transfer_failure_after_both_bumps_preserves_history_and_export_tokens(
    db, monkeypatch,
):
    """Both source/destination snapshots survive a failed real downstream repair."""
    old_id, new_id = _project(db, "old"), _project(db, "new")
    source_id = _order(db, "transfer")
    client = _client(_user(db, "r3b_transfer_admin"))
    initial = _assign(client, old_id, source_id)
    assert initial.status_code == 200, initial.text
    original = initial.json()["assignments"][0]
    before = {
        old_id: _export_baseline(client, old_id, [source_id]),
        new_id: _export_baseline(client, new_id, []),
    }
    audit_count = _count(db, MaintenanceProjectAuditLog)
    entered = []
    real_reconcile = warehouse.reconcile_project_assignment_links

    def fail_after_real_repair(session, **kwargs):
        real_reconcile(session, **kwargs)
        revisions = {
            row.project_id: row.revision
            for row in session.scalars(select(MaintenanceProjectWorkbookState))
        }
        assert revisions == {old_id: 2, new_id: 1}
        assert session.scalar(select(func.count()).select_from(
            MaintenanceSourceOrderAssignment)) == 2
        entered.append(True)
        raise RuntimeError("synthetic late warehouse repair failure")

    monkeypatch.setattr(warehouse, "reconcile_project_assignment_links", fail_after_real_repair)
    failed = _assign(client, new_id, source_id, original)
    assert failed.status_code == 500, failed.text
    assert entered == [True]
    assert _count(db, MaintenanceSourceOrderAssignment) == 1
    assert _count(db, MaintenanceProjectAuditLog) == audit_count
    assert _state(db, old_id) == (1, before[old_id])
    # Export itself may INSERT revision zero but its read request does not commit.
    assert _state(db, new_id) is None
    assert _export_baseline(client, old_id, [source_id]) == before[old_id]
    assert _export_baseline(client, new_id, []) == before[new_id]

    monkeypatch.setattr(warehouse, "reconcile_project_assignment_links", real_reconcile)
    retried = _assign(client, new_id, source_id, original)
    assert retried.status_code == 200, retried.text
    current = retried.json()["assignments"][0]
    assert current["project_id"] == new_id
    assert current["assignment_id"] != original["assignment_id"]
    assert _count(db, MaintenanceProjectAuditLog) == audit_count + 2
    assert _state(db, old_id)[0] == 2
    assert _state(db, new_id)[0] == 1
    assert _export_baseline(client, old_id, []) != before[old_id]
    assert _export_baseline(client, new_id, [source_id]) != before[new_id]
    db.expire_all()
    archived = db.get(MaintenanceSourceOrderAssignment, original["assignment_id"])
    assert (archived.is_active, archived.version, archived.archived_by) == (
        False, 2, "r3b_transfer_admin",
    )
    directory = client.get(ORDERS, params={"assignment_status": "assigned"})
    assert directory.status_code == 200, directory.text
    assert directory.json()["total"] == 1
    assert directory.json()["rows"][0]["assigned_project"]["project_id"] == new_id
    # An obsolete generation cannot undo the committed transfer or bump either side.
    stale = _assign(client, old_id, source_id, original)
    assert stale.status_code == 409, stale.text
    assert (_state(db, old_id)[0], _state(db, new_id)[0]) == (2, 1)
    assert _count(db, MaintenanceProjectAuditLog) == audit_count + 2


@pytest.mark.parametrize("failure_type,status_code", [(RuntimeError, 500),
    (viewers.MaintenanceProjectAssignmentError, 400)])
def test_mixed_patch_late_viewer_failure_rolls_back_first_bump_then_retry_is_fresh(
    db, monkeypatch, failure_type, status_code,
):
    project_id = _project(db, "mixed")
    viewer = _user(db, "r3b_late_viewer", role="purchaser")
    client = _client(_user(db, "r3b_mixed_admin"))
    body = {
        "version": 1, "project_note": "独立备注", "no_return_note": "人工提醒",
        "visible_usernames": [viewer.username], "reason": "混合写入末端故障",
    }
    real_sync = viewers.sync_project_viewers
    reached = []

    def fail_after_sync(session, **kwargs):
        result = real_sync(session, **kwargs)
        assert result[1] is True
        assert session.get(MaintenanceProjectWorkbookState, project_id).revision == 1
        assert session.scalar(select(func.count()).select_from(
            MaintenanceProjectUserAssignment)) == 1
        assert session.get(MaintenanceProject, project_id).project_note == "独立备注"
        reached.append(True)
        raise failure_type("synthetic failure after actual viewer audit")

    monkeypatch.setattr(viewers, "sync_project_viewers", fail_after_sync)
    failed = client.patch(f"{PROJECTS}/{project_id}", json=body)
    assert failed.status_code == status_code, failed.text
    assert reached == [True]
    assert _state(db, project_id) is None
    assert _count(db, MaintenanceProjectAuditLog) == 0
    assert _count(db, MaintenanceProjectUserAssignment) == 0
    db.expire_all()
    project = db.get(MaintenanceProject, project_id)
    assert (project.version, project.project_note, project.no_return_note) == (1, None, None)

    monkeypatch.setattr(viewers, "sync_project_viewers", real_sync)
    retry = client.patch(f"{PROJECTS}/{project_id}", json=body)
    assert retry.status_code == 200, retry.text
    assert retry.json()["version"] == 2
    assert _state(db, project_id)[0] == 1
    assert _count(db, MaintenanceProjectAuditLog) == 2
    # Same nominal request is now stale, even though its values match current facts.
    stale = client.patch(f"{PROJECTS}/{project_id}", json=body)
    assert stale.status_code == 409, stale.text
    assert _state(db, project_id)[0] == 1
    assert _count(db, MaintenanceProjectAuditLog) == 2
    fresh = client.patch(f"{PROJECTS}/{project_id}", json={
        "version": 2, "visible_usernames": [], "reason": "下一独立请求撤销可见性",
    })
    assert fresh.status_code == 200, fresh.text
    assert fresh.json()["version"] == 3
    assert _state(db, project_id)[0] == 2
    db.expire_all()
    project = db.get(MaintenanceProject, project_id)
    assert (project.project_note, project.no_return_note) == ("独立备注", "人工提醒")


def test_real_invalid_viewer_rolls_back_master_and_export_baseline(db):
    project_id = _project(db, "invalid-viewer")
    client = _client(_user(db, "r3b_invalid_viewer_admin"))
    baseline = _export_baseline(client, project_id)
    failed = client.patch(f"{PROJECTS}/{project_id}", json={
        "version": 1, "project_note": "不该部分落库",
        "visible_usernames": ["r3b-absent-account"], "reason": "实际业务校验拒绝",
    })
    assert failed.status_code == 400, failed.text
    assert "r3b-absent-account" in failed.json()["detail"]
    assert _state(db, project_id) is None
    assert _count(db, MaintenanceProjectAuditLog) == 0
    assert _count(db, MaintenanceProjectUserAssignment) == 0
    assert db.get(MaintenanceProject, project_id).project_note is None
    assert _export_baseline(client, project_id) == baseline


def test_revoked_viewer_old_token_cannot_enter_state_writes_or_export(db, monkeypatch):
    project_id = _project(db, "revoked-viewer")
    user = _user(db, "r3b_scoped_writer", role="purchaser", permissions={
        "page_maintenance": True, "action_maintenance_project_manage": True,
        "data_profit": True, "data_purchase_cost": True,
    })
    admin = _client(_user(db, "r3b_revoke_admin"))
    assigned = admin.patch(f"{PROJECTS}/{project_id}", json={
        "version": 1, "visible_usernames": [user.username], "reason": "显式赋予项目可见性",
    })
    assert assigned.status_code == 200, assigned.text
    client = _client(user)
    assert _export_baseline(client, project_id) == _state(db, project_id)[1]
    revoked = admin.patch(f"{PROJECTS}/{project_id}", json={
        "version": 2, "visible_usernames": [], "reason": "撤销项目可见性",
    })
    assert revoked.status_code == 200, revoked.text
    before = _state(db, project_id)
    audit_count = _count(db, MaintenanceProjectAuditLog)

    def forbidden(*_args, **_kwargs):
        pytest.fail("revoked scope reached the state primitive")

    for name in ("lock_workbook_states", "get_or_create_workbook_state"):
        monkeypatch.setattr(state, name, forbidden)
    denied_write = client.patch(f"{PROJECTS}/{project_id}", json={
        "version": 3, "project_note": "越权", "reason": "沿用旧 token 猜项目 ID",
    })
    assert denied_write.status_code == 403, denied_write.text
    assert "无权访问" in denied_write.json()["detail"]
    denied_export = client.get(f"{PROJECTS}/{project_id}/workbook-v3.xlsx")
    assert denied_export.status_code == 403, denied_export.text
    assert _state(db, project_id) == before
    assert _count(db, MaintenanceProjectAuditLog) == audit_count
    assert db.get(MaintenanceProject, project_id).project_note is None


def test_live_money_permission_revocation_precedes_v3_state_creation(db, monkeypatch):
    project_id = _project(db, "money-revoked")
    user = _user(db, "r3b_money_reader", role="boss", permissions={
        "page_maintenance": True, "data_profit": True, "data_purchase_cost": True,
    })
    client = _client(user)
    _export_baseline(client, project_id)
    user.permissions = {**user.permissions, "data_purchase_cost": False}
    db.commit()

    def forbidden(*_args, **_kwargs):
        pytest.fail("live financial permission denial reached workbook state creation")

    monkeypatch.setattr(state, "get_or_create_workbook_state", forbidden)
    denied = client.get(f"{PROJECTS}/{project_id}/workbook-v3.xlsx")
    assert denied.status_code == 403, denied.text
    assert "采购成本" in denied.json()["detail"]
    assert _state(db, project_id) is None
    assert _count(db, MaintenanceProjectAuditLog) == 0


@pytest.mark.parametrize("case", ["incomplete-expectation", "note-501", "note-control"])
def test_invalid_http_input_does_not_open_a_workbook_state(db, monkeypatch, case):
    project_id = _project(db, "input")
    client = _client(_user(db, "r3b_input_admin"))

    def forbidden(*_args, **_kwargs):
        pytest.fail("invalid HTTP body reached state creation/locking")

    for name in ("get_or_create_workbook_state", "lock_workbook_states"):
        monkeypatch.setattr(state, name, forbidden)
    if case == "incomplete-expectation":
        response = client.post(f"{ORDERS}/assign", json={
            "project_id": project_id,
            "items": [{"source_order_id": "not-looked-up", "expected_version": 1}],
            "reason": "不完整的乐观锁输入",
        })
    else:
        response = client.patch(f"{PROJECTS}/{project_id}", json={
            "version": 1, "project_note": "字" * 501 if case == "note-501" else "含\x01控制符",
            "reason": "当前备注校验边界",
        })
    assert response.status_code == 422, response.text
    assert _state(db, project_id) is None
    assert _count(db, MaintenanceProjectAuditLog) == 0


def test_missing_source_rejects_after_state_creation_then_valid_retry_is_not_deduped(db):
    project_id = _project(db, "missing-source")
    source_id = _order(db, "missing-source")
    client = _client(_user(db, "r3b_missing_source_admin"))
    failed = client.post(f"{ORDERS}/assign", json={
        "project_id": project_id,
        "items": [{"source_order_id": source_id}, {"source_order_id": "absent"}],
        "reason": "整组校验不能留下首行挂靠",
    })
    assert failed.status_code == 400, failed.text
    assert "来源维保单不存在" in failed.json()["detail"]
    assert _state(db, project_id) is None
    assert _count(db, MaintenanceSourceOrderAssignment) == 0
    assert _count(db, MaintenanceProjectAuditLog) == 0
    retry = _assign(client, project_id, source_id)
    assert retry.status_code == 200, retry.text
    assert _state(db, project_id)[0] == 1
    assert _count(db, MaintenanceSourceOrderAssignment) == 1
    assert _count(db, MaintenanceProjectAuditLog) == 1
