"""06 锁后授权事务反例：覆盖 CREATE、缺行 VOID、同号合单、地址和免返编辑。

混合修改 03 和 02；后续 02 等锁时撤权，必须回滚已 flush 的事实、事件和回执。
使用真实 HTTP、PostgreSQL、鉴权和锁。
"""
from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
from datetime import date
from decimal import Decimal
from io import BytesIO
import json
from threading import Event
from time import monotonic
from types import SimpleNamespace
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import func, select, text
from sqlalchemy.orm import Session

from app import config
from app.api import maintenance_project_master_workbook as workbook_api
from app.auth import hash_password, verify_token
from app.db import Base, SessionLocal, engine
from app.main import app
from app.models.maintenance import FMaintenanceLine
from app.models.maintenance_manager import MaintenanceCollectionMilestone
from app.models.maintenance_project import MaintenanceProjectContract
from app.models.maintenance_project_operations import (
    MaintenanceProjectOperationAudit, MaintenanceProjectWorkbookOperation,
    MaintenanceSiteIssue, MaintenanceSiteIssueLine,
)
from app.models.system import SysUser
from app.services import maintenance_project_master_workbook as master
from tests.test_maintenance_expense_collection_workbook import _client, _PASSWORD, _XLSX
from tests.test_maintenance_project_master_v2_editable import _make_project_with_line, _site_issue


_ADMIN_PASSWORD = "synthetic-transaction-crosscheck-admin-1"
_RIGHTS = {
    "page_maintenance": True, "page_maintenance_beta": True,
    "data_purchase_cost": True, "data_profit": True,
    "action_maintenance_expense_collection_upload": True,
    "action_maintenance_site_issue_manage": False,
}


def _business_snapshot():
    # 快照包含事件、合单引用、计数器、保护字段、版本、回执和导入批次。
    # 访问与安全审计不属于经营事实；预检记录也不视为应用写入。
    tables = sorted([
        table for name, table in Base.metadata.tables.items()
        if (name.startswith(("maintenance_", "f_maintenance"))
            or name in {"f_project_expense", "sys_import_batch"})
        and name != "maintenance_project_workbook_validation"
    ], key=lambda table: table.name)
    with SessionLocal() as db:
        return {
            table.name: sorted(json.dumps(dict(row), sort_keys=True, default=str)
                               for row in db.execute(select(table)).mappings())
            for table in tables
        }


def _login(client, username, password):
    response = client.post("/api/auth/login", json={"username": username, "password": password})
    assert response.status_code == 200, response.text
    client.headers["Authorization"] = "Bearer " + response.json()["token"]


def _send(case):
    return case.client.post(case.route + "/apply",
                           files={"file": ("independent-actions.xlsx", case.data, _XLSX)},
                           data={"site_issue_merge_token": case.merge_token})


def _set_rights(case, allowed):
    response = case.admin.put(f"/api/accounts/{case.username}", json={
        "overrides": {**_RIGHTS, "action_maintenance_expense_collection_upload": allowed},
    })
    assert response.status_code == 200, response.text


@contextmanager
def _scenario(db, monkeypatch, action):
    monkeypatch.setattr(workbook_api.get_settings(), "maintenance_project_master_v2_enabled", True)
    project, part, _order, demand = _make_project_with_line(db)
    issue, line = _site_issue(db, project, part, issue_no="CROSS-ORIGINAL")
    issue.site_location = "Synthetic original address"
    issues = [issue.issue_id]
    lines = [line.issue_line_id]
    if action == "merge":
        second, second_line = _site_issue(db, project, part, issue_no="CROSS-SECOND")
        second.site_location = issue.site_location
        issues.append(second.issue_id)
        lines.append(second_line.issue_line_id)
    contract = db.scalar(select(MaintenanceProjectContract).where(
        MaintenanceProjectContract.project_id == project.project_id))
    milestone_id = str(uuid4())
    db.add(MaintenanceCollectionMilestone(
        milestone_id=milestone_id, project_id=project.project_id,
        project_contract_id=contract.project_contract_id, sequence=1,
        planned_date=date(2026, 10, 31), planned_amount=Decimal("800.00"),
        completeness_state="complete", source="project_master_v2", version=1,
    ))
    db.add(SysUser(username="crosscheck-admin", role="admin", display_name="Synthetic admin",
                   password_hash=hash_password(_ADMIN_PASSWORD), is_active=True))
    db.commit()
    username = "crosscheck-uploader"
    client = _client(db, username=username, overrides=_RIGHTS)
    admin = TestClient(app)
    _login(admin, "crosscheck-admin", _ADMIN_PASSWORD)
    route = f"/api/maintenance/projects/stable/{project.project_id}/master-workbook"
    exported = client.get(route + ".xlsx", params={"sheets": ",".join((
        master.V2_SHEET_PARTS, master.V2_SHEET_SITE, master.V2_SHEET_PLAN))})
    assert exported.status_code == 200, exported.text
    book = load_workbook(BytesIO(exported.content))
    site = book[master.V2_SHEET_SITE]
    columns = {cell.value: cell.column for cell in site[1] if cell.value}
    row_ids = {site.cell(row, columns["实体ID"]).value: row
               for row in range(2, site.max_row + 1)
               if site.cell(row, columns["实体ID"]).value}
    if action == "create":
        values = {"领用日期": "2026-10-23", "PN": part.pn_std,
                  "SN": "INDEPENDENT-NEW-SERIAL", "领用数量": "1.75",
                  "是否应返还": "是", "领用地址": "Synthetic new address"}
        site.append([values.get(cell.value) for cell in site[1]])
    elif action == "void":
        site.delete_rows(row_ids[line.issue_line_id])
    elif action == "merge":
        for row in row_ids.values():
            site.cell(row, columns["领用单号"], "CROSS-MERGED")
    else:
        assert action == "address_and_no_return"
        row = row_ids[line.issue_line_id]
        site.cell(row, columns["地址操作"], "清空")
        site.cell(row, columns["是否应返还"], "否")
    parts = book[master.V2_SHEET_PARTS]
    parts_columns = {cell.value: cell.column for cell in parts[1] if cell.value}
    parts.cell(2, parts_columns["备注"], "Mixed 03 mutation must be atomic")
    plan = book[master.V2_SHEET_PLAN]
    plan_columns = {cell.value: cell.column for cell in plan[1] if cell.value}
    row = next(row for row in range(2, plan.max_row + 1)
               if plan.cell(row, plan_columns["实体ID"]).value == milestone_id)
    plan.cell(row, plan_columns["计划回款金额（含税）"], "702.37")
    out = BytesIO()
    book.save(out)
    book.close()
    data = out.getvalue()
    preview = client.post(route + "/validate", files={"file": ("actions.xlsx", data, _XLSX)})
    assert preview.status_code == 200, preview.text
    assert not preview.json()["conflicts"], preview.text
    key = {"create": "site_creates", "void": "site_voids", "merge": "site_issue_merges",
           "address_and_no_return": "site_updates"}[action]
    assert preview.json()["summary"][key] >= 1, preview.text
    case = SimpleNamespace(
        action=action, project_id=project.project_id, issue_ids=issues, line_ids=lines,
        demand_id=demand.id, milestone_id=milestone_id, username=username,
        client=client, admin=admin, route=route, data=data,
        merge_token=preview.json().get("site_issue_merge_token") or "",
    )
    db.rollback()
    try:
        yield case
    finally:
        client.close()
        admin.close()


def _observe_block(observer, holder_pid, future, expected):
    deadline = monotonic() + 15
    while monotonic() < deadline:
        rows = observer.execute(text("""
            SELECT pid, wait_event_type, query FROM pg_stat_activity
            WHERE datname = current_database() AND :pid = ANY(pg_blocking_pids(pid))
        """), {"pid": holder_pid}).mappings().all()
        if rows:
            assert len(rows) == 1, rows
            assert rows[0]["wait_event_type"] == "Lock"
            assert expected in rows[0]["query"], rows[0]
            return rows[0]["pid"]
        assert not future.done(), f"request did not wait: {future.result().text}"
        Event().wait(0.01)
    pytest.fail("PostgreSQL did not report the expected real lock wait")


def _assert_applied(case):
    with SessionLocal() as current:
        assert current.get(FMaintenanceLine, case.demand_id).line_note == "Mixed 03 mutation must be atomic"
        assert current.get(MaintenanceCollectionMilestone, case.milestone_id).planned_amount == Decimal("702.37")
        if case.action == "create":
            added = current.scalars(select(MaintenanceSiteIssueLine).where(
                MaintenanceSiteIssueLine.serial_number == "INDEPENDENT-NEW-SERIAL")).one()
            assert added.quantity == Decimal("1.750")
            parent = current.get(MaintenanceSiteIssue, added.issue_id)
            assert parent.display_no and parent.site_location == "Synthetic new address"
        elif case.action == "void":
            assert current.get(MaintenanceSiteIssueLine, case.line_ids[0]).is_active is False
            assert current.get(MaintenanceSiteIssue, case.issue_ids[0]).normalized_status == "void"
        elif case.action == "merge":
            rows = [current.get(MaintenanceSiteIssueLine, value) for value in case.line_ids]
            assert len({row.issue_id for row in rows}) == 1
            assert all(row.is_active for row in rows)
            assert sorted(row.line_no for row in rows) == [1, 2]
            assert current.get(MaintenanceSiteIssue, rows[0].issue_id).issue_no == "CROSS-MERGED"
            assert sorted(current.get(MaintenanceSiteIssue, value).normalized_status
                          for value in case.issue_ids) == ["confirmed", "void"]
        else:
            assert current.get(MaintenanceSiteIssueLine, case.line_ids[0]).no_return is True
            assert current.get(MaintenanceSiteIssue, case.issue_ids[0]).site_location is None
        assert current.scalar(select(func.count()).select_from(MaintenanceProjectWorkbookOperation)) == 1
        assert current.scalar(select(func.count()).select_from(MaintenanceProjectOperationAudit).where(
            MaintenanceProjectOperationAudit.action == "apply_completed")) == 1


@pytest.mark.parametrize("action", ["create", "void", "merge", "address_and_no_return"])
@pytest.mark.parametrize("wait_at", ["before_facts", "after_site_writes"])
def test_complex_site_actions_revoke_atomically_then_retry_and_replay(db, monkeypatch, action, wait_at):
    with _scenario(db, monkeypatch, action) as case:
        before = _business_snapshot()
        with SessionLocal() as blocker, engine.connect().execution_options(isolation_level="AUTOCOMMIT") as observer:
            holder_pid = blocker.scalar(text("SELECT pg_backend_pid()"))
            if wait_at == "before_facts":
                blocker.execute(select(func.pg_advisory_xact_lock(config.DATA_CHANGE_ADVISORY_LOCK_KEY)))
                expected = "pg_advisory_xact_lock"
            else:
                blocker.execute(select(MaintenanceCollectionMilestone).where(
                    MaintenanceCollectionMilestone.milestone_id == case.milestone_id).with_for_update()).scalar_one()
                expected = "maintenance_collection_milestone"
            with ThreadPoolExecutor(max_workers=1) as pool:
                future = pool.submit(_send, case)
                try:
                    waiter_pid = _observe_block(observer, holder_pid, future, expected)
                    _set_rights(case, False)
                    assert holder_pid in observer.scalar(text("SELECT pg_blocking_pids(:pid)"), {"pid": waiter_pid})
                    assert _business_snapshot() == before
                finally:
                    blocker.rollback()
                rejected = future.result(timeout=30)
        assert rejected.status_code == 401, rejected.text
        assert _business_snapshot() == before
        _set_rights(case, True)
        _login(case.client, case.username, _PASSWORD)
        first = _send(case)
        assert first.status_code == 200, first.text
        assert first.json()["replayed"] is False
        _assert_applied(case)
        after = _business_snapshot()
        replay = _send(case)
        assert replay.status_code == 200, replay.text
        assert replay.json() == {**first.json(), "replayed": True}
        assert _business_snapshot() == after


def test_http_locked_receipt_replay_checks_revocation_after_another_writer_commits(db, monkeypatch):
    """第二个真实 HTTP 请求在回执产生前开始等待。

    测试使用同键的会话 advisory lock，使第二个请求等待首个写入者真实提交，
    不替换应用的锁、回执或鉴权逻辑。
    """
    with _scenario(db, monkeypatch, "create") as case:
        identity = verify_token(case.client.headers["Authorization"].removeprefix("Bearer "))
        with engine.connect() as connection, engine.connect().execution_options(isolation_level="AUTOCOMMIT") as observer:
            holder_pid = connection.scalar(text("SELECT pg_backend_pid()"))
            connection.execute(select(func.pg_advisory_lock(config.DATA_CHANGE_ADVISORY_LOCK_KEY)))
            connection.commit()
            try:
                with Session(bind=connection, autoflush=False, expire_on_commit=False) as first_db:
                    plan = master.validate_project_master_v2(first_db, project_id=case.project_id, data=case.data)
                    first_db.rollback()
                    with ThreadPoolExecutor(max_workers=1) as pool:
                        waiting = pool.submit(_send, case)
                        try:
                            waiter_pid = _observe_block(observer, holder_pid, waiting, "pg_advisory_xact_lock")
                            first = master.apply_project_master_v2(
                                first_db, plan, operated_by=case.username, import_batch_id=str(uuid4()),
                                reauthorize=lambda: workbook_api._reauthorize_site_workbook(
                                    first_db, ident=identity, plan=plan, operated_by=case.username),
                            )
                            assert first["replayed"] is False
                            _assert_applied(case)
                            committed = _business_snapshot()
                            _set_rights(case, False)
                            assert holder_pid in observer.scalar(text("SELECT pg_blocking_pids(:pid)"), {"pid": waiter_pid})
                        finally:
                            first_db.rollback()
                            connection.execute(select(func.pg_advisory_unlock(config.DATA_CHANGE_ADVISORY_LOCK_KEY)))
                            connection.commit()
                        denied = waiting.result(timeout=30)
                assert denied.status_code == 401, denied.text
                assert _business_snapshot() == committed
                _set_rights(case, True)
                _login(case.client, case.username, _PASSWORD)
                replay = _send(case)
                assert replay.status_code == 200, replay.text
                assert replay.json() == {**first, "replayed": True}
                assert _business_snapshot() == committed
            finally:
                # 会话锁不会随事务回滚释放，连接归还池前必须显式解锁。
                connection.rollback()
                connection.execute(select(func.pg_advisory_unlock_all()))
                connection.commit()
