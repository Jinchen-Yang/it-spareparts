"""Agent 文件归属/越权（PR-C）：上传与生成文件归创建者，他人不可读/下，admin 例外。"""
from io import BytesIO
import struct
import xml.etree.ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image
import pytest

from app import security
from app.agent import tools
from app.services import agent_files


def _ctx(user_id, role="sales"):
    return security.UserContext(user_id=user_id, role=role, is_authenticated=True)


def test_upload_records_owner_and_owns_logic(db):
    up = agent_files.save_upload(b"hello", "a.txt", "alice")
    fid = up["file_id"]
    assert agent_files.owner_of(fid) == "alice"
    assert tools._owns(_ctx("alice"), fid) is True            # 本人
    assert tools._owns(_ctx("bob"), fid) is False             # 他人
    assert tools._owns(_ctx("bob", role="admin"), fid) is True  # admin 例外
    assert tools._owns(_ctx("bob"), None) is True             # 无 file_id 放行


def test_read_tools_block_other_users_file(db):
    fid = agent_files.save_upload(b"secret quote", "q.txt", "alice")["file_id"]
    # 他人通过读工具访问 → 拒
    assert tools._read_document(db, {"file_id": fid}, _ctx("bob")) == tools._NO_ACCESS
    assert tools._inspect_file(db, {"file_id": fid}, _ctx("bob")) == tools._NO_ACCESS
    assert tools._read_file_rows(db, {"file_id": fid}, _ctx("bob")) == tools._NO_ACCESS
    assert tools._write_excel(db, {"base_file_id": fid}, _ctx("bob")) == tools._NO_ACCESS
    # 本人读自己的 txt 正常（非越权拒绝）
    assert tools._read_document(db, {"file_id": fid}, _ctx("alice")) != tools._NO_ACCESS


def test_readonly_cannot_read_others_file(db):
    """收紧（2026-06-15）：readonly 不再是文件全量角色——共享口令回退把非 admin 一律发成
    readonly，若放行会让任何知道 ADMIN_PASSWORD 的人凭 file_id 读他人上传的报价/合同(IDOR)。"""
    fid = agent_files.save_upload(b"secret quote", "q.txt", "alice")["file_id"]
    assert tools._owns(_ctx("bob", role="readonly"), fid) is False        # readonly 也需本人
    assert tools._read_document(db, {"file_id": fid}, _ctx("bob", role="readonly")) == tools._NO_ACCESS
    # admin / boss 仍可越权读（运维/审计），readonly 不行
    assert tools._owns(_ctx("bob", role="admin"), fid) is True
    assert tools._owns(_ctx("bob", role="boss"), fid) is True


def test_generated_report_owned_by_creator(db):
    out = tools._write_report(db, {"headers": ["型号"], "rows": [["PN-A"]], "title": "报价"},
                              _ctx("carol"))
    fid = out["file_id"]
    assert agent_files.owner_of(fid) == "carol"      # 生成文件归创建者 → 本人能下
    assert tools._owns(_ctx("carol"), fid) is True
    assert tools._owns(_ctx("dave"), fid) is False   # 他人不可下


def _template_with_image(payload: bytes, extension: str) -> bytes:
    """构造关系、媒体类型一致的真实模板；TIFF 绕过 openpyxl 构造时的 PNG 编码。"""
    with Image.open(BytesIO(payload)) as source:
        placeholder = BytesIO()
        Image.new("RGB", source.size).save(placeholder, format="PNG")
    wb = Workbook()
    wb.active["A1"] = "before"
    image = ExcelImage(BytesIO(payload if extension in {"png", "jpeg"} else placeholder.getvalue()))
    image.width, image.height = 100, 35
    wb.create_sheet("插图").add_image(image, "C4")
    original = BytesIO()
    wb.save(original)
    wb.close()
    if extension in {"png", "jpeg"}:
        return original.getvalue()

    output = BytesIO()
    with ZipFile(BytesIO(original.getvalue())) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            name, data = entry.filename, source.read(entry.filename)
            if name == "xl/media/image1.png":
                name, data = f"xl/media/image1.{extension}", payload
            elif name == "xl/drawings/_rels/drawing1.xml.rels":
                tree = ET.fromstring(data)
                for relation in tree:
                    if relation.attrib.get("Target", "").endswith("/image1.png"):
                        relation.set("Target", relation.attrib["Target"].removesuffix("png") + extension)
                data = ET.tostring(tree)
            elif name == "[Content_Types].xml":
                tree = ET.fromstring(data)
                ET.SubElement(tree, "{http://schemas.openxmlformats.org/package/2006/content-types}Default",
                              Extension=extension, ContentType=f"image/{extension}")
                data = ET.tostring(tree)
            target.writestr(name, data)
    return output.getvalue()


def _refill_image_template(template: bytes):
    uploaded = agent_files.save_upload(template, "template.xlsx", "alice")
    result = agent_files.write_excel(uploaded["file_id"], None,
                                    [{"row": 1, "col": 1, "value": "after"}], "filled.xlsx", "alice")
    original_path, _ = agent_files.get_download(uploaded["file_id"])
    assert original_path.read_bytes() == template
    output_path, _ = agent_files.get_download(result["file_id"])
    wb = load_workbook(output_path)
    assert wb.active["A1"].value == "after"
    assert len(wb["插图"]._images) == 1
    image = wb["插图"]._images[0]
    assert (image.anchor._from.row, image.anchor._from.col) == (3, 2)
    assert (image.anchor.ext.cx, image.anchor.ext.cy) == (952500, 333375)
    assert (image.width, image.height) == (10, 1)
    wb.close()
    return output_path, image


@pytest.mark.parametrize("mode", ["I", "I;32S", "F", "F-extreme"])
def test_write_excel_preserves_numeric_tiff_image(mode):
    values = [-2147483648, -1, 0, 1, 254, 255, 256, 65535, 65536, 2147483647]
    if mode == "I;32S":
        source = Image.frombytes("I", (10, 1), struct.pack("=10i", *values), "raw", mode)
    else:
        source = Image.new("F" if mode.startswith("F") else mode, (10, 1))
        if mode == "F":
            values = [-1.5, -1, 0, 0.5, 1.5, 254.9, 255.9, 256.1, 65535.9, 100000]
        elif mode == "F-extreme":
            values = [float("nan"), float("-inf"), -1e20, -1.5, 0.5, 65535.9,
                      2147483648.0, 1e20, float("inf"), 1.5]
        source.putdata(values)
    payload = BytesIO()
    source.save(payload, format="TIFF")
    with Image.open(BytesIO(payload.getvalue())) as check:
        assert check.mode == ("F" if mode.startswith("F") else "I")
    template = _template_with_image(payload.getvalue(), "tiff")
    output_path, image = _refill_image_template(template)
    assert image.format == "png"
    expected = ([0, 0, 0, 1, 254, 255, 256, 65535, 65535, 65535] if mode != "F"
                else [0, 0, 0, 0, 1, 254, 255, 256, 65535, 65535])
    if mode == "F-extreme":
        expected = [0, 0, 0, 0, 0, 65535, 65535, 65535, 65535, 1]
    with ZipFile(output_path) as archive:
        assert [name for name in archive.namelist() if name.startswith("xl/media/")] == ["xl/media/image1.png"]
        with Image.open(BytesIO(archive.read("xl/media/image1.png"))) as saved:
            assert saved.mode == "I;16"
            assert [saved.getpixel((x, 0)) for x in range(10)] == expected


@pytest.mark.parametrize("format", ["PNG", "JPEG"])
def test_write_excel_preserves_png_jpeg_image_bytes(format):
    payload = BytesIO()
    Image.new("RGB", (10, 1), (20, 100, 200)).save(payload, format=format)
    template = _template_with_image(payload.getvalue(), format.lower())
    output_path, image = _refill_image_template(template)
    assert image.format == format.lower()
    with ZipFile(output_path) as archive:
        saved_bytes = archive.read(f"xl/media/image1.{format.lower()}")
        assert saved_bytes == payload.getvalue()
        with Image.open(BytesIO(saved_bytes)) as saved:
            assert saved.mode == "RGB" and saved.format == format
