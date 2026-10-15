"""Project reminders use the controlled master interface, never return facts."""

from concurrent.futures import ThreadPoolExecutor
from threading import Barrier

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import event, func, select, text
from sqlalchemy.exc import DBAPIError
from alembic import command

from app.db import SessionLocal, engine
from app.main import app
from app.auth import _make_token
from app.models.maintenance_project import MaintenanceProject
from app.models.system import SysUser
from app.services import maintenance_project_catalog as catalog

from app.models.maintenance_project import MaintenanceProjectAuditLog
from app.models.maintenance_project_operations import MaintenanceProjectWorkbookState
from tests.test_maintenance_project_management_api import _admin_client, _permission_client
from tests.test_maintenance_project_foundation import _alembic_cfg


def test_note_patch_reads_back_and_records_one_versioned_change(db):
    client = _admin_client(db, "note_admin")
    project = client.post("/api/maintenance/projects/stable", json={
        "project_code": "NOTE-PROJECT", "display_name": "标注测试项目",
        "reason": "建立合成项目",
    }).json()
    pid = project["project_id"]
    changed = client.patch(f"/api/maintenance/projects/stable/{pid}", json={
        "version": 1, "no_return_note": "  硬盘免返\n其他按合同  ",
        "reason": "按合同补充提醒",
    })
    assert changed.status_code == 200, changed.text
    assert changed.json()["no_return_note"] == "硬盘免返\n其他按合同"
    assert changed.json()["version"] == 2
    assert client.get(f"/api/maintenance/projects/stable/{pid}").json()["project"]["no_return_note"] == "硬盘免返\n其他按合同"
    assert client.get("/api/maintenance/projects/stable").json()["rows"][0]["no_return_note"] == "硬盘免返\n其他按合同"
    state = db.scalar(select(MaintenanceProjectWorkbookState))
    assert state.revision == 1
    audit = db.scalar(select(MaintenanceProjectAuditLog).where(
        MaintenanceProjectAuditLog.action == "update"))
    assert audit.before_json["no_return_note"] is None
    assert audit.after_json["no_return_note"] == "硬盘免返\n其他按合同"
    assert audit.operated_by == "note_admin"
    assert audit.reason == "按合同补充提醒"
    assert client.patch(f"/api/maintenance/projects/stable/{pid}", json={
        "version": 1, "no_return_note": "陈旧覆盖", "reason": "旧客户端重试",
    }).status_code == 409


@pytest.mark.parametrize("second_field", ["no_return_note", "display_name"])
def test_concurrent_note_changes_refresh_preloaded_project_before_cas(db, second_field):
    client = _admin_client(db, "note_race_admin")
    pid = client.post("/api/maintenance/projects/stable", json={
        "project_code": "NOTE-RACE", "display_name": "并发提醒项目",
        "reason": "建立合成项目",
    }).json()["project_id"]
    barrier = Barrier(2, timeout=30)

    def save(updates):
        with SessionLocal() as session:
            # 两个调用者均持有旧 ORM 快照；取得行锁后仍必须读最新版本。
            project = session.get(MaintenanceProject, pid)
            assert project.version == 1
            barrier.wait()
            try:
                catalog.update_project(session, project_id=pid, version=1,
                    updates=updates, reason="并发更新提醒", operated_by="note_race_admin")
                session.commit()
                return "saved"
            except catalog.MaintenanceProjectCatalogConflict:
                session.rollback()
                return "conflict"

    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(save, [{"no_return_note": "甲的提醒"}, {second_field: "乙的提醒"}]))
    assert sorted(results) == ["conflict", "saved"]
    project = client.get(f"/api/maintenance/projects/stable/{pid}").json()["project"]
    assert project["version"] == 2
    assert project["no_return_note"] in ({"甲的提醒", "乙的提醒"} if second_field == "no_return_note" else {"甲的提醒", None})


def test_note_and_workbook_apply_keep_both_changes_and_bumps(db):
    import io
    from uuid import uuid4
    from openpyxl import load_workbook
    from app.services import maintenance_project_master_workbook as master
    from tests.test_maintenance_project_master_v2_editable import _make_project_with_line, _save, _site_issue

    project, part, _, _ = _make_project_with_line(db)
    _, line = _site_issue(db, project, part, qty="2")
    pid = project.project_id
    wb = load_workbook(io.BytesIO(master.build_project_master_v2(db, project_id=pid, sheets=(master.V2_SHEET_SITE,))))
    ws = wb[master.V2_SHEET_SITE]
    headers = {cell.value: cell.column for cell in ws[1]}
    ws.cell(2, headers["备注"]).value = "工作簿并发修改"
    db.commit()
    with SessionLocal() as validator:
        plan = master.validate_project_master_v2(validator, project_id=pid, data=_save(wb))
    initial_revision = db.scalar(select(MaintenanceProjectWorkbookState.revision))
    db.rollback()
    barrier = Barrier(2, timeout=30)

    def save_note():
        with SessionLocal() as session:
            barrier.wait()
            catalog.update_project(session, project_id=pid, version=1,
                updates={"no_return_note": "并发提醒"}, reason="并发标注", operated_by="synthetic_operator")
            session.commit()

    def apply_book():
        with SessionLocal() as session:
            barrier.wait()
            master.apply_project_master_v2(session, plan, operated_by="synthetic_operator", import_batch_id=str(uuid4()))
            session.commit()

    with ThreadPoolExecutor(max_workers=2) as pool:
        futures = [pool.submit(save_note), pool.submit(apply_book)]
        for future in futures:
            future.result(timeout=90)
    db.expire_all()
    assert db.get(MaintenanceProject, pid).no_return_note == "并发提醒"
    assert line.remark == "工作簿并发修改"
    assert db.scalar(select(MaintenanceProjectWorkbookState.revision)) == initial_revision + 2


def _project(db):
    client = _admin_client(db, "note_contract_admin")
    project = client.post("/api/maintenance/projects/stable", json={
        "project_code": "NOTE-CONTRACT", "display_name": "提醒契约项目",
        "reason": "建立合成项目",
    }).json()
    return client, project["project_id"]


def _patch(client, pid, version, **updates):
    return client.patch(f"/api/maintenance/projects/stable/{pid}", json={
        "version": version, "reason": "维护收货提醒", **updates,
    })


@pytest.mark.parametrize("clear", [None, "", " \n\t "])
def test_omitted_note_preserved_explicit_clear_and_noop_have_exact_increments(db, clear):
    client, pid = _project(db)
    assert _patch(client, pid, 1, no_return_note="原提醒").status_code == 200
    assert _patch(client, pid, 2, display_name="更名保留提醒").json()["no_return_note"] == "原提醒"
    assert _patch(client, pid, 3, no_return_note=" 原提醒 ").json()["version"] == 3
    cleared = _patch(client, pid, 3, no_return_note=clear)
    assert cleared.json()["no_return_note"] is None
    assert cleared.json()["version"] == 4
    assert _patch(client, pid, 4, no_return_note=clear).json()["version"] == 4
    db.expire_all()
    assert db.scalar(select(MaintenanceProjectWorkbookState.revision)) == 3
    assert db.scalar(select(func.count()).select_from(MaintenanceProjectAuditLog).where(
        MaintenanceProjectAuditLog.action == "update")) == 3


@pytest.mark.parametrize("value,status", [
    ("中" * 2000, 200), ("😀" * 2000, 200), ("中" * 2001, 422),
    (" " * 2001, 422), (17, 422), ({"text": "免返"}, 422),
])
def test_note_api_enforces_text_and_original_unicode_character_limit(db, value, status):
    client, pid = _project(db)
    response = _patch(client, pid, 1, no_return_note=value)
    assert response.status_code == status, response.text
    if status == 200:
        assert response.json()["no_return_note"] == value
    else:
        assert client.get(f"/api/maintenance/projects/stable/{pid}").json()["project"]["version"] == 1


def test_direct_service_also_rejects_invalid_type_and_raw_length(db):
    client, pid = _project(db)
    for value in ["中" * 2001, " " * 2001, 17, {"text": "免返"}]:
        with pytest.raises(catalog.MaintenanceProjectCatalogError):
            catalog.update_project(db, project_id=pid, version=1,
                updates={"no_return_note": value}, reason="测试服务门禁", operated_by="note_contract_admin")
        db.rollback()
    assert _patch(client, pid, 1, no_return_default=True).status_code == 422
    assert _patch(client, pid, 1, unknown_note="免返").status_code == 422


@pytest.mark.parametrize("missing", ["page_maintenance", "data_profit", "action_maintenance_project_manage"])
def test_note_keeps_existing_master_permission_matrix(db, missing):
    _, pid = _project(db)
    perms = {"page_maintenance": True, "page_maintenance_beta": True,
             "data_purchase_cost": True, "data_profit": True,
             "action_maintenance_project_manage": True, missing: False}
    client = _permission_client(db, username="note_limited", permissions=perms)
    assert _patch(client, pid, 1, no_return_note="越权提醒").status_code == 403
    if missing != "page_maintenance":
        assert client.get(f"/api/maintenance/projects/stable/{pid}").status_code == 200


def test_shared_identity_and_unowned_project_cannot_write_note(db):
    client, pid = _project(db)
    token, _ = _make_token("admin", "shared-admin", None,
                           fallback=True, perms={"action_maintenance_project_manage": True})
    shared = TestClient(app)
    shared.headers["Authorization"] = f"Bearer {token}"
    assert _patch(shared, pid, 1, no_return_note="共享身份写入").status_code == 403
    limited = _permission_client(db, username="note_own_only", permissions={
        "page_maintenance": True, "data_profit": True,
        "action_maintenance_project_manage": True, "own_maintenance_projects_only": True,
    })
    # boss 本来是全范围角色；改用真实销售角色核对项目行级范围。
    account = db.scalar(select(SysUser).where(SysUser.username == "note_own_only"))
    account.role = "sales"
    db.commit()
    login = limited.post("/api/auth/login", json={
        "username": "note_own_only", "password": "synthetic-password-123",
    })
    limited.headers["Authorization"] = f"Bearer {login.json()['token']}"
    assert _patch(limited, pid, 1, no_return_note="越项目写入").status_code == 403
    assert client.get(f"/api/maintenance/projects/stable/{pid}").json()["project"]["no_return_note"] is None


def test_note_is_independent_of_return_flags_quantities_costs_and_receipts(db):
    from app.services import maintenance_site_return_requirements as requirements
    from app.services.maintenance_return_receipts import receipt_summary, register_receipt
    from tests.test_maintenance_project_master_v2_editable import _make_project_with_line, _site_issue

    client = _admin_client(db, "note_fact_admin")
    project, part, _, _ = _make_project_with_line(db)
    issue, line = _site_issue(db, project, part, qty="3")
    line.no_return = False
    register_receipt(db, project_id=project.project_id, pn="ANOTHER-PN", qty=2,
                     operated_by="synthetic_receiver")
    db.commit()
    original_requirements = requirements.line_details(db, [line], project=project)
    original_receipts = receipt_summary(db, project_id=project.project_id)
    before = (line.no_return, line.quantity, line.cost_amount_inc_tax, line.version, issue.version,
              project.no_return_default)
    changed = _patch(client, project.project_id, project.version, no_return_note="全部免返")
    assert changed.status_code == 200, changed.text
    db.expire_all()
    assert requirements.line_details(db, [line], project=project) == original_requirements
    assert receipt_summary(db, project_id=project.project_id) == original_receipts
    assert (line.no_return, line.quantity, line.cost_amount_inc_tax, line.version, issue.version,
            project.no_return_default) == before


def test_note_and_viewers_and_name_are_one_revision_and_archived_note_is_rejected(db):
    client, pid = _project(db)
    _permission_client(db, username="note_viewer", permissions={"page_maintenance": True})
    changed = _patch(client, pid, 1, no_return_note="提醒", display_name="新名",
                     visible_usernames=["note_viewer"])
    assert changed.status_code == 200, changed.text
    assert changed.json()["version"] == 2
    assert [viewer["username"] for viewer in changed.json()["visible_usernames"]] == ["note_viewer"]
    assert db.scalar(select(MaintenanceProjectWorkbookState.revision)) == 1
    assert _patch(client, pid, 2, no_return_note="提醒", visible_usernames=["note_viewer"]).json()["version"] == 2
    assert client.post(f"/api/maintenance/projects/stable/{pid}/archive", json={"version": 2, "reason": "归档测试"}).status_code == 200
    assert _patch(client, pid, 3, no_return_note="归档后修改").status_code == 400


def test_note_audit_failure_rolls_back_master_and_workbook_revision(db):
    client, pid = _project(db)

    def fail_audit(_mapper, _connection, target):
        if target.action == "update":
            raise RuntimeError("synthetic audit storage failure")

    event.listen(MaintenanceProjectAuditLog, "before_insert", fail_audit)
    try:
        safe_client = TestClient(app, raise_server_exceptions=False)
        safe_client.headers.update(client.headers)
        assert _patch(safe_client, pid, 1, no_return_note="不能孤立落库").status_code == 500
    finally:
        event.remove(MaintenanceProjectAuditLog, "before_insert", fail_audit)
    project = client.get(f"/api/maintenance/projects/stable/{pid}").json()["project"]
    assert project["no_return_note"] is None
    assert project["version"] == 1
    assert db.scalar(select(MaintenanceProjectWorkbookState.revision)) is None
    assert db.scalar(select(func.count()).select_from(MaintenanceProjectAuditLog)) == 1


def test_note_migration_roundtrip_check_and_nonempty_downgrade_guard(db):
    client, pid = _project(db)
    db.close()
    cfg = _alembic_cfg()
    try:
        command.downgrade(cfg, "c8f3d1a7e2b9")
        with engine.connect() as connection:
            assert connection.scalar(text("SELECT count(*) FROM information_schema.columns WHERE table_name='maintenance_project' AND column_name='no_return_note'")) == 0
        command.upgrade(cfg, "head")
        command.check(cfg)
        assert _patch(client, pid, 1, no_return_note="已保存必须保留").status_code == 200
        with pytest.raises(DBAPIError, match="downgrade blocked"):
            command.downgrade(cfg, "c8f3d1a7e2b9")
        with engine.connect() as connection:
            assert connection.scalar(text("SELECT version_num FROM alembic_version")) == "f307604fa4d7"
            assert connection.scalar(text("SELECT no_return_note FROM maintenance_project WHERE project_id=:pid"), {"pid": pid}) == "已保存必须保留"
    finally:
        command.upgrade(cfg, "head")
