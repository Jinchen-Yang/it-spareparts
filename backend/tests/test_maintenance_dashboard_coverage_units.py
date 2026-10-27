"""覆盖/单位/无日期契约（coveragefix 轮 B-F06~B-F12）真实 PG 测试。

数字锚点来自独立审查 ``lane-b-independent-coverage-feedback.md`` 的真实反例
夹具（不是按实现回算）：
- B-F07：采购 3/4=75%、销售 3/4=75%、成本 1/4=25%；成本参考自身的领用
  已计覆盖 1/10=10% 是另一口径（issued_lines），不得冒充 25%。
- B-F09：件 2 已计（未税20/含税22.60）+ 台 3 未计，不能合成数量 5/覆盖40%。
- B-F11：全部未计成本的金额键不存在（不是 0.00）；真实 0 元已计保留 0.00。
- B-F12：无日期有效下单 7/单数 1 单列，不无声消失、不冒充入某天。
"""

import csv
import io
import uuid
from datetime import date, datetime, timezone
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models.maintenance import FMaintenanceLine, FMaintenanceOrder
from app.services import maintenance_dashboard as dash
from tests.test_maintenance_analytics import _line
from tests.test_maintenance_dashboard_api import (
    BASE,
    _assign_manager,
    _forbidden_keys,
    _part,
    _project,
    _purchase_ref,
    client_for,
)
from tests.test_maintenance_dashboard_valuation import _sale
from tests.test_maintenance_return_metrics import _issue, _receipt

TODAY = date(2026, 10, 21)
UTC = timezone.utc
RANGE = {"range": "custom", "date_from": "2026-10-16", "date_to": "2026-10-22"}


@pytest.fixture(autouse=True)
def _freeze_today(monkeypatch):
    monkeypatch.setattr(dash, "business_today", lambda: TODAY)


def _costed(db, project, part, qty, *, issue_date=date(2026, 10, 17),
            unit_cost_ex=Decimal("25.00"), cost_source="direct_purchase"):
    line = _issue(db, project, part, qty, issue_date=issue_date)
    unit_inc = (unit_cost_ex * Decimal("1.13")).quantize(Decimal("0.01"))
    line.unit_cost = unit_cost_ex
    line.unit_cost_ex_tax = unit_cost_ex
    line.unit_cost_inc_tax = unit_inc
    line.cost_amount = (Decimal(qty) * unit_cost_ex).quantize(Decimal("0.01"))
    line.cost_amount_ex_tax = line.cost_amount
    line.cost_amount_inc_tax = (Decimal(qty) * unit_inc).quantize(Decimal("0.01"))
    line.cost_source = cost_source
    db.flush()
    return line


def _raw_no(db, order_no: str) -> str:
    return db.scalar(select(FMaintenanceOrder.raw_order_id).where(
        FMaintenanceOrder.order_no == order_no))


def _extra_demand_line(db, first_line, part, qty, *, line_no):
    """同一需求单追加第二行（B-F10 整单选择夹具）。"""
    db.add(FMaintenanceLine(
        raw_line_id=f"rl-{uuid.uuid4()}", order_id=first_line.order_id,
        line_no=line_no, part_id=part.id, pn_std=part.pn_std,
        pn_raw=part.pn_std, description=part.description, qty=Decimal(qty),
        return_qty=Decimal("0"), is_active=True,
        import_batch_id=first_line.import_batch_id,
    ))
    db.flush()


def _csv_rows(content: bytes) -> list[list[str]]:
    return list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))


# ---------- B-F06 ----------

def test_bf06_all_pending_project_and_real_zero_ratio(db):
    """全待估项目不得 ready + 0.00；有合格参考的显式 ratio=0 是真实 0。"""
    pa = _project(db, "MDCV-P06A")
    no_unit = _part(db, "MDCV-P06-NOUNIT", unit=None)
    _costed(db, pa, no_unit, 2, unit_cost_ex=Decimal("175.00"))
    _receipt(db, pa, no_unit, 4, condition="坏品",
             occurred_at=datetime(2026, 10, 18, 8, tzinfo=UTC))
    pb = _project(db, "MDCV-P06B")
    zero_part = _part(db, "MDCV-P06-ZERO", unit="件")
    _purchase_ref(db, zero_part)
    _receipt(db, pb, zero_part, 2, condition="坏品",
             occurred_at=datetime(2026, 10, 18, 9, tzinfo=UTC))
    db.commit()

    client = client_for(db, username="mdcu-f06", role="admin")
    params = {**RANGE, "value_retention_ratio": "0"}
    rows = {
        row["project_code"]: row
        for row in client.get(f"{BASE}/projects", params=params).json()["rows"]
    }
    pa_val = rows["MDCV-P06A"]["valuation"]
    assert pa_val["state"] == "pending"
    assert pa_val["pending_bad_qty"] == "4.000"
    assert "bad_estimated_value_inc_tax" not in pa_val
    assert "bad_estimated_value_ex_tax" not in pa_val
    pb_val = rows["MDCV-P06B"]["valuation"]
    assert pb_val["state"] == "ready"
    assert pb_val["bad_estimated_value_inc_tax"] == "0.00"
    assert pb_val["bad_estimated_value_ex_tax"] == "0.00"
    assert pb_val["coverage_pct"] == 100.0

    export = client.get(
        f"{BASE}/export", params={"kind": "projects", "format": "csv", **params}
    )
    csv_rows = _csv_rows(export.content)
    header = csv_rows[0]
    code_col, inc_col = header.index("项目编码"), header.index("坏件试算价值-含税")
    by_code = {row[code_col]: row for row in csv_rows[1:]}
    assert by_code["MDCV-P06A"][inc_col] == ""      # 待估 → 空，不是 0.00
    assert by_code["MDCV-P06B"][inc_col] == "0.00"  # 显式 0 的真实 0 保留

    xlsx = client.get(
        f"{BASE}/export", params={"kind": "projects", "format": "xlsx", **params}
    )
    sheet = load_workbook(io.BytesIO(xlsx.content)).active
    xheader = [cell.value for cell in sheet[1]]
    x_code, x_inc = xheader.index("项目编码"), xheader.index("坏件试算价值-含税")
    x_cells = {
        row[x_code].value: row[x_inc].value for row in sheet.iter_rows(min_row=2)
    }
    assert x_cells["MDCV-P06A"] is None
    assert x_cells["MDCV-P06B"] == 0


# ---------- B-F07 ----------

def test_bf07_three_domain_return_valuation_coverage(db):
    """三域各自返件估值覆盖（审查夹具数字）：75/75/25 与 10% 分别命名。"""
    pv = _project(db, "MDCV-F07")
    part_m = _part(db, "MDCV-F07-M", unit="件")
    part_c = _part(db, "MDCV-F07-C", unit="件")
    _purchase_ref(db, part_m)  # 113.00 含 / 100.00 未
    _sale(db, part_m, "mdcu-f07-xs", price=Decimal("226.00"))  # 226 / 200
    _receipt(db, pv, part_m, 3, condition="坏品",
             occurred_at=datetime(2026, 10, 18, 8, tzinfo=UTC))
    _costed(db, pv, part_c, 1, unit_cost_ex=Decimal("175.00"))  # 175 / 197.75
    _issue(db, pv, part_c, 9, issue_date=date(2026, 10, 17))    # 未计 9
    _receipt(db, pv, part_c, 1, condition="坏品",
             occurred_at=datetime(2026, 10, 18, 9, tzinfo=UTC))
    db.commit()

    client = client_for(db, username="mdcu-f07", role="admin")
    card = client.get(f"{BASE}/summary", params=RANGE).json()["cards"]["valuation"]
    domains = card["domains"]
    assert domains["purchase"]["coverage_pct"] == 75.0
    assert domains["purchase"]["bad_qty"] == "4.000"
    assert domains["purchase"]["estimated_bad_qty"] == "3.000"
    assert domains["purchase"]["pending_bad_qty"] == "1.000"
    assert domains["purchase"]["coverage_basis"] == "bad_receipts"
    assert domains["sales"]["coverage_pct"] == 75.0
    assert domains["cost"]["coverage_pct"] == 25.0
    assert domains["cost"]["estimated_bad_qty"] == "1.000"
    assert domains["cost"]["pending_bad_qty"] == "3.000"
    # 域金额小计：采购 3×113×0.5；销售 3×226×0.5；成本 1×197.75×0.5=98.88。
    assert domains["purchase"]["value_inc_tax"] == "169.50"
    assert domains["sales"]["value_inc_tax"] == "339.00"
    assert domains["cost"]["value_inc_tax"] == "98.88"

    parts = client.get(f"{BASE}/parts", params=RANGE).json()["rows"]
    m_row = next(row for row in parts if row["pn"] == "MDCV-F07-M")
    c_row = next(row for row in parts if row["pn"] == "MDCV-F07-C")
    m_domains = m_row["valuation"]["domains"]
    assert m_domains["purchase"]["coverage_pct"] == 100.0
    assert m_domains["sales"]["coverage_pct"] == 100.0
    assert m_domains["cost"]["coverage_pct"] == 0.0
    assert m_domains["cost"]["pending_bad_qty"] == "3.000"
    assert "value_inc_tax" not in m_domains["cost"]  # 缺样域无金额键
    c_domains = c_row["valuation"]["domains"]
    assert c_domains["purchase"]["coverage_pct"] == 0.0
    assert c_domains["purchase"]["pending_bad_qty"] == "1.000"
    assert c_domains["cost"]["coverage_pct"] == 100.0
    assert c_domains["cost"]["value_inc_tax"] == "98.88"
    # 成本参考「领用已计覆盖」与成本域「返件估值覆盖」分别命名。
    cost_ref = c_row["references"]["cost"]
    assert cost_ref["coverage_pct"] == 10.0
    assert cost_ref["coverage_basis"] == "issued_lines"

    project_row = next(
        row for row in client.get(f"{BASE}/projects", params=RANGE).json()["rows"]
        if row["project_code"] == "MDCV-F07"
    )
    project_domains = project_row["valuation"]["domains"]
    assert project_domains["purchase"]["coverage_pct"] == 75.0
    assert project_domains["cost"]["coverage_pct"] == 25.0

    valuation_cards = client.get(f"{BASE}/valuation", params=RANGE).json()["cards"]
    assert valuation_cards["coverage_pct"] == 75.0
    card_domains = valuation_cards["domains"]
    assert card_domains["purchase"]["coverage_pct"] == 75.0
    assert card_domains["sales"]["coverage_pct"] == 75.0
    assert card_domains["cost"]["coverage_pct"] == 25.0

    username = "mdcu-f07-mgr"
    manager = client_for(db, username=username, role="maintenance_manager")
    _assign_manager(db, username, pv)
    for endpoint in ("/summary", "/projects", "/parts"):
        response = manager.get(f"{BASE}{endpoint}", params=RANGE)
        assert response.status_code == 200
        assert _forbidden_keys(response.json()) == set()
    assert manager.get(f"{BASE}/valuation", params=RANGE).status_code == 403


def test_bf07_bf09_mixed_unit_denominators_split(db):
    """件/台分母分开：平面数量/覆盖 null + unit_breakdown；金额可小计。"""
    pm = _project(db, "MDCV-MIX")
    part_j = _part(db, "MDCV-MIX-J", unit="件")
    part_t = _part(db, "MDCV-MIX-T", unit="台")
    _purchase_ref(db, part_j)
    _receipt(db, pm, part_j, 2, condition="坏品",
             occurred_at=datetime(2026, 10, 18, 8, tzinfo=UTC))
    _receipt(db, pm, part_t, 1, condition="坏品",
             occurred_at=datetime(2026, 10, 18, 9, tzinfo=UTC))
    db.commit()

    client = client_for(db, username="mdcu-f07mix", role="admin")
    summary = client.get(f"{BASE}/summary", params=RANGE).json()
    card = summary["cards"]["valuation"]
    assert card["coverage_pct"] is None
    assert card["estimated_bad_qty"] is None
    assert card["pending_bad_qty"] is None
    breakdown = {entry["unit"]: entry for entry in card["unit_breakdown"]}
    assert set(breakdown) == {"件", "台"}
    assert breakdown["件"]["bad_qty"] == "2.000"
    assert breakdown["件"]["coverage_pct"] == 100.0
    assert breakdown["台"]["bad_qty"] == "1.000"
    assert breakdown["台"]["coverage_pct"] == 0.0
    purchase = card["domains"]["purchase"]
    assert purchase["bad_qty"] is None
    assert purchase["coverage_pct"] is None
    assert purchase["value_inc_tax"] == "113.00"  # 2×113×0.5 跨单位小计
    domain_split = {entry["unit"]: entry for entry in purchase["unit_breakdown"]}
    assert domain_split["件"]["coverage_pct"] == 100.0
    assert domain_split["台"]["coverage_pct"] == 0.0
    # 数量卡同规则：混合 → 平面 null + 分桶。
    assert summary["cards"]["receipt_qty"]["total"] is None
    assert summary["cards"]["receipt_qty"]["unit_breakdown"][0]["unit"] == "件"
    trend_day = next(
        item for item in summary["trend"] if item["date"] == "2026-10-18"
    )
    assert trend_day["receipt_qty"] is None
    assert trend_day["unit_breakdown"][0]["receipt_qty"] == "2.000"
    condition = next(
        item for item in summary["condition_structure"]
        if item["test_result"] == "坏品"
    )
    assert condition["qty"] is None
    assert condition["unit_breakdown"][0]["qty"] == "2.000"

    project_row = next(
        row for row in client.get(f"{BASE}/projects", params=RANGE).json()["rows"]
        if row["project_code"] == "MDCV-MIX"
    )
    assert project_row["receipt_qty"] is None
    assert project_row["bad_receipt_qty"] is None
    assert project_row["issued_qty"] is None
    assert project_row["return_rate_pct"] is None
    assert project_row["unit"] is None
    project_split = {
        entry["unit"]: entry for entry in project_row["unit_breakdown"]
    }
    assert project_split["件"]["receipt_qty"] == "2.000"
    assert project_split["台"]["receipt_qty"] == "1.000"
    assert project_row["valuation"]["coverage_pct"] is None


# ---------- B-F08 ----------

def test_bf08_scoped_unassigned_receipts_kept(db):
    """受限但有 A 项目范围：A 的未关联 4 保留，无权项目不出现。"""
    pa = _project(db, "MDCV-F08A")
    px = _project(db, "MDCV-F08X")
    part = _part(db, "MDCV-F08-PN", unit="件")
    _receipt(db, pa, part, 4,
             occurred_at=datetime(2026, 10, 17, 8, tzinfo=UTC))
    _receipt(db, px, part, 7,
             occurred_at=datetime(2026, 10, 17, 9, tzinfo=UTC))
    db.commit()

    username = "mdcu-f08"
    manager = client_for(db, username=username, role="maintenance_manager")
    _assign_manager(db, username, pa)
    scoped = manager.get(f"{BASE}/quantity-checks", params=RANGE).json()
    assert scoped["cards"]["unassigned_receipt_qty"] == "4.000"
    assert [item["project_code"] for item in scoped["unassigned_receipts"]] == [
        "MDCV-F08A"
    ]
    assert scoped["unassigned_receipts"][0]["received_qty"] == "4.000"
    assert scoped["cards"]["checked_demand_qty"] == "0.000"  # 不摊给需求单

    full = client_for(db, username="mdcu-f08-admin", role="admin").get(
        f"{BASE}/quantity-checks", params=RANGE
    ).json()
    assert full["cards"]["unassigned_receipt_qty"] == "11.000"
    assert {item["project_code"] for item in full["unassigned_receipts"]} == {
        "MDCV-F08A", "MDCV-F08X"
    }


# ---------- B-F09 ----------

def test_bf09_consumption_unit_buckets_and_exports(db):
    """件 2 已计 100% / 台 3 未计 0%：不合 5/40%；导出按单位展开。"""
    pc = _project(db, "MDCV-F09")
    part_p = _part(db, "MDCV-F09-P", unit="件")
    part_t = _part(db, "MDCV-F09-T", unit="台")
    _costed(db, pc, part_p, 2, unit_cost_ex=Decimal("10.00"))
    _issue(db, pc, part_t, 3, issue_date=date(2026, 10, 17))
    db.commit()

    client = client_for(db, username="mdcu-f09", role="admin")
    body = client.get(f"{BASE}/consumption", params=RANGE).json()
    (row,) = body["rows"]
    assert row["issue_date"] == "2026-10-17"
    assert row["issued_qty"] is None
    assert row["costed_qty"] is None
    assert row["uncosted_qty"] is None
    assert row["coverage_pct"] is None
    assert row["line_count"] == 2
    assert row["cost_amount_ex_tax"] == "20.00"
    assert row["cost_amount_inc_tax"] == "22.60"
    split = {entry["unit"]: entry for entry in row["unit_breakdown"]}
    assert set(split) == {"件", "台"}
    assert split["件"]["issued_qty"] == "2.000"
    assert split["件"]["costed_qty"] == "2.000"
    assert split["件"]["uncosted_qty"] == "0.000"
    assert split["件"]["coverage_pct"] == 100.0
    assert split["件"]["cost_amount_inc_tax"] == "22.60"
    assert split["台"]["issued_qty"] == "3.000"
    assert split["台"]["costed_qty"] == "0.000"
    assert split["台"]["uncosted_qty"] == "3.000"
    assert split["台"]["coverage_pct"] == 0.0
    assert "cost_amount_inc_tax" not in split["台"]  # 全未计单位无金额键
    cards = body["cards"]
    assert cards["issued_qty"] is None
    assert cards["coverage_pct"] is None
    assert cards["uncosted_qty"] is None
    assert cards["cost_amount_inc_tax"] == "22.60"
    card_split = {entry["unit"]: entry for entry in cards["unit_breakdown"]}
    assert card_split["件"]["issued_qty"] == "2.000"
    assert card_split["台"]["issued_qty"] == "3.000"

    issues = client.get(
        f"{BASE}/consumption/issues",
        params={**RANGE, "issue_date": "2026-10-17"},
    ).json()["rows"]
    assert {row["pn"]: row["unit"] for row in issues} == {
        "MDCV-F09-P": "件", "MDCV-F09-T": "台"
    }

    export = client.get(
        f"{BASE}/export", params={"kind": "consumption_daily", "format": "csv", **RANGE}
    )
    csv_rows = _csv_rows(export.content)
    header = csv_rows[0]
    unit_col = header.index("单位")
    inc_col = header.index("已计成本-含税")
    coverage_col = header.index("已计成本覆盖率%")
    by_unit = {row[unit_col]: row for row in csv_rows[1:]}
    assert set(by_unit) == {"件", "台"}  # 按单位展开，无混合 40% 行
    assert by_unit["件"][inc_col] == "22.60"
    assert by_unit["件"][coverage_col] == "100.0"
    assert by_unit["台"][inc_col] == ""
    assert by_unit["台"][coverage_col] == "0.0"

    username = "mdcu-f09-mgr"
    manager = client_for(db, username=username, role="maintenance_manager")
    _assign_manager(db, username, pc)
    manager_body = manager.get(f"{BASE}/consumption", params=RANGE).json()
    assert _forbidden_keys(manager_body) == set()
    (manager_row,) = manager_body["rows"]
    assert manager_row["issued_qty"] is None
    for entry in manager_row["unit_breakdown"]:
        assert set(entry) == {"unit", "issued_qty"}  # 数量账号条目无成本键


def test_bf09_orders_mixed_unit_buckets(db):
    """M01 下单混合单位：order_qty null + 分桶；导出单位列「混合」。"""
    pd = _project(db, "MDCV-ORDMIX")
    part_j = _part(db, "MDCV-ORD-J", unit="件")
    part_t = _part(db, "MDCV-ORD-T", unit="台")
    _line(db, pd, part_j, order_no="MDCV-OM-1", order_date=date(2026, 10, 17),
          qty=Decimal("4"))
    _line(db, pd, part_t, order_no="MDCV-OM-2", order_date=date(2026, 10, 18),
          qty=Decimal("5"))
    db.commit()

    client = client_for(db, username="mdcu-ordmix", role="admin")
    body = client.get(f"{BASE}/orders", params=RANGE).json()
    (row,) = body["rows"]
    assert row["order_qty"] is None
    assert row["unit"] is None
    assert row["order_count"] == 2
    split = {entry["unit"]: entry for entry in row["unit_breakdown"]}
    assert split["件"]["order_qty"] == "4.000"
    assert split["台"]["order_qty"] == "5.000"
    assert body["cards"]["order_qty"] is None
    assert body["cards"]["order_count"] == 2

    export = client.get(
        f"{BASE}/export", params={"kind": "orders", "format": "csv", **RANGE}
    )
    csv_rows = _csv_rows(export.content)
    header = csv_rows[0]
    unit_col = header.index("单位")
    qty_col = header.index("期间下单量（原qty，不扣退货/免返）")
    # 静态审查补充：导出按 项目×单位 展开，数量不丢（不再输出 混合+空）。
    rows_data = csv_rows[1:]
    assert [row[unit_col] for row in rows_data] == ["件", "台"]
    assert rows_data[0][qty_col] == "4.000"
    assert rows_data[1][qty_col] == "5.000"


# ---------- B-F10 ----------

def test_bf10_quantity_checks_q_selects_whole_orders(db):
    """q=PN 只选含该 PN 的完整有效需求单；整单 N/R 不裁局部、不忽略 q。"""
    pq = _project(db, "MDCV-F10")
    part_x = _part(db, "MDCV-F10-X", unit="件")
    part_y = _part(db, "MDCV-F10-Y", unit="件")
    part_z = _part(db, "MDCV-F10-Z", unit="件")
    first = _line(db, pq, part_x, order_no="MDCV-F10-1",
                  order_date=date(2026, 10, 17), qty=Decimal("3"))
    _extra_demand_line(db, first, part_y, 7, line_no=2)  # 单1：X3+Y7 → N10
    _line(db, pq, part_x, order_no="MDCV-F10-2",
          order_date=date(2026, 10, 17), qty=Decimal("10"))
    _line(db, pq, part_y, order_no="MDCV-F10-3",
          order_date=date(2026, 10, 17), qty=Decimal("100"))  # 哨兵单3
    _receipt(db, pq, part_y, 12, condition="坏品",
             occurred_at=datetime(2026, 10, 17, 8, tzinfo=UTC),
             source_order_id=_raw_no(db, "MDCV-F10-1"))
    _receipt(db, pq, part_x, 5, condition="坏品",
             occurred_at=datetime(2026, 10, 17, 9, tzinfo=UTC),
             source_order_id=_raw_no(db, "MDCV-F10-2"))
    _receipt(db, pq, part_y, 99, condition="坏品",
             occurred_at=datetime(2026, 10, 17, 10, tzinfo=UTC),
             source_order_id=_raw_no(db, "MDCV-F10-3"))
    _receipt(db, pq, part_z, "4",
             occurred_at=datetime(2026, 10, 17, 11, tzinfo=UTC))  # 未关联
    db.commit()

    client = client_for(db, username="mdcu-f10", role="admin")
    filtered = client.get(
        f"{BASE}/quantity-checks", params={**RANGE, "q": "MDCV-F10-X"}
    ).json()
    assert {row["order_no"] for row in filtered["rows"]} == {
        "MDCV-F10-1", "MDCV-F10-2"
    }
    rows = {row["order_no"]: row for row in filtered["rows"]}
    assert rows["MDCV-F10-1"]["demand_qty"] == "10.000"  # 整单 N（X3+Y7）
    assert rows["MDCV-F10-1"]["received_qty"] == "12.000"
    assert rows["MDCV-F10-1"]["difference_qty"] == "0.000"
    assert rows["MDCV-F10-1"]["over_received_qty"] == "2.000"
    assert rows["MDCV-F10-2"]["difference_qty"] == "5.000"
    assert filtered["cards"]["checked_demand_qty"] == "20.000"
    assert filtered["cards"]["over_received_qty"] == "2.000"
    assert filtered["cards"]["unassigned_receipt_qty"] == "0.000"  # Z 不命中
    assert filtered["unassigned_receipts"] == []

    unfiltered = client.get(f"{BASE}/quantity-checks", params=RANGE).json()
    assert len(unfiltered["rows"]) == 3
    assert unfiltered["cards"]["checked_demand_qty"] == "120.000"
    assert unfiltered["cards"]["unassigned_receipt_qty"] == "4.000"

    export = client.get(
        f"{BASE}/export",
        params={"kind": "quantity_checks", "format": "csv", **RANGE,
                "q": "MDCV-F10-X"},
    )
    csv_rows = _csv_rows(export.content)
    no_col = csv_rows[0].index("需求单号")
    assert {row[no_col] for row in csv_rows[1:]} == {"MDCV-F10-1", "MDCV-F10-2"}


# ---------- B-F11 ----------

def test_bf11_all_uncosted_amounts_absent_real_zero_kept(db):
    """全部未计 → 金额键不存在；真实 0 元已计行输出 0.00。"""
    pu = _project(db, "MDCV-F11")
    part = _part(db, "MDCV-F11-PN", unit="件")
    _issue(db, pu, part, 3, issue_date=date(2026, 10, 17))  # 全未计
    _costed(db, pu, part, 2, issue_date=date(2026, 10, 18),
            unit_cost_ex=Decimal("0.00"))  # 真实 0 元已计
    db.commit()

    client = client_for(db, username="mdcu-f11", role="admin")
    body = client.get(f"{BASE}/consumption", params=RANGE).json()
    rows = {row["issue_date"]: row for row in body["rows"]}
    unknown = rows["2026-10-17"]
    assert unknown["uncosted_qty"] == "3.000"
    assert unknown["coverage_pct"] == 0.0
    assert "cost_amount_ex_tax" not in unknown
    assert "cost_amount_inc_tax" not in unknown
    zero = rows["2026-10-18"]
    assert zero["cost_amount_ex_tax"] == "0.00"
    assert zero["cost_amount_inc_tax"] == "0.00"
    assert zero["coverage_pct"] == 100.0
    cards = body["cards"]
    assert cards["cost_amount_ex_tax"] == "0.00"  # 已知小计（0 元已计在）
    assert cards["coverage_pct"] == 40.0


def test_bf11_cards_all_uncosted_window_amounts_absent(db):
    """整窗全未计：卡片金额键不存在（数量/未计/覆盖 0% 照给）。"""
    pu = _project(db, "MDCV-F11B")
    part = _part(db, "MDCV-F11B-PN", unit="件")
    _issue(db, pu, part, 3, issue_date=date(2026, 10, 17))
    db.commit()

    client = client_for(db, username="mdcu-f11b", role="admin")
    body = client.get(f"{BASE}/consumption", params=RANGE).json()
    cards = body["cards"]
    assert cards["issued_qty"] == "3.000"
    assert cards["uncosted_qty"] == "3.000"
    assert cards["coverage_pct"] == 0.0
    assert "cost_amount_ex_tax" not in cards
    assert "cost_amount_inc_tax" not in cards
    (row,) = body["rows"]
    assert "cost_amount_ex_tax" not in row

    export = client.get(
        f"{BASE}/export", params={"kind": "consumption_daily", "format": "csv", **RANGE}
    )
    csv_rows = _csv_rows(export.content)
    header = csv_rows[0]
    inc_col = header.index("已计成本-含税")
    (data_row,) = csv_rows[1:]
    assert data_row[inc_col] == ""


# ---------- B-F12 ----------

def test_bf12_undated_orders_bucket(db):
    """无日期有效下单 7/1 单列（未计入期间）；range=all 仍单列。"""
    po = _project(db, "MDCV-F12")
    part = _part(db, "MDCV-F12-PN", unit="件")
    _line(db, po, part, order_no="MDCV-F12-DATED",
          order_date=date(2026, 10, 17), qty=Decimal("5"))
    _line(db, po, part, order_no="MDCV-F12-NODATE",
          order_date=None, qty=Decimal("7"))
    db.commit()

    client = client_for(db, username="mdcu-f12", role="admin")
    body = client.get(f"{BASE}/orders", params=RANGE).json()
    assert body["cards"]["order_qty"] == "5.000"
    assert body["cards"]["order_count"] == 1
    undated = body["cards"]["undated_orders"]
    assert undated["order_qty"] == "7.000"
    assert undated["order_count"] == 1
    assert "未计入期间" in undated["note"]
    (row,) = body["rows"]
    assert row["order_qty"] == "5.000"
    assert row["undated_order_qty"] == "7.000"
    assert row["undated_order_count"] == 1

    # range=all：dated 全量（仍只含有日期单），undated 仍单列，不并入。
    all_body = client.get(f"{BASE}/orders", params={"range": "all"}).json()
    assert all_body["cards"]["order_qty"] == "5.000"
    assert all_body["cards"]["order_count"] == 1
    assert all_body["cards"]["undated_orders"]["order_qty"] == "7.000"
    assert all_body["cards"]["undated_orders"]["order_count"] == 1

    # v3：无日期卡同形单位桶（单一单位 → 平面照给 + 桶恰 1 条）。
    assert body["cards"]["undated_orders"]["unit_breakdown"] == [
        {"unit": "件", "order_qty": "7.000"}
    ]
    assert all_body["cards"]["undated_orders"]["unit_breakdown"] == [
        {"unit": "件", "order_qty": "7.000"}
    ]

    export = client.get(
        f"{BASE}/export", params={"kind": "orders", "format": "csv", **RANGE}
    )
    csv_rows = _csv_rows(export.content)
    header = csv_rows[0]
    assert "无制单日期下单量（未计入期间）" in header
    assert "单位" in header
    undated_col = header.index("无制单日期下单量（未计入期间）")
    (data_row,) = csv_rows[1:]
    assert data_row[undated_col] == "7.000"


# ---------- v3 契约同步（root contract-sync 轮）----------


def test_v3_summary_return_rate_breakdown_mixed_units(db):
    """v3 §5.2：混合单位率平面 null + 每单位分子/分母/率容器。

    数字锚点来自 root 同步要求：件 1/2=50%、台 1/3=33.3%；套只实收不领用
    （无分母 → 率 null，不是 0%）；只只领用不实收（真实 0%）。
    单位排序 = 名称升序（件 U+4EF6 < 只 U+53EA < 台 U+53F0 < 套 U+5957）。
    """
    project = _project(db, "MDCV-V3S")
    part_jian = _part(db, "MDCV-V3S-J", unit="件")
    part_zhi = _part(db, "MDCV-V3S-Z", unit="只")
    part_tai = _part(db, "MDCV-V3S-T", unit="台")
    part_tao = _part(db, "MDCV-V3S-A", unit="套")
    _receipt(db, project, part_jian, 1,
             occurred_at=datetime(2026, 10, 17, 8, tzinfo=UTC))
    _receipt(db, project, part_tai, 1,
             occurred_at=datetime(2026, 10, 17, 9, tzinfo=UTC))
    _receipt(db, project, part_tao, 5,
             occurred_at=datetime(2026, 10, 18, 8, tzinfo=UTC))
    _issue(db, project, part_jian, 2, issue_date=date(2026, 10, 17))
    _issue(db, project, part_tai, 3, issue_date=date(2026, 10, 17))
    _issue(db, project, part_zhi, 4, issue_date=date(2026, 10, 18))
    db.commit()

    admin = client_for(db, username="mdcu-v3s-admin", role="admin")
    cards = admin.get(f"{BASE}/summary", params=RANGE).json()["cards"]
    assert cards["return_rate_pct"] is None       # 混合单位：平面不可加总
    assert cards["rate_numerator"] is None
    assert cards["rate_denominator"] is None
    assert cards["return_rate_breakdown"] == [
        {"unit": "件", "rate_numerator": "1.000",
         "rate_denominator": "2.000", "return_rate_pct": 50.0},
        {"unit": "只", "rate_numerator": "0.000",
         "rate_denominator": "4.000", "return_rate_pct": 0.0},
        {"unit": "台", "rate_numerator": "1.000",
         "rate_denominator": "3.000", "return_rate_pct": 33.3},
        {"unit": "套", "rate_numerator": "5.000",
         "rate_denominator": "0.000", "return_rate_pct": None},
    ]

    # 数量账号：新键全在数量域，递归无金额泄漏（valuation 卡不出现）。
    username = "mdcu-v3s"
    manager = client_for(db, username=username, role="maintenance_manager")
    _assign_manager(db, username, project)
    scoped = manager.get(f"{BASE}/summary", params=RANGE).json()
    assert scoped["cards"]["return_rate_breakdown"] == [
        {"unit": "件", "rate_numerator": "1.000",
         "rate_denominator": "2.000", "return_rate_pct": 50.0},
        {"unit": "只", "rate_numerator": "0.000",
         "rate_denominator": "4.000", "return_rate_pct": 0.0},
        {"unit": "台", "rate_numerator": "1.000",
         "rate_denominator": "3.000", "return_rate_pct": 33.3},
        {"unit": "套", "rate_numerator": "5.000",
         "rate_denominator": "0.000", "return_rate_pct": None},
    ]
    assert "valuation" not in scoped["cards"]
    assert _forbidden_keys(scoped) == set()

    # 空范围（未授权项目）：容器 []，平面数量 0、无分母率 null。
    other = _project(db, "MDCV-V3E")
    db.commit()
    empty = manager.get(f"{BASE}/summary", params={
        **RANGE, "project": other.project_id}).json()
    assert empty["cards"]["return_rate_breakdown"] == []
    assert empty["cards"]["rate_numerator"] == "0.000"
    assert empty["cards"]["rate_denominator"] == "0.000"
    assert empty["cards"]["return_rate_pct"] is None


def test_v3_unassigned_receipts_unit_buckets(db):
    """v3 §3：未关联实收按单位分桶；混合平面 null；B sentinel 不出现。"""
    pa = _project(db, "MDCV-V3UA")   # 授权内：件 3 + 台 2（混合）
    pb = _project(db, "MDCV-V3UB")   # sentinel：不在授权内（件 5）
    pc = _project(db, "MDCV-V3UC")   # 授权内：件 4（单一单位兼容）
    part_jian = _part(db, "MDCV-V3U-J", unit="件")
    part_tai = _part(db, "MDCV-V3U-T", unit="台")
    _receipt(db, pa, part_jian, 3,
             occurred_at=datetime(2026, 10, 17, 8, tzinfo=UTC))
    _receipt(db, pa, part_tai, 2,
             occurred_at=datetime(2026, 10, 17, 9, tzinfo=UTC))
    _receipt(db, pb, part_jian, 5,
             occurred_at=datetime(2026, 10, 17, 10, tzinfo=UTC))
    _receipt(db, pc, part_jian, 4,
             occurred_at=datetime(2026, 10, 18, 8, tzinfo=UTC))
    db.commit()

    username = "mdcu-v3u"
    manager = client_for(db, username=username, role="maintenance_manager")
    _assign_manager(db, username, pa)
    _assign_manager(db, username, pc)
    scoped = manager.get(f"{BASE}/quantity-checks", params=RANGE).json()
    by_code = {row["project_code"]: row
               for row in scoped["unassigned_receipts"]}
    assert set(by_code) == {"MDCV-V3UA", "MDCV-V3UC"}  # B 不出现（交集不扩）
    assert by_code["MDCV-V3UA"]["received_qty"] is None  # 混合 → null
    assert by_code["MDCV-V3UA"]["unit_breakdown"] == [
        {"unit": "件", "received_qty": "3.000"},
        {"unit": "台", "received_qty": "2.000"},
    ]
    assert by_code["MDCV-V3UC"]["received_qty"] == "4.000"  # 单一单位照给
    assert by_code["MDCV-V3UC"]["unit_breakdown"] == [
        {"unit": "件", "received_qty": "4.000"},
    ]
    # 卡：跨项目混合（件 7 / 台 2）→ 平面 null + 按单位合计。
    assert scoped["cards"]["unassigned_receipt_qty"] is None
    assert scoped["cards"]["unassigned_receipt_unit_breakdown"] == [
        {"unit": "件", "received_qty": "7.000"},
        {"unit": "台", "received_qty": "2.000"},
    ]
    assert _forbidden_keys(scoped) == set()

    # 全范围：B 进入；平面仍 null（件 12 / 台 2 跨项目按单位）。
    full = client_for(db, username="mdcu-v3u-admin", role="admin").get(
        f"{BASE}/quantity-checks", params=RANGE).json()
    full_by_code = {row["project_code"]: row
                    for row in full["unassigned_receipts"]}
    assert set(full_by_code) == {"MDCV-V3UA", "MDCV-V3UB", "MDCV-V3UC"}
    assert full_by_code["MDCV-V3UB"]["received_qty"] == "5.000"
    assert full["cards"]["unassigned_receipt_qty"] is None
    assert full["cards"]["unassigned_receipt_unit_breakdown"] == [
        {"unit": "件", "received_qty": "12.000"},
        {"unit": "台", "received_qty": "2.000"},
    ]


def _purchase_ref_unit(db, part, unit):
    """同单位采购参考（_purchase_ref 固定 件；台件参考需 unit=台）。"""
    from decimal import Decimal
    import uuid as _uuid

    from app.models.purchase import FPurchaseLine, FPurchaseOrder
    from app.models.system import SysImportBatch

    batch = SysImportBatch(
        filename="mdcu-cg.xlsx", file_type="purchase",
        file_hash=_uuid.uuid4().hex, status="success",
    )
    db.add(batch)
    db.flush()
    order = FPurchaseOrder(
        raw_order_id=f"mdcu-cg-{_uuid.uuid4().hex[:8]}",
        order_no="MDCU-CG-1", purchaser="t", data_status="已生效",
        order_date=date(2026, 10, 18), tax_rate=Decimal("0.13"),
        is_tax_inclusive=True, import_batch_id=batch.id,
    )
    db.add(order)
    db.flush()
    db.add(FPurchaseLine(
        raw_line_id=f"mdcu-cg-l-{_uuid.uuid4().hex[:8]}", order_id=order.id,
        line_no=1, part_id=part.id, pn_std=part.pn_std, qty=Decimal(2),
        unit_price=Decimal("113.00"), unit=unit, import_batch_id=batch.id,
    ))
    db.flush()
    return order



# ---------- 静态审查补充（feedback §98-114） ----------

def test_bf_static_quantity_checks_unit_buckets(db):
    """整单 N/R/差额/超收按单位分桶（不延期）。

    混合单位（件3+台7）不得给 10 的确定平面合计；待核单位未证明可比。
    平面 null + unit_breakdown 原始量；差额/超收按单位求和不跨单位抵消。
    """
    pu = _project(db, "MDCV-ST1")
    part_x = _part(db, "MDCV-ST1-X", unit="件")
    part_t = _part(db, "MDCV-ST1-T", unit="台")
    part_n = _part(db, "MDCV-ST1-N")  # unit=None → 待核
    part_s = _part(db, "MDCV-ST1-S", unit="件")
    first = _line(db, pu, part_x, order_no="MDCV-ST1-1",
                  order_date=date(2026, 10, 17), qty=Decimal("3"))
    _extra_demand_line(db, first, part_t, 7, line_no=2)  # 单1：件3+台7
    _line(db, pu, part_n, order_no="MDCV-ST1-2",
          order_date=date(2026, 10, 17), qty=Decimal("6"))  # 单2：待核
    _line(db, pu, part_s, order_no="MDCV-ST1-3",
          order_date=date(2026, 10, 17), qty=Decimal("2"))  # 单3：件2
    raw1 = _raw_no(db, "MDCV-ST1-1")
    _receipt(db, pu, part_x, 4,
             occurred_at=datetime(2026, 10, 17, 8, tzinfo=UTC),
             source_order_id=raw1)
    _receipt(db, pu, part_t, 5,
             occurred_at=datetime(2026, 10, 17, 9, tzinfo=UTC),
             source_order_id=raw1)
    _receipt(db, pu, None, 2, pn="MDCV-ST1-RAW",
             occurred_at=None, source_order_id=raw1)  # 无日期未解析 → 待核
    _receipt(db, pu, part_n, 2,
             occurred_at=datetime(2026, 10, 17, 10, tzinfo=UTC),
             source_order_id=_raw_no(db, "MDCV-ST1-2"))
    _receipt(db, pu, part_s, 1,
             occurred_at=datetime(2026, 10, 17, 11, tzinfo=UTC),
             source_order_id=_raw_no(db, "MDCV-ST1-3"))
    db.commit()

    client = client_for(db, username="mdcu-st1", role="admin")
    body = client.get(f"{BASE}/quantity-checks", params=RANGE).json()
    rows = {row["order_no"]: row for row in body["rows"]}

    mixed = rows["MDCV-ST1-1"]
    for key in ("demand_qty", "received_qty", "difference_qty",
                "over_received_qty", "undated_receipt_qty"):
        assert mixed[key] is None, key
    assert mixed["unit"] is None
    assert mixed["unit_reliability"] == "mixed_units"
    buckets = {entry["unit"]: entry for entry in mixed["unit_breakdown"]}
    assert set(buckets) == {"件", "台", None}
    assert buckets["件"] == {
        "unit": "件", "demand_qty": "3.000", "received_qty": "4.000",
        "difference_qty": "0.000", "over_received_qty": "1.000",
        "undated_receipt_qty": "0.000",
    }
    assert buckets["台"] == {
        "unit": "台", "demand_qty": "7.000", "received_qty": "5.000",
        "difference_qty": "2.000", "over_received_qty": "0.000",
        "undated_receipt_qty": "0.000",
    }
    assert buckets[None]["received_qty"] == "2.000"
    assert buckets[None]["undated_receipt_qty"] == "2.000"
    assert buckets[None]["demand_qty"] == "0.000"

    # 仅待核桶：数量平面兼容照给（§0），unit=None + 未证明标注（不猜「件」）。
    unverified = rows["MDCV-ST1-2"]
    assert unverified["demand_qty"] == "6.000"
    assert unverified["received_qty"] == "2.000"
    assert unverified["difference_qty"] == "4.000"
    assert unverified["unit"] is None
    assert unverified["unit_reliability"] == "unverified_unit"
    (entry,) = unverified["unit_breakdown"]
    assert entry["unit"] is None
    assert entry["demand_qty"] == "6.000"
    assert entry["received_qty"] == "2.000"
    assert entry["difference_qty"] == "4.000"

    cards = body["cards"]
    for key in ("checked_demand_qty", "received_qty", "known_difference_qty",
                "over_received_qty"):
        assert cards[key] is None, key
    assert cards["unit_reliability"] == "mixed_units"
    card_buckets = {entry["unit"]: entry for entry in cards["unit_breakdown"]}
    assert card_buckets["件"]["checked_demand_qty"] == "5.000"  # 件跨单合法合计 3+2
    assert card_buckets["件"]["over_received_qty"] == "1.000"
    assert card_buckets["台"]["known_difference_qty"] == "2.000"
    assert card_buckets[None]["received_qty"] == "4.000"  # 待核跨单合计 2+2

    # 单一已证单位单照旧平面兼容（q 选择同参同集）。
    only3 = client.get(
        f"{BASE}/quantity-checks", params={**RANGE, "q": "MDCV-ST1-S"}
    ).json()
    row3 = only3["rows"][0]
    assert row3["demand_qty"] == "2.000"
    assert row3["received_qty"] == "1.000"
    assert row3["difference_qty"] == "1.000"
    assert row3["over_received_qty"] == "0.000"
    assert row3["unit"] == "件"
    assert row3["unit_reliability"] == "single_unit"

    # 导出按 单×单位 展开：混合单 3 行（件/台/待核），不丢待核量。
    export = client.get(
        f"{BASE}/export",
        params={"kind": "quantity_checks", "format": "csv",
                **RANGE, "q": "MDCV-ST1-X"},
    )
    csv_rows = _csv_rows(export.content)
    header = csv_rows[0]
    unit_col = header.index("单位")
    no_col = header.index("需求单号")
    demand_col = header.index("需求量N")
    over_col = header.index("超收max(R-N,0)")
    rows1 = [row for row in csv_rows[1:] if row[no_col] == "MDCV-ST1-1"]
    assert [row[unit_col] for row in rows1] == ["件", "台", "待核"]
    unit_jian = next(
        row for row in rows1 if row[unit_col] == "件")
    assert unit_jian[demand_col] == "3.000"
    assert unit_jian[over_col] == "1.000"
    unit_pending = rows1[2]
    assert unit_pending[demand_col] == "0.000"
    assert unit_pending[over_col] == "2.000"


def test_bf_static_projects_orders_export_expand_units(db):
    """projects/orders 导出按 项目×单位 展开：数量不丢失、单位金额用本桶值，
    去重单数只在首单位行出现一次（不可相加元信息）。"""
    pa = _project(db, "MDCV-ST2")
    part_a = _part(db, "MDCV-ST2-A", unit="件")
    part_b = _part(db, "MDCV-ST2-B", unit="台")
    _purchase_ref(db, part_a)
    _purchase_ref_unit(db, part_b, "台")
    _receipt(db, pa, part_a, 2, condition="坏品",
             occurred_at=datetime(2026, 10, 17, 8, tzinfo=UTC))
    _receipt(db, pa, part_b, 3, condition="坏品",
             occurred_at=datetime(2026, 10, 17, 9, tzinfo=UTC))
    _issue(db, pa, part_a, 1, issue_date=date(2026, 10, 18))
    # 混合单位需求单（件10+台4）+ 无日期单（件7）。
    first = _line(db, pa, part_a, order_no="MDCV-ST2-O1",
                  order_date=date(2026, 10, 17), qty=Decimal("10"))
    _extra_demand_line(db, first, part_b, 4, line_no=2)
    _line(db, pa, part_a, order_no="MDCV-ST2-O2",
          order_date=None, qty=Decimal("7"))
    db.commit()

    client = client_for(db, username="mdcu-st2", role="admin")

    projects = client.get(
        f"{BASE}/export", params={"kind": "projects", "format": "csv", **RANGE}
    )
    csv_rows = _csv_rows(projects.content)
    header = csv_rows[0]
    code_col = header.index("项目编码")
    unit_col = header.index("单位")
    receipt_col = header.index("返件实收合计")
    issued_col = header.index("实际领用（未采集件况）")
    value_col = header.index("坏件试算价值-含税")
    value_ex_col = header.index("坏件试算价值-不含税")
    rows_st2 = [row for row in csv_rows[1:] if row[code_col] == "MDCV-ST2"]
    assert [row[unit_col] for row in rows_st2] == ["件", "台"]
    row_jian, row_tai = rows_st2
    assert row_jian[receipt_col] == "2.000"
    assert row_jian[issued_col] == "1.000"
    assert row_tai[receipt_col] == "3.000"
    assert row_tai[issued_col] == "0.000"
    # 单位行金额 = 本桶试算值（件 2×113×0.5=113.00；台 3×113×0.5=169.50），
    # 不是父项总额重复。
    assert row_jian[value_col] == "113.00"
    assert row_jian[value_ex_col] == "100.00"
    assert row_tai[value_col] == "169.50"
    assert row_tai[value_ex_col] == "150.00"

    orders = client.get(
        f"{BASE}/export", params={"kind": "orders", "format": "csv", **RANGE}
    )
    csv_rows = _csv_rows(orders.content)
    header = csv_rows[0]
    code_col = header.index("项目编码")
    unit_col = header.index("单位")
    qty_col = header.index("期间下单量（原qty，不扣退货/免返）")
    count_col = header.index("期间下单单数")
    undated_col = header.index("无制单日期下单量（未计入期间）")
    undated_count_col = header.index("无制单日期单数")
    rows_st2 = [row for row in csv_rows[1:] if row[code_col] == "MDCV-ST2"]
    assert [row[unit_col] for row in rows_st2] == ["件", "台"]
    row_jian, row_tai = rows_st2
    assert row_jian[qty_col] == "10.000"
    assert row_tai[qty_col] == "4.000"
    # 去重单数不可按单位行相加：只在首单位行出现一次。
    assert row_jian[count_col] == "1"
    assert row_tai[count_col] == ""
    # 无日期单按单位给量与单数（同样只计一次）。
    assert row_jian[undated_col] == "7.000"
    assert row_jian[undated_count_col] == "1"
    assert row_tai[undated_col] == "0.000"
    assert row_tai[undated_count_col] == ""


def test_project_export_includes_period_orders_single_unit(db):
    """T17：项目对比导出补期间下单量/单数（单一单位，CSV 与 XLSX 都验）。"""
    project = _project(db, "MDCV-T17A")
    part = _part(db, "MDCV-T17A-PN", unit="件")
    _receipt(db, project, part, 2, condition="坏品",
             occurred_at=datetime(2026, 10, 18, 8, tzinfo=UTC))
    _line(db, project, part, order_no="MDCV-T17A-O1",
          order_date=date(2026, 10, 17), qty=Decimal("10"))
    _line(db, project, part, order_no="MDCV-T17A-O2",
          order_date=date(2026, 10, 18), qty=Decimal("5"))
    db.commit()
    client = client_for(db, username="mdcv-t17a", role="admin")

    order_row = next(
        row for row in client.get(f"{BASE}/orders", params=RANGE).json()["rows"]
        if row["project_id"] == project.project_id
    )
    assert order_row["order_qty"] == "15.000"
    assert order_row["order_count"] == 2

    export = client.get(
        f"{BASE}/export", params={"kind": "projects", "format": "csv", **RANGE}
    )
    rows = _csv_rows(export.content)
    header = rows[0]
    code_col = header.index("项目编码")
    unit_col = header.index("单位")
    qty_col = header.index("期间下单量（原qty）")
    count_col = header.index("期间下单单数")
    # 列顺序：单位 → 期间下单量 → 期间下单单数 → 返件实收合计。
    assert (qty_col, count_col, header.index("返件实收合计")) == (
        unit_col + 1, unit_col + 2, unit_col + 3
    )
    row = next(
        row for row in rows[1:] if row[code_col] == "MDCV-T17A"
    )
    assert row[unit_col] == "件"
    assert row[qty_col] == order_row["order_qty"]
    assert row[count_col] == str(order_row["order_count"])

    xlsx = client.get(
        f"{BASE}/export", params={"kind": "projects", "format": "xlsx", **RANGE}
    )
    sheet = load_workbook(io.BytesIO(xlsx.content)).active
    xheader = [cell.value for cell in sheet[1]]
    x_code = xheader.index("项目编码")
    x_unit = xheader.index("单位")
    x_qty = xheader.index("期间下单量（原qty）")
    x_count = xheader.index("期间下单单数")
    xrow = next(
        row for row in sheet.iter_rows(min_row=2)
        if row[x_code].value == "MDCV-T17A"
    )
    assert xrow[x_unit].value == "件"
    # XLSX 数值列写数值单元格（不是文本）。
    assert isinstance(xrow[x_qty].value, (int, float))
    assert xrow[x_qty].value == 15
    assert isinstance(xrow[x_count].value, (int, float))
    assert xrow[x_count].value == 2


def test_project_export_zero_order_columns_without_orders(db):
    """T17：无下单的项目，下单量 0、单数 0（不是空白）。"""
    project = _project(db, "MDCV-T17B")
    part = _part(db, "MDCV-T17B-PN", unit="件")
    _receipt(db, project, part, 2, condition="坏品",
             occurred_at=datetime(2026, 10, 18, 8, tzinfo=UTC))
    db.commit()
    client = client_for(db, username="mdcv-t17b", role="admin")

    export = client.get(
        f"{BASE}/export", params={"kind": "projects", "format": "csv", **RANGE}
    )
    rows = _csv_rows(export.content)
    header = rows[0]
    row = next(
        row for row in rows[1:]
        if row[header.index("项目编码")] == "MDCV-T17B"
    )
    assert row[header.index("期间下单量（原qty）")] == "0.000"
    assert row[header.index("期间下单单数")] == "0"


def test_project_export_orders_by_unit_and_order_only_unit(db):
    """T17：下单落同单位行；只在下单出现的单位追加一行，其他列留空；
    去重单数只在项目首行出现一次。"""
    project = _project(db, "MDCV-T17C")
    part_j = _part(db, "MDCV-T17C-J", unit="件")
    part_t = _part(db, "MDCV-T17C-T", unit="台")
    part_s = _part(db, "MDCV-T17C-S", unit="套")
    _receipt(db, project, part_j, 2, condition="坏品",
             occurred_at=datetime(2026, 10, 17, 8, tzinfo=UTC))
    _receipt(db, project, part_t, 3, condition="坏品",
             occurred_at=datetime(2026, 10, 17, 9, tzinfo=UTC))
    # 一张单 件10+台4（单数 1），另一张单 套7（单数 2）。
    first = _line(db, project, part_j, order_no="MDCV-T17C-O1",
                  order_date=date(2026, 10, 17), qty=Decimal("10"))
    _extra_demand_line(db, first, part_t, 4, line_no=2)
    _line(db, project, part_s, order_no="MDCV-T17C-O2",
          order_date=date(2026, 10, 18), qty=Decimal("7"))
    db.commit()
    client = client_for(db, username="mdcv-t17c", role="admin")

    export = client.get(
        f"{BASE}/export", params={"kind": "projects", "format": "csv", **RANGE}
    )
    rows = _csv_rows(export.content)
    header = rows[0]
    code_col = header.index("项目编码")
    unit_col = header.index("单位")
    qty_col = header.index("期间下单量（原qty）")
    count_col = header.index("期间下单单数")
    rows_c = [row for row in rows[1:] if row[code_col] == "MDCV-T17C"]
    assert [row[unit_col] for row in rows_c] == ["件", "台", "套"]
    row_j, row_t, row_s = rows_c
    assert row_j[header.index("返件实收合计")] == "2.000"
    assert row_t[header.index("返件实收合计")] == "3.000"
    assert row_j[qty_col] == "10.000"
    assert row_t[qty_col] == "4.000"
    assert row_s[qty_col] == "7.000"  # 只在下单出现的单位照填下单量
    # 去重单数不可相加：只在项目首行出现一次。
    assert row_j[count_col] == "2"
    assert row_t[count_col] == ""
    assert row_s[count_col] == ""
    # 追加行的其他数量列与金额列留空（不是 0）。
    for column in ("返件实收合计", "其中坏件实收",
                   "实际领用（未采集件况）",
                   "坏件试算价值-含税", "坏件试算价值-不含税"):
        assert row_s[header.index(column)] == ""


def test_project_export_orders_columns_without_amount_permission(db):
    """T17：无金额权账号仍导出两列数量；金额列头仍不存在。"""
    project = _project(db, "MDCV-T17D")
    part = _part(db, "MDCV-T17D-PN", unit="件")
    _receipt(db, project, part, 2, condition="坏品",
             occurred_at=datetime(2026, 10, 18, 8, tzinfo=UTC))
    _line(db, project, part, order_no="MDCV-T17D-O1",
          order_date=date(2026, 10, 18), qty=Decimal("6"))
    db.commit()
    username = "mdcv-t17d-mgr"
    client = client_for(db, username=username, role="maintenance_manager")
    _assign_manager(db, username, project)

    export = client.get(
        f"{BASE}/export", params={"kind": "projects", "format": "csv", **RANGE}
    )
    assert export.status_code == 200
    header = _csv_rows(export.content)[0]
    assert "期间下单量（原qty）" in header
    assert "期间下单单数" in header
    for forbidden in ("坏件试算价值-含税", "估价政策版本", "价值保留比例"):
        assert forbidden not in header
    row = next(
        row for row in _csv_rows(export.content)[1:]
        if row[header.index("项目编码")] == "MDCV-T17D"
    )
    assert row[header.index("期间下单量（原qty）")] == "6.000"
    assert row[header.index("期间下单单数")] == "1"


def test_project_export_empty_row_single_order_unit(db):
    """T17-M1：项目只有下单、无领用/返件时，首个下单单位直接填入空桶行。"""
    project = _project(db, "MDCV-T17E")
    part = _part(db, "MDCV-T17E-PN", unit="件")
    _line(db, project, part, order_no="MDCV-T17E-O1",
          order_date=date(2026, 10, 18), qty=Decimal("8"))
    db.commit()
    client = client_for(db, username="mdcv-t17e", role="admin")

    export = client.get(
        f"{BASE}/export", params={"kind": "projects", "format": "csv", **RANGE}
    )
    rows = _csv_rows(export.content)
    header = rows[0]
    code_col = header.index("项目编码")
    unit_col = header.index("单位")
    qty_col = header.index("期间下单量（原qty）")
    count_col = header.index("期间下单单数")
    rows_e = [row for row in rows[1:] if row[code_col] == "MDCV-T17E"]
    assert len(rows_e) == 1
    row = rows_e[0]
    assert row[unit_col] == "件"
    assert row[qty_col] == "8.000"
    assert row[count_col] == "1"
    # 空桶行原有项目对比数量列保持原值（0），不被下单流程改写。
    assert row[header.index("返件实收合计")] == "0.000"


def test_project_export_empty_row_two_order_units(db):
    """T17-M1：空桶 + 两个下单单位 → 两行，单数只在首行。"""
    project = _project(db, "MDCV-T17F")
    part_j = _part(db, "MDCV-T17F-J", unit="件")
    part_t = _part(db, "MDCV-T17F-T", unit="台")
    _line(db, project, part_j, order_no="MDCV-T17F-O1",
          order_date=date(2026, 10, 17), qty=Decimal("3"))
    _line(db, project, part_t, order_no="MDCV-T17F-O2",
          order_date=date(2026, 10, 18), qty=Decimal("4"))
    db.commit()
    client = client_for(db, username="mdcv-t17f", role="admin")

    export = client.get(
        f"{BASE}/export", params={"kind": "projects", "format": "csv", **RANGE}
    )
    rows = _csv_rows(export.content)
    header = rows[0]
    code_col = header.index("项目编码")
    unit_col = header.index("单位")
    qty_col = header.index("期间下单量（原qty）")
    count_col = header.index("期间下单单数")
    rows_f = [row for row in rows[1:] if row[code_col] == "MDCV-T17F"]
    assert [row[unit_col] for row in rows_f] == ["件", "台"]
    row_j, row_t = rows_f
    assert row_j[qty_col] == "3.000"
    assert row_t[qty_col] == "4.000"
    assert row_j[count_col] == "2"
    assert row_t[count_col] == ""


def test_project_export_empty_row_unresolved_order_unit(db):
    """T17-M1：空桶 + 下单单位为 None → 一行，单位列写「待核」。"""
    project = _project(db, "MDCV-T17G")
    part = _part(db, "MDCV-T17G-PN", unit=None)
    _line(db, project, part, order_no="MDCV-T17G-O1",
          order_date=date(2026, 10, 18), qty=Decimal("5"))
    db.commit()
    client = client_for(db, username="mdcv-t17g", role="admin")

    export = client.get(
        f"{BASE}/export", params={"kind": "projects", "format": "csv", **RANGE}
    )
    rows = _csv_rows(export.content)
    header = rows[0]
    code_col = header.index("项目编码")
    rows_g = [row for row in rows[1:] if row[code_col] == "MDCV-T17G"]
    assert len(rows_g) == 1
    assert rows_g[0][header.index("单位")] == "待核"
    assert rows_g[0][header.index("期间下单量（原qty）")] == "5.000"
    assert rows_g[0][header.index("期间下单单数")] == "1"


def test_b3_disjoint_known_rate_units_flat_null(db):
    """B3 独立验收 05:10（disjoint_known_rate_units）：只收件 2、只领台 3。

    两侧各自单桶、但单位并集跨 件+台 → 分子/分母跨单位不可比，
    平面率三字段 null（不得 66.7 / "2.000" / "3.000"）；两张数量卡
    保持各自真实单桶值；每单位率在 return_rate_breakdown 照实（M11/M12）。
    """
    project = _project(db, "MDCV-B3R")
    part_jian = _part(db, "MDCV-B3R-J", unit="件")
    part_tai = _part(db, "MDCV-B3R-T", unit="台")
    _receipt(db, project, part_jian, 2, condition="坏品",
             occurred_at=datetime(2026, 10, 17, 8, tzinfo=UTC))
    _issue(db, project, part_tai, 3, issue_date=date(2026, 10, 17))
    db.commit()

    body = client_for(db, username="mdcu-b3r-admin", role="admin").get(
        f"{BASE}/summary", params=RANGE).json()
    cards = body["cards"]
    # 数量卡：各自身单桶真实值（不互相污染、不因率不可比清零）。
    assert cards["receipt_qty"]["total"] == "2.000"
    assert cards["receipt_qty"]["unit_breakdown"] == [
        {"unit": "件", "total": "2.000", "part": "2.000", "machine": "0.000"}
    ]
    assert cards["issued_qty"]["total"] == "3.000"
    assert cards["issued_qty"]["unit_breakdown"] == [
        {"unit": "台", "total": "3.000"}
    ]
    # 平面率三字段：单位并集（件+台）跨单位 → 全 null。
    assert cards["return_rate_pct"] is None
    assert cards["rate_numerator"] is None
    assert cards["rate_denominator"] is None
    # 每单位率容器照实：件 无分母 null；台 真实 0%。
    assert cards["return_rate_breakdown"] == [
        {"unit": "件", "rate_numerator": "2.000",
         "rate_denominator": "0.000", "return_rate_pct": None},
        {"unit": "台", "rate_numerator": "0.000",
         "rate_denominator": "3.000", "return_rate_pct": 0.0},
    ]


def test_b3_export_three_domain_dual_tax_amounts(db):
    """B3 05:12（independently_averaged_cost_tax）：金额导出三域各给
    含税/不含税双列（M10/R34-06 双税分别计算，不从含税反推）；混合单位
    项目展开用**本单位**域值，不复制父域总额；quantity 角色无金额列。"""
    from decimal import Decimal as D

    pm = _project(db, "MDCV-B3X")
    part_all = _part(db, "MDCV-B3X-A", unit="件")   # 三域全有参考
    _purchase_ref(db, part_all)                       # 113.00 / 100.00
    _sale(db, part_all, "mdcu-b3x-xs", price=D("226.00"))  # 226.00 / 200.00
    _costed(db, pm, part_all, 2, unit_cost_ex=D("175.00"))  # 197.75 / 175.00
    part_tai = _part(db, "MDCV-B3X-T", unit="台")    # 无参考：三域全待估
    _receipt(db, pm, part_all, 2, condition="坏品",
             occurred_at=datetime(2026, 10, 18, 8, tzinfo=UTC))
    _receipt(db, pm, part_tai, 3, condition="坏品",
             occurred_at=datetime(2026, 10, 18, 9, tzinfo=UTC))
    db.commit()

    client = client_for(db, username="mdcu-b3x-admin", role="admin")
    # 锚点（默认比例 0.5）：2×113×0.5、2×226×0.5、2×197.75×0.5，
    # 不含税侧独立从各自未税价计算（不从含税反推）。
    api_domains = next(
        row for row in client.get(
            f"{BASE}/parts", params=RANGE).json()["rows"]
        if row["pn"] == "MDCV-B3X-A")["valuation"]["domains"]
    assert api_domains["purchase"]["value_inc_tax"] == "113.00"
    assert api_domains["purchase"]["value_ex_tax"] == "100.00"
    assert api_domains["sales"]["value_inc_tax"] == "226.00"
    assert api_domains["sales"]["value_ex_tax"] == "200.00"
    assert api_domains["cost"]["value_inc_tax"] == "197.75"
    assert api_domains["cost"]["value_ex_tax"] == "175.00"

    # parts CSV：三域各双税列 ↔ API 同值。
    export = client.get(
        f"{BASE}/export", params={"kind": "parts", "format": "csv", **RANGE})
    csv_rows = _csv_rows(export.content)
    header = csv_rows[0]
    pn_col = header.index("PN")
    (part_row,) = [r for r in csv_rows[1:] if r[pn_col] == "MDCV-B3X-A"]
    for label, expect in (
        ("三域-采购-坏件价值-含税", "113.00"),
        ("三域-采购-坏件价值-不含税", "100.00"),
        ("三域-销售-坏件价值-含税", "226.00"),
        ("三域-销售-坏件价值-不含税", "200.00"),
        ("三域-领用已计成本-价值-含税", "197.75"),
        ("三域-领用已计成本-价值-不含税", "175.00"),
    ):
        col = header.index(label)
        assert part_row[col] == expect, label

    # parts XLSX：同列同值（数值单元格）。
    xlsx = load_workbook(io.BytesIO(client.get(
        f"{BASE}/export",
        params={"kind": "parts", "format": "xlsx", **RANGE}).content)).active
    xheader = [cell.value for cell in xlsx[1]]
    x_pn = xheader.index("PN")
    (x_row,) = [r for r in xlsx.iter_rows(min_row=2)
                if r[x_pn].value == "MDCV-B3X-A"]
    for label, expect in (
        ("三域-采购-坏件价值-不含税", "100.00"),
        ("三域-销售-坏件价值-不含税", "200.00"),
        ("三域-领用已计成本-价值-不含税", "175.00"),
    ):
        cell = x_row[xheader.index(label)].value
        assert D(str(cell)) == D(expect), label

    # receipts CSV：行级三域双税；无参考行（台）待估 → 空，不是 0.00。
    r_export = client.get(
        f"{BASE}/export", params={"kind": "receipts", "format": "csv", **RANGE})
    r_rows = _csv_rows(r_export.content)
    r_header = r_rows[0]
    r_pn = r_header.index("PN")
    r_a, = [r for r in r_rows[1:] if r[r_pn] == "MDCV-B3X-A"]
    r_t, = [r for r in r_rows[1:] if r[r_pn] == "MDCV-B3X-T"]
    for label, expect in (
        ("三域-采购-行价值-含税", "113.00"),
        ("三域-采购-行价值-不含税", "100.00"),
        ("三域-销售-行价值-含税", "226.00"),
        ("三域-销售-行价值-不含税", "200.00"),
        ("三域-领用已计成本-价值-含税", "197.75"),
        ("三域-领用已计成本-价值-不含税", "175.00"),
    ):
        col = r_header.index(label)
        assert r_a[col] == expect, label
        assert r_t[col] == "", label          # 待估 → 空格

    # projects（混合单位 件+台 展开）：本单位域值；待估单位不得复制父总额。
    p_export = client.get(
        f"{BASE}/export", params={"kind": "projects", "format": "csv", **RANGE})
    p_rows = _csv_rows(p_export.content)
    p_header = p_rows[0]
    p_code, p_unit = p_header.index("项目编码"), p_header.index("单位")
    by_unit = {r[p_unit]: r for r in p_rows[1:] if r[p_code] == "MDCV-B3X"}
    for label, jian_expect in (
        ("三域-采购-坏件价值-含税", "113.00"),
        ("三域-采购-坏件价值-不含税", "100.00"),
        ("三域-销售-坏件价值-含税", "226.00"),
        ("三域-销售-坏件价值-不含税", "200.00"),
        ("三域-领用已计成本-价值-含税", "197.75"),
        ("三域-领用已计成本-价值-不含税", "175.00"),
    ):
        col = p_header.index(label)
        assert by_unit["件"][col] == jian_expect, label
        assert by_unit["台"][col] == "", label   # 无已估本桶 → 空，非父值

    # projects XLSX：同展开同值/同空。
    p_xlsx = load_workbook(io.BytesIO(client.get(
        f"{BASE}/export",
        params={"kind": "projects", "format": "xlsx", **RANGE}).content)).active
    px_header = [cell.value for cell in p_xlsx[1]]
    px_code, px_unit = px_header.index("项目编码"), px_header.index("单位")
    px_by_unit = {r[px_unit].value: r for r in p_xlsx.iter_rows(min_row=2)
                  if r[px_code].value == "MDCV-B3X"}
    for label, expect in (
        ("三域-采购-坏件价值-不含税", "100.00"),
        ("三域-销售-坏件价值-不含税", "200.00"),
        ("三域-领用已计成本-价值-不含税", "175.00"),
    ):
        col = px_header.index(label)
        assert D(str(px_by_unit["件"][col].value)) == D(expect), label
        assert px_by_unit["台"][col].value is None, label

    # quantity 角色：无任何金额/三域列（列头都没有）。
    username = "mdcu-b3x-mgr"
    manager = client_for(db, username=username, role="maintenance_manager")
    _assign_manager(db, username, pm)
    m_rows = _csv_rows(manager.get(
        f"{BASE}/export", params={"kind": "parts", "format": "csv", **RANGE}
    ).content)
    assert not any("三域" in h or "价值" in h for h in m_rows[0])
