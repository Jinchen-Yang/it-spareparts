"""P2-1: source status constrains effective views, while history survives."""
import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.config import get_settings
from app.models.dimensions import DimPart
from app.models.maintenance import FMaintenanceLine, FMaintenanceOrder
from app.models.maintenance_project import MaintenanceProject
from app.models.maintenance_source_assignment import MaintenanceSourceOrderAssignment
from app.services import maintenance_analytics
from app.services import maintenance_project_master_workbook as master
from app.services import maintenance_project_workbook_v3 as workbook_v3
from tests.boss_board_helpers import assign, client_for, import_wbdd, make_project
from tests.test_maintenance_return_metrics import _issue, _receipt
from tests.test_maintenance_wbdd_import_api import _upload
from tests.test_maintenance_project_workbook_v3 import _seed as seed_v3
from tests.wbdd_fixtures import COLUMNS_91, make_rows, write_workbook


@pytest.mark.parametrize("cutover", [False, True])
@pytest.mark.parametrize("status", ["已取消", "草稿", "进行中", "", None])
def test_non_effective_header_exits_rows_exports_and_board_but_keeps_history(
    db, tmp_path, monkeypatch, cutover, status,
):
    monkeypatch.setattr(get_settings(), "maintenance_cutover_enabled", cutover)
    monkeypatch.setattr(get_settings(), "maintenance_project_master_v2_enabled", False)
    project = make_project(db)
    order, = import_wbdd(db, tmp_path)
    assign(db, order, project)
    line = db.scalar(select(FMaintenanceLine))
    line.unit_cost_ex_tax = Decimal("100")
    line.unit_cost_inc_tax = Decimal("113")
    line.cost_amount_ex_tax = Decimal("300")
    line.cost_amount_inc_tax = Decimal("339")
    line.cost_source = "manual"
    db.commit()
    client = client_for(db, username="effective-read-admin", role="admin")
    rows_url = f"/api/maintenance/projects/stable/{project.project_id}/master-workbook/rows"
    assert client.get(rows_url, params={"sheet": master.SHEET_PARTS}).json()["total"] == 1

    order.data_status = status
    db.commit()
    rows = client.get(rows_url, params={"sheet": master.SHEET_PARTS})
    assert rows.status_code == 200, rows.text
    assert rows.json()["total"] == 0

    workbook = load_workbook(io.BytesIO(master.build_project_master(
        db, project_id=project.project_id,
        sheets=(master.SHEET_OVERVIEW, master.SHEET_PARTS),
    )))
    assert workbook[master.SHEET_PARTS].max_row == 1
    metrics = {row[0]: row[1] for row in workbook[master.SHEET_OVERVIEW].values if row[0]}
    assert metrics["项目已计成本(含税)"] == 0
    assert metrics["缺失成本行数"] == 0

    monkeypatch.setattr(get_settings(), "maintenance_project_master_v2_enabled", True)
    v2_rows = client.get(rows_url, params={"sheet": master.V2_SHEET_PARTS})
    assert v2_rows.status_code == 200, v2_rows.text
    assert v2_rows.json()["total"] == 0
    v2_book = load_workbook(io.BytesIO(master.build_project_master_v2(
        db, project_id=project.project_id, sheets=(master.V2_SHEET_PARTS,),
    )))
    v2_sheet = v2_book[master.V2_SHEET_PARTS]
    # V2 always retains blank CREATE rows and an example; no exported entity
    # may survive the status filter.
    entity_column = master.V2_PART_HEADERS.index("实体ID") + 1
    assert all(row[0] is None for row in v2_sheet.iter_rows(
        min_row=2, min_col=entity_column, max_col=entity_column, values_only=True,
    ))
    monkeypatch.setattr(get_settings(), "maintenance_project_master_v2_enabled", False)

    global_book = load_workbook(io.BytesIO(master.build_global_lines(
        db, preset="custom", date_from=order.order_date, date_to=order.order_date,
    )))
    assert global_book.active.max_row == 1
    board = client.get(f"/api/maintenance/boss-board/projects/{project.project_id}/orders")
    assert board.status_code == 200, board.text
    assert board.json()["total"] == 0

    history = client.post("/api/maintenance/demands/search", json={"q": order.order_no})
    assert history.status_code == 200, history.text
    assert history.json()["total"] == 1
    db.refresh(line)
    assert line.is_active is True
    assert line.cost_amount_inc_tax == Decimal("339")

    # A source status can become effective again; system VOID remains dominant.
    order.data_status = "已生效"
    db.commit()
    assert client.get(rows_url, params={"sheet": master.SHEET_PARTS}).json()["total"] == 1
    line.is_active = False
    db.commit()
    assert client.get(rows_url, params={"sheet": master.SHEET_PARTS}).json()["total"] == 0


def test_cancelled_upload_with_bad_child_keeps_actual_issue_and_receipt(db, tmp_path, monkeypatch):
    monkeypatch.setattr(get_settings(), "maintenance_cutover_enabled", False)
    client = client_for(db, username="cancelled-upload-admin", role="admin")
    rows = make_rows(orders=1, lines_per_order=1)
    rows[0]["销售订单"] = ""
    first = write_workbook(str(tmp_path / "effective.xlsx"), COLUMNS_91, rows)
    assert _upload(client, first).status_code == 200
    order = db.scalar(select(FMaintenanceOrder))
    line = db.scalar(select(FMaintenanceLine))
    assignment = db.scalar(select(MaintenanceSourceOrderAssignment).where(
        MaintenanceSourceOrderAssignment.source_order_id == order.raw_order_id,
        MaintenanceSourceOrderAssignment.is_active.is_(True),
    ))
    if assignment is None:
        project = make_project(db)
        assign(db, order, project)
    else:
        project = db.get(MaintenanceProject, assignment.project_id)
    part = db.get(DimPart, line.part_id)
    site_line = _issue(db, project, part, "2", source_order_id=order.raw_order_id)
    site_line.source_line_id = line.raw_line_id
    receipt = _receipt(db, project, part, "1", source_order_id=order.raw_order_id)
    db.commit()

    rows[0]["数据状态"] = "已取消"
    rows[0]["需求明细.需供货产品"] = ""
    second = write_workbook(str(tmp_path / "cancelled.xlsx"), COLUMNS_91, rows)
    response = _upload(client, second)
    assert response.status_code == 200, response.text
    result = client.get(
        f"/api/maintenance/projects/stable/{project.project_id}/master-workbook/rows",
        params={"sheet": master.SHEET_PARTS},
    )
    assert result.status_code == 200, result.text
    assert result.json()["total"] == 0
    summary = maintenance_analytics.pn_ranking(db, range_="all", can_cost=True)["summary"]
    assert Decimal(summary["total_effective_qty"]) == 0
    assert Decimal(summary["total_issued_qty"]) == 2
    assert Decimal(summary["total_receipt_qty"]) == 1
    db.refresh(site_line)
    db.refresh(receipt)
    assert site_line.is_active is True
    assert receipt.line_status == "active"


def test_v3_export_uses_effective_headers_and_preserves_history(db, monkeypatch):
    monkeypatch.setattr(get_settings(), "maintenance_cutover_enabled", False)
    seed_v3(db)
    order = db.scalar(select(FMaintenanceOrder))
    order.data_status = "已取消"
    db.commit()
    exported = load_workbook(io.BytesIO(workbook_v3.build_project_workbook(db, "wb3-project-1")))
    assert order.order_no not in {row[0] for row in exported["03_备件订单"].values}
    assert db.scalar(select(FMaintenanceLine)).is_active is True
    order.data_status = "已生效"
    db.commit()
    effective = load_workbook(io.BytesIO(workbook_v3.build_project_workbook(db, "wb3-project-1")))
    assert order.order_no in {row[0] for row in effective["03_备件订单"].values}
