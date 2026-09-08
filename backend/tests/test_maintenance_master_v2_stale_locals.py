"""04 / 05 / 06 三路合并后仍用合并前解析的局部变量 → 静默回退他人改动（2026-09-08 审查 P1）。

`_v2_merge_row` 返回**已 rebase 的整行**：用户没碰的基线字段换成服务端现值，下游拿整行
与服务端逐一比对时就不会拿导出时的旧值盖掉别人的改动（D-02）。03_备件明细 用对了
（`tests/test_maintenance_master_v2_row_merge.py::test_untouched_field_not_reverted_when_server_changed_it`），
另外三张表没有：

- **04_费用报销**、**06_领用返还**：`if not _v2_merge_row(...): continue` —— 返回的 rebase 行
  被**整个丢弃**，下游既内联读原始 `row`，又用合并前解析出的局部量。
- **05_实收回款**：有 `row = merged_row`，但 `amount`（累计实收金额）是在合并**之前**解析的，
  构造 `CollectionOp` 时用的还是旧值。同一构造里的凭证号 / 备注 / 状态是合并后读的，所以
  只有钱是错的。

`merge.guard` 救不了：它拿 **server_values**（新值）登记，应用期 CAS 比对自然通过，
然后把旧值写进去。用户侧表现是「我只改了个备注」，结果同事改过的金额 / 数量 / 报销人 /
日期被悄悄改回下载时的样子，预览不报冲突、不给提示。

`_v2_merge_row` 的 docstring 此前明确写着 04/06「不在本次修复范围」，且错误地宣称 05 已覆盖。
本次三处一起修，docstring 同步改成正面规则。
"""

import io
import uuid
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from app.services import maintenance_project_master_workbook as master
from app.services import maintenance_project_operations as operations
from tests.test_maintenance_project_master_v2_editable import (
    _make_project_with_expense,
    _make_project_with_line,
    _save,
    _site_issue,
)


def _download(db, project_id: str, sheet: str):
    content = master.build_project_master_v2(
        db, project_id=project_id, sheets=(sheet,))
    workbook = load_workbook(io.BytesIO(content))
    worksheet = workbook[sheet]
    headers = {cell.value: cell.column for cell in worksheet[1]}
    return workbook, worksheet, headers


def _row_for(worksheet, headers, entity_id) -> int:
    column = headers["实体ID"]
    return next(
        row_no
        for row_no in range(2, worksheet.max_row + 1)
        if worksheet.cell(row_no, column).value == entity_id
    )


def _apply(db, project_id, workbook, *, operated_by="stale-locals-test"):
    plan = master.validate_project_master_v2(
        db, project_id=project_id, data=_save(workbook))
    result = master.apply_project_master_v2(
        db, plan, operated_by=operated_by, import_batch_id=str(uuid.uuid4()))
    db.flush()
    return plan, result


# ---------- 04_费用报销 ----------

def test_expense_untouched_amount_is_not_reverted_by_a_remark_only_edit(db):
    """同事把未税金额 500 → 900；用户拿旧文件只改备注上传，900 必须留住。"""

    project, _order, _line, expense = _make_project_with_expense(db)
    workbook, worksheet, headers = _download(
        db, project.project_id, master.V2_SHEET_EXPENSE)
    row_no = _row_for(worksheet, headers, expense.raw_line_id)

    expense.amount = Decimal("900.00")
    expense.amount_ex_tax = Decimal("900.00")
    expense.amount_inc_tax = Decimal("1017.00")
    expense.person = "同事改的报销人"
    operations.bump_workbook_revision(db, project_id=project.project_id)
    db.commit()

    worksheet.cell(row_no, headers["操作"], "UPDATE")
    worksheet.cell(row_no, headers["备注"], "只改了个备注")

    plan, _result = _apply(db, project.project_id, workbook)
    db.refresh(expense)

    assert expense.amount_ex_tax == Decimal("900.00"), (
        f"只改备注却把未税金额回退成了 {expense.amount_ex_tax}")
    assert expense.person == "同事改的报销人", (
        f"只改备注却把报销人回退成了 {expense.person!r}")
    assert not plan.conflicts, plan.conflicts


def test_expense_touched_amount_still_wins(db):
    """防误伤：用户真改了金额，就得按用户的写。"""

    project, _order, _line, expense = _make_project_with_expense(db)
    workbook, worksheet, headers = _download(
        db, project.project_id, master.V2_SHEET_EXPENSE)
    row_no = _row_for(worksheet, headers, expense.raw_line_id)
    worksheet.cell(row_no, headers["操作"], "UPDATE")
    worksheet.cell(row_no, headers["未税金额"], 600)

    _apply(db, project.project_id, workbook)
    db.refresh(expense)
    assert expense.amount_ex_tax == Decimal("600.00")


# ---------- 05_实收回款 ----------

def _seed_receipt(db, project):
    from app.models.maintenance_project import MaintenanceProjectContract
    from app.models.maintenance_project_operations import (
        MaintenanceCollectionSnapshot,
    )

    contract = MaintenanceProjectContract(
        project_contract_id=str(uuid.uuid4()),
        project_id=project.project_id,
        contract_id=f"c-{uuid.uuid4().hex[:8]}",
        contract_no=f"XSDD-{uuid.uuid4().hex[:8].upper()}",
        amount_inc_tax=Decimal("1000000.00"),
        contract_status="执行中",
        status_mapping_state="mapped",
        status_mapping_version="v1",
        included_in_total=True,
        effective_from=date(2020, 1, 1),
        source="synthetic_test",
        version=1,
    )
    db.add(contract)
    db.flush()
    snapshot = MaintenanceCollectionSnapshot(
        collection_id=str(uuid.uuid4()),
        project_id=project.project_id,
        project_contract_id=contract.project_contract_id,
        report_month=date(2026, 6, 1),
        cumulative_amount=Decimal("100.00"),
        status="confirmed", source="direct_api", version=1,
    )
    db.add(snapshot)
    db.commit()
    return contract, snapshot


def test_receipt_untouched_amount_is_not_reverted_by_a_remark_only_edit(db):
    """同事把累计实收 100 → 200；用户只改备注上传，200 必须留住。"""

    project, _part, _order, _line = _make_project_with_line(db)
    _contract, snapshot = _seed_receipt(db, project)
    workbook, worksheet, headers = _download(
        db, project.project_id, master.V2_SHEET_RECEIPTS)
    row_no = _row_for(worksheet, headers, snapshot.collection_id)

    snapshot.cumulative_amount = Decimal("200.00")
    snapshot.version += 1
    operations.bump_workbook_revision(db, project_id=project.project_id)
    db.commit()

    worksheet.cell(row_no, headers["备注"], "只改了个备注")

    plan, _result = _apply(db, project.project_id, workbook)
    db.refresh(snapshot)

    assert snapshot.cumulative_amount == Decimal("200.00"), (
        f"只改备注却把累计实收回退成了 {snapshot.cumulative_amount}")
    assert not plan.conflicts, plan.conflicts


def test_receipt_touched_amount_still_wins(db):
    """防误伤：用户真改了累计实收，就得按用户的写。"""

    project, _part, _order, _line = _make_project_with_line(db)
    _contract, snapshot = _seed_receipt(db, project)
    workbook, worksheet, headers = _download(
        db, project.project_id, master.V2_SHEET_RECEIPTS)
    row_no = _row_for(worksheet, headers, snapshot.collection_id)
    worksheet.cell(row_no, headers["累计实收金额（含税）"], 300)

    _apply(db, project.project_id, workbook)
    db.refresh(snapshot)
    assert snapshot.cumulative_amount == Decimal("300.00")


# ---------- 06_领用返还 ----------

def test_site_untouched_quantity_is_not_reverted_by_a_remark_only_edit(db):
    """同事把领用数量 2 → 7；用户只改备注上传，7 必须留住。"""

    project, part, _order, _line = _make_project_with_line(db)
    _issue, line = _site_issue(db, project, part, qty="2")
    workbook, worksheet, headers = _download(
        db, project.project_id, master.V2_SHEET_SITE)
    row_no = _row_for(worksheet, headers, line.issue_line_id)

    line.quantity = Decimal("7")
    operations.bump_workbook_revision(db, project_id=project.project_id)
    db.commit()

    worksheet.cell(row_no, headers["备注"], "只改了个备注")

    plan, _result = _apply(db, project.project_id, workbook)
    db.refresh(line)

    assert line.quantity == Decimal("7.000"), (
        f"只改备注却把领用数量回退成了 {line.quantity}")
    assert not plan.conflicts, plan.conflicts


def test_site_touched_quantity_still_wins(db):
    """防误伤：用户真改了数量，就得按用户的写。"""

    project, part, _order, _line = _make_project_with_line(db)
    _issue, line = _site_issue(db, project, part, qty="2")
    workbook, worksheet, headers = _download(
        db, project.project_id, master.V2_SHEET_SITE)
    row_no = _row_for(worksheet, headers, line.issue_line_id)
    worksheet.cell(row_no, headers["领用数量"], 9)

    _apply(db, project.project_id, workbook)
    db.refresh(line)
    assert line.quantity == Decimal("9.000")


# ---------- 三张表同时改 ----------

def test_one_workbook_touching_all_three_sheets_rebases_each_of_them(db):
    """三处补丁共享同一个 _V2MergeContext，合起来跑一次，别互相踩。"""

    project, part, _order, _line = _make_project_with_line(db)
    _issue, site_line = _site_issue(db, project, part, qty="2")
    _contract, snapshot = _seed_receipt(db, project)
    from app.models.maintenance import FProjectExpense
    from tests.test_maintenance_project_master_v2_editable import (
        _attribute_expense, _batch,
    )
    expense = FProjectExpense(
        raw_line_id=f"bxd-{uuid.uuid4()}", import_batch_id=_batch(db),
        bxd_no="BXD-ALL3-001", line_no=1, expense_date=date(2026, 8, 3),
        person="测试员", linked_sales_order_no="XSDD-EDIT-001",
        amount=Decimal("500.00"), tax_basis="ex",
        amount_ex_tax=Decimal("500.00"), amount_inc_tax=Decimal("565.00"),
    )
    db.add(expense)
    _attribute_expense(db, project, expense)
    db.commit()

    content = master.build_project_master_v2(
        db, project_id=project.project_id,
        sheets=(master.V2_SHEET_EXPENSE, master.V2_SHEET_RECEIPTS,
                master.V2_SHEET_SITE))
    workbook = load_workbook(io.BytesIO(content))

    # 同事在导出之后各改一处
    expense.amount_ex_tax = Decimal("900.00")
    expense.amount = Decimal("900.00")
    expense.amount_inc_tax = Decimal("1017.00")
    snapshot.cumulative_amount = Decimal("200.00")
    snapshot.version += 1
    site_line.quantity = Decimal("7")
    operations.bump_workbook_revision(db, project_id=project.project_id)
    db.commit()

    for sheet_name, entity_id, extra in (
        (master.V2_SHEET_EXPENSE, expense.raw_line_id, {"操作": "UPDATE"}),
        (master.V2_SHEET_RECEIPTS, snapshot.collection_id, {}),
        (master.V2_SHEET_SITE, site_line.issue_line_id, {}),
    ):
        worksheet = workbook[sheet_name]
        headers = {cell.value: cell.column for cell in worksheet[1]}
        row_no = _row_for(worksheet, headers, entity_id)
        for name, value in extra.items():
            worksheet.cell(row_no, headers[name], value)
        worksheet.cell(row_no, headers["备注"], "只改了个备注")

    plan, _result = _apply(db, project.project_id, workbook)
    db.refresh(expense)
    db.refresh(snapshot)
    db.refresh(site_line)

    assert expense.amount_ex_tax == Decimal("900.00")
    assert snapshot.cumulative_amount == Decimal("200.00")
    assert site_line.quantity == Decimal("7.000")
    assert not plan.conflicts, plan.conflicts
