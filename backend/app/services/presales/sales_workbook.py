"""销售自有询价的只读 PN 汇总；金额和处理人先经售前可见性出口。"""

from decimal import Decimal
from io import BytesIO

from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from app.models.dimensions import DimPart
from app.models.presales import PresalesInquiry
from app.services.excel_text_safety import SafeTextWorkbook, excel_safe
from app.services.presales import common, visibility
from app.services.workbook.schema import ColumnSpec, SheetSpec, WorkbookError, WorkbookSchema, column_allowed
from app.services.workflow.registry import Forbidden, InvalidTransition, WorkflowError


SUMMARY_SCHEMA = WorkbookSchema("presales_sales_summary", 1, [
    SheetSpec("summary", "销售询价汇总", [
        ColumnSpec("pn", "PN", "code", "reference", "按 PN 汇总；身份列始终保留"),
        ColumnSpec("description", "描述", "text", "reference", "来源行描述，重复描述合并"),
        ColumnSpec("qty", "合计数量", "decimal", "reference", "来源行数量相加", decimals=3),
        ColumnSpec("inquiry_nos", "来源单号", "text", "reference", "逐张列出来源询价单"),
        ColumnSpec("require_exact_pn", "要求 PN 一致", "bool", "reference", "任一来源要求一致时为是"),
        ColumnSpec("handler_min_prices", "各处理人最低价", "text", "reference", "仅统计可见的已发布报价；含税口径分别列示"),
        ColumnSpec("adopted_prices", "采纳价", "text", "reference", "成交时钉住的报价版本；含税口径分别列示"),
    ]),
], instructions=["本表为只读汇总，不含往返令牌，不能导入。客户信息不在汇总范围内。"])


def can_access(db, ctx, subject_type, subject_id):
    try:
        _, ctx = common.resolve_actor(db, ctx)
        common.require(ctx, "page_presales")
        common.require(ctx, "action_presales_inquire")
        return subject_type is None and subject_id is None
    except WorkflowError:
        return False


def reject_import(*args, **kwargs):
    raise WorkbookError("销售汇总为只读报表，不支持导入或应用")


def _ids(inquiry_ids):
    if (not isinstance(inquiry_ids, list) or not 1 <= len(inquiry_ids) <= 200
            or any(type(value) is not int or value < 1 for value in inquiry_ids)):
        raise InvalidTransition("请选择 1–200 张询价单")
    return list(dict.fromkeys(inquiry_ids))


def _tax_label(value):
    return "含税" if value is True else "不含税" if value is False else "税口径未填"


def _offer_prices(views):
    prices = {}
    for view in views:
        if view.get("status") != "published" or view.get("no_stock") or view.get("unit_price") is None:
            continue
        reporter = view.get("reporter") or {}
        key = (reporter.get("id"), reporter.get("name", "其他采购"), view.get("is_tax_inclusive"))
        price = Decimal(str(view["unit_price"]))
        prices[key] = min(prices.get(key, price), price)
    return "\n".join(f"{name}：{price:.2f}（{_tax_label(tax)}）"
                     for (_, name, tax), price in sorted(prices.items(), key=lambda item: (item[0][1], str(item[0][2]))))


def summary_rows(db, inquiry_ids, ctx):
    user, ctx = common.resolve_actor(db, ctx)
    common.require(ctx, "page_presales")
    common.require(ctx, "action_presales_inquire")
    ids = _ids(inquiry_ids)
    inquiries = list(db.scalars(select(PresalesInquiry).where(PresalesInquiry.id.in_(ids))
                               .order_by(PresalesInquiry.id)))
    if (len(inquiries) != len(ids) or any(row.created_by != user.id or row.kind != "sales_inquiry"
                                         for row in inquiries)):
        raise Forbidden("只能汇总本人发起的销售询价单")
    result = {}
    for inquiry in inquiries:
        detail = visibility.inquiry_view(db, inquiry, ctx)
        won = detail.get("won_offer")
        for line in detail["lines"]:
            key = ("part", line["part_id"]) if line["part_id"] is not None else ("raw", line["pn_raw"])
            if key not in result:
                part = db.get(DimPart, line["part_id"]) if line["part_id"] is not None else None
                result[key] = {"pn": part.pn_std if part else line["pn_raw"], "qty": Decimal("0"),
                               "require_exact_pn": False, "descriptions": [], "sources": [],
                               "offers": [], "adopted": []}
            row = result[key]
            row["qty"] += Decimal(str(line["qty"]))
            row["require_exact_pn"] |= line["require_exact_pn"]
            if line.get("description") and line["description"] not in row["descriptions"]:
                row["descriptions"].append(line["description"])
            if inquiry.inquiry_no not in row["sources"]:
                row["sources"].append(inquiry.inquiry_no)
            row["offers"].extend(offer for offer in detail["offers"] if offer["line_id"] == line["id"])
            row["offers"].extend(line.get("linked_offers", []))
            if line.get("ref_offer"):
                row["offers"].append(line["ref_offer"])
            row["offers"].extend(response["offer"] for response in detail["line_responses"]
                                 if response["line_id"] == line["id"] and response.get("offer"))
            if (won and won.get("unit_price") is not None and won.get("part_id") == line["part_id"]
                    and won.get("part_id") is not None):
                adopted = f"{inquiry.inquiry_no}：{Decimal(str(won['unit_price'])):.2f}（{_tax_label(won.get('is_tax_inclusive'))}）"
                if adopted not in row["adopted"]:
                    row["adopted"].append(adopted)
    return [{"pn": row["pn"], "description": "\n".join(row["descriptions"]), "qty": row["qty"],
             "require_exact_pn": row["require_exact_pn"], "inquiry_nos": "\n".join(row["sources"]),
             "handler_min_prices": _offer_prices(row["offers"]), "adopted_prices": "\n".join(row["adopted"])}
            for row in result.values()]


def _write(cell, value, *, text=False):
    value = excel_safe(value)
    if isinstance(value, str) and len(value) > 32767:
        raise WorkbookError("汇总单元格超过 Excel 文本上限，请减少来源单据")
    if isinstance(value, Decimal):
        value = format(value, "f")
    if text and value is not None:
        value = str(value)
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"
    if text:
        cell.number_format = "@"


def export_sales_summary(db, inquiry_ids, ctx, visible_columns=None):
    _, ctx = common.resolve_actor(db, ctx)
    rows = summary_rows(db, inquiry_ids, ctx)
    if visible_columns is not None and (not isinstance(visible_columns, dict)
            or any(not isinstance(values, list) or any(not isinstance(key, str) for key in values)
                   for values in visible_columns.values())):
        raise WorkbookError("visible_columns 必须按工作表列出列键")
    selected = None if visible_columns is None else visible_columns.get("summary", [])
    sheet = SUMMARY_SCHEMA.sheets[0]
    columns = [column for column in sheet.columns if column_allowed(column, ctx)
               and (selected is None or column.key in selected or column.key == "pn")]
    # 只读汇总不调用往返导出器，因此不存在令牌、签名元数据或可导入的实体身份。
    wb = SafeTextWorkbook()
    instructions = wb.active
    instructions.title = "填写说明"
    for index, instruction in enumerate(SUMMARY_SCHEMA.instructions, 1):
        _write(instructions.cell(index, 1), instruction, text=True)
    ws = wb.create_sheet(sheet.title)
    ws.freeze_panes = "B2"
    ws.protection.sheet = True
    for index, column in enumerate(columns, 1):
        cell = ws.cell(1, index)
        _write(cell, column.header, text=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
        cell.font = Font(bold=True)
        cell.comment = Comment(column.hint, "系统")
        ws.column_dimensions[get_column_letter(index)].width = column.width or 24
    for index, row in enumerate(rows, 2):
        for col_index, column in enumerate(columns, 1):
            value = row[column.key]
            if column.type == "bool":
                value = "是" if value else "否"
            cell = ws.cell(index, col_index)
            _write(cell, value, text=column.type in {"code", "text"})
            cell.fill = PatternFill("solid", fgColor="DDEBF7")
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
