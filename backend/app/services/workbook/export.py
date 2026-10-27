"""仅使用各办公软件共同支持的 XLSX 样式与保护功能。"""

from __future__ import annotations

from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO

from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.services.excel_text_safety import SafeTextWorkbook, excel_safe
from app.services.workbook.protocol import MAX_SHEET_COLUMNS, MAX_SHEET_ROWS, TOKEN_HEADER, row_token, write_metadata
from app.services.workbook.schema import WorkbookError, column_allowed, json_value
from app.services.workbook.values import clean_text, normalize

COLORS = {"locked": "D9D9D9", "required": "FFF2CC", "optional": "E2F0D9", "reference": "DDEBF7", "custom": "FFFFFF"}
BLANK_ROWS = 50


def _write(cell, value, *, text=False):
    value = excel_safe(value)
    if text and value is not None:
        value = str(value)
    if isinstance(value, str) and len(value) > 32767:
        raise WorkbookError("单元格文本超过 Excel 的 32767 字符上限")
    cell.value = value
    if isinstance(value, str):
        # 前导等号、加号等是用户文本，不作为公式执行。
        cell.data_type = "s"
    if text:
        cell.number_format = "@"


def _options(column):
    if column.type == "bool":
        return ["是", "否"]
    choices = column.choices or []
    return list(choices.values()) if isinstance(choices, dict) else list(choices)


def _hint(column):
    options = _options(column)
    return column.hint + ("；可选值：" + "、".join(map(str, options)) if options else "")


def export_workbook(schema, data, *, export_id, actor, visible_columns=None, custom_fields=None, ctx) -> bytes:
    if not str(export_id) or not str(actor):
        raise WorkbookError("导出编号和导出人不能为空")
    wb = SafeTextWorkbook()
    instructions = wb.active
    instructions.title = "填写说明"
    instructions.append(["颜色", "填写规则"])
    for edit, label in (("locked", "系统字段，锁定"), ("required", "必填"), ("optional", "选填"), ("reference", "参考信息，只读"), ("custom", "自定义字段，表头以 + 开头")):
        instructions.append([label, "灰底" if edit == "locked" else {"required": "黄底", "optional": "绿底", "reference": "蓝底", "custom": "白底"}[edit]])
        instructions.cell(instructions.max_row, 1).fill = PatternFill("solid", fgColor=COLORS[edit])
    for text in [
        "编码（如 PN）必须设为文本后再粘贴。错误示例：0231 变成 231、2.31E+07、长数字截断或 2310000.0；系统不会猜原值。",
        "数字支持千分位和人民币符号；超出约定小数位会报错，不四舍五入。日期请使用 YYYY-MM-DD 或 YYYY/MM/DD。",
        "不要复制隐藏的行令牌来新增行；新增行保持令牌为空。可调整行顺序、插入空行，系统按令牌识别。",
        "移除导出行会产生删除候选，必须在预览后逐条确认；未确认的行不会删除。",
        "只上传 XLSX；不要合并单元格，不要使用外部链接；公式须有办公软件保存的缓存结果。",
        *schema.instructions,
    ]:
        instructions.append([text])
        _write(instructions.cell(instructions.max_row, 1), text, text=True)
    manifest = {"sheets": {}}
    for sheet in schema.sheets:
        ws = wb.create_sheet(sheet.title)
        selected = None if visible_columns is None else visible_columns.get(sheet.key, [])
        columns = [c for c in sheet.columns if column_allowed(c, ctx) and (selected is None or c.key in selected or c.mandatory)]
        custom = (custom_fields or {}).get(sheet.key, [])
        if isinstance(custom, dict):
            custom = list(custom)
        names = [clean_text(str(name).removeprefix("+")) for name in custom]
        if names and not sheet.custom_columns:
            raise WorkbookError(f"{sheet.title} 不允许自定义列")
        if len(names) != len(set(names)) or any(not name for name in names):
            raise WorkbookError("自定义字段名称不能为空或重复")
        if len(columns) + len(names) + 1 > MAX_SHEET_COLUMNS:
            raise WorkbookError("工作表列数超出上限")
        rows = data.get(sheet.key, [])
        if len(rows) + BLANK_ROWS + 1 > MAX_SHEET_ROWS:
            raise WorkbookError("工作表行数超出上限")
        headers = [TOKEN_HEADER, *(c.header for c in columns), *("+" + name for name in names)]
        for index, header in enumerate(headers, 1):
            cell = ws.cell(1, index)
            _write(cell, header, text=True)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor=COLORS[columns[index - 2].edit] if 1 < index <= len(columns) + 1 else COLORS["locked" if index == 1 else "custom"])
            cell.comment = Comment(_hint(columns[index - 2]) if 1 < index <= len(columns) + 1 else ("系统行令牌，请勿改动或复制" if index == 1 else "自定义文本，可自由填写"), "系统")
        ws.column_dimensions["A"].hidden = True
        ws.column_dimensions["A"].fill = PatternFill("solid", fgColor=COLORS["locked"])
        ws.freeze_panes = "B2"
        ws.protection.sheet = True
        ws.protection.insertRows = not sheet.allow_insert
        ws.protection.deleteRows = not sheet.allow_delete
        ws.protection.insertColumns = not sheet.custom_columns
        ws.protection.formatCells = False
        ws.protection.sort = False
        sheet_manifest = {"columns": [c.key for c in columns], "custom_columns": names, "rows": []}
        seen = set()
        for row_index, row in enumerate(rows, 2):
            entity_id = row["entity_id"]
            version = row["base_version"]
            if entity_id is None or isinstance(version, bool) or not isinstance(version, int) or version < 0:
                raise WorkbookError("导出行必须有 entity_id 和非负整数 base_version")
            if entity_id in seen:
                raise WorkbookError("同一工作表不能重复导出同一实体")
            seen.add(entity_id)
            token = row_token(schema, str(export_id), sheet.key, entity_id, version)
            _write(ws.cell(row_index, 1), token, text=True)
            baseline = {}
            for col_index, column in enumerate(columns, 2):
                raw = row.get("values", {}).get(column.key)
                if column.type in {"code", "text"} and raw is not None:
                    raw = str(excel_safe(raw))
                try:
                    normalized = normalize(column, raw)
                except ValueError as exc:
                    raise WorkbookError(f"{sheet.title} 第 {row_index} 行 {column.header}：{exc}") from None
                baseline[column.key] = json_value(normalized)
                value = normalized
                if column.type == "enum" and isinstance(column.choices, dict):
                    value = column.choices.get(normalized, normalized)
                # 超过 Excel 数字精度时使用十进制文本，解析仍按声明类型处理。
                if isinstance(value, Decimal) and (len(value.as_tuple().digits) > 15 or "e" in str(value).lower()):
                    value = format(value, "f")
                elif isinstance(value, int) and not isinstance(value, bool) and len(str(abs(value))) > 15:
                    value = str(value)
                _write(ws.cell(row_index, col_index), value, text=column.type in {"code", "text"})
            custom_values = {}
            for name in names:
                raw = row.get("custom_fields", {}).get(name)
                custom_values[name] = clean_text(str(raw)) or None if raw is not None else None
            for index, name in enumerate(names, len(columns) + 2):
                _write(ws.cell(row_index, index), custom_values[name], text=True)
            sheet_manifest["rows"].append({"entity_id": entity_id, "base_version": version, "token": token, "values": baseline, "custom_fields": custom_values})
        last_row = max(2, len(rows) + 1) + (BLANK_ROWS if sheet.allow_insert else 0)
        for row_index in range(2, last_row + 1):
            ws.cell(row_index, 1).fill = PatternFill("solid", fgColor=COLORS["locked"])
            for index, column in enumerate(columns, 2):
                cell = ws.cell(row_index, index)
                cell.fill = PatternFill("solid", fgColor=COLORS[column.edit])
                cell.protection = Protection(locked=not column.editable)
                if column.type in {"text", "code"}:
                    cell.number_format = "@"
                elif column.type == "date":
                    cell.number_format = "yyyy-mm-dd"
                elif column.type == "decimal" and column.decimals is not None:
                    cell.number_format = "0" + ("." + "0" * column.decimals if column.decimals else "")
            for index in range(len(columns) + 2, len(headers) + 1):
                cell = ws.cell(row_index, index)
                cell.fill = PatternFill("solid", fgColor=COLORS["custom"])
                cell.protection = Protection(locked=False)
                cell.number_format = "@"
        for index, column in enumerate(columns, 2):
            letter = get_column_letter(index)
            ws.column_dimensions[letter].width = column.width or 22
            ws.column_dimensions[letter].protection = Protection(locked=not column.editable)
            ws.column_dimensions[letter].fill = PatternFill("solid", fgColor=COLORS[column.edit])
            if column.type in {"text", "code"}:
                ws.column_dimensions[letter].number_format = "@"
            options = _options(column)
            inline = ",".join(map(str, options))
            if options and len(inline) + 2 <= 255 and not any("," in str(v) or '"' in str(v) for v in options):
                dv = DataValidation(type="list", formula1='"' + inline + '"', allow_blank=column.edit != "required")
                dv.showDropDown = False
                ws.add_data_validation(dv)
                dv.add(f"{letter}2:{letter}{MAX_SHEET_ROWS}")
            instructions.append([f"{sheet.title} · {column.header}", _hint(column)])
            _write(instructions.cell(instructions.max_row, 2), _hint(column), text=True)
        for index in range(len(columns) + 2, len(headers) + 1):
            dimension = ws.column_dimensions[get_column_letter(index)]
            dimension.width = 24
            dimension.protection = Protection(locked=False)
            dimension.fill = PatternFill("solid", fgColor=COLORS["custom"])
            dimension.number_format = "@"
        instructions.append([sheet.title, f"新增行：{'允许（令牌留空）' if sheet.allow_insert else '不允许'}；删行：{'允许，上传后逐条确认' if sheet.allow_delete else '不允许'}"])
        manifest["sheets"][sheet.key] = sheet_manifest
    instructions.column_dimensions["A"].width = 65
    instructions.column_dimensions["B"].width = 90
    instructions.freeze_panes = "A2"
    write_metadata(wb, {"schema_name": schema.name, "schema_version": str(schema.version), "export_id": str(export_id), "actor": str(actor), "exported_at": datetime.now(timezone.utc).isoformat(), "date_system": "1900"}, manifest)
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
