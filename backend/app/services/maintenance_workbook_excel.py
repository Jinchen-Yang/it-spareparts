"""Excel value and metadata codecs shared by workbook adapters.

These functions do not query a Session. They depend on openpyxl's epoch handling
and the existing field-protection quantity normalization, not just the stdlib.
"""
from __future__ import annotations

import hashlib
import json
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

from openpyxl.utils.datetime import WINDOWS_EPOCH, from_excel

from app.services import maintenance_field_protection as field_protection
from app.services.maintenance_workbook_plans import WorkbookError
from app.services.maintenance_workbook_protocol import V2_SHEET_META, _META_CHUNK_SEP
from app.services.maintenance_workbook_values import _v2_hash_value


def _decimal(value, *, label: str, row_no: int) -> Decimal:
    try:
        parsed = Decimal(str(value).strip().replace(",", ""))
    except (InvalidOperation, AttributeError, ValueError):
        raise WorkbookError("invalid_amount", f"第 {row_no} 行{label}不是合法数字")
    if parsed < 0:
        raise WorkbookError("invalid_amount", f"第 {row_no} 行{label}不能为负")
    return parsed.quantize(Decimal("0.01"))


def _v2_date(value, *, row_no: int, label: str, epoch: datetime = WINDOWS_EPOCH) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        # WPS/Excel copy-paste can discard the date style while retaining the
        # serial. Use the workbook's epoch, never Unix time or the current year.
        try:
            if not Decimal(str(value)).is_finite() or value < 1:
                raise ValueError
            converted = from_excel(float(value), epoch=epoch)
            if isinstance(converted, datetime):
                return converted.date()
        except (ValueError, OverflowError, InvalidOperation):
            pass
        raise WorkbookError("invalid_date", f"第 {row_no} 行{label}日期序列值无效：{value!r}")
    raw = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S", "%Y-%m"):
        try:
            parsed = datetime.strptime(raw, fmt).date()
            return parsed.replace(day=1) if fmt == "%Y-%m" else parsed
        except ValueError:
            continue
    raise WorkbookError("invalid_date", f"第 {row_no} 行{label}日期格式无效")


def _v2_decimal(
    value,
    *,
    row_no: int,
    label: str,
    required: bool = False,
    allow_negative: bool = False,
    demand_quantity: bool = False,
) -> Decimal | None:
    if value in (None, ""):
        if required:
            raise WorkbookError("missing_amount", f"第 {row_no} 行{label}不能为空")
        return None
    try:
        parsed = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        raise WorkbookError("invalid_amount", f"第 {row_no} 行{label}不是合法数字")
    normalized = (
        field_protection.normalized_value("qty", parsed) if demand_quantity
        else parsed.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    )
    if (
        not normalized.is_finite()
        or abs(normalized) >= Decimal("1000000000000")
        or (normalized < 0 and not allow_negative)
    ):
        raise WorkbookError("invalid_amount", f"第 {row_no} 行{label}超出允许范围")
    return normalized


def _v2_meta(wb) -> dict[str, str]:
    if V2_SHEET_META not in wb.sheetnames:
        raise WorkbookError("template_version_mismatch", "工作簿缺少 V2 元数据，请重新下载当前项目总表")
    raw: dict[str, str] = {}
    for row in wb[V2_SHEET_META].iter_rows(min_col=1, max_col=2):
        if not row[0].value:
            continue
        key = str(row[0].value).strip()
        raw[key] = str(row[1].value or "").strip()
    # 合并 2.7.1 分块（key@2/key@3... 按数字序拼接——字典序会让 @10 排在 @2 前）；
    # 签名在发射端按逻辑值计算，因此必须在任何使用/验签之前完成合并。
    merged: dict[str, str] = {}
    chunks: dict[str, list[tuple[int, str]]] = {}
    for key, value in raw.items():
        if (_META_CHUNK_SEP in key
                and key.rsplit(_META_CHUNK_SEP, 1)[-1].isdigit()):
            base, suffix = key.rsplit(_META_CHUNK_SEP, 1)
            chunks.setdefault(base, []).append((int(suffix), value))
        else:
            merged[key] = value
    for base, items in chunks.items():
        if base not in merged:
            merged[base] = ""
        merged[base] += "".join(value for _n, value in sorted(items))
    return merged


def _v2_readonly_sheet_hash(ws) -> str:
    rows = []
    for cells in ws.iter_rows(values_only=True):
        values = [_v2_hash_value(value) for value in cells]
        while values and values[-1] == "":
            values.pop()
        rows.append(values)
    while rows and not rows[-1]:
        rows.pop()
    return hashlib.sha256(json.dumps(rows, ensure_ascii=False, separators=(",", ":")).encode()).hexdigest()
