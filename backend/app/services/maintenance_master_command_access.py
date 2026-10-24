"""Master upload parsing and exact conditional contracts in one read snapshot."""
from dataclasses import asdict, dataclass
import hashlib
import json
from io import BytesIO
from openpyxl import load_workbook

from sqlalchemy import select

from app.authz import service
from app.authz.context import create_context_batch, read_snapshot, relation_revision
from app.authz.context_registry import ENTRY_CONTRACTS
from app.authz.errors import authz_error
from app.models.maintenance import FMaintenanceOrder, FMaintenanceLine
from app.models.maintenance_source_assignment import MaintenanceSourceOrderAssignment
from app.services import maintenance_authorization as foundation
from app.services import maintenance_project_master_workbook as master
from app.services.maintenance_master_workbook_access import complete_alternative
from app.services.maintenance_workflow_authorization import require_complete

VALIDATE = 'POST /api/maintenance/projects/stable/{project_id}/master-workbook/validate'
APPLY = 'POST /api/maintenance/projects/stable/{project_id}/master-workbook/apply'
RESOURCES = ('maintenance_managed_workbook', 'maintenance_general_workbook_write',
             'maintenance_contract_amount_write', 'maintenance_project')


def selected_branches(entry_id, plan):
    amount = getattr(plan, 'contract_amount_change', None) is not None
    note = getattr(plan, 'project_note_change', None) is not None
    assignment = bool(getattr(plan, 'assignment_changes', ()))
    result = []
    for name, contract in ENTRY_CONTRACTS[entry_id]['branches'].items():
        if name.startswith('replay_'):
            continue  # Persisted receipts resolve their registered branch independently.
        req = contract['requirements']
        has_amount = ('maintenance_contract_amount_write' in req or
                      'update_contract_amount' in req.get('maintenance_managed_workbook', ()))
        has_note = 'update' in req.get('maintenance_project', ())
        has_assignment = 'reassign_source_orders' in req.get('maintenance_general_workbook_write', ())
        if (has_amount, has_note, has_assignment) == (amount, note, assignment):
            result.append(name)
    if not result:
        raise ValueError('实际工作簿操作没有对应的封闭合同')
    return tuple(result)


def plan_digest(plan):
    return relation_revision(json.loads(json.dumps(asdict(plan), default=str, sort_keys=True)))


def _resolve(db, identity, *, entry_id, params, branch=None, read_snapshot=None, locked_user=None, _principal=None):
    plan = params.get('_parsed_plan')
    if branch not in selected_branches(entry_id, plan):
        raise ValueError('工作簿合同分支与实际解析操作不一致')
    project_id = params['project_id']
    if plan is not None and plan.project_id != project_id:
        raise authz_error(409, 'AUTHZ_CONTEXT_STALE', '工作簿与实际项目不一致')
    changes = getattr(plan, 'assignment_changes', ())
    source_evidence = []
    parent_ids = {project_id}
    for change in changes:
        source = db.execute(select(FMaintenanceOrder.id, FMaintenanceOrder.raw_order_id,
            FMaintenanceOrder.order_no).where(FMaintenanceOrder.raw_order_id == change.source_order_id)).one_or_none()
        current = db.scalar(select(MaintenanceSourceOrderAssignment).where(
            MaintenanceSourceOrderAssignment.source_order_id == change.source_order_id,
            MaintenanceSourceOrderAssignment.is_active.is_(True)).execution_options(populate_existing=True))
        if (source is None or (current.assignment_id if current else None) != change.expected_assignment_id
                or (current.version if current else None) != change.expected_version
                or (current.project_id if current else None) != change.previous_project_id):
            raise authz_error(409, 'AUTHZ_CONTEXT_STALE', '来源单当前项目归属已变化，请重新校验')
        if current is not None:
            parent_ids.add(current.project_id)
        source_evidence.append([source.raw_order_id, source.id, source.order_no,
            current.assignment_id if current else None, current.version if current else None,
            current.project_id if current else None, project_id])
    if _principal is None:
        principal, evidence = foundation.resolve_project_subjects(db, identity, sorted(parent_ids),
            read_snapshot=read_snapshot, locked_user=locked_user)
    else:
        # Internal write capture runs after domain locks, before taking any actor
        # lock. This creates evidence only; the final core decision is mandatory.
        principal = _principal
        evidence = foundation.project_evidence(db, principal, sorted(parent_ids))
    selection = {'file_sha256': params['file_sha256']}
    if plan is not None:
        selection['command_plan_sha256'] = plan_digest(plan)
    carrier = foundation.slot_ref(entry_id, 'context_0', project_id, selection=selection)
    slots = {'context_0': (carrier,)}
    if changes:
        slots['source_orders'] = tuple(foundation.slot_ref(entry_id, 'source_orders', change.source_order_id)
                                      for change in sorted(changes, key=lambda change: change.source_order_id))
        slots['source_projects'] = tuple(foundation.slot_ref(entry_id, 'source_projects', pid) for pid in sorted(parent_ids))
    req = ENTRY_CONTRACTS[entry_id]['branches'][branch]['requirements']
    subjects = {resource: tuple(evidence[pid].subject for pid in
        (sorted(parent_ids) if resource == 'maintenance_general_workbook_write' and changes else [project_id]))
        for resource in req}
    context = foundation.resolved_context(db, entry_id=entry_id, branch=branch, carrier=carrier,
        slots=slots, subjects=subjects, revision=relation_revision({'carrier': carrier.to_dict(),
            'projects': {pid: value.authorization_revision for pid, value in evidence.items()},
            'sources': source_evidence}))
    return create_context_batch(db, [context])


def _input_only_reference(*args, **kwargs):
    # Neither a caller hash nor an old project ref can reconstruct an uploaded
    # plan. Persistent outputs must use a registered actual-carrier contract.
    raise authz_error(403, 'AUTHZ_ARTIFACT_UNAVAILABLE', '上传授权需要重新提交实际文件')


def _entry_context(db, identity, **kwargs):
    if kwargs['entry_id'] == APPLY and str(kwargs.get('branch', '')).startswith('replay_'):
        from app.services.maintenance_master_apply_access import _resolve as resolve_receipt
        return resolve_receipt(db, identity, **kwargs)
    return _resolve(db, identity, **kwargs)


def _entry_reference(db, identity, ref, **kwargs):
    if kwargs['entry_id'] == APPLY and ref.kind == 'master_workbook_operation':
        from app.services.maintenance_master_apply_access import _reference
        return _reference(db, identity, ref, **kwargs)
    return _input_only_reference()


for entry in (VALIDATE, APPLY):
    foundation.register_entry_resolver(entry, _entry_context, reference_resolver=_entry_reference)


@dataclass(frozen=True)
class _DiscoveredInput:
    project_id: str
    assignment_changes: tuple


def discover_source_targets(db, project_id, content, proof):
    """Read only input and minimal actual parent identifiers before business facts."""
    from app.services.query_filters import active_beta_maintenance_orders
    try:
        book = load_workbook(BytesIO(content), data_only=True)
    except Exception:
        raise master.WorkbookError('invalid_file', '无法读取 .xlsx 工作簿') from None
    try:
        master._v2_verify_meta(db, book, project_id, read_snapshot=proof)
        if master.V2_SHEET_PARTS not in book.sheetnames:
            return _DiscoveredInput(project_id, ())
        sheet = book[master.V2_SHEET_PARTS]
        index = {str(cell.value or ''): i for i, cell in enumerate(sheet[1])}
        if not {'实体ID', '维保单号', 'PN'}.issubset(index):
            return _DiscoveredInput(project_id, ())  # Parser returns the schema error without facts.
        orders = set()
        line_ids = set()
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or master._is_example_row(row):
                continue
            raw_id = master._cell(row, index, '实体ID')
            if raw_id not in (None, ''):
                try:
                    line_ids.add(int(raw_id))
                except (ValueError, TypeError):
                    raise master.WorkbookError('invalid_line_id', '备件实体ID无效') from None
            else:
                order_no = str(master._cell(row, index, '维保单号') or '').strip()
                if order_no:
                    orders.add(order_no)
        if line_ids:
            parents = list(db.execute(select(FMaintenanceLine.id, MaintenanceSourceOrderAssignment.project_id)
                .join(FMaintenanceOrder, FMaintenanceOrder.id == FMaintenanceLine.order_id)
                .outerjoin(MaintenanceSourceOrderAssignment,
                    (MaintenanceSourceOrderAssignment.source_order_id == FMaintenanceOrder.raw_order_id) &
                    MaintenanceSourceOrderAssignment.is_active.is_(True))
                .where(FMaintenanceLine.id.in_(line_ids))))
            if {row.id for row in parents} != line_ids or any(row.project_id != project_id for row in parents):
                raise authz_error(403, 'AUTHZ_FORBIDDEN', '上传的备件实体不属于当前项目')
        changes = []
        for order_no in sorted(orders):
            statement = select(FMaintenanceOrder.raw_order_id).where(FMaintenanceOrder.order_no == order_no)
            sources = list(db.scalars(active_beta_maintenance_orders(statement, FMaintenanceOrder)))
            if len(sources) != 1:
                raise master.WorkbookError('order_not_found', '上传的维保单号无法唯一确定实际来源单')
            source_id = sources[0]
            current = db.execute(select(MaintenanceSourceOrderAssignment.assignment_id,
                MaintenanceSourceOrderAssignment.version, MaintenanceSourceOrderAssignment.project_id)
                .where(MaintenanceSourceOrderAssignment.source_order_id == source_id,
                    MaintenanceSourceOrderAssignment.is_active.is_(True))).one_or_none()
            if current is None or current.project_id != project_id:
                changes.append(master.SourceOrderAssignmentChange(source_id, order_no,
                    current.assignment_id if current else None, current.version if current else None,
                    current.project_id if current else None, None))
        return _DiscoveredInput(project_id, tuple(changes))
    finally:
        book.close()


def validate_upload(session_factory, identity, *, project_id, content, force_takeover=False,
                    authorization_view=False, authorization_precondition=None, entry_id=VALIDATE):
    if entry_id not in (VALIDATE, APPLY):
        raise ValueError("上传入口必须是精确 validate/apply")
    with read_snapshot(session_factory) as (db, proof):
        params = {'project_id': project_id, 'file_sha256': hashlib.sha256(content).hexdigest()}
        initial = complete_alternative(db, identity, entry_id=entry_id, params=params, proof=proof,
            branches=selected_branches(entry_id, None))
        require_complete(initial)
        from app.config import get_settings
        if get_settings().maintenance_project_master_v2_enabled:
            discovered = discover_source_targets(db, project_id, content, proof)
            if discovered.assignment_changes:
                source_accesses = complete_alternative(db, identity, entry_id=entry_id,
                    params={**params, '_parsed_plan': discovered}, proof=proof,
                    branches=selected_branches(entry_id, discovered))
                require_complete(source_accesses)
            # The complete conditional plan is rechecked before any preview leaves.
            plan = master.validate_project_master_v2(db, project_id=project_id, data=content,
                user_ctx=None, force_takeover=force_takeover)
        else:
            plan = master.validate(db, project_id=project_id, data=content)
        accesses = complete_alternative(db, identity, entry_id=entry_id,
            params={**params, '_parsed_plan': plan}, proof=proof,
            branches=selected_branches(entry_id, plan))
        if authorization_view:
            return service.context_resolution(accesses, accesses.batch.contexts[0].ref)
        require_complete(accesses)
        service.check_context_precondition(authorization_precondition, accesses, accesses.batch.contexts[0].ref)
        return plan, accesses.principal, accesses.wire(accesses.batch.contexts[0].ref)


def validation_payload(plan, project_id, operator):
    return {
        "valid": True,
        "protocol_id": master.V2_PROTOCOL_ID,
        "template_version": master.V2_TEMPLATE_VERSION,
        "project_id": project_id,
        "sheets": list(plan.sheets),
        **plan.summary,
        "summary": plan.summary,
        "site_issue_merge_count": len(plan.site_issue_merges),
        "site_issue_merges": [dict(item) for item in plan.site_issue_merges],
        "site_issue_merge_token": (
            master.site_issue_merge_token(plan, operator)
            if plan.site_issue_merges else None
        ),
        # 2.7.0：字段级改动预览 + 行级冲突预览（前端据此提供接管入口）
        "changes": [dict(item) for item in plan.field_changes],
        "conflicts": [dict(item) for item in plan.conflicts],
        "overridden": [dict(item) for item in plan.overridden],
        # #265 契约：作废预览（03 显式 VOID + 04 显式 VOID/缺行），
        # apply 前对用户可见（前端 WorkbookRoundTrip 两阶段确认）。
        "will_void_rows": [dict(r) for r in plan.will_void_rows],
        "will_reassign_orders": [{
            "source_order_id": change.source_order_id,
            "order_no": change.order_no,
            "from_project_id": change.previous_project_id,
            "from_project_name": change.previous_project_name,
            "to_project_id": project_id,
        } for change in plan.assignment_changes],
        "warnings": list(plan.warnings),
        # D-02 作废优先：落在已作废行/单上的改动不生效、也不整本拒绝，
        # 行级列明作废人/时间（apply 回执同字段）。
        "voided_rows": [dict(item) for item in plan.voided_rows],
    }
