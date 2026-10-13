"""Bounded reading of large business exports, including malformed dimensions.

Keep openpyxl's value/date decoding, but never construct its full Cell map or a
per-cell merge expansion. The caller remains responsible for the upload and ZIP
structure checks. This reader additionally limits actual worksheet coordinates.
"""

from __future__ import annotations

import io
import unicodedata
from dataclasses import dataclass
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries

_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
MAX_UNCOMPRESSED_BYTES = 512 * 1024 * 1024


class StreamingWorkbookError(ValueError):
    def __init__(self, message: str, *, issues: list[dict] | None = None):
        super().__init__(message)
        self.issues = issues or []


@dataclass(frozen=True)
class _SheetMetadata:
    max_row: int
    max_column: int
    # (first row, last row, first column, last column), all 1-based.
    merges: tuple[tuple[int, int, int, int], ...]


def _text(value: Any) -> str:
    return unicodedata.normalize("NFKC", "" if value is None else str(value)).strip()


def _metadata(
    archive: ZipFile,
    path: str,
    *,
    max_source_rows: int,
    max_columns: int,
    max_cells: int,
) -> _SheetMetadata:
    last_row = max_row = max_column = cell_count = 0
    merges: list[tuple[int, int, int, int]] = []
    with archive.open(path) as source:
        for _event, element in ET.iterparse(source, events=("end",)):
            if element.tag == _NS + "row":
                row = int(element.get("r", last_row + 1))
                if row <= last_row:
                    raise StreamingWorkbookError("工作表行号重复或顺序错误，请重新导出")
                last_row = row
                max_row = max(max_row, row)
                if row > max_source_rows:
                    raise StreamingWorkbookError(f"源文件行数超过安全上限 {max_source_rows}")
                for cell in element:
                    if cell.tag != _NS + "c":
                        continue
                    cell_count += 1
                    cell_row, column = coordinate_to_tuple(cell.get("r", ""))
                    if cell_row != row:
                        raise StreamingWorkbookError("单元格坐标与所在行不一致，请重新导出")
                    max_column = max(max_column, column)
                    if column > max_columns:
                        raise StreamingWorkbookError(f"源文件列数超过安全上限 {max_columns}")
                    if cell_count > max_cells:
                        raise StreamingWorkbookError(f"源文件单元格数超过安全上限 {max_cells}")
                # Clearing only the outer worksheet leaves a detached sheetData
                # tree alive. Clear each row so its cells/text are released now.
                element.clear()
            elif element.tag == _NS + "mergeCell":
                first_col, first_row, last_col, end_row = range_boundaries(
                    element.get("ref", "")
                )
                if first_row < 1 or first_col < 1 or end_row < first_row or last_col < first_col:
                    raise StreamingWorkbookError("合并单元格范围无效，请重新导出")
                if end_row > max_source_rows or last_col > max_columns:
                    raise StreamingWorkbookError("合并单元格范围超过源文件安全上限")
                # Independently bound metadata, including workbooks with few cells
                # but a maliciously large merge list.
                if len(merges) >= min(max_cells, 1_000_000):
                    raise StreamingWorkbookError("合并单元格数量超过安全上限")
                merges.append((first_row, end_row, first_col, last_col))
                max_row = max(max_row, end_row)
                max_column = max(max_column, last_col)
                element.clear()
    return _SheetMetadata(max_row, max_column, tuple(sorted(merges)))


def detect_streaming(
    data: bytes,
    adapters: Any,
    *,
    header_scan_rows: int = 20,
    max_source_rows: int = 250_000,
    max_columns: int = 256,
    max_cells: int = 8_000_000,
) -> tuple[Any, dict]:
    """Return the chosen adapter and keyword arguments for ``DetectedSheet``.

    Only recognized fields survive in data rows; header tuples remain complete.
    Limits include physical header/blank rows and cells, before business filtering.
    """
    workbook = None
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise StreamingWorkbookError("本次导入的 ZIP 解压总量超过 512 MiB 安全上限")
            workbook = load_workbook(
                io.BytesIO(data), read_only=True, data_only=True, keep_links=False
            )
            registered = list(adapters.values() if isinstance(adapters, dict) else adapters)
            candidates = []
            metadata = {}
            for sheet in workbook.worksheets:
                if sheet.sheet_state != "visible" or sheet.title.lower().startswith("hidden_"):
                    continue
                # ReadOnlyWorksheet has no public path accessor. The same path
                # openpyxl uses is needed to inspect true coordinates/merge ranges.
                info = _metadata(
                    archive, sheet._worksheet_path,
                    max_source_rows=max_source_rows, max_columns=max_columns,
                    max_cells=max_cells,
                )
                metadata[sheet.title] = info
                sheet.reset_dimensions()
                band = [
                    tuple(_text(value) for value in row)
                    for row in sheet.iter_rows(
                        min_row=1, max_row=min(header_scan_rows, info.max_row),
                        max_col=info.max_column or 1, values_only=True,
                    )
                ]
                for row_no, headers in enumerate(band, start=1):
                    system_headers = band[row_no - 2] if row_no > 1 else ()
                    for adapter in registered:
                        result = adapter.recognize(headers, system_headers or None)
                        if result is not None:
                            score, indexes, matches = result
                            candidates.append((
                                score, sheet.title, row_no, adapter, headers,
                                system_headers, indexes, matches,
                            ))
            if not candidates:
                raise StreamingWorkbookError(
                    "无法识别表单：未找到销售合同额或收款单所需表头",
                    issues=[{"code": "unknown_form", "registered_forms": [a.key for a in registered]}],
                )
            candidates.sort(key=lambda item: (-item[0], item[1], item[2], item[3].key))
            best = [item for item in candidates if item[0] == candidates[0][0]]
            if len({(item[1], item[2], item[3].key) for item in best}) != 1:
                raise StreamingWorkbookError(
                    "工作簿中存在多个同等匹配的业务表，无法自动选择",
                    issues=[
                        {"sheet": s, "header_row": r, "form_type": a.key}
                        for _score, s, r, a, *_rest in best[:20]
                    ],
                )
            _score, name, header_row, adapter, headers, system_headers, indexes, matches = best[0]
            info = metadata[name]
            guarded = {
                index + 1: field for field, index in indexes.items()
                if field in getattr(adapter, "merge_guarded_fields", frozenset())
            }
            ranges = []
            for start, end, first_col, last_col in info.merges:
                if end <= header_row or (start == end and first_col == last_col):
                    continue
                for column, field in guarded.items():
                    if first_col <= column <= last_col:
                        cell_range = (
                            f"{get_column_letter(first_col)}{start}:"
                            f"{get_column_letter(last_col)}{end}"
                        )
                        raise StreamingWorkbookError(
                            f"第 {start}-{end} 行的「{headers[column - 1]}」是合并单元格；"
                            "逐行金额列不接受合并（铺开会把一笔钱算成多笔），请取消合并后重新导出",
                            issues=[{"code": "merged_amount_column", "cell_range": cell_range,
                                     "canonical_field": field}],
                        )
                ranges.append((start, end, first_col, last_col))

            sheet = workbook[name]
            selected = set(indexes.values())
            rows = []
            active: list[tuple[int, int, int, int, Any]] = []
            range_index = expanded = 0
            # Starting at row 1 also handles a merge anchored in a header that
            # extends into data, matching the existing ordinary reader.
            for row_no, raw in enumerate(sheet.iter_rows(
                min_row=1, max_row=info.max_row, max_col=info.max_column or 1,
                values_only=True,
            ), start=1):
                active = [merge for merge in active if merge[1] >= row_no]
                while range_index < len(ranges) and ranges[range_index][0] == row_no:
                    start, end, first_col, last_col = ranges[range_index]
                    anchor = raw[first_col - 1]
                    if anchor is not None:
                        active.append((start, end, first_col, last_col, anchor))
                        expanded += (end - start + 1) * (last_col - first_col + 1) - 1
                    range_index += 1
                if row_no <= header_row or all(_text(value) == "" for value in raw):
                    continue
                values = [value if index in selected else None for index, value in enumerate(raw)]
                for start, _end, first_col, last_col, anchor in active:
                    for index in selected:
                        if first_col - 1 <= index < last_col and not (
                            row_no == start and index == first_col - 1
                        ):
                            values[index] = anchor
                rows.append((row_no, tuple(values)))
            if not rows:
                raise StreamingWorkbookError("识别到表头，但没有数据行")
            return adapter, {
                "name": name, "header_row": header_row,
                "header_rows": ((header_row - 1, header_row) if system_headers else (header_row,)),
                "headers": headers, "system_headers": system_headers,
                "field_indexes": indexes, "field_matches": matches,
                "rows": tuple(rows), "merged_cells_expanded": expanded,
            }
    except StreamingWorkbookError:
        raise
    except Exception as exc:
        raise StreamingWorkbookError(f"无法读取 .xlsx：{type(exc).__name__}") from exc
    finally:
        if workbook is not None:
            workbook.close()
