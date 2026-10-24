"""Whole uploaded global line workbook; exact rows and actual parent closure."""
import io
from hashlib import sha256
from uuid import uuid4
from openpyxl import load_workbook
from sqlalchemy import select,func
from sqlalchemy.orm import Session
from app.authz import service
from app.authz.context import PolicySubject,create_context_batch,relation_revision
from app.authz.errors import authz_error
from app.config import DATA_CHANGE_ADVISORY_LOCK_KEY
from app.models.maintenance import FMaintenanceLine,FMaintenanceOrder
from app.models.maintenance_project import MaintenanceProject
from app.models.maintenance_source_assignment import MaintenanceSourceOrderAssignment as Assignment
from app.services import maintenance_authorization as foundation
from app.services import maintenance_project_master_workbook as business
from app.services import maintenance_project_operations as operations

VALIDATE='POST /api/maintenance/spare-part-lines/validate'
APPLY='POST /api/maintenance/spare-part-lines/apply'
RESOURCE='maintenance_global_workbook'


def uploaded_row_ids(data):
    """Only user-provided bytes; no stored business facts before authorization."""
    try:wb=load_workbook(io.BytesIO(data),data_only=True,read_only=True)
    except Exception as exc:raise business.WorkbookError('invalid_file','无法读取 .xlsx') from exc
    try:
        if business.GLOBAL_SHEET not in wb.sheetnames:raise business.WorkbookError('missing_sheet','缺少全局备件工作表')
        rows=[];col=len(business._GLOBAL_HEADERS)
        for values in wb[business.GLOBAL_SHEET].iter_rows(min_row=2,values_only=True):
            if not values or all(v is None or str(v).strip()=='' for v in values):continue
            try:ident=int(str(values[col]).strip())
            except (IndexError,ValueError,TypeError):raise business.WorkbookError('line_not_recognized','工作簿含无效的备件行标识') from None
            if ident<=0 or ident>=2**53 or ident in rows:raise business.WorkbookError('duplicate_line','工作簿包含无效或重复行')
            rows.append(ident)
        return sorted(rows)
    finally:wb.close()


def _batch(db,principal,entry_id,params):
    line_ids=params['line_ids']
    edges=list(db.execute(select(FMaintenanceLine.id,FMaintenanceOrder.raw_order_id,Assignment.assignment_id,
        Assignment.project_id,Assignment.version).join(FMaintenanceOrder,FMaintenanceOrder.id==FMaintenanceLine.order_id)
        .join(Assignment,(Assignment.source_order_id==FMaintenanceOrder.raw_order_id)&Assignment.is_active.is_(True))
        .where(FMaintenanceLine.id.in_(line_ids)).order_by(FMaintenanceLine.id)))
    if {r.id for r in edges}!=set(line_ids):raise authz_error(403,'AUTHZ_FORBIDDEN','工作簿包含无有效项目归属的备件行')
    projects=sorted({r.project_id for r in edges});evidence=foundation.project_evidence(db,principal,projects)
    digest=relation_revision({'file':params['file_sha256'],'rows':[tuple(r) for r in edges]})
    discovery=foundation.slot_ref(entry_id,'context_0','maintenance_global_workbook_content',selection={'file_sha256':params['file_sha256'],'command_plan_sha256':digest})
    carrier=foundation.slot_ref(entry_id,'context_1',digest,selection={'command_plan_sha256':digest})
    slots={'context_0':(discovery,),'context_1':(carrier,)}
    if projects:slots['context_projects']=tuple(foundation.slot_ref(entry_id,'context_projects',pid) for pid in projects)
    subjects=tuple(evidence[p].subject for p in projects) or (PolicySubject(carrier.id,dict.fromkeys(('is_manager','is_viewer','is_sales','is_active'),False)),)
    context=foundation.resolved_context(db,entry_id=entry_id,branch='default' if projects else 'no_targets',carrier=carrier,
        slots=slots,subjects={RESOURCE:subjects},revision=relation_revision({'input':digest,
            'actor':[principal.id,principal.authorization_epoch],'projects':[evidence[p].authorization_revision for p in projects]}))
    return create_context_batch(db,[context]),set(projects)


def resolve(db,identity,*,entry_id,params,branch=None,read_snapshot=None,locked_user=None):
    principal=foundation._named_principal(db,identity,read_snapshot=read_snapshot,locked_user=locked_user)
    return _batch(db,principal,entry_id,params)[0]
for _entry in (VALIDATE,APPLY):foundation.register_entry_resolver(_entry,resolve)


class GlobalLinesAuthorization:
    """Typed callback scoped to one full uploaded plan, never a reusable allow."""
    def __init__(self,identity,params,original,principal):
        self.identity=identity;self.params=params;self.original=original;self.principal=principal;self.wire=None
    def authorize_after(self,db):
        db.flush();actor,_=foundation.lock_final_users(db,self.identity)
        batch=resolve(db,self.identity,entry_id=APPLY,params=self.params,locked_user=actor)
        accesses=service.context_accesses(db,self.identity,batch,locked_user=actor)
        accesses.require(batch.contexts[0].ref,RESOURCE,'apply');self.wire=accesses.wire(batch.contexts[0].ref)


def execute(db,identity,*,entry_id,data,view=None,precondition=None):
    if view not in (None,'authorization'):raise authz_error(422,'AUTHZ_INVALID_REQUEST','未知授权视图')
    params={'file_sha256':sha256(data).hexdigest(),'line_ids':uploaded_row_ids(data)};action='validate' if entry_id==VALIDATE else 'apply'
    with foundation.read_context_snapshot(lambda:Session(db.get_bind()),identity,entry_id=entry_id,params=params) as (snapshot,_,batch,accesses):
        original=batch.contexts[0];ref=original.ref
        if view=='authorization':return service.context_resolution(accesses,ref)
        access=accesses.require(ref,RESOURCE,action)
        service.check_context_precondition(foundation.decode_precondition(precondition),accesses,ref)
        projects={r.id for r in original.slots.get('context_projects',())}
        plan=business.validate_global(snapshot,data=data,allowed_project_ids=projects)
        if entry_id==VALIDATE:return {'valid':True,**plan.summary,'authorization_context':accesses.wire(ref)}
    try:
        db.execute(select(func.pg_advisory_xact_lock(DATA_CHANGE_ADVISORY_LOCK_KEY)))
        operations.lock_workbook_states(db,project_ids=projects)
        list(db.scalars(select(MaintenanceProject).where(MaintenanceProject.project_id.in_(projects))
            .order_by(MaintenanceProject.project_id).with_for_update().execution_options(populate_existing=True)))
        business._lock_global_refill_rows(db,line_ids=params['line_ids'],allowed_project_ids=projects,prelocked_project_ids=projects)
        stable,_=_batch(db,access.principal,entry_id,params)
        if stable.contexts[0].authorization_revision!=original.authorization_revision:raise authz_error(409,'AUTHZ_CONTEXT_CHANGED','工作簿完整归属已变化')
        guard=GlobalLinesAuthorization(identity,params,original,access.principal)
        result=business.apply_global_lines(db,plan,operated_by=access.principal.username,import_batch_id=str(uuid4()),
            allowed_project_ids=projects,_global_authorization=guard)
        if guard.wire is None:raise RuntimeError('global workbook omitted final authorization')
        return {**result,'authorization_context':guard.wire}
    except Exception:db.rollback();raise
