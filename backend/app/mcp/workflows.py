"""Bounded file workflows with native business rules and durable receipts."""

import hashlib
import io
import os
import zipfile
from dataclasses import asdict
from pathlib import Path

from fastapi.encoders import jsonable_encoder
from openpyxl import load_workbook
from sqlalchemy import select

from app.api.maintenance_project_scope import (
    enforce_maintenance_project_access,
)
from app.config import get_settings
from app.etl import pipeline, precheck
from app.mcp.core import (
    McpError,
    digest,
    event,
    now,
    owned,
    record,
    refresh_actor,
    require,
    uid,
)
from app.models.maintenance_project import MaintenanceProject
from app.models.mcp import McpRecord
from app.services import import_safety
from app.services import maintenance_acceptance_checklist as checklist
from app.services import maintenance_expense_collection_workbook as expense
from app.services import maintenance_project_operations as operations

FAMILIES = ("legacy_trade", "acceptance_checklist", "expense_collection")


def project_access(db, actor, project_id, *, expense_data=False, write=False):
    require(actor, "mcp.projects.read", "page_maintenance")
    enforce_maintenance_project_access(db, project_id=project_id, ctx=actor.ctx)
    if not db.get(MaintenanceProject, project_id):
        raise McpError("not_found", "项目不存在", 404)
    if expense_data:
        perms = actor.ctx.permissions or {}
        if actor.ctx.role != "admin" and not perms.get("data_profit"):
            raise McpError("permission_denied", "无金额查看权限", 403)
        if (
            write
            and actor.ctx.role != "admin"
            and not perms.get("action_maintenance_expense_collection_upload")
        ):
            raise McpError("permission_denied", "无费用回款上传权限", 403)


def family_access(db, actor, family, project_id=None):
    if family not in FAMILIES:
        raise McpError("unsupported_family", "该单据族尚未开放")
    if family == "legacy_trade":
        require(actor, "mcp.import.preview", "page_import")
    else:
        project_access(
            db,
            actor,
            project_id,
            expense_data=family == "expense_collection",
            write=True,
        )


def root():
    p = Path(get_settings().raw_file_dir).resolve() / "mcp-files"
    p.mkdir(mode=0o700, parents=True, exist_ok=True)
    return p


def data_path(id):
    # UUID only, never a filename or arbitrary client path.
    import uuid

    try:
        if str(uuid.UUID(id)) != id:
            raise ValueError()
    except ValueError as exc:
        raise McpError("not_found", "文件不存在", 404) from exc
    p = root() / id
    if p.is_symlink():
        raise McpError("invalid_file", "文件存储状态异常")
    return p


def save_bytes(id, data):
    path = data_path(id)
    fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW, 0o600)
    with os.fdopen(fd, "wb") as f:
        f.write(data)
        f.flush()
        os.fsync(f.fileno())


def file_bytes(db, actor, id, kind="file"):
    r = owned(db, actor, id, kind)
    if r.state != "ready":
        raise McpError("file_not_ready", "文件尚未就绪", 409)
    p = data_path(id)
    try:
        if p.stat().st_size > get_settings().mcp_max_file_bytes:
            raise McpError("file_too_large", "文件超出限制", 413)
        with os.fdopen(os.open(p, os.O_RDONLY | os.O_NOFOLLOW), "rb") as f:
            data = f.read(get_settings().mcp_max_file_bytes + 1)
    except OSError as exc:
        raise McpError("not_found", "文件不存在", 404) from exc
    if hashlib.sha256(data).hexdigest() != r.payload["sha256"]:
        raise McpError("invalid_file", "文件完整性校验失败", 409)
    return r, data


def safe_xlsx(data):
    import_safety.validate_xlsx_zip(data, max_bytes=get_settings().mcp_max_file_bytes)
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if sum(i.file_size for i in z.infolist()) > 100 * 1024 * 1024:
            raise McpError("file_too_large", "工作簿解压后过大", 413)
        for n in z.namelist():
            if "vbaproject" in n.lower() or n.startswith("xl/externalLinks/"):
                raise McpError("invalid_file", "不支持宏或外链工作簿")
            if n.endswith(".rels") and b'TargetMode="External"' in z.read(n):
                raise McpError("invalid_file", "不支持外部链接")
        if (
            "[Content_Types].xml" not in z.namelist()
            or "xl/workbook.xml" not in z.namelist()
        ):
            raise McpError("invalid_file", "不是有效工作簿")


def inspect_document(db, actor, id):
    require(actor, "mcp.files.read")
    r, data = file_bytes(db, actor, id)
    ext = Path(r.payload["name"]).suffix.lower()
    result = {
        "file_id": id,
        "filename": r.payload["name"],
        "sha256": r.payload["sha256"],
        "detected_family": None,
        "confidence": "unknown",
        "warnings": [],
    }
    if ext != ".xlsx":
        result["warnings"] = ["仅暂存附件，未提取或确认业务事实"]
        return result
    safe_xlsx(data)
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    try:
        result["sheets"] = [
            {"name": s.title, "rows": s.max_row, "columns": s.max_column} for s in wb
        ]
        # Only classify, never return arbitrary cell contents as instructions.
        head = {
            str(v).strip()
            for s in wb
            for row in s.iter_rows(max_row=10, max_col=12, values_only=True)
            for v in row
            if v is not None
        }
        if "验收需求" in head:
            result.update(
                detected_family="acceptance_checklist",
                confidence="header_match",
                protocol="checklist-v1",
            )
        elif any("费用报销" in s.title or "实收回款" in s.title for s in wb):
            result["warnings"].append(
                "疑似往返工作簿；请按原导出类型选择，不能根据名称认定版本"
            )
        else:
            result["warnings"].append("可尝试通用单据预检；WBDD 与其他模板不得直接套用")
    finally:
        wb.close()
    return result


def preview_one(db, actor, file_id, family, project_id, mode):
    family_access(db, actor, family, project_id)
    file, data = file_bytes(db, actor, file_id)
    if not file.payload["name"].lower().endswith(".xlsx"):
        raise McpError("unsupported_family", "业务预检仅接收 xlsx")
    safe_xlsx(data)
    base = {
        "file_id": file_id,
        "sha256": file.payload["sha256"],
        "filename": file.payload["name"],
        "family": family,
        "project_id": project_id,
        "mode": mode,
        "auth_version": actor.version,
        "actor": actor.snapshot(),
    }
    if family == "legacy_trade":
        result = precheck.inspect_file(
            str(data_path(file_id)), file.payload["name"], mode=mode
        )
        matches = pipeline.successful_batch_ids_by_hash(db, {file.payload["sha256"]})
        precheck.apply_exact_success_matches(
            [(result, file.payload["sha256"])], matches, mode
        )
        return dict(
            base,
            summary=result,
            can_apply=False,
            warnings=["通用单据首版仅预检；正式导入仍使用原业务页面，不会自动写入"],
            transaction_scope="per_file",
        )
    if family == "acceptance_checklist":
        # Same project row lock is used by the native apply path.
        db.execute(
            select(MaintenanceProject)
            .where(MaintenanceProject.project_id == project_id)
            .with_for_update()
        )
        if len(data) > checklist.MAX_CHECKLIST_BYTES:
            raise McpError("file_too_large", "验收清单超过8MiB限制", 413)
        parsed = checklist.parse_checklist_workbook(data, file.payload["name"])
        # Native parser caps rows. Reject oversize instead of silently truncating.
        wb = load_workbook(io.BytesIO(data), read_only=True)
        try:
            if any(
                (s.max_row or 0)
                > checklist.MAX_CHECKLIST_ROWS + checklist.HEADER_SCAN_ROWS
                for s in wb
            ):
                raise McpError("file_too_large", "清单超过支持行数，未截断导入")
        finally:
            wb.close()
        current = checklist.project_checklist(db, project_id)["current"]
        bid = checklist.store_preview(
            db,
            parsed,
            project_id=project_id,
            uploaded_by=actor.name,
            idempotency_key=uid(),
        )
        return dict(
            base,
            native_batch_id=bid,
            baseline=current["batch_id"] if current else None,
            summary={
                "item_rows": parsed["item_rows"],
                "issue_rows": parsed["issue_rows"],
                "will_replace_rows": current["item_rows"] if current else 0,
            },
            issues=[
                {"row": i["row_no"], "messages": i["issues"]}
                for i in parsed["items"]
                if i["issues"]
            ],
            changes=parsed["items"],
            can_apply=not parsed["issue_rows"],
            transaction_scope="whole_checklist",
        )
    if not get_settings().maintenance_boss_dashboard_enabled:
        raise McpError("feature_disabled", "维保功能未启用", 404)
    state = operations.get_or_create_workbook_state(
        db, project_id=project_id, lock=True
    )
    plan = expense.validate(db, project_id=project_id, data=data)
    if len(plan.expense_updates) + len(plan.collection_ops) > 2000:
        raise McpError("file_too_large", "单份工作簿变更超过2000条，请分批处理", 413)
    return dict(
        base,
        baseline=state.data_version,
        plan_hash=digest(asdict(plan)),
        summary=plan.summary,
        changes=jsonable_encoder(asdict(plan)),
        can_apply=True,
        transaction_scope="whole_workbook",
    )


def make_preview(db, actor, args):
    require(actor, "mcp.import.preview")
    require(actor, "mcp.files.read")
    family = args["family"]
    family_access(db, actor, family, args.get("project_id"))
    items = []
    for fid in args["file_ids"]:
        try:
            with db.begin_nested():
                item = preview_one(
                    db,
                    actor,
                    fid,
                    family,
                    args.get("project_id"),
                    args.get("mode", "skip"),
                )
                p = record(db, actor, "preview", item, hours=0.5)
                item = {
                    "preview_id": p.id,
                    "state": "ready" if item["can_apply"] else "needs_review",
                    "file_id": fid,
                    "summary": item["summary"],
                    "can_apply": item["can_apply"],
                }
        except McpError as exc:
            if exc.status in (401, 403):
                raise
            item = {
                "file_id": fid,
                "state": "failed",
                "error": {"code": exc.code, "message": exc.message},
            }
        except (
            checklist.ChecklistParseError,
            expense.WorkbookError,
            import_safety.UploadSafetyError,
        ) as exc:
            item = {
                "file_id": fid,
                "state": "failed",
                "error": {"code": "invalid_template", "message": str(exc)[:600]},
            }
        items.append(item)
    return {
        "items": items,
        "can_apply": any(i.get("can_apply") for i in items),
        "transaction_scope": "per_file",
        "status": "partial" if any(i["state"] == "failed" for i in items) else "ready",
    }


def apply_preview(db, actor, preview_id, expected_hash, request_id):
    """Web-only confirmation + native write + durable receipt/audit in ONE transaction."""
    if not get_settings().mcp_confirm_enabled:
        raise McpError("feature_disabled", "正式确认暂未启用", 403)
    actor = refresh_actor(db, actor)
    require(actor, "mcp.import.preview")
    p = owned(db, actor, preview_id, "preview", lock=True)
    body = p.payload
    from app.mcp.core import Actor

    snap = body["actor"]
    actor = refresh_actor(
        db,
        Actor(
            actor.name,
            actor.ctx,
            set(snap["scopes"]),
            snap["version"],
            snap.get("credential_id"),
        ),
    )
    require(actor, "mcp.import.preview")
    family_access(db, actor, body["family"], body.get("project_id"))
    if digest(body) != expected_hash:
        raise McpError("stale_preview", "预览内容已变化", 409)
    if p.state == "applied":
        receipt = db.scalar(
            select(McpRecord).where(
                McpRecord.kind == "receipt",
                McpRecord.owner == actor.name,
                McpRecord.key == p.id,
            )
        )
        event(
            db,
            actor,
            "confirm_import",
            "replayed",
            {
                "preview_id": p.id,
                "receipt_id": receipt.id,
                "project_id": body["project_id"],
            },
            request_id,
        )
        return receipt.payload
    if body["auth_version"] != actor.version or not body["can_apply"]:
        raise McpError("stale_preview", "预览不能提交，请重新检查", 409)
    file, data = file_bytes(db, actor, body["file_id"])
    if file.payload["sha256"] != body["sha256"]:
        raise McpError("stale_preview", "原件已变化", 409)
    if body["family"] == "acceptance_checklist":
        db.execute(
            select(MaintenanceProject)
            .where(MaintenanceProject.project_id == body["project_id"])
            .with_for_update()
        )
        current = checklist.project_checklist(db, body["project_id"])["current"]
        if (current["batch_id"] if current else None) != body["baseline"]:
            raise McpError("stale_preview", "当前清单已变化，请重新预览", 409)
        result = checklist.apply_batch(
            db, body["native_batch_id"], operated_by=actor.name
        )
    elif body["family"] == "expense_collection":
        if not get_settings().maintenance_boss_dashboard_enabled:
            raise McpError("feature_disabled", "维保功能未启用", 404)
        state = operations.get_or_create_workbook_state(
            db, project_id=body["project_id"], lock=True
        )
        plan = expense.validate(db, project_id=body["project_id"], data=data)
        if (
            state.data_version != body["baseline"]
            or digest(asdict(plan)) != body["plan_hash"]
        ):
            raise McpError("stale_preview", "业务数据已变化，请重新预览", 409)
        result = expense.apply(
            db,
            plan,
            operated_by=actor.name,
            import_batch_id=p.id,
            commit=False,
            workbook_state=state,
        )
    else:
        raise McpError("unsupported_family", "该单据类型尚未开放正式提交")
    receipt = record(
        db,
        actor,
        "receipt",
        {"status": "succeeded", "preview_id": p.id, "result": result},
        key=p.id,
        request=body,
        hours=24 * 365,
    )
    receipt.payload = {**receipt.payload, "receipt_id": receipt.id}
    p.state = "applied"
    event(
        db,
        actor,
        "confirm_import",
        "approved",
        {"preview_id": p.id, "preview_hash": expected_hash},
        request_id,
    )
    event(
        db,
        actor,
        "confirm_import",
        "completed",
        {
            "preview_id": p.id,
            "receipt_id": receipt.id,
            "project_id": body["project_id"],
        },
        request_id,
    )
    return receipt.payload


def export_bytes(db, actor, args):
    require(actor, "mcp.export")
    kind = args["export_kind"]
    project_id = (args.get("project_ids") or [None])[0]
    if (
        len(args.get("project_ids", [])) > 1
        or args.get("fields")
        or args.get("date_from")
        or args.get("date_to")
    ):
        raise McpError("invalid_export", "当前完整模板不支持多项目、字段裁剪或日期过滤")
    project_access(db, actor, project_id, expense_data=kind == "expense_collection")
    if kind == "acceptance_checklist":
        current = checklist.project_checklist(db, project_id)["current"]
        if current:
            from openpyxl import Workbook

            wb = Workbook()
            ws = wb.active
            ws.title = "验收清单"
            ws.append(["验收需求", "是否完成"])
            for item in current["items"]:
                ws.append([item["requirement"], "是" if item["done"] else "否"])
                ws.cell(ws.max_row, 1).data_type = "s"
            out = io.BytesIO()
            wb.save(out)
            wb.close()
            data = out.getvalue()
        else:
            data = checklist.build_template()
        name = "acceptance-checklist.xlsx"
    elif kind == "expense_collection":
        if not get_settings().maintenance_boss_dashboard_enabled:
            raise McpError("feature_disabled", "维保功能未启用", 404)
        data = expense.build_workbook(db, project_id=project_id)
        name = "expense-collection.xlsx"
    else:
        raise McpError("unsupported_family", "该导出类型尚未开放")
    if data is None:
        raise McpError("not_found", "项目不存在", 404)
    if len(data) > get_settings().mcp_max_file_bytes:
        raise McpError("file_too_large", "导出文件超出限制", 413)
    r = record(
        db,
        actor,
        "artifact",
        {
            "name": name,
            "sha256": hashlib.sha256(data).hexdigest(),
            "size_bytes": len(data),
            "project_id": project_id,
            "export_kind": kind,
            "auth_version": actor.version,
            "credential_id": actor.credential_id,
        },
        hours=24 * 7,
    )
    save_bytes(r.id, data)
    return {
        "artifact_id": r.id,
        "filename": name,
        "size_bytes": len(data),
        "sha256": r.payload["sha256"],
        "roundtrip": True,
    }


def check_artifact(db, actor, id):
    require(actor, "mcp.files.download")
    r = owned(db, actor, id, "artifact")
    project_access(
        db,
        actor,
        r.payload["project_id"],
        expense_data=r.payload["export_kind"] == "expense_collection",
    )
    if (
        r.payload["export_kind"] == "expense_collection"
        and not get_settings().maintenance_boss_dashboard_enabled
    ):
        raise McpError("feature_disabled", "维保功能未启用", 404)
    if actor.version != r.payload["auth_version"]:
        raise McpError("permission_denied", "账号授权已变化，请重新导出", 403)
    if r.payload.get("credential_id"):
        from app.models.mcp import McpCredential

        cred = db.get(McpCredential, r.payload["credential_id"])
        if not cred or cred.revoked_at or cred.expires_at <= now():
            raise McpError("permission_denied", "原导出连接授权已失效", 403)
    return r
