"""Annotate incomplete site dates without changing any workbook values or identities.

Usage: python scripts/annotate_site_dates.py source.xlsx review.xlsx
The output is a review copy, not a successfully validated import.
"""
from __future__ import annotations

import argparse
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

from app.services.maintenance_project_master_workbook import _v2_date, _is_example_row, WorkbookError


def annotate(source: Path, destination: Path) -> list[int]:
    if source.resolve() == destination.resolve():
        raise ValueError("输出文件必须与原文件分开")
    wb = load_workbook(source)
    ws = wb["06_领用返还"]
    index = {c.value: c.column for c in ws[1]}
    marked = []
    for cells in ws.iter_rows(min_row=2):
        values = tuple(c.value for c in cells)
        if not any(v not in (None, "") for v in values) or _is_example_row(values):
            continue
        cell = cells[index["领用日期"] - 1]
        try:
            parsed = _v2_date(cell.value, row_no=cell.row, label="领用", epoch=wb.epoch)
            incomplete = parsed is None or bool(re.fullmatch(r"\d{4}-\d{2}", str(cell.value).strip()))
        except WorkbookError:
            incomplete = True
        if not incomplete:
            continue
        note = f"待业务确认：原值 {cell.value!r}。请补全真实领用日期 YYYY-MM-DD；不能用 WBDD 制单日期代替。"
        if isinstance(cell.value, str) and re.fullmatch(r"\d{1,2}月\d{1,2}日?", cell.value):
            note += " 此处缺少年份，请查原始领用记录确认。"
        cell.comment = Comment(note, "导入核对")
        cell.fill = PatternFill("solid", fgColor="FFC7CE")
        marked.append(cell.row)
    # Exclusive creation prevents overwriting another review or the original.
    with destination.open("xb") as output:
        wb.save(output)
    wb.close()
    return marked


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    args = parser.parse_args()
    print("标记待补日期行：", annotate(args.source, args.destination))
